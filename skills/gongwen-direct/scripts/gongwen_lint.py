# -*- coding: utf-8 -*-
"""公文机检引擎（gongwen_lint.py）。

源文件位于仓库 shared/scripts/，各 skill 中的副本由 tools/sync_shared.py 生成，请在源文件处修改。
接口见 docs/开发规范.md 第 8 节，其他 skill 只依赖命令行接口和 JSON 输出格式：

    python scripts/gongwen_lint.py <文件> [--doc-type 文书类型] [--checks 类别,类别] [--json] [--strict]

- 只依赖标准库；读取 .docx 时需要 python-docx。
- 规则数据：同目录下的 lint_rules.json（正则规则 + 函数规则登记）和 style_patterns.json（AI 味、防御性词表）。
- 遇到单独成行的【起草说明】即停止读取（兼容“## 【起草说明】”“**【起草说明】**”，与 fact_diff.py、gongwen-format 一致）。
- 退出码：0 正常；--strict 且存在 P0 时为 1；输入错误为 2。
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Any, Callable, Dict, List, Optional, Pattern, Sequence, Set, Tuple

VERSION = "0.1.0"
SCRIPT_DIR = Path(__file__).resolve().parent
DEFAULT_RULES = SCRIPT_DIR / "lint_rules.json"
DEFAULT_PATTERNS = SCRIPT_DIR / "style_patterns.json"

# 检查类别（顺序即输出顺序），与开发规范第 8 节一致
CHECK_ORDER = (
    "residue", "punct", "number", "format", "structure",
    "style", "defensive", "privacy", "legal", "jianfu", "marker", "length",
)
CHECK_NAMES = {
    "residue": "残留",
    "punct": "标点",
    "number": "数字",
    "format": "格式要素",
    "structure": "结构",
    "style": "AI味",
    "defensive": "防御性",
    "privacy": "隐私与保密",
    "legal": "法规",
    "jianfu": "减负",
    "marker": "待补标记",
    "length": "篇幅",
}
LEVELS = ("P0", "P1", "P2")
LEVEL_RANK = {"P0": 0, "P1": 1, "P2": 2}

# 统一词表中的文书类型（开发规范第 6 节）
DOC_TYPES = (
    "决议", "决定", "命令", "公报", "公告", "通告", "意见", "通知", "通报", "报告", "请示",
    "批复", "议案", "函", "纪要", "签报",
    "总结", "计划", "工作要点", "汇报材料", "调研报告", "经验材料", "述职报告", "实施方案",
    "讲话稿", "信息", "其他",
)
# 法定文种（条例第八条）：标题文种与 --doc-type 不一致的提示只在两者都是法定文种时给出
STATUTORY_TYPES = {
    "决议", "决定", "命令", "公报", "公告", "通告", "意见", "通知", "通报", "报告", "请示", "批复", "议案", "函", "纪要",
}
# 有主送机关的常见文种：标题后一行缺冒号时才判为主送机关缺冒号
ZHUSONG_TYPES = {"通知", "通报", "报告", "请示", "批复", "函", "意见", "决定", "议案", "签报"}
# 篇幅上限适用的下行文件（《整治形式主义为基层减负若干规定》第2条）
LENGTH_LIMIT_TYPES = {
    "决议", "决定", "命令", "公报", "公告", "通告", "意见", "通知", "通报", "批复",
    "实施方案", "工作要点", "计划",
}
# 标题结尾 → 文书类型（按顺序匹配，长的在前）
TITLE_DOC_SUFFIX = (
    ("请示报告", "请示"), ("调研报告", "调研报告"), ("述责述廉报告", "述职报告"),
    ("述职述廉报告", "述职报告"), ("述职报告", "述职报告"), ("实施方案", "实施方案"),
    ("工作方案", "实施方案"), ("方案", "实施方案"), ("工作要点", "工作要点"), ("要点", "工作要点"),
    ("汇报材料", "汇报材料"), ("经验材料", "经验材料"), ("工作总结", "总结"), ("总结", "总结"),
    ("调查报告", "调研报告"), ("工作计划", "计划"), ("计划", "计划"), ("讲话稿", "讲话稿"), ("讲话", "讲话稿"),
    ("发言", "讲话稿"), ("致辞", "讲话稿"), ("请示", "请示"), ("报告", "报告"), ("通知", "通知"),
    ("通报", "通报"), ("通告", "通告"), ("公告", "公告"), ("公报", "公报"), ("批复", "批复"),
    ("意见", "意见"), ("决定", "决定"), ("决议", "决议"), ("命令", "命令"), ("议案", "议案"),
    ("纪要", "纪要"), ("签报", "签报"), ("复函", "函"), ("函", "函"),
)

# 【起草说明】分界行：去掉行首空白后以【起草说明】开头；兼容 Markdown 标题、加粗包裹的写法。
# 与 fact_diff.py 的 DRAFT_NOTE_RE、gongwen-format 的 gw_layout.is_draft_note_line 保持一致。
DRAFT_NOTE_RE = re.compile(r"^\s*(?:#{1,6}\s*)?(?:(?:\*\*|__)\s*)?【起草说明】")
PUNCT_TAIL = "。，；：、！？.,;:!?"

# ---------------------------------------------------------------------------
# 版头、标题、落款、版记的识别
# ---------------------------------------------------------------------------
FENHAO_RE = re.compile(r"^\s*\d{6}\s*$")
MIJI_LINE_RE = re.compile(r"^\s*(?:绝密|机密|秘密)(?:\s*★.*)?\s*$")
INNER_MARK_RE = re.compile(
    r"^\s*[（(【\[]?\s*(?:内部资料|内部文件|内部材料|内部参阅|工作秘密|内部使用|不得外传|注意保密|内部传阅)"
    r"[\s　，,、]*(?:注意保存|妥善保管|不得外传|请勿外传|注意保密)?\s*[）)】\]]?\s*$"
)
# 【待补：】【待核实：】标记：正文唯一允许的提示内容，不参与冒号、文风类检查
MARKER_RE = re.compile(r"【(?:待补|待核实)[：:][^】\n]*】")
# summary.markers 计数用（与 MRK-001、MRK-002 的写法一致：全角冒号）
TODO_MARK_RE = re.compile(r"【待补：[^】\n]*】")
VERIFY_MARK_RE = re.compile(r"【待核实：[^】\n]*】")
JINJI_RE = re.compile(r"^\s*(?:特急|加急|特提|平急)\s*$")
BIAOZHI_RE = re.compile(r"^\s*\S{2,40}(?:文件|命令|令)\s*$")
WENHAO_RE = re.compile(
    r"[一-鿿A-Za-z]{1,15}\s*[〔\[［(（【﹝{｛〖]\s*(?:19|20)?\d{2}\s*[〕\]］)）】﹞}｝〗]\s*第?\s*\d+\s*号"
)
QIANFA_RE = re.compile(r"签发人\s*[：:]")
HEADING_START_RE = re.compile(r"^\s*(?:[一二三四五六七八九十]{1,3}\s*[、，,.．]|[（(]\s*[一二三四五六七八九十\d]{1,3}\s*[）)]|\d{1,2}\s*[.．、])")
TITLE_END_RE = re.compile(
    r"(?:决议|决定|命令|公报|公告|通告|意见|通知|通报|报告|请示|批复|议案|函|纪要|签报|总结|计划|要点|材料|方案"
    r"|讲话|发言|致辞|信息|简报|办法|规定|细则|规则|制度|章程|预案|规划|说明|汇报|提纲)$"
)
DATE_ONLY_RE = re.compile(
    r"^\s*[\d〇○零一二三四五六七八九OＯ]{4}\s*年\s*[\d一二三四五六七八九十]{1,3}\s*月\s*[\d一二三四五六七八九十]{1,3}\s*日\s*$"
)
CN_DATE_ONLY_RE = re.compile(
    r"^\s*[〇○零一二三四五六七八九OＯ]{4}\s*年\s*[一二三四五六七八九十]{1,3}\s*月\s*[一二三四五六七八九十]{1,3}\s*日\s*$"
)
BANJI_RE = re.compile(r"^\s*(?:抄送|抄报|主送)\s*[：:]|印发\s*$")
ATTACH_DESC_RE = re.compile(r"^\s*附件\s*[：:]")
ATTACH_PAGE_RE = re.compile(r"^\s*附件\s*\d{0,2}\s*$")
ORG_SUFFIX_RE = re.compile(
    r"(?:单位|部门|政府|委员会|党委|党组|机关|学校|医院|公司|集团|办公室|办公厅|办事处|局|厅|委|办|处|科|室|院|校"
    r"|中心|街道|乡镇|社区|村|镇|乡|区|县|市|省|部|会|组|所|站|队)$"
)

# 面向群众的文书：主送或称呼行写给居民、家长等，或标题为公示、告知书等
PUBLIC_AUDIENCE_RE = re.compile(r"居民|村民|群众|市民|家长|业主|住户|商户|患者|游客|乘客|消费者|朋友们|村（居）民")
PUBLIC_TITLE_RE = re.compile(r"(?:公示|告知书|一封信|倡议书|温馨提示|告居民书|公开信|须知|告示)$")
# 法规类规则的“已写明依据”判断：同一句中命中位置之前出现这些字样时不报
LEGAL_BASIS_RE = re.compile(
    r"依法|依照|《[^》]{2,40}(?:法|条例|规定|办法)》|(?:根据|按照)[^。；！？]{0,30}(?:法|条例)"
)

# 层次序数：level, 正则（第 1 组为序号）
HEADING_PATTERNS = (
    # “、”后可以直接是数字（如“三、2027年工作打算”）；其他误用的点号后紧跟数字的不算标题
    (1, re.compile(r"^\s*([一二三四五六七八九十]{1,3})\s*(?:、|[，,.．:：](?!\d))"
                   r"(?![一二三四五六七八九十]+(?:年级|产业|季度|等奖|类|级))")),
    (2, re.compile(r"^\s*[（(]\s*([一二三四五六七八九十]{1,3})\s*[）)]")),
    (3, re.compile(r"^\s*(\d{1,2})\s*(?:[.．](?!\d)|、(?!\d))")),
    (4, re.compile(r"^\s*[（(]\s*(\d{1,2})\s*[）)]")),
)
CN_DIGIT = {"零": 0, "〇": 0, "一": 1, "二": 2, "两": 2, "三": 3, "四": 4, "五": 5, "六": 6, "七": 7, "八": 8, "九": 9}

# 常见姓氏（用于隐私检查中判断“像姓名”，减少误报）
SURNAMES = set(
    "王李张刘陈杨黄赵吴周徐孙马朱胡郭何高林罗郑梁谢宋唐许韩冯邓曹彭曾肖田董袁潘于蒋蔡余杜叶程苏魏吕丁任沈"
    "姚卢姜崔钟谭陆汪范金石廖贾夏韦付方白邹孟熊秦邱江尹薛闫段雷侯龙史陶黎贺顾毛郝龚邵万钱严覃武戴莫孔向汤常"
    "温康施文牛樊葛邢安齐易乔伍庞颜倪庄聂章鲁岳翟殷詹申欧耿关兰焦俞左柳甘祝包宁尚符舒阮柯纪梅童凌毕单季裴霍"
    "涂成苗谷盛曲翁冉骆蓝路游辛靳管柴蒙鲍华喻祁蒲房滕屈饶解牟艾尤阳时穆农司卓古吉缪简车项连芦麦褚娄窦戚岑景党"
    "宫费卜冷晏席卫米柏宗瞿桂全佟应臧闵苟邬边卞姬师和仇栾隋商刁沙荣巫寇桑郎甄丛仲虞敖巩明佘池查麻苑迟邝"
)
ROLE_WORDS = ("患者", "病人", "患儿", "学生", "幼儿", "考生", "当事人", "居民", "村民", "职工", "老人", "家属", "儿童",
              "服务对象", "群众", "同志", "护士", "医生", "教师", "老师")
MASK_CHARS = set("某×X*＊")


class InputError(Exception):
    """输入错误：文件不存在、格式不支持、依赖缺失等，退出码 2。"""


# ---------------------------------------------------------------------------
# 数据结构
# ---------------------------------------------------------------------------

@dataclass
class Para:
    index: int          # 在 paras 中的序号（从 0 开始）
    line: int           # 原文行号；docx 为段落序号（从 1 开始）
    text: str
    in_table: bool = False


@dataclass
class Issue:
    id: str
    level: str
    check: str
    category: str
    gate: str
    line: int
    excerpt: str
    message: str
    suggestion: str
    basis: str
    context: str = ""
    start: int = -1
    end: int = -1
    count: Optional[int] = None
    para_index: int = -1

    def to_dict(self) -> Dict[str, Any]:
        data: Dict[str, Any] = {
            "id": self.id,
            "level": self.level,
            "category": self.category,
            "line": self.line,
            "excerpt": self.excerpt,
            "message": self.message,
            "suggestion": self.suggestion,
            "basis": self.basis,
            "check": self.check,
            "gate": self.gate,
        }
        if self.context:
            data["context"] = self.context
        if self.count is not None:
            data["count"] = self.count
        return data


@dataclass
class RegexRule:
    id: str
    check: str
    category: str
    gate: str
    level: str
    regex: Pattern[str]
    message: str
    suggestion: str
    basis: str
    doc_types: Tuple[str, ...]
    exempt_doc_types: Tuple[str, ...]
    yield_to: Tuple[str, ...]
    once: bool
    skip_line: Optional[Pattern[str]]
    fixes: Tuple[Tuple[Pattern[str], str], ...]
    skip_table: bool = False
    zone: str = ""
    also_public: bool = False
    public_only: bool = False
    skip_if_basis: bool = False
    skip_in_quotes: bool = False


@dataclass
class StylePattern:
    id: str
    group: str
    category: str
    regex: Pattern[str]
    level: str
    message: str
    suggestion: str
    density: Optional[float]
    doc_types: Tuple[str, ...]
    exempt_doc_types: Tuple[str, ...]


def applies_to(doc_type: str, doc_types: Sequence[str], exempt: Sequence[str]) -> bool:
    """doc_types 非空时为正向限定；exempt_doc_types 为豁免。"""
    if doc_types and doc_type not in doc_types:
        return False
    return doc_type not in exempt


def nonspace_len(text: str) -> int:
    return len(re.sub(r"\s", "", text))


def cn_to_int(text: str) -> int:
    """把“一”“十二”“二十”等汉字序数或阿拉伯数字转为整数。"""
    text = text.strip()
    if text.isdigit():
        return int(text)
    if "十" in text:
        head, _, tail = text.partition("十")
        tens = CN_DIGIT.get(head, 1) if head else 1
        ones = CN_DIGIT.get(tail, 0) if tail else 0
        return tens * 10 + ones
    value = 0
    for ch in text:
        value = value * 10 + CN_DIGIT.get(ch, 0)
    return value


# 行首的 Markdown 标题、加粗符号：结构识别时先去掉（残留类规则仍按原文报告）
MD_HEADING_PREFIX_RE = re.compile(r"^\s*(?:#{1,6}\s*)?(?:\*\*|__)?")


def heading_of(text: str) -> Optional[Tuple[int, int]]:
    """识别行首层次序数，返回 (层级, 序号)。“**三、时间安排**”“## 一、总体要求”按一级标题识别。"""
    text = MD_HEADING_PREFIX_RE.sub("", text, count=1)
    for level, regex in HEADING_PATTERNS:
        match = regex.match(text)
        if match:
            return level, cn_to_int(match.group(1))
    return None


def clip(text: str, limit: int = 60) -> str:
    text = text.strip()
    return text if len(text) <= limit else text[:limit]


def context_of(text: str, start: int, end: int, width: int = 12) -> str:
    return text[max(0, start - width):min(len(text), end + width)].strip()


# ---------------------------------------------------------------------------
# 输入读取
# ---------------------------------------------------------------------------

def read_lines(path: Path) -> Tuple[List[Tuple[int, str, bool]], List[str]]:
    """读取文件，返回 [(行号, 文本, 是否表格内)] 和提示信息。"""
    notes: List[str] = []
    suffix = path.suffix.lower()
    if suffix in (".txt", ".md", ".markdown", ".text"):
        try:
            raw = path.read_bytes()
        except OSError as exc:
            raise InputError(f"无法读取文件：{exc}") from exc
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                text = raw.decode("gb18030")
                notes.append("文件不是 UTF-8 编码，已按 GB18030 读取。")
            except UnicodeDecodeError as exc:
                raise InputError("无法识别文件编码，请另存为 UTF-8 后重试。") from exc
        text = text.replace("\r\n", "\n").replace("\r", "\n")
        return [(i, line, False) for i, line in enumerate(text.split("\n"), start=1)], notes
    if suffix == ".docx":
        return read_docx(path), notes
    raise InputError(f"不支持的文件类型：{path.suffix or '无扩展名'}。只支持 .txt、.md、.docx。")


def read_docx(path: Path) -> List[Tuple[int, str, bool]]:
    try:
        import docx  # type: ignore
        from docx.oxml.ns import qn  # type: ignore
        from docx.text.paragraph import Paragraph  # type: ignore
    except ImportError as exc:
        raise InputError("读取 .docx 需要 python-docx，请先运行：pip install python-docx") from exc
    try:
        document = docx.Document(str(path))
    except Exception as exc:  # 损坏的 docx、加密文件等
        raise InputError(f"无法打开 docx 文件：{exc}") from exc
    lines: List[Tuple[int, str, bool]] = []
    number = 0
    p_tag, tbl_tag = qn("w:p"), qn("w:tbl")
    for child in document.element.body.iterchildren():
        if child.tag == p_tag:
            number += 1
            lines.append((number, Paragraph(child, document).text, False))
        elif child.tag == tbl_tag:
            for p_el in child.iter(p_tag):
                number += 1
                lines.append((number, Paragraph(p_el, document).text, True))
    return lines


# ---------------------------------------------------------------------------
# 文档结构
# ---------------------------------------------------------------------------

class Document:
    """正文（【起草说明】之前）的段落及结构识别结果。"""

    def __init__(self, file_name: str, lines: Sequence[Tuple[int, str, bool]], unit: str = "行") -> None:
        self.file = file_name
        self.unit = unit
        self.truncated_line: Optional[int] = None
        paras: List[Para] = []
        for number, text, in_table in lines:
            if DRAFT_NOTE_RE.match(text):
                self.truncated_line = number
                break
            if text.strip():
                paras.append(Para(len(paras), number, text.rstrip(), in_table))
        self.paras = paras
        self.char_count = sum(nonspace_len(p.text) for p in paras)
        self.head: Set[int] = set()
        self.title: List[int] = []
        self.subtitle: Optional[int] = None
        self.zhusong: Optional[int] = None
        self.banji_start: Optional[int] = None
        self.signoff_date: Optional[int] = None
        self.attach_desc: Optional[int] = None
        self.wenhao_present = False
        self.qianfa_present = False
        self.public_facing = False
        self.first_body: Optional[int] = None
        self._analyze()

    # -- 结构识别 -----------------------------------------------------------
    def _analyze(self) -> None:
        paras = self.paras
        if not paras:
            return
        i = 0
        while i < len(paras) and i < 8:
            text = paras[i].text.strip()
            nxt = paras[i + 1].text.strip() if i + 1 < len(paras) else ""
            if self._is_wenhao_line(text):
                self.wenhao_present = True
                if QIANFA_RE.search(text):
                    self.qianfa_present = True
                self.head.add(i)
            elif QIANFA_RE.search(text) and len(text) <= 30:
                self.qianfa_present = True
                self.head.add(i)
            elif (FENHAO_RE.match(text) or MIJI_LINE_RE.match(text) or JINJI_RE.match(text)
                  or BIAOZHI_RE.match(text) or INNER_MARK_RE.match(text)):
                self.head.add(i)
            elif (len(text) <= 40 and nxt and (self._is_wenhao_line(nxt) or QIANFA_RE.search(nxt))
                  and not self._title_shaped(text)):
                self.head.add(i)  # 发文机关标志（不带“文件”二字）；标题后紧跟发文字号的不算
            else:
                break
            i += 1
        if i < len(paras) and self._title_like(paras[i].text):
            self.title.append(i)
            j = i + 1
            while (j < len(paras) and len(self.title) < 3
                   and not TITLE_END_RE.search(self._title_text().rstrip(PUNCT_TAIL))
                   and self._title_continuation(paras[j].text)):
                self.title.append(j)
                j += 1
            if j < len(paras) and paras[j].text.strip().startswith("——"):
                self.subtitle = j
                j += 1
            if j < len(paras) and self._is_wenhao_line(paras[j].text.strip()):
                # 发文字号写在标题之下（AI 稿常见）：按版头要素处理，不当作主送或正文
                self.wenhao_present = True
                if QIANFA_RE.search(paras[j].text):
                    self.qianfa_present = True
                self.head.add(j)
                j += 1
            if j < len(paras):
                cand = paras[j].text.strip()
                if len(cand) <= 100 and "。" not in cand and (cand.endswith("：") or cand.endswith(":")):
                    self.zhusong = j
        # 版记：从末尾向前找抄送、印发行
        for k in range(len(paras) - 1, max(-1, len(paras) - 7), -1):
            if k in self.head or k in self.title:
                break
            if BANJI_RE.search(paras[k].text):
                self.banji_start = k
        end = self.banji_start if self.banji_start is not None else len(paras)
        for k in range(end - 1, max(-1, end - 10), -1):
            if DATE_ONLY_RE.match(paras[k].text):
                self.signoff_date = k
                break
        for para in paras:
            if ATTACH_DESC_RE.match(para.text) and para.index not in self.head:
                self.attach_desc = para.index
                break
        title = self.title_text.rstrip(PUNCT_TAIL)
        audience = paras[self.zhusong].text if self.zhusong is not None else ""
        if not audience:
            # 没有标题时，开头几段中以冒号结尾的短行视为称呼（如“各位居民：”）
            for para in paras[:3]:
                text = para.text.strip()
                if len(text) <= 40 and text.endswith(("：", ":")):
                    audience = text
                    break
        self.public_facing = bool(PUBLIC_TITLE_RE.search(title) or PUBLIC_AUDIENCE_RE.search(audience)
                                  or re.search(r"致[^，。]{0,10}(?:居民|村民|家长|业主|市民)", title))
        for k in self.body_indices():
            text = paras[k].text
            # 正文第一段：不是层次标题、附件说明，且是成句的段落（缺冒号的主送行等短行跳过）
            if (not heading_of(text) and not ATTACH_DESC_RE.match(text) and not paras[k].in_table
                    and re.search(r"[。；，！？：]", text)):
                self.first_body = k
                break

    @staticmethod
    def _is_wenhao_line(text: str) -> bool:
        """整行基本就是发文字号（可带签发人），而不是正文中引用文号的句子。"""
        m = WENHAO_RE.search(text)
        if not m:
            return False
        rest = (text[:m.start()] + text[m.end():]).strip()
        return len(rest) <= 6 or rest.startswith("签发人") or bool(QIANFA_RE.search(rest))

    def _title_text(self) -> str:
        return "".join(self.paras[k].text.strip() for k in self.title)

    @staticmethod
    def _title_shaped(text: str) -> bool:
        """以文种结尾或以“关于”开头：是标题，不是发文机关标志。"""
        text = text.strip()
        return text.startswith("关于") or bool(TITLE_END_RE.search(text.rstrip(PUNCT_TAIL)))

    @staticmethod
    def _title_like(text: str) -> bool:
        text = text.strip()
        if not text or len(text) > 80 or HEADING_START_RE.match(text):
            return False
        if text.endswith("：") or text.endswith(":") or text.startswith("——"):
            return False
        if "，" in text or "；" in text or "：" in text:
            return False
        if "。" not in text:
            return True
        # 以句号结尾的，只有像标题（以文种结尾或以“关于”开头）才算标题，避免把单句正文当标题
        body = text[:-1]
        return text.count("。") == 1 and text.endswith("。") and (
            bool(TITLE_END_RE.search(body)) or body.startswith("关于"))

    @staticmethod
    def _title_continuation(text: str) -> bool:
        text = text.strip()
        if not text or len(text) > 40 or HEADING_START_RE.match(text) or text.startswith("——"):
            return False
        if DATE_ONLY_RE.match(text) or re.search(r"[。；：:，,（(]", text):
            return False
        return True

    # -- 查询 -----------------------------------------------------------------
    @property
    def title_text(self) -> str:
        return self._title_text()

    def body_indices(self) -> List[int]:
        """正文段落：去掉版头、标题、副标题、主送和版记。"""
        skip = set(self.head) | set(self.title)
        if self.subtitle is not None:
            skip.add(self.subtitle)
        if self.zhusong is not None:
            skip.add(self.zhusong)
        end = self.banji_start if self.banji_start is not None else len(self.paras)
        return [p.index for p in self.paras[:end] if p.index not in skip]

    def body_char_count(self) -> int:
        """篇幅统计：标题之后、【起草说明】之前，不含空白和版记。"""
        skip = set(self.head) | set(self.title)
        if self.subtitle is not None:
            skip.add(self.subtitle)
        end = self.banji_start if self.banji_start is not None else len(self.paras)
        return sum(nonspace_len(p.text) for p in self.paras[:end] if p.index not in skip)

    def infer_doc_type(self) -> Optional[str]:
        title = self.title_text.strip().rstrip(PUNCT_TAIL)
        if title:
            for suffix, doc_type in TITLE_DOC_SUFFIX:
                if title.endswith(suffix):
                    return doc_type
        if self.subtitle is not None and re.search(r"(讲话|发言|致辞)\s*$", self.paras[self.subtitle].text):
            return "讲话稿"
        return None


# ---------------------------------------------------------------------------
# 引擎
# ---------------------------------------------------------------------------

def _compile(pattern: str, flags: int = 0) -> Pattern[str]:
    return re.compile(pattern, flags)


class Linter:
    """加载规则数据，对 Document 执行检查。"""

    def __init__(self, rules_path: Optional[Path] = None, patterns_path: Optional[Path] = None) -> None:
        self.warnings: List[str] = []
        self.rules: List[RegexRule] = []
        self.proc_meta: Dict[str, Dict[str, Any]] = {}
        self.patterns: List[StylePattern] = []
        self.whitelist: List[Tuple[str, Pattern[str]]] = []
        self._load_rules(rules_path or DEFAULT_RULES)
        self._load_patterns(patterns_path or DEFAULT_PATTERNS)

    # -- 加载 -----------------------------------------------------------------
    def _load_json(self, path: Path, label: str) -> Optional[Dict[str, Any]]:
        if not path.is_file():
            self.warnings.append(f"未找到 {label}（{path.name}），相关检查已跳过。")
            return None
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, ValueError) as exc:
            self.warnings.append(f"{label}（{path.name}）无法解析：{exc}，相关检查已跳过。")
            return None
        if not isinstance(data, dict):
            self.warnings.append(f"{label}（{path.name}）顶层不是对象，相关检查已跳过。")
            return None
        return data

    def _load_rules(self, path: Path) -> None:
        data = self._load_json(path, "规则文件")
        if data is None:
            return
        for item in data.get("rules", []):
            try:
                flags = re.IGNORECASE if item.get("flags") == "IGNORECASE" else 0
                skip = item.get("skip_line")
                fixes = tuple((_compile(p), r) for p, r in item.get("fix", []))
                self.rules.append(RegexRule(
                    id=item["id"], check=item["check"],
                    category=item.get("category") or CHECK_NAMES.get(item["check"], item["check"]),
                    gate=item.get("gate", ""), level=item.get("level", "P2"),
                    regex=_compile(item["pattern"], flags),
                    message=item.get("message", ""), suggestion=item.get("suggestion", ""),
                    basis=item.get("basis", ""),
                    doc_types=tuple(item.get("doc_types") or ()),
                    exempt_doc_types=tuple(item.get("exempt_doc_types") or ()),
                    yield_to=tuple(item.get("yield_to") or ()),
                    once=bool(item.get("once")),
                    skip_line=_compile(skip) if skip else None,
                    fixes=fixes,
                    skip_table=bool(item.get("skip_table")),
                    zone=item.get("zone", ""),
                    also_public=bool(item.get("also_public")),
                    public_only=bool(item.get("public_only")),
                    skip_if_basis=bool(item.get("skip_if_basis")),
                    skip_in_quotes=bool(item.get("skip_in_quotes")),
                ))
            except (KeyError, re.error, TypeError, ValueError) as exc:
                self.warnings.append(f"规则 {item.get('id', '?')} 无效，已跳过：{exc}")
        for item in data.get("procedural", []):
            if isinstance(item, dict) and item.get("id"):
                self.proc_meta[item["id"]] = item

    def _load_patterns(self, path: Path) -> None:
        data = self._load_json(path, "文风词表")
        if data is None:
            return
        for item in data.get("patterns", []):
            try:
                density = item.get("density_per_1000")
                self.patterns.append(StylePattern(
                    id=item["id"], group=item.get("group", "style"),
                    category=item.get("category", ""), regex=_compile(item["pattern"]),
                    level=item.get("level", "P2"), message=item.get("message", ""),
                    suggestion=item.get("suggestion", ""),
                    density=float(density) if density is not None else None,
                    doc_types=tuple(item.get("doc_types") or ()),
                    exempt_doc_types=tuple(item.get("exempt_doc_types") or ()),
                ))
            except (KeyError, re.error, TypeError, ValueError) as exc:
                self.warnings.append(f"词表规则 {item.get('id', '?')} 无效，已跳过：{exc}")
        for item in data.get("whitelist", []):
            try:
                self.whitelist.append((item["id"], _compile(item["pattern"])))
            except (KeyError, re.error, TypeError) as exc:
                self.warnings.append(f"白名单 {item.get('id', '?')} 无效，已跳过：{exc}")

    # -- 入口 -----------------------------------------------------------------
    def lint(self, doc: Document, doc_type: str, checks: Set[str]) -> List[Issue]:
        issues: List[Issue] = []
        issues.extend(self._run_regex_rules(doc, doc_type, checks))
        for check in CHECK_ORDER:
            if check not in checks:
                continue
            for func in self._procedures().get(check, ()):
                issues.extend(func(doc, doc_type))
        issues.extend(self._run_style(doc, doc_type, checks))
        issues = self._resolve_yields(issues)
        issues = self._apply_once(issues)
        issues.sort(key=lambda x: (LEVEL_RANK.get(x.level, 9), x.line, x.start, x.id))
        return issues

    def _procedures(self) -> Dict[str, Sequence[Callable[[Document, str], List[Issue]]]]:
        return {
            "punct": (self._check_quotes, self._check_colon_nesting),
            "number": (self._check_date_styles,),
            "format": (self._check_date_padding, self._check_signoff_date, self._check_zhusong),
            "structure": (self._check_title, self._check_heading_sequence, self._check_count_summary,
                          self._check_qingshi, self._check_baogao, self._check_han, self._check_shangxing),
            "privacy": (self._check_privacy,),
            "length": (self._check_length,),
        }

    # -- 构造问题 -----------------------------------------------------------
    def _proc_issue(self, rid: str, para: Optional[Para], start: int, end: int, *,
                    message: Optional[str] = None, level: Optional[str] = None,
                    suggestion: Optional[str] = None, excerpt: Optional[str] = None,
                    context: Optional[str] = None, count: Optional[int] = None) -> Optional[Issue]:
        meta = self.proc_meta.get(rid)
        if meta is None:
            return None  # 未在 lint_rules.json 的 procedural 中登记的检查不输出
        text = para.text if para else ""
        if excerpt is None:
            excerpt = clip(text[start:end], 80) if para and start >= 0 else ""
        if context is None:
            context = context_of(text, start, end) if para and start >= 0 else ""
        return Issue(
            id=rid, level=level or meta.get("level", "P2"), check=meta.get("check", ""),
            category=meta.get("category", CHECK_NAMES.get(meta.get("check", ""), "")),
            gate=meta.get("gate", ""), line=para.line if para else 0, excerpt=excerpt,
            message=message or meta.get("message", ""),
            suggestion=suggestion if suggestion is not None else meta.get("suggestion", ""),
            basis=meta.get("basis", ""), context=context if context != excerpt else "",
            start=start, end=end, count=count, para_index=para.index if para else -1,
        )

    def _proc_applies(self, rid: str, doc_type: str) -> bool:
        meta = self.proc_meta.get(rid)
        if meta is None:
            return False
        return applies_to(doc_type, meta.get("doc_types") or (), meta.get("exempt_doc_types") or ())

    # -- 正则规则 -----------------------------------------------------------
    @staticmethod
    def _rule_applies(rule: RegexRule, doc: Document, doc_type: str) -> bool:
        """文种限定：public_only 只查面向群众的文书；also_public 在限定文种之外也查面向群众的文书。"""
        if doc_type in rule.exempt_doc_types:
            return False
        if rule.public_only:
            return doc.public_facing
        if rule.doc_types and doc_type not in rule.doc_types:
            return rule.also_public and doc.public_facing
        return True

    def _run_regex_rules(self, doc: Document, doc_type: str, checks: Set[str]) -> List[Issue]:
        issues: List[Issue] = []
        for rule in self.rules:
            if rule.check not in checks or not self._rule_applies(rule, doc, doc_type):
                continue
            if rule.zone == "first_body":
                paras = [doc.paras[doc.first_body]] if doc.first_body is not None else []
            else:
                paras = doc.paras
            for para in paras:
                if rule.skip_table and para.in_table:
                    continue
                if rule.skip_line and rule.skip_line.search(para.text):
                    continue
                for match in rule.regex.finditer(para.text):
                    start, end = match.span()
                    if start == end:
                        continue
                    if rule.skip_in_quotes and self._inside_quote(para.text, start):
                        continue  # 引号内是举例或引述，不是稿件本身的标注
                    if rule.skip_if_basis:
                        sent_start = max(para.text.rfind(ch, 0, start) for ch in "。！？；") + 1
                        if LEGAL_BASIS_RE.search(para.text[sent_start:start]):
                            continue
                    excerpt = clip(match.group(0), 80)
                    suggestion = rule.suggestion
                    if rule.fixes:
                        fixed = match.group(0).strip()
                        for fix_re, repl in rule.fixes:
                            fixed = fix_re.sub(repl, fixed)
                        if fixed != match.group(0).strip():
                            suggestion = f"改为“{fixed}”" if fixed else "删除"
                    ctx = context_of(para.text, start, end)
                    issues.append(Issue(
                        id=rule.id, level=rule.level, check=rule.check, category=rule.category,
                        gate=rule.gate, line=para.line, excerpt=excerpt, message=rule.message,
                        suggestion=suggestion, basis=rule.basis, context=ctx if ctx != excerpt else "",
                        start=start, end=end, para_index=para.index,
                    ))
        return issues

    def _resolve_yields(self, issues: List[Issue]) -> List[Issue]:
        yields = {r.id: set(r.yield_to) for r in self.rules if r.yield_to}
        if not yields:
            return issues
        # 按（段落, 规则）分组，只与让位目标规则的命中比较，避免同段问题很多时退化为二次方
        by_para_rule: Dict[Tuple[int, str], List[Issue]] = {}
        for issue in issues:
            by_para_rule.setdefault((issue.para_index, issue.id), []).append(issue)
        kept: List[Issue] = []
        for issue in issues:
            targets = yields.get(issue.id)
            if targets:
                overlapped = any(
                    other.start < issue.end and issue.start < other.end
                    for target in targets
                    for other in by_para_rule.get((issue.para_index, target), ())
                )
                if overlapped:
                    continue
            kept.append(issue)
        return kept

    def _apply_once(self, issues: List[Issue]) -> List[Issue]:
        once_ids = {r.id for r in self.rules if r.once}
        if not once_ids:
            return issues
        first: Dict[str, Issue] = {}
        counts: Dict[str, int] = {}
        kept: List[Issue] = []
        for issue in sorted(issues, key=lambda x: (x.line, x.start)):
            if issue.id in once_ids:
                counts[issue.id] = counts.get(issue.id, 0) + 1
                if issue.id in first:
                    continue
                first[issue.id] = issue
            kept.append(issue)
        for rid, issue in first.items():
            if counts[rid] > 1:
                issue.count = counts[rid]
                issue.message = f"{issue.message}（全文共 {counts[rid]} 处，只列第一处）"
        return kept

    # -- AI 味、防御性 --------------------------------------------------------
    def _whitelist_spans(self, para: Para) -> List[Tuple[int, int]]:
        """词表白名单命中区间（【待补：】【待核实：】标记见 WL-018）。"""
        return [m.span() for _wid, regex in self.whitelist
                for m in regex.finditer(para.text) if m.end() > m.start()]

    def style_hits(self, doc: Document, pat: StylePattern,
                   wl_spans: Optional[Callable[[Para], List[Tuple[int, int]]]] = None) -> List[Tuple[Para, int, int]]:
        """逐段匹配一条词表规则，剔除与白名单重叠的命中，返回 (段落, 起, 止)。不做密度判断。"""
        if wl_spans is None:
            wl_spans = self._whitelist_spans
        hits: List[Tuple[Para, int, int]] = []
        for para in doc.paras:
            for match in pat.regex.finditer(para.text):
                start, end = match.span()
                if start == end:
                    continue
                if any(ws < end and start < we for ws, we in wl_spans(para)):
                    continue  # 与白名单重叠的命中不报告
                # 去掉正则为定位句首而带上的前导标点；命中本身就是标点（如顿号、冒号）时保留
                trimmed = start
                while trimmed < end and para.text[trimmed] in "。！？；，,：:、 \t　":
                    trimmed += 1
                hits.append((para, trimmed if trimmed < end else start, end))
        return hits

    def _run_style(self, doc: Document, doc_type: str, checks: Set[str]) -> List[Issue]:
        wanted = [p for p in self.patterns if p.group in checks]
        if not wanted:
            return []
        wl_cache: Dict[int, List[Tuple[int, int]]] = {}

        def wl_spans(para: Para) -> List[Tuple[int, int]]:
            if para.index not in wl_cache:
                wl_cache[para.index] = self._whitelist_spans(para)
            return wl_cache[para.index]

        issues: List[Issue] = []
        total = max(doc.char_count, 1)
        for pat in wanted:
            if not applies_to(doc_type, pat.doc_types, pat.exempt_doc_types):
                continue
            hits = self.style_hits(doc, pat, wl_spans)
            if not hits:
                continue
            check_name = pat.category or CHECK_NAMES.get(pat.group, pat.group)
            basis_file = "ai-tone.md" if pat.group == "style" else "defensive.md"
            basis = f"文风词表 {pat.id}（{pat.category}），见 references/_shared/style/{basis_file}"
            if pat.density is not None:
                per_k = len(hits) * 1000.0 / total
                if len(hits) < 2 or per_k <= pat.density:
                    continue
                para, start, end = hits[0]
                issues.append(Issue(
                    id=pat.id, level=pat.level, check=pat.group, category=check_name, gate="文风关",
                    line=para.line, excerpt=clip(para.text[start:end], 80),
                    message=f"{pat.message}（全文 {len(hits)} 处，每千字 {per_k:.1f} 处，阈值 {pat.density:g}）",
                    suggestion=pat.suggestion, basis=basis,
                    context=context_of(para.text, start, end), start=start, end=end,
                    count=len(hits), para_index=para.index,
                ))
            else:
                for para, start, end in hits:
                    ctx = context_of(para.text, start, end)
                    excerpt = clip(para.text[start:end], 80)
                    issues.append(Issue(
                        id=pat.id, level=pat.level, check=pat.group, category=check_name, gate="文风关",
                        line=para.line, excerpt=excerpt, message=pat.message,
                        suggestion=pat.suggestion, basis=basis,
                        context=ctx if ctx != excerpt else "", start=start, end=end,
                        para_index=para.index,
                    ))
        return issues

    # -- 标点（函数） ---------------------------------------------------------
    def _check_quotes(self, doc: Document, doc_type: str) -> List[Issue]:
        """PUN-009 引号嵌套：外层双引号内又出现双引号。"""
        issues: List[Issue] = []
        for para in doc.paras:
            text = para.text
            if "“" not in text:
                continue
            depth, outer = 0, -1
            for pos, ch in enumerate(text):
                if ch == "“":
                    if depth == 0:
                        outer = pos
                    elif depth >= 1:
                        closes = [m.start() for m in re.finditer("”", text[pos:])]
                        end = pos + (closes[1] if len(closes) > 1 else (closes[0] if closes else 0)) + 1
                        end = min(max(end, pos + 1), outer + 60, len(text))
                        issue = self._proc_issue("PUN-009", para, outer, end)
                        if issue:
                            issues.append(issue)
                        break
                    depth += 1
                elif ch == "”":
                    depth = max(0, depth - 1)
        return issues

    FIELD_LABEL_RE = re.compile(
        r"(?:主送|抄送|附件|联系人|联系电话|电话|传真|电子邮箱|邮箱|地址|签发人|时间|地点|主持人?|出席|列席|参加人员|"
        r"记录人?|议题|牵头单位|责任单位|配合单位|责任人|完成时限|时间节点|报送方式|邮编|网址|开户行|账号|户名)\s*$"
    )

    def _check_colon_nesting(self, doc: Document, doc_type: str) -> List[Issue]:
        """PUN-014 一句之内套用冒号。"""
        issues: List[Issue] = []
        skip = set(doc.head) | set(doc.title)
        if doc.banji_start is not None:
            skip |= set(range(doc.banji_start, len(doc.paras)))
        for para in doc.paras:
            if para.index in skip or para.in_table or "：" not in para.text:
                continue
            text = para.text
            marker_spans = [mk.span() for mk in MARKER_RE.finditer(text)]
            sent_start = 0
            for m in re.finditer(r"[。！？；]|$", text):
                sent_end = m.start()
                colons = []
                for c in re.finditer("：", text[sent_start:sent_end]):
                    pos = sent_start + c.start()
                    if any(s <= pos < e for s, e in marker_spans):
                        continue
                    if pos > 0 and pos + 1 < len(text) and text[pos - 1].isdigit() and text[pos + 1].isdigit():
                        continue
                    if self.FIELD_LABEL_RE.search(text[max(sent_start, pos - 8):pos]):
                        continue
                    colons.append(pos)
                if len(colons) >= 2:
                    issue = self._proc_issue("PUN-014", para, sent_start, min(sent_end + 1, len(text)),
                                             excerpt=clip(text[sent_start:sent_end + 1], 60))
                    if issue:
                        issues.append(issue)
                sent_start = m.end()
                if sent_start >= len(text):
                    break
        return issues

    # -- 数字（函数） ---------------------------------------------------------
    DATE_STYLES = (
        ("年月日", re.compile(r"\d{4}\s*年\s*\d{1,2}\s*月\s*\d{1,2}\s*日")),
        ("短横线", re.compile(r"(?<![\d\-])\d{4}-\d{1,2}-\d{1,2}(?![\d\-])")),
        ("点号", re.compile(r"(?<![\d.])\d{4}\.\d{1,2}\.\d{1,2}(?![\d.])")),
        ("斜线", re.compile(r"(?<![\d/])\d{4}/\d{1,2}/\d{1,2}(?![\d/])")),
    )

    def _check_date_styles(self, doc: Document, doc_type: str) -> List[Issue]:
        """NUM-010 日期写法统一。"""
        found: Dict[str, List[Tuple[Para, int, int]]] = {}
        end_idx = doc.banji_start if doc.banji_start is not None else len(doc.paras)
        for para in doc.paras[:end_idx]:
            for name, regex in self.DATE_STYLES:
                for m in regex.finditer(para.text):
                    found.setdefault(name, []).append((para, m.start(), m.end()))
        issues: List[Issue] = []
        for name in ("点号", "斜线"):
            if name in found:
                para, s, e = found[name][0]
                issue = self._proc_issue(
                    "NUM-010", para, s, e,
                    message=f"日期用{name}分隔（如“{para.text[s:e]}”），不在 GB/T 15835 列举的写法中（共 {len(found[name])} 处）",
                    count=len(found[name]))
                if issue:
                    issues.append(issue)
        if "年月日" in found and "短横线" in found:
            para, s, e = found["短横线"][0]
            issue = self._proc_issue(
                "NUM-010", para, s, e,
                message=f"同一文件日期写法不统一：“年月日”写法 {len(found['年月日'])} 处，短横线写法 {len(found['短横线'])} 处",
                count=len(found["短横线"]))
            if issue:
                issues.append(issue)
        return issues

    # -- 格式要素（函数） -----------------------------------------------------
    FULL_DATE_RE = re.compile(r"(\d{4})\s*年\s*(\d{1,2})\s*月(?:\s*(\d{1,2})\s*日)?")
    MD_DATE_RE = re.compile(r"(?<![\d年])(\d{1,2})\s*月\s*(\d{1,2})\s*日")

    def _check_date_padding(self, doc: Document, doc_type: str) -> List[Issue]:
        """FMT-005 月日虚位。"""
        issues: List[Issue] = []
        strict_idx = set()
        if doc.signoff_date is not None:
            strict_idx.add(doc.signoff_date)
        if doc.banji_start is not None:
            strict_idx |= {p.index for p in doc.paras[doc.banji_start:] if "印发" in p.text}
        for para in doc.paras:
            spans: List[Tuple[int, int]] = []
            for m in self.FULL_DATE_RE.finditer(para.text):
                month, day = m.group(2), m.group(3)
                if (len(month) == 2 and month.startswith("0")) or (day and len(day) == 2 and day.startswith("0")):
                    spans.append(m.span())
            for m in self.MD_DATE_RE.finditer(para.text):
                if any(s <= m.start() < e for s, e in spans):
                    continue
                month, day = m.group(1), m.group(2)
                if (len(month) == 2 and month.startswith("0")) or (len(day) == 2 and day.startswith("0")):
                    spans.append(m.span())
            for s, e in spans:
                level = "P0" if para.index in strict_idx else "P1"
                fixed = re.sub(r"(?<!\d)0(\d)(?=\s*[月日])", r"\1", para.text[s:e])
                issue = self._proc_issue("FMT-005", para, s, e, level=level, suggestion=f"改为“{fixed}”")
                if issue:
                    issues.append(issue)
        return issues

    def _check_signoff_date(self, doc: Document, doc_type: str) -> List[Issue]:
        """FMT-006 成文日期用阿拉伯数字。"""
        k = doc.signoff_date
        if k is None or not CN_DATE_ONLY_RE.match(doc.paras[k].text):
            return []
        para = doc.paras[k]
        text = para.text.strip()
        start = para.text.index(text)
        issue = self._proc_issue("FMT-006", para, start, start + len(text))
        return [issue] if issue else []

    def _check_zhusong(self, doc: Document, doc_type: str) -> List[Issue]:
        """FMT-007 主送机关后用全角冒号。"""
        issues: List[Issue] = []
        if doc.zhusong is not None:
            para = doc.paras[doc.zhusong]
            text = para.text.rstrip()
            if text.endswith(":"):
                issue = self._proc_issue("FMT-007", para, 0, len(text), excerpt=clip(text.strip(), 80),
                                         message="主送机关后用了半角冒号", suggestion="改为全角冒号“：”")
                if issue:
                    issues.append(issue)
            return issues
        if not doc.title or doc_type not in ZHUSONG_TYPES:
            return issues
        j = doc.title[-1] + 1
        if doc.subtitle is not None:
            j = doc.subtitle + 1
        if j >= len(doc.paras) - 1:
            return issues
        para = doc.paras[j]
        text = para.text.strip()
        if (len(text) <= 60 and not re.search(r"[。；：:！？]", text) and ORG_SUFFIX_RE.search(text)
                and not HEADING_START_RE.match(text) and not DATE_ONLY_RE.match(text)):
            issue = self._proc_issue("FMT-007", para, 0, len(para.text), excerpt=clip(text, 80),
                                     message="主送机关后缺少全角冒号")
            if issue:
                issues.append(issue)
        return issues

    # -- 结构（函数） ---------------------------------------------------------
    def _check_title(self, doc: Document, doc_type: str) -> List[Issue]:
        """STR-005 标题末尾标点；STR-010 文种混用或与参数不一致。"""
        issues: List[Issue] = []
        if not doc.title:
            return issues
        last = doc.paras[doc.title[-1]]
        text = last.text.rstrip()
        if text and text[-1] in PUNCT_TAIL:
            issue = self._proc_issue("STR-005", last, len(text) - 1, len(text),
                                     excerpt=clip(text.strip(), 80),
                                     suggestion=f"删除标题末尾的“{text[-1]}”")
            if issue:
                issues.append(issue)
        title = doc.title_text.rstrip(PUNCT_TAIL)
        first = doc.paras[doc.title[0]]
        if title.endswith("请示报告"):
            issue = self._proc_issue("STR-010", first, 0, len(first.text), excerpt=clip(first.text.strip(), 80),
                                     message="标题文种写作“请示报告”，请示和报告不能合用")
            if issue:
                issues.append(issue)
        else:
            inferred = doc.infer_doc_type()
            if (inferred and doc_type != inferred and doc_type in STATUTORY_TYPES
                    and inferred in STATUTORY_TYPES):
                issue = self._proc_issue(
                    "STR-010", first, 0, len(first.text), excerpt=clip(first.text.strip(), 80), level="P2",
                    message=f"标题文种为“{inferred}”，与指定的文书类型“{doc_type}”不一致",
                    suggestion="核对文种：请求批准用请示，汇报工作用报告，布置工作用通知")
                if issue:
                    issues.append(issue)
        return issues

    def _attach_list_indices(self, doc: Document) -> Set[int]:
        """附件说明后续行（“2. 某某表”）不参与层次序号检查。"""
        result: Set[int] = set()
        if doc.attach_desc is None:
            return result
        k = doc.attach_desc + 1
        while k < len(doc.paras) and re.match(r"^\s*\d{1,2}\s*[.．、]", doc.paras[k].text):
            result.add(k)
            k += 1
        return result

    def _check_heading_sequence(self, doc: Document, doc_type: str) -> List[Issue]:
        """STR-001 层次序号连续性与逐级使用。"""
        issues: List[Issue] = []
        skip = self._attach_list_indices(doc)
        if doc.attach_desc is not None:
            skip.add(doc.attach_desc)
        last: Dict[int, Optional[int]] = {1: None, 2: None, 3: None, 4: None}
        cur_level = 0
        seen_any = False
        jump_issue: Optional[Issue] = None
        jump_count = 0
        list_reset = False
        for k in doc.body_indices():
            para = doc.paras[k]
            if para.in_table or k in skip:
                continue
            if ATTACH_PAGE_RE.match(para.text):
                last = {1: None, 2: None, 3: None, 4: None}
                cur_level, seen_any = 0, False
                continue
            found = heading_of(para.text)
            if not found:
                # 以冒号结尾的引导句（如“一等助学金：”）之后，阿拉伯数字序号可以重新从 1 编起
                if para.text.rstrip().endswith(("：", ":")):
                    list_reset = True
                continue
            level, num = found
            # 引导句之后从 1 重新编号的，视为新列表；接着上一组继续编号的，照常检查连续性
            if list_reset and level >= 3 and num == 1:
                for lv in range(level, 5):
                    last[lv] = None
            list_reset = False
            head_text = para.text.strip()[:30]
            if seen_any and level > cur_level + 1 and cur_level > 0:
                jump_count += 1
                if jump_issue is None:
                    jump_issue = self._proc_issue(
                        "STR-001", para, 0, len(para.text), level="P2", excerpt=clip(head_text, 30),
                        message="层次序数越级使用（如“一、”之后直接用“1.”，跳过“（一）”）",
                        suggestion="按 一、（一）1.（1） 逐级使用")
            prev = last[level]
            expected = 1 if prev is None else prev + 1
            if num != expected:
                if prev is None:
                    msg = f"本级序号未从 1 开始（出现的是第 {num} 项）"
                    lvl = "P2" if not seen_any else "P1"
                elif num == prev:
                    msg, lvl = f"层次序号重复（第 {num} 项出现两次）", "P1"
                elif num < prev:
                    msg, lvl = f"层次序号乱序（第 {prev} 项之后出现第 {num} 项）", "P1"
                else:
                    msg, lvl = f"层次序号跳号（第 {prev} 项之后直接是第 {num} 项）", "P1"
                issue = self._proc_issue("STR-001", para, 0, len(para.text), level=lvl,
                                         excerpt=clip(head_text, 30), message=msg)
                if issue:
                    issues.append(issue)
            last[level] = num
            for deeper in range(level + 1, 5):
                last[deeper] = None
            cur_level = level
            seen_any = True
        if jump_issue:
            if jump_count > 1:
                jump_issue.message += f"（全文共 {jump_count} 处，只列第一处）"
                jump_issue.count = jump_count
            issues.append(jump_issue)
        return issues

    COUNT_RE = re.compile(
        r"(?<![第\d一二三四五六七八九十])(十[一二三四五六七八九]?|[两二三四五六七八九])"
        r"(?:(?:个|项|点|条|件|大)(?:方面|措施|任务|重点工作|重点任务|工作|意见|建议|体会|要求|问题|环节|阶段|举措"
        r"|原则|目标|机制|制度|实事|行动|工程|做法|经验|特点|成效|不足|短板|方向|内容|事项|领域|步骤|变化)|方面)"
        r"(?=(?:的|方面的)?[一-鿿]{0,6}(?:[：:，,。；]|$))"
    )
    INLINE_MARKERS = (
        re.compile(r"(?<![一-鿿])([一二三四五六七八九十]{1,3})是|(?<=[。；，：\s])([一二三四五六七八九十]{1,3})是"),
        re.compile(r"第([一二三四五六七八九十]{1,3})[，,、是]"),
        re.compile(r"[（(]([一二三四五六七八九十]{1,3})[）)]"),
        re.compile(r"[（(](\d{1,2})[）)]"),
    )

    @classmethod
    def _inline_count(cls, text: str) -> int:
        """统计“一是……二是……”等段内列举的连续条数。"""
        best = 0
        for regex in cls.INLINE_MARKERS:
            nums = []
            for m in regex.finditer(text):
                raw = next(g for g in m.groups() if g)
                nums.append(cn_to_int(raw))
            if not nums or nums[0] != 1:
                continue
            count = 0
            for n in nums:
                if n == count + 1:
                    count = n
            best = max(best, count)
        return best

    @staticmethod
    def _inside_quote(text: str, pos: int) -> bool:
        before = text[:pos]
        return before.count("“") > before.count("”") or before.count("\"") % 2 == 1

    def _check_count_summary(self, doc: Document, doc_type: str) -> List[Issue]:
        """STR-002 数字概括与实际条数不一致（启发式，P2）。"""
        issues: List[Issue] = []
        body = doc.body_indices()
        body_set = set(body)
        heading_cache = {k: heading_of(doc.paras[k].text) for k in body}
        for pos_in_body, k in enumerate(body):
            para = doc.paras[k]
            if para.in_table:
                continue
            for m in self.COUNT_RE.finditer(para.text):
                if self._inside_quote(para.text, m.start()):
                    continue
                n = cn_to_int(m.group(1))
                if n < 2:
                    continue
                tail = para.text[m.end():]
                found = self._inline_count(tail)
                if found < 2:
                    found = self._following_count(doc, body, pos_in_body, heading_cache, body_set)
                if found >= 2 and found != n:
                    issue = self._proc_issue(
                        "STR-002", para, m.start(), m.end(),
                        message=f"数字概括为“{m.group(0)}”，其后实际列出 {found} 项",
                        suggestion="核对条数，改正概括数字或补删条目")
                    if issue:
                        issues.append(issue)
        return issues

    def _following_count(self, doc: Document, body: List[int], pos: int,
                         heading_cache: Dict[int, Optional[Tuple[int, int]]], body_set: Set[int]) -> int:
        """在后续段落中统计下一级标题或“一是”段的条数。"""
        k0 = body[pos]
        own = heading_cache.get(k0)
        ctx_level = own[0] if own else 0
        if not own:
            for back in range(pos - 1, -1, -1):
                h = heading_cache.get(body[back])
                if h:
                    ctx_level = h[0]
                    break
        child_level: Optional[int] = None
        count = 0
        yishi = 0
        for idx in range(pos + 1, min(len(body), pos + 40)):
            k = body[idx]
            h = heading_cache.get(k)
            text = doc.paras[k].text.strip()
            if h:
                level = h[0]
                if level <= ctx_level:
                    break
                if child_level is None:
                    child_level = level
                if level < child_level:
                    break
                if level == child_level:
                    count += 1
                continue
            if child_level is None:
                m = re.match(r"^([一二三四五六七八九十]{1,3})是", text)
                if m and cn_to_int(m.group(1)) == yishi + 1:
                    yishi += 1
                elif idx > pos + 1 and yishi == 0:
                    break  # 概括句之后紧跟的不是条目，放弃判断
        return count if child_level is not None else yishi

    QINGSHI_END_RE = re.compile(
        r"(?:妥否|当否|可否|是否妥当)[，,]?\s*请(?:批示|批复|审批|审定|示)|请予以?(?:审批|批准|批复|审定|批示|支持)"
        r"|请(?:批示|批复|审批|审定)|恳请(?:批准|审批|批复|审定)|以上请示[，,]?\s*请"
    )
    REQUEST_RE = re.compile(
        r"(?:妥否|当否|可否)[，,]?\s*请(?:批示|批复|审批|示)|请予以?(?:审批|批准|批复|审定)"
        r"|(?<![报提申呈])请(?:批示|审批|批准|批复)(?![件人])"
    )
    PLEA_RE = re.compile(r"恳请[^。；\n]{0,20}(?:批准|审批|批复|同意|拨付|核拨|解决)")

    def _check_qingshi(self, doc: Document, doc_type: str) -> List[Issue]:
        """STR-003 请示结尾；STR-006 主送；STR-007 特此报告。"""
        issues: List[Issue] = []
        if doc_type != "请示":
            if doc_type == "报告":
                issues.extend(self._check_main_recipient(doc, doc_type))
            return issues
        body = doc.body_indices()
        if body and self._proc_applies("STR-003", doc_type):
            if not any(self.QINGSHI_END_RE.search(doc.paras[k].text) for k in body):
                last_k = body[-1]
                for k in reversed(body):
                    if not DATE_ONLY_RE.match(doc.paras[k].text) and not ATTACH_DESC_RE.match(doc.paras[k].text):
                        last_k = k
                        if len(doc.paras[k].text.strip()) > 20:
                            break
                para = doc.paras[last_k]
                issue = self._proc_issue("STR-003", para, 0, len(para.text), excerpt=clip(para.text, 40))
                if issue:
                    issues.append(issue)
        for k in body:
            para = doc.paras[k]
            m = re.search(r"特此报告", para.text)
            if m and self._proc_applies("STR-007", doc_type):
                issue = self._proc_issue("STR-007", para, m.start(), m.end())
                if issue:
                    issues.append(issue)
        issues.extend(self._check_main_recipient(doc, doc_type))
        return issues

    def _check_main_recipient(self, doc: Document, doc_type: str) -> List[Issue]:
        if doc.zhusong is None or not self._proc_applies("STR-006", doc_type):
            return []
        para = doc.paras[doc.zhusong]
        text = para.text.strip().rstrip("：:")
        if re.search(r"(?:书记|市长|县长|区长|省长|局长|主任|院长|校长|董事长|总经理|厅长|部长|处长|同志)$", text):
            issue = self._proc_issue("STR-006", para, 0, len(para.text), excerpt=clip(para.text.strip(), 80),
                                     message=f"{doc_type}主送领导个人")
            return [issue] if issue else []
        if doc_type == "请示" and re.search(r"[、，,]", text):
            issue = self._proc_issue("STR-006", para, 0, len(para.text), excerpt=clip(para.text.strip(), 80),
                                     message="请示主送了多个机关")
            return [issue] if issue else []
        return []

    def _check_baogao(self, doc: Document, doc_type: str) -> List[Issue]:
        """STR-004 报告夹带请示事项。"""
        if not self._proc_applies("STR-004", doc_type):
            return []
        issues: List[Issue] = []
        for k in doc.body_indices():
            para = doc.paras[k]
            spans: List[Tuple[int, int]] = []
            for m in self.REQUEST_RE.finditer(para.text):
                spans.append(m.span())
                issue = self._proc_issue("STR-004", para, m.start(), m.end())
                if issue:
                    issues.append(issue)
            for m in self.PLEA_RE.finditer(para.text):
                if any(s < m.end() and m.start() < e for s, e in spans):
                    continue
                issue = self._proc_issue("STR-004", para, m.start(), m.end(), level="P1",
                                         message="报告中出现疑似请示事项（请求批准、拨付等）",
                                         suggestion="需要上级批准的事项另行请示；确属反映困难的，改为陈述情况")
                if issue:
                    issues.append(issue)
        return issues

    def _check_han(self, doc: Document, doc_type: str) -> List[Issue]:
        """STR-008 函中用“请批示”；STR-011 下行文用请示结束语。"""
        issues: List[Issue] = []
        if self._proc_applies("STR-008", doc_type):
            for k in doc.body_indices():
                para = doc.paras[k]
                for m in re.finditer(r"请(?:批示|批复)", para.text):
                    issue = self._proc_issue("STR-008", para, m.start(), m.end())
                    if issue:
                        issues.append(issue)
        if self._proc_applies("STR-011", doc_type):
            for k in doc.body_indices():
                para = doc.paras[k]
                for m in self.REQUEST_RE.finditer(para.text):
                    issue = self._proc_issue("STR-011", para, m.start(), m.end())
                    if issue:
                        issues.append(issue)
        return issues

    def _check_shangxing(self, doc: Document, doc_type: str) -> List[Issue]:
        """STR-009 上行文有发文字号却没有签发人。"""
        if not self._proc_applies("STR-009", doc_type) or not doc.wenhao_present or doc.qianfa_present:
            return []
        for k in sorted(doc.head):
            para = doc.paras[k]
            m = WENHAO_RE.search(para.text)
            if m:
                issue = self._proc_issue("STR-009", para, m.start(), m.end())
                return [issue] if issue else []
        return []

    # -- 隐私与保密（函数） ---------------------------------------------------
    ID18_RE = re.compile(
        r"(?<![0-9A-Za-z])([1-9]\d{5})((?:18|19|20)\d{2})(0[1-9]|1[0-2])(0[1-9]|[12]\d|3[01])(\d{3})([\dXx])(?![0-9A-Za-z])"
    )
    MOBILE_RE = re.compile(r"(?<![\d])(?:\+?86[-\s]?)?(1[3-9]\d)[-\s]?(\d{4})[-\s]?(\d{4})(?![\d])")
    BANK_RE = re.compile(r"(?<![\d])(\d{4}(?:[\s-]?\d{4}){3}(?:[\s-]?\d{1,3})?)(?![\d])")
    LABELED_RE = re.compile(
        r"(病案号|住院号|门诊号|就诊卡号|医保卡号|社保卡号|学号|学籍号|考生号|准考证号|护照号|工号|员工号|职工号)"
        r"\s*(?:为|是)?\s*[：:]?\s*([A-Za-z0-9][A-Za-z0-9\-]{3,})(?![A-Za-z0-9*＊])"
    )
    ADDR_RE = re.compile(
        r"(家庭住址|家庭地址|现住址|住址|户籍地址|居住地址|现居住地|户籍所在地)\s*(?:为|是|在)?\s*[：:]?\s*([^\s，。；,;）)]{4,40})"
    )
    NAME_GENDER_RE = re.compile(r"([一-鿿]{2,6})\s*[，,（(、]\s*(男|女)\s*[，,、）)]?\s*(\d{1,3})\s*(?:岁|周岁)")
    BED_NAME_RE = re.compile(r"(\d{1,3})\s*(?:床|号床)(?![位单旁边头尾日次垫上下])\s*[：:]?\s*([一-鿿]{2,3})")
    CONTACT_HINT_RE = re.compile(r"联系|电话|手机|热线|咨询|值班|拨打|致电|手机号|联系方式|Tel|TEL|tel")
    PERSONAL_BANK_HINT = re.compile(r"银行卡|卡号|储蓄卡|信用卡|借记卡|社保卡")
    ACCOUNT_HINT = re.compile(r"账号|账户|帐号|帐户|开户")

    @staticmethod
    def _id_checksum_ok(id17: str, check: str) -> bool:
        weights = (7, 9, 10, 5, 8, 4, 2, 1, 6, 3, 7, 9, 10, 5, 8, 4, 2)
        total = sum(int(d) * w for d, w in zip(id17, weights))
        return "10X98765432"[total % 11] == check.upper()

    @staticmethod
    def _luhn_ok(digits: str) -> bool:
        total = 0
        for i, ch in enumerate(reversed(digits)):
            d = int(ch)
            if i % 2 == 1:
                d *= 2
                if d > 9:
                    d -= 9
            total += d
        return total % 10 == 0

    @staticmethod
    def _label_context(text: str, start: int) -> str:
        """隐私问题的 context：只取号码前的标签文字，不含号码本身。"""
        return text[max(0, start - 12):start].strip()

    @staticmethod
    def _looks_like_name(name: str) -> bool:
        if not name or any(ch in MASK_CHARS for ch in name):
            return False
        if name in ROLE_WORDS or any(name.endswith(w) for w in ROLE_WORDS):
            return False
        return name[0] in SURNAMES

    def _check_privacy(self, doc: Document, doc_type: str) -> List[Issue]:
        issues: List[Issue] = []
        for para in doc.paras:
            text = para.text
            taken: List[Tuple[int, int]] = []

            def overlaps(s: int, e: int) -> bool:
                return any(a < e and s < b for a, b in taken)

            def add(rid: str, s: int, e: int, masked: str, **kw: Any) -> None:
                issue = self._proc_issue(rid, para, s, e, excerpt=masked,
                                         context=self._label_context(text, s), **kw)
                if issue:
                    issues.append(issue)
                    taken.append((s, e))

            # 身份证号
            for m in self.ID18_RE.finditer(text):
                full = m.group(0)
                try:
                    date(int(m.group(2)), int(m.group(3)), int(m.group(4)))
                except ValueError:
                    continue
                masked = full[:6] + "********" + full[-4:]
                if self._id_checksum_ok(full[:17], full[17]):
                    add("PRI-001", m.start(), m.end(), masked)
                else:
                    add("PRI-001", m.start(), m.end(), masked, level="P1",
                        message="疑似身份证号（校验位不符，可能是录入错误）")
            # 带标签的编号
            for m in self.LABELED_RE.finditer(text):
                s, e = m.span(2)
                if overlaps(s, e):
                    continue
                value = m.group(2)
                masked = value[:2] + "****" + value[-2:] if len(value) > 4 else "****"
                level = "P1" if m.group(1) in ("工号", "员工号", "职工号") else None
                add("PRI-004", s, e, f"{m.group(1)}：{masked}", level=level,
                    message=f"出现{m.group(1)}，属于可识别个人的编号")
            # 银行卡号
            for m in self.BANK_RE.finditer(text):
                s, e = m.span(1)
                if overlaps(s, e):
                    continue
                digits = re.sub(r"\D", "", m.group(1))
                if not 16 <= len(digits) <= 19:
                    continue
                before = text[max(0, s - 15):s]
                masked = digits[:4] + " **** **** " + digits[-4:]
                if self.PERSONAL_BANK_HINT.search(before):
                    add("PRI-003", s, e, masked)
                elif self.ACCOUNT_HINT.search(before):
                    add("PRI-003", s, e, masked, level="P2",
                        message="出现银行账号，如为个人账户须脱敏；单位对公账户可保留",
                        suggestion="确认是否为单位对公账户；个人账户删除或脱敏")
                elif self._luhn_ok(digits) and digits[:1] in ("4", "5", "6") and (digits[0] != "6" or digits[:2] == "62"):
                    add("PRI-003", s, e, masked, level="P1", message="疑似银行卡号（无标签，Luhn 校验通过）")
            # 家庭住址
            for m in self.ADDR_RE.finditer(text):
                s, e = m.span(2)
                if overlaps(s, e):
                    continue
                value = m.group(2)
                if not re.search(r"\d|号|室|幢|单元|栋|楼|村|组|弄|巷", value):
                    continue
                add("PRI-008", s, e, f"{m.group(1)}：{value[:4]}****")
            # 姓名 + 性别 + 年龄
            for m in self.NAME_GENDER_RE.finditer(text):
                raw = m.group(1)
                for role in ROLE_WORDS:
                    if role in raw:
                        raw = raw.split(role)[-1]
                candidates = [raw[-3:], raw[-2:]] if len(raw) >= 3 else [raw]
                name = next((c for c in candidates if self._looks_like_name(c)), "")
                if not name or any(ch in MASK_CHARS for ch in raw):
                    continue
                s = m.start(1) + m.group(1).rfind(name)
                e = s + len(name)
                if overlaps(s, e):
                    continue
                add("PRI-009", s, e, name[0] + "*" * (len(name) - 1) + f"，{m.group(2)}，{m.group(3)}岁")
            # 床号 + 姓名
            for m in self.BED_NAME_RE.finditer(text):
                name = m.group(2)
                if not self._looks_like_name(name):
                    continue
                s, e = m.span(2)
                if overlaps(s, e):
                    continue
                add("PRI-010", s, e, f"{m.group(1)}床{name[0]}{'*' * (len(name) - 1)}")
            # 手机号
            for m in self.MOBILE_RE.finditer(text):
                s, e = m.span()
                if overlaps(s, e):
                    continue
                if re.match(r"\s*[元万亿千百吨件人次个户]", text[e:e + 2]):
                    continue
                digits = re.sub(r"\D", "", m.group(0))[-11:]
                masked = digits[:3] + "****" + digits[-4:]
                window = text[max(0, s - 20):s] + text[e:e + 8]
                with_identity = any(i.para_index == para.index and i.id in ("PRI-001", "PRI-004", "PRI-008", "PRI-009")
                                    for i in issues)
                if with_identity:
                    add("PRI-002", s, e, masked, message="出现个人手机号（与姓名、身份证号、病案号或住址同段）")
                elif self.CONTACT_HINT_RE.search(window):
                    add("PRI-002", s, e, masked, level="P2",
                        message="出现手机号码（工作联系方式）。对外张贴、公开发布的，确认本人同意或改用工作电话",
                        suggestion="对外版本改用座机或工作电话；内部报送可保留")
                else:
                    add("PRI-002", s, e, masked)
        return issues

    # -- 篇幅（函数） ---------------------------------------------------------
    def _check_length(self, doc: Document, doc_type: str) -> List[Issue]:
        n = doc.body_char_count()
        anchor = doc.paras[doc.title[0]] if doc.title else (doc.paras[0] if doc.paras else None)
        if anchor is None:
            return []
        excerpt = clip(anchor.text, 40)
        if doc_type in LENGTH_LIMIT_TYPES:
            if n > 5000:
                msg = (f"正文约 {n} 字，超过 5000 字。除部署综合性工作外，地方和部门文件一般不超过 5000 字；"
                       "部署综合性工作的未设上限，仍须坚持短实新")
            elif n > 4000:
                msg = f"正文约 {n} 字，超过 4000 字。部署专项工作或者具体任务的文件一般不超过 4000 字"
            else:
                return []
            issue = self._proc_issue("LEN-001", anchor, 0, len(anchor.text), excerpt=excerpt, message=msg)
            return [issue] if issue else []
        if doc_type == "讲话稿" and n > 12000:
            issue = self._proc_issue("LEN-002", anchor, 0, len(anchor.text), excerpt=excerpt,
                                     message=f"讲话稿约 {n} 字，按每分钟约 200 字估算超过 1 小时")
            return [issue] if issue else []
        return []


# ---------------------------------------------------------------------------
# 对外函数
# ---------------------------------------------------------------------------

def parse_checks(value: Optional[str]) -> List[str]:
    """解析 --checks，返回按标准顺序排列的类别列表；非法值抛出 InputError。"""
    if not value or value.strip().lower() == "all":
        return list(CHECK_ORDER)
    wanted = [v.strip().lower() for v in re.split(r"[,，\s]+", value) if v.strip()]
    if "all" in wanted:
        return list(CHECK_ORDER)
    bad = [v for v in wanted if v not in CHECK_NAMES]
    if bad:
        raise InputError(f"未知的检查类别：{'、'.join(bad)}。可选：all、{'、'.join(CHECK_ORDER)}")
    return [c for c in CHECK_ORDER if c in wanted]


def resolve_doc_type(doc: Document, given: Optional[str]) -> Tuple[str, str]:
    if given and given.strip():
        return given.strip(), "参数"
    inferred = doc.infer_doc_type()
    if inferred:
        return inferred, "标题推断"
    return "其他", "默认"


def lint_document(doc: Document, doc_type: Optional[str] = None, checks: Optional[str] = None,
                  linter: Optional[Linter] = None, notes: Optional[List[str]] = None) -> Dict[str, Any]:
    linter = linter or Linter()
    check_list = parse_checks(checks)
    dtype, source = resolve_doc_type(doc, doc_type)
    all_notes = list(notes or [])
    if dtype not in DOC_TYPES:
        all_notes.append(f"文书类型“{dtype}”不在统一词表中，按文种限定的规则不会生效。")
    issues = linter.lint(doc, dtype, set(check_list))
    summary: Dict[str, Any] = {lv: sum(1 for i in issues if i.level == lv) for lv in LEVELS}
    # 待补、待核实标记单独计数（开发规范第 8 节），不论是否检查 marker 类
    summary["markers"] = {
        "待补": sum(len(TODO_MARK_RE.findall(p.text)) for p in doc.paras),
        "待核实": sum(len(VERIFY_MARK_RE.findall(p.text)) for p in doc.paras),
    }
    if doc.truncated_line is not None:
        all_notes.append(f"第 {doc.truncated_line} {doc.unit}为【起草说明】，其后内容未检查。")
    if "marker" in check_list:
        todo = sum(1 for i in issues if i.id == "MRK-001")
        verify = sum(1 for i in issues if i.id == "MRK-002")
        bad = sum(1 for i in issues if i.id == "MRK-003")
        if todo or verify or bad:
            all_notes.append(
                f"正文含【待补】{todo} 处、【待核实】{verify} 处"
                + (f"、格式不规范的标记 {bad} 处" if bad else "")
                + "，定稿前须逐条补齐或核实，并在【起草说明】中汇总。")
    if "length" in check_list:
        all_notes.append(f"篇幅口径：标题之后、【起草说明】之前的正文，不含空白和版记，约 {doc.body_char_count()} 字。")
    all_notes.extend(linter.warnings)
    return {
        "file": doc.file,
        "doc_type": dtype,
        "doc_type_source": source,
        "char_count": doc.char_count,
        "summary": summary,
        "issues": [i.to_dict() for i in issues],
        "checks": check_list,
        "notes": all_notes,
        "engine_version": VERSION,
    }


def lint_text(text: str, doc_type: Optional[str] = None, checks: Optional[str] = None,
              file_name: str = "<text>", linter: Optional[Linter] = None) -> Dict[str, Any]:
    """检查一段文本（供其他脚本和测试调用）。"""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    lines = [(i, line, False) for i, line in enumerate(text.split("\n"), start=1)]
    return lint_document(Document(file_name, lines), doc_type, checks, linter)


def lint_file(path: Path, doc_type: Optional[str] = None, checks: Optional[str] = None,
              linter: Optional[Linter] = None) -> Dict[str, Any]:
    if not path.is_file():
        raise InputError(f"文件不存在：{path}")
    lines, notes = read_lines(path)
    unit = "段" if path.suffix.lower() == ".docx" else "行"
    return lint_document(Document(path.name, lines, unit), doc_type, checks, linter, notes)


def render_text(result: Dict[str, Any]) -> str:
    unit = "段" if str(result["file"]).lower().endswith(".docx") else "行"
    s = result["summary"]
    src = "" if result["doc_type_source"] == "参数" else f"，{result['doc_type_source']}"
    out = [
        f"机检结果：{result['file']}（文书类型：{result['doc_type']}{src}；正文 {result['char_count']} 字）",
        f"P0 {s['P0']} 处｜P1 {s['P1']} 处｜P2 {s['P2']} 处",
    ]
    markers = s.get("markers") or {}
    if markers.get("待补") or markers.get("待核实"):
        out.append(f"待补 {markers.get('待补', 0)} 处、待核实 {markers.get('待核实', 0)} 处，定稿前补齐")
    if not result["issues"]:
        out.append("")
        out.append("未发现问题。机检只覆盖可规则化的项目，内容、提法、行文规则仍需人工审读。")
    for issue in result["issues"]:
        out.append("")
        out.append(f"[{issue['level']}] 第{issue['line']}{unit}  {issue['id']}  {issue['category']}：{issue['message']}")
        line = f"    原文：{issue['excerpt']}"
        if issue.get("context") and len(issue["excerpt"]) < 8:
            line += f"（上下文：{issue['context']}）"
        out.append(line)
        if issue.get("suggestion"):
            out.append(f"    建议：{issue['suggestion']}")
        if issue.get("basis"):
            out.append(f"    依据：{issue['basis']}")
    if result["notes"]:
        out.append("")
        out.extend(f"提示：{n}" for n in result["notes"])
    return "\n".join(out)


class _ChineseHelpFormatter(argparse.HelpFormatter):
    def add_usage(self, usage, actions, groups, prefix=None):  # type: ignore[override]
        return super().add_usage(usage, actions, groups, prefix="用法：")


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        prog="gongwen_lint.py",
        description="公文机检：残留、标点、数字、格式要素、结构、AI 味、防御性、隐私与保密、待补标记、篇幅。",
        epilog="退出码：0 正常；--strict 且存在 P0 时为 1；输入错误为 2。docx 输入时“行号”为段落序号。",
        add_help=False,
        formatter_class=_ChineseHelpFormatter,
    )
    parser._positionals.title = "位置参数"
    parser._optionals.title = "选项"
    parser.add_argument("file", help="待检查文件（.txt、.md、.docx），遇到【起草说明】行即停止读取")
    parser.add_argument("--doc-type", default=None,
                        help="文书类型，如 通知、请示、报告、总结、讲话稿；不填时按标题推断，推断不出按“其他”")
    parser.add_argument("--checks", default="all",
                        help="检查类别，逗号分隔：all（默认）、" + "、".join(CHECK_ORDER))
    parser.add_argument("--json", action="store_true", help="输出 JSON")
    parser.add_argument("--strict", action="store_true", help="存在 P0 时返回 1")
    parser.add_argument("-h", "--help", action="help", help="显示帮助并退出")
    return parser


def main(argv: Optional[Sequence[str]] = None) -> int:
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8")  # type: ignore[attr-defined]
        except Exception:
            pass
    parser = build_parser()
    args = parser.parse_args(argv)
    try:
        result = lint_file(Path(args.file), args.doc_type, args.checks)
    except InputError as exc:
        print(f"错误：{exc}", file=sys.stderr)
        return 2
    if args.json:
        print(json.dumps(result, ensure_ascii=False, indent=2))
    else:
        print(render_text(result))
    if args.strict and result["summary"]["P0"] > 0:
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
