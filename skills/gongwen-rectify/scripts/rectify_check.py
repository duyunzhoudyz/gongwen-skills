# -*- coding: utf-8 -*-
"""整改台账与整改报告对账。

属于 gongwen-rectify。台账是整改材料的唯一事实源，本脚本逐条比对报告与台账的编号和整改状态，
核对报告中的统计数和完成率，并对台账本身做完整性检查。

用法：
    python scripts/rectify_check.py <整改台账.md|.xlsx> [<报告.txt|.md|.docx>]
                                    [--feedback <反馈原文>] [--date YYYY-MM-DD] [--plan] [--json] [--strict]

只给台账时，只做台账自检；给了报告时，再做逐条对账和统计核对。
第二个文件是整改方案时加 --plan：只查是否逐条覆盖、问题表述是否弱化、有没有把未完成写成已完成，
不要求每条写明状态。

台账识别：
    Markdown 取第一张同时有"编号（或序号）"列和"整改状态（或状态、完成情况）"列的表格；
    Excel 在各工作表前 10 行中找同样的表头行。编号为空的行视为上一条的续行（合并单元格）。
    状态口径：已完成、立行立改、持续推进、长期坚持，另允许临时状态"未启动"。
    "已整改""正在推进""整改中"等同义写法按对应口径比对，并提示统一写法；
    "基本完成""部分完成"不是状态口径，提示改为"持续推进"或"已完成"。

报告中怎样认出编号：
    "问题3""问题编号3""问题 3-1""台账第3项""第3项问题"，以及台账中带前缀的编号原样出现
    （字母前缀如 XC-03，汉字前缀如 巡-1、审-1、自-1）。
    列举合写的也逐个认出："问题1、4、5""问题1和2""问题2至4""第1、4项问题"，这些编号共用后面的状态文字。
    报告只写到大项（问题3）而台账拆了子项（3-1、3-2）的，按子项汇总后的状态比对：子项全部完成才算完成。
    只有数字后面再接"-数字"的才是子项（3-1、XC-3-1）；XC-01、XC-02、巡-1 是各自独立的问题。

对账规则（级别与 gongwen_lint 一致：P0 必改，P1 建议改，P2 提示）：
    P0  台账未完成，报告该条写成"已完成整改""整改到位""整改完毕""已销号"，或主语是本条问题的
        "该问题已完成。""该问题已整改。""整改措施已落实到位。"等；
        报告的完成数、完成率高于台账（"已完成整改5个""已有5个完成整改""完成整改的有5个""6个问题已整改5个"等）；
        台账未全部完成，报告写"全部完成整改""所有问题已整改完毕""6个问题均已完成整改"等。
    P1  台账已完成，报告该条写成未完成；台账为持续推进，报告写成"长期坚持"；
        台账未完成，段中有"已完成整改"一类说法、段末又明确写了"该问题正在推进"（前后矛盾，降为 P1 提示核对）；
        未完成的台账条目报告未提及（已完成的未提及为 P2）；报告提到台账中没有的编号；统计数与台账不符；
        报告使用"基本完成"；
        台账状态为空或不在口径内、编号重复、关键栏为空、已完成却没有佐证、
        未完成条目的完成时限不是具体日期（"长期""尽快""近期"）、超过完成时限仍未完成（需 --date）；
        给了 --feedback 时，台账问题表现与反馈原文对不上（RC-301，可能被改写、弱化），
        或删减了对应条目中带数字、金额、限定词、引号内容的分句（RC-302，删数字、截短）；
        报告"问题N："后的问题表述比台账多了"有待""进一步"等弱化词。
    P2  报告该条没有写明状态；同义写法提示统一；--plan 时没写时限的目标句（"反馈问题全部整改到位"），
        以及分句里没有"已"的完成说法（"……，完成整改。"）；其他需要人工判断的差异。
    "本月新完成整改1个"这类本期数不与台账累计数比对，高于累计完成数的仍报出。
    "将""计划""预计""力争""确保""届时""……前"之后的"完成整改"是计划，不算已完成；句首"到2026年12月31日，"引出的
    也是目标；"未完成整改"不算已完成；"已完成整改方案制定""整改到位的标准""督促3个科室完成整改"
    "部分科室已完成整改"是名词性用法或部分完成，不算本条已完成；
    照录的问题表述（如"仍有1个社区未整改到位"）不当作状态说法。

报告只读【起草说明】之前的正文。--date 用于判断是否超期，缺省不判断。

退出码：0 正常；--strict 时存在 P0 返回 1；输入错误返回 2（文件不存在、格式不支持、台账缺编号列或状态列、
读取 .xlsx 缺 openpyxl、读取 .docx 缺 python-docx）。
"""
from __future__ import annotations

import argparse
import datetime as _dt
import json
import re
import sys
from collections import OrderedDict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Tuple

BOUNDARY = "【起草说明】"
DRAFT_NOTE_RE = re.compile(r"^\s*(?:#{1,6}\s*)?(?:(?:\*\*|__)\s*)?" + re.escape(BOUNDARY))
MARKER_RE = re.compile(r"【(待补|待核实)(?:[：:][^】\n]*)?】")
FULLWIDTH = str.maketrans("０１２３４５６７８９％．－—–ＡＢＣＤＥＦＧＨＩＪＫＬＭＮＯＰＱＲＳＴＵＶＷＸＹＺ",
                          "0123456789%.---ABCDEFGHIJKLMNOPQRSTUVWXYZ")

# ---------------------------------------------------------------------------
# 状态口径
# ---------------------------------------------------------------------------

DONE = "已完成"
IMMEDIATE = "立行立改"
ONGOING = "持续推进"
LONGTERM = "长期坚持"
NOT_STARTED = "未启动"
STATUSES = (DONE, IMMEDIATE, ONGOING, LONGTERM, NOT_STARTED)
DONE_SET = (DONE, IMMEDIATE)

SYNONYMS: Dict[str, str] = {
    "已完成": DONE, "完成": DONE, "已整改": DONE, "已整改完成": DONE, "整改完成": DONE,
    "已完成整改": DONE, "完成整改": DONE, "已整改到位": DONE, "整改到位": DONE, "已办结": DONE,
    "办结": DONE, "已销号": DONE, "销号": DONE,
    "立行立改": IMMEDIATE, "已立行立改": IMMEDIATE, "立行立改完成": IMMEDIATE, "已立改": IMMEDIATE,
    "持续推进": ONGOING, "正在推进": ONGOING, "推进中": ONGOING, "整改中": ONGOING,
    "正在整改": ONGOING, "进行中": ONGOING, "正在进行": ONGOING, "持续推进中": ONGOING,
    "长期坚持": LONGTERM, "需长期坚持": LONGTERM, "长期推进": LONGTERM, "常态化推进": LONGTERM,
    "长期整改": LONGTERM,
    "未启动": NOT_STARTED, "尚未启动": NOT_STARTED, "未开始": NOT_STARTED, "尚未开始": NOT_STARTED,
}
NONSTANDARD: Dict[str, str] = {
    "基本完成": "“基本完成”不是状态口径：措施全部落实、成效达标并有佐证的记“已完成”，否则记“持续推进”",
    "基本整改": "“基本整改”不是状态口径，按实际记“已完成”或“持续推进”",
    "大部分完成": "“大部分完成”不是状态口径，记“持续推进”",
    "部分完成": "“部分完成”不是状态口径，记“持续推进”",
    "初步完成": "“初步完成”不是状态口径，记“持续推进”",
    "未完成": "“未完成”不够具体：已动手的记“持续推进”，尚未动手的记“未启动”",
    "持续整改": "“持续整改”是审计整改要求类型，不是状态口径：未完成的记“持续推进”，确需长期抓的记“长期坚持”",
    "分阶段整改": "“分阶段整改”是审计整改要求类型，不是状态口径：未完成的记“持续推进”",
}


def normalize_status(raw: str) -> Tuple[Optional[str], str]:
    """返回 (口径, 说明)。口径为 None 表示无法认定；说明非空表示需提示。"""
    text = (raw or "").strip()
    if not text:
        return None, "empty"
    if MARKER_RE.search(text):
        return None, "marker"
    text = re.sub(r"[（(][^）)]*[）)]", "", text)
    text = re.sub(r"[\s*_`]", "", text).rstrip("。；;，,")
    head = re.split(r"[，,；;。:：]", text, maxsplit=1)[0]
    for candidate in (text, head):
        if candidate in NONSTANDARD:
            return None, "nonstandard:" + candidate
        if candidate in SYNONYMS:
            canon = SYNONYMS[candidate]
            return canon, "" if candidate == canon else "synonym:" + candidate
    return None, "unknown"


def aggregate_status(statuses: List[Optional[str]]) -> Optional[str]:
    """多个子项汇总为一个状态：全部完成才算完成。"""
    known = [s for s in statuses if s]
    if not known or len(known) < len(statuses):
        return None
    if all(s in DONE_SET for s in known):
        return IMMEDIATE if all(s == IMMEDIATE for s in known) else DONE
    if all(s == NOT_STARTED for s in known):
        return NOT_STARTED
    if all(s in DONE_SET or s == LONGTERM for s in known):
        return LONGTERM
    return ONGOING


# ---------------------------------------------------------------------------
# 数据结构
# ---------------------------------------------------------------------------

class InputError(Exception):
    """输入错误：退出码 2。"""


@dataclass
class Item:
    raw_id: str
    key: str
    row: int
    fields: Dict[str, str] = field(default_factory=dict)
    status_raw: List[str] = field(default_factory=list)
    status: Optional[str] = None


@dataclass
class Ledger:
    path: Path
    items: "OrderedDict[str, Item]"
    columns: Dict[str, str]
    where: str  # "第 1 张表" 或 "工作表 Sheet1"
    issues: List[dict] = field(default_factory=list)


FIELD_KEYWORDS: List[Tuple[str, Tuple[str, ...]]] = [
    ("id", ("问题编号", "编号", "问题序号", "序号")),
    ("status", ("整改状态", "完成状态", "状态", "整改进度", "完成情况")),
    ("problem", ("问题表现", "反馈问题", "问题描述", "具体问题", "存在问题", "问题")),
    ("measures", ("整改措施", "措施")),
    ("leader", ("责任领导", "牵头领导", "分管领导")),
    ("dept", ("责任部门", "责任单位", "牵头部门", "牵头单位")),
    ("deadline", ("完成时限", "整改时限", "时限", "完成时间")),
    ("standard", ("成效标准", "验收标准", "完成标准", "整改目标")),
    ("evidence", ("佐证材料", "佐证", "证明材料", "销号材料")),
    ("progress", ("进展情况", "整改进展", "推进情况", "整改情况")),
]
FIELD_NAMES = {
    "id": "编号", "status": "整改状态", "problem": "问题表现", "measures": "整改措施",
    "leader": "责任领导", "dept": "责任部门", "deadline": "完成时限", "standard": "成效标准",
    "evidence": "佐证材料", "progress": "进展情况",
}
REQUIRED_FIELDS = ("measures", "leader", "dept", "deadline", "standard")


def _norm_header(text: str) -> str:
    return re.sub(r"[\s*_`：:（(][^）)]*[）)]|[\s*_`：:]", "", str(text or ""))


def map_columns(header: List[str]) -> Dict[str, int]:
    """把表头映射到字段，先精确匹配，再包含匹配；同一列只用一次。"""
    names = [_norm_header(h) for h in header]
    mapping: Dict[str, int] = {}
    used: set = set()
    for fname, keywords in FIELD_KEYWORDS:
        found = None
        for kw in keywords:
            for idx, name in enumerate(names):
                if idx not in used and name == kw:
                    found = idx
                    break
            if found is not None:
                break
        if found is None:
            for kw in keywords:
                for idx, name in enumerate(names):
                    # "问题来源""问题简称"不是问题表现列
                    if fname == "problem" and re.search(r"来源|简称|类别|分类", name):
                        continue
                    if idx not in used and kw in name:
                        found = idx
                        break
                if found is not None:
                    break
        if found is not None:
            mapping[fname] = found
            used.add(found)
    return mapping


def normalize_id(raw: str) -> str:
    """编号归一：全角转半角，去空白和"问题""第""项"等字样，数字段去掉前导 0，字母大写。"""
    text = str(raw or "").translate(FULLWIDTH).strip()
    if re.fullmatch(r"\d+\.0", text):  # Excel 中的数值编号
        text = text[:-2]
    text = re.sub(r"\s+", "", text)
    text = re.sub(r"^(问题|第)", "", text)
    text = re.sub(r"(项|号|条)$", "", text)
    text = text.replace("．", ".").replace("_", "-")
    parts = re.split(r"([-.])", text)
    out = []
    for part in parts:
        if part.isdigit():
            out.append(str(int(part)))
        else:
            out.append(part.upper())
    return "".join(out)


def parent_key(key: str) -> Optional[str]:
    """子项的大项编号：3-1→3，XC-3-1→XC-3。分隔符前不是数字的（XC-3、巡-1）是来源前缀，不算子项。"""
    match = re.match(r"^(.*\d)[-.]\d+$", key)
    return match.group(1) if match else None


# ---------------------------------------------------------------------------
# 读取文件
# ---------------------------------------------------------------------------

def read_text(path: Path) -> str:
    data = path.read_bytes()
    for enc in ("utf-8-sig", "gb18030"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    raise InputError(f"无法按 UTF-8 或 GB18030 读取文件：{path}")


_PIPE_PLACEHOLDER = "\ue001"  # 私用区字符，暂存单元格内转义的竖线


def _split_md_row(line: str) -> List[str]:
    text = line.strip().replace("\\|", _PIPE_PLACEHOLDER)
    if text.startswith("|"):
        text = text[1:]
    if text.endswith("|"):
        text = text[:-1]
    return [c.strip().replace(_PIPE_PLACEHOLDER, "|") for c in text.split("|")]


_DELIM_RE = re.compile(r"^\s*\|?\s*:?-+:?\s*(\|\s*:?-+:?\s*)*\|?\s*$")


def parse_md_tables(text: str) -> List[Tuple[int, List[str], List[Tuple[int, List[str]]]]]:
    """返回 [(表头行号, 表头, [(行号, 单元格)])]，跳过代码块。"""
    lines = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    tables = []
    in_fence = False
    i = 0
    while i < len(lines):
        line = lines[i]
        if re.match(r"^\s{0,3}(```|~~~)", line):
            in_fence = not in_fence
            i += 1
            continue
        if not in_fence and "|" in line and i + 1 < len(lines) and _DELIM_RE.match(lines[i + 1]):
            header = _split_md_row(line)
            if len(_split_md_row(lines[i + 1])) == len(header):
                rows = []
                j = i + 2
                while j < len(lines) and lines[j].strip() and "|" in lines[j]:
                    cells = _split_md_row(lines[j])
                    cells = (cells + [""] * len(header))[: len(header)]
                    cells = [re.sub(r"<br\s*/?>", "\n", c, flags=re.I) for c in cells]
                    rows.append((j + 1, cells))
                    j += 1
                tables.append((i + 1, header, rows))
                i = j
                continue
        i += 1
    return tables


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (_dt.datetime, _dt.date)):
        return value.strftime("%Y-%m-%d")
    return str(value).strip()


def read_xlsx_tables(path: Path):
    try:
        import openpyxl
    except ImportError as exc:
        raise InputError("读取 .xlsx 台账需要 openpyxl，请先运行：pip install openpyxl") from exc
    try:
        wb = openpyxl.load_workbook(str(path), data_only=True)
    except Exception as exc:  # 文件损坏或不是 xlsx
        raise InputError(f"无法读取 Excel 文件：{path}（{exc}）") from exc
    tables = []
    for ws in wb.worksheets:
        rows = [[_cell_text(v) for v in r] for r in ws.iter_rows(values_only=True)]
        for h_index in range(min(10, len(rows))):
            mapping = map_columns(rows[h_index])
            if "id" in mapping and "status" in mapping:
                body = []
                for r_index in range(h_index + 1, len(rows)):
                    cells = rows[r_index]
                    if not any(c for c in cells):
                        if body:
                            break
                        continue
                    body.append((r_index + 1, cells + [""] * (len(rows[h_index]) - len(cells))))
                tables.append((f"工作表“{ws.title}”", h_index + 1, rows[h_index], body))
                break
    return tables


def load_ledger(path: Path) -> Ledger:
    if not path.is_file():
        raise InputError(f"未找到台账文件：{path}")
    suffix = path.suffix.lower()
    candidates = []
    if suffix in (".md", ".txt"):
        for index, (line, header, rows) in enumerate(parse_md_tables(read_text(path)), start=1):
            candidates.append((f"第 {index} 张表（第 {line} 行）", line, header, rows))
    elif suffix in (".xlsx", ".xlsm"):
        candidates = read_xlsx_tables(path)
    else:
        raise InputError(f"台账只支持 .md、.txt、.xlsx：{path}")
    chosen = None
    for where, line, header, rows in candidates:
        mapping = map_columns(header)
        if "id" in mapping and "status" in mapping:
            chosen = (where, header, rows, mapping)
            break
    if chosen is None:
        if candidates:
            raise InputError("台账表格缺少“编号”列或“整改状态”列，无法对账。请在表头写明这两列。")
        raise InputError("台账中没有找到表格（Markdown 表格需要表头行和 | --- | 分隔行）。")
    where, header, rows, mapping = chosen
    columns = {FIELD_NAMES[k]: str(header[v]).strip() for k, v in mapping.items()}

    items: "OrderedDict[str, Item]" = OrderedDict()
    issues: List[dict] = []
    current: Optional[Item] = None
    for row_no, cells in rows:
        values = {k: str(cells[v]).strip() if v < len(cells) else "" for k, v in mapping.items()}
        raw_id = values.get("id", "")
        if not raw_id:
            if current is None or not any(values.values()):
                continue
            # 续行（合并单元格）：状态并入上一条
            if values.get("status"):
                current.status_raw.append(values["status"])
            for k, v in values.items():
                if v and k not in ("id", "status"):
                    current.fields[k] = (current.fields.get(k, "") + "\n" + v).strip()
            continue
        key = normalize_id(raw_id)
        if key in items:
            issues.append(_issue("RC-001", "P1", "台账", "编号重复", key, row_no, raw_id,
                                 f"台账编号“{raw_id}”重复出现",
                                 "每条问题或子项使用唯一编号，报告才能逐条对账"))
            continue
        current = Item(raw_id=raw_id, key=key, row=row_no,
                       fields={k: v for k, v in values.items() if k not in ("id", "status")},
                       status_raw=[values.get("status", "")])
        items[key] = current
    if not items:
        raise InputError("台账表格没有数据行。")
    return Ledger(path=path, items=items, columns=columns, where=where, issues=issues)


def load_report(path: Path) -> List[Tuple[int, str]]:
    """返回【起草说明】之前的正文行 [(行号, 文字)]。"""
    if not path.is_file():
        raise InputError(f"未找到报告文件：{path}")
    suffix = path.suffix.lower()
    if suffix in (".txt", ".md"):
        raw_lines = read_text(path).replace("\r\n", "\n").replace("\r", "\n").split("\n")
    elif suffix == ".docx":
        raw_lines = _read_docx_lines(path)
    else:
        raise InputError(f"报告只支持 .txt、.md、.docx：{path}")
    out = []
    for index, line in enumerate(raw_lines, start=1):
        if DRAFT_NOTE_RE.match(line):
            break
        out.append((index, line))
    return out


def _read_docx_lines(path: Path) -> List[str]:
    try:
        import docx  # type: ignore
    except ImportError as exc:
        raise InputError("读取 .docx 需要 python-docx，请先运行：pip install python-docx") from exc
    try:
        document = docx.Document(str(path))
    except Exception as exc:
        raise InputError(f"无法读取 Word 文件：{path}（{exc}）") from exc
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append(" ".join(cell.text for cell in row.cells))
    return lines


# ---------------------------------------------------------------------------
# 台账自检
# ---------------------------------------------------------------------------

def _issue(code, level, where, category, item, line, excerpt, message, suggestion) -> dict:
    return OrderedDict([
        ("id", code), ("level", level), ("where", where), ("category", category),
        ("item", item), ("line", line), ("excerpt", excerpt),
        ("message", message), ("suggestion", suggestion),
    ])


DATE_RES = [
    (re.compile(r"(\d{4})\s*[年\-/.]\s*(\d{1,2})\s*[月\-/.]\s*(\d{1,2})"), "ymd"),
    (re.compile(r"(\d{4})\s*年\s*(\d{1,2})\s*月\s*(底|末)"), "ym_end"),
    (re.compile(r"(\d{4})\s*年\s*年?\s*(底|末)"), "y_end"),  # 含"2026年年底前"
    (re.compile(r"(\d{4})\s*年\s*(\d{1,2})\s*月"), "ym_end"),
]


def parse_deadline(text: str) -> Optional[_dt.date]:
    """取时限中的最后一个日期；"长期"等无日期的返回 None。"""
    text = (text or "").translate(FULLWIDTH)
    best: Optional[_dt.date] = None
    for regex, kind in DATE_RES:
        for m in regex.finditer(text):
            try:
                year = int(m.group(1))
                if kind == "ymd":
                    d = _dt.date(year, int(m.group(2)), int(m.group(3)))
                elif kind == "y_end":
                    d = _dt.date(year, 12, 31)
                else:
                    month = int(m.group(2))
                    nxt = _dt.date(year + (month == 12), month % 12 + 1, 1)
                    d = nxt - _dt.timedelta(days=1)
            except ValueError:
                continue
            if best is None or d > best:
                best = d
        if best is not None and kind == "ymd":
            break
    return best


_MONTH_DAY_RE = re.compile(r"\d{1,2}\s*月\s*(?:\d{1,2}\s*日|底|末)")


def has_concrete_date(text: str) -> bool:
    """时限里有没有具体日期：带年份的日期，或"11月30日""11月底"这类当年日期。"长期""尽快""年内"没有。"""
    text = (text or "").translate(FULLWIDTH)
    return parse_deadline(text) is not None or bool(_MONTH_DAY_RE.search(text))


def check_ledger(ledger: Ledger, as_of: Optional[_dt.date]) -> List[dict]:
    issues = list(ledger.issues)
    has_evidence_col = "佐证材料" in ledger.columns
    missing_evidence_col_noted = False
    for item in ledger.items.values():
        statuses = []
        for raw in item.status_raw:
            canon, note = normalize_status(raw)
            statuses.append(canon)
            if note == "empty":
                issues.append(_issue("RC-002", "P1", "台账", "状态为空", item.key, item.row, item.raw_id,
                                     f"编号 {item.raw_id} 的整改状态为空",
                                     "按状态口径填写：已完成、立行立改、持续推进、长期坚持；尚未动手的填“未启动”"))
            elif note == "marker":
                issues.append(_issue("RC-002", "P2", "台账", "状态待补", item.key, item.row, raw,
                                     f"编号 {item.raw_id} 的整改状态尚待补充，本条不参与对账",
                                     "补齐进展和佐证后按状态口径填写"))
            elif note.startswith("nonstandard:"):
                word = note.split(":", 1)[1]
                issues.append(_issue("RC-003", "P1", "台账", "状态口径", item.key, item.row, raw,
                                     NONSTANDARD[word], "改为四种口径之一，本条暂不参与对账"))
            elif note == "unknown":
                issues.append(_issue("RC-003", "P1", "台账", "状态口径", item.key, item.row, raw,
                                     f"编号 {item.raw_id} 的状态“{raw}”不在口径内",
                                     "状态栏只写已完成、立行立改、持续推进、长期坚持（或临时的未启动），进展写在进展情况栏"))
            elif note.startswith("synonym:"):
                issues.append(_issue("RC-003", "P2", "台账", "状态口径", item.key, item.row, raw,
                                     f"“{raw}”按“{canon}”比对",
                                     f"统一写成“{canon}”"))
        item.status = aggregate_status(statuses)

        if item.status in DONE_SET:
            if has_evidence_col:
                evidence = item.fields.get("evidence", "")
                if not evidence or MARKER_RE.search(evidence):
                    issues.append(_issue("RC-004", "P1", "台账", "佐证材料", item.key, item.row, item.raw_id,
                                         f"编号 {item.raw_id} 记为“{item.status}”，但佐证材料为空或待补",
                                         "补上佐证材料（文件名、文号、记录名称）；拿不出佐证的，状态改为“持续推进”"))
            elif not missing_evidence_col_noted:
                missing_evidence_col_noted = True
                issues.append(_issue("RC-004", "P2", "台账", "佐证材料", "", 0, "",
                                     "台账没有“佐证材料”列，无法核对已完成事项是否有佐证",
                                     "增加“佐证材料”列，已完成、立行立改的事项逐条填写"))

        empty = [FIELD_NAMES[f] for f in REQUIRED_FIELDS
                 if FIELD_NAMES[f] in ledger.columns and not item.fields.get(f, "").strip()]
        if empty:
            issues.append(_issue("RC-005", "P1", "台账", "关键栏为空", item.key, item.row, item.raw_id,
                                 f"编号 {item.raw_id} 缺：{'、'.join(empty)}",
                                 "按用户提供的情况填写；用户没有提供的写【待补：说明】，不留空"))

        deadline_text = item.fields.get("deadline", "").strip()
        if ("完成时限" in ledger.columns and deadline_text and not MARKER_RE.search(deadline_text)
                and item.status not in DONE_SET and not has_concrete_date(deadline_text)):
            issues.append(_issue("RC-007", "P1", "台账", "时限不具体", item.key, item.row, deadline_text[:30],
                                 f"编号 {item.raw_id} 的完成时限“{deadline_text}”不是具体日期",
                                 "写具体日期，如“2026年12月31日前”；需要长期坚持的写“长期坚持，阶段目标于××前完成”；"
                                 "用户没给的写【待补：完成时限】"))

        if as_of and item.status in (ONGOING, NOT_STARTED):
            deadline = parse_deadline(item.fields.get("deadline", ""))
            if deadline and deadline < as_of:
                issues.append(_issue("RC-006", "P1", "台账", "超期未完成", item.key, item.row,
                                     item.fields.get("deadline", ""),
                                     f"编号 {item.raw_id} 完成时限为 {deadline.isoformat()}，截至 {as_of.isoformat()} 仍为“{item.status}”",
                                     "报告中如实写明超期及原因，提出调整后的完成时间；不得改写为已完成"))
    return issues


def ledger_markers(ledger: Ledger) -> Dict[str, int]:
    counts = {"待补": 0, "待核实": 0}
    for item in ledger.items.values():
        for text in list(item.fields.values()) + item.status_raw:
            for m in MARKER_RE.finditer(text or ""):
                counts[m.group(1)] += 1
    return counts


# ---------------------------------------------------------------------------
# 统计
# ---------------------------------------------------------------------------

def count_statuses(statuses: List[Optional[str]]) -> Dict[str, object]:
    counts = OrderedDict((s, 0) for s in STATUSES)
    unknown = 0
    for s in statuses:
        if s:
            counts[s] += 1
        else:
            unknown += 1
    total = len(statuses)
    done = counts[DONE] + counts[IMMEDIATE]
    return OrderedDict([
        ("total", total), ("status_count", counts), ("unknown", unknown),
        ("done", done), ("done_rate", round(done / total * 100, 1) if total else 0.0),
    ])


def group_view(ledger: Ledger) -> Optional[Dict[str, List[str]]]:
    """台账有子项（3-1、3-2）时，按大项分组；没有子项返回 None。"""
    groups: "OrderedDict[str, List[str]]" = OrderedDict()
    has_sub = False
    for key in ledger.items:
        parent = parent_key(key)
        if parent and parent not in ledger.items:
            has_sub = True
            groups.setdefault(parent, []).append(key)
        else:
            groups.setdefault(key, []).append(key)
    return groups if has_sub else None


def ledger_summary(ledger: Ledger) -> Dict[str, object]:
    summary = count_statuses([i.status for i in ledger.items.values()])
    groups = group_view(ledger)
    if groups:
        summary["groups"] = count_statuses(
            [aggregate_status([ledger.items[k].status for k in keys]) for keys in groups.values()]
        )
    return summary


# ---------------------------------------------------------------------------
# 报告对账
# ---------------------------------------------------------------------------

# 编号前缀：字母（XC-3、SJ02）或一两个汉字加连接号（巡-1、审-1、自-1）
ID_PREFIX = r"(?:[A-Za-z]{0,6}-?|[一-鿿]{1,2}-)"
ID_TOKEN = r"(" + ID_PREFIX + r"\d+(?:[-.]\d+)*)"
NUM_GUARD = r"(?!\d)(?!\s*(?:个|项|条|件|处|名|人|万|亿|元|%|次|家|户))"
_LIST_ALT = "以及|、|和|及|与"
_RANGE_ALT = "至|到|~|～"
# 每项为 (类别, 正则)：plain 是"问题N"，item 是"第N项""台账第N项"（列举中的最后一个编号后面跟"项"）
MENTION_RES = [
    ("plain", re.compile(r"问题\s*(?:编号\s*)?[:：]?\s*" + ID_TOKEN + NUM_GUARD)),
    ("item", re.compile(r"台账\s*(?:第|编号|序号)\s*" + ID_TOKEN + r"\s*(?:项|号|条)?" + r"(?!\d)")),
    ("item", re.compile(r"第\s*(\d+(?:[-.]\d+)*)\s*(?:号|项)\s*问题")),
    # "第1、4、5项问题""第2至4项问题"：先认出第一个编号，其余由列举续接
    ("item", re.compile(r"第\s*(\d{1,3}(?:[-.]\d+)*)(?=(?:\s*(?:" + _LIST_ALT + "|" + _RANGE_ALT
                        + r")\s*\d{1,3}(?:[-.]\d+)*)+\s*(?:号|项)\s*问题)")),
]
# 列举续接："问题1、4、5""问题1和2""问题2至4""问题1至问题3"。编号最多三位，后面跟量词、日期的不算编号；
# item 类的列举以"项"结尾，所以不排除"项""条""号"
_CONT_UNITS = "个|件|处|名|人|万|亿|元|%|次|家|户|月|日|年|季度|周|天|笔|宗|批|台|岁|倍|分"
CONT_RES = {
    kind: re.compile(
        r"(?:(?P<range>\s*(?:" + _RANGE_ALT + r")\s*(?:问题\s*)?)|\s*(?:" + _LIST_ALT + r")\s*)"
        r"(?P<id>" + ID_PREFIX + r"\d{1,3}(?:[-.]\d{1,3})*)(?!\d)(?!\s*(?:" + units + "))"
    )
    for kind, units in (("plain", "项|条|号|" + _CONT_UNITS), ("item", _CONT_UNITS))
}
SEPARATOR_RE = re.compile(r"^[\s、，,和及与以至到~～“”\"（）()]*$")
HEADING_RE = re.compile(r"^\s*(?:[一二三四五六七八九十]+、|（[一二三四五六七八九十]+）)")

FUTURE_RE = re.compile(r"(将|拟|计划|预计|力争|确保|争取|预期|目标|要|应当|须|需在|必须|(?<![目此之])前|以后|之后|待|后续"
                       r"|届时|即可|方可|才能)")
NEGATION_RE = re.compile(r"(未|尚未|没有|不能|无法|未能|难以|不到)$")
# 后面接这些字的是名词性用法（"已完成整改方案制定""整改到位的标准""已整改到位的有2个社区"），不是状态说法；
# 后面紧跟数量的（"已整改完成12处"）是部分完成
_DONE_TAIL = (r"(?!率|情况|时限|时间|标准|后|工作|任务清单|方案|措施|台账|计划|报告|清单|要求|的)"
              r"(?!\s*(?:\d|[一二两三四五六七八九十百]+\s*(?:个|项|条|处|家|户|所|名|件|宗|台|人|批)))")
DONE_CLAIM_RE = re.compile(
    r"(已完成整改|整改已完成|已整改完成|完成了整改|完成整改|整改完成|已整改到位|整改已到位|整改到位|整改完毕|"
    r"已全部整改|全部整改到位|已对账销号|已销号|已办结|已彻底解决|问题已解决|已立行立改|立行立改完成|已全面完成整改)"
    + _DONE_TAIL
)
# 不带"整改"二字的完成说法（"该问题已完成。""问题6：已整改。""整改措施已落实到位。"）：
# 只在主语是本条问题本身时算，"盘点已完成，系统正在部署"不算
BARE_DONE_RE = re.compile(r"已经?(?:完成|整改|落实到位|解决)(?=\s*(?:[。；;！!，,]|$))")
_ID_BODY = ID_PREFIX + r"\d+(?:[-.]\d+)*"
_ID_LIST = _ID_BODY + r"(?:(?:" + _LIST_ALT + "|" + _RANGE_ALT + r")(?:问题)?" + _ID_BODY + r")*"
_BARE_SUBJECT_RE = re.compile(
    r"^(?:[一二三四五六七八九十]+、|（[一二三四五六七八九十]+）|\d+[.．、])?\s*"
    r"(?:(?:该|本|此|这个|这一|上述|相关)?(?:问题|项|条|事项)?|问题" + _ID_LIST
    + r"|第" + _ID_LIST + r"(?:号|项)问题"
    r"|(?:[A-Za-z]{1,6}-?|[一-鿿]{1,2}-)\d+(?:[-.]\d+)*)"
    r"(?:的)?(?:各项)?(?:整改)?(?:工作|任务|措施)?(?:均|都|也|目前|现|全部)?$"
)
ONGOING_RE = re.compile(r"(正在推进|持续推进|正在整改|整改中|推进中|正在实施|尚未完成|未完成|部分完成|仍在|正在开展)")
LONG_RE = re.compile(r"(需长期坚持|长期坚持|长期推进|常态化推进)")
NOT_STARTED_RE = re.compile(r"(尚未启动|未启动|尚未开展|尚未开始)")
# 段末明确写了未完成的状态句（"该问题正在推进"）：同段的"已完成……"多半说的是其中一部分
STATUS_SENTENCE_RE = re.compile(
    r"(?:该|本|此)(?:问题|项|条)\s*(?:目前|当前|仍|仍在|还在|整改|整改工作)?\s*"
    r"(?:正在推进|持续推进|正在整改|整改中|推进中|尚未完成|尚未启动|未启动|需长期坚持|取得阶段性成效)"
)
# "督促3个科室完成整改""部分科室已完成整改"：说的是下属单位或其中一部分，不是本条问题整体完成
PARTIAL_RE = re.compile(
    r"(?:\d+|[一二两三四五六七八九十]+)\s*(?:个|家|户|所|名|处|件|宗|台|批|项|条)[^，。；]{0,8}$"
    r"|(?:部分|其中|个别|少数|有的)[^，。；]{0,8}$"
)
# "第1、4、5项问题""第3项问题"是编号，不是"3个科室"那样的部分数量，也不是统计数
ORDINAL_TAIL_RE = re.compile(r"第[\d\s、和及与以至到~～.\-]*\d\s*(?:号|项)(?:\s*问题)?[^，,。；;]{0,8}$")
ORDINAL_LEAD_RE = re.compile(r"第[\d\s、和及与以至到~～.\-]*$")
DECLARE_RE = re.compile(r"(?:整改状态|状态)\s*[:：为]\s*([^\s，。；,;）)]{2,6})")
# 方案目标句"到2026年12月31日，反馈问题全部整改到位"：句首的"到……日，"表示计划
PLAN_LEAD_RE = re.compile(
    r"^\s*(?:[一二三四五六七八九十]+、|（[一二三四五六七八九十]+）|\d+[.．、])?\s*"
    r"(?:力争|确保|争取|计划|预计|要)?\s*(?:到|至)\s*\d{4}\s*年[^，,。；;]{0,10}?(?:日|底|末|前|月)"
    r"(?:前|以前|之前)?\s*[，,]"
)


def _clause_prefix(text: str, pos: int, width: int = 14, stops: str = "。；;，,\n") -> str:
    """取 pos 之前同一分句内的文字（遇句号、分号、逗号截断）。"""
    start = max(0, pos - width)
    prefix = text[start:pos]
    cut = max(prefix.rfind(ch) for ch in stops)
    return prefix[cut + 1:] if cut >= 0 else prefix


def _sentence_prefix(text: str, pos: int, width: int = 60) -> str:
    """取 pos 之前同一句内的文字（遇句号、分号、换行截断，逗号不截断）。"""
    return _clause_prefix(text, pos, width, "。；;\n")


def is_future(prefix: str, word: str = "") -> bool:
    """"将""计划""确保""……前"等表示计划；分句里有"已"的仍按已发生处理（如"已于10月20日前完成整改"）。"""
    if word.startswith("已") or "已" in prefix:
        return False
    return bool(FUTURE_RE.search(prefix))


def is_planned(text: str, pos: int, word: str = "") -> bool:
    """在 is_future 之外，再看整句：句首"到2026年12月31日，"引出的是目标，句中有"已"的除外。"""
    if is_future(_clause_prefix(text, pos), word):
        return True
    sentence = _sentence_prefix(text, pos)
    if word.startswith("已") or "已" in sentence:
        return False
    return bool(PLAN_LEAD_RE.match(sentence))


def _range_keys(first: str, last: str) -> List[str]:
    """"问题2至4"展开为 3、4（起点 2 已单独记下）；前缀不同或跨度过大的只记终点。"""
    a = re.match(r"^(.*?)(\d+)$", first)
    b = re.match(r"^(.*?)(\d+)$", last)
    if a and b and a.group(1) == b.group(1):
        lo, hi = int(a.group(2)), int(b.group(2))
        if lo < hi <= lo + 50:
            return [a.group(1) + str(n) for n in range(lo + 1, hi + 1)]
    return [last]


def _id_pattern(key: str) -> re.Pattern:
    """台账中带前缀的编号原样匹配：XC-1 也认 XC-01、xc－1，巡-1 也认 巡－01。"""
    parts = []
    for part in re.split(r"([-.])", key):
        if part.isdigit():
            parts.append("0*" + part)
        elif part in ("-", "."):
            parts.append("[-－.．]")
        else:
            parts.append(re.escape(part))
    return re.compile(r"(?<![A-Za-z0-9])" + "".join(parts) + r"(?![0-9])", re.IGNORECASE)


def literal_ids(ledger: Ledger) -> List[Tuple[str, re.Pattern]]:
    """不是纯数字的台账编号（XC-03、巡-1、自-2），逐个生成字面匹配模式；长编号在前，XC-1-1 不会被认成 XC-1。"""
    keys = [k for k in ledger.items if not re.fullmatch(r"\d+(?:[-.]\d+)*", k)]
    return [(k, _id_pattern(k)) for k in sorted(keys, key=len, reverse=True)]


def find_mentions(text: str, ledger: Ledger, alpha_ids: List[Tuple[str, re.Pattern]]) -> List[Tuple[int, int, List[str]]]:
    """返回 [(起点, 终点, [归一编号])]，按位置排序、去重。

    "问题1、4、5""问题2至4""第1、4项问题"这类列举，后面的编号各占一处；范围中间的编号并到终点那一处。
    """
    found: Dict[int, list] = {}
    primaries = []
    for kind, regex in MENTION_RES:
        for m in regex.finditer(text):
            if m.start() not in found:
                found[m.start()] = [m.start(), m.end(), [normalize_id(m.group(1))]]
                primaries.append((kind, found[m.start()]))
    for kind, entry in sorted(primaries, key=lambda p: p[1][0]):
        pos, prev = entry[1], entry[2][-1]
        while True:
            c = CONT_RES[kind].match(text, pos)
            if not c:
                break
            key = normalize_id(c.group("id"))
            keys = _range_keys(prev, key) if c.group("range") else [key]
            start, end = c.start("id"), c.end("id")
            hit = next((f for f in found.values() if f[0] < end and start < f[1]), None)
            if hit is not None:
                # "问题1至问题3"的"问题3"已单独认出：把中间的编号并进去，后面的列举由它自己续接
                hit[2] = [k for k in keys[:-1] if k not in hit[2]] + hit[2]
                break
            found[start] = [start, end, keys]
            pos, prev = c.end(), key
    for key, regex in alpha_ids:
        for m in regex.finditer(text):
            if not any(f[0] <= m.start() < f[1] for f in found.values()):
                found.setdefault(m.start(), [m.start(), m.end(), [key]])
    return [(found[k][0], found[k][1], found[k][2]) for k in sorted(found)]


def segment_report(lines: List[Tuple[int, str]], ledger: Ledger, heads: Optional[Dict[str, Tuple[int, str]]] = None):
    """按编号把报告切成段：返回 {编号: [(行号, 文字)]} 和全部提及的编号。

    heads 不为 None 时，记下每个编号第一次出现后紧跟的问题表述（到句号为止），供弱化检查用。
    """
    alpha_ids = literal_ids(ledger)
    blocks: "OrderedDict[str, List[Tuple[int, str]]]" = OrderedDict()
    mentioned: "OrderedDict[str, int]" = OrderedDict()
    current: List[str] = []
    for line_no, raw in lines:
        text = raw.translate(FULLWIDTH)
        if not text.strip():
            continue
        mentions = find_mentions(text, ledger, alpha_ids)
        if not mentions:
            if HEADING_RE.match(text):  # 不带编号的一、（一）级标题：上一条到此为止
                current = []
                continue
            for key in current:
                blocks.setdefault(key, []).append((line_no, text))
            continue
        groups: List[list] = []  # [编号组, 段起点, 最后一个编号的终点]
        last_end = 0
        for start, end, keys in mentions:
            for key in keys:
                mentioned.setdefault(key, line_no)
            # "问题3、问题5需长期坚持""问题1、4、5已完成整改"：只隔着顿号、"和"等的编号共用后面的文字
            if groups and SEPARATOR_RE.match(text[last_end:start]):
                groups[-1][0].extend(k for k in keys if k not in groups[-1][0])
                groups[-1][2] = end
            else:
                groups.append([list(keys), start, end])
            last_end = end
        if heads is not None:
            for keys, _, end in groups:
                head = re.split(r"[。；;\n]", text[end:].lstrip("：:，,、 （("), maxsplit=1)[0]
                for key in keys:
                    heads.setdefault(key, (line_no, head.rstrip("）)").strip()))
        for index, (keys, start, _) in enumerate(groups):
            seg_start = 0 if index == 0 else start
            seg_end = groups[index + 1][1] if index + 1 < len(groups) else len(text)
            segment = text[seg_start:seg_end]
            for key in keys:
                blocks.setdefault(key, []).append((line_no, segment))
        current = groups[-1][0]
    return blocks, mentioned


def detect_claims(text: str) -> Dict[str, List[str]]:
    """done 为全部完成说法；done_weak 是其中分句里没有"已"的（如"……，完成整改。"），整改方案中按 P2 提示。"""
    claims: Dict[str, List[str]] = {"done": [], "done_weak": [], "ongoing": [], "long": [], "not_started": [],
                                    "declared": []}
    for m in DECLARE_RE.finditer(text):
        canon, _ = normalize_status(m.group(1))
        if canon:
            claims["declared"].append(canon)
    for m in DONE_CLAIM_RE.finditer(text):
        prefix = _clause_prefix(text, m.start())
        if NEGATION_RE.search(text[max(0, m.start() - 3):m.start()]):
            claims["ongoing"].append(m.group(0))
            continue
        if is_planned(text, m.start(), m.group(0)) or (PARTIAL_RE.search(prefix)
                                                        and not ORDINAL_TAIL_RE.search(prefix)):
            continue
        claims["done"].append(m.group(0))
        if "已" not in m.group(0) and "已" not in prefix:
            claims["done_weak"].append(m.group(0))
    for m in BARE_DONE_RE.finditer(text):
        subject = re.sub(r"\s", "", _clause_prefix(text, m.start(), 20, "。；;，,：:\n"))
        if _BARE_SUBJECT_RE.match(subject):
            claims["done"].append(m.group(0))
    claims["ongoing"] += [m.group(0) for m in ONGOING_RE.finditer(text)]
    claims["long"] += [m.group(0) for m in LONG_RE.finditer(text)]
    claims["not_started"] += [m.group(0) for m in NOT_STARTED_RE.finditer(text)]
    return claims


def _excerpt(text: str, word: str, radius: int = 18) -> str:
    """取 word 前后各 radius 字，不跨行。"""
    pos = text.find(word) if word else -1
    if pos < 0:
        return text.split("\n", 1)[0][: radius * 2].strip()
    line_start = text.rfind("\n", 0, pos) + 1
    line_end = text.find("\n", pos)
    line_end = len(text) if line_end < 0 else line_end
    start = max(line_start, pos - radius)
    return text[start: min(line_end, pos + len(word) + radius)].strip()


def resolve_key(key: str, ledger: Ledger) -> Tuple[Optional[str], List[str]]:
    """把报告中的编号对到台账：返回 (显示用编号, 对应的台账编号列表)。"""
    if key in ledger.items:
        return key, [key]
    subs = [k for k in ledger.items if k.startswith(key + "-") or k.startswith(key + ".")]
    if subs:
        return key, subs
    return None, []


_ITEM_PREFIX_RE = re.compile(
    r"^\s*(?:[一二三四五六七八九十]+、|（[一二三四五六七八九十]+）|\d+[.．、])?\s*"
    r"(?:问题\s*(?:编号\s*)?" + _ID_LIST + r"|第" + _ID_LIST + r"(?:号|项)问题"
    r"|台账\s*(?:第|编号|序号)\s*\S+?(?:项|号|条))?"
)


def strip_problem(text: str, problem_norm: str) -> str:
    """去掉照录的问题表述（如"仍有1个社区未整改到位"），只在其余文字中找状态说法。"""
    if not problem_norm:
        return text
    kept = []
    for part in re.split(r"(?<=[。；;])", text):
        core = _PUNCT_RE.sub("", _ITEM_PREFIX_RE.sub("", part, count=1))
        if len(core) >= 4 and core in problem_norm:
            continue
        kept.append(part)
    return "".join(kept)


def compare_items(blocks, mentioned, ledger: Ledger, plan: bool = False) -> List[dict]:
    """逐条比对。plan 为 True（整改方案）时不要求每条写明状态。"""
    issues: List[dict] = []
    covered: set = set()
    for key, first_line in mentioned.items():
        shown, keys = resolve_key(key, ledger)
        if not keys:
            issues.append(_issue("RC-106", "P1", "报告", "编号不在台账", key, first_line, f"问题{key}",
                                 f"报告提到的编号 {key} 在台账中不存在",
                                 "核对编号；报告逐条整改情况的编号与台账保持一致"))
            continue
        covered.update(keys)
        status = aggregate_status([ledger.items[k].status for k in keys])
        if status is None:
            continue  # 台账状态待定，已在台账自检中报出
        segments = blocks.get(key, [])
        problem_norm = _PUNCT_RE.sub("", "".join(ledger.items[k].fields.get("problem", "") for k in keys))
        text = "".join(strip_problem(seg, problem_norm) for _, seg in segments)
        line = segments[0][0] if segments else first_line
        claims = detect_claims(text)
        declared = claims["declared"][0] if claims["declared"] else None
        label = f"问题{shown}"
        done_words = list(claims["done"])
        if plan:
            # 方案里没有"已"的"……，完成整改。"多是计划，只作 P2 提示
            for weak in claims["done_weak"]:
                done_words.remove(weak)
        if status not in DONE_SET:
            claimed_done = done_words or (declared in DONE_SET if declared else False)
            if not claimed_done and claims["done_weak"]:
                word = claims["done_weak"][0]
                issues.append(_issue("RC-101", "P2", "报告", "计划表述", key, line, _excerpt(text, word),
                                     f"{label}：方案中的“{word}”前面没有“于……前”“将”等计划字样，容易读成已经完成"
                                     f"（台账为“{status}”）",
                                     "写成计划句，如“于2026年12月31日前完成整改”"))
                continue
            if claimed_done:
                word = done_words[0] if done_words else declared
                explicit = STATUS_SENTENCE_RE.search(text)
                if explicit and declared not in DONE_SET:
                    # 同段明确写了"该问题正在推进"：前面的"已完成……"多半说的是其中一项措施
                    issues.append(_issue("RC-101", "P1", "报告", "状态不一致", key, line, _excerpt(text, word),
                                         f"{label}：段中有“{word}”，段末又写“{explicit.group(0)}”（台账为“{status}”），"
                                         "读者可能误以为整条已完成",
                                         "写明完成的是哪一项措施，如“盘点已于10月20日完成”；不单独使用“已完成整改”"))
                    continue
                issues.append(_issue("RC-101", "P0", "报告", "状态不一致", key, line, _excerpt(text, word),
                                     f"{label}：报告写成“{word}”，台账状态为“{status}”",
                                     "按台账如实写：已完成哪些、还差什么、预计何时完成；台账状态只能凭新的进展和佐证更新"))
                continue
            if status == ONGOING and (claims["long"] or declared == LONGTERM) and not claims["ongoing"]:
                word = claims["long"][0] if claims["long"] else LONGTERM
                issues.append(_issue("RC-103", "P1", "报告", "状态不一致", key, line, _excerpt(text, word),
                                     f"{label}：报告写成“{word}”，台账为“持续推进”",
                                     "有明确完成标准的问题不能以“长期坚持”代替“持续推进”，写明预计完成时间"))
                continue
            if status == NOT_STARTED and (claims["ongoing"] or declared == ONGOING) and not claims["not_started"]:
                issues.append(_issue("RC-103", "P1", "报告", "状态不一致", key, line,
                                     _excerpt(text, claims["ongoing"][0] if claims["ongoing"] else ""),
                                     f"{label}：报告写成正在推进，台账为“未启动”",
                                     "如实写明尚未启动的原因和计划启动时间，或补充进展后更新台账"))
                continue
            if status == LONGTERM and (claims["ongoing"] or declared == ONGOING) and not (claims["long"] or declared == LONGTERM):
                issues.append(_issue("RC-103", "P2", "报告", "状态不一致", key, line,
                                     _excerpt(text, claims["ongoing"][0] if claims["ongoing"] else ""),
                                     f"{label}：报告写成正在推进，台账为“长期坚持”",
                                     "两种写法都表示未完成，按台账口径统一"))
                continue
            if not plan and not (claims["ongoing"] or claims["long"] or claims["not_started"] or declared):
                issues.append(_issue("RC-104", "P2", "报告", "未写明状态", key, line, text[:30].strip(),
                                     f"{label}：报告没有写明整改状态（台账为“{status}”）",
                                     "在该条末尾写明状态，如“该问题正在推进，预计12月31日前完成”"))
        elif not plan:
            if declared and declared not in DONE_SET:
                issues.append(_issue("RC-102", "P1", "报告", "状态不一致", key, line, _excerpt(text, "状态"),
                                     f"{label}：报告写明状态为“{declared}”，台账为“{status}”",
                                     "核对后统一；台账为已完成的，报告写“已完成整改”并概述佐证"))
            elif not claims["done"] and not declared:
                words = claims["ongoing"] + claims["not_started"] + claims["long"]
                if words:
                    issues.append(_issue("RC-102", "P1", "报告", "状态不一致", key, line, _excerpt(text, words[0]),
                                         f"{label}：报告写成“{words[0]}”，台账为“{status}”",
                                         "核对后统一；台账为已完成的，报告写“已完成整改”"))
                else:
                    issues.append(_issue("RC-104", "P2", "报告", "未写明状态", key, line, text[:30].strip(),
                                         f"{label}：报告没有写明整改状态（台账为“{status}”）",
                                         "在该条末尾写明“该问题已完成整改”"))
    # 漏报：台账有、报告没提到
    groups = group_view(ledger) or OrderedDict((k, [k]) for k in ledger.items)
    for parent, keys in groups.items():
        if not any(k in covered for k in keys):
            item = ledger.items[keys[0]]
            status = aggregate_status([ledger.items[k].status for k in keys])
            if status in DONE_SET:
                # 进展报告可以不重复写已完成、本期无变化的条目；情况报告仍要逐条写
                issues.append(_issue("RC-105", "P2", "报告", "漏报", parent, 0, item.raw_id,
                                     f"台账编号 {parent}（{status}）在报告中没有提到",
                                     "进展报告中本期无变化的已完成条目可以不写；整改情况报告要逐条写明"))
            else:
                issues.append(_issue("RC-105", "P1", "报告", "漏报", parent, 0, item.raw_id,
                                     f"台账编号 {parent}（{status or '状态待定'}）在报告中没有提到",
                                     "未完成的问题要逐条写明进展、还差什么和预计完成时间"))
    return issues


# 统计句
CN_DIGITS = {"零": 0, "〇": 0, "一": 1, "二": 2, "两": 2, "三": 3, "四": 4, "五": 5, "六": 6,
             "七": 7, "八": 8, "九": 9}
NUM = r"(\d+|[零〇一二两三四五六七八九十百]+)"
UNIT = r"\s*(?:个|项|条)"
UNIT_END = r"(?=\s*(?:问题|整改事项|整改任务)?\s*(?:[，,。；;、（(）)]|$|其中|，))"


def cn_to_int(text: str) -> Optional[int]:
    if text.isdigit():
        return int(text)
    total, current = 0, 0
    for ch in text:
        if ch in CN_DIGITS:
            current = CN_DIGITS[ch]
        elif ch == "十":
            total += (current or 1) * 10
            current = 0
        elif ch == "百":
            total += (current or 1) * 100
            current = 0
        else:
            return None
    return total + current


# "8个问题已完成整改6个"：动词后面还跟着数的，前面的 8 是总数，不是完成数
NO_NUM_AFTER = r"(?!(?:整改|到位)?\s*(?:共\s*)?[\d零〇一二两三四五六七八九十百])"
STAT_PATTERNS = {
    "total": [
        re.compile(r"(?:反馈|指出|交办|查出|发现|梳理|移交|涉及)(?:的|了)?(?:具体)?问题\s*(?:共\s*)?" + NUM + UNIT),
        re.compile(r"(?:反馈|指出|交办|查出|梳理出|共|涉及)([^。；，,\d]{0,10}?)" + NUM + UNIT + r"\s*(?:具体)?(?:问题|整改事项|整改任务)"),
    ],
    "done": [
        # "已完成整改5个""已完成整改的有5个""共整改完成5个问题""6个问题已整改5个"
        re.compile(r"(?:已完成整改|已整改完成|已完成|完成整改|整改完成|整改完毕|已整改到位|整改到位|已整改|已销号|已办结)"
                   r"(?:的)?(?:问题)?\s*(?:有|共有|共|为|达)?\s*" + NUM + UNIT + UNIT_END),
        # "5个问题已完成整改""6个问题均已完成整改"
        re.compile(NUM + UNIT + r"(?:问题)?\s*(?:均|都|全部)?\s*"
                   r"(?:已完成整改|已整改完成|已完成|已整改到位|已整改完毕|已整改|已销号|已办结)" + NO_NUM_AFTER),
        # "已有5个完成整改""有3个问题整改到位"
        re.compile(r"(?:已有|有)\s*" + NUM + UNIT + r"\s*(?:问题)?\s*(?:已经?)?"
                   r"(?:完成整改|整改完成|整改完毕|整改到位|销号|办结)(?!率)" + NO_NUM_AFTER),
    ],
    "immediate": [
        re.compile(r"(?:已)?立行立改(?:的)?(?:问题)?\s*" + NUM + UNIT + UNIT_END),
        re.compile(NUM + UNIT + r"(?:问题)?\s*(?:已)?立行立改" + NO_NUM_AFTER),
    ],
    "ongoing": [
        re.compile(r"(?:持续推进|正在推进|正在整改|整改中|推进中|尚未完成)(?:的)?(?:问题)?\s*" + NUM + UNIT + UNIT_END),
        re.compile(NUM + UNIT + r"(?:问题)?\s*(?:正在|持续)推进" + NO_NUM_AFTER),
    ],
    "long": [
        re.compile(r"(?:需)?长期坚持(?:的)?(?:问题)?\s*" + NUM + UNIT + UNIT_END),
        re.compile(NUM + UNIT + r"(?:问题)?\s*(?:需)?长期坚持" + NO_NUM_AFTER),
    ],
    "not_started": [
        re.compile(r"(?:尚未启动|未启动)(?:的)?(?:问题)?\s*" + NUM + UNIT + UNIT_END),
        re.compile(NUM + UNIT + r"(?:问题)?\s*(?:尚未|未)启动" + NO_NUM_AFTER),
    ],
}
RATE_RE = re.compile(r"完成率\s*(?:为|达到|达|是)?\s*(\d+(?:\.\d+)?)\s*[%％]")
ALL_DONE_RE = re.compile(
    r"(?:全部|全面)(?:完成整改|整改到位|整改完毕|销号)|整改(?:任务|工作)?(?:已)?全部完成|"
    r"(?:问题|任务)(?:已)?全部(?:完成)?整改|已全部整改|"
    r"(?:所有|全部)(?:的)?(?:反馈)?(?:问题|整改事项|整改任务)(?:均|都|全部)?(?:已经?)?"
    r"(?:完成整改|整改完毕|整改到位|整改完成|销号)|"
    r"(?:所有|全部)(?:的)?(?:反馈)?(?:问题|整改事项|整改任务)(?:均|都|全部)?已经?(?:整改|完成)(?=\s*(?:[。；;，,]|$))"
)
# "6个问题均已完成整改""反馈的6个问题，均已整改到位"：主语是问题总体时才算，"3个科室均已完成整改"不算
ALL_DONE_EACH_RE = re.compile(r"(?:均|都)已经?(?:完成整改|整改到位|整改完毕|整改完成|销号)")
_EACH_SUBJECT_RE = re.compile(r"(?:问题|整改事项|整改任务|反馈意见)\s*[，,]?\s*$")
_EACH_COUNT_RE = re.compile(NUM + r"\s*(?:个|项|条)\s*(?:具体)?(?:问题|整改事项|整改任务)")
BASIC_RE = re.compile(r"基本完成|基本整改到位|大部分完成")
PERIOD_RE = re.compile(r"本月|本期|本周|本季度|本阶段|当月|本轮|新增|新")
CUMULATIVE_RE = re.compile(r"截至|截止|累计|共|总|目前")
IDIOM_AFTER_RE = re.compile(r"\s*[、，,]?\s*(?:销号|验收|核销|办结)\s*[一1]")
STAT_EXCLUDE_GAP = re.compile(r"完成|整改|推进|坚持|启动|销号|其中")


def _stat_values(text: str, kind: str) -> List[Tuple[int, int, str]]:
    values = []
    seen = set()
    for regex in STAT_PATTERNS[kind]:
        for m in regex.finditer(text):
            groups = [g for g in m.groups() if g is not None]
            if kind == "total" and len(groups) == 2:
                if STAT_EXCLUDE_GAP.search(groups[0]):
                    continue
                num_text = groups[1]
            else:
                num_text = groups[-1]
            if kind == "done" and is_planned(text, m.start(), m.group(0)):
                continue
            if kind == "done" and IDIOM_AFTER_RE.match(text, m.end()):
                continue  # "整改完成一项、销号一项"是工作要求，不是统计
            if ORDINAL_LEAD_RE.search(text[max(0, m.start() - 20):m.start()]):
                continue  # "第1、4、5项问题已完成整改"：5 是编号
            value = cn_to_int(num_text)
            if value is None or m.start() in seen:
                continue
            seen.add(m.start())
            values.append((value, m.start(), m.group(0)))
    return values


def _line_of(lines: List[Tuple[int, str]], offsets: List[int], pos: int) -> int:
    line_no = lines[0][0] if lines else 0
    for (no, _), start in zip(lines, offsets):
        if start <= pos:
            line_no = no
        else:
            break
    return line_no


def _all_done_claims(body: str, summary: Dict[str, object]) -> List[re.Match]:
    """找出"全部完成整改"一类说法，按位置排序。"均已完成整改"只在主语是问题总体时算。"""
    found = {m.start(): m for m in ALL_DONE_RE.finditer(body)}
    for m in ALL_DONE_EACH_RE.finditer(body):
        if any(s <= m.start() < f.end() for s, f in found.items()):
            continue
        sentence = _sentence_prefix(body, m.start())
        if not _EACH_SUBJECT_RE.search(sentence) or re.search(r"其中|部分|个别", sentence):
            continue
        counts = [cn_to_int(c) for c in _EACH_COUNT_RE.findall(sentence)]
        if counts and counts[-1] is not None and counts[-1] != summary["total"]:
            continue  # "立行立改的3个问题均已整改到位"：说的是其中一部分，由统计核对处理
        found[m.start()] = m
    return [found[k] for k in sorted(found)]


def check_statistics(lines: List[Tuple[int, str]], ledger: Ledger, summary: Dict[str, object],
                     plan: bool = False) -> List[dict]:
    issues: List[dict] = []
    offsets, parts, pos = [], [], 0
    for _, text in lines:
        offsets.append(pos)
        t = text.translate(FULLWIDTH)
        parts.append(t)
        pos += len(t) + 1
    body = "\n".join(parts)
    views = [summary] + ([summary["groups"]] if "groups" in summary else [])

    def expect(kind: str) -> List[int]:
        out = []
        for v in views:
            c = v["status_count"]
            if kind == "total":
                out.append(v["total"])
            elif kind == "done":
                out += [v["done"], c[DONE]]
            elif kind == "immediate":
                out.append(c[IMMEDIATE])
            elif kind == "ongoing":
                out += [c[ONGOING], c[ONGOING] + c[NOT_STARTED]]
            elif kind == "long":
                out.append(c[LONGTERM])
            elif kind == "not_started":
                out.append(c[NOT_STARTED])
        return sorted(set(out))

    names = {"total": "问题总数", "done": "完成数", "immediate": "立行立改数", "ongoing": "持续推进数",
             "long": "长期坚持数", "not_started": "未启动数"}
    for kind in ("total", "done", "immediate", "ongoing", "long", "not_started"):
        allowed = expect(kind)
        for value, start, matched in _stat_values(body, kind):
            if value in allowed:
                continue
            if kind in ("done", "immediate") and value <= max(allowed):
                # "本月新完成整改1个"是本期数，不与累计数比对；"截至本月已完成整改8个"仍是累计数
                prefix = _clause_prefix(body, start, 10)
                if PERIOD_RE.search(prefix) and not CUMULATIVE_RE.search(prefix):
                    continue
            over = kind == "done" and value > max(allowed)
            level = "P0" if over else "P1"
            c = summary["status_count"]
            if kind == "done":
                expected = f"{summary['done']}（已完成 {c[DONE]}，立行立改 {c[IMMEDIATE]}）"
            else:
                expected = str(expect(kind)[0]) if kind != "total" else str(summary["total"])
            if "groups" in summary:
                expected += "；按大项汇总另有口径"
            issues.append(_issue("RC-201", level, "报告", "统计不符", "", _line_of(lines, offsets, start), matched,
                                 f"报告中的{names[kind]}为 {value}，台账为 {expected}",
                                 "统计数一律从台账计算；完成数＝已完成＋立行立改；统计的是部分问题的，写明口径"))
    for m in RATE_RE.finditer(body):
        value = float(m.group(1))
        decimals = len(m.group(1).split(".")[1]) if "." in m.group(1) else 0
        tol = 0.5 * (10 ** -decimals) + 1e-9
        rates = []
        for v in views:
            if v["total"]:
                rates.append(v["done"] / v["total"] * 100)
        if not rates or any(abs(value - r) <= tol for r in rates):
            continue
        level = "P0" if value > max(rates) + tol else "P1"
        issues.append(_issue("RC-202", level, "报告", "完成率不符", "", _line_of(lines, offsets, m.start()), m.group(0),
                             f"报告完成率为 {m.group(1)}%，按台账计算为 {round(summary['done_rate'], 1)}%",
                             "完成率＝（已完成＋立行立改）÷问题总数×100%，保留一位小数；长期坚持、持续推进不计入"))
    all_done = summary["done"] == summary["total"] and not summary["unknown"]
    if not all_done:
        for m in _all_done_claims(body, summary):
            if is_planned(body, m.start(), m.group(0)):
                continue
            past = "已" in m.group(0) or "已" in _sentence_prefix(body, m.start())
            if plan and not past:
                # 整改方案里的目标句没写时限，如"反馈问题全部整改到位"：提示写明"到……前"，不当作夸大完成
                issues.append(_issue("RC-203", "P2", "报告", "目标表述", "", _line_of(lines, offsets, m.start()),
                                     _excerpt(body, m.group(0)),
                                     f"方案中的“{m.group(0)}”没有写明时限，容易读成已经完成",
                                     "写成目标句：到2026年12月31日前，反馈问题全部整改到位"))
                continue
            issues.append(_issue("RC-203", "P0", "报告", "夸大完成", "", _line_of(lines, offsets, m.start()),
                                 _excerpt(body, m.group(0)),
                                 f"台账完成 {summary['done']} 条、共 {summary['total']} 条，报告写“{m.group(0)}”",
                                 "改为统计句：已完成整改 n 个（其中立行立改 m 个），持续推进 x 个，需长期坚持 y 个"))
    for m in BASIC_RE.finditer(body):
        issues.append(_issue("RC-204", "P1", "报告", "状态口径", "", _line_of(lines, offsets, m.start()),
                             _excerpt(body, m.group(0)),
                             f"“{m.group(0)}”不是整改状态口径，容易把未完成写成完成",
                             "按台账写明已完成几个、持续推进几个"))
    return issues


# ---------------------------------------------------------------------------
# 反馈原文核对
# ---------------------------------------------------------------------------

_PUNCT_RE = re.compile(r"[\s，。、；：！？,.;:!?“”‘’\"'（）()《》〈〉【】\[\]—\-…·]")


def _plain(text: str) -> str:
    """比对用：全角转半角，去掉空白和标点。"""
    return _PUNCT_RE.sub("", (text or "").translate(FULLWIDTH))


# 反馈原文的条目编号："1.""（1）""一、""第3条""问题3："
FEEDBACK_ENTRY_RE = re.compile(
    r"^\s*(?:第\s*([一二三四五六七八九十百\d]+)\s*[条项]\s*[，,、：:.]?"
    r"|[（(]\s*([一二三四五六七八九十百\d]+)\s*[)）]"
    r"|(\d+)\s*[.、](?!\d)"
    r"|([一二三四五六七八九十百]+)\s*、"
    r"|问题\s*(\d+)\s*[：:])\s*"
)
# 删了就会改变问题性质或程度的限定词
QUALIFIER_RE = re.compile(r"部分|个别|少数|多数|大部分|一些|有的|长期|多次|反复|屡次|普遍|至今|一直|严重|大量|较多|仍然|仍")
FACT_NUM_RE = re.compile(r"\d+(?:\.\d+)?\s*(?:%|万元|亿元|[^\d\s，。、；：！？,.;:!?“”\"《》（）()]?)")
FACT_QUOTE_RE = re.compile(r"“([^”]{1,40})”|\"([^\"]{1,40})\"|「([^」]{1,40})」|《([^》]{1,60})》")


def split_feedback(text: str) -> List[Tuple[Optional[int], str]]:
    """把反馈原文拆成条目 [(编号, 文字)]。

    有编号的行是一条，没有句末标点就断行的下一行并入该条（软换行）；
    没有编号的段落按句拆开，每句一条、编号为 None。
    """
    entries: List[List] = []
    for raw in text.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        line = raw.translate(FULLWIDTH).strip()
        if not line:
            continue
        m = FEEDBACK_ENTRY_RE.match(line)
        if m:
            num_text = next(g for g in m.groups() if g)
            entries.append([cn_to_int(num_text), line[m.end():].strip()])
        elif entries and entries[-1][0] is not None and not re.search(r"[。；;！!？?：:]$", entries[-1][1]):
            entries[-1][1] += line
        else:
            for sentence in re.split(r"(?<=[。；;！!？?])", line):
                if sentence.strip():
                    entries.append([None, sentence.strip()])
    return [(n, t) for n, t in entries if t]


def _bigrams(text: str) -> set:
    return {text[i:i + 2] for i in range(len(text) - 1)}


def _match_entry(key: str, problem: str, entries: List[Tuple[Optional[int], str]]) -> Optional[int]:
    """为台账一条找反馈原文中对应的条目：编号相同且重合度够的优先，否则取重合度最高的。"""
    grams = _bigrams(_plain(problem))
    if not grams:
        return None
    number = re.search(r"\d+", key)
    number = int(number.group(0)) if number else None
    scored = []
    for index, (no, text) in enumerate(entries):
        score = len(grams & _bigrams(_plain(text))) / len(grams)
        scored.append((score, no == number and number is not None, index))
    same = [s for s in scored if s[1] and s[0] >= 0.5]
    if same:
        return max(same)[2]
    best = max(scored) if scored else None
    return best[2] if best and best[0] >= 0.6 else None


def _facts(clause: str) -> List[str]:
    """分句中的数字（连同后面一个字的单位）、引号和书名号里的内容、限定词。"""
    facts = [m.group(0).strip() for m in FACT_NUM_RE.finditer(clause)]
    facts += [next(g for g in m.groups() if g) for m in FACT_QUOTE_RE.finditer(clause)]
    facts += QUALIFIER_RE.findall(clause)
    return facts


def check_feedback(ledger: Ledger, feedback_text: str) -> List[dict]:
    """台账问题表现与反馈原文核对。

    RC-301：台账的分句在反馈原文中找不到（被改写、弱化）；
    RC-302：台账对应的那一条反馈原文里，有数字、金额、限定词或引号内容的分句没有录进台账（被删减、截短）。
    同一条拆成子项（4-1、4-2）的，合起来核对。
    """
    issues: List[dict] = []
    if "问题表现" not in ledger.columns:
        return [_issue("RC-301", "P2", "台账", "问题表现", "", 0, "",
                       "台账没有“问题表现”列，无法与反馈原文核对", "增加“问题表现”列，照录反馈原文")]
    source = _plain(feedback_text)
    entries = split_feedback(feedback_text)
    siblings: Dict[str, List[str]] = {}
    for key in ledger.items:
        parent = parent_key(key)
        if parent and parent not in ledger.items:
            siblings.setdefault(parent, []).append(key)
    reported_entries: set = set()
    for item in ledger.items.values():
        problem = item.fields.get("problem", "")
        if not problem or MARKER_RE.search(problem):
            continue
        parts = [p for p in re.split(r"[；;。\n]", problem) if len(_plain(p)) >= 6]
        missing = [p.strip() for p in parts if _plain(p) not in source]
        if missing:
            issues.append(_issue("RC-301", "P1", "台账", "问题表述", item.key, item.row, missing[0][:40],
                                 f"编号 {item.raw_id} 的问题表现与反馈原文对不上，可能被改写或弱化",
                                 "问题表现照录反馈原文，不改字、不删限定和数字；确需概括的另设“问题简称”列"))

        index = _match_entry(item.key, problem, entries)
        if index is None or index in reported_entries:
            continue
        group = siblings.get(parent_key(item.key) or "", [item.key])
        if item.key not in group:
            group = [item.key]
        recorded = _plain("".join(ledger.items[k].fields.get("problem", "") for k in group))
        lost = []
        for clause in re.split(r"[，,；;。：:！!？?]", entries[index][1]):
            plain = _plain(clause)
            if len(plain) < 2 or plain in recorded:
                continue
            if any(_plain(f) and _plain(f) not in recorded for f in _facts(clause)):
                lost.append(clause.strip())
        if lost:
            reported_entries.add(index)
            more = f"等 {len(lost)} 处" if len(lost) > 1 else ""
            label = "、".join(ledger.items[k].raw_id for k in group)
            issues.append(_issue("RC-302", "P1", "台账", "问题表述", item.key, item.row, lost[0][:40],
                                 f"编号 {label} 的问题表现删减了反馈原文中的“{lost[0]}”{more}",
                                 "问题表现照录反馈原文，数字、金额、限定词和引号里的内容都要保留；"
                                 "需要短名的另设“问题简称”列"))
    return issues


WEAK_RE = re.compile(r"有待|进一步|需加强|需要加强|还需|仍需|不够|不足|个别|瑕疵|欠缺|不尽|还不|尚不|稍有|略有|需改进")


def check_weakening(heads: Dict[str, Tuple[int, str]], ledger: Ledger) -> List[dict]:
    """报告中"问题N："后面的问题表述，与台账问题表现对不上且多了弱化词的，提示可能弱化。"""
    issues: List[dict] = []
    if "问题表现" not in ledger.columns:
        return issues
    for key, (line, head) in heads.items():
        shown, keys = resolve_key(key, ledger)
        if not keys or len(_PUNCT_RE.sub("", head)) < 4:
            continue
        problem = "".join(ledger.items[k].fields.get("problem", "") for k in keys)
        if not problem or _PUNCT_RE.sub("", head) in _PUNCT_RE.sub("", problem):
            continue
        added = [w for w in WEAK_RE.findall(head) if w not in problem]
        if added:
            issues.append(_issue("RC-108", "P1", "报告", "问题表述", key, line, head[:40],
                                 f"问题{shown}的表述与台账问题表现不一致，多了“{added[0]}”，可能弱化了问题",
                                 "照录台账（反馈原文）的问题表现；要用短名的，用台账“问题简称”"))
    return issues


# ---------------------------------------------------------------------------
# 汇总输出
# ---------------------------------------------------------------------------

LEVEL_ORDER = {"P0": 0, "P1": 1, "P2": 2}


def build_report(ledger_path: Path, report_path: Optional[Path] = None,
                 feedback_path: Optional[Path] = None, as_of: Optional[_dt.date] = None,
                 plan: bool = False) -> dict:
    ledger = load_ledger(ledger_path)
    issues = check_ledger(ledger, as_of)
    summary = ledger_summary(ledger)
    result: "OrderedDict[str, object]" = OrderedDict()
    result["ledger"] = OrderedDict([("file", str(ledger_path)), ("table", ledger.where),
                                    ("columns", ledger.columns), ("summary", summary)])
    result["report"] = None
    result["feedback"] = None
    result["as_of"] = as_of.isoformat() if as_of else None
    if feedback_path is not None:
        if not feedback_path.is_file():
            raise InputError(f"未找到反馈原文：{feedback_path}")
        feedback_lines = load_report(feedback_path)
        issues += check_feedback(ledger, "\n".join(t for _, t in feedback_lines))
        result["feedback"] = str(feedback_path)
    if report_path is not None:
        lines = load_report(report_path)
        heads: Dict[str, Tuple[int, str]] = OrderedDict()
        blocks, mentioned = segment_report(lines, ledger, heads)
        if not mentioned:
            issues.append(_issue("RC-107", "P1", "报告", "无编号", "", 0, "",
                                 "报告中没有找到台账编号，无法逐条对账",
                                 "逐条整改情况写明“问题1”“问题2”……，与台账编号一致"))
        else:
            issues += compare_items(blocks, mentioned, ledger, plan)
            issues += check_weakening(heads, ledger)
        issues += check_statistics(lines, ledger, summary, plan)
        result["report"] = OrderedDict([
            ("file", str(report_path)),
            ("char_count", sum(len(re.sub(r"\s", "", t)) for _, t in lines)),
            ("mentioned", list(mentioned.keys())),
        ])
    issues.sort(key=lambda i: (LEVEL_ORDER[i["level"]], 0 if i["where"] == "报告" else 1, i["line"] or 0))
    counts = {"P0": 0, "P1": 0, "P2": 0}
    for issue in issues:
        counts[issue["level"]] += 1
    result["summary"] = OrderedDict([("P0", counts["P0"]), ("P1", counts["P1"]), ("P2", counts["P2"]),
                                     ("markers", ledger_markers(ledger))])
    result["issues"] = issues
    return result


def render_text(result: dict) -> str:
    led = result["ledger"]["summary"]
    c = led["status_count"]
    lines = [
        f"台账：{result['ledger']['file']}（{result['ledger']['table']}）",
        f"  共 {led['total']} 条：已完成 {c[DONE]}，立行立改 {c[IMMEDIATE]}，持续推进 {c[ONGOING]}，"
        f"长期坚持 {c[LONGTERM]}，未启动 {c[NOT_STARTED]}" + (f"，状态待定 {led['unknown']}" if led["unknown"] else "")
        + f"；完成 {led['done']} 条，完成率 {led['done_rate']}%",
    ]
    if "groups" in led:
        g = led["groups"]
        lines.append(f"  按大项汇总：共 {g['total']} 个，完成 {g['done']} 个，完成率 {g['done_rate']}%")
    if result.get("report"):
        lines.append(f"报告：{result['report']['file']}（提到编号 {len(result['report']['mentioned'])} 个）")
    if result.get("feedback"):
        lines.append(f"反馈原文：{result['feedback']}")
    s = result["summary"]
    markers = s["markers"]
    lines.append(f"P0 {s['P0']} 条，P1 {s['P1']} 条，P2 {s['P2']} 条；台账中待补 {markers['待补']} 处、待核实 {markers['待核实']} 处")
    for issue in result["issues"]:
        where = issue["where"] + (f"第 {issue['line']} 行" if issue["line"] else "")
        lines.append(f"[{issue['level']}] {issue['id']} {issue['category']}（{where}）：{issue['message']}")
        if issue["excerpt"]:
            lines.append(f"      原文：{issue['excerpt']}")
        lines.append(f"      建议：{issue['suggestion']}")
    if not result["issues"]:
        lines.append("未发现问题。")
    return "\n".join(lines)


def _setup_console() -> None:
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8")
        except Exception:
            pass


class _ChineseHelpFormatter(argparse.RawDescriptionHelpFormatter):
    def add_usage(self, usage, actions, groups, prefix=None):
        return super().add_usage(usage, actions, groups, "用法：" if prefix is None else prefix)


def make_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        prog="rectify_check.py",
        description="整改台账与整改报告对账：逐条比对编号和整改状态，核对统计数和完成率，并检查台账是否完整。"
                    "只给台账时只做台账自检。",
        epilog="退出码：0 正常；1 使用 --strict 且存在 P0；2 输入错误。",
        formatter_class=_ChineseHelpFormatter,
        add_help=False,
    )
    parser._optionals.title = "选项"
    parser._positionals.title = "参数"
    parser.add_argument("-h", "--help", action="help", help="显示本帮助信息并退出")
    parser.add_argument("ledger", help="整改台账（.md、.txt、.xlsx）")
    parser.add_argument("report", nargs="?", help="整改报告或进展报告（.txt、.md、.docx），可省略")
    parser.add_argument("--feedback", help="反馈意见原文（.txt、.md、.docx），用于核对问题表现是否照录")
    parser.add_argument("--date", help="对账日期 YYYY-MM-DD，用于判断是否超过完成时限")
    parser.add_argument("--json", action="store_true", help="以 JSON 格式输出")
    parser.add_argument("--plan", action="store_true",
                        help="第二个文件是整改方案：只查覆盖、问题表述和夸大完成，不要求每条写明状态")
    parser.add_argument("--strict", action="store_true", help="存在 P0 时返回 1")
    return parser


def main(argv: Optional[Sequence[str]] = None) -> int:
    _setup_console()
    parser = make_parser()
    try:
        args = parser.parse_args(argv)
    except SystemExit as exc:
        return 0 if exc.code == 0 else 2
    as_of = None
    if args.date:
        try:
            as_of = _dt.date.fromisoformat(args.date)
        except ValueError:
            print(f"错误：--date 须为 YYYY-MM-DD 格式：{args.date}", file=sys.stderr)
            return 2
    try:
        result = build_report(Path(args.ledger), Path(args.report) if args.report else None,
                              Path(args.feedback) if args.feedback else None, as_of, plan=args.plan)
    except InputError as exc:
        print(f"错误：{exc}", file=sys.stderr)
        return 2
    if args.json:
        print(json.dumps(result, ensure_ascii=False, indent=2))
    else:
        print(render_text(result))
    if args.strict and result["summary"]["P0"]:
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
