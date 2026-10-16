#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""信息、简报、新闻稿、公众号推文的专项检查（gongwen-brief）。

只做 gongwen_lint.py 没有覆盖、本类稿件最容易出问题的几项：
    BRF-001  篇幅超出或低于常用范围（按稿件类型，或 --min、--max 指定）
    BRF-002  标题中的数字概括（“三项举措”“‘四个到位’”）与正文实际分项条数不一致
    BRF-003  引号内的原话、提法在用户材料中找不到（需 --source）；短引语不是材料中的原话、
             说话人与材料对不上、荣誉称号无出处
    BRF-004  长引语没有交代说话人
    BRF-005  导语先写铺垫（“为深入贯彻……”“随着……”）
    BRF-006  导语没有交代时间（信息、新闻稿）
    BRF-007  无依据的最高级、首创、领先、“全市率先”“走在全国前列”类说法
    BRF-008  疗效保证类说法（治愈率、无副作用等；“根治”“立竿见影”等只在医疗语境按疗效保证处理，
             “根治术”等手术名称不算）
    BRF-009  宣传升学率、状元、高分考生（批评“片面追求升学率”的和内部报送稿只作提示）
    BRF-010  标题党、自媒体腔用语，正文连续感叹号
    BRF-011  感叹号过多（引语之外全文超过一处）
    BRF-012～014  简报报头（期号、编印单位、日期）和报尾（报送范围）要素
    BRF-015  对外稿涉及照片、视频，提醒确认授权
    BRF-016  标题空词（“多措并举”“扎实推进”“取得明显成效”等）

标点、数字、格式、隐私、AI 腔等通用检查由 gongwen_lint.py 负责，二者配合使用。

用法：
    python scripts/brief_check.py <稿件> [--type 信息|简报|新闻稿|公众号|紧急信息|社区信息] [--source 材料 ...]
                                  [--external] [--unit 单位类型] [--min 字数] [--max 字数] [--json] [--strict]

只检查【起草说明】之前的正文。支持 .txt、.md、.docx（docx 需要 python-docx）。
简报按报头、报尾之间的各篇分别核对标题数字概括。--unit 为公立医院（含“医院”“卫生”“医疗”字样）时，全文按医疗语境检查疗效保证。
退出码：0 正常；--strict 时存在 P0 返回 1；输入错误返回 2。只依赖标准库，不访问网络，不写任何文件。
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Tuple, Union

BOUNDARY = "【起草说明】"
DRAFT_NOTE_RE = re.compile(r"^\s*(?:#{1,6}\s*)?(?:(?:\*\*|__)\s*)?" + re.escape(BOUNDARY))
SUPPORTED_SUFFIXES = (".txt", ".md", ".markdown", ".text", ".docx")

TYPES = ("信息", "简报", "新闻稿", "公众号")
TYPE_ALIASES = {
    "政务信息": "信息", "工作信息": "信息", "经验信息": "信息", "经验做法信息": "信息",
    "工作简报": "简报", "会议简报": "简报",
    "新闻": "新闻稿", "消息": "新闻稿",
    "推文": "公众号", "公众号推文": "公众号", "微信推文": "公众号",
    "紧急信息": "信息", "社区信息": "信息",
}
# 默认篇幅（工具约定，单位画像或上级信息部门另有要求的，用 --min、--max 覆盖）
DEFAULT_RANGES: Dict[str, Tuple[Optional[int], Optional[int]]] = {
    "信息": (500, 800),
    "简报": (None, 3000),
    "新闻稿": (500, 1000),
    "公众号": (800, 1500),
}
# 按别名另定篇幅的类型：紧急信息越短越好，只设上限
ALIAS_RANGES: Dict[str, Tuple[Optional[int], Optional[int]]] = {
    "紧急信息": (None, 500),
    "社区信息": (300, 500),
}
EXTERNAL_TYPES = ("新闻稿", "公众号")
BASIS_LENGTH = "references/types/ 对应文件的篇幅约定；单位画像或上级信息部门另有要求的从其要求"

CN_DIGITS = {"一": 1, "二": 2, "两": 2, "三": 3, "四": 4, "五": 5, "六": 6, "七": 7, "八": 8, "九": 9, "十": 10}

# ---------------------------------------------------------------------------
# 正则
# ---------------------------------------------------------------------------

MARKER_RE = re.compile(r"【(?:待补|待核实)[^】]*】")
SPACE_RE = re.compile(r"\s+")
# 用于引语比对的归一化：去掉空白和中英文标点
PUNCT_RE = re.compile(
    r"[\s\u3000，。、；：？！…—\-–·“”‘’「」『』《》〈〉（）()\[\]【】〔〕,.;:?!'\"~～/\\|]+"
)
FULLWIDTH_TABLE = str.maketrans(
    "０１２３４５６７８９ＡＢＣＤＥＦＧＨＩＪＫＬＭＮＯＰＱＲＳＴＵＶＷＸＹＺａｂｃｄｅｆｇｈｉｊｋｌｍｎｏｐｑｒｓｔｕｖｗｘｙｚ％",
    "0123456789ABCDEFGHIJKLMNOPQRSTUVWXYZabcdefghijklmnopqrstuvwxyz%",
)

TITLE_END_PUNCT = tuple("。！？；…")
BYLINE_RE = re.compile(
    r"^\s*(?:[（(](?:通讯员|记者|作者|撰稿|供稿|文|图|摄影?|编辑)[^）)]{0,30}[）)]"
    r"|(?:来源|作者|撰稿|供稿|编辑|审核|责任编辑|文|图|摄影?|通讯员|记者|摘要|导读)\s*[：:/／])"
)
DATELINE_RE = re.compile(r"^\s*(?:本报讯|本网讯|本刊讯|讯)\s*(?:[（(][^）)]{0,30}[）)])?\s*")
SUBTITLE_RE = re.compile(r"^\s*[—－-]{1,2}")

# 标题中的数字概括
_N = r"(?P<n>[一二两三四五六七八九十]|10|[1-9])(?![0-9])"
QUOTED_SUMMARY_RE = re.compile(r"[“‘\"「『]" + _N + r"(?:个|张|项|条|招|步|大|化)[\u4e00-\u9fff]{0,6}[”’\"」』]")
PLAIN_SUMMARY_RE = re.compile(
    r"(?<![第0-9一二两三四五六七八九十百千万])" + _N +
    r"(?:项|条|招|步(?=工作法|走|法)|大(?=[\u4e00-\u9fff]{2})|个(?=方面|环节|到位|坚持|聚焦|结合|强化|突出|优先|统一|机制|平台|载体|抓手|一))"
)

# 固定提法中的数字不是对正文分项的概括，不参与比对
FIXED_TERMS_RE = re.compile(
    r"两个确立|两个维护|四个意识|四个自信|两个务必|三个务必|四个全面|五位一体|两个毫不动摇|两个结合|三个代表|"
    r"八项规定|两个大局|两个一百年|三个区分开来|四下基层|四个服务|四个面向|四个最严|两个健康|三个一批|两学一做|"
    r"三会一课|三重一大|四风|五个必由之路|六个必须坚持|十个坚持|两个先行|三个转变"
)

HEADING1_RE = re.compile(r"^\s*([一二三四五六七八九十]{1,3})、")
HEADING2_RE = re.compile(r"^\s*[（(]([一二三四五六七八九十]{1,3})[）)]")
ARABIC_ITEM_RE = re.compile(r"^\s*([1-9][0-9]?)[.．、](?![0-9])")
SHI_RE = re.compile(r"(?<![统唯同单划均专第])([一二三四五六七八九十])是")
LEAD_BYLINE_RE = re.compile(r"^[（(](?:通讯员|记者|特约记者|作者)[^）)]{0,30}[）)]\s*")

# 引号：中文双引号、直角引号、成对的英文直引号、单独使用的中文单引号（嵌在双引号内的随外层一起处理）
QUOTE_RE = re.compile(r"“([^“”\n]{1,300})”|「([^「」\n]{1,300})」|\"([^\"\n]{1,300})\"|‘([^‘’\n]{1,300})’")
SPEECH_PUNCT_RE = re.compile(r"[，。！？；…,!?;]")
# 口语痕迹：人称代词、语气词。名称、口号一般不带这些字
SPEECH_WORD_RE = re.compile(r"[我咱俺你您]|了|吧|呢|啊|呀|嘛|哇")
SPEAKER_BEFORE_RE = re.compile(
    r"(说|表示|介绍|告诉|坦言|感慨|讲|谈到|指出|强调|提到|(?<![简名号统俗又昵])称(?![为作号])|评价|回忆|感叹|直言|笑着|"
    r"写道|留言|反映|(?:笑|问|答|喊|叹)道)[^“”]{0,6}[：:，,]?\s*$"
)
# 引号紧跟在冒号后，视为交代了说话人，如 李大爷竖起大拇指：“……”
COLON_BEFORE_RE = re.compile(r"[：:]\s*$")
# 引号前是这些词的，引号内是名称、口号、提法，不按原话处理
NAME_INTRO_RE = re.compile(
    r"(?:简称|名称|称为|称作|统称|俗称|又称|号称|命名为|叫做|叫作|誉为|评为|获评|授予|冠名|主题为|题为|口号|标语|"
    r"品牌|栏目|活动|工作法|模式|机制|打造|推出|开展|实施|创建|建立|推行|启动)[^“”，。]{0,4}$"
)
SPEAKER_AFTER_RE = re.compile(
    r"^\s*[，,]?\s*(?P<who>[\u4e00-\u9fff某*＊]{0,10}?)(说|表示|介绍|告诉|坦言|感慨|评价|感叹|留言)"
)
# 说话人短语的切分符；代词、泛称不做说话人核对
SPEAKER_SPLIT_RE = re.compile(r"[，。；：！？、,.;:!?\s“”\"「」‘’（）()]")
SPEAKER_SKIP_RE = re.compile(r"^(?:他|她|他们|她们|大家|有人|对方|记者|笔者|我|我们|其|该负责人|有关负责人)$")
SPEAKER_NAME_RE = re.compile(r"[\u4e00-\u9fff]某某?")
# 短引语：去标点后不足这个字数的，子串匹配容易误中材料叙述中的片段
SHORT_QUOTE_LEN = 8
# 核对说话人时，在材料命中处前后各取的字数（去标点后）
SPEAKER_WINDOW_BEFORE = 30
SPEAKER_WINDOW_AFTER = 16
# 荣誉称号：引号前是这些词的，引号内按荣誉、称号处理，材料中找不到的升为 P1
HONOR_INTRO_RE = re.compile(r"(?:评为|获评|授予|荣获|命名为|获得|入选|认定为)[^“”\"，。]{0,4}$")

PADDING_RE = re.compile(
    r"^(?:为(?:深入|全面|认真|进一步|切实|扎实|更好地?|持续)|随着|当前，|近年来，随着|众所周知|"
    r"在[^。，]{2,30}(?:正确领导|坚强领导|关心支持|大力支持)下)"
)
TIME_RE = re.compile(
    r"\d{4}年|\d{1,2}月|\d{1,2}日|[一二三四五六七八九十]{1,3}月|近日|日前|近期|今年以来|年初以来|入汛以来|开学以来|"
    r"上半年|下半年|前三季度|一季度|二季度|三季度|四季度|本周|上周|今天|昨天|当天|当日|连日来|此前|"
    r"【待补[^】]*(?:时间|日期)[^】]*】|【待核实[^】]*(?:时间|日期)[^】]*】"
)

ABSOLUTE_RE = re.compile(
    r"(?:全国|全省|全市|全区|全县|国内|国际|业内|行业内?|同行业|同类)(?:范围内)?"
    r"(?:领先|首创|首家|唯一|一流|顶尖|第一|率先|最(?!美)[\u4e00-\u9fff]{1,2})"
    r"|(?:排名|位居|位列)[^，。；\n]{0,6}第一|首创|独创|独家|唯一|第一家|首家|填补[^，。；\n]{0,10}空白"
    r"|最先进|最权威|最佳|顶级|顶尖|(?:国际|国内)先进水平|零差错|零投诉|满意率(?:达到?|为)?100%|100%满意|首例"
    r"|(?:走在|位居|位列|处于)[^，。；\n]{0,8}前列"
)
# 疗效保证分两组：第一组本身是医疗用语，任何语境都按疗效保证处理；
# 第二组在机关稿件中常是工作提法（如“根治欠薪”“立竿见影”），只在医疗语境按疗效保证处理
EFFICACY_MED_RE = re.compile(r"治愈率|包治|药到病除|无副作用|保证(?:疗效|治愈)|彻底治愈|永不复发")
EFFICACY_CTX_RE = re.compile(r"根治|有效率|零风险|保证(?:有效|效果)|立竿见影|彻底(?:摆脱|告别)")
# 固定工作提法，不按效果保证处理
EFFICACY_EXCEPT_RE = re.compile(
    r"根治(?:欠薪|拖欠|[\u4e00-\u9fff]{0,6}(?:顽疾|顽瘴痼疾|乱象|水患|污染|腐败))"
    # 手术、治疗名称：胃癌根治术、根治性切除术、根治性放疗等
    r"|根治术|根治手术|根治性(?:切除|手术|放疗|化疗|放化疗|治疗|清扫)"
)
# “有效率”只在指比率时命中：前面是“治疗、临床、总”，或后面跟“达、为、数字”
RATE_BEFORE_RE = re.compile(r"(?:治疗|临床|总)$")
RATE_AFTER_RE = re.compile(r"^\s*(?:达到?|为|是|超过|高达|[0-9０-９])")
# 医疗语境：同段出现以下词语
MEDICAL_RE = re.compile(
    r"患者|病人|病例|病种|病情|疾病|治疗|疗效|疗程|药物|药品|用药|服药|手术|康复|医院|医疗|医生|医师|诊疗|门诊|住院|临床|"
    r"(?<![毛通弊])病(?!险)"
)
MEDICAL_UNIT_RE = re.compile(r"医院|医疗|卫生|卫健")
EDU_HYPE_RE = re.compile(
    r"(?:高考|中考|文科|理科|全省|全市|全区|全县)状元|升学率|名校率|清北率|考取清华北大|高分考生|(?:高考|中考)喜报|"
    r"一本率|重本率|(?:高考|中考|本科|一本|二本|重点|特控|特招)上线率"
)
# 批评性语境：“片面追求升学率”“按升学率考核教师”“严禁宣传高考状元”，不是宣传
EDU_CRITICAL_BEFORE_RE = re.compile(
    r"(?:片面追求|单纯追求|追求|唯|不得|严禁|禁止|不准|不再|不以|不搞|不能|不宜|反对|纠正|扭转|克服|破除|淡化)"
    r"[^，。；！？\n]{0,4}$"
)
EDU_CRITICAL_AFTER_RE = re.compile(r"^\s*(?:来|作为)?(?:考核|评价|挂钩|奖惩|论英雄)")
# “按升学率排名”是批评，“升学率排名全市第一”是宣传，排名只在前有“按、以、用、拿”时算批评
EDU_RANK_BEFORE_RE = re.compile(r"(?:按|以|用|拿)$")
EDU_RANK_AFTER_RE = re.compile(r"^\s*(?:来)?(?:排名|排队)")
# 标题空词：放在任何单位的任何工作上都成立，与 references/types/xinxi.md 第三部分第 4 条一致
EMPTY_TITLE_RE = re.compile(
    r"多措并举|扎实推进|取得了?(?:明显|显著|积极|良好|阶段性)?成效|成效(?:明显|显著|初显)|亮点纷呈|亮点频现|"
    r"精准施策|持续发力|再上新台阶|迈上新台阶|开创新局面|谱写新篇章"
)
TITLE_BAIT_RE = re.compile(
    r"震惊|刚刚|速看|必看|速转|紧急扩散|转发扩散|重磅|炸裂|绝了|竟然|居然|万万没想到|吐血|看哭|泪目|破防|太[\u4e00-\u9fff]{1,2}了|"
    r"[！!]{2,}|[？?][！!]|[！!][？?]"
)
# 正文中的连续感叹号、问叹连用，按标题党同样处理
BODY_BANG_RE = re.compile(r"[！!]{2,}|[？?][！!]|[！!][？?]")
SLANG_RE = re.compile(
    r"宝子们?|集美们?|绝绝子|冲鸭|安排上了?|种草|拔草|干货满满|码住|速来|手慢无|吃瓜|整活|氛围感|松弛感|显眼包|"
    r"(?<![A-Za-z])[Ee][Mm][Oo](?![A-Za-z])|家人们|小伙伴们|亲们"
)
EXCLAIM_RE = re.compile(r"[！!]")
# 引语之外全文最多一处感叹号（邀请、提示类结尾句），与 references/types/gongzhonghao.md 第五部分一致
EXCLAIM_LIMIT = 1
PHOTO_RE = re.compile(r"图为|照片|合影|配图|视频|摄影|【待补[^】]*图[^】]*】")

# 简报报头、报尾
ISSUE_NO_RE = re.compile(r"第\s*(?:[0-9０-９]+|[一二三四五六七八九十〇零百]+|【待(?:补|核实)[^】]*】)\s*期")
EDITOR_RE = re.compile(r"编|印发|主办|办公室|编辑部|秘书处|中心")
DATE_RE = re.compile(r"\d{4}\s*年\s*\d{1,2}\s*月\s*\d{1,2}\s*日|【待(?:补|核实)[^】]*(?:日期|时间)[^】]*】")
DISTRIBUTION_RE = re.compile(r"^\s*(?:报|报送|送|发|分送|抄送)\s*[：:]")
SEPARATOR_RE = re.compile(r"^\s*[─━—\-=＝_＿·.]{3,}\s*$")
BRIEF_NAME_RE = re.compile(r"(?:简报|动态|信息|专报|快报|通讯|工作交流)\s*$")
FOOTER_MISC_RE = re.compile(r"^\s*(?:[（(]?共印\s*\S*\s*份[）)]?|(?:责任编辑|审核|签发|编辑|联系人|联系电话)\s*[：:])")
# 简报分篇：篇目标题是不带句末标点的短行；括号注（“（某街道供稿）”）、“一是”开头的行、层次标题不算
ARTICLE_TITLE_MAX = 50
ARTICLE_TITLE_END = tuple("。；：:，、！？!?…")
PAREN_NOTE_RE = re.compile(r"^\s*[（(][^（）()]{0,40}[）)]\s*$")
SHI_START_RE = re.compile(r"^\s*[一二三四五六七八九十]是")
EDITOR_NOTE_RE = re.compile(r"^\s*(?:编者按|按语|编者的话)")


class InputError(Exception):
    """输入文件无法读取或参数错误。"""


# ---------------------------------------------------------------------------
# 读取
# ---------------------------------------------------------------------------

def read_text_lines(path: Path) -> List[str]:
    if not path.exists():
        raise InputError(f"文件不存在：{path}")
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise InputError(f"不支持的文件类型：{path.suffix or '无扩展名'}。只支持 .txt、.md、.docx。")
    if suffix == ".docx":
        try:
            import docx  # type: ignore
        except ImportError as exc:
            raise InputError("读取 .docx 需要 python-docx，请先运行：pip install python-docx") from exc
        try:
            document = docx.Document(str(path))
        except Exception as exc:  # 损坏或加密的 docx
            raise InputError(f"无法打开 docx 文件：{exc}") from exc
        lines = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                lines.append("　".join(cell.text for cell in row.cells))
        return lines
    try:
        raw = path.read_bytes()
    except OSError as exc:
        raise InputError(f"无法读取文件：{exc}") from exc
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise InputError("无法识别文件编码，请另存为 UTF-8 后重试。")
    return text.replace("\r\n", "\n").replace("\r", "\n").split("\n")


def body_lines(lines: Sequence[str]) -> List[Tuple[int, str]]:
    """返回【起草说明】之前的 (行号, 文本)。"""
    result: List[Tuple[int, str]] = []
    for number, line in enumerate(lines, start=1):
        if DRAFT_NOTE_RE.match(line):
            break
        result.append((number, line.rstrip()))
    return result


def normalize(text: str) -> str:
    text = MARKER_RE.sub("", text)
    return PUNCT_RE.sub("", text.translate(FULLWIDTH_TABLE))


def visible_len(text: str) -> int:
    return len(SPACE_RE.sub("", text))


def cn_int(token: str) -> int:
    if token.isdigit():
        return int(token)
    if token in CN_DIGITS:
        return CN_DIGITS[token]
    if token.startswith("十"):
        return 10 + CN_DIGITS.get(token[1:], 0)
    if "十" in token:
        head, _, tail = token.partition("十")
        return CN_DIGITS.get(head, 0) * 10 + CN_DIGITS.get(tail, 0)
    return 0


# ---------------------------------------------------------------------------
# 结构
# ---------------------------------------------------------------------------

class Draft:
    """稿件结构：标题行、正文段落、报头报尾（简报）。"""

    def __init__(self, lines: Sequence[Tuple[int, str]], doc_type: str) -> None:
        self.doc_type = doc_type
        self.nonempty = [(n, t.strip()) for n, t in lines if t.strip()]
        self.header: List[Tuple[int, str]] = []
        self.footer: List[Tuple[int, str]] = []
        self.titles: List[Tuple[int, str]] = []
        self.bylines: List[Tuple[int, str]] = []
        self.paragraphs: List[Tuple[int, str]] = []
        if doc_type == "简报":
            self._split_brief()
        else:
            self._split_article()

    def _split_article(self) -> None:
        rest = list(self.nonempty)
        # 标题：开头连续的短行（引题、主题、副题，最多 3 行），遇到句末标点、署名行、层次标题即停
        while rest and len(self.titles) < 3:
            number, text = rest[0]
            if BYLINE_RE.match(text) or DATELINE_RE.match(text) and len(text) > 8:
                break
            if HEADING1_RE.match(text) or HEADING2_RE.match(text) or ARABIC_ITEM_RE.match(text):
                break
            # 第一行即使以感叹号、问号结尾也按标题处理（标题党多是这样），后续行遇句末标点即停
            ends = ("。", "；") if not self.titles else TITLE_END_PUNCT
            if text.endswith(ends) or len(text) > 50:
                break
            if SUBTITLE_RE.match(text) or not self.titles or len(text) <= 40:
                self.titles.append((number, text))
                rest.pop(0)
                continue
            break
        for number, text in rest:
            if BYLINE_RE.match(text) and len(text) <= 60:
                self.bylines.append((number, text))
            else:
                self.paragraphs.append((number, text))

    @staticmethod
    def _is_header_line(index: int, text: str) -> bool:
        if ISSUE_NO_RE.search(text) or SEPARATOR_RE.match(text):
            return True
        if index <= 1 and len(text) <= 40 and BRIEF_NAME_RE.search(text):
            return True
        if len(text) <= 60 and DATE_RE.search(text) and (EDITOR_RE.search(text) or len(text) <= 16):
            return True
        if len(text) <= 30 and re.search(r"(?:编印|编发|主办|编)\s*$|^\s*(?:编印|编发|主办)", text):
            return True
        return bool(len(text) <= 30 and re.search(r"内部资料|注意保存|秘密|机密", text))

    def _split_brief(self) -> None:
        rest = list(self.nonempty)
        # 报头：开头 8 行内的简报名称、期号、编印单位、日期、分隔线
        index = 0
        while index < len(rest) and index < 8 and self._is_header_line(index, rest[index][1]):
            self.header.append(rest[index])
            index += 1
        rest = rest[index:]
        # 报尾：报、送、发、共印、责任编辑、分隔线
        while rest:
            number, text = rest[-1]
            if (DISTRIBUTION_RE.match(text) or SEPARATOR_RE.match(text) or FOOTER_MISC_RE.match(text)):
                self.footer.insert(0, (number, text))
                rest.pop()
                continue
            break
        self.paragraphs = rest

    @property
    def title_text(self) -> str:
        return "".join(t for _, t in self.titles)

    @property
    def body_text(self) -> str:
        return "\n".join(t for _, t in self.paragraphs)

    def char_count(self) -> int:
        return sum(visible_len(t) for _, t in self.paragraphs)

    def lead(self) -> Optional[Tuple[int, str]]:
        for number, text in self.paragraphs:
            stripped = LEAD_BYLINE_RE.sub("", DATELINE_RE.sub("", text, count=1), count=1)
            if stripped:
                return number, stripped
        return None

    @staticmethod
    def _is_article_title(text: str) -> bool:
        if len(text) > ARTICLE_TITLE_MAX or text.endswith(ARTICLE_TITLE_END):
            return False
        if HEADING1_RE.match(text) or HEADING2_RE.match(text) or ARABIC_ITEM_RE.match(text):
            return False
        if SHI_START_RE.match(text) or EDITOR_NOTE_RE.match(text):
            return False
        return not (PAREN_NOTE_RE.match(text) or BYLINE_RE.match(text))

    def sections(self) -> List["Section"]:
        """简报按篇切分：每篇为篇目标题（可含栏目名、副题）加其后的正文段落。"""
        result: List[Section] = []
        current: Optional[Section] = None
        for number, text in self.paragraphs:
            if self._is_article_title(text):
                # 连续的短行（栏目名、引题、主题、副题）合并为同一篇的标题
                if current is not None and not current.paragraphs:
                    current.titles.append((number, text))
                    continue
                current = Section([(number, text)], [])
                result.append(current)
                continue
            if current is None:
                current = Section([], [])
                result.append(current)
            current.paragraphs.append((number, text))
        return result


class Section:
    """简报中的一篇：与 Draft 同样提供 titles、paragraphs、title_text、body_text。"""

    def __init__(self, titles: List[Tuple[int, str]], paragraphs: List[Tuple[int, str]]) -> None:
        self.titles = titles
        self.paragraphs = paragraphs

    @property
    def title_text(self) -> str:
        return "".join(t for _, t in self.titles)

    @property
    def body_text(self) -> str:
        return "\n".join(t for _, t in self.paragraphs)


# ---------------------------------------------------------------------------
# 检查
# ---------------------------------------------------------------------------

def issue(rule: str, level: str, category: str, line: Optional[int], excerpt: str, message: str,
          suggestion: str, basis: str) -> dict:
    return {
        "id": rule, "level": level, "category": category, "line": line,
        "excerpt": excerpt[:80], "message": message, "suggestion": suggestion, "basis": basis,
    }


def check_length(draft: Draft, low: Optional[int], high: Optional[int], doc_type: str) -> List[dict]:
    count = draft.char_count()
    issues: List[dict] = []
    if high is not None and count > high:
        level = "P1" if doc_type in ("信息", "简报") else "P2"
        issues.append(issue(
            "BRF-001", level, "篇幅", None, f"正文约 {count} 字",
            f"正文约 {count} 字，超过上限 {high} 字",
            "按“铺垫、复述、空泛表态、重复举例”的顺序删减，事实和数据保留",
            BASIS_LENGTH,
        ))
    if low is not None and count < low:
        issues.append(issue(
            "BRF-001", "P2", "篇幅", None, f"正文约 {count} 字",
            f"正文约 {count} 字，低于常用下限 {low} 字",
            "素材不足时按实际篇幅交付，不加背景和表态凑字；缺的事实标【待补】",
            BASIS_LENGTH,
        ))
    return issues


def title_summaries(title: str) -> List[Tuple[int, str]]:
    found: List[Tuple[int, str]] = []
    spans: List[Tuple[int, int]] = [m.span() for m in FIXED_TERMS_RE.finditer(title)]
    for match in QUOTED_SUMMARY_RE.finditer(title):
        if any(s <= match.start() + 1 < e for s, e in spans):
            continue
        found.append((cn_int(match.group("n")), match.group(0)))
        spans.append(match.span())
    for match in PLAIN_SUMMARY_RE.finditer(title):
        if any(s <= match.start() < e for s, e in spans):
            continue
        tail = title[match.end():match.end() + 6]
        found.append((cn_int(match.group("n")), match.group(0) + tail))
    return [(n, text) for n, text in found if 2 <= n <= 10]


def item_counts(draft: Union["Draft", "Section"]) -> Tuple[List[int], str]:
    """返回正文可计数的分项条数（可能有多组）及计数方式。"""
    level1 = [cn_int(m.group(1)) for _, t in draft.paragraphs for m in [HEADING1_RE.match(t)] if m]
    if level1:
        return [len(set(level1))], "一级标题“一、”"
    level2 = [cn_int(m.group(1)) for _, t in draft.paragraphs for m in [HEADING2_RE.match(t)] if m]
    if level2:
        return [len(set(level2))], "二级标题“（一）”"
    runs: List[int] = []
    current = 0
    for match in SHI_RE.finditer(draft.body_text):
        value = cn_int(match.group(1))
        if value == 1:
            if current:
                runs.append(current)
            current = 1
        elif current and value == current + 1:
            current = value
    if current:
        runs.append(current)
    if runs:
        return runs, "“一是……二是……”"
    arabic = [int(m.group(1)) for _, t in draft.paragraphs for m in [ARABIC_ITEM_RE.match(t)] if m]
    if arabic:
        return [len(set(arabic))], "段首序号“1.”"
    return [], ""


def check_title_numbers(draft: Union["Draft", "Section"]) -> List[dict]:
    if not draft.titles:
        return []
    summaries = title_summaries(draft.title_text)
    if not summaries:
        return []
    counts, how = item_counts(draft)
    line = draft.titles[0][0]
    issues: List[dict] = []
    for value, text in summaries:
        if not counts:
            issues.append(issue(
                "BRF-002", "P2", "标题", line, text,
                f"标题用了数字概括“{text}”，正文没有可计数的分项（一、（一）、一是、1.）",
                "人工核对正文是否正好写了这么多项；对不上的改标题，不为凑数补写做法",
                "gongwen-brief SKILL.md 第四步；数字概括与实际条数一致",
            ))
        elif value not in counts:
            shown = "、".join(str(c) for c in counts)
            issues.append(issue(
                "BRF-002", "P1", "标题", line, text,
                f"标题概括为 {value} 项，正文按{how}计为 {shown} 项",
                "按正文实际条数改标题里的数字，或去掉数字概括；不为凑数补写做法",
                "gongwen-brief SKILL.md 第四步；数字概括与实际条数一致",
            ))
    return issues


def iter_quotes(draft: Draft):
    """逐个给出引号内容：(行号, 段落, 起点, 终点, 内容, 是否有说话人)。"""
    for number, text in draft.titles + draft.paragraphs:
        marker_spans = [m.span() for m in MARKER_RE.finditer(text)]
        for match in QUOTE_RE.finditer(text):
            # 【待补】【待核实】标记内引用用户原话的引号是说明文字，不当作稿中引语
            if any(lo <= match.start() and match.end() <= hi for lo, hi in marker_spans):
                continue
            content = next(g for g in match.groups() if g is not None)
            before = text[max(0, match.start() - 30):match.start()]
            after = text[match.end():match.end() + 16]
            if NAME_INTRO_RE.search(before):
                has_speaker = False
            else:
                has_speaker = bool(
                    SPEAKER_BEFORE_RE.search(before) or COLON_BEFORE_RE.search(before) or SPEAKER_AFTER_RE.match(after)
                )
            yield number, text, match.start(), match.end(), content, has_speaker


def is_speech(content: str, has_speaker: bool) -> bool:
    """判断引号内是原话（引语）还是名称、提法。"""
    if not normalize(content):
        return False
    if has_speaker or SPEECH_PUNCT_RE.search(content):
        return True
    return len(normalize(content)) >= 6 and bool(SPEECH_WORD_RE.search(content))


def speaker_of(before: str, after: str) -> str:
    """从引号前后取说话人短语，如“村民王某说：”取“村民王某”；取不到或是代词的返回空串。"""
    who = ""
    match = SPEAKER_BEFORE_RE.search(before)
    if match:
        who = SPEAKER_SPLIT_RE.split(before[:match.start()])[-1][-10:]
    else:
        match = SPEAKER_AFTER_RE.match(after)
        if match:
            who = match.group("who")
    who = normalize(who)
    if len(who) < 2 or SPEAKER_SKIP_RE.match(who):
        return ""
    return who


def speaker_tokens(who: str) -> List[str]:
    """说话人的比对用词：全称、“王某”式化名、末尾 2～3 字（多为职务或称谓）。"""
    tokens = {who, who[-2:]}
    if len(who) >= 3:
        tokens.add(who[-3:])
    tokens.update(SPEAKER_NAME_RE.findall(who))
    return [t for t in tokens if len(t) >= 2]


def quoted_in_sources(needle: str, sources: Sequence[str]) -> bool:
    """材料中引号内、或“某某：”记录格式冒号后有这段话，才算材料中的原话。"""
    for source in sources:
        for match in QUOTE_RE.finditer(source):
            content = next(g for g in match.groups() if g is not None)
            if needle in normalize(content):
                return True
        for line in source.split("\n"):
            head, sep, tail = line.partition("：")
            if not sep:
                head, sep, tail = line.partition(":")
            if sep and len(head.strip()) <= 12 and needle in normalize(tail):
                return True
    return False


def speaker_status(needle: str, who: str, haystack: Sequence[str]) -> str:
    """说话人核对结果：near（命中处附近有）、far（材料别处有）、none（材料中没有）。"""
    tokens = speaker_tokens(who)
    anywhere = False
    for h in haystack:
        if any(t in h for t in tokens):
            anywhere = True
        start = h.find(needle)
        while start != -1:
            window = h[max(0, start - SPEAKER_WINDOW_BEFORE):start + len(needle) + SPEAKER_WINDOW_AFTER]
            if any(t in window for t in tokens):
                return "near"
            start = h.find(needle, start + 1)
    return "far" if anywhere else "none"


def check_quotes(draft: Draft, sources: Optional[List[str]]) -> List[dict]:
    issues: List[dict] = []
    quotes = [q for q in iter_quotes(draft) if normalize(q[4])]
    if not quotes:
        return issues
    if sources is None:
        issues.append(issue(
            "BRF-003", "P2", "引语", quotes[0][0], quotes[0][4],
            f"稿中有 {len(quotes)} 处引号内容，未提供材料文件，无法自动溯源",
            "用 --source 传入用户材料再查一次；或逐条对照材料原文人工核对",
            "references/_shared/guardrails.md 第二条第 1 款",
        ))
        return issues
    haystack = [normalize(s) for s in sources]
    for number, text, start, end, content, has_speaker in quotes:
        needle = normalize(content)
        before = text[max(0, start - 30):start]
        after = text[end:end + 16]
        found = any(needle in h for h in haystack)
        speech = is_speech(content, has_speaker)
        if not found:
            if speech:
                issues.append(issue(
                    "BRF-003", "P0", "引语", number, content,
                    "引语在用户材料中找不到原文",
                    "删去，或改为【待补：某某原话，须本人同意引用】；节选只能取材料中连续的一段，不改字、不拼接",
                    "references/_shared/guardrails.md 第二条第 1、5 款",
                ))
            elif HONOR_INTRO_RE.search(before):
                issues.append(issue(
                    "BRF-003", "P1", "引语", number, content,
                    "荣誉称号在用户材料中找不到出处",
                    "核对授予文件或通报，照原名称写；材料没有的删去，或改为【待核实：荣誉名称和授予单位】",
                    "references/_shared/guardrails.md 第二条第 1、4 款",
                ))
            else:
                issues.append(issue(
                    "BRF-003", "P2", "引语", number, content,
                    "引号内的提法或名称在用户材料中找不到",
                    "核对是否为材料中的正式名称或提法；不是的去掉引号，或标【待核实】",
                    "references/_shared/guardrails.md 第二条第 4 款",
                ))
            continue
        if not speech:
            continue
        # 材料中找得到的原话，再核对两件事：短引语是不是材料里的原话；说话人对不对得上
        problems: List[str] = []
        level = "P2"
        if len(needle) < SHORT_QUOTE_LEN and not quoted_in_sources(needle, sources):
            problems.append("短引语在材料中只出现在叙述里，不是引号内或记录中的原话，可能是从转述中截取的")
            level = "P1"
        who = speaker_of(before, after)
        if who:
            status = speaker_status(needle, who, haystack)
            if status == "none":
                problems.append(f"说话人“{who}”在材料中找不到")
                level = "P1"
            elif status == "far":
                problems.append(f"说话人“{who}”在材料中与这段话相距较远")
        if problems:
            issues.append(issue(
                "BRF-003", level, "引语", number, content,
                "；".join(problems) + "，需人工核对",
                "对照材料核对原话和说话人：材料是转述的，去掉引号改为转述；说话人对不上的照材料改，"
                "材料没有记录谁说的，改为【待补：某某原话，须本人同意引用】",
                "references/_shared/guardrails.md 第二条第 1、5 款；references/ledger.md 第三部分",
            ))
    return issues


def check_speakers(draft: Draft) -> List[dict]:
    issues: List[dict] = []
    for number, _text, _start, _end, content, has_speaker in iter_quotes(draft):
        if has_speaker or not is_speech(content, has_speaker):
            continue
        issues.append(issue(
            "BRF-004", "P2", "引语", number, content,
            "引语没有交代说话人",
            "写明说话人身份（如“居民张某”“社区负责人”），说话人只来自材料",
            "references/ledger.md 第三部分",
        ))
    return issues


def check_lead(draft: Draft, doc_type: str) -> List[dict]:
    if doc_type == "简报":
        return []
    lead = draft.lead()
    if lead is None:
        return []
    number, text = lead
    issues: List[dict] = []
    if PADDING_RE.match(text):
        issues.append(issue(
            "BRF-005", "P1", "导语", number, text[:40],
            "导语先写铺垫，结论被推到后面",
            "第一句直接写时间、主体、做法和效果；铺垫删去，确需交代依据的压成半句",
            "《整治形式主义为基层减负若干规定》第2条；references/_shared/rules/doc-types.md 4.10",
        ))
    if doc_type in ("信息", "新闻稿") and not TIME_RE.search(text):
        issues.append(issue(
            "BRF-006", "P2", "导语", number, text[:40],
            "导语没有交代时间，时效无从判断",
            "写明“近日”“10月12日”“今年以来”等时间；材料没有的标【待补：时间】",
            "references/types/ 对应文件“导语”部分",
        ))
    return issues


def scan(draft: Draft, regex: "re.Pattern[str]", include_titles: bool = True):
    pool = (draft.titles if include_titles else []) + draft.paragraphs
    for number, text in pool:
        for match in regex.finditer(text):
            yield number, text, match


EFFICACY_BASIS = "《中华人民共和国广告法》第十六条第一款第（一）、（二）项（参照）；references/_shared/profiles/units/yiyuan.md"
PHRASE_TAIL_RE = re.compile(r"[^\s，。；：、！？,.;:!?“”‘’（）()【】]{0,4}")


def is_medical(draft: Draft, number: int, text: str, medical_unit: bool) -> bool:
    """同段出现医疗词语，或 --unit 为医院时，按医疗语境处理；标题行看标题和导语。"""
    if medical_unit:
        return True
    if any(n == number for n, _ in draft.titles):
        lead = draft.lead()
        text = draft.title_text + (lead[1] if lead else "")
    return bool(MEDICAL_RE.search(text))


def check_efficacy(draft: Draft, external: bool, medical_unit: bool = False,
                   sources: Optional[List[str]] = None) -> List[dict]:
    issues: List[dict] = []
    haystack = [normalize(s) for s in sources] if sources else []
    for number, text, match in scan(draft, EFFICACY_MED_RE):
        word = match.group(0)
        issues.append(issue(
            "BRF-008", "P0" if external else "P1", "宣传用语", number, word,
            f"“{word}”属于疗效保证",
            "删去；改写为具体病例数、随访结果等可核实的事实，并以材料为准",
            EFFICACY_BASIS,
        ))
    excepted = {(n, m.start()) for n, _t, m in scan(draft, EFFICACY_EXCEPT_RE)}
    for number, text, match in scan(draft, EFFICACY_CTX_RE):
        word = match.group(0)
        if (number, match.start()) in excepted:
            continue
        if word == "有效率" and not (RATE_BEFORE_RE.search(text[:match.start()])
                                    or RATE_AFTER_RE.match(text[match.end():])):
            continue
        if not is_medical(draft, number, text, medical_unit):
            issues.append(issue(
                "BRF-008", "P2", "宣传用语", number, word,
                f"“{word}”是效果保证类用语，所在段落不是医疗语境",
                "工作提法原文（如“根治欠薪”）照用；不是提法而是对效果的断言的，改写为可核实的事实",
                "references/external-check.md 第二部分第 6 项",
            ))
            continue
        level = "P0" if external else "P1"
        message = f"“{word}”在医疗语境中属于疗效保证"
        phrase = normalize(word + PHRASE_TAIL_RE.match(text, match.end()).group(0))
        if any(phrase in h for h in haystack):
            level = "P1" if level == "P0" else "P2"
            message += "；材料中有同样说法，核对是否为提法原文（如专项行动名称）"
        issues.append(issue(
            "BRF-008", level, "宣传用语", number, word, message,
            "确属疗效保证的删去，改写为具体病例数、随访结果等可核实的事实；是提法原文的照用并在起草说明注明出处",
            EFFICACY_BASIS,
        ))
    return issues


def check_claims(draft: Draft, external: bool, medical_unit: bool = False,
                 sources: Optional[List[str]] = None) -> List[dict]:
    issues: List[dict] = []
    for number, _text, match in scan(draft, ABSOLUTE_RE):
        word = match.group(0)
        level = "P2" if word == "首例" or not external else "P1"
        issues.append(issue(
            "BRF-007", level, "宣传用语", number, word,
            f"“{word}”须有认定依据（评审结果、排名通报、鉴定意见等）",
            "用户提供了依据的，在起草说明写明出处；没有的删去，或改为【待核实：认定依据】",
            "references/external-check.md 第二部分；《中华人民共和国广告法》第九条第（三）项（参照）",
        ))
    issues += check_efficacy(draft, external, medical_unit, sources)
    for number, text, match in scan(draft, EDU_HYPE_RE):
        word = match.group(0)
        before = text[max(0, match.start() - 10):match.start()]
        after = text[match.end():match.end() + 6]
        if is_edu_critical(before, after):
            issues.append(issue(
                "BRF-009", "P2", "宣传用语", number, word,
                f"“{word}”出现在批评性语境中（如“片面追求升学率”“按升学率考核”）",
                "核对是反映问题而不是宣传的，照常保留；对外稿避免出现具体学校的升学数据",
                "教育部关于严禁宣传炒作中高考状元和升学率的有关要求（以当年文件为准）",
            ))
            continue
        issues.append(issue(
            "BRF-009", "P1" if external else "P2", "宣传用语", number, word,
            f"“{word}”属于宣传升学率、状元一类内容" + ("" if external else "；内部报送稿只作提示"),
            "对外稿删去，改写学生综合发展、教育教学改进等材料中有的事实；内部报送的统计数据核对口径后保留",
            "教育部关于严禁宣传炒作中高考状元和升学率的有关要求（以当年文件为准）",
        ))
    return issues


def is_edu_critical(before: str, after: str) -> bool:
    """“片面追求升学率”“按升学率考核教师”一类批评性说法，不是宣传。"""
    if EDU_CRITICAL_BEFORE_RE.search(before) or EDU_CRITICAL_AFTER_RE.match(after):
        return True
    return bool(EDU_RANK_BEFORE_RE.search(before) and EDU_RANK_AFTER_RE.match(after))


def check_title_words(titles: Sequence[Tuple[int, str]]) -> List[dict]:
    issues: List[dict] = []
    for number, text in titles:
        for match in EMPTY_TITLE_RE.finditer(text):
            issues.append(issue(
                "BRF-016", "P1", "标题", number, match.group(0),
                f"标题用了空词“{match.group(0)}”，放在任何单位的任何工作上都成立",
                "标题写主体和做法（或进展），有成效数据的写数据；没有成效数据的不为标题补效果",
                "references/types/xinxi.md 第三部分第 4 条",
            ))
    return issues


def check_tone(draft: Draft) -> List[dict]:
    issues: List[dict] = []
    for number, text in draft.titles:
        for match in TITLE_BAIT_RE.finditer(text):
            issues.append(issue(
                "BRF-010", "P1", "语体", number, match.group(0),
                "标题党用语，单位官方稿件不用",
                "标题写清是什么事、和读者有什么关系",
                "references/types/gongzhonghao.md 第二部分；国办发〔2018〕123号 三、（一）",
            ))
    for number, _text, match in scan(draft, SLANG_RE):
        issues.append(issue(
            "BRF-010", "P1", "语体", number, match.group(0),
            "自媒体腔、网络流行语，单位官方稿件不用",
            "改为通俗、规范的说法",
            "references/types/gongzhonghao.md 第五部分",
        ))
    # 引号内是照录的原话，感叹号不计入、不改
    unquoted = [(n, QUOTE_RE.sub("", t)) for n, t in draft.paragraphs]
    for number, text in unquoted:
        for match in BODY_BANG_RE.finditer(text):
            issues.append(issue(
                "BRF-010", "P1", "语体", number, match.group(0),
                "正文连续感叹号或问叹连用，单位官方稿件不用",
                "陈述事实用句号；确需感叹的只用一个感叹号",
                "references/types/gongzhonghao.md 第五部分",
            ))
    count = sum(len(EXCLAIM_RE.findall(t)) for _, t in unquoted)
    if count > EXCLAIM_LIMIT:
        first = next((n for n, t in unquoted if EXCLAIM_RE.search(t)), None)
        issues.append(issue(
            "BRF-011", "P2", "语体", first, f"感叹号 {count} 处",
            f"引语之外用了 {count} 处感叹号",
            f"陈述事实用句号；全文最多保留 {EXCLAIM_LIMIT} 处（邀请、提示类结尾句，如“欢迎各位居民前来就餐！”），引语中的照录不改",
            "references/types/gongzhonghao.md 第五部分",
        ))
    return issues


def check_brief_format(draft: Draft) -> List[dict]:
    issues: List[dict] = []
    head = "\n".join(t for _, t in draft.header) or "\n".join(t for _, t in draft.nonempty[:8])
    if not ISSUE_NO_RE.search(head):
        issues.append(issue(
            "BRF-012", "P1", "简报格式", None, "报头",
            "报头没有期号（第×期）",
            "按本单位顺序编号写“第×期”，期号不知道的写【待补：期号】",
            "references/types/jianbao.md 第二部分",
        ))
    if not (EDITOR_RE.search(head) and DATE_RE.search(head)):
        issues.append(issue(
            "BRF-013", "P2", "简报格式", None, "报头",
            "报头缺编印单位或印发日期",
            "期号下一行写编印单位和日期，日期不知道的写【待补：印发日期】",
            "references/types/jianbao.md 第二部分",
        ))
    if not any(DISTRIBUTION_RE.match(t) for _, t in draft.footer + draft.nonempty[-6:]):
        issues.append(issue(
            "BRF-014", "P2", "简报格式", None, "报尾",
            "报尾没有报送范围（报：送：发：）",
            "按用户给的范围写；没给的写【待补：报送范围】",
            "references/types/jianbao.md 第四部分",
        ))
    return issues


def check_photos(draft: Draft, external: bool) -> List[dict]:
    if not external:
        return []
    for number, _text, match in scan(draft, PHOTO_RE):
        return [issue(
            "BRF-015", "P2", "授权", number, match.group(0),
            "对外稿涉及照片或视频",
            "发布前确认肖像授权：群众、患者须本人同意，不满十四周岁的学生须监护人同意",
            "《中华人民共和国民法典》第一千零一十九条；《中华人民共和国个人信息保护法》第二十八条、第三十一条",
        )]
    return []


# ---------------------------------------------------------------------------
# 汇总与输出
# ---------------------------------------------------------------------------

def normalize_type(raw: str) -> str:
    raw = (raw or "信息").strip()
    raw = TYPE_ALIASES.get(raw, raw)
    if raw not in TYPES:
        raise InputError(f"不支持的稿件类型：{raw}。可选：{'、'.join(TYPES + tuple(ALIAS_RANGES))}。")
    return raw


def build_report(path: Path, doc_type: str = "信息", sources: Optional[Sequence[Path]] = None,
                 external: bool = False, low: Optional[int] = None, high: Optional[int] = None,
                 unit: Optional[str] = None) -> dict:
    alias = (doc_type or "").strip()
    doc_type = normalize_type(doc_type)
    lines = read_text_lines(path)
    draft = Draft(body_lines(lines), doc_type)
    if not draft.nonempty:
        raise InputError("正文为空（【起草说明】之前没有内容）。")
    external = external or doc_type in EXTERNAL_TYPES
    default_low, default_high = ALIAS_RANGES.get(alias, DEFAULT_RANGES[doc_type])
    # 用户只给“不超过N字”（只传 --max）时，N 是上限，不再套用默认下限
    low = (None if high is not None else default_low) if low is None else low
    high = default_high if high is None else high
    if low is not None and high is not None and low > high:
        raise InputError("--min 不能大于 --max。")
    source_texts: Optional[List[str]] = None
    if sources:
        source_texts = ["\n".join(read_text_lines(Path(s))) for s in sources]

    issues: List[dict] = []
    medical_unit = bool(unit and MEDICAL_UNIT_RE.search(unit))
    issues += check_length(draft, low, high, doc_type)
    if doc_type == "简报":
        for section in draft.sections():
            issues += check_title_numbers(section)
            issues += check_title_words(section.titles)
    else:
        issues += check_title_numbers(draft)
        issues += check_title_words(draft.titles)
    issues += check_quotes(draft, source_texts)
    issues += check_speakers(draft)
    issues += check_lead(draft, doc_type)
    issues += check_claims(draft, external, medical_unit, source_texts)
    issues += check_tone(draft)
    if doc_type == "简报":
        issues += check_brief_format(draft)
    issues += check_photos(draft, external)

    order = {"P0": 0, "P1": 1, "P2": 2}
    issues.sort(key=lambda it: (order[it["level"]], it["line"] or 0, it["id"]))
    summary = {level: sum(1 for it in issues if it["level"] == level) for level in ("P0", "P1", "P2")}
    notes = [
        "篇幅口径：标题、署名行" + ("、报头和报尾" if doc_type == "简报" else "") + "之外的正文，不含空白，至【起草说明】为止。",
        "本脚本只做专项检查，标点、数字、隐私、AI 腔等请另运行 gongwen_lint.py。",
    ]
    if sources is None:
        notes.append("未提供 --source，引语未做溯源。")
    return {
        "file": path.name,
        "type": doc_type,
        "external": external,
        "char_count": draft.char_count(),
        "length_range": [low, high],
        "title": draft.title_text,
        "summary": summary,
        "issues": issues,
        "notes": notes,
    }


def render_text(report: dict) -> str:
    low, high = report["length_range"]
    span = f"{low or '不限'}～{high or '不限'}"
    lines = [
        f"文件：{report['file']}　类型：{report['type']}　{'对外发布' if report['external'] else '内部报送'}",
        f"正文约 {report['char_count']} 字（常用范围 {span} 字）",
        f"问题：P0 {report['summary']['P0']} 条，P1 {report['summary']['P1']} 条，P2 {report['summary']['P2']} 条",
    ]
    for item in report["issues"]:
        where = f"第 {item['line']} 行" if item["line"] else "全文"
        lines.append(f"[{item['level']}] {item['id']} {item['category']}｜{where}｜{item['excerpt']}")
        lines.append(f"    问题：{item['message']}")
        lines.append(f"    建议：{item['suggestion']}")
        lines.append(f"    依据：{item['basis']}")
    if not report["issues"]:
        lines.append("未发现专项问题。")
    lines.extend(f"说明：{note}" for note in report["notes"])
    return "\n".join(lines)


class ChineseHelpFormatter(argparse.HelpFormatter):
    """把 argparse 的 usage 前缀改为中文。"""

    def add_usage(self, usage, actions, groups, prefix=None):
        return super().add_usage(usage, actions, groups, "用法：" if prefix is None else prefix)


def make_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        prog="brief_check.py",
        description="信息、简报、新闻稿、公众号推文的专项检查（篇幅、标题数字、引语溯源、导语、宣传用语、简报要素）。",
        formatter_class=ChineseHelpFormatter,
        add_help=False,
    )
    parser._optionals.title = "选项"
    parser._positionals.title = "参数"
    parser.add_argument("-h", "--help", action="help", help="显示本帮助信息并退出")
    parser.add_argument("draft", help="稿件文件（.txt、.md、.docx），只检查【起草说明】之前的正文")
    parser.add_argument("--type", default="信息", metavar="类型", help="稿件类型：信息、简报、新闻稿、公众号，另可用紧急信息、社区信息（默认：信息）")
    parser.add_argument("--source", nargs="+", metavar="材料", help="用户提供的材料文件，用于引语溯源")
    parser.add_argument("--external", action="store_true", help="对外发布（新闻稿、公众号默认按对外处理）")
    parser.add_argument("--unit", metavar="单位类型", help="单位类型，如 公立医院；为医院时全文按医疗语境检查疗效保证")
    parser.add_argument("--min", type=int, dest="low", metavar="字数", help="篇幅下限，覆盖默认值")
    parser.add_argument("--max", type=int, dest="high", metavar="字数", help="篇幅上限，覆盖默认值")
    parser.add_argument("--json", action="store_true", help="以 JSON 格式输出")
    parser.add_argument("--strict", action="store_true", help="存在 P0 时返回 1")
    return parser


def main(argv: Optional[Sequence[str]] = None) -> int:
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8")  # type: ignore[attr-defined]
        except Exception:
            pass
    parser = make_parser()
    try:
        args = parser.parse_args(argv)
    except SystemExit as exc:
        return 0 if exc.code == 0 else 2
    try:
        report = build_report(
            Path(args.draft), args.type,
            [Path(s) for s in args.source] if args.source else None,
            args.external, args.low, args.high, args.unit,
        )
    except InputError as exc:
        print(f"输入错误：{exc}", file=sys.stderr)
        return 2
    if args.json:
        print(json.dumps(report, ensure_ascii=False, indent=2))
    else:
        print(render_text(report))
    if args.strict and report["summary"]["P0"]:
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
