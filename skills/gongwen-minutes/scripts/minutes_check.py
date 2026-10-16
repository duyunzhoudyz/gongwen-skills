# -*- coding: utf-8 -*-
"""会议纪要议定事项检查与督办清单对账（gongwen-minutes）。

用法：
    python scripts/minutes_check.py <纪要正文> [--list 督办清单.md|.xlsx] [--emit-list 输出.md] [--json] [--strict]

做三件事：
  1. 识别纪要中的议定事项，逐条检查：责任单位、完成时限是否写明或已标【待补】；
     "原则同意"是否写了修改要求；"暂缓"是否写了原因或再议条件。另查口语残留、转写标签残留、出席名单。
  2. --list：与督办清单逐条对账。清单中的事项能否在纪要中找到；责任单位、配合单位、完成时限是否与纪要一致；
     纪要中需要落实的事项是否漏列。比对的是从该条识别出的牵头（负责）单位、配合单位和时限，不是"在原文中出现过"：
     牵头与配合对调、时限挂到同一条中其他单位的任务上，都报不一致。一条中有几个"××单位＋时限"的，按子句配对。
  3. --emit-list：按识别结果生成督办清单草表（Markdown），供人工逐行核对后使用。一条中有几个单位各有任务的拆成几行；
     任职、审议通过等无需落实的结论不填【待补】，备注提示确认是否列入。输出文件已存在时加 -v2、-v3，不覆盖。

议定事项的识别范围：
  - "会议议定：""会议决定：""会议明确："等单独成行时，其后的各段（到下一个标题、出席名单或其他"会议……"提示语为止）；
  - 以"会议议定：……""会议决定：……"开头、内容写在同一段的，段内"一是……二是……"拆成多条；
  - 以"会议同意""会议原则同意""会议决定""会议暂缓"等开头的段落；
  - 议题标题（"一、关于……"）下没有"会议议定："时，以"××事项暂缓""××另行研究"之类开头的段落。
  "会议听取了""会议认为""会议指出""会议强调""会议要求"引出的段落不算议定事项。
  "同意暂不……"按同意处理；任职、审议通过、维持现状等没有写责任单位和时限的结论视为无需落实，不报缺项。
  一段中有几类结论的（如"会议同意……；爽约限制暂缓实施"），按句号、分号拆成几条，各自判定结论类型；
  暂缓、不同意、另行研究的条目中以"由××……"开头的准备任务（如"由财务科先行测算"）单列一条。

只检查【起草说明】之前的正文。只依赖标准库；读取 .docx 需要 python-docx，读取 .xlsx 需要 openpyxl。
退出码：0 正常；--strict 且存在 P0 时返回 1；输入错误返回 2。
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Tuple

DRAFT_NOTE_RE = re.compile(r"^\s*(?:#{1,6}\s*)?(?:(?:\*\*|__)\s*)?【起草说明】")
TEXT_SUFFIXES = (".txt", ".md")
LEVEL_ORDER = ("P0", "P1", "P2")

# ---------------------------------------------------------------------------
# 正则
# ---------------------------------------------------------------------------

CN_NUM = "一二三四五六七八九十"
# 层次序号：一、（一）1.（1）
LABEL_RE = re.compile(
    rf"^\s*(?P<label>[{CN_NUM}]{{1,3}}、|[（(][{CN_NUM}]{{1,3}}[)）]|\d{{1,2}}[.．、]|[（(]\d{{1,2}}[)）])\s*"
)
H1_PREFIX_RE = re.compile(rf"^\s*([{CN_NUM}]{{1,3}}、)")
ATTEND_RE = re.compile(r"^\s*(?:出\s*席|请\s*假|列\s*席|缺\s*席)\s*[：:]")
# "经研究，""经认真讨论，""经党委会前置研究，"
LEAD_IN = r"^\s*(?:经[^，,。；;：:]{0,20}?(?:研究|讨论|审议)[，,]\s*)?"
# "会议议定："等单独成行，引出下面的各条
BLOCK_START_RE = re.compile(
    LEAD_IN + r"会议(?:研究|审议)?(?:议定|决定|明确|确定)(?:了)?(?:如下|以下)?(?:事项)?(?:如下)?[：:]\s*$"
)
# "会议议定：……"内容在同一段
INLINE_RE = re.compile(
    LEAD_IN + r"会议(?:研究|审议)?(?:议定|决定|明确|确定)(?:如下|以下事项)?[：:，,]\s*(?P<content>\S.*)$"
)
# "会议同意……""会议原则同意……""会议决定暂缓……"开头的段落
SENTENCE_RE = re.compile(
    LEAD_IN + r"会议(?:研究|审议|讨论)?(?:并|后)?(?:原则上?|一致)?"
    r"(?:同意|通过|不同意|不予同意|决定|议定|明确|确定|暂缓|审议通过)"
)
# "会议听取了……，原则同意……""关于……事项。会议同意……"：结论写在过程交代之后
PROCESS_CUE_RE = re.compile(r"^\s*(?:会议(?:听取|研究了|审议了|讨论了)|关于)")
MID_DECISION_RE = re.compile(
    r"[。；;，,]\s*(?=(?:会议)?(?:研究|审议)?(?:后)?(?:原则上?|一致)?(?:同意|不同意|不予同意|决定|议定|暂缓|明确))"
)
# "二、议定事项""（二）会议要求"之类的标题，下面各条按议定事项或会议要求识别
DECISION_HEADING_RE = re.compile(r"议定事项|决定事项|会议决定|会议议定")
REQ_HEADING_RE = re.compile(r"会议要求|工作要求")
# "会议要求："单独成行，或"会议要求，……"
REQ_BLOCK_RE = re.compile(r"^\s*会议(?:还|同时|进一步)?要求[：:]\s*$")
REQ_INLINE_RE = re.compile(r"^\s*会议(?:还|同时|进一步)?要求[：:，,]?\s*(?P<content>\S.*)$")
# 其他提示语：遇到时结束"会议议定："引出的块
OTHER_CUE_RE = re.compile(r"^\s*会议(?:听取|认为|指出|强调|号召|提出|传达|学习|讨论了|研究了|审议了)")
DECISION_WORD_RE = re.compile(r"同意|决定|暂缓|明确|另行研究|由[一-鿿]")
YISHI_RE = re.compile(rf"[{CN_NUM}]是")
# 识别责任单位前去掉的段首结论词，避免把"同意××科"整体当成单位名称
LEAD_DECISION_RE = re.compile(
    r"^\s*(?:会议)?(?:研究|审议|讨论)?(?:后)?(?:原则上?|一致)?(?:同意|通过|决定|议定|明确|确定)[：:，,]?\s*"
)

ORG_SUFFIX = (
    "办公室|委员会|居委会|村委会|管委会|领导小组|工作专班|年级组|教研组|中心|小组|专班|学院|街道|社区|公司|集团|部门|单位|"
    "院长|校长|书记|主任|局长|经理|总监|处长|科长|部长|镇长|站长|所长|"
    "科|处|部|室|局|委|办|组|院|系|队|所|站|镇|乡|村|厅|署|会"
)
NAME = rf"[\u4e00-\u9fffA-Za-z0-9]{{1,15}}?(?:{ORG_SUFFIX})"
NAMES = rf"{NAME}(?:[、和及与]{NAME})*"
OWNER_VERBS = (
    "牵头|负责|承办|组织|会同|落实|办理|实施|具体|统筹|研究|测算|修改|起草|制定|完善|开展|做好|对接|协调|压缩|"
    "拟定|编制|会商|梳理|核实|调整|报送|提交|印发|启动|推进|跟进|准备|采购|招标|整理|按|在|于|对|向|进一步|抓紧|尽快|结合|"
    "[，,。；;]"
)
OWNER_PATTERNS = (
    re.compile(rf"由(?P<o>{NAMES})(?=(?:{OWNER_VERBS}))"),
    re.compile(r"(?:责任单位|牵头单位|责任部门|牵头部门|承办单位|责任科室|牵头科室|责任领导)[：:为是]\s*(?P<o>[^，,。；;\s]{2,30})"),
    re.compile(rf"(?:^|[，,。；;：:])(?P<o>{NAMES})(?=牵头|负责)"),
    # "具体措施各子公司10月底前报财务部"
    re.compile(rf"(?P<o>各{NAME})(?=\d|于|在|报|向|要|按|结合|对)"),
    # "，财务科11月底前完成测算"
    re.compile(rf"(?:^|[，,。；;：:])(?P<o>{NAMES})(?=(?:\d{{4}}年|今年|明年)?\d{{1,2}}月|于\d|月底|年底|本月|本周|下周)"),
    re.compile(rf"由(?P<o>{NAMES})"),
)
COOP_PATTERNS = (
    re.compile(rf"(?:^|[，,。；;：:、])(?P<c>{NAMES})(?=配合|协助|参与)"),
    re.compile(rf"会同(?P<c>{NAMES})"),
)
OWNER_MARK_RE = re.compile(r"【待(?:补|核实)[：:][^】]*(?:责任|牵头|单位|部门|科室|承办)[^】]*】")

CN_DIGITS = "\\d一二两三四五六七八九十"
YEAR_PREFIX = r"(?:\d{4}年|今年|明年|当年|本年)?"
# "今年12月开工""12月份启动"：月份后直接跟落实动作的，也是时限
DEADLINE_ACTIONS = (
    "开工|启动|实施|完成|上线|印发|开始|投用|投入使用|建成|竣工|召开|报送|交付|到位|执行|运行|落实|"
    "出台|发布|签订|验收|整改|开诊|开放|招标|提交"
)
DEADLINE_PATTERNS = (
    re.compile(YEAR_PREFIX + r"\d{1,2}月\d{1,2}日(?:前|之前|以前|起|内)?"),
    re.compile(YEAR_PREFIX + r"\d{1,2}月(?:底|末|上旬|中旬|下旬|份)(?:前|之前|以前)?"),
    re.compile(YEAR_PREFIX + r"\d{1,2}月(?:前|之前|以前)"),
    re.compile(YEAR_PREFIX + rf"\d{{1,2}}月(?=(?:内|中|起)?(?:{DEADLINE_ACTIONS}))"),
    re.compile(r"(?:\d{4}年|今年|明年|本年度?)?(?:年底|年末|年内)(?:前|之前)?"),
    re.compile(r"(?:本|下|这个|下个)(?:月|季度|周)(?:底|末|内)?(?:前|之前)?"),
    # 不带数字的"月底前""月末""月中旬前"
    re.compile(r"(?<![\d一二两三四五六七八九十个])月(?:底|末|上旬|中旬|下旬)(?:前|之前|以前)?"),
    re.compile(r"(?<![\d一二两三四五六七八九十个])月(?:初|中)(?:前|之前|以前)"),
    # "周五前""本周三""下星期一之前"
    re.compile(r"(?:本|下|这)?(?:周|星期)[一二三四五六日天](?:前|之前|以前)?"),
    re.compile(r"下次(?:会议|例会|办公会|会)(?:前|之前)?"),
    re.compile(r"(?:节后|会后|开学后|假期后|春节后|国庆后)第[一二三四五12345]周"),
    re.compile(rf"[{CN_DIGITS}]+个?(?:工作日|日|天|周|星期|个月|月)(?:内|之内)"),
    re.compile(r"每(?:月|季度|季|周|半年|年|天|日)"),
    re.compile(r"(?:长期坚持|长期|立即|即日起|立行立改)"),
    re.compile(r"(?:第[一二三四]季度|[一二三四]季度|上半年|下半年)(?:内|前|末)?"),
    re.compile(r"(?:春节|元旦|开学|寒假|暑假|国庆)(?:前|之前)"),
)
DEADLINE_MARK_RE = re.compile(r"【待(?:补|核实)[：:][^】]*(?:时限|时间|日期|期限|完成)[^】]*】")
VAGUE_TIME_RE = re.compile(r"尽快|抓紧|及时|适时|择机|近期")
MARK_RE = re.compile(r"【待(?:补|核实)[：:][^】]*】")

REVISION_HINT_RE = re.compile(
    r"修改|完善|调整|补充|压缩|增加|减少|删去|删除|核对|核实|细化|优化|重新|进一步|按照?会议|会议提出|会议意见|【待(?:补|核实)"
)
POSTPONE_HINT_RE = re.compile(r"待|等|因|鉴于|由于|条件|之后|以后|后再|再议|再提交|再研究|【待(?:补|核实)")

COLLOQUIAL_RE = re.compile(r"嗯|呃|哎呀|那个那个|就是说|对吧|是吧|我觉得|我个人|咱们|这个事儿|搞一下|弄一下|整一下")
SPEAKER_TAG_RE = re.compile(r"(?:发言人|说话人)\s*\d+")
QUOTE_RE = re.compile(r"“[^”]*”|‘[^’]*’|《[^》]*》")

PUNCT_RE = re.compile(r"[\s，,。．.；;：:、！!？?“”‘’\"'（）()《》〈〉【】\[\]…—\-·/|]")
YEAR_BEFORE_MONTH_RE = re.compile(r"(?:\d{4}年|今年|当年|本年)(?=\d{1,2}月)")
DEADLINE_SUFFIX_RE = re.compile(r"(?:之前|以前|前|之内|内)$")

DECISION_EXEMPT = ("暂缓", "不同意", "另行研究")


class InputError(Exception):
    """输入文件不存在、格式不支持或读取失败。"""


@dataclass
class Item:
    no: int
    line: int
    position: str
    section: str
    text: str
    decision: str
    kind: str = "议定"          # 议定：议定事项；要求：会议要求中的任务（只参与督办清单对账）
    owner: str = ""
    owner_marked: bool = False
    deadline: str = ""
    deadline_marked: bool = False
    deadline_vague: bool = False
    cooperate: str = ""
    no_follow: str = ""         # 无需落实的结论类别：任职、审议通过、维持现状；空串表示需要落实


@dataclass
class Task:
    """一条议定事项中某个单位承担的任务：责任单位、该单位的时限、所属子句。"""
    owner: str
    deadlines: List[str] = field(default_factory=list)
    clauses: List[str] = field(default_factory=list)
    marked: bool = False        # 该任务的子句中有时限的【待补】【待核实】

    @property
    def text(self) -> str:
        return "".join(self.clauses).strip().rstrip("，,；;")


@dataclass
class Issue:
    id: str
    level: str
    line: Optional[int]
    excerpt: str
    message: str
    suggestion: str
    item: Optional[int] = None
    row: Optional[int] = None


@dataclass
class ListRow:
    index: int
    text: str
    owner: str
    deadline: str
    position: str = ""
    cooperate: str = ""


@dataclass
class ListTable:
    rows: List[ListRow] = field(default_factory=list)
    missing_columns: List[str] = field(default_factory=list)


# ---------------------------------------------------------------------------
# 读取
# ---------------------------------------------------------------------------

def setup_console() -> None:
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8")
        except Exception:  # 被重定向为 StringIO 等对象时没有 reconfigure
            pass


def read_lines(path: Path) -> List[str]:
    """读取文件为行列表（docx 为段落列表）。"""
    if not path.is_file():
        raise InputError(f"未找到文件：{path}")
    suffix = path.suffix.lower()
    if suffix in TEXT_SUFFIXES:
        try:
            text = path.read_text(encoding="utf-8-sig")
        except UnicodeDecodeError:
            raise InputError(f"文件不是 UTF-8 编码：{path}")
        return text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    if suffix == ".docx":
        try:
            import docx  # type: ignore
        except ImportError:
            raise InputError("读取 .docx 需要 python-docx，请先运行：pip install python-docx；"
                             "或把正文另存为 .txt 后再检查")
        try:
            document = docx.Document(str(path))
        except Exception as exc:  # 损坏或加密的文件
            raise InputError(f"无法读取 docx：{path}（{exc}）")
        return [p.text for p in document.paragraphs]
    raise InputError(f"不支持的文件格式：{path.suffix or '无扩展名'}（支持 .txt、.md、.docx）")


def body_lines(lines: Sequence[str]) -> List[Tuple[int, str]]:
    """【起草说明】之前的非空行，带 1 起的行号。"""
    result: List[Tuple[int, str]] = []
    for index, raw in enumerate(lines, start=1):
        if DRAFT_NOTE_RE.match(raw):
            break
        text = raw.strip()
        if text:
            result.append((index, text))
    return result


# ---------------------------------------------------------------------------
# 议定事项识别
# ---------------------------------------------------------------------------

def is_heading(text: str) -> bool:
    """层次标题行：带序号、较短、句末没有句号分号。"""
    if not LABEL_RE.match(text):
        return False
    return len(text) <= 40 and not re.search(r"[。；;]$", text) and "。" not in text


def split_label(text: str) -> Tuple[str, str]:
    m = LABEL_RE.match(text)
    if not m:
        return "", text.strip()
    return m.group("label"), text[m.end():].strip()


def split_yishi(content: str) -> List[str]:
    """"一是……二是……"拆成多条；只在段首或标点后出现的"×是"才算列举。"""
    starts = [m.start() for m in YISHI_RE.finditer(content)
              if m.start() == 0 or content[m.start() - 1] in "；;。，,：:"]
    if len(starts) < 2:
        return [content]
    parts = []
    for i, start in enumerate(starts):
        end = starts[i + 1] if i + 1 < len(starts) else len(content)
        part = content[start + 2:end].strip().rstrip("；;，,")
        if part:
            parts.append(part)
    head = content[:starts[0]].strip().rstrip("：:，,")
    if head:
        parts = [head + "：" + parts[0]] + parts[1:] if parts else [head]
    return parts


NOT_AGREE_RE = re.compile(r"不予?同意|不予通过|未予通过|否决")
# "会议同意暂不调整收费标准"是同意维持现状，不是暂缓
POSTPONE_RE = re.compile(r"暂缓|(?<!同意)暂不(?=[一-鿿])")
# 无需落实的结论：没有写责任单位和时限时，不报缺项，督办清单草表不填【待补】
APPOINT_RE = re.compile(
    r"任命|聘任|聘为|免去|免职|"
    r"同意[一-鿿]{1,4}?(?:同志)?(?:任|担任)[一-鿿]{0,12}?(?:主任|处长|科长|部长|院长|校长|经理|书记|主席|职务)"
)
KEEP_RE = re.compile(r"同意暂不|维持现状|维持原")
PASSED_RE = re.compile(r"审议通过|通过了")
# 议题标题下不带"会议"的暂缓、另行研究段落："爽约限制暂缓实施，待……"
BARE_DEFER_RE = re.compile(
    r"^(?!会议)[^，,。；;：:]{2,40}?(?:暂缓|暂不(?=[一-鿿])|另行(?:研究|审议)|下次会议(?:再议|研究|审议)|再议)"
)
RETURN_RE = re.compile(r"退回|修改后重新(?:提交|上会)")
LATER_RE = re.compile(r"另行(?:研究|审议)|再行研究|下次(?:会议)?(?:再议|研究|审议)|择期(?:再议|研究)|再议")
# 子句开头的同意类结论词："会议同意……""同意……""原则同意……"
AGREE_START_RE = re.compile(
    LEAD_IN + r"(?:会议)?(?:研究|审议|讨论)?(?:并|后)?(?P<p>原则上?)?(?:一致)?(?:同意|通过|审议通过)"
)
# 拆分时的结论分组：同组的子句并入同一条，不同组的拆开
GROUP_OF = {"同意": "yes", "原则同意": "yes", "决定": "yes", "退回修改": "yes",
            "暂缓": "later", "另行研究": "later", "不同意": "no"}
# 暂缓等条目中，以责任单位开头的准备任务："由财务科先行测算""财务科会后测算"
TASK_START_RE = re.compile(
    rf"^\s*(?:并|同时|先|会后)?(?:由(?:{NAMES})|(?:{NAMES})(?=先|会后|负责|牵头|抓紧|进一步))"
)
# 再议条件或原因：留在原条目
CONDITION_RE = re.compile(r"^\s*(?:待|等|因|鉴于|由于|视|如|若)|再(?:提交|研究|审议|议)|另行|下次|择期")
CLAUSE_RE = re.compile(r"[^，,；;。]+[，,；;。]*")
SENTENCE_RE_SPLIT = re.compile(r"[^；;。]+[；;。]*")


def _strip_quotes(text: str) -> str:
    return QUOTE_RE.sub(lambda m: "　" * len(m.group(0)), text)


def _mask(text: str) -> str:
    """引号、书名号内的文字和【待补】【待核实】标记换成等长空白，位置不变。"""
    text = _strip_quotes(text)
    return MARK_RE.sub(lambda m: "　" * len(m.group(0)), text)


def classify(text: str) -> str:
    text = _mask(text)
    if re.search(r"原则上?(?:同意|通过)", text):
        return "原则同意"
    if NOT_AGREE_RE.search(text):
        return "不同意"
    if POSTPONE_RE.search(text):
        return "暂缓"
    if RETURN_RE.search(text):
        return "退回修改"
    if LATER_RE.search(text):
        return "另行研究"
    if re.search(r"同意(?!书|票)|通过", text):  # "知情同意书""同意票"不是结论
        return "同意"
    return "决定"


def _clause_signal(clause: str) -> str:
    """子句自带的结论类型；没有明确结论词的返回空串。同意类只认子句开头的。"""
    masked = _mask(clause)
    if NOT_AGREE_RE.search(masked):
        return "不同意"
    if POSTPONE_RE.search(masked):
        return "暂缓"
    if LATER_RE.search(masked):
        return "另行研究"
    m = AGREE_START_RE.match(masked)
    if m:
        return "原则同意" if m.group("p") else "同意"
    return ""


def _segments(text: str, pattern: "re.Pattern[str]") -> List[str]:
    """按标点切分，切分位置在屏蔽引号和标记后的文本上确定，返回原文片段。"""
    masked = _mask(text)
    return [text[m.start():m.end()] for m in pattern.finditer(masked) if text[m.start():m.end()].strip()]


def _join(parts: Sequence[str]) -> str:
    return "".join(parts).strip().rstrip("，,；;")


def _split_tasks(piece: str) -> List[str]:
    """暂缓、不同意、另行研究的条目中，把"由××……"开头的准备任务拆成单独一条。"""
    clauses = _segments(piece, CLAUSE_RE)
    parent: List[str] = []
    tasks: List[List[str]] = []
    in_task = False
    for i, clause in enumerate(clauses):
        masked = _mask(clause)
        if i > 0 and TASK_START_RE.match(masked) and not CONDITION_RE.search(masked):
            tasks.append([clause])
            in_task = True
        elif in_task and not CONDITION_RE.search(masked):
            tasks[-1].append(clause)
        else:
            parent.append(clause)
            in_task = False
    if not tasks:
        return [piece]
    return [_join(parent)] + [_join(t) for t in tasks]


def _split_inline(piece: str) -> List[str]:
    """以"同意"开头、逗号后又有暂缓或不同意的一句，把暂缓、不同意的子句（连同其后的条件）拆出来。"""
    clauses = _segments(piece, CLAUSE_RE)
    parent: List[str] = []
    subs: List[List[str]] = []
    in_sub = False
    for i, clause in enumerate(clauses):
        masked = _mask(clause)
        signal = _clause_signal(clause) if i > 0 else ""
        if in_sub and CONDITION_RE.search(masked):
            subs[-1].append(clause)
        elif signal and GROUP_OF[signal] != "yes":
            subs.append([clause])
            in_sub = True
        else:
            parent.append(clause)
            in_sub = False
    if not subs or not parent:
        return [piece]
    return [_join(parent)] + [_join(s) for s in subs]


def split_mixed(text: str) -> List[str]:
    """一段中有几类结论的拆成几条：同意与暂缓、不同意、另行研究分开；暂缓等条目中的准备任务单列。"""
    sentences = _segments(text, SENTENCE_RE_SPLIT)
    if not sentences:
        return [text]
    pieces: List[List[str]] = [[sentences[0]]]
    group = GROUP_OF[classify(sentences[0])]
    last_yes = 0 if group == "yes" else -1
    for sentence in sentences[1:]:
        signal = _clause_signal(sentence)
        if signal and GROUP_OF[signal] != group:
            pieces.append([sentence])
            group = GROUP_OF[signal]
            if group == "yes":
                last_yes = len(pieces) - 1
        elif (not signal and group != "yes" and last_yes >= 0
              and TASK_START_RE.match(_mask(sentence)) and not CONDITION_RE.search(_mask(sentence))):
            # "同意……；爽约限制暂缓实施。由门诊部牵头……"：暂缓句之后的责任单位、时限属于前面的同意事项
            pieces[last_yes].append(sentence)
        else:
            pieces[-1].append(sentence)
    result: List[str] = []
    for piece in pieces:
        joined = _join(piece)
        if GROUP_OF[classify(joined)] in ("later", "no"):
            clauses = _segments(joined, CLAUSE_RE)
            first = clauses[0] if clauses else joined
            if AGREE_START_RE.match(_mask(first)):
                result.extend(_split_inline(joined))
            else:
                result.extend(_split_tasks(joined))
        else:
            result.append(joined)
    return [p for p in result if p]


def _plain(text: str) -> str:
    """去掉引号书名号内的文字、待补标记和段首结论词，供识别单位和时限用。"""
    return LEAD_DECISION_RE.sub("", MARK_RE.sub("，", _strip_quotes(text)))


def find_owner(text: str) -> str:
    plain = _plain(text)
    for pattern in OWNER_PATTERNS:
        m = pattern.search(plain)
        if m:
            return m.group("o").strip()
    return ""


def find_cooperate(text: str) -> str:
    plain = _plain(text)
    found: List[str] = []
    for pattern in COOP_PATTERNS:
        for m in pattern.finditer(plain):
            value = m.group("c").strip()
            if value and value not in found:
                found.append(value)
    return "、".join(found)


def find_deadline(text: str) -> str:
    plain = MARK_RE.sub("，", _strip_quotes(text))
    hits: List[Tuple[int, str]] = []
    for pattern in DEADLINE_PATTERNS:
        for m in pattern.finditer(plain):
            hits.append((m.start(), m.group(0)))
    if not hits:
        return ""
    hits.sort(key=lambda h: (h[0], -len(h[1])))
    # 去掉被更长命中覆盖的短命中
    result: List[str] = []
    last_end = -1
    for start, value in hits:
        if start < last_end:
            continue
        result.append(value)
        last_end = start + len(value)
    return "、".join(result)


def make_item(no: int, line: int, position: str, section: str, text: str, kind: str = "议定") -> Item:
    decision = classify(text) if kind == "议定" else "会议要求"
    item = Item(no=no, line=line, position=position, section=section, text=text, decision=decision, kind=kind)
    item.owner = find_owner(text)
    item.owner_marked = bool(OWNER_MARK_RE.search(text))
    item.deadline = find_deadline(text)
    item.deadline_marked = bool(DEADLINE_MARK_RE.search(text))
    item.deadline_vague = bool(VAGUE_TIME_RE.search(text))
    item.cooperate = find_cooperate(text)
    if kind == "议定" and not (item.owner or item.owner_marked or item.deadline or item.deadline_marked):
        item.no_follow = no_follow_kind(text)
    return item


def no_follow_kind(text: str) -> str:
    """任职、维持现状、审议通过这类无需落实的结论，返回类别；其他返回空串。"""
    masked = _mask(text)
    if APPOINT_RE.search(masked):
        return "任职"
    if KEEP_RE.search(masked):
        return "维持现状"
    if PASSED_RE.search(masked):
        return "审议通过"
    return ""


# 会议机构不是任务的承办单位："由党委会研究决定"不另起一项任务
MEETING_BODY_RE = re.compile(
    r"(?:党委|党组|常委|办公|局务|董事|职代|教代|股东|代表大)会$|会议$"
)


def _units_match(a: str, b: str) -> bool:
    x, y = normalize(a), normalize(b)
    return bool(x) and bool(y) and (x == y or x in y or y in x)


def _any_unit(unit: str, candidates: Sequence[str]) -> bool:
    return any(_units_match(unit, c) for c in candidates)


def _same_owner(a: str, b: str) -> bool:
    return any(_any_unit(u, split_units(b)) for u in split_units(a))


def split_tasks(text: str) -> List[Task]:
    """按子句把责任单位和时限配对。

    时限归它所在子句或此前最近子句中的责任单位；第一个责任单位之前的子句和时限归第一个责任单位。
    "由后勤处于12月10日前完成，财务科11月20日前拨付经费"拆成两项任务，各带各的时限。
    """
    tasks: List[Task] = []
    lead: List[str] = []
    lead_deadlines: List[str] = []
    lead_marked = False
    current: Optional[Task] = None
    for clause in _segments(text, CLAUSE_RE):
        owner = find_owner(clause)
        if owner and not MEETING_BODY_RE.search(owner):
            found = next((t for t in tasks if _same_owner(t.owner, owner)), None)
            if found is None:
                found = Task(owner=owner)
                tasks.append(found)
            current = found
        deadline = find_deadline(clause)
        tokens = deadline.split("、") if deadline else []
        marked = bool(DEADLINE_MARK_RE.search(clause))
        if current is None:
            lead.append(clause)
            lead_deadlines.extend(tokens)
            lead_marked = lead_marked or marked
            continue
        current.clauses.append(clause)
        current.deadlines.extend(t for t in tokens if t not in current.deadlines)
        current.marked = current.marked or marked
    if not tasks:
        return [Task(owner="", deadlines=lead_deadlines, clauses=lead, marked=lead_marked)]
    first = tasks[0]
    first.clauses = lead + first.clauses
    first.deadlines = [t for t in lead_deadlines if t not in first.deadlines] + first.deadlines
    first.marked = first.marked or lead_marked
    return tasks


def heading_block(content: str) -> str:
    """标题本身是"议定事项""会议要求"时，返回其下各条的类别。"""
    if DECISION_HEADING_RE.search(content):
        return "议定"
    if REQ_HEADING_RE.search(content):
        return "要求"
    return ""


def extract_items(body: Sequence[Tuple[int, str]]) -> List[Item]:
    """识别议定事项和会议要求中的任务，按出现顺序编号。"""
    items: List[Item] = []
    section_text = ""
    section_ord = ""
    block = ""            # ""：不在块内；"议定"：会议议定引出的块；"要求"：会议要求引出的块
    counts = {"议定": 0, "要求": 0}

    def add(line_no: int, label: str, text: str, kind: str) -> None:
        parts = split_yishi(text)
        single = len(parts) == 1
        if kind == "议定":
            # 一段中有几类结论的拆开；拆出的各条共用原来的纪要位置
            parts = [piece for part in parts for piece in split_mixed(part)]
        for part in parts:
            counts[kind] += 1
            n = counts[kind]
            if kind == "要求":
                pos = f"{section_ord}会议要求第{n}条"
            elif label and single:
                pos = label if H1_PREFIX_RE.match(label) else f"{section_ord}{label}"
            else:
                pos = f"{section_ord}第{n}项"
            items.append(make_item(len(items) + 1, line_no, pos, section_text, part, kind))

    for line_no, text in body:
        if ATTEND_RE.match(text):
            block = ""
            continue
        label, content = split_label(text)
        heading = is_heading(text)
        h1 = H1_PREFIX_RE.match(text)
        if heading and h1:
            # 单议题纪要"会议议定以下事项："之后的"一、同意……"（不带句号）仍算条目
            if block and not section_text and DECISION_WORD_RE.search(content):
                add(line_no, label, content, block)
                continue
            block = heading_block(content)
            section_ord, section_text = h1.group(1), text
            counts = {"议定": 0, "要求": 0}
            continue
        if heading and (not block or heading_block(content)):
            block = heading_block(content) or block
            continue
        if BLOCK_START_RE.match(content):
            block = "议定"
            continue
        if REQ_BLOCK_RE.match(content):
            block = "要求"
            continue
        inline = INLINE_RE.match(content)
        if inline:
            block = ""
            add(line_no, label, inline.group("content"), "议定")
            continue
        if SENTENCE_RE.match(content):
            add(line_no, label, content, "议定")
            continue
        req = REQ_INLINE_RE.match(content)
        if req:
            block = ""
            add(line_no, label, req.group("content"), "要求")
            continue
        if PROCESS_CUE_RE.match(content):
            mid = MID_DECISION_RE.search(content)
            if mid:
                block = ""
                add(line_no, label, content[mid.end():], "议定")
                continue
        if OTHER_CUE_RE.match(content):
            block = ""
            continue
        if block:
            add(line_no, label, content, block)
        elif section_text and not heading and BARE_DEFER_RE.match(_mask(content)):
            # 多议题纪要中某议题没写"会议议定："，"××事项暂缓，待……"仍按议定事项识别
            add(line_no, label, content, "议定")
    return items


# ---------------------------------------------------------------------------
# 纪要检查
# ---------------------------------------------------------------------------

def excerpt_of(text: str, limit: int = 40) -> str:
    return text if len(text) <= limit else text[:limit] + "……"


def check_items(items: Sequence[Item]) -> List[Issue]:
    issues: List[Issue] = []
    for it in items:
        if it.kind != "议定":
            continue
        ex = excerpt_of(it.text)
        if it.decision not in DECISION_EXEMPT and not it.no_follow:
            if not it.owner and not it.owner_marked:
                issues.append(Issue(
                    "MIN-001", "P1", it.line, ex,
                    "议定事项没有写责任单位，也没有标【待补】",
                    "按记录写明牵头单位；记录没说的写【待补：责任单位】。任职、审议通过、纯表态等无需落实的结论可以不写",
                    item=it.no))
            if it.decision != "退回修改" and not it.deadline and not it.deadline_marked:
                if it.deadline_vague:
                    message = "完成时限只写了“尽快”“抓紧”之类，没有具体时限"
                    suggestion = "照记录写具体时限；会上只说了“尽快”的，写“尽快完成【待补：具体完成时限】”"
                else:
                    message = "议定事项没有写完成时限，也没有标【待补】"
                    suggestion = ("按记录写完成时限；记录没说的写【待补：完成时限】，不自行补日期。"
                                  "任职、审议通过、纯表态等无需落实的结论不必写时限，也不加【待补】")
                issues.append(Issue("MIN-002", "P1", it.line, ex, message, suggestion, item=it.no))
        if it.decision == "原则同意" and not REVISION_HINT_RE.search(it.text):
            issues.append(Issue(
                "MIN-003", "P2", it.line, ex,
                "“原则同意”没有写出会上提出的修改要求",
                "写明改什么、谁修改；会上没说具体要求的，标【待核实：会上提出的具体修改意见】",
                item=it.no))
        if it.decision == "暂缓" and not POSTPONE_HINT_RE.search(it.text.replace("暂缓", "")):
            issues.append(Issue(
                "MIN-004", "P2", it.line, ex,
                "“暂缓”没有写原因或再议条件",
                "照记录写暂缓原因或“待……后再提交研究”；记录没有的标【待补：暂缓原因或再议条件】",
                item=it.no))
    return issues


def check_body(body: Sequence[Tuple[int, str]], items: Sequence[Item]) -> List[Issue]:
    issues: List[Issue] = []
    for line_no, text in body:
        plain = _strip_quotes(text)
        m = COLLOQUIAL_RE.search(plain)
        if m:
            issues.append(Issue(
                "MIN-005", "P1", line_no, excerpt_of(text),
                f"口语残留：“{m.group(0)}”",
                "改成书面表述，意思、数字、条件和态度强弱不变"))
        tag = SPEAKER_TAG_RE.search(plain)
        if tag:
            issues.append(Issue(
                "MIN-008", "P1", line_no, excerpt_of(text),
                f"转写标签残留：“{tag.group(0)}”",
                "纪要按议题归纳，不逐人列发言；需要保留发言过程的整理成会议记录"))
    if not any(it.kind == "议定" for it in items):
        line = body[0][0] if body else None
        issues.append(Issue(
            "MIN-006", "P2", line, "",
            "没有识别到议定事项（未找到“会议议定：”“会议同意……”等写法）",
            "纪要要写明议定事项；会议确实没有形成结论的，在起草说明中写明"))
    title = body[0][1] if body else ""
    if title.endswith("纪要") and not any(re.match(r"^\s*出\s*席\s*[：:]", t) for _, t in body):
        issues.append(Issue(
            "MIN-007", "P2", body[0][0], title,
            "没有出席人员名单",
            "正文后空一行写“出席：”，名单只用用户提供的；没有的写“出席：【待补：出席人员名单】”（GB/T 9704-2012 10.3）"))
    return issues


# ---------------------------------------------------------------------------
# 督办清单读取与对账
# ---------------------------------------------------------------------------

COLUMN_KEYS = {
    "text": ("议定事项", "督办事项", "事项", "任务", "督办内容", "工作内容"),
    "owner": ("责任单位", "牵头单位", "责任部门", "牵头部门", "承办单位", "责任科室", "牵头科室"),
    "deadline": ("完成时限", "办结时限", "完成时间", "时限", "截止时间", "期限"),
    "position": ("纪要位置", "纪要条目", "对应纪要", "来源"),
    "cooperate": ("配合单位", "配合部门", "协办单位"),
}
COLUMN_NAMES = {"text": "议定事项", "owner": "责任单位", "deadline": "完成时限"}


def _match_columns(header: Sequence[str]) -> Dict[str, int]:
    mapping: Dict[str, int] = {}
    cells = [re.sub(r"\s", "", str(h or "")) for h in header]
    for key, names in COLUMN_KEYS.items():
        for name in names:
            for idx, cell in enumerate(cells):
                if idx in mapping.values():
                    continue
                if cell == name or (name in cell and key != "text"):
                    mapping[key] = idx
                    break
            if key in mapping:
                break
    if "text" not in mapping:
        for idx, cell in enumerate(cells):
            if idx not in mapping.values() and ("事项" in cell or "任务" in cell):
                mapping["text"] = idx
                break
    return mapping


def _split_md_row(line: str) -> List[str]:
    text = line.strip()
    if text.startswith("|"):
        text = text[1:]
    if text.endswith("|") and not text.endswith("\\|"):
        text = text[:-1]
    cells = re.split(r"(?<!\\)\|", text)
    return [c.replace("\\|", "|").strip() for c in cells]


def _rows_to_table(header: Sequence[str], data: Sequence[Sequence[str]]) -> ListTable:
    mapping = _match_columns(header)
    table = ListTable()
    for key, name in COLUMN_NAMES.items():
        if key not in mapping:
            table.missing_columns.append(name)
    if table.missing_columns:
        return table

    def cell(row: Sequence[str], key: str) -> str:
        idx = mapping.get(key)
        if idx is None or idx >= len(row):
            return ""
        return str(row[idx] or "").strip()

    for i, row in enumerate(data, start=1):
        if not any(str(c or "").strip() for c in row):
            continue
        table.rows.append(ListRow(
            index=i, text=cell(row, "text"), owner=cell(row, "owner"), deadline=cell(row, "deadline"),
            position=cell(row, "position"), cooperate=cell(row, "cooperate")))
    return table


def read_list(path: Path) -> ListTable:
    if not path.is_file():
        raise InputError(f"未找到督办清单：{path}")
    suffix = path.suffix.lower()
    if suffix in TEXT_SUFFIXES:
        lines = read_lines(path)
        tables: List[List[List[str]]] = []
        current: List[List[str]] = []
        for line in lines + [""]:
            if line.strip().startswith("|"):
                current.append(_split_md_row(line))
            elif current:
                tables.append(current)
                current = []
        for table in tables:
            if len(table) < 2:
                continue
            header, rest = table[0], table[1:]
            if rest and all(re.fullmatch(r":?-{2,}:?", c.replace(" ", "")) for c in rest[0] if c):
                rest = rest[1:]
            mapping = _match_columns(header)
            if "owner" in mapping or "deadline" in mapping:
                return _rows_to_table(header, rest)
        raise InputError(f"督办清单中没有找到含“责任单位”或“完成时限”列的 Markdown 表格：{path}")
    if suffix == ".xlsx":
        try:
            import openpyxl  # type: ignore
        except ImportError:
            raise InputError("读取 .xlsx 需要 openpyxl，请先运行：pip install openpyxl；或改用 Markdown 清单")
        try:
            workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        except Exception as exc:
            raise InputError(f"无法读取 xlsx：{path}（{exc}）")
        try:
            rows = [[("" if v is None else str(v)) for v in row] for row in workbook.worksheets[0].iter_rows(values_only=True)]
        finally:
            workbook.close()
        for idx, row in enumerate(rows[:10]):
            mapping = _match_columns(row)
            if "owner" in mapping or "deadline" in mapping:
                return _rows_to_table(row, rows[idx + 1:])
        raise InputError(f"督办清单前 10 行中没有找到含“责任单位”或“完成时限”的表头：{path}")
    raise InputError(f"不支持的督办清单格式：{path.suffix or '无扩展名'}（支持 .md、.txt、.xlsx）")


def normalize(text: str) -> str:
    text = MARK_RE.sub("", text)
    text = YEAR_BEFORE_MONTH_RE.sub("", text)
    return PUNCT_RE.sub("", text)


def bigrams(text: str) -> set:
    norm = normalize(text)
    return {norm[i:i + 2] for i in range(len(norm) - 1)}


def similarity(row_text: str, item_text: str) -> float:
    a = bigrams(row_text)
    if not a:
        return 0.0
    return len(a & bigrams(item_text)) / len(a)


def is_marked(value: str) -> bool:
    return bool(MARK_RE.search(value)) or value.strip() in ("", "无", "—", "-", "／", "/")


def split_units(value: str) -> List[str]:
    parts = re.split(r"[、，,；;/\s]+|(?<=[\u4e00-\u9fff])[和及与](?=[\u4e00-\u9fff])", MARK_RE.sub("", value))
    return [p for p in (s.strip() for s in parts) if len(p) >= 2]


def norm_deadline(token: str) -> str:
    """时限比对用的规范形式：去掉年份前缀、标点和"前""之前""内"等后缀。"""
    return DEADLINE_SUFFIX_RE.sub("", normalize(token))


def deadline_tokens(value: str) -> List[str]:
    found = find_deadline(value)
    tokens = found.split("、") if found else []
    return [n for n in (norm_deadline(t) for t in tokens) if n]


def _issue(id_: str, level: str, best: Item, excerpt: str, row: ListRow, message: str, suggestion: str) -> Issue:
    return Issue(id_, level, best.line, excerpt, message, suggestion, item=best.no, row=row.index)


def _check_owner(row: ListRow, best: Item, tasks: Sequence[Task], excerpt: str) -> List[Issue]:
    """清单责任单位必须是该条识别出的牵头（负责）单位，不能只是在原文中出现过。"""
    if is_marked(row.owner):
        if best.owner and not best.owner_marked:
            return [_issue("MIN-107", "P1", best, excerpt, row,
                           f"督办清单第 {row.index} 行的责任单位为空或标了待补，纪要中写的是“{best.owner}”",
                           "按纪要补齐责任单位")]
        return []
    item_norm = normalize(best.text)
    units = split_units(row.owner)
    owners = [t.owner for t in tasks if t.owner]
    if best.owner and not any(_same_owner(best.owner, o) for o in owners):
        owners.append(best.owner)
    if not owners:
        if best.owner_marked or not units or all(normalize(u) not in item_norm for u in units):
            return [_issue("MIN-106", "P0", best, excerpt, row,
                           f"督办清单第 {row.index} 行写了责任单位“{row.owner}”，纪要中这条没有写明（或标了待补）",
                           "清单不新增纪要没有的责任单位；先按记录补纪要，再同步清单")]
        missing = [u for u in units if normalize(u) not in item_norm]
        if missing:
            return [_issue("MIN-102", "P0", best, excerpt, row,
                           f"督办清单第 {row.index} 行的责任单位“{'、'.join(missing)}”与纪要不一致",
                           "责任单位逐字照纪要写")]
        return []
    owner_units = [u for o in owners for u in split_units(o)]
    coop_units = split_units(best.cooperate)
    hit = [u for u in units if _any_unit(u, owner_units)]
    extra = [u for u in units if u not in hit]
    as_coop = [u for u in extra if _any_unit(u, coop_units)]
    unknown = [u for u in extra if u not in as_coop and normalize(u) not in item_norm]
    parts: List[str] = []
    if not hit:
        parts.append(f"纪要中这条由“{'、'.join(owners)}”牵头或负责")
    if as_coop:
        parts.append(f"“{'、'.join(as_coop)}”在纪要中是配合单位")
    if unknown:
        parts.append(f"“{'、'.join(unknown)}”在纪要这一条中找不到")
    if not parts:
        return []
    return [_issue("MIN-102", "P0", best, excerpt, row,
                   f"督办清单第 {row.index} 行的责任单位“{row.owner}”与纪要不一致：" + "；".join(parts),
                   "责任单位、配合单位逐字照纪要写，牵头和配合不对调")]


def _row_tasks(row: ListRow, tasks: Sequence[Task]) -> List[Task]:
    """清单这一行的责任单位在该条中承担的任务。"""
    if is_marked(row.owner):
        return []
    units = split_units(row.owner)
    return [t for t in tasks if t.owner and any(_any_unit(u, split_units(t.owner)) for u in units)]


def _check_deadline(row: ListRow, best: Item, tasks: Sequence[Task], excerpt: str) -> List[Issue]:
    """清单时限与该单位所承担任务的时限比对；该单位的任务没写时限时，再与整条比对。"""
    own = _row_tasks(row, tasks)
    own_deadlines = [d for t in own for d in t.deadlines]
    scope = own_deadlines or [d for t in tasks for d in t.deadlines]
    if is_marked(row.deadline):
        if scope and (own_deadlines or not best.deadline_marked):
            return [_issue("MIN-107", "P1", best, excerpt, row,
                           f"督办清单第 {row.index} 行的完成时限为空或标了待补，纪要中写的是“{'、'.join(scope)}”",
                           "按纪要补齐完成时限")]
        return []
    if not best.deadline or (own and not own_deadlines and any(t.marked for t in own)):
        return [_issue("MIN-106", "P0", best, excerpt, row,
                       f"督办清单第 {row.index} 行写了完成时限“{row.deadline}”，纪要中这条（这项任务）没有写明或标了待补",
                       "清单不新增纪要没有的时限；先按记录补纪要，再同步清单")]
    allowed = {norm_deadline(d) for d in scope}
    tokens = deadline_tokens(row.deadline)
    if tokens:
        bad = [t for t in tokens if t not in allowed]
    else:
        bad = [] if normalize(row.deadline) in normalize(best.text) else [row.deadline]
    if not bad:
        return []
    every = {norm_deadline(d) for t in tasks for d in t.deadlines}
    if own_deadlines and all(t in every for t in bad):
        message = (f"督办清单第 {row.index} 行的完成时限“{row.deadline}”在纪要中属于同一条中其他单位的任务，"
                   f"“{row.owner}”的时限是“{'、'.join(own_deadlines)}”")
    else:
        message = f"督办清单第 {row.index} 行的完成时限“{row.deadline}”与纪要（“{'、'.join(scope)}”）不一致"
    return [_issue("MIN-103", "P0", best, excerpt, row, message, "完成时限逐字照纪要写，跟着所属任务走")]


def _check_cooperate(row: ListRow, best: Item, tasks: Sequence[Task], excerpt: str) -> List[Issue]:
    """配合单位不能是这项任务的牵头单位，也不能是纪要这一条中没有的单位。"""
    if is_marked(row.cooperate):
        return []
    units = split_units(row.cooperate)
    own = _row_tasks(row, tasks)
    leaders = [u for t in (own or tasks) for u in split_units(t.owner)]
    if best.owner and not own:
        leaders += split_units(best.owner)
    swapped = [u for u in units if _any_unit(u, leaders)]
    item_norm = normalize(best.text)
    unknown = [u for u in units if u not in swapped and normalize(u) not in item_norm]
    parts: List[str] = []
    if swapped:
        parts.append(f"“{'、'.join(swapped)}”在纪要中是牵头或负责单位")
    if unknown:
        parts.append(f"“{'、'.join(unknown)}”在纪要这一条中找不到")
    if not parts:
        return []
    return [_issue("MIN-108", "P1", best, excerpt, row,
                   f"督办清单第 {row.index} 行的配合单位“{row.cooperate}”与纪要不一致：" + "；".join(parts),
                   "配合单位照纪要写，牵头和配合不对调")]


def match_list(items: Sequence[Item], table: ListTable, threshold: float = 0.5) -> List[Issue]:
    issues: List[Issue] = []
    if table.missing_columns:
        issues.append(Issue(
            "MIN-105", "P1", None, "",
            "督办清单缺少必需列：" + "、".join(table.missing_columns),
            "列为 序号｜议定事项｜纪要位置｜责任单位｜配合单位｜完成时限｜办理状态｜备注"))
        return issues
    matched_items = set()
    task_cache: Dict[int, List[Task]] = {}
    for row in table.rows:
        best: Optional[Item] = None
        best_score = 0.0
        best_key = (0.0, 0.0)
        for it in items:
            raw = similarity(row.text, it.text)
            score = raw
            if row.position and it.position and normalize(row.position) == normalize(it.position):
                score = max(raw, 0.6 if raw >= 0.3 else raw)
            # 同一纪要位置拆出几条时，位置加分相同，再比文字相似度
            if (score, raw) > best_key:
                best, best_score, best_key = it, score, (score, raw)
        excerpt = excerpt_of(row.text)
        if best is None or best_score < threshold:
            issues.append(Issue(
                "MIN-101", "P0", None, excerpt,
                f"督办清单第 {row.index} 行的事项在纪要中找不到对应的议定事项",
                "督办清单只从纪要取内容；纪要没有的事项删去，或先确认纪要是否漏写", row=row.index))
            continue
        matched_items.add(best.no)
        if best.no not in task_cache:
            task_cache[best.no] = split_tasks(best.text)
        tasks = task_cache[best.no]
        issues += _check_owner(row, best, tasks, excerpt)
        issues += _check_deadline(row, best, tasks, excerpt)
        issues += _check_cooperate(row, best, tasks, excerpt)
    for it in items:
        if it.kind != "议定" or it.no_follow:
            continue
        needs_follow = it.decision not in DECISION_EXEMPT or (it.decision == "另行研究" and it.owner)
        if needs_follow and it.no not in matched_items:
            issues.append(Issue(
                "MIN-104", "P1", it.line, excerpt_of(it.text),
                f"纪要第 {it.no} 条议定事项（{it.position}）没有列入督办清单",
                "需要落实的议定事项逐条列入清单；确属纯表态、无需落实的可以不列", item=it.no))
    return issues


# ---------------------------------------------------------------------------
# 生成督办清单草表
# ---------------------------------------------------------------------------

def _cell(text: str) -> str:
    return text.replace("|", "\\|").replace("\n", " ")


def render_list(items: Sequence[Item]) -> str:
    lines = [
        "| 序号 | 议定事项 | 纪要位置 | 责任单位 | 配合单位 | 完成时限 | 办理状态 | 备注 |",
        "|---|---|---|---|---|---|---|---|",
    ]
    rows: List[Tuple[str, str, str, str, str, str]] = []  # 事项、位置、责任单位、配合单位、时限、备注
    for it in items:
        if it.decision in ("暂缓", "不同意") or (it.decision == "另行研究" and not it.owner):
            continue
        if it.kind == "要求" and not it.owner:
            continue
        if it.no_follow == "维持现状":
            continue
        if it.no_follow:
            # 任职、审议通过：无需落实，不填【待补】，由人工确认是否列入
            rows.append((it.text, it.position, "", "", "", f"{it.no_follow}事项，请确认是否列入"))
            continue
        tasks = [t for t in split_tasks(it.text) if t.owner]
        if len(tasks) >= 2:
            # 一条中几个单位各有任务的，拆成几行，各带各的时限
            for t in tasks:
                deadline = "、".join(t.deadlines) or (
                    DEADLINE_MARK_RE.search(t.text).group(0) if t.marked else "【待补：完成时限】")
                note = "" if t.deadlines or t.marked else "纪要未写时限"
                rows.append((t.text, it.position, t.owner, find_cooperate(t.text), deadline, note))
            continue
        owner = it.owner or (MARK_RE.search(it.text).group(0) if it.owner_marked else "【待补：责任单位】")
        deadline = it.deadline or (DEADLINE_MARK_RE.search(it.text).group(0) if it.deadline_marked
                                   else "【待补：完成时限】")
        notes = []
        if not it.owner and not it.owner_marked:
            notes.append("纪要未写责任单位")
        if not it.deadline and not it.deadline_marked:
            notes.append("纪要未写时限")
        rows.append((it.text, it.position, owner, it.cooperate, deadline, "；".join(notes)))
    for no, (text, position, owner, coop, deadline, note) in enumerate(rows, start=1):
        lines.append("| {} | {} | {} | {} | {} | {} | 待办理 | {} |".format(
            no, _cell(text), _cell(position), _cell(owner), _cell(coop), _cell(deadline), _cell(note)))
    return "\n".join(lines) + "\n"


def unique_path(path: Path) -> Path:
    if not path.exists():
        return path
    for n in range(2, 1000):
        candidate = path.with_name(f"{path.stem}-v{n}{path.suffix}")
        if not candidate.exists():
            return candidate
    raise InputError(f"同名文件过多，无法生成新文件名：{path}")


# ---------------------------------------------------------------------------
# 汇总输出
# ---------------------------------------------------------------------------

def build_report(minutes: Path, list_file: Optional[Path] = None, emit: Optional[Path] = None) -> dict:
    body = body_lines(read_lines(minutes))
    items = extract_items(body)
    issues = check_items(items) + check_body(body, items)
    if list_file is not None:
        issues += match_list(items, read_list(list_file))
    emitted = None
    if emit is not None:
        target = unique_path(emit)
        try:
            if not target.parent.exists():
                target.parent.mkdir(parents=True)
            target.write_text(render_list(items), encoding="utf-8")
        except OSError as exc:
            raise InputError(f"无法写入督办清单草表：{target}（{exc}）")
        emitted = str(target)
    issues.sort(key=lambda i: (LEVEL_ORDER.index(i.level), i.line or 0, i.id))
    summary = {"items": len(items), "P0": 0, "P1": 0, "P2": 0}
    for issue in issues:
        summary[issue.level] += 1
    return {
        "file": str(minutes),
        "list_file": str(list_file) if list_file else None,
        "summary": summary,
        "items": [asdict(it) for it in items],
        "issues": [asdict(i) for i in issues],
        "emitted": emitted,
    }


def render_text(report: dict) -> str:
    s = report["summary"]
    out = [f"文件：{report['file']}"]
    if report.get("list_file"):
        out.append(f"督办清单：{report['list_file']}")
    out.append(f"议定事项 {s['items']} 条；问题 P0 {s['P0']} 条、P1 {s['P1']} 条、P2 {s['P2']} 条")
    if report["items"]:
        out.append("")
        out.append("议定事项：")
        for it in report["items"]:
            owner = it["owner"] or ("已标待补" if it["owner_marked"] else "未写")
            deadline = it["deadline"] or ("已标待补" if it["deadline_marked"] else "未写")
            out.append(f"  {it['no']}. [{it['decision']}] {it['position']}（第 {it['line']} 行）"
                       f" 责任单位：{owner}；时限：{deadline}")
    if report["issues"]:
        out.append("")
        out.append("问题：")
        for i in report["issues"]:
            if i["line"]:
                where = f"第 {i['line']} 行"
            elif i.get("row"):
                where = f"清单第 {i['row']} 行"
            elif i["id"].startswith("MIN-1"):
                where = "清单"
            else:
                where = "全文"
            out.append(f"  [{i['level']}] {i['id']} {where}：{i['message']}")
            if i["excerpt"]:
                out.append(f"        原文：{i['excerpt']}")
            out.append(f"        建议：{i['suggestion']}")
    if report.get("emitted"):
        out.append("")
        out.append(f"已生成督办清单草表：{report['emitted']}（请逐行核对后使用）")
    return "\n".join(out)


class _ChineseHelpFormatter(argparse.RawDescriptionHelpFormatter):
    def add_usage(self, usage, actions, groups, prefix=None):
        return super().add_usage(usage, actions, groups, "用法：" if prefix is None else prefix)


def make_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        prog="minutes_check.py",
        description="检查会议纪要的议定事项是否写明责任单位和完成时限，与督办清单逐条对账，"
                    "或生成督办清单草表。只检查【起草说明】之前的正文。",
        epilog="退出码：0 正常；--strict 且存在 P0 时为 1；输入错误为 2。",
        formatter_class=_ChineseHelpFormatter,
        add_help=False,
    )
    parser._optionals.title = "选项"
    parser._positionals.title = "参数"
    parser.add_argument("-h", "--help", action="help", help="显示本帮助信息并退出")
    parser.add_argument("minutes", help="纪要正文（.txt、.md、.docx）")
    parser.add_argument("--list", dest="list_file", help="督办清单（.md 中的 Markdown 表格，或 .xlsx）")
    parser.add_argument("--emit-list", dest="emit", help="按纪要生成督办清单草表（Markdown），文件已存在时加 -v2")
    parser.add_argument("--json", action="store_true", help="以 JSON 格式输出")
    parser.add_argument("--strict", action="store_true", help="存在 P0 时返回 1")
    return parser


def main(argv: Optional[Sequence[str]] = None) -> int:
    setup_console()
    parser = make_parser()
    try:
        args = parser.parse_args(argv)
    except SystemExit as exc:
        return 0 if exc.code == 0 else 2
    try:
        report = build_report(
            Path(args.minutes),
            Path(args.list_file) if args.list_file else None,
            Path(args.emit) if args.emit else None,
        )
    except InputError as exc:
        print(f"错误：{exc}", file=sys.stderr)
        return 2
    if args.json:
        print(json.dumps(report, ensure_ascii=False, indent=2))
    else:
        print(render_text(report))
    if args.strict and report["summary"]["P0"]:
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
