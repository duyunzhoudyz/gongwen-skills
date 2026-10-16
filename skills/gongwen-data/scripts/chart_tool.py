#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""公文图表：用 openpyxl 生成 Excel 原生图表（黑白灰配色，标注单位和数据来源）；装有 matplotlib 时可另出 PNG。

用法：
  python scripts/chart_tool.py <表格> --type bar|barh|line|pie|stacked [--sheet 名称] [--header-row N]
         [--x 分类列] [--y 数值列[,数值列]] [--title 图题] [--unit 单位] [--source 数据来源] [--note 注]
         [--sort none|desc|asc] [--top N] [--highlight 名称] [--out 图表.xlsx] [--png] [--json]

图表类型：bar 柱状图（比较各单位）；barh 条形图（类别多、名称长）；line 折线图（按时间的变化）；
         pie 饼图（构成，最多 6 类）；stacked 堆积柱状图（总量及其构成）。
约定：
  - 本地生成，不联网，不调用任何在线图表或渲染服务。
  - 合计、小计、"其中"、平均行不入图；空值留空，不按 0 画。
  - 纵轴从 0 开始；不画双坐标轴，不同单位的指标分开出图。
  - 输出文件已存在时加 -v2、-v3，不覆盖。
  - 姓名、手机号等个人信息列不作分类；材料带密级或内部标注时停止（DT-SEC-01），用户说明后加 --marking-ok。
退出码：0 成功；2 输入错误（文件或列不存在、图表类型不适用、个人信息列作分类、缺少 openpyxl、
       材料带密级或内部标注而未加 --marking-ok）。
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from decimal import Decimal
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence

sys.path.insert(0, str(Path(__file__).resolve().parent))

import data_tool as dt  # noqa: E402

VERSION = dt.VERSION
GRAYS = ("404040", "A6A6A6", "D9D9D9", "7F7F7F", "BFBFBF", "000000")
LINE_STYLES = (("000000", "solid", "circle"), ("595959", "dash", "triangle"), ("8C8C8C", "sysDot", "square"),
               ("000000", "lgDash", "diamond"), ("7F7F7F", "dashDot", "x"))
HATCHES = ("", "///", "...", "xxx", "\\\\\\", "---")
PIE_MAX = 6
TIME_RE = re.compile(r"年|月|季度|季|日期|时间|周|学期|学年")
PERIOD_VALUE_RE = re.compile(r"(?:19|20)\d{2}|\d{1,2}月|第?[一二三四1-4]季度|[上下]半年|\d{4}[-/.]\d{1,2}")
CJK_FONTS = ("SimSun", "SimHei", "Microsoft YaHei", "Noto Sans CJK SC", "Noto Serif CJK SC", "Source Han Sans SC",
             "Source Han Serif SC", "WenQuanYi Micro Hei", "PingFang SC", "Heiti SC", "STSong", "FangSong")


class ChartError(dt.InputError):
    pass


def pick_columns(table: dt.Table, x: Optional[str], y: Optional[str], kind: str = "bar",
                 warnings: Optional[List[str]] = None) -> tuple:
    if x:
        xcol = table.column(x)
        if xcol is None:
            raise ChartError(f"表中没有名为“{x}”的列；现有列：{'、'.join(c.name for c in table.columns)}")
    else:
        if table.key_col is None:
            candidates = [c for c in table.columns if c.kind in ("period", "text", "id")]
            if not candidates:
                raise ChartError("找不到分类列，请用 --x 指定")
            xcol = candidates[0]
        else:
            xcol = table.columns[table.key_col]
    if xcol.sensitive or dt.SENSITIVE_NAME_RE.search(xcol.name):
        # 姓名、手机号等不作分类：脱敏后会重名（如两个"张某"），不脱敏又泄露个人信息
        raise ChartError(f"“{xcol.name}”是个人信息列，不作图表分类；请用 --x 指定科室、类别等其他列，"
                         "或先按类别汇总（data_tool.py count）再出图")
    if y:
        ycols = []
        for name in re.split(r"[,，、]", y):
            name = name.strip()
            if not name:
                continue
            col = table.column(name)
            if col is None:
                raise ChartError(f"表中没有名为“{name}”的列；现有列：{'、'.join(c.name for c in table.columns)}")
            ycols.append(col)
    else:
        nums = [c for c in table.columns if c.kind in ("number", "rate", "avg")]
        if not nums:
            raise ChartError("找不到数值列，请用 --y 指定")
        first = nums[0]
        # 折线图默认画全部同单位的数值列（如 2026年、2025年 两条线）；其他图只画第一列
        if kind == "line":
            ycols = [c for c in nums if c.unit == first.unit and (c.kind == "rate") == (first.kind == "rate")]
        else:
            ycols = [first]
        rest = [c for c in nums if c not in ycols]
        if rest and warnings is not None:
            warnings.append(f"未指定 --y，图中只画了“{'、'.join(c.base for c in ycols)}”；表中还有"
                            + "、".join(f"“{c.base}”" for c in rest)
                            + "，需要时用 --y 指定（同期对比两列都写，如 --y 2026年,2025年；不同单位分开出图）")
    for col in ycols:
        if col.kind not in ("number", "rate", "avg"):
            raise ChartError(f"“{col.name}”不是数值列，不能出图")
    return xcol, ycols


TITLE_TAIL_RE = re.compile(r"(?:汇总表|统计表|情况表|一览表|明细表|统计|汇总|情况|表)$")
YEAR_RE = re.compile(r"^(?:19|20)\d{2}年?$")
# 只表示时段、不表示指标的列名（如"本年累计""上年同期"），图题里的指标改取表题
PERIODISH_RE = re.compile(r"^(?:(?:19|20)\d{2}年?|本年累计|上年同期累计|本年|上年|本年度|上年度|本月|上月|当月|"
                          r"上年同月|本期|上期|基期|累计|同期|上年同期|本季度?|上季度?)$")


def auto_title(table: dt.Table, xcol: dt.Column, ycols: List[dt.Column], kind: str, sort: str,
               top: Optional[int]) -> str:
    """没给 --title 时按"时段＋（各分类）＋指标"拼图题，如"2026年9月各科室出院人次（前10位）"。"""
    period = table.periods[0] if table.periods else ""
    names = [c.base for c in ycols]
    if all(PERIODISH_RE.match(n) for n in names):
        # 列名只表示时段（2026年、本年累计）：指标取表题中时段后面的部分
        metric = ""
        if table.title and period and period in table.title:
            metric = TITLE_TAIL_RE.sub("", table.title.split(period, 1)[1]).strip()
        metric = metric or "数据"
        if len(names) > 1 and all(YEAR_RE.match(n) for n in names):
            span = re.sub(r"^(?:19|20)\d{2}年度?", "", period)
            years = "与".join(n if n.endswith("年") else n + "年" for n in names)
            title = f"{years}{span}{metric}对比"
        else:
            title = f"{period}{metric}" + ("构成" if kind == "pie" else "") \
                + (f"（{'与'.join(names)}）" if len(names) > 1 else "")
    else:
        metric = "、".join(names)
        group = ""
        if kind not in ("line", "pie") and xcol.kind != "period" and not TIME_RE.search(xcol.base) \
                and len(xcol.base) <= 6 and not xcol.base.startswith("各"):
            group = "各" + xcol.base
        title = f"{period}{group}{metric}" + ("构成" if kind == "pie" else "")
    if top:
        title += {"desc": f"（前{top}位）", "asc": f"（后{top}位）"}.get(sort, f"（前{top}项）")
    return title


def label_places(series: List[List[Optional[Decimal]]], is_rate: bool) -> int:
    """数据标签的小数位：比率或有小数的保留 1 位，整数不带小数，与 Excel 图表的数字格式一致。"""
    if is_rate or any(v is not None and v != v.to_integral_value() for s in series for v in s):
        return 1
    return 0


def label_text(value: float, places: int) -> str:
    """PNG 数据标签：不用科学计数法（286450000 不写成 2.8645e+08）；空值不标。"""
    if value != value:  # NaN
        return ""
    return dt.fmt_num(dt.to_decimal(value), places)  # 十进制四舍五入，与正文一致


def build_data(table: dt.Table, xcol: dt.Column, ycols: List[dt.Column], warnings: List[str]) -> tuple:
    categories: List[str] = []
    series: List[List[Optional[Decimal]]] = [[] for _ in ycols]
    excluded = [r for r in table.rows if r.kind in ("total", "subtotal", "of_which", "avg")]
    if excluded:
        warnings.append("合计、小计、“其中”和平均行未入图：" + "、".join(
            table.display_label(r.label) or f"第{r.excel_row}行" for r in excluded[:dt.MAX_LIST]))
    for row in table.data_rows():
        cell = table.cell(row, xcol)
        label = cell.text if cell.text else f"第{row.excel_row}行"
        if xcol.sensitive or (table.key_col is not None and xcol.index == table.key_col):
            label = table.display_label(label)
        categories.append(label)
        for i, col in enumerate(ycols):
            value = table.cell(row, col)
            if value.is_number and "approx" not in value.flags:
                series[i].append(value.value)
            else:
                series[i].append(None)
                warnings.append(f"{label} 的“{col.name}”为“{value.text or '空'}”，图中留空，不按 0 画")
    if not categories:
        raise ChartError("没有可以出图的数据行")
    return categories, series


def validate(kind: str, xcol: dt.Column, ycols: List[dt.Column], categories: List[str],
             series: List[List[Optional[Decimal]]], unit: str, warnings: List[str]) -> None:
    units = {c.unit for c in ycols if c.unit}
    kinds = {"rate" if c.kind == "rate" else "amount" for c in ycols}
    if len(units) > 1 or len(kinds) > 1:
        raise ChartError("不同单位的指标不画在同一张图上（不用双坐标轴）：" +
                         "、".join(f"{c.name}（{c.unit or '无单位'}）" for c in ycols) + "，请分开出图")
    if kind == "pie":
        if len(ycols) > 1:
            raise ChartError("饼图只能画一个指标的构成")
        if any(v is not None and v < 0 for v in series[0]):
            raise ChartError("有负数，不能画饼图，请改用柱状图")
        if ycols[0].kind in ("rate", "avg"):
            raise ChartError("比率和平均数不能画饼图（各部分相加没有意义），请改用柱状图")
        blanks = [c for c, v in zip(categories, series[0]) if v is None]
        if blanks:
            raise ChartError("饼图表示构成，有空值时各部分占比会失真（空值不能按0画）：" + "、".join(blanks[:dt.MAX_LIST])
                             + "。请补齐数据后再画，或改用柱状图")
        if len(categories) > PIE_MAX:
            warnings.append(f"饼图有 {len(categories)} 类，超过 {PIE_MAX} 类时难以辨认，建议改用条形图，或把小类并为“其他”")
    if kind == "stacked" and any(c.kind in ("rate", "avg") for c in ycols):
        raise ChartError("比率和平均数不能堆积，请改用柱状图")
    if kind == "line":
        looks_time = bool(TIME_RE.search(xcol.name)) or \
            sum(1 for c in categories if PERIOD_VALUE_RE.search(c)) >= max(1, len(categories) // 2)
        if not looks_time:
            warnings.append("折线图适合按时间排列的数据；比较各单位请用柱状图或条形图")
    if kind == "bar" and len(categories) > 12:
        warnings.append(f"有 {len(categories)} 个类别，竖向柱状图标签会挤在一起，建议用条形图（--type barh）")
    if not unit and ycols[0].kind != "rate":
        warnings.append("没有计量单位，图上未标注单位；请用 --unit 补充")
    values = [v for s in series for v in s if v is not None]
    if values and min(values) < 0 and kind in ("bar", "barh", "stacked"):
        warnings.append("有负数，纵轴不从 0 开始截断，负值画在 0 以下")


def apply_sort(kind: str, categories: List[str], series: List[List[Optional[Decimal]]], sort: str, top: Optional[int],
               warnings: List[str]) -> tuple:
    if sort != "none":
        if kind == "line":
            warnings.append("折线图按时间顺序排列，已忽略 --sort")
        else:
            order = sorted(range(len(categories)),
                           key=lambda i: (series[0][i] is None, series[0][i] if series[0][i] is not None else 0),
                           reverse=False)
            if sort == "desc":
                with_value = [i for i in order if series[0][i] is not None][::-1]
                order = with_value + [i for i in order if series[0][i] is None]
            categories = [categories[i] for i in order]
            series = [[s[i] for i in order] for s in series]
    if top:
        if top < len(categories):
            warnings.append(f"只显示前 {top} 项（共 {len(categories)} 项），图注中写明")
        categories = categories[:top]
        series = [s[:top] for s in series]
    return categories, series


# ---------------------------------------------------------------------------
# Excel 原生图表
# ---------------------------------------------------------------------------

def _rich(text: str, size: int, bold: bool = False, font: str = "宋体"):
    from openpyxl.chart.text import RichText, Text  # type: ignore
    from openpyxl.drawing.text import (CharacterProperties, Font as DFont, Paragraph,  # type: ignore
                                       ParagraphProperties, RegularTextRun)
    props = CharacterProperties(sz=size, b=bold, latin=DFont(typeface=font), ea=DFont(typeface=font),
                                solidFill="000000")
    para = Paragraph(pPr=ParagraphProperties(defRPr=props), r=[RegularTextRun(rPr=props, t=text)])
    return Text(rich=RichText(p=[para]))


def _axis_text(size: int = 1000):
    from openpyxl.chart.text import RichText  # type: ignore
    from openpyxl.drawing.text import CharacterProperties, Font as DFont, Paragraph, ParagraphProperties  # type: ignore
    props = CharacterProperties(sz=size, latin=DFont(typeface="宋体"), ea=DFont(typeface="宋体"), solidFill="000000")
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=props), endParaRPr=props, r=[])])


def _title(text: str, size: int, bold: bool):
    from openpyxl.chart.title import Title  # type: ignore
    return Title(tx=_rich(text, size, bold, "黑体" if bold else "宋体"), overlay=False)


def write_xlsx(out: Path, kind: str, title: str, unit: str, source: str, note: str, xname: str,
               ynames: List[str], categories: List[str], series: List[List[Optional[Decimal]]],
               highlight: Optional[str], is_rate: bool, warnings: List[str]) -> Path:
    openpyxl = dt.need_openpyxl()
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference  # type: ignore
    from openpyxl.chart.label import DataLabelList  # type: ignore
    from openpyxl.chart.marker import DataPoint, Marker  # type: ignore
    from openpyxl.chart.shapes import GraphicalProperties  # type: ignore
    from openpyxl.drawing.line import LineProperties  # type: ignore
    from openpyxl.styles import Alignment, Border, Font, Side  # type: ignore

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "图表"
    ws["A1"] = title
    ws["A1"].font = Font(name="黑体", size=12, bold=True)
    ws["A2"] = f"单位：{unit}" if unit else ""
    ws["A2"].font = Font(name="宋体", size=10)
    header_row = 3
    ws.cell(header_row, 1, xname)
    for j, name in enumerate(ynames, start=2):
        ws.cell(header_row, j, name)
    thin = Side(style="thin", color="000000")
    for i, cat in enumerate(categories):
        r = header_row + 1 + i
        ws.cell(r, 1, cat)
        for j, s in enumerate(series, start=2):
            v = s[i]
            ws.cell(r, j, None if v is None else (int(v) if v == v.to_integral_value() else float(v)))
    last = header_row + len(categories)
    for r in range(header_row, last + 1):
        for c in range(1, len(ynames) + 2):
            cell = ws.cell(r, c)
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.font = Font(name="宋体", size=10, bold=(r == header_row))
            cell.alignment = Alignment(horizontal="center" if r == header_row else None, vertical="center")
    ws.cell(last + 2, 1, f"数据来源：{source}" if source else "数据来源：【待补：数据来源】")
    ws.cell(last + 2, 1).font = Font(name="宋体", size=9)
    if note:
        ws.cell(last + 3, 1, f"注：{note}")
        ws.cell(last + 3, 1).font = Font(name="宋体", size=9)
    ws.column_dimensions["A"].width = max(10, min(30, max(len(c) for c in categories + [xname]) * 2 + 2))
    for j in range(2, len(ynames) + 2):
        ws.column_dimensions[dt.col_letter(j - 1)].width = max(10, len(ynames[j - 2]) * 2 + 2)

    data_ref = Reference(ws, min_col=2, max_col=len(ynames) + 1, min_row=header_row, max_row=last)
    cats_ref = Reference(ws, min_col=1, min_row=header_row + 1, max_row=last)
    number_format = "0.0" if label_places(series, is_rate) == 1 else "0"

    if kind == "pie":
        chart = PieChart()
        chart.add_data(Reference(ws, min_col=2, min_row=header_row, max_row=last), titles_from_data=True)
        chart.set_categories(cats_ref)
        s = chart.series[0]
        for i in range(len(categories)):
            pt = DataPoint(idx=i)
            pt.graphicalProperties = GraphicalProperties(solidFill=GRAYS[i % len(GRAYS)])
            pt.graphicalProperties.line.solidFill = "FFFFFF"
            s.dPt.append(pt)
        s.dLbls = DataLabelList()
        s.dLbls.showCatName = True
        s.dLbls.showPercent = True
        s.dLbls.showVal = False
        s.dLbls.showSerName = False
        s.dLbls.showLeaderLines = True
        s.dLbls.showLegendKey = False
        s.dLbls.numFmt = "0.0%"
        s.dLbls.txPr = _axis_text(900)
        chart.legend = None
    elif kind == "line":
        chart = LineChart()
        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(cats_ref)
        for i, s in enumerate(chart.series):
            color, dash, symbol = LINE_STYLES[i % len(LINE_STYLES)]
            s.graphicalProperties.line = LineProperties(solidFill=color, w=22225, prstDash=dash)
            s.marker = Marker(symbol=symbol, size=6)
            s.marker.graphicalProperties = GraphicalProperties(solidFill=color)
            s.marker.graphicalProperties.line.solidFill = color
            s.smooth = False
    else:
        chart = BarChart()
        chart.type = "bar" if kind == "barh" else "col"
        if kind == "barh":
            # 条形图按表中顺序从上到下排列，数值轴留在下方
            chart.x_axis.scaling.orientation = "maxMin"
            chart.y_axis.crosses = "max"
        if kind == "stacked":
            chart.grouping = "stacked"
            chart.overlap = 100
        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(cats_ref)
        chart.gapWidth = 80
        for i, s in enumerate(chart.series):
            fill = GRAYS[i % len(GRAYS)]
            s.graphicalProperties = GraphicalProperties(solidFill=fill)
            s.graphicalProperties.line.solidFill = "000000"
            s.graphicalProperties.line.width = 6350
        if highlight and len(series) == 1:
            found = [i for i, c in enumerate(categories) if c == highlight or highlight in c]
            if not found:
                warnings.append(f"没有找到要突出的“{highlight}”")
            else:
                s = chart.series[0]
                s.graphicalProperties.solidFill = "BFBFBF"
                for i in found:
                    pt = DataPoint(idx=i)
                    pt.graphicalProperties = GraphicalProperties(solidFill="262626")
                    pt.graphicalProperties.line.solidFill = "000000"
                    s.dPt.append(pt)
        elif highlight:
            warnings.append("--highlight 只用于单个指标的图，已忽略")

    if kind != "pie":
        chart.dataLabels = DataLabelList()
        chart.dataLabels.showVal = True
        chart.dataLabels.showSerName = False
        chart.dataLabels.showCatName = False
        chart.dataLabels.showLegendKey = False
        chart.dataLabels.numFmt = number_format
        chart.dataLabels.txPr = _axis_text(900)
        for axis in (chart.x_axis, chart.y_axis):
            axis.delete = False
            axis.txPr = _axis_text(1000)
        values = [v for s in series for v in s if v is not None]
        if values and min(values) >= 0:
            chart.y_axis.scaling.min = 0
        chart.y_axis.numFmt = "General"
        chart.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="D9D9D9", w=6350))
        if unit:
            try:
                chart.y_axis.title = _title(f"单位：{unit}", 1000, False)
            except Exception:  # noqa: BLE001  个别 openpyxl 版本不支持富文本标题时退回纯文本
                chart.y_axis.title = f"单位：{unit}"
        if len(series) > 1:
            chart.legend.position = "b"
        else:
            chart.legend = None
    try:
        chart.title = _title(title, 1400, True)
    except Exception:  # noqa: BLE001
        chart.title = title
    chart.style = 2
    chart.width = 18 if len(categories) <= 12 else 24
    chart.height = 10 if kind != "barh" else max(10, min(24, 0.7 * len(categories) + 3))
    anchor = f"{dt.col_letter(len(ynames) + 2)}1"
    ws.add_chart(chart, anchor)
    out = dt.unique_path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))
    return out


# ---------------------------------------------------------------------------
# PNG（可选）
# ---------------------------------------------------------------------------

def write_png(out: Path, kind: str, title: str, unit: str, source: str, note: str, ynames: List[str],
              categories: List[str], series: List[List[Optional[Decimal]]], highlight: Optional[str],
              warnings: List[str], is_rate: bool = False) -> Optional[Path]:
    try:
        import matplotlib  # type: ignore
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt  # type: ignore
        from matplotlib import font_manager  # type: ignore
    except ImportError:
        warnings.append("未安装 matplotlib，未生成 PNG；可在 Excel 或 WPS 中右键图表另存为图片，"
                        "需要时运行 pip install matplotlib")
        return None
    available = {f.name for f in font_manager.fontManager.ttflist}
    font = next((f for f in CJK_FONTS if f in available), None)
    if font:
        plt.rcParams["font.sans-serif"] = [font]
        plt.rcParams["font.family"] = "sans-serif"
    else:
        warnings.append("系统中没有找到中文字体，PNG 中的汉字可能显示为方框；Excel 图表不受影响")
    plt.rcParams["axes.unicode_minus"] = False
    fig, ax = plt.subplots(figsize=(8, 4.8) if kind != "barh" else (8, max(4.8, 0.35 * len(categories) + 1.5)),
                           dpi=200)
    vals = [[float(v) if v is not None else float("nan") for v in s] for s in series]
    places = label_places(series, is_rate)
    n = len(series)
    idx = list(range(len(categories)))
    if kind == "pie":
        data = [v if v == v else 0 for v in vals[0]]
        colors = ["#" + GRAYS[i % len(GRAYS)] for i in idx]
        ax.pie(data, labels=categories, colors=colors, autopct="%1.1f%%", startangle=90, counterclock=False,
               wedgeprops={"edgecolor": "white"}, textprops={"fontsize": 9})
        ax.axis("equal")
    elif kind == "line":
        for i, s in enumerate(vals):
            color, dash, _ = LINE_STYLES[i % len(LINE_STYLES)]
            style = {"solid": "-", "dash": "--", "sysDot": ":", "lgDash": "--", "dashDot": "-."}[dash]
            ax.plot(idx, s, linestyle=style, color="#" + color, marker="osd^x"[i % 5], label=ynames[i])
            for x, y in zip(idx, s):
                if y == y:
                    ax.annotate(label_text(y, places), (x, y), textcoords="offset points", xytext=(0, 5), ha="center", fontsize=8)
        ax.set_xticks(idx)
        ax.set_xticklabels(categories)
    else:
        width = 0.8 / (1 if kind == "stacked" else n)
        bottom = [0.0] * len(categories)
        for i, s in enumerate(vals):
            colors = ["#" + GRAYS[i % len(GRAYS)]] * len(categories)
            if highlight and n == 1:
                colors = ["#262626" if (c == highlight or highlight in c) else "#BFBFBF" for c in categories]
            pos = idx if kind == "stacked" else [x - 0.4 + width * (i + 0.5) for x in idx]
            plot = ax.barh if kind == "barh" else ax.bar
            kwargs = {"left" if kind == "barh" else "bottom": bottom} if kind == "stacked" else {}
            bars = plot(pos, [v if v == v else 0 for v in s], width, color=colors, edgecolor="black",
                        linewidth=0.5, hatch=HATCHES[i % len(HATCHES)], label=ynames[i], **kwargs)
            if kind != "stacked":
                ax.bar_label(bars, labels=[label_text(v, places) for v in s], fontsize=8, padding=2)
            else:
                bottom = [b + (v if v == v else 0) for b, v in zip(bottom, s)]
        if kind == "barh":
            ax.set_yticks(idx)
            ax.set_yticklabels(categories)
            ax.invert_yaxis()
        else:
            ax.set_xticks(idx)
            ax.set_xticklabels(categories, rotation=0 if len(categories) <= 8 else 45,
                               ha="center" if len(categories) <= 8 else "right")
    if kind != "pie":
        for side in ("top", "right"):
            ax.spines[side].set_visible(False)
        (ax.set_xlabel if kind == "barh" else ax.set_ylabel)(f"单位：{unit}" if unit else "")
        ax.grid(axis="x" if kind == "barh" else "y", color="#D9D9D9", linewidth=0.5)
        ax.set_axisbelow(True)
        if n > 1:
            ax.legend(frameon=False, loc="upper center", bbox_to_anchor=(0.5, -0.12), ncol=n)
    ax.set_title(title, fontsize=13, fontweight="bold")
    foot = f"数据来源：{source}" if source else "数据来源：【待补：数据来源】"
    if note:
        foot += f"　注：{note}"
    fig.text(0.01, 0.01, foot, fontsize=8, ha="left", va="bottom")
    fig.tight_layout(rect=(0, 0.04, 1, 1))
    png = dt.unique_path(out.with_suffix(".png"))
    fig.savefig(str(png), facecolor="white")
    plt.close(fig)
    return png


# ---------------------------------------------------------------------------
# 命令行
# ---------------------------------------------------------------------------

def build_parser() -> argparse.ArgumentParser:
    parser = dt.localize(argparse.ArgumentParser(
        prog="chart_tool.py",
        description="生成公文风格的 Excel 原生图表（黑白灰，标注单位和数据来源）。本地生成，不联网。",
        epilog="退出码：0 成功；2 输入错误。",
        formatter_class=dt.ChineseHelpFormatter,
        add_help=False,
    ))
    parser.add_argument("file", help="数据表格（.csv 或 .xlsx）")
    parser.add_argument("--type", required=True, choices=("bar", "barh", "line", "pie", "stacked"),
                        help="bar 柱状图；barh 条形图；line 折线图；pie 饼图；stacked 堆积柱状图")
    parser.add_argument("--sheet", help="工作表名称")
    parser.add_argument("--header-row", type=int, help="表头所在行号（从 1 起）")
    parser.add_argument("--x", help="分类列（默认取第一列文字列）")
    parser.add_argument("--y", help="数值列，多个用逗号分隔（默认：折线图取全部同单位数值列，其他图取第一个数值列）")
    parser.add_argument("--title", help="图题（默认按“时段＋指标”自动生成，请核对）")
    parser.add_argument("--unit", help="计量单位（默认取列的单位）")
    parser.add_argument("--source", help="数据来源（默认取表注中的“数据来源”）")
    parser.add_argument("--note", help="图注")
    parser.add_argument("--sort", choices=("none", "desc", "asc"), default="none", help="按第一个指标排序（折线图不排序）")
    parser.add_argument("--top", type=int, help="只显示排序后的前 N 项")
    parser.add_argument("--highlight", help="突出显示的类别（如本单位），只用于单个指标的柱状图、条形图")
    parser.add_argument("--out", help="输出 .xlsx 路径（默认为“<输入文件名>-图表.xlsx”），已存在时加 -v2")
    parser.add_argument("--png", action="store_true", help="另出 PNG（需要 matplotlib，没有时只出 Excel）")
    parser.add_argument("--json", action="store_true", help="输出 JSON")
    parser.add_argument("--marking-ok", action="store_true", help=dt.MARKING_OK_HELP)
    return parser


def main(argv: Optional[Sequence[str]] = None) -> int:
    dt.setup_console()
    parser = build_parser()
    try:
        args = parser.parse_args(argv)
    except SystemExit as exc:
        return 0 if exc.code == 0 else 2
    warnings: List[str] = []
    try:
        dt.need_openpyxl()
        sheets = dt.load_sheets(args.file, args.sheet)
        dt.guard_marking(sheets, args.marking_ok)
        table = dt.analyze(sheets[0], "表A", args.header_row)
        dt.inspect_table(table, args.marking_ok)
        p0 = [i for i in table.issues if i.level == "P0" and i.category != "个人信息"]
        if p0:
            warnings.append(f"表中有 {len(p0)} 个 P0 问题（如合计不符、单位不一致），出图前请先用 data_tool.py inspect 查清")
        xcol, ycols = pick_columns(table, args.x, args.y, args.type, warnings)
        categories, series = build_data(table, xcol, ycols, warnings)
        unit = args.unit if args.unit is not None else (ycols[0].unit or "")
        validate(args.type, xcol, ycols, categories, series, unit, warnings)
        if args.top is not None and args.top <= 0:
            raise ChartError("--top 须大于 0")
        categories, series = apply_sort(args.type, categories, series, args.sort, args.top, warnings)
        title = args.title
        if not title:
            title = auto_title(table, xcol, ycols, args.type, args.sort, args.top)
            warnings.append(f"未给 --title，图题按“时段＋指标”自动生成为“{title}”，请核对后用 --title 修改")
        source = args.source if args.source is not None else table.source_note
        if not source:
            warnings.append("未提供数据来源，表格和图下已写“数据来源：【待补：数据来源】”，请补充")
        ynames = [c.base for c in ycols]
        out = Path(args.out) if args.out else Path(args.file).with_name(Path(args.file).stem + "-图表.xlsx")
        if out.suffix.lower() != ".xlsx":
            out = out.with_suffix(".xlsx")
        is_rate = any(c.kind == "rate" for c in ycols)
        saved = write_xlsx(out, args.type, title, unit, source, args.note or "", xcol.base, ynames, categories,
                           series, args.highlight, is_rate, warnings)
        png = None
        if args.png:
            png = write_png(saved, args.type, title, unit, source, args.note or "", ynames, categories, series,
                            args.highlight, warnings, is_rate)
    except dt.InputError as exc:
        print(f"错误：{exc}", file=sys.stderr)
        return 2
    except OSError as exc:
        print(f"错误：无法读写文件（{exc.strerror or exc}）：{exc.filename or ''}。"
              f"请检查路径是否存在、文件是否被 Excel 占用；Windows 下路径超过 260 个字符时请换到短一些的目录",
              file=sys.stderr)
        return 2
    result: Dict[str, Any] = {
        "tool": "chart_tool", "version": VERSION, "out": str(saved), "png": str(png) if png else None,
        "type": args.type, "title": title, "unit": unit, "source": source or "",
        "categories": categories,
        "series": [{"name": n, "values": [None if v is None else float(v) for v in s]} for n, s in zip(ynames, series)],
        "warnings": warnings,
    }
    if args.json:
        print(json.dumps(result, ensure_ascii=False, indent=2))
    else:
        print(f"已生成：{saved}" + (f"；PNG：{png}" if png else ""))
        print(f"图题：{title}；单位：{unit or '无'}；数据来源：{source or '【待补】'}；类别 {len(categories)} 个")
        for w in warnings:
            print(f"提示：{w}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
