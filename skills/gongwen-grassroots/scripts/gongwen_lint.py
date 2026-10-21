# -*- coding: utf-8 -*-
"""公文机检引擎（gongwen_lint.py）。

源文件位于仓库 shared/scripts/，各 skill 中的副本由 tools/sync_shared.py 生成，请在源文件处修改。
接口见 docs/开发规范.md 第 8 节，其他 skill 只依赖命令行接口和 JSON 输出格式：

    python scripts/gongwen_lint.py <文件> [--doc-type 文书类型] [--checks 类别,类别] [--json] [--strict]
                                   [--banned 禁用词.md] [--allow 规则ID,规则ID] [--evidence-allow 词,词] [--baseline [baseline.json]]

- 只依赖标准库；读取 .docx 时需要 python-docx。
- 规则数据：同目录下的 lint_rules.json（正则规则 + 函数规则登记）和 style_patterns.json（AI 味、防御性词表）。
- 遇到单独成行的【起草说明】即停止读取（兼容“## 【起草说明】”“**【起草说明】**”，与 fact_diff.py、gongwen-format 一致）。
- --banned：读取本单位禁用词表（.gongwen/禁用词.md），“禁用”命中报 P1（BAN-001），“慎用”命中报 P2（BAN-002）。
- --allow：本单位制度明文规定的写法（如企业用方括号的发文字号），列出的规则命中时降为 P2，不计入 --strict。
- --evidence-allow：评价词白名单。EVD 类检查（评价词与证据强度匹配）对其中的词不再报告；固定短语白名单在 lint_rules.json 的 evidence.whitelist。
- --baseline：对已识别文种，把篇幅、一级标题数、“一是”组数、顿号密度与公开公文语料基线（references/_shared/style/baseline.json，
  仓库内为 shared/style/baseline.json）比较，偏离 P10～P90 区间的报 P2 提示（BSL-001～004），只提示不判错；找不到基线文件时跳过并提示。
- 要素完整性（ELM-001～007，归 format 类）：按文种查标题、主送、结语、落款、成文日期、附件说明与正文引用、纪要出席列席；
  【待补：…】占位视为作者已注意到，报 P2。评价证据（EVD-001～003，归 style 类）：强评价（显著、大幅……）同句或相邻句无数字、对比报 P1，
  中评价无任何事实报 P2；只查回顾性文书（总结、报告、通报等），计划、方案、讲话稿不查。
- 退出码：0 正常；--strict 且存在 P0 时为 1；输入错误为 2。
"""
from __future__ import annotations

import argparse
import json
import os
import re
import sys
from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from pathlib import Path
from typing import Any, Callable, Dict, List, Optional, Pattern, Sequence, Set, Tuple

VERSION = "1.0.0"
SCRIPT_DIR = Path(__file__).resolve().parent
DEFAULT_RULES = SCRIPT_DIR / "lint_rules.json"
DEFAULT_PATTERNS = SCRIPT_DIR / "style_patterns.json"

# 检查类别（顺序即输出顺序），与开发规范第 8 节一致
CHECK_ORDER = (
    "residue", "punct", "number", "format", "structure",
    "style", "defensive", "privacy", "legal", "jianfu", "marker", "length",
)
CHECK_NAMES = {
    "residue": "残留",
    "punct": "标点",
    "number": "数字",
    "format": "格式要素",
    "structure": "结构",
    "style": "AI味",
    "defensive": "防御性",
    "privacy": "隐私与保密",
    "legal": "法规",
    "jianfu": "减负",
    "marker": "待补标记",
    "length": "篇幅",
}
LEVELS = ("P0", "P1", "P2")
LEVEL_RANK = {"P0": 0, "P1": 1, "P2": 2}

# 统一词表中的文书类型（开发规范第 6 节）。扩充类型的规则适用范围：
# - 贯彻落实方案、整改方案：同实施方案（篇幅上限、部署类文风规则）；整改方案的报送、考核要求来自巡视巡察和审计整改规定，
#   不查 JFU-002、JFU-004；
# - 党课讲稿、研讨发言：同讲话稿，放宽口语化句式；
# - 简报、新闻稿、公众号推文：同信息；
# - 整改报告、自查报告、自评报告：同报告（报告夹带请示、主送、签发人检查），回顾性材料不查转嫁责任类减负规则；
# - 议题材料、会议记录、整改报告、自查报告：字段冒号多，不查冒号密度；
# - 对照检查材料：不适用请示、报告的结构检查，查空泛问题表述（AI-053）。
DOC_TYPES = (
    "决议", "决定", "命令", "公报", "公告", "通告", "意见", "通知", "通报", "报告", "请示",
    "批复", "议案", "函", "纪要", "签报",
    "总结", "计划", "工作要点", "汇报材料", "调研报告", "经验材料", "述职报告", "实施方案",
    "讲话稿", "信息", "其他",
    "对照检查材料", "党课讲稿", "心得体会", "研讨发言", "会议记录", "议题材料", "简报", "新闻稿",
    "整改方案", "整改报告", "自查报告", "贯彻落实方案", "自评报告", "公众号推文",
)
# 报告类：报告夹带请示（STR-004）、主送（STR-006）等检查按“报告”处理
REPORT_TYPES = {"报告", "整改报告", "自查报告", "自评报告"}
# 法定文种（条例第八条）：标题文种与 --doc-type 不一致的提示只在两者都是法定文种时给出
STATUTORY_TYPES = {
    "决议", "决定", "命令", "公报", "公告", "通告", "意见", "通知", "通报", "报告", "请示", "批复", "议案", "函", "纪要",
}
# 有主送机关的常见文种：标题后一行缺冒号时才判为主送机关缺冒号
ZHUSONG_TYPES = {"通知", "通报", "报告", "请示", "批复", "函", "意见", "决定", "议案", "签报"}
# 篇幅上限适用的下行文件（《整治形式主义为基层减负若干规定》第2条）
LENGTH_LIMIT_TYPES = {
    "决议", "决定", "命令", "公报", "公告", "通告", "意见", "通知", "通报", "批复",
    "实施方案", "工作要点", "计划", "贯彻落实方案", "整改方案",
}
# 标题 → 扩充文书类型（先于 TITLE_DOC_SUFFIX 匹配；“请示报告”留给 STR-010）
TITLE_DOC_PATTERNS = (
    (re.compile(r"整改[^，。；]{0,16}(?<!请示)报告$"), "整改报告"),
    (re.compile(r"自查[^，。；]{0,12}(?<!请示)报告$"), "自查报告"),
    (re.compile(r"自评[^，。；]{0,12}(?<!请示)报告$"), "自评报告"),
    (re.compile(r"整改[^，。；]{0,12}方案$"), "整改方案"),
    (re.compile(r"贯彻落实[^，。；]{0,40}方案$"), "贯彻落实方案"),
    (re.compile(r"(?:对照检查|检视剖析)[^，。；]{0,6}(?:材料|发言提纲)$"), "对照检查材料"),
)
# 标题结尾 → 文书类型（按顺序匹配，长的在前）
TITLE_DOC_SUFFIX = (
    ("对照检查材料", "对照检查材料"), ("检视剖析材料", "对照检查材料"), ("党课讲稿", "党课讲稿"), ("党课", "党课讲稿"),
    ("心得体会", "心得体会"), ("学习体会", "心得体会"), ("研讨发言提纲", "研讨发言"), ("研讨发言", "研讨发言"),
    ("会议记录", "会议记录"), ("议题材料", "议题材料"), ("新闻通稿", "新闻稿"), ("新闻稿", "新闻稿"),
    ("简报", "简报"), ("推文", "公众号推文"),
    ("请示报告", "请示"), ("调研报告", "调研报告"), ("述责述廉报告", "述职报告"),
    ("述职述廉报告", "述职报告"), ("述职报告", "述职报告"), ("实施方案", "实施方案"),
    ("工作方案", "实施方案"), ("方案", "实施方案"), ("工作要点", "工作要点"), ("要点", "工作要点"),
    ("汇报材料", "汇报材料"), ("经验材料", "经验材料"), ("工作总结", "总结"), ("总结", "总结"),
    ("调查报告", "调研报告"), ("工作计划", "计划"), ("计划", "计划"), ("讲话稿", "讲话稿"), ("讲话", "讲话稿"),
    ("发言", "讲话稿"), ("致辞", "讲话稿"), ("请示", "请示"), ("报告", "报告"), ("通知", "通知"),
    ("通报", "通报"), ("通告", "通告"), ("公告", "公告"), ("公报", "公报"), ("批复", "批复"),
    ("意见", "意见"), ("决定", "决定"), ("决议", "决议"), ("命令", "命令"), ("议案", "议案"),
    ("纪要", "纪要"), ("签报", "签报"), ("复函", "函"), ("函", "函"),
)

# 【起草说明】分界行：去掉行首空白后以【起草说明】开头；兼容 Markdown 标题、加粗包裹的写法。
# 与 fact_diff.py 的 DRAFT_NOTE_RE、gongwen-format 的 gw_layout.is_draft_note_line 保持一致。
DRAFT_NOTE_RE = re.compile(r"^\s*(?:#{1,6}\s*)?(?:(?:\*\*|__)\s*)?【起草说明】")
PUNCT_TAIL = "。，；：、！？.,;:!?"

# ---------------------------------------------------------------------------
# 版头、标题、落款、版记的识别
# ---------------------------------------------------------------------------
QIANBAO_HEADERS = ("签报", "呈批件")  # 稿首表头
FENHAO_RE = re.compile(r"^\s*\d{6}\s*$")
MIJI_LINE_RE = re.compile(r"^\s*(?:绝密|机密|秘密)(?:\s*★.*)?\s*$")
INNER_MARK_RE = re.compile(
    r"^\s*[（(【\[]?\s*(?:内部资料|内部文件|内部材料|内部参阅|工作秘密|内部使用|不得外传|注意保密|内部传阅)"
    r"[\s　，,、]*(?:注意保存|妥善保管|不得外传|请勿外传|注意保密)?\s*[）)】\]]?\s*$"
)
# 【待补：】【待核实：】标记：正文唯一允许的提示内容，不参与冒号、文风类检查
MARKER_RE = re.compile(r"【(?:待补|待核实)[：:][^】\n]*】")
# summary.markers 计数用（与 MRK-001、MRK-002 的写法一致：全角冒号）
ZHUSONG_MARKER_RE = re.compile(r"【(?:待补|待核实)[^】]*】")  # 主送行里的标记：冒号、括号写法不限
TODO_MARK_RE = re.compile(r"【待补：[^】\n]*】")
VERIFY_MARK_RE = re.compile(r"【待核实：[^】\n]*】")
# 要素完整性、评价证据检查用：任何写法的【待补…】【待核实…】占位
ELM_MARKER_RE = re.compile(r"【(?:待补|待核实)[^】\n]*】")
JINJI_RE = re.compile(r"^\s*(?:特急|加急|特提|平急)\s*$")
BIAOZHI_RE = re.compile(r"^\s*\S{2,40}(?:文件|命令|令)\s*$")
WENHAO_RE = re.compile(
    r"[一-鿿A-Za-z]{1,15}\s*[〔\[［(（【﹝{｛〖]\s*(?:19|20)?\d{2}\s*[〕\]］)）】﹞}｝〗]\s*第?\s*\d+\s*号"
)
QIANFA_RE = re.compile(r"签发人\s*[：:]")
HEADING_START_RE = re.compile(r"^\s*(?:[一二三四五六七八九十]{1,3}\s*[、，,.．]|[（(]\s*[一二三四五六七八九十\d]{1,3}\s*[）)]|\d{1,2}\s*[.．、])")
TITLE_END_RE = re.compile(
    r"(?:决议|决定|命令|公报|公告|通告|意见|通知|通报|报告|请示|批复|议案|函|纪要|签报|总结|计划|要点|材料|方案"
    r"|讲话|发言|致辞|信息|简报|办法|规定|细则|规则|制度|章程|预案|规划|说明|汇报|提纲"
    r"|讲稿|讲话稿|发言稿|新闻稿|通稿|心得体会|学习体会|会议记录|推文|党课)$"
)
DATE_ONLY_RE = re.compile(
    r"^\s*[\d〇○零一二三四五六七八九OＯ]{4}\s*年\s*[\d一二三四五六七八九十]{1,3}\s*月\s*[\d一二三四五六七八九十]{1,3}\s*日\s*$"
)
CN_DATE_ONLY_RE = re.compile(
    r"^\s*[〇○零一二三四五六七八九OＯ]{4}\s*年\s*[一二三四五六七八九十]{1,3}\s*月\s*[一二三四五六七八九十]{1,3}\s*日\s*$"
)
BANJI_RE = re.compile(r"^\s*(?:抄送|抄报|主送)\s*[：:]|印发\s*$")
ATTACH_DESC_RE = re.compile(r"^\s*附件\s*[：:]")
ATTACH_PAGE_RE = re.compile(r"^\s*附件\s*\d{0,2}\s*$")
ORG_SUFFIX_RE = re.compile(
    r"(?:单位|部门|政府|委员会|党委|党组|机关|学校|医院|公司|集团|办公室|办公厅|办事处|局|厅|委|办|处|科|室|院|校"
    r"|中心|街道|乡镇|社区|村|镇|乡|区|县|市|省|部|会|组|所|站|队)$"
)

# 面向群众的文书：主送或称呼行写给居民、家长等，或标题为公示、告知书等
PUBLIC_AUDIENCE_RE = re.compile(r"居民|村民|群众|市民|家长|业主|住户|商户|患者|游客|乘客|消费者|朋友们|村（居）民")
PUBLIC_TITLE_RE = re.compile(r"(?:公示|告知书|一封信|倡议书|温馨提示|告居民书|公开信|须知|告示)$")
# 法规类规则的“已写明依据”判断：同一句中命中位置之前出现这些字样时不报
LEGAL_BASIS_RE = re.compile(
    r"依法|依照|《[^》]{2,40}(?:法|条例|规定|办法)》|(?:根据|按照)[^。；！？]{0,30}(?:法|条例)"
)
# 减负类规则的否定语境（skip_if_negated）：命中词之前 8 字内出现这些字样时不报
NEGATION_RE = re.compile(
    r"不要求|不得|禁止|严禁|不再|杜绝|避免|减少|取消|不搞|不以|不把|不将|无需|不需要?|不必|防止|纠正|不准|不允许"
)

# 层次序数：level, 正则（第 1 组为序号）
HEADING_PATTERNS = (
    # “、”后可以直接是数字（如“三、2027年工作打算”）；其他误用的点号后紧跟数字的不算标题。
    # 否定前瞻：后面紧跟“X年级”“X季度”等的是并列写法（“一、二年级学生”“三、四季度要抓好……”），不是标题。
    # 例外：“、”前后是同一个数（“四、四季度安排”“二、二年级工作”）不可能是并列写法，按一级标题识别。
    (1, re.compile(r"^\s*([一二三四五六七八九十]{1,3})\s*(?:"
                   r"、(?:(?=\1(?:年级|产业|季度|等奖|类|级))|(?![一二三四五六七八九十]+(?:年级|产业|季度|等奖|类|级)))"
                   r"|[，,.．:：](?!\d)(?![一二三四五六七八九十]+(?:年级|产业|季度|等奖|类|级)))")),
    (2, re.compile(r"^\s*[（(]\s*([一二三四五六七八九十]{1,3})\s*[）)]")),
    (3, re.compile(r"^\s*(\d{1,2})\s*(?:[.．](?!\d)|、(?!\d))")),
    (4, re.compile(r"^\s*[（(]\s*(\d{1,2})\s*[）)]")),
)
CN_DIGIT = {"零": 0, "〇": 0, "一": 1, "二": 2, "两": 2, "三": 3, "四": 4, "五": 5, "六": 6, "七": 7, "八": 8, "九": 9}

# 常见姓氏（用于隐私检查中判断“像姓名”，减少误报）
SURNAMES = set(
    "王李张刘陈杨黄赵吴周徐孙马朱胡郭何高林罗郑梁谢宋唐许韩冯邓曹彭曾肖田董袁潘于蒋蔡余杜叶程苏魏吕丁任沈"
    "姚卢姜崔钟谭陆汪范金石廖贾夏韦付方白邹孟熊秦邱江尹薛闫段雷侯龙史陶黎贺顾毛郝龚邵万钱严覃武戴莫孔向汤常"
    "温康施文牛樊葛邢安齐易乔伍庞颜倪庄聂章鲁岳翟殷詹申欧耿关兰焦俞左柳甘祝包宁尚符舒阮柯纪梅童凌毕单季裴霍"
    "涂成苗谷盛曲翁冉骆蓝路游辛靳管柴蒙鲍华喻祁蒲房滕屈饶解牟艾尤阳时穆农司卓古吉缪简车项连芦麦褚娄窦戚岑景党"
    "宫费卜冷晏席卫米柏宗瞿桂全佟应臧闵苟邬边卞姬师和仇栾隋商刁沙荣巫寇桑郎甄丛仲虞敖巩明佘池查麻苑迟邝"
)
ROLE_WORDS = ("患者", "病人", "患儿", "学生", "幼儿", "考生", "当事人", "居民", "村民", "职工", "老人", "家属", "儿童",
              "服务对象", "群众", "同志", "护士", "医生", "教师", "老师")
MASK_CHARS = set("某×X*＊")


class InputError(Exception):
    """输入错误：文件不存在、格式不支持、依赖缺失等，退出码 2。"""


def long_path(path: Path) -> Path:
    """Windows 下完整路径接近 260 字符上限时加长路径前缀，避免“文件不存在”误报（与 overlap_check.py 一致）。"""
    if os.name != "nt":
        return path
    try:
        full = str(path.resolve())
    except OSError:
        return path
    prefix = "\\\\?\\"
    if len(full) < 240 or full.startswith(prefix):
        return path
    if full.startswith("\\\\"):
        return Path(prefix + "UNC\\" + full[2:])
    return Path(prefix + full)


# ---------------------------------------------------------------------------
# 数据结构
# ---------------------------------------------------------------------------

@dataclass
class Para:
    index: int          # 在 paras 中的序号（从 0 开始）
    line: int           # 原文行号；docx 为段落序号（从 1 开始）
    text: str
    in_table: bool = False


@dataclass
class Issue:
    id: str
    level: str
    check: str
    category: str
    gate: str
    line: int
    excerpt: str
    message: str
    suggestion: str
    basis: str
    context: str = ""
    start: int = -1
    end: int = -1
    count: Optional[int] = None
    para_index: int = -1

    def to_dict(self) -> Dict[str, Any]:
        data: Dict[str, Any] = {
            "id": self.id,
            "level": self.level,
            "category": self.category,
            "line": self.line,
            "excerpt": self.excerpt,
            "message": self.message,
            "suggestion": self.suggestion,
            "basis": self.basis,
            "check": self.check,
            "gate": self.gate,
        }
        if self.context:
            data["context"] = self.context
        if self.count is not None:
            data["count"] = self.count
        return data


@dataclass
class RegexRule:
    id: str
    check: str
    category: str
    gate: str
    level: str
    regex: Pattern[str]
    message: str
    suggestion: str
    basis: str
    doc_types: Tuple[str, ...]
    exempt_doc_types: Tuple[str, ...]
    yield_to: Tuple[str, ...]
    once: bool
    skip_line: Optional[Pattern[str]]
    fixes: Tuple[Tuple[Pattern[str], str], ...]
    skip_table: bool = False
    zone: str = ""
    also_public: bool = False
    public_only: bool = False
    skip_if_basis: bool = False
    skip_in_quotes: bool = False
    skip_if_negated: bool = False


@dataclass
class StylePattern:
    id: str
    group: str
    category: str
    regex: Pattern[str]
    level: str
    message: str
    suggestion: str
    density: Optional[float]
    doc_types: Tuple[str, ...]
    exempt_doc_types: Tuple[str, ...]
    exclude: Tuple[Pattern[str], ...] = ()


@dataclass
class BannedTerm:
    """本单位禁用词表中的一条（--banned）。"""
    term: str
    regex: Pattern[str]
    banned: bool            # True：禁用（BAN-001，P1）；False：慎用（BAN-002，P2）
    reason: str = ""
    suggestion: str = ""
    scope: str = ""
    source: str = ""        # 禁用词文件名，写入 basis


def applies_to(doc_type: str, doc_types: Sequence[str], exempt: Sequence[str]) -> bool:
    """doc_types 非空时为正向限定；exempt_doc_types 为豁免。"""
    if doc_types and doc_type not in doc_types:
        return False
    return doc_type not in exempt


def nonspace_len(text: str) -> int:
    return len(re.sub(r"\s", "", text))


def cn_to_int(text: str) -> int:
    """把“一”“十二”“二十”等汉字序数或阿拉伯数字转为整数。"""
    text = text.strip()
    if text.isdigit():
        return int(text)
    if "十" in text:
        head, _, tail = text.partition("十")
        tens = CN_DIGIT.get(head, 1) if head else 1
        ones = CN_DIGIT.get(tail, 0) if tail else 0
        return tens * 10 + ones
    value = 0
    for ch in text:
        value = value * 10 + CN_DIGIT.get(ch, 0)
    return value


# 行首的 Markdown 标题、加粗符号：结构识别时先去掉（残留类规则仍按原文报告）
MD_HEADING_PREFIX_RE = re.compile(r"^\s*(?:#{1,6}\s*)?(?:\*\*|__)?")


def heading_of(text: str) -> Optional[Tuple[int, int]]:
    """识别行首层次序数，返回 (层级, 序号)。“**三、时间安排**”“## 一、总体要求”按一级标题识别。"""
    text = MD_HEADING_PREFIX_RE.sub("", text, count=1)
    for level, regex in HEADING_PATTERNS:
        match = regex.match(text)
        if match:
            return level, cn_to_int(match.group(1))
    return None


def clip(text: str, limit: int = 60) -> str:
    text = text.strip()
    return text if len(text) <= limit else text[:limit]


def context_of(text: str, start: int, end: int, width: int = 12) -> str:
    return text[max(0, start - width):min(len(text), end + width)].strip()


# ---------------------------------------------------------------------------
# 表格合计、小计核算（NUM-023）
# ---------------------------------------------------------------------------
# 核算范围：同一张表格里首列为“合计”“总计”“总数”“小计”的行，按列与分项行的数值之和比对。
#   - 小计：与所在分组（紧邻的连续明细行，上方没有就取下方）之和比对；
#   - 合计、总计、总数：与全部明细行之和、各小计行之和二者之一一致即算一致（有小计时两种写法都常见）。
# 只核算可解析为数字的单元格（支持千分位逗号、带单位的写法：1,234.5万元、300人次）；
# 以下情形整列跳过，不报错（宁可漏报，不误报 P0）：
#   - 表头含“率、%、占比、比重、比例、增长、增幅、增速、同比、环比”的百分比、比率、增长率列，
#     以及“平均、人均、均值、单价、排名、名次、序号、编号”这类不能相加的列；
#   - 参与核算的单元格中有空单元格、【待补】【待核实】单元格、无法解析的文字，或带百分号的数值；
#   - 同一列的单位不一致（万元与元混写）。
#   “—”“-”“/”“无”按 0 计。
# 误差：按参与核算单元格的最大小数位数取一个末位（1 位小数允许差 0.1，2 位小数允许差 0.01）；
#   整数列只有金额、面积等可修约的计量单位（元、万元、亿元、吨、亩、米等）允许差 1，
#   人次、人数、件数等计数列不允许有误差。
TOTAL_LABELS = frozenset({"合计", "总计", "总数"})
SUBTOTAL_LABEL = "小计"
NON_ADDITIVE_HEADER_RE = re.compile(r"率|%|％|占比|比重|比例|增长|增幅|增速|同比|环比|平均|人均|均值|单价|排名|名次|序号|编号")
ROUNDABLE_RE = re.compile(r"元|万|亿|吨|公顷|亩|米|公里|千克|公斤|升")
TABLE_SEP_RE = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(?:\|\s*:?-{2,}:?\s*)*\|?\s*$")
CELL_NUMBER_RE = re.compile(r"^([+-]?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?)\s*([^\d\s,.+\-]{0,6})$")
DASH_CELL_RE = re.compile(r"^(?:[-—–－/／]+|无|暂无)$")
_UNIT_ALIASES = {"万元": "万", "亿元": "亿", "％": "%"}


def parse_table_number(cell: str) -> Optional[Tuple[Decimal, int, str]]:
    """把表格单元格解析为 (数值, 小数位数, 单位)；不是单纯的数字（带单位可以）返回 None。

    “1,234.50万元”→(1234.50, 2, "万")，“95%”→(95, 0, "%")；单元格里的 ** 加粗符号先去掉。
    """
    text = re.sub(r"\*\*|__", "", cell).strip().replace("，", ",").replace("．", ".")
    match = CELL_NUMBER_RE.match(text)
    if not match:
        return None
    raw, unit = match.group(1), match.group(2).strip()
    digits = raw.replace(",", "")
    decimals = len(digits.split(".")[1]) if "." in digits else 0
    return Decimal(digits), decimals, _UNIT_ALIASES.get(unit, unit)


def _classify_cell(cell: str) -> Tuple[str, Optional[Tuple[Decimal, int, str]]]:
    """返回 (类别, 解析结果)：num 数字、dash 无数据（按 0 计）、blank 空、todo 待补待核实、text 其他文字。"""
    text = re.sub(r"\*\*|__", "", cell).strip()
    if not text:
        return "blank", None
    if "【待补" in text or "【待核实" in text:
        return "todo", None
    parsed = parse_table_number(text)
    if parsed:
        return "num", parsed
    if DASH_CELL_RE.match(text):
        return "dash", None
    return "text", None


def _row_label(cells: Sequence[str]) -> str:
    return re.sub(r"[\s　*_`]", "", cells[0]).rstrip("：:") if cells else ""


def split_table_row(text: str) -> List[str]:
    body = text.strip()
    if body.startswith("|"):
        body = body[1:]
    if body.endswith("|"):
        body = body[:-1]
    return [c.strip() for c in body.split("|")]


def find_markdown_tables(paras: Sequence["Para"]) -> List[Tuple[List[str], List[Tuple[Optional["Para"], List[str]]]]]:
    """找出段落序列中的 Markdown 表格，返回 [(表头单元格, [(段落, 单元格)])]（行不含表头和分隔行）。

    以“表头行 + 分隔行（|---|---|）”为起点，之后连续含“|”的行为表体；表体中遇到下一张表的表头（下一行是分隔行）即止。
    """
    def is_sep(k: int) -> bool:
        return k < len(paras) and "|" in paras[k].text and bool(TABLE_SEP_RE.match(paras[k].text))

    tables = []
    i = 0
    while i < len(paras):
        text = paras[i].text
        if "|" in text and not paras[i].in_table and not TABLE_SEP_RE.match(text) and is_sep(i + 1):
            rows: List[Tuple[Optional[Para], List[str]]] = []
            j = i + 2
            while j < len(paras) and "|" in paras[j].text and not paras[j].in_table and not is_sep(j + 1):
                rows.append((paras[j], split_table_row(paras[j].text)))
                j += 1
            tables.append((split_table_row(text), rows))
            i = j
        else:
            i += 1
    return tables


def _sum_cells(cells: Sequence[str]) -> Optional[Tuple[Decimal, int, Set[str]]]:
    """一组单元格求和：全部是数字（或“—”按 0 计）时返回 (和, 最大小数位数, 单位集合)，否则 None。"""
    total, decimals, units = Decimal(0), 0, set()
    for cell in cells:
        kind, parsed = _classify_cell(cell)
        if kind == "num" and parsed:
            total += parsed[0]
            decimals = max(decimals, parsed[1])
            if parsed[2]:
                units.add(parsed[2])
        elif kind != "dash":
            return None
    return total, decimals, units


def table_total_findings(header: Sequence[str],
                         rows: Sequence[Tuple[Optional["Para"], Sequence[str]]]) -> List[Tuple[int, str]]:
    """核算一张表格，返回 [(行下标, 说明)]，行下标对应 rows。核算规则见本节开头的说明。"""
    labels = [_row_label(cells) for _, cells in rows]
    is_total = [lab in TOTAL_LABELS for lab in labels]
    is_sub = [lab == SUBTOTAL_LABEL for lab in labels]
    is_special = [t or s for t, s in zip(is_total, is_sub)]
    findings: List[Tuple[int, str]] = []

    def cell(row: int, col: int) -> str:
        cells = rows[row][1]
        return cells[col] if col < len(cells) else ""

    def run_of_details(start: int, step: int) -> List[int]:
        out, k = [], start
        while 0 <= k < len(rows) and not is_special[k]:
            out.append(k)
            k += step
        return out

    for i, label in enumerate(labels):
        if not is_special[i]:
            continue
        # 找出本行对应的分项行：candidates = [(名称, 行下标列表)]
        if is_sub[i]:
            group = run_of_details(i - 1, -1) or run_of_details(i + 1, 1)
            candidates = [("同组分项之和", group)]
        else:
            lo = max([k for k in range(i) if is_total[k]], default=-1) + 1
            above = list(range(lo, i))
            hi = next((k for k in range(i + 1, len(rows)) if is_total[k]), len(rows))
            span = above if any(not is_special[k] or is_sub[k] for k in above) else list(range(i + 1, hi))
            details = [k for k in span if not is_special[k]]
            subs = [k for k in span if is_sub[k]]
            candidates = [("分项之和", details)]
            if subs:
                candidates.append(("各小计之和", subs))
        candidates = [(name, members) for name, members in candidates if members]
        if not candidates:
            continue
        for col in range(1, len(header)):
            name = header[col].strip() or f"第{col + 1}列"
            if NON_ADDITIVE_HEADER_RE.search(name):
                continue
            kind, stated = _classify_cell(cell(i, col))
            if kind != "num" or not stated or stated[2] == "%":
                continue
            sums = []
            for cand_name, members in candidates:
                result = _sum_cells([cell(k, col) for k in members])
                if result is not None:
                    sums.append((cand_name, result))
            if not sums:
                continue
            units = {stated[2]} - {""}
            for _, (_, _, u) in sums:
                units |= u
            if "%" in units or len(units) > 1:
                continue
            decimals = max([stated[1]] + [d for _, (_, d, _) in sums])
            if decimals > 0:
                tolerance = Decimal(1).scaleb(-decimals)
            else:
                tolerance = Decimal(1) if ROUNDABLE_RE.search("".join(units) + name) else Decimal(0)
            if any(abs(stated[0] - total) <= tolerance for _, (total, _, _) in sums):
                continue
            shown = "；".join(f"{cand_name}为 {total:.{decimals}f}" for cand_name, (total, _, _) in sums)
            findings.append((i, f"表格“{labels[i]}”行“{name}”列填写 {cell(i, col).strip()}，与{shown}不一致"))
    return findings


# ---------------------------------------------------------------------------
# 输入读取
# ---------------------------------------------------------------------------

def read_lines(path: Path) -> Tuple[List[Tuple[int, str, bool]], List[str]]:
    """读取文件，返回 [(行号, 文本, 是否表格内)] 和提示信息。"""
    notes: List[str] = []
    suffix = path.suffix.lower()
    if suffix in (".txt", ".md", ".markdown", ".text"):
        try:
            raw = path.read_bytes()
        except OSError as exc:
            raise InputError(f"无法读取文件：{exc}") from exc
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                text = raw.decode("gb18030")
                notes.append("文件不是 UTF-8 编码，已按 GB18030 读取。")
            except UnicodeDecodeError as exc:
                raise InputError("无法识别文件编码，请另存为 UTF-8 后重试。") from exc
        text = text.replace("\r\n", "\n").replace("\r", "\n")
        return [(i, line, False) for i, line in enumerate(text.split("\n"), start=1)], notes
    if suffix == ".docx":
        return read_docx(path), notes
    raise InputError(f"不支持的文件类型：{path.suffix or '无扩展名'}。只支持 .txt、.md、.docx。")


class DocxLines(list):
    """read_docx 的返回值：行列表，另带 tables 属性（表格的行列结构，供表格合计核算用）。"""
    tables: List[List[Tuple[int, List[str]]]]


def read_docx(path: Path) -> List[Tuple[int, str, bool]]:
    try:
        import docx  # type: ignore
        from docx.oxml.ns import qn  # type: ignore
        from docx.text.paragraph import Paragraph  # type: ignore
    except ImportError as exc:
        raise InputError("读取 .docx 需要 python-docx，请先运行：pip install python-docx") from exc
    try:
        document = docx.Document(str(path))
    except Exception as exc:  # 损坏的 docx、加密文件等
        raise InputError(f"无法打开 docx 文件：{exc}") from exc
    lines = DocxLines()
    lines.tables = []
    number = 0
    p_tag, tbl_tag, tr_tag, tc_tag = qn("w:p"), qn("w:tbl"), qn("w:tr"), qn("w:tc")
    for child in document.element.body.iterchildren():
        if child.tag == p_tag:
            number += 1
            lines.append((number, Paragraph(child, document).text, False))
        elif child.tag == tbl_tag:
            numbers = {}
            for p_el in child.iter(p_tag):
                number += 1
                numbers[p_el] = number
                lines.append((number, Paragraph(p_el, document).text, True))
            rows: List[Tuple[int, List[str]]] = []
            for tr in child.findall(tr_tag):
                first = next(tr.iter(p_tag), None)
                if first is None:
                    continue
                cells = ["".join(Paragraph(p_el, document).text for p_el in tc.iter(p_tag)).strip()
                         for tc in tr.findall(tc_tag)]
                rows.append((numbers[first], cells))
            lines.tables.append(rows)
    return lines


# ---------------------------------------------------------------------------
# 文档结构
# ---------------------------------------------------------------------------

class Document:
    """正文（【起草说明】之前）的段落及结构识别结果。"""

    def __init__(self, file_name: str, lines: Sequence[Tuple[int, str, bool]], unit: str = "行") -> None:
        self.file = file_name
        self.unit = unit
        self.truncated_line: Optional[int] = None
        paras: List[Para] = []
        for number, text, in_table in lines:
            if DRAFT_NOTE_RE.match(text):
                self.truncated_line = number
                break
            if text.strip():
                paras.append(Para(len(paras), number, text.rstrip(), in_table))
        self.paras = paras
        self.char_count = sum(nonspace_len(p.text) for p in paras)
        self.head: Set[int] = set()
        self.title: List[int] = []
        self.subtitle: Optional[int] = None
        self.zhusong: Optional[int] = None
        self.banji_start: Optional[int] = None
        self.signoff_date: Optional[int] = None
        self.attach_desc: Optional[int] = None
        # docx 表格的行列结构：[[(首段落行号, [单元格文字])]]，表头在前；其他输入为空（Markdown 表格从段落中识别）
        self.docx_tables: List[List[Tuple[int, List[str]]]] = []
        self.wenhao_present = False
        self.qianfa_present = False
        self.public_facing = False
        self.first_body: Optional[int] = None
        self._analyze()

    # -- 结构识别 -----------------------------------------------------------
    def _analyze(self) -> None:
        paras = self.paras
        if not paras:
            return
        i = 0
        while i < len(paras) and i < 8:
            text = paras[i].text.strip()
            nxt = paras[i + 1].text.strip() if i + 1 < len(paras) else ""
            if self._is_wenhao_line(text):
                self.wenhao_present = True
                if QIANFA_RE.search(text):
                    self.qianfa_present = True
                self.head.add(i)
            elif QIANFA_RE.search(text) and len(text) <= 30:
                self.qianfa_present = True
                self.head.add(i)
            elif (FENHAO_RE.match(text) or MIJI_LINE_RE.match(text) or JINJI_RE.match(text)
                  or BIAOZHI_RE.match(text) or INNER_MARK_RE.match(text)):
                self.head.add(i)
            elif (len(text) <= 40 and nxt and (self._is_wenhao_line(nxt) or QIANFA_RE.search(nxt))
                  and not self._title_shaped(text)):
                self.head.add(i)  # 发文机关标志（不带“文件”二字）；标题后紧跟发文字号的不算
            else:
                break
            i += 1
        if i < len(paras) and self._title_like(paras[i].text):
            self.title.append(i)
            j = i + 1
            while (j < len(paras) and len(self.title) < 3
                   and not TITLE_END_RE.search(self._title_text().rstrip(PUNCT_TAIL))
                   and self._title_continuation(paras[j].text)):
                self.title.append(j)
                j += 1
            if j < len(paras) and paras[j].text.strip().startswith("——"):
                self.subtitle = j
                j += 1
            if j < len(paras) and self._is_wenhao_line(paras[j].text.strip()):
                # 发文字号写在标题之下（AI 稿常见）：按版头要素处理，不当作主送或正文
                self.wenhao_present = True
                if QIANFA_RE.search(paras[j].text):
                    self.qianfa_present = True
                self.head.add(j)
                j += 1
            if j < len(paras):
                cand = paras[j].text.strip()
                if len(cand) <= 100 and "。" not in cand and (cand.endswith("：") or cand.endswith(":")):
                    self.zhusong = j
        # 版记：从末尾向前找抄送、印发行
        for k in range(len(paras) - 1, max(-1, len(paras) - 7), -1):
            if k in self.head or k in self.title:
                break
            if BANJI_RE.search(paras[k].text):
                self.banji_start = k
        end = self.banji_start if self.banji_start is not None else len(paras)
        for k in range(end - 1, max(-1, end - 10), -1):
            if DATE_ONLY_RE.match(paras[k].text):
                self.signoff_date = k
                break
        for para in paras:
            if ATTACH_DESC_RE.match(para.text) and para.index not in self.head:
                self.attach_desc = para.index
                break
        title = self.title_text.rstrip(PUNCT_TAIL)
        audience = paras[self.zhusong].text if self.zhusong is not None else ""
        if not audience:
            # 没有标题时，开头几段中以冒号结尾的短行视为称呼（如“各位居民：”）
            for para in paras[:3]:
                text = para.text.strip()
                if len(text) <= 40 and text.endswith(("：", ":")):
                    audience = text
                    break
        self.public_facing = bool(PUBLIC_TITLE_RE.search(title) or PUBLIC_AUDIENCE_RE.search(audience)
                                  or re.search(r"致[^，。]{0,10}(?:居民|村民|家长|业主|市民)", title))
        for k in self.body_indices():
            text = paras[k].text
            # 正文第一段：不是层次标题、附件说明，且是成句的段落（缺冒号的主送行等短行跳过）
            if (not heading_of(text) and not ATTACH_DESC_RE.match(text) and not paras[k].in_table
                    and re.search(r"[。；，！？：]", text)):
                self.first_body = k
                break

    @staticmethod
    def _is_wenhao_line(text: str) -> bool:
        """整行基本就是发文字号（可带签发人），而不是正文中引用文号的句子。"""
        m = WENHAO_RE.search(text)
        if not m:
            return False
        rest = (text[:m.start()] + text[m.end():]).strip()
        return len(rest) <= 6 or rest.startswith("签发人") or bool(QIANFA_RE.search(rest))

    def _title_text(self) -> str:
        return "".join(self.paras[k].text.strip() for k in self.title)

    @staticmethod
    def _title_shaped(text: str) -> bool:
        """以文种结尾或以“关于”开头：是标题，不是发文机关标志。"""
        text = text.strip()
        return text.startswith("关于") or bool(TITLE_END_RE.search(text.rstrip(PUNCT_TAIL)))

    @staticmethod
    def _title_like(text: str) -> bool:
        text = text.strip()
        if not text or len(text) > 80 or HEADING_START_RE.match(text):
            return False
        if text.endswith("：") or text.endswith(":") or text.startswith("——"):
            return False
        if "，" in text or "；" in text or "：" in text:
            return False
        if "。" not in text:
            return True
        # 以句号结尾的，只有像标题（以文种结尾或以“关于”开头）才算标题，避免把单句正文当标题
        body = text[:-1]
        return text.count("。") == 1 and text.endswith("。") and (
            bool(TITLE_END_RE.search(body)) or body.startswith("关于"))

    @staticmethod
    def _title_continuation(text: str) -> bool:
        text = text.strip()
        if not text or len(text) > 40 or HEADING_START_RE.match(text) or text.startswith("——"):
            return False
        if DATE_ONLY_RE.match(text) or re.search(r"[。；：:，,（(]", text):
            return False
        return True

    # -- 查询 -----------------------------------------------------------------
    @property
    def title_text(self) -> str:
        return self._title_text()

    def body_indices(self) -> List[int]:
        """正文段落：去掉版头、标题、副标题、主送和版记。"""
        skip = set(self.head) | set(self.title)
        if self.subtitle is not None:
            skip.add(self.subtitle)
        if self.zhusong is not None:
            skip.add(self.zhusong)
        end = self.banji_start if self.banji_start is not None else len(self.paras)
        return [p.index for p in self.paras[:end] if p.index not in skip]

    def body_char_count(self) -> int:
        """篇幅统计：标题之后、【起草说明】之前，不含空白和版记。"""
        skip = set(self.head) | set(self.title)
        if self.subtitle is not None:
            skip.add(self.subtitle)
        end = self.banji_start if self.banji_start is not None else len(self.paras)
        return sum(nonspace_len(p.text) for p in self.paras[:end] if p.index not in skip)

    def has_qianbao_header(self) -> bool:
        """稿首（前 3 个非空段落）单独一行写着“签报”“呈批件”表头（容许 Markdown 标题、加粗、字间空格）。"""
        for para in self.paras[:3]:
            head = MD_HEADING_PREFIX_RE.sub("", para.text, count=1)
            head = re.sub(r"[\s　]+|(?:\*\*|__)$", "", head).rstrip(PUNCT_TAIL)
            if head in QIANBAO_HEADERS:
                return True
        return False

    def qianbao_marked(self) -> bool:
        """是否带签报标志：稿首有“签报”“呈批件”表头，或全文有“会签情况”栏（用于免查请示主送领导个人）。"""
        return self.has_qianbao_header() or any("会签情况" in p.text for p in self.paras)

    def infer_doc_type(self) -> Optional[str]:
        if self.has_qianbao_header():
            return "签报"  # 签报按规定写“关于……的请示”，只能靠稿首表头识别
        title = self.title_text.strip().rstrip(PUNCT_TAIL)
        if title:
            for regex, doc_type in TITLE_DOC_PATTERNS:
                if regex.search(title):
                    return doc_type
            for suffix, doc_type in TITLE_DOC_SUFFIX:
                if title.endswith(suffix):
                    return doc_type
        if self.subtitle is not None:
            sub = self.paras[self.subtitle].text.strip()
            if "党课" in sub:
                return "党课讲稿"
            if re.search(r"研讨[^，。]{0,20}发言\s*$|研讨发言", sub):
                return "研讨发言"
            if re.search(r"(讲话|发言|致辞|讲稿)\s*$", sub):
                return "讲话稿"
        return None


# ---------------------------------------------------------------------------
# 引擎
# ---------------------------------------------------------------------------

def _compile(pattern: str, flags: int = 0) -> Pattern[str]:
    return re.compile(pattern, flags)


class Linter:
    """加载规则数据，对 Document 执行检查。"""

    def __init__(self, rules_path: Optional[Path] = None, patterns_path: Optional[Path] = None) -> None:
        self.warnings: List[str] = []
        self.rules: List[RegexRule] = []
        self.proc_meta: Dict[str, Dict[str, Any]] = {}
        self.patterns: List[StylePattern] = []
        self.whitelist: List[Tuple[str, Pattern[str]]] = []
        self.evidence: Dict[str, Any] = {}          # 评价词分级表（lint_rules.json 的 evidence 段）
        self.evidence_allow: Set[str] = set()      # 本次检查的评价词白名单（--evidence-allow）
        self._load_rules(rules_path or DEFAULT_RULES)
        self._load_patterns(patterns_path or DEFAULT_PATTERNS)

    # -- 加载 -----------------------------------------------------------------
    def _load_json(self, path: Path, label: str) -> Optional[Dict[str, Any]]:
        if not path.is_file():
            self.warnings.append(f"未找到 {label}（{path.name}），相关检查已跳过。")
            return None
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, ValueError) as exc:
            self.warnings.append(f"{label}（{path.name}）无法解析：{exc}，相关检查已跳过。")
            return None
        if not isinstance(data, dict):
            self.warnings.append(f"{label}（{path.name}）顶层不是对象，相关检查已跳过。")
            return None
        return data

    def _load_rules(self, path: Path) -> None:
        data = self._load_json(path, "规则文件")
        if data is None:
            return
        for item in data.get("rules", []):
            try:
                flags = re.IGNORECASE if item.get("flags") == "IGNORECASE" else 0
                skip = item.get("skip_line")
                fixes = tuple((_compile(p), r) for p, r in item.get("fix", []))
                self.rules.append(RegexRule(
                    id=item["id"], check=item["check"],
                    category=item.get("category") or CHECK_NAMES.get(item["check"], item["check"]),
                    gate=item.get("gate", ""), level=item.get("level", "P2"),
                    regex=_compile(item["pattern"], flags),
                    message=item.get("message", ""), suggestion=item.get("suggestion", ""),
                    basis=item.get("basis", ""),
                    doc_types=tuple(item.get("doc_types") or ()),
                    exempt_doc_types=tuple(item.get("exempt_doc_types") or ()),
                    yield_to=tuple(item.get("yield_to") or ()),
                    once=bool(item.get("once")),
                    skip_line=_compile(skip) if skip else None,
                    fixes=fixes,
                    skip_table=bool(item.get("skip_table")),
                    zone=item.get("zone", ""),
                    also_public=bool(item.get("also_public")),
                    public_only=bool(item.get("public_only")),
                    skip_if_basis=bool(item.get("skip_if_basis")),
                    skip_in_quotes=bool(item.get("skip_in_quotes")),
                    skip_if_negated=bool(item.get("skip_if_negated")),
                ))
            except (KeyError, re.error, TypeError, ValueError) as exc:
                self.warnings.append(f"规则 {item.get('id', '?')} 无效，已跳过：{exc}")
        for item in data.get("procedural", []):
            if isinstance(item, dict) and item.get("id"):
                self.proc_meta[item["id"]] = item
        self._load_evidence(data.get("evidence"))

    def _load_evidence(self, cfg: Any) -> None:
        """评价词分级表：strong、medium 为正则词条，weak 为限定词（出现在评价词前 4 字内则不查），
        whitelist 为不算评价的固定短语，doc_types 为适用文书类型（空表示全部）。"""
        if not isinstance(cfg, dict):
            return
        try:
            def alt(key: str) -> Pattern[str]:
                words = [w for w in cfg.get(key, []) if isinstance(w, str) and w]
                words.sort(key=len, reverse=True)
                return re.compile("|".join(words) if words else r"(?!x)x")
            self.evidence = {
                "strong": alt("strong"), "medium": alt("medium"), "weak": alt("weak"),
                "whitelist": [re.compile(w) for w in cfg.get("whitelist", []) if isinstance(w, str) and w],
                "doc_types": tuple(cfg.get("doc_types") or ()),
            }
        except re.error as exc:
            self.warnings.append(f"评价词表无效，评价证据检查已跳过：{exc}")
            self.evidence = {}

    def _load_patterns(self, path: Path) -> None:
        data = self._load_json(path, "文风词表")
        if data is None:
            return
        for item in data.get("patterns", []):
            try:
                density = item.get("density_per_1000")
                self.patterns.append(StylePattern(
                    id=item["id"], group=item.get("group", "style"),
                    category=item.get("category", ""), regex=_compile(item["pattern"]),
                    level=item.get("level", "P2"), message=item.get("message", ""),
                    suggestion=item.get("suggestion", ""),
                    density=float(density) if density is not None else None,
                    doc_types=tuple(item.get("doc_types") or ()),
                    exempt_doc_types=tuple(item.get("exempt_doc_types") or ()),
                    exclude=tuple(_compile(p) for p in item.get("exclude") or ()),
                ))
            except (KeyError, re.error, TypeError, ValueError) as exc:
                self.warnings.append(f"词表规则 {item.get('id', '?')} 无效，已跳过：{exc}")
        for item in data.get("whitelist", []):
            try:
                self.whitelist.append((item["id"], _compile(item["pattern"])))
            except (KeyError, re.error, TypeError) as exc:
                self.warnings.append(f"白名单 {item.get('id', '?')} 无效，已跳过：{exc}")

    # -- 入口 -----------------------------------------------------------------
    def lint(self, doc: Document, doc_type: str, checks: Set[str],
             banned: Optional[Sequence[BannedTerm]] = None) -> List[Issue]:
        issues: List[Issue] = []
        issues.extend(self._run_regex_rules(doc, doc_type, checks))
        for check in CHECK_ORDER:
            if check not in checks:
                continue
            for func in self._procedures().get(check, ()):
                issues.extend(func(doc, doc_type))
        issues.extend(self._run_style(doc, doc_type, checks))
        if banned and "style" in checks:
            issues.extend(self._check_banned(doc, doc_type, banned))
        issues = self._resolve_yields(issues)
        issues = self._apply_once(issues)
        issues.sort(key=lambda x: (LEVEL_RANK.get(x.level, 9), x.line, x.start, x.id))
        return issues

    def _procedures(self) -> Dict[str, Sequence[Callable[[Document, str], List[Issue]]]]:
        return {
            "punct": (self._check_quotes, self._check_colon_nesting),
            "number": (self._check_date_styles, self._check_quantity_points, self._check_table_totals),
            "format": (self._check_date_padding, self._check_signoff_date, self._check_zhusong, self._check_elements),
            "structure": (self._check_title, self._check_heading_sequence, self._check_count_summary,
                          self._check_qingshi, self._check_baogao, self._check_han, self._check_shangxing),
            "style": (self._check_evidence,),
            "privacy": (self._check_privacy,),
            "length": (self._check_length,),
        }

    # -- 构造问题 -----------------------------------------------------------
    def _proc_issue(self, rid: str, para: Optional[Para], start: int, end: int, *,
                    message: Optional[str] = None, level: Optional[str] = None,
                    suggestion: Optional[str] = None, excerpt: Optional[str] = None,
                    context: Optional[str] = None, count: Optional[int] = None) -> Optional[Issue]:
        meta = self.proc_meta.get(rid)
        if meta is None:
            return None  # 未在 lint_rules.json 的 procedural 中登记的检查不输出
        text = para.text if para else ""
        if excerpt is None:
            excerpt = clip(text[start:end], 80) if para and start >= 0 else ""
        if context is None:
            context = context_of(text, start, end) if para and start >= 0 else ""
        return Issue(
            id=rid, level=level or meta.get("level", "P2"), check=meta.get("check", ""),
            category=meta.get("category", CHECK_NAMES.get(meta.get("check", ""), "")),
            gate=meta.get("gate", ""), line=para.line if para else 0, excerpt=excerpt,
            message=message or meta.get("message", ""),
            suggestion=suggestion if suggestion is not None else meta.get("suggestion", ""),
            basis=meta.get("basis", ""), context=context if context != excerpt else "",
            start=start, end=end, count=count, para_index=para.index if para else -1,
        )

    def _proc_applies(self, rid: str, doc_type: str) -> bool:
        meta = self.proc_meta.get(rid)
        if meta is None:
            return False
        return applies_to(doc_type, meta.get("doc_types") or (), meta.get("exempt_doc_types") or ())

    # -- 正则规则 -----------------------------------------------------------
    @staticmethod
    def _rule_applies(rule: RegexRule, doc: Document, doc_type: str) -> bool:
        """文种限定：public_only 只查面向群众的文书；also_public 在限定文种之外也查面向群众的文书。"""
        if doc_type in rule.exempt_doc_types:
            return False
        if rule.public_only:
            return doc.public_facing
        if rule.doc_types and doc_type not in rule.doc_types:
            return rule.also_public and doc.public_facing
        return True

    def _run_regex_rules(self, doc: Document, doc_type: str, checks: Set[str]) -> List[Issue]:
        issues: List[Issue] = []
        for rule in self.rules:
            if rule.check not in checks or not self._rule_applies(rule, doc, doc_type):
                continue
            if rule.zone == "first_body":
                paras = [doc.paras[doc.first_body]] if doc.first_body is not None else []
            else:
                paras = doc.paras
            for para in paras:
                if rule.skip_table and para.in_table:
                    continue
                if rule.skip_line and rule.skip_line.search(para.text):
                    continue
                for match in rule.regex.finditer(para.text):
                    start, end = match.span()
                    if start == end:
                        continue
                    if rule.skip_in_quotes and self._inside_quote(para.text, start):
                        continue  # 引号内是举例或引述，不是稿件本身的标注
                    if rule.skip_if_basis or rule.skip_if_negated:
                        sent_start = max(para.text.rfind(ch, 0, start) for ch in "。！？；") + 1
                        if rule.skip_if_basis and LEGAL_BASIS_RE.search(para.text[sent_start:start]):
                            continue
                        if rule.skip_if_negated and NEGATION_RE.search(para.text[max(sent_start, start - 8):start]):
                            continue  # “不要求拍照留痕”“严禁签订责任状”是在落实减负规定，不是违反
                    excerpt = clip(match.group(0), 80)
                    suggestion = rule.suggestion
                    if rule.fixes:
                        fixed = match.group(0).strip()
                        for fix_re, repl in rule.fixes:
                            fixed = fix_re.sub(repl, fixed)
                        if fixed != match.group(0).strip():
                            suggestion = f"改为“{fixed}”" if fixed else "删除"
                    ctx = context_of(para.text, start, end)
                    issues.append(Issue(
                        id=rule.id, level=rule.level, check=rule.check, category=rule.category,
                        gate=rule.gate, line=para.line, excerpt=excerpt, message=rule.message,
                        suggestion=suggestion, basis=rule.basis, context=ctx if ctx != excerpt else "",
                        start=start, end=end, para_index=para.index,
                    ))
        return issues

    def _resolve_yields(self, issues: List[Issue]) -> List[Issue]:
        yields = {r.id: set(r.yield_to) for r in self.rules if r.yield_to}
        if not yields:
            return issues
        # 按（段落, 规则）分组，只与让位目标规则的命中比较，避免同段问题很多时退化为二次方
        by_para_rule: Dict[Tuple[int, str], List[Issue]] = {}
        for issue in issues:
            by_para_rule.setdefault((issue.para_index, issue.id), []).append(issue)
        kept: List[Issue] = []
        for issue in issues:
            targets = yields.get(issue.id)
            if targets:
                overlapped = any(
                    other.start < issue.end and issue.start < other.end
                    for target in targets
                    for other in by_para_rule.get((issue.para_index, target), ())
                )
                if overlapped:
                    continue
            kept.append(issue)
        return kept

    def _apply_once(self, issues: List[Issue]) -> List[Issue]:
        once_ids = {r.id for r in self.rules if r.once}
        if not once_ids:
            return issues
        first: Dict[str, Issue] = {}
        counts: Dict[str, int] = {}
        kept: List[Issue] = []
        for issue in sorted(issues, key=lambda x: (x.line, x.start)):
            if issue.id in once_ids:
                counts[issue.id] = counts.get(issue.id, 0) + 1
                if issue.id in first:
                    continue
                first[issue.id] = issue
            kept.append(issue)
        for rid, issue in first.items():
            if counts[rid] > 1:
                issue.count = counts[rid]
                issue.message = f"{issue.message}（全文共 {counts[rid]} 处，只列第一处）"
        return kept

    # -- AI 味、防御性 --------------------------------------------------------
    def _whitelist_spans(self, para: Para) -> List[Tuple[int, int]]:
        """词表白名单命中区间（【待补：】【待核实：】标记见 WL-018）。"""
        return [m.span() for _wid, regex in self.whitelist
                for m in regex.finditer(para.text) if m.end() > m.start()]

    def style_hits(self, doc: Document, pat: StylePattern,
                   wl_spans: Optional[Callable[[Para], List[Tuple[int, int]]]] = None) -> List[Tuple[Para, int, int]]:
        """逐段匹配一条词表规则，剔除与白名单重叠的命中，返回 (段落, 起, 止)。不做密度判断。"""
        if wl_spans is None:
            wl_spans = self._whitelist_spans
        hits: List[Tuple[Para, int, int]] = []
        for para in doc.paras:
            own_spans: Optional[List[Tuple[int, int]]] = None
            for match in pat.regex.finditer(para.text):
                start, end = match.span()
                if start == end:
                    continue
                if any(ws < end and start < we for ws, we in wl_spans(para)):
                    continue  # 与白名单重叠的命中不报告
                if pat.exclude:
                    if own_spans is None:  # 本规则专用的排除区间（如 AI-013 的字段标签、称呼行）
                        own_spans = [m.span() for regex in pat.exclude for m in regex.finditer(para.text)
                                     if m.end() > m.start()]
                    if any(ws < end and start < we for ws, we in own_spans):
                        continue
                # 去掉正则为定位句首而带上的前导标点；命中本身就是标点（如顿号、冒号）时保留
                trimmed = start
                while trimmed < end and para.text[trimmed] in "。！？；，,：:、 \t　":
                    trimmed += 1
                hits.append((para, trimmed if trimmed < end else start, end))
        return hits

    def _run_style(self, doc: Document, doc_type: str, checks: Set[str]) -> List[Issue]:
        wanted = [p for p in self.patterns if p.group in checks]
        if not wanted:
            return []
        wl_cache: Dict[int, List[Tuple[int, int]]] = {}

        def wl_spans(para: Para) -> List[Tuple[int, int]]:
            if para.index not in wl_cache:
                wl_cache[para.index] = self._whitelist_spans(para)
            return wl_cache[para.index]

        issues: List[Issue] = []
        total = max(doc.char_count, 1)
        for pat in wanted:
            if not applies_to(doc_type, pat.doc_types, pat.exempt_doc_types):
                continue
            hits = self.style_hits(doc, pat, wl_spans)
            if not hits:
                continue
            check_name = pat.category or CHECK_NAMES.get(pat.group, pat.group)
            basis_file = "ai-tone.md" if pat.group == "style" else "defensive.md"
            basis = f"文风词表 {pat.id}（{pat.category}），见 references/_shared/style/{basis_file}"
            if pat.density is not None:
                per_k = len(hits) * 1000.0 / total
                if len(hits) < 2 or per_k <= pat.density:
                    continue
                para, start, end = hits[0]
                issues.append(Issue(
                    id=pat.id, level=pat.level, check=pat.group, category=check_name, gate="文风关",
                    line=para.line, excerpt=clip(para.text[start:end], 80),
                    message=f"{pat.message}（全文 {len(hits)} 处，每千字 {per_k:.1f} 处，阈值 {pat.density:g}）",
                    suggestion=pat.suggestion, basis=basis,
                    context=context_of(para.text, start, end), start=start, end=end,
                    count=len(hits), para_index=para.index,
                ))
            else:
                for para, start, end in hits:
                    ctx = context_of(para.text, start, end)
                    excerpt = clip(para.text[start:end], 80)
                    issues.append(Issue(
                        id=pat.id, level=pat.level, check=pat.group, category=check_name, gate="文风关",
                        line=para.line, excerpt=excerpt, message=pat.message,
                        suggestion=pat.suggestion, basis=basis,
                        context=ctx if ctx != excerpt else "", start=start, end=end,
                        para_index=para.index,
                    ))
        return issues

    # -- 标点（函数） ---------------------------------------------------------
    def _check_quotes(self, doc: Document, doc_type: str) -> List[Issue]:
        """PUN-009 引号嵌套：外层双引号内又出现双引号。"""
        issues: List[Issue] = []
        for para in doc.paras:
            text = para.text
            if "“" not in text:
                continue
            depth, outer = 0, -1
            for pos, ch in enumerate(text):
                if ch == "“":
                    if depth == 0:
                        outer = pos
                    elif depth >= 1:
                        closes = [m.start() for m in re.finditer("”", text[pos:])]
                        end = pos + (closes[1] if len(closes) > 1 else (closes[0] if closes else 0)) + 1
                        end = min(max(end, pos + 1), outer + 60, len(text))
                        issue = self._proc_issue("PUN-009", para, outer, end)
                        if issue:
                            issues.append(issue)
                        break
                    depth += 1
                elif ch == "”":
                    depth = max(0, depth - 1)
        return issues

    FIELD_LABEL_RE = re.compile(
        r"(?:主送|抄送|附件|联系人|联系电话|电话|传真|电子邮箱|邮箱|地址|签发人|时间|地点|主持人?|出席|列席|参加人员|"
        r"记录人?|议题|牵头单位|责任单位|配合单位|责任人|完成时限|时间节点|报送方式|邮编|网址|开户行|账号|户名)\s*$"
    )

    # “提交部门：某科　汇报人：李某”这类字段行：每个冒号前都是行首或空格后的短标签
    FIELD_SEG_RE = re.compile(r"(?:^|(?<=[ 　\t]))[^\s：:，,。；;！？!?、]{1,8}[：:]")

    @classmethod
    def _is_field_line(cls, text: str) -> bool:
        text = text.strip()
        if re.search(r"[。！？!?；;](?=.)", text):
            return False  # 句末之外还有句号、分号的，是成句的正文
        colons = [m.start() for m in re.finditer("[：:]", text)
                  if not (0 < m.start() < len(text) - 1 and text[m.start() - 1].isdigit()
                          and text[m.start() + 1].isdigit())]
        if len(colons) < 2:
            return False
        labels = {m.end() - 1 for m in cls.FIELD_SEG_RE.finditer(text)}
        return all(pos in labels for pos in colons)

    def _check_colon_nesting(self, doc: Document, doc_type: str) -> List[Issue]:
        """PUN-014 一句之内套用冒号。"""
        issues: List[Issue] = []
        skip = set(doc.head) | set(doc.title)
        if doc.banji_start is not None:
            skip |= set(range(doc.banji_start, len(doc.paras)))
        for para in doc.paras:
            if para.index in skip or para.in_table or "：" not in para.text:
                continue
            if self._is_field_line(para.text):
                continue
            text = para.text
            marker_spans = [mk.span() for mk in MARKER_RE.finditer(text)]
            sent_start = 0
            for m in re.finditer(r"[。！？；]|$", text):
                sent_end = m.start()
                colons = []
                for c in re.finditer("：", text[sent_start:sent_end]):
                    pos = sent_start + c.start()
                    if any(s <= pos < e for s, e in marker_spans):
                        continue
                    if pos > 0 and pos + 1 < len(text) and text[pos - 1].isdigit() and text[pos + 1].isdigit():
                        continue
                    if self.FIELD_LABEL_RE.search(text[max(sent_start, pos - 8):pos]):
                        continue
                    colons.append(pos)
                if len(colons) >= 2:
                    issue = self._proc_issue("PUN-014", para, sent_start, min(sent_end + 1, len(text)),
                                             excerpt=clip(text[sent_start:sent_end + 1], 60))
                    if issue:
                        issues.append(issue)
                sent_start = m.end()
                if sent_start >= len(text):
                    break
        return issues

    # -- 数字（函数） ---------------------------------------------------------
    DATE_STYLES = (
        ("年月日", re.compile(r"\d{4}\s*年\s*\d{1,2}\s*月\s*\d{1,2}\s*日")),
        ("短横线", re.compile(r"(?<![\d\-])\d{4}-\d{1,2}-\d{1,2}(?![\d\-])")),
        ("点号", re.compile(r"(?<![\d.])\d{4}\.\d{1,2}\.\d{1,2}(?![\d.])")),
        ("斜线", re.compile(r"(?<![\d/])\d{4}/\d{1,2}/\d{1,2}(?![\d/])")),
    )

    def _check_date_styles(self, doc: Document, doc_type: str) -> List[Issue]:
        """NUM-010 日期写法统一。"""
        found: Dict[str, List[Tuple[Para, int, int]]] = {}
        end_idx = doc.banji_start if doc.banji_start is not None else len(doc.paras)
        for para in doc.paras[:end_idx]:
            for name, regex in self.DATE_STYLES:
                for m in regex.finditer(para.text):
                    found.setdefault(name, []).append((para, m.start(), m.end()))
        issues: List[Issue] = []
        for name in ("点号", "斜线"):
            if name in found:
                para, s, e = found[name][0]
                issue = self._proc_issue(
                    "NUM-010", para, s, e,
                    message=f"日期用{name}分隔（如“{para.text[s:e]}”），不在 GB/T 15835 列举的写法中（共 {len(found[name])} 处）",
                    count=len(found[name]))
                if issue:
                    issues.append(issue)
        if "年月日" in found and "短横线" in found:
            para, s, e = found["短横线"][0]
            issue = self._proc_issue(
                "NUM-010", para, s, e,
                message=f"同一文件日期写法不统一：“年月日”写法 {len(found['年月日'])} 处，短横线写法 {len(found['短横线'])} 处",
                count=len(found["短横线"]))
            if issue:
                issues.append(issue)
        return issues

    # 数量变化用了百分点（NUM-020，data-expression.md DE-08）
    POINT_CHANGE_RE = re.compile(r"(?:增长|增加|下降|减少|提高|降低|上升|回落)了?\s*\d+(?:\.\d+)?\s*个?百分点")
    QTY_WORD_RE = re.compile(
        r"人次|人数|件数|户数|台次|例数|万元|亿元|千元|元|件|户|台|例|吨|辆|床日|数量|金额|总量|总额|收入|产值"
        r"|销售额|营业额|交易额|利润|支出|投资额|门诊量|业务量|办件量|工作量|诊疗量|客流量|用量|销量|产量"
    )
    RATE_WORD_RE = re.compile(r"率|占比|比重|比例|增速|增幅|指数|份额|度|水平|系数|密度|进度|差距")

    def _check_table_totals(self, doc: Document, doc_type: str) -> List[Issue]:
        """NUM-023 表格合计、小计核算：Markdown 表格从段落中识别，docx 表格用读入时记下的行列结构。核算规则见 table_total_findings。"""
        if not self._proc_applies("NUM-023", doc_type):
            return []
        tables = find_markdown_tables(doc.paras)
        by_line = {p.line: p for p in doc.paras}
        for rows in doc.docx_tables:
            if len(rows) < 2 or (doc.truncated_line is not None and rows[0][0] >= doc.truncated_line):
                continue
            tables.append((rows[0][1], [(by_line.get(line), cells) for line, cells in rows[1:]]))
        issues: List[Issue] = []
        for header, rows in tables:
            for row_index, message in table_total_findings(header, rows):
                para, cells = rows[row_index]
                excerpt = clip(" | ".join(c for c in cells if c), 80)
                issue = self._proc_issue("NUM-023", para, 0, len(para.text) if para else 0,
                                         message=message, excerpt=excerpt)
                if issue:
                    issues.append(issue)
        return issues

    def _check_quantity_points(self, doc: Document, doc_type: str) -> List[Issue]:
        """NUM-020 数量指标的变化用了“百分点”。只看同一分句内、数量词之后 15 字以内的写法。"""
        issues: List[Issue] = []
        for para in doc.paras:
            if para.in_table or "百分点" not in para.text:
                continue
            text = para.text
            for m in self.POINT_CHANGE_RE.finditer(text):
                clause_start = max(text.rfind(ch, 0, m.start()) for ch in "，。；！？：,;:") + 1
                window_start = max(clause_start, m.start() - 20)
                qty = None
                for q in self.QTY_WORD_RE.finditer(text, window_start, m.start()):
                    qty = q
                if qty is None or m.start() - qty.end() > 15:
                    continue
                # 数量词之后（含其前两字，排除“比例”“效率”之类）到动词之间出现比率字样的，是比率的变化
                if self.RATE_WORD_RE.search(text[max(clause_start, qty.start() - 2):m.start()]):
                    continue
                issue = self._proc_issue("NUM-020", para, qty.start(), m.end())
                if issue:
                    issues.append(issue)
        return issues

    # -- 格式要素（函数） -----------------------------------------------------
    FULL_DATE_RE = re.compile(r"(\d{4})\s*年\s*(\d{1,2})\s*月(?:\s*(\d{1,2})\s*日)?")
    MD_DATE_RE = re.compile(r"(?<![\d年])(\d{1,2})\s*月\s*(\d{1,2})\s*日")

    def _check_date_padding(self, doc: Document, doc_type: str) -> List[Issue]:
        """FMT-005 月日虚位。"""
        issues: List[Issue] = []
        strict_idx = set()
        if doc.signoff_date is not None:
            strict_idx.add(doc.signoff_date)
        if doc.banji_start is not None:
            strict_idx |= {p.index for p in doc.paras[doc.banji_start:] if "印发" in p.text}
        for para in doc.paras:
            spans: List[Tuple[int, int]] = []
            for m in self.FULL_DATE_RE.finditer(para.text):
                month, day = m.group(2), m.group(3)
                if (len(month) == 2 and month.startswith("0")) or (day and len(day) == 2 and day.startswith("0")):
                    spans.append(m.span())
            for m in self.MD_DATE_RE.finditer(para.text):
                if any(s <= m.start() < e for s, e in spans):
                    continue
                month, day = m.group(1), m.group(2)
                if (len(month) == 2 and month.startswith("0")) or (len(day) == 2 and day.startswith("0")):
                    spans.append(m.span())
            for s, e in spans:
                level = "P0" if para.index in strict_idx else "P1"
                fixed = re.sub(r"(?<!\d)0(\d)(?=\s*[月日])", r"\1", para.text[s:e])
                issue = self._proc_issue("FMT-005", para, s, e, level=level, suggestion=f"改为“{fixed}”")
                if issue:
                    issues.append(issue)
        return issues

    def _check_signoff_date(self, doc: Document, doc_type: str) -> List[Issue]:
        """FMT-006 成文日期用阿拉伯数字。"""
        k = doc.signoff_date
        if k is None or not CN_DATE_ONLY_RE.match(doc.paras[k].text):
            return []
        para = doc.paras[k]
        text = para.text.strip()
        start = para.text.index(text)
        issue = self._proc_issue("FMT-006", para, start, start + len(text))
        return [issue] if issue else []

    def _check_zhusong(self, doc: Document, doc_type: str) -> List[Issue]:
        """FMT-007 主送机关后用全角冒号。"""
        issues: List[Issue] = []
        if doc.zhusong is not None:
            para = doc.paras[doc.zhusong]
            text = para.text.rstrip()
            if text.endswith(":"):
                issue = self._proc_issue("FMT-007", para, 0, len(text), excerpt=clip(text.strip(), 80),
                                         message="主送机关后用了半角冒号", suggestion="改为全角冒号“：”")
                if issue:
                    issues.append(issue)
            return issues
        if not doc.title or doc_type not in ZHUSONG_TYPES:
            return issues
        j = doc.title[-1] + 1
        if doc.subtitle is not None:
            j = doc.subtitle + 1
        if j >= len(doc.paras) - 1:
            return issues
        para = doc.paras[j]
        text = para.text.strip()
        if (len(text) <= 60 and not re.search(r"[。；：:！？]", text) and ORG_SUFFIX_RE.search(text)
                and not HEADING_START_RE.match(text) and not DATE_ONLY_RE.match(text)):
            issue = self._proc_issue("FMT-007", para, 0, len(para.text), excerpt=clip(text, 80),
                                     message="主送机关后缺少全角冒号")
            if issue:
                issues.append(issue)
        return issues

    # -- 结构（函数） ---------------------------------------------------------
    def _check_title(self, doc: Document, doc_type: str) -> List[Issue]:
        """STR-005 标题末尾标点；STR-010 文种混用或与参数不一致。"""
        issues: List[Issue] = []
        if not doc.title:
            return issues
        last = doc.paras[doc.title[-1]]
        text = last.text.rstrip()
        if text and text[-1] in PUNCT_TAIL:
            issue = self._proc_issue("STR-005", last, len(text) - 1, len(text),
                                     excerpt=clip(text.strip(), 80),
                                     suggestion=f"删除标题末尾的“{text[-1]}”")
            if issue:
                issues.append(issue)
        title = doc.title_text.rstrip(PUNCT_TAIL)
        first = doc.paras[doc.title[0]]
        if title.endswith("请示报告"):
            issue = self._proc_issue("STR-010", first, 0, len(first.text), excerpt=clip(first.text.strip(), 80),
                                     message="标题文种写作“请示报告”，请示和报告不能合用")
            if issue:
                issues.append(issue)
        else:
            inferred = doc.infer_doc_type()
            if (inferred and doc_type != inferred and doc_type in STATUTORY_TYPES
                    and inferred in STATUTORY_TYPES):
                issue = self._proc_issue(
                    "STR-010", first, 0, len(first.text), excerpt=clip(first.text.strip(), 80), level="P2",
                    message=f"标题文种为“{inferred}”，与指定的文书类型“{doc_type}”不一致",
                    suggestion="核对文种：请求批准用请示，汇报工作用报告，布置工作用通知")
                if issue:
                    issues.append(issue)
        return issues

    def _attach_list_indices(self, doc: Document) -> Set[int]:
        """附件说明后续行（“2. 某某表”）不参与层次序号检查。"""
        result: Set[int] = set()
        if doc.attach_desc is None:
            return result
        k = doc.attach_desc + 1
        while k < len(doc.paras) and re.match(r"^\s*\d{1,2}\s*[.．、]", doc.paras[k].text):
            result.add(k)
            k += 1
        return result

    def _check_heading_sequence(self, doc: Document, doc_type: str) -> List[Issue]:
        """STR-001 层次序号连续性与逐级使用。"""
        issues: List[Issue] = []
        skip = self._attach_list_indices(doc)
        if doc.attach_desc is not None:
            skip.add(doc.attach_desc)
        last: Dict[int, Optional[int]] = {1: None, 2: None, 3: None, 4: None}
        cur_level = 0
        seen_any = False
        jump_issue: Optional[Issue] = None
        jump_count = 0
        list_reset = False
        for k in doc.body_indices():
            para = doc.paras[k]
            if para.in_table or k in skip:
                continue
            if ATTACH_PAGE_RE.match(para.text):
                last = {1: None, 2: None, 3: None, 4: None}
                cur_level, seen_any = 0, False
                continue
            found = heading_of(para.text)
            if not found:
                # 以冒号结尾的引导句（如“一等助学金：”）之后，阿拉伯数字序号可以重新从 1 编起
                if para.text.rstrip().endswith(("：", ":")):
                    list_reset = True
                continue
            level, num = found
            # 引导句之后从 1 重新编号的，视为新列表；接着上一组继续编号的，照常检查连续性
            if list_reset and level >= 3 and num == 1:
                for lv in range(level, 5):
                    last[lv] = None
            list_reset = False
            head_text = para.text.strip()[:30]
            if seen_any and level > cur_level + 1 and cur_level > 0:
                jump_count += 1
                if jump_issue is None:
                    jump_issue = self._proc_issue(
                        "STR-001", para, 0, len(para.text), level="P2", excerpt=clip(head_text, 30),
                        message="层次序数越级使用（如“一、”之后直接用“1.”，跳过“（一）”）",
                        suggestion="按 一、（一）1.（1） 逐级使用")
            prev = last[level]
            expected = 1 if prev is None else prev + 1
            if num != expected:
                if prev is None:
                    msg = f"本级序号未从 1 开始（出现的是第 {num} 项）"
                    lvl = "P2" if not seen_any else "P1"
                elif num == prev:
                    msg, lvl = f"层次序号重复（第 {num} 项出现两次）", "P1"
                elif num < prev:
                    msg, lvl = f"层次序号乱序（第 {prev} 项之后出现第 {num} 项）", "P1"
                else:
                    msg, lvl = f"层次序号跳号（第 {prev} 项之后直接是第 {num} 项）", "P1"
                issue = self._proc_issue("STR-001", para, 0, len(para.text), level=lvl,
                                         excerpt=clip(head_text, 30), message=msg)
                if issue:
                    issues.append(issue)
            last[level] = num
            for deeper in range(level + 1, 5):
                last[deeper] = None
            cur_level = level
            seen_any = True
        if jump_issue:
            if jump_count > 1:
                jump_issue.message += f"（全文共 {jump_count} 处，只列第一处）"
                jump_issue.count = jump_count
            issues.append(jump_issue)
        return issues

    COUNT_RE = re.compile(
        r"(?<![第\d一二三四五六七八九十])(十[一二三四五六七八九]?|[两二三四五六七八九])"
        r"(?:(?:个|项|点|条|件|大)(?:方面|措施|任务|重点工作|重点任务|工作|意见|建议|体会|要求|问题|环节|阶段|举措"
        r"|原则|目标|机制|制度|实事|行动|工程|做法|经验|特点|成效|不足|短板|方向|内容|事项|领域|步骤|变化)|方面)"
        r"(?=(?:的|方面的)?[一-鿿]{0,6}(?:[：:，,。；]|$))"
    )
    INLINE_MARKERS = (
        re.compile(r"(?<![一-鿿])([一二三四五六七八九十]{1,3})是|(?<=[。；，：\s])([一二三四五六七八九十]{1,3})是"),
        re.compile(r"第([一二三四五六七八九十]{1,3})[，,、是]"),
        re.compile(r"[（(]([一二三四五六七八九十]{1,3})[）)]"),
        re.compile(r"[（(](\d{1,2})[）)]"),
    )

    @classmethod
    def _inline_count(cls, text: str) -> int:
        """统计“一是……二是……”等段内列举的连续条数。"""
        best = 0
        for regex in cls.INLINE_MARKERS:
            nums = []
            for m in regex.finditer(text):
                raw = next(g for g in m.groups() if g)
                nums.append(cn_to_int(raw))
            if not nums or nums[0] != 1:
                continue
            count = 0
            for n in nums:
                if n == count + 1:
                    count = n
            best = max(best, count)
        return best

    @staticmethod
    def _inside_quote(text: str, pos: int) -> bool:
        before = text[:pos]
        return before.count("“") > before.count("”") or before.count("\"") % 2 == 1

    def _check_count_summary(self, doc: Document, doc_type: str) -> List[Issue]:
        """STR-002 数字概括与实际条数不一致（启发式，P2）。"""
        issues: List[Issue] = []
        body = doc.body_indices()
        body_set = set(body)
        heading_cache = {k: heading_of(doc.paras[k].text) for k in body}
        for pos_in_body, k in enumerate(body):
            para = doc.paras[k]
            if para.in_table:
                continue
            for m in self.COUNT_RE.finditer(para.text):
                if self._inside_quote(para.text, m.start()):
                    continue
                n = cn_to_int(m.group(1))
                if n < 2:
                    continue
                tail = para.text[m.end():]
                found = self._inline_count(tail)
                if found < 2 and not self._external_aspects(doc, body, pos_in_body, heading_cache, para.text, m.start()):
                    found = self._following_count(doc, body, pos_in_body, heading_cache, body_set)
                if found >= 2 and found != n:
                    issue = self._proc_issue(
                        "STR-002", para, m.start(), m.end(),
                        message=f"数字概括为“{m.group(0)}”，其后实际列出 {found} 项",
                        suggestion="核对条数，改正概括数字或补删条目")
                    if issue:
                        issues.append(issue)
        return issues

    # 对照检查材料常见“围绕通知规定的四个方面，对照检查如下”：方面来自上级通知，不对应下文标题
    INSPECT_WORDS_RE = re.compile(r"对照检查|查摆|检视|剖析")
    SCOPE_LEAD_RE = re.compile(r"(?:围绕|按照|对照|聚焦)[^。；！？]{0,30}$")

    @classmethod
    def _external_aspects(cls, doc: Document, body: List[int], pos: int,
                          heading_cache: Dict[int, Optional[Tuple[int, int]]], text: str, start: int) -> bool:
        """数字概括所指的方面来自上级通知、不对应后续标题时返回 True（只比对同段“一是……”）。"""
        sent_start = max(text.rfind(ch, 0, start) for ch in "。！？；") + 1
        sent_end_candidates = [i for i in (text.find(ch, start) for ch in "。！？；") if i >= 0]
        sentence = text[sent_start:min(sent_end_candidates) if sent_end_candidates else len(text)]
        if cls.INSPECT_WORDS_RE.search(sentence):
            return True
        if cls.SCOPE_LEAD_RE.search(text[sent_start:start]):
            # 只在全文开头（还没有任何层次标题）时适用；章节内部的“围绕三个方面”照常与下一级标题比对
            return not any(heading_cache.get(body[back]) for back in range(pos + 1))
        return False

    def _following_count(self, doc: Document, body: List[int], pos: int,
                         heading_cache: Dict[int, Optional[Tuple[int, int]]], body_set: Set[int]) -> int:
        """在后续段落中统计下一级标题或“一是”段的条数。"""
        k0 = body[pos]
        own = heading_cache.get(k0)
        ctx_level = own[0] if own else 0
        if not own:
            for back in range(pos - 1, -1, -1):
                h = heading_cache.get(body[back])
                if h:
                    ctx_level = h[0]
                    break
        child_level: Optional[int] = None
        count = 0
        yishi = 0
        for idx in range(pos + 1, min(len(body), pos + 40)):
            k = body[idx]
            h = heading_cache.get(k)
            text = doc.paras[k].text.strip()
            if h:
                level = h[0]
                if level <= ctx_level:
                    break
                if child_level is None:
                    child_level = level
                if level < child_level:
                    break
                if level == child_level:
                    count += 1
                continue
            if child_level is None:
                m = re.match(r"^([一二三四五六七八九十]{1,3})是", text)
                if m and cn_to_int(m.group(1)) == yishi + 1:
                    yishi += 1
                elif idx > pos + 1 and yishi == 0:
                    break  # 概括句之后紧跟的不是条目，放弃判断
        return count if child_level is not None else yishi

    QINGSHI_END_RE = re.compile(
        r"(?:妥否|当否|可否|是否妥当)[，,]?\s*请(?:批示|批复|审批|审定|示)|请予以?(?:审批|批准|批复|审定|批示|支持)"
        r"|请(?:批示|批复|审批|审定)|恳请(?:批准|审批|批复|审定)|以上请示[，,]?\s*请"
    )
    REQUEST_RE = re.compile(
        r"(?:妥否|当否|可否)[，,]?\s*请(?:批示|批复|审批|示)|请予以?(?:审批|批准|批复|审定)"
        r"|(?<![报提申呈])请(?:批示|审批|批准|批复)(?![件人])"
    )
    PLEA_RE = re.compile(r"恳请[^。；\n]{0,20}(?:批准|审批|批复|同意|拨付|核拨|解决)")

    def _check_qingshi(self, doc: Document, doc_type: str) -> List[Issue]:
        """STR-003 请示结尾；STR-006 主送；STR-007 特此报告。"""
        issues: List[Issue] = []
        if doc_type != "请示":
            if doc_type in REPORT_TYPES:
                issues.extend(self._check_main_recipient(doc, doc_type))
            return issues
        body = doc.body_indices()
        if body and self._proc_applies("STR-003", doc_type):
            if not any(self.QINGSHI_END_RE.search(doc.paras[k].text) for k in body):
                last_k = body[-1]
                for k in reversed(body):
                    if not DATE_ONLY_RE.match(doc.paras[k].text) and not ATTACH_DESC_RE.match(doc.paras[k].text):
                        last_k = k
                        if len(doc.paras[k].text.strip()) > 20:
                            break
                para = doc.paras[last_k]
                issue = self._proc_issue("STR-003", para, 0, len(para.text), excerpt=clip(para.text, 40))
                if issue:
                    issues.append(issue)
        for k in body:
            para = doc.paras[k]
            m = re.search(r"特此报告", para.text)
            if m and self._proc_applies("STR-007", doc_type):
                issue = self._proc_issue("STR-007", para, m.start(), m.end())
                if issue:
                    issues.append(issue)
        issues.extend(self._check_main_recipient(doc, doc_type))
        return issues

    def _check_main_recipient(self, doc: Document, doc_type: str) -> List[Issue]:
        if doc.zhusong is None or not self._proc_applies("STR-006", doc_type):
            return []
        # 签报是内部文书，按规定主送领导职务，不适用“请示不向领导个人报送”（也不查多个主送）：
        # doc_type 为签报，或稿首有“签报”“呈批件”表头，或全文有“会签情况”栏
        if doc_type == "签报" or doc.qianbao_marked():
            return []
        para = doc.paras[doc.zhusong]
        text = para.text.strip().rstrip("：:")
        # 先剔除【待补…】【待核实…】标记再判断（标记里的逗号、顿号不是多个主送机关）；剔除后为空说明主送待补，不查
        text = ZHUSONG_MARKER_RE.sub("", text).strip().rstrip("：:").strip()
        if not text:
            return []
        if re.search(r"(?:书记|市长|县长|区长|省长|局长|主任|院长|校长|董事长|总经理|厅长|部长|处长|同志)$", text):
            issue = self._proc_issue("STR-006", para, 0, len(para.text), excerpt=clip(para.text.strip(), 80),
                                     message=f"{doc_type}主送领导个人")
            return [issue] if issue else []
        if doc_type == "请示" and re.search(r"[、，,]", text):
            issue = self._proc_issue("STR-006", para, 0, len(para.text), excerpt=clip(para.text.strip(), 80),
                                     message="请示主送了多个机关")
            return [issue] if issue else []
        return []

    def _check_baogao(self, doc: Document, doc_type: str) -> List[Issue]:
        """STR-004 报告夹带请示事项。"""
        if not self._proc_applies("STR-004", doc_type):
            return []
        issues: List[Issue] = []
        for k in doc.body_indices():
            para = doc.paras[k]
            spans: List[Tuple[int, int]] = []
            for m in self.REQUEST_RE.finditer(para.text):
                spans.append(m.span())
                issue = self._proc_issue("STR-004", para, m.start(), m.end())
                if issue:
                    issues.append(issue)
            for m in self.PLEA_RE.finditer(para.text):
                if any(s < m.end() and m.start() < e for s, e in spans):
                    continue
                issue = self._proc_issue("STR-004", para, m.start(), m.end(), level="P1",
                                         message="报告中出现疑似请示事项（请求批准、拨付等）",
                                         suggestion="需要上级批准的事项另行请示；确属反映困难的，改为陈述情况")
                if issue:
                    issues.append(issue)
        return issues

    def _check_han(self, doc: Document, doc_type: str) -> List[Issue]:
        """STR-008 函中用“请批示”；STR-011 下行文用请示结束语。"""
        issues: List[Issue] = []
        if self._proc_applies("STR-008", doc_type):
            for k in doc.body_indices():
                para = doc.paras[k]
                for m in re.finditer(r"请(?:批示|批复)", para.text):
                    issue = self._proc_issue("STR-008", para, m.start(), m.end())
                    if issue:
                        issues.append(issue)
        if self._proc_applies("STR-011", doc_type):
            for k in doc.body_indices():
                para = doc.paras[k]
                for m in self.REQUEST_RE.finditer(para.text):
                    issue = self._proc_issue("STR-011", para, m.start(), m.end())
                    if issue:
                        issues.append(issue)
        return issues

    def _check_shangxing(self, doc: Document, doc_type: str) -> List[Issue]:
        """STR-009 上行文有发文字号却没有签发人。"""
        if not self._proc_applies("STR-009", doc_type) or not doc.wenhao_present or doc.qianfa_present:
            return []
        for k in sorted(doc.head):
            para = doc.paras[k]
            m = WENHAO_RE.search(para.text)
            if m:
                issue = self._proc_issue("STR-009", para, m.start(), m.end())
                return [issue] if issue else []
        return []

    # -- 要素完整性（函数） ---------------------------------------------------
    ELM_TITLE_TYPES = STATUTORY_TYPES | REPORT_TYPES
    ELM_REASON_TYPES = {"通知", "通报", "请示", "报告", "整改报告", "自查报告", "自评报告", "批复", "函", "意见", "决定", "议案", "通告"}
    ELM_ZHUSONG_TYPES = {"通知", "请示", "报告", "整改报告", "自查报告", "自评报告", "批复", "函", "议案"}
    ELM_SIGN_TYPES = {"通知", "通报", "请示", "报告", "整改报告", "自查报告", "自评报告", "批复", "函", "议案",
                      "意见", "决定", "公告", "通告"}
    ELM_TITLE_KIND_RE = re.compile(r"(?:决议|决定|命令|令|公报|公告|通告|意见|通知|通报|报告|请示|批复|议案|函|纪要)$")
    ELM_ORG_ONLY_RE = re.compile(
        r"^[^，。；：、]{2,16}?(?:局|厅|委员会|办公室|办公厅|政府|党委|党组|学校|医院|公司|集团|大学|学院|中心|街道|"
        r"办事处|协会|研究院|研究所|团委|党支部|支部)$")
    ELM_CLOSING_RES = {
        "报告": re.compile(r"特此(?:报告|汇报|呈报)|专此(?:报告|汇报)|以上(?:报告|情况|汇报)|请(?:予)?(?:审阅|审议|阅示|阅)"),
        "函": re.compile(r"特此(?:函告|函复|复函|函达)|专此(?:函告|函复|函达)|盼复|请(?:予)?(?:函复|复函|研究函复)|"
                         r"(?<![一-鿿])此复|请予(?:以)?(?:支持|协助|配合|办理|研究|同意|批准|审批)|为盼|为荷|是荷|敬请|烦请|恳请|"
                         r"谢谢|感谢"),
        "批复": re.compile(r"此复|特此(?:批复|函复)|请(?:遵照|按照|认真)(?:执行|落实|办理)|遵照执行"),
        "公告": re.compile(r"特此公告|现予公告|予以公告|特此通告"),
        "通告": re.compile(r"特此通告|特此公告|现予通告"),
    }
    ELM_CLOSING_ALIAS = {"整改报告": "报告", "自查报告": "报告", "自评报告": "报告"}
    ELM_CLOSING_HINT = {
        "报告": "结尾写“特此报告。”", "函": "结尾写“特此函告。”“请予支持。”或复函的“特此函复。”",
        "批复": "结尾写“此复。”或“特此批复。”", "公告": "结尾写“特此公告。”", "通告": "结尾写“特此通告。”",
    }
    ELM_DATE_TAIL_RE = re.compile(
        r"(?:(?:19|20)\d{2}\s*年\s*\d{1,2}\s*月\s*\d{1,2}\s*日|[〇○零一二三四五六七八九OＯ]{4}\s*年\s*[一二三四五六七八九十]{1,3}\s*月\s*"
        r"[一二三四五六七八九十]{1,3}\s*日|(?:19|20)\d{2}[.\-/]\d{1,2}[.\-/]\d{1,2})\s*$")
    ELM_ORG_EXTRA_RE = re.compile(
        r"(?:学|园|馆|社|台|庭|行|厂|矿|场|团|盟|联|总会|协会|学会|商会|工会|居委会|村委会|支部|总支|常委会|管委会|公司|银行|法院|检察院)$")
    ELM_SIGN_STRIP_RE = re.compile(r"[（(]\s*(?:盖章|公章|章|印|单位盖章)\s*[）)]|[\s　]+")
    ELM_ATTACH_REF_RE = re.compile(r"附件\s*([0-9]{1,2}|[一二三四五六七八九十]{1,2})?(?![说页])")

    @staticmethod
    def _elm_clean_title(text: str) -> str:
        return re.sub(r"^[#*_\s　]+|[*_\s　]+$", "", text).rstrip(PUNCT_TAIL)

    @staticmethod
    def _elm_marker(doc: Document, keywords: str) -> Optional[Tuple[Para, "re.Match[str]"]]:
        """全文中是否有与该要素相关的【待补：…】标记（视为作者已注意到，降为 P2 提示）。"""
        pat = re.compile(keywords)
        for para in doc.paras:
            for m in ELM_MARKER_RE.finditer(para.text):
                if pat.search(m.group(0)):
                    return para, m
        return None

    @staticmethod
    def _elm_zhusong_like(text: str) -> bool:
        stripped = text.strip().rstrip("：:").strip()
        return (0 < len(stripped) <= 60 and "。" not in stripped
                and not re.search(r"如下|以下|通知|报告|情况|事项|^现将|^为|^根据|^按照|^经|^兹|^特", stripped))

    def _elm_find_signoff(self, doc: Document) -> Tuple[Optional[int], Optional[int]]:
        """返回（落款单位所在段落序号，成文日期所在段落序号）；找不到为 None。"""
        paras = doc.paras
        end = doc.banji_start if doc.banji_start is not None else len(paras)
        skip = set(doc.head) | set(doc.title)
        date_k: Optional[int] = doc.signoff_date
        if date_k is None:
            for k in range(end - 1, max(-1, end - 12), -1):
                if k in skip:
                    continue
                text = paras[k].text.strip()
                if len(text) <= 60 and self.ELM_DATE_TAIL_RE.search(text) and not HEADING_START_RE.match(text):
                    date_k = k
                    break

        def org_like(text: str, loose: bool = False) -> bool:
            cleaned = self.ELM_SIGN_STRIP_RE.sub("", text)
            if not cleaned or len(cleaned) > 40 or re.search(r"[。；：:！？，,]", cleaned):
                return False
            if HEADING_START_RE.match(cleaned) or ATTACH_DESC_RE.match(cleaned) or "附件" in cleaned:
                return False
            if (ORG_SUFFIX_RE.search(cleaned) or self.ELM_ORG_EXTRA_RE.search(cleaned)
                    or ELM_MARKER_RE.search(cleaned)):
                return True
            # 紧挨成文日期上一行的短行（无标点、不是附注）：多半是署名，机关名称写法特殊时也认
            return loose and 3 <= len(cleaned) <= 30 and not cleaned.startswith(("（", "(", "联系", "电话", "注"))

        if date_k is not None:
            line = paras[date_k].text.strip()
            rest = self.ELM_DATE_TAIL_RE.sub("", line).strip()
            if rest and org_like(rest):
                return date_k, date_k
            for n, k in enumerate((date_k - 1, date_k - 2)):
                if k >= 0 and k not in skip and org_like(paras[k].text.strip(), loose=(n == 0)):
                    return k, date_k
            return None, date_k
        for k in range(end - 1, max(-1, end - 8), -1):
            if k not in skip and org_like(paras[k].text.strip()):
                return k, None
        return None, None

    def _check_elements(self, doc: Document, doc_type: str) -> List[Issue]:
        """ELM-001～ELM-007 按文种检查必备要素：标题、主送、结语、落款、成文日期、附件说明、纪要出席列席。
        【待补：…】占位视为作者已注意到，报 P2 提示；正文不足 120 字的片段不检查。"""
        if not doc.paras or doc_type not in DOC_TYPES:
            return []
        if doc.body_char_count() < 120:
            return []
        issues: List[Issue] = []

        def emit(rid: str, para: Optional[Para], start: int, end: int, **kw: Any) -> None:
            if not self._proc_applies(rid, doc_type):
                return
            issue = self._proc_issue(rid, para, start, end, **kw)
            if issue:
                issues.append(issue)

        anchor = doc.paras[doc.title[0]] if doc.title else doc.paras[0]
        body = doc.body_indices()
        # 标题不是“关于……的报告”形式的“报告”（如“××分析报告”“××评估报告”）多为内部材料，不按公文报告查主送、结语、落款和日期
        informal_report = (doc_type == "报告" and bool(doc.title)
                           and not re.search(r"关于.*报告$|的报告$", self._elm_clean_title(doc.title_text)))

        def marker_or(rid: str, keywords: str, message: str, suggestion: str, para: Optional[Para] = None,
                      start: int = 0, end: int = 0, excerpt: str = "") -> None:
            found = self._elm_marker(doc, keywords)
            if found:
                mpara, m = found
                emit(rid, mpara, m.start(), m.end(), level="P2",
                     message=message + "（已用【待补】标注，定稿前补齐）", suggestion="按【待补】提示补齐后删去标记")
            else:
                emit(rid, para, start, end, message=message, suggestion=suggestion, excerpt=excerpt)

        # ELM-001 标题：发文机关、事由、文种
        if doc_type in self.ELM_TITLE_TYPES:
            if not doc.title:
                marker_or("ELM-001", "标题|事由|文种|发文机关", f"{doc_type}缺少标题",
                          "标题由发文机关名称、事由和文种组成，如“关于……的" + doc_type + "”",
                          anchor, 0, len(anchor.text), clip(anchor.text, 40))
            else:
                title = self._elm_clean_title(doc.title_text)
                tpara = doc.paras[doc.title[0]]
                if ELM_MARKER_RE.search(title):
                    m = ELM_MARKER_RE.search(tpara.text)
                    emit("ELM-001", tpara, m.start() if m else 0, m.end() if m else len(tpara.text), level="P2",
                         message="标题含【待补】，定稿前补齐", suggestion="按【待补】提示补齐标题要素后删去标记")
                else:
                    km = self.ELM_TITLE_KIND_RE.search(title)
                    if not km:
                        emit("ELM-001", tpara, 0, len(tpara.text), excerpt=clip(title, 60),
                             message=f"标题没有文种（{doc_type}应以“{doc_type}”结尾）",
                             suggestion="标题补全文种，如“关于……的" + doc_type + "”")
                    elif doc_type in self.ELM_REASON_TYPES:
                        rest = re.sub(r"的$", "", title[:km.start()])
                        reason = ""
                        if "关于" in rest:
                            reason = rest.split("关于", 1)[1]
                        elif not self.ELM_ORG_ONLY_RE.match(rest):
                            reason = rest
                        if len(reason.strip()) < 2:
                            emit("ELM-001", tpara, 0, len(tpara.text), excerpt=clip(title, 60),
                                 message="标题缺少事由（只有发文机关加文种，或“关于”后没有内容）",
                                 suggestion="补写事由，如“关于……的" + doc_type + "”")

        # ELM-002 主送机关
        if doc_type in self.ELM_ZHUSONG_TYPES and not (doc_type == "通知" and doc.public_facing) and not informal_report:
            has, para_z = False, None
            if doc.zhusong is not None and self._elm_zhusong_like(doc.paras[doc.zhusong].text):
                has, para_z = True, doc.paras[doc.zhusong]
            elif any(re.match(r"^\s*主送\s*[：:]", p.text) for p in doc.paras):
                has = True
            else:
                start_k = (doc.title[-1] + 1) if doc.title else 0
                for p in doc.paras[start_k:start_k + 4]:
                    if p.index in doc.head:
                        continue
                    if p.text.strip().endswith(("：", ":")) and self._elm_zhusong_like(p.text):
                        has, para_z = True, p
                        break
            if has and para_z is not None and not ZHUSONG_MARKER_RE.sub("", para_z.text).strip().rstrip("：:").strip():
                emit("ELM-002", para_z, 0, len(para_z.text), level="P2", excerpt=clip(para_z.text.strip(), 40),
                     message="主送机关为【待补】，定稿前补齐", suggestion="按【待补】提示补齐主送机关")
            elif not has:
                marker_or("ELM-002", "主送|受文|报送机关|机关", f"{doc_type}缺少主送机关",
                          "标题下空一行顶格写主送机关，后标全角冒号", anchor, 0, len(anchor.text), clip(anchor.text, 40))

        # ELM-003 规范结语（请示的结语由 STR-003 检查）
        closing_type = self.ELM_CLOSING_ALIAS.get(doc_type, doc_type)
        closing_re = self.ELM_CLOSING_RES.get(closing_type)
        if closing_re is not None and body and not informal_report:
            if not any(closing_re.search(doc.paras[k].text) for k in body):
                last = None
                for k in reversed(body):
                    text = doc.paras[k].text
                    if (not DATE_ONLY_RE.match(text) and not ATTACH_DESC_RE.match(text)
                            and len(text.strip()) > 12 and re.search(r"[。！？；]", text)):
                        last = doc.paras[k]
                        break
                if last is not None:
                    emit("ELM-003", last, 0, len(last.text), excerpt=clip(last.text, 40),
                         message=f"{doc_type}缺少规范结语", suggestion=self.ELM_CLOSING_HINT.get(closing_type, ""))

        # ELM-004 落款、ELM-005 成文日期
        if doc_type in self.ELM_SIGN_TYPES and body and not informal_report:
            org_k, date_k = self._elm_find_signoff(doc)
            tail = None
            for k in reversed(body):
                if len(doc.paras[k].text.strip()) > 2:
                    tail = doc.paras[k]
                    break
            tail = tail or doc.paras[body[-1]]
            if org_k is None:
                marker_or("ELM-004", "落款|署名|发文机关|单位", f"{doc_type}缺少落款（发文机关署名）",
                          "正文（或附件说明）之后写发文机关全称或规范化简称，位于成文日期之上",
                          tail, 0, len(tail.text), clip(tail.text, 40))
            if date_k is None:
                marker_or("ELM-005", "日期|成文", f"{doc_type}缺少成文日期",
                          "落款之下写成文日期，如“2026年10月16日”（阿拉伯数字，年月日标全）",
                          tail, 0, len(tail.text), clip(tail.text, 40))
            elif ELM_MARKER_RE.search(doc.paras[date_k].text) and not self.ELM_DATE_TAIL_RE.search(doc.paras[date_k].text):
                m = ELM_MARKER_RE.search(doc.paras[date_k].text)
                emit("ELM-005", doc.paras[date_k], m.start(), m.end(), level="P2",
                     message="成文日期为【待补】，定稿前补齐", suggestion="按【待补】提示补齐")

        # ELM-006 附件说明与正文引用一致
        self._elm_attachments(doc, doc_type, emit, marker_or)

        # ELM-007 纪要的出席、列席人员
        if doc_type == "纪要":
            text = "\n".join(doc.paras[k].text for k in body)
            if not re.search(r"出席|与会人员|参会人员|参加人员|到会", text):
                marker_or("ELM-007", "出席|与会|参会", "纪要未列出席人员",
                          "会议概况或文末列出出席人员（单位、姓名、职务），有请假的写明请假人员",
                          anchor, 0, len(anchor.text), clip(anchor.text, 40))
            elif not re.search(r"列席|请假|缺席", text) and not ELM_MARKER_RE.search(text):
                emit("ELM-007", anchor, 0, len(anchor.text), level="P2", excerpt=clip(anchor.text, 40),
                     message="纪要未见列席或请假人员（没有列席人员的，宜注明“无”）",
                     suggestion="出席人员之后写列席、请假人员；确无的可注明“无”")
        return issues

    def _elm_attachments(self, doc: Document, doc_type: str, emit: Callable[..., None],
                         marker_or: Callable[..., None]) -> None:
        if doc_type in ("讲话稿", "党课讲稿", "研讨发言", "心得体会", "会议记录", "简报", "其他"):
            return
        paras = doc.paras
        desc_k = doc.attach_desc
        listed = 0
        list_idx: Set[int] = set()
        if desc_k is not None:
            list_idx.add(desc_k)
            head_text = ATTACH_DESC_RE.sub("", paras[desc_k].text, count=1).strip()
            if re.match(r"(\d{1,2})\s*[.．、]", head_text):
                listed = 1
                k = desc_k + 1
                while k < len(paras):
                    m2 = re.match(r"^\s*(\d{1,2})\s*[.．、]", paras[k].text)
                    if not m2 or int(m2.group(1)) != listed + 1:
                        break
                    listed += 1
                    list_idx.add(k)
                    k += 1
            else:
                listed = 1 if head_text else 0
        page_k = next((p.index for p in paras if ATTACH_PAGE_RE.match(p.text) and p.index not in doc.head
                       and (desc_k is None or p.index > desc_k)), None)
        refs: List[Tuple[Para, int, int, Optional[int]]] = []
        for k in doc.body_indices():
            if k in list_idx or (page_k is not None and k >= page_k) or paras[k].in_table:
                continue
            if ATTACH_DESC_RE.match(paras[k].text) or ATTACH_PAGE_RE.match(paras[k].text):
                continue
            for m in self.ELM_ATTACH_REF_RE.finditer(paras[k].text):
                num_txt = m.group(1)
                number: Optional[int] = None
                if num_txt:
                    number = int(num_txt) if num_txt.isdigit() else cn_to_int(num_txt)
                refs.append((paras[k], m.start(), m.end(), number))
        if refs and desc_k is None and page_k is None:
            para, s, e, _ = refs[0]
            marker_or("ELM-006", "附件", "正文提到附件，但没有附件说明",
                      "正文下空一行写“附件：附件名称”，多个附件标顺序号；或删去“见附件”",
                      para, s, e, para.text[max(0, s - 4):e + 4])
            return
        if desc_k is not None and listed >= 1:
            for para, s, e, number in refs:
                if number is not None and number > listed:
                    emit("ELM-006", para, s, e, message=f"正文引用“附件{number}”，附件说明只列了 {listed} 个附件",
                         suggestion="核对附件序号，补写附件说明或改正文引用")
                    break
            if listed >= 2 and refs:
                cited = {n for _, _, _, n in refs if n is not None}
                if cited:
                    missing = [n for n in range(1, listed + 1) if n not in cited]
                    if missing:
                        para = paras[desc_k]
                        emit("ELM-006", para, 0, len(para.text), level="P2", excerpt=clip(para.text.strip(), 40),
                             message="附件" + "、".join(str(n) for n in missing) + "在正文中未被引用",
                             suggestion="正文中说明各附件的用途（如“详见附件2”），或核对附件说明")

    # -- 评价词与证据强度匹配（函数） -----------------------------------------
    EVD_DATE_RE = re.compile(
        r"(?:19|20)?\d{2}\s*年(?:\s*\d{1,2}\s*月(?:\s*\d{1,2}\s*[日号])?)?|\d{1,2}\s*月(?:\s*\d{1,2}\s*[日号])?|\d{1,2}\s*[日号]"
        r"|[〔\[（(]\s*(?:19|20)\d{2}\s*[〕\]）)]\s*\d*\s*号?")
    EVD_NUM_RE = re.compile(
        r"(?<![\d.])\d[\d,]*(?:\.\d+)?\s*(?:%|％|‰|个百分点|百分点|倍|万|亿|千|百|元|人次|人|次|件|个|户|家|项|起|天|小时|分钟|"
        r"平方米|平方公里|公里|吨|例|场|批|台|套|所|名|条|篇|期|届|位|岁|分|级|类|辆|张|份|宗|处|间|座|亩|公顷|斤|米|克|毫米|度|"
        r"课时|门|本|册|周|个月|周年|年)"
        r"|[一二三四五六七八九十百千]+(?:余|多)?(?:个|次|项|家|批|件|人|场|期|倍|成|条|所|类|名|户|起|台|套)")
    EVD_ANYNUM_RE = re.compile(r"\d|[一二三四五六七八九十百千万]+(?:余|多)?(?:个|次|项|家|批|件|人|场|期|倍|成|条|所|类|名|户|起)")
    EVD_COMPARE_RE = re.compile(
        r"同比|环比|较(?:上年|去年|上月|上季|年初|同期|上期|往年|以往|前)|比(?:上年|去年|上月|上季|年初|同期|上期|往年|以往)|"
        r"较(?!大|高|好|快|多|强|低|小|为|明显|显著|有)[^，。；]{0,8}(?:增长|增加|提高|提升|下降|减少|降低|上升|增幅|降幅)|"
        r"对比|相比|相较|高于|低于|超过|超出|翻了?一番|翻番|"
        r"增幅|降幅|增速|位居|位列|排名|第[一二三四五六七八九十\d]+名|名列前茅|居[全省市县区校院]{1,3}第|"
        r"(?:提高|提升|增长|增加|下降|减少|降低|缩短|压缩)了|(?:从|由).{1,20}(?:提高|提升|增至|降至|下降|上升|增加|减少|缩短|压缩)|"
        r"降至|增至|升至|提高到|提升到|下降到|减少到|缩短到|压缩到")
    EVD_CASE_RE = re.compile(r"例如|比如|譬如|以[^，。；]{1,20}为例|案例|个案|典型|如[“《「]")
    EVD_SENT_RE = re.compile(r"[^。！？；]*[。！？；]+[”’）)]*|[^。！？；]+$")
    EVD_DIRECTIVE_STRICT = re.compile(
        r"力争|争取|确保|务必|务求|力求|必须|计划|拟|目标|预期|预计|须|应当|应该|需要|要求|力促|以期|争创|确保")
    EVD_DIRECTIVE_LOOSE = re.compile(
        r"切实|着力|努力|坚决|(?<![主重需概纲首摘必不想紧纲])要(?![求素点领害闻案])|将|为了|旨在|进一步|继续|加快|大力|全力|尽快|尽早")

    EVD_DIRECTIVE_PREV = re.compile(r"切实|着力|坚决|各[^，。；]{1,8}要|要(?=[^求素点领害闻案])|必须|将")

    def _evd_units(self, doc: Document) -> List[Tuple[Para, int, int, str, bool]]:
        """正文按句切分，返回 [(段落, 句首, 句尾, 句子, 是否层次标题或以冒号结尾)]。引号内的句末标点不切。"""
        units: List[Tuple[Para, int, int, str, bool]] = []
        for k in doc.body_indices():
            para = doc.paras[k]
            if para.in_table or ATTACH_DESC_RE.match(para.text) or ATTACH_PAGE_RE.match(para.text):
                continue
            text = para.text
            starts: List[Tuple[int, int]] = []
            depth, begin = 0, 0
            for pos, ch in enumerate(text):
                if ch in "“‘「『":
                    depth += 1
                elif ch in "”’」』" and depth:
                    depth -= 1
                elif ch in "。！？；" and depth == 0:
                    end = pos + 1
                    while end < len(text) and text[end] in "”’）)":
                        end += 1
                    starts.append((begin, end))
                    begin = end
            if begin < len(text):
                starts.append((begin, len(text)))
            single = len(starts) == 1
            for s, e in starts:
                piece = text[s:e]
                if not piece.strip():
                    continue
                heading = (single and len(piece.strip()) <= 40 and not re.search(r"[。！？；]$", piece.strip())) \
                    or piece.rstrip().endswith(("：", ":"))
                units.append((para, s, e, piece, heading))
        return units

    def _check_evidence(self, doc: Document, doc_type: str) -> List[Issue]:
        """EVD-001～EVD-003 评价词与证据强度匹配：强评价（显著、大幅、全面完成……）同句或相邻句须有数字或对比，
        中评价（明显、有效降低……）须有事实支撑；弱评价（一定、初步）不查。"""
        cfg = self.evidence
        if not cfg or not doc.paras or not (self._proc_applies("EVD-001", doc_type) or self._proc_applies("EVD-002", doc_type)):
            return []
        if cfg.get("doc_types") and doc_type not in cfg["doc_types"]:
            return []
        units = self._evd_units(doc)
        issues: List[Issue] = []
        allow = self.evidence_allow
        for idx, (para, s, e, piece, heading) in enumerate(units):
            hit = None
            for tier in ("strong", "medium"):
                for m in cfg[tier].finditer(piece):
                    if self._evd_hit_ok(piece, m, cfg, allow):
                        hit = (tier, m)
                        break
                if hit:
                    break
            if not hit:
                continue
            tier, m = hit
            hi = idx + (5 if heading else 3)
            window = units[max(0, idx - 1):hi]
            wtext = "".join(w[3] for w in window)
            if ELM_MARKER_RE.search(wtext):
                continue  # 【待补：数据】等占位：作者已注意到
            plain = self.EVD_DATE_RE.sub(" ", wtext)
            has_num = bool(self.EVD_NUM_RE.search(plain))
            has_cmp = bool(self.EVD_COMPARE_RE.search(wtext))
            has_case = bool(self.EVD_CASE_RE.search(wtext))
            start, end = s + m.start(), s + m.end()
            if tier == "strong":
                if has_num or has_cmp:
                    continue
                if has_case:
                    issue = self._proc_issue(
                        "EVD-003", para, start, end,
                        context=clip(piece.strip(), 60))
                else:
                    issue = self._proc_issue("EVD-001", para, start, end, context=clip(piece.strip(), 60))
            else:
                if has_num or has_cmp or has_case or self.EVD_ANYNUM_RE.search(wtext):
                    continue
                issue = self._proc_issue("EVD-002", para, start, end, context=clip(piece.strip(), 60))
            if issue:
                issues.append(issue)
        return issues

    def _evd_hit_ok(self, piece: str, m: "re.Match[str]", cfg: Dict[str, Any], allow: Set[str]) -> bool:
        """命中是否算一次“评价”：排除白名单短语、用户白名单、引号内、弱化限定、要求和目标语境。"""
        term = m.group(0)
        if any(a and (a in term or term in a) for a in allow):
            return False
        for wl in cfg["whitelist"]:
            for w in wl.finditer(piece):
                if w.start() < m.end() and m.start() < w.end():
                    return False
        before = piece[:m.start()]
        if before.count("“") + before.count("‘") + before.count("「") > before.count("”") + before.count("’") + before.count("」"):
            return False  # 引述他人的评价，不是本文的断言
        if cfg["weak"].search(before[-6:]):
            return False  # “一定程度上明显”“初步见效”：已作限定
        clause_start = max(before.rfind(c) for c in "，、；：") + 1
        clause = piece[clause_start:m.start()]
        prev_clause = ""
        if clause_start > 0:
            earlier = piece[:clause_start - 1]
            prev_clause = earlier[max(earlier.rfind(c) for c in "，、；：") + 1:]
        if self.EVD_DIRECTIVE_STRICT.search(clause) or self.EVD_DIRECTIVE_STRICT.search(prev_clause):
            return False
        if self.EVD_DIRECTIVE_LOOSE.search(clause) or self.EVD_DIRECTIVE_PREV.search(prev_clause):
            return False
        return True

    # -- 隐私与保密（函数） ---------------------------------------------------
    ID18_RE = re.compile(
        r"(?<![0-9A-Za-z])([1-9]\d{5})((?:18|19|20)\d{2})(0[1-9]|1[0-2])(0[1-9]|[12]\d|3[01])(\d{3})([\dXx])(?![0-9A-Za-z])"
    )
    MOBILE_RE = re.compile(r"(?<![\d])(?:\+?86[-\s]?)?(1[3-9]\d)[-\s]?(\d{4})[-\s]?(\d{4})(?![\d])")
    BANK_RE = re.compile(r"(?<![\d])(\d{4}(?:[\s-]?\d{4}){3}(?:[\s-]?\d{1,3})?)(?![\d])")
    LABELED_RE = re.compile(
        r"(病案号|住院号|门诊号|就诊卡号|医保卡号|社保卡号|学号|学籍号|考生号|准考证号|护照号|工号|员工号|职工号)"
        r"\s*(?:为|是)?\s*[：:]?\s*([A-Za-z0-9][A-Za-z0-9\-]{3,})(?![A-Za-z0-9*＊])"
    )
    ADDR_RE = re.compile(
        r"(家庭住址|家庭地址|现住址|住址|户籍地址|居住地址|现居住地|户籍所在地)\s*(?:为|是|在)?\s*[：:]?\s*([^\s，。；,;）)]{4,40})"
    )
    NAME_GENDER_RE = re.compile(r"([一-鿿]{2,6})\s*[，,（(、]\s*(男|女)\s*[，,、）)]?\s*(\d{1,3})\s*(?:岁|周岁)")
    BED_NAME_RE = re.compile(r"(\d{1,3})\s*(?:床|号床)(?![位单旁边头尾日次垫上下])\s*[：:]?\s*([一-鿿]{2,3})")
    CONTACT_HINT_RE = re.compile(r"联系|电话|手机|热线|咨询|值班|拨打|致电|手机号|联系方式|Tel|TEL|tel")
    PERSONAL_BANK_HINT = re.compile(r"银行卡|卡号|储蓄卡|信用卡|借记卡|社保卡")
    ACCOUNT_HINT = re.compile(r"账号|账户|帐号|帐户|开户")

    @staticmethod
    def _id_checksum_ok(id17: str, check: str) -> bool:
        weights = (7, 9, 10, 5, 8, 4, 2, 1, 6, 3, 7, 9, 10, 5, 8, 4, 2)
        total = sum(int(d) * w for d, w in zip(id17, weights))
        return "10X98765432"[total % 11] == check.upper()

    @staticmethod
    def _luhn_ok(digits: str) -> bool:
        total = 0
        for i, ch in enumerate(reversed(digits)):
            d = int(ch)
            if i % 2 == 1:
                d *= 2
                if d > 9:
                    d -= 9
            total += d
        return total % 10 == 0

    @staticmethod
    def _label_context(text: str, start: int) -> str:
        """隐私问题的 context：只取号码前的标签文字，不含号码本身。"""
        return text[max(0, start - 12):start].strip()

    @staticmethod
    def _looks_like_name(name: str) -> bool:
        if not name or any(ch in MASK_CHARS for ch in name):
            return False
        if name in ROLE_WORDS or any(name.endswith(w) for w in ROLE_WORDS):
            return False
        return name[0] in SURNAMES

    def _check_privacy(self, doc: Document, doc_type: str) -> List[Issue]:
        issues: List[Issue] = []
        for para in doc.paras:
            text = para.text
            taken: List[Tuple[int, int]] = []

            def overlaps(s: int, e: int) -> bool:
                return any(a < e and s < b for a, b in taken)

            def add(rid: str, s: int, e: int, masked: str, **kw: Any) -> None:
                issue = self._proc_issue(rid, para, s, e, excerpt=masked,
                                         context=self._label_context(text, s), **kw)
                if issue:
                    issues.append(issue)
                    taken.append((s, e))

            # 身份证号
            for m in self.ID18_RE.finditer(text):
                full = m.group(0)
                try:
                    date(int(m.group(2)), int(m.group(3)), int(m.group(4)))
                except ValueError:
                    continue
                masked = full[:6] + "********" + full[-4:]
                if self._id_checksum_ok(full[:17], full[17]):
                    add("PRI-001", m.start(), m.end(), masked)
                else:
                    add("PRI-001", m.start(), m.end(), masked, level="P1",
                        message="疑似身份证号（校验位不符，可能是录入错误）")
            # 带标签的编号
            for m in self.LABELED_RE.finditer(text):
                s, e = m.span(2)
                if overlaps(s, e):
                    continue
                value = m.group(2)
                masked = value[:2] + "****" + value[-2:] if len(value) > 4 else "****"
                level = "P1" if m.group(1) in ("工号", "员工号", "职工号") else None
                add("PRI-004", s, e, f"{m.group(1)}：{masked}", level=level,
                    message=f"出现{m.group(1)}，属于可识别个人的编号")
            # 银行卡号
            for m in self.BANK_RE.finditer(text):
                s, e = m.span(1)
                if overlaps(s, e):
                    continue
                digits = re.sub(r"\D", "", m.group(1))
                if not 16 <= len(digits) <= 19:
                    continue
                before = text[max(0, s - 15):s]
                masked = digits[:4] + " **** **** " + digits[-4:]
                if self.PERSONAL_BANK_HINT.search(before):
                    add("PRI-003", s, e, masked)
                elif self.ACCOUNT_HINT.search(before):
                    add("PRI-003", s, e, masked, level="P2",
                        message="出现银行账号，如为个人账户须脱敏；单位对公账户可保留",
                        suggestion="确认是否为单位对公账户；个人账户删除或脱敏")
                elif self._luhn_ok(digits) and digits[:1] in ("4", "5", "6") and (digits[0] != "6" or digits[:2] == "62"):
                    add("PRI-003", s, e, masked, level="P1", message="疑似银行卡号（无标签，Luhn 校验通过）")
            # 家庭住址
            for m in self.ADDR_RE.finditer(text):
                s, e = m.span(2)
                if overlaps(s, e):
                    continue
                value = m.group(2)
                if not re.search(r"\d|号|室|幢|单元|栋|楼|村|组|弄|巷", value):
                    continue
                add("PRI-008", s, e, f"{m.group(1)}：{value[:4]}****")
            # 姓名 + 性别 + 年龄
            for m in self.NAME_GENDER_RE.finditer(text):
                raw = m.group(1)
                for role in ROLE_WORDS:
                    if role in raw:
                        raw = raw.split(role)[-1]
                candidates = [raw[-3:], raw[-2:]] if len(raw) >= 3 else [raw]
                name = next((c for c in candidates if self._looks_like_name(c)), "")
                if not name or any(ch in MASK_CHARS for ch in raw):
                    continue
                s = m.start(1) + m.group(1).rfind(name)
                e = s + len(name)
                if overlaps(s, e):
                    continue
                add("PRI-009", s, e, name[0] + "*" * (len(name) - 1) + f"，{m.group(2)}，{m.group(3)}岁")
            # 床号 + 姓名
            for m in self.BED_NAME_RE.finditer(text):
                name = m.group(2)
                if not self._looks_like_name(name):
                    continue
                s, e = m.span(2)
                if overlaps(s, e):
                    continue
                add("PRI-010", s, e, f"{m.group(1)}床{name[0]}{'*' * (len(name) - 1)}")
            # 手机号
            for m in self.MOBILE_RE.finditer(text):
                s, e = m.span()
                if overlaps(s, e):
                    continue
                if re.match(r"\s*[元万亿千百吨件人次个户]", text[e:e + 2]):
                    continue
                digits = re.sub(r"\D", "", m.group(0))[-11:]
                masked = digits[:3] + "****" + digits[-4:]
                window = text[max(0, s - 20):s] + text[e:e + 8]
                with_identity = any(i.para_index == para.index and i.id in ("PRI-001", "PRI-004", "PRI-008", "PRI-009")
                                    for i in issues)
                if with_identity:
                    add("PRI-002", s, e, masked, message="出现个人手机号（与姓名、身份证号、病案号或住址同段）")
                elif self.CONTACT_HINT_RE.search(window):
                    add("PRI-002", s, e, masked, level="P2",
                        message="出现手机号码（工作联系方式）。对外张贴、公开发布的，确认本人同意或改用工作电话",
                        suggestion="对外版本改用座机或工作电话；内部报送可保留")
                else:
                    add("PRI-002", s, e, masked)
        return issues

    # -- 本单位禁用词（--banned） ---------------------------------------------
    @staticmethod
    def _scope_applies(scope: str, doc: Document, doc_type: str) -> bool:
        """禁用词表“适用范围”一栏：空或“全部”都查；否则按文书类型、面向群众、对外发布判断。"""
        tokens = [t.strip() for t in re.split(r"[、，,/／;；\s]+", scope or "") if t.strip()]
        if not tokens or any(t in ("全部", "全文", "所有", "通用", "不限") for t in tokens):
            return True
        for token in tokens:
            if "群众" in token or "居民" in token:
                if doc.public_facing:
                    return True
            elif "对外" in token or "公开" in token:
                if doc.public_facing or doc_type in ("新闻稿", "公众号推文", "信息", "简报", "公告", "通告"):
                    return True
            elif token in doc_type or (len(doc_type) >= 2 and doc_type in token):
                return True
        return False

    def _check_banned(self, doc: Document, doc_type: str, terms: Sequence[BannedTerm]) -> List[Issue]:
        """BAN-001 禁用表述（P1）；BAN-002 慎用表述（P2）。长的词条优先，同一处只报一次。"""
        issues: List[Issue] = []
        active = [t for t in terms if self._scope_applies(t.scope, doc, doc_type)]
        active.sort(key=lambda t: len(t.term), reverse=True)
        for para in doc.paras:
            marker_spans = [mk.span() for mk in MARKER_RE.finditer(para.text)]
            taken: List[Tuple[int, int]] = []
            for term in active:
                for m in term.regex.finditer(para.text):
                    s, e = m.span()
                    if s == e or any(a < e and s < b for a, b in taken + marker_spans):
                        continue
                    taken.append((s, e))
                    level_word = "禁用" if term.banned else "慎用"
                    message = f"使用了本单位{level_word}表述“{term.term}”"
                    if term.reason:
                        message += f"，原因：{term.reason}"
                    suggestion = term.suggestion or ("按本单位禁用词表修改" if term.banned else "由用户决定是否保留")
                    if term.banned:
                        issue = self._proc_issue("BAN-001", para, s, e, message=message, suggestion=suggestion)
                    else:
                        issue = self._proc_issue("BAN-002", para, s, e, message=message, suggestion=suggestion)
                    if issue:
                        if term.source:
                            issue.basis = f"本单位禁用词表（{term.source}）"
                        issues.append(issue)
        return issues

    # -- 篇幅（函数） ---------------------------------------------------------
    def _check_length(self, doc: Document, doc_type: str) -> List[Issue]:
        n = doc.body_char_count()
        anchor = doc.paras[doc.title[0]] if doc.title else (doc.paras[0] if doc.paras else None)
        if anchor is None:
            return []
        excerpt = clip(anchor.text, 40)
        if doc_type in LENGTH_LIMIT_TYPES:
            if n > 5000:
                msg = (f"正文约 {n} 字，超过 5000 字。除部署综合性工作外，地方和部门文件一般不超过 5000 字；"
                       "部署综合性工作的未设上限，仍须坚持短实新")
            elif n > 4000:
                msg = f"正文约 {n} 字，超过 4000 字。部署专项工作或者具体任务的文件一般不超过 4000 字"
            else:
                return []
            issue = self._proc_issue("LEN-001", anchor, 0, len(anchor.text), excerpt=excerpt, message=msg)
            return [issue] if issue else []
        if doc_type == "讲话稿" and n > 12000:
            issue = self._proc_issue("LEN-002", anchor, 0, len(anchor.text), excerpt=excerpt,
                                     message=f"讲话稿约 {n} 字，按每分钟约 200 字估算超过 1 小时")
            return [issue] if issue else []
        return []


# ---------------------------------------------------------------------------
# 对外函数
# ---------------------------------------------------------------------------

def parse_checks(value: Optional[str]) -> List[str]:
    """解析 --checks，返回按标准顺序排列的类别列表；非法值抛出 InputError。"""
    if not value or value.strip().lower() == "all":
        return list(CHECK_ORDER)
    wanted = [v.strip().lower() for v in re.split(r"[,，\s]+", value) if v.strip()]
    if "all" in wanted:
        return list(CHECK_ORDER)
    bad = [v for v in wanted if v not in CHECK_NAMES]
    if bad:
        raise InputError(f"未知的检查类别：{'、'.join(bad)}。可选：all、{'、'.join(CHECK_ORDER)}")
    return [c for c in CHECK_ORDER if c in wanted]


# ---------------------------------------------------------------------------
# 本单位禁用词表（--banned）
# 格式见 shared/profiles/profile-template.md 第三部分：表格第一列（或表头为“词语或表述”的一列）取词，
# “级别”一列为“禁用”的报 P1、“慎用”的报 P2；没有表格时按列表项取词，级别看列表项中的“（慎用）”或所在小标题。
# ---------------------------------------------------------------------------
BANNED_SKIP_SECTION_RE = re.compile(r"维护记录|更新记录|变更记录|修改记录|填写说明|使用说明|备注|附注|^说明$")
BANNED_LIST_ITEM_RE = re.compile(r"^\s*(?:[-*+•·]|\d{1,3}[.、)）]|[（(]\d{1,3}[）)])\s*(.+?)\s*$")
BANNED_SEPARATOR_RE = re.compile(r"^\s*:?-{2,}:?\s*$")
BANNED_WILDCARDS = ("……", "…", "×××", "XXX")


def _banned_level(text: str) -> Optional[bool]:
    """“慎用”返回 False，“禁用”“不用”“停用”返回 True，看不出返回 None。"""
    if "慎用" in text or "少用" in text:
        return False
    if "禁用" in text or "不用" in text or "停用" in text or "不得使用" in text:
        return True
    return None


def _clean_term(text: str) -> str:
    text = re.sub(r"[`*]+", "", text).strip()
    return text.strip("“”\"'‘’「」『』 　")


def _term_regex(term: str) -> Pattern[str]:
    """词条转正则：词中空白可有可无，“……”“×××”作通配。"""
    parts: List[str] = []
    for token in re.split(r"(……|…|×××|XXX|\s+)", term):
        if not token:
            continue
        if token in BANNED_WILDCARDS:
            parts.append(r"[^。；！？\n]{0,20}?")
        elif token.isspace():
            parts.append(r"\s*")
        else:
            parts.append(re.escape(token))
    return re.compile("".join(parts))


def _split_terms(cell: str) -> List[str]:
    terms = []
    for piece in re.split(r"[、/／;；]", cell):
        term = _clean_term(piece)
        if term and len(term) <= 30 and not re.fullmatch(r"[-—–\s]*|【待补[^】]*】|无", term):
            terms.append(term)
    return terms


def load_banned(path: Path) -> List[BannedTerm]:
    """读取禁用词文件；文件不存在或无法读取时抛出 InputError（退出码 2）。"""
    real = long_path(path)
    if not real.is_file():
        raise InputError(f"禁用词文件不存在：{path}（可用 gongwen-profile 建立 .gongwen/禁用词.md）")
    try:
        raw = real.read_bytes()
    except OSError as exc:
        raise InputError(f"无法读取禁用词文件：{exc}") from exc
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            text = raw.decode("gb18030")
        except UnicodeDecodeError as exc:
            raise InputError("无法识别禁用词文件的编码，请另存为 UTF-8 后重试。") from exc
    lines = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    source = path.name
    table_terms: List[BannedTerm] = []
    list_terms: List[BannedTerm] = []
    section_level: Optional[bool] = None
    skip_section = False
    cols: Optional[Dict[str, Optional[int]]] = None

    def is_row(s: str) -> bool:
        return s.startswith("|") and s.count("|") >= 2

    def cells_of(s: str) -> List[str]:
        return [c.strip() for c in s.strip().strip("|").split("|")]

    for idx, line in enumerate(lines):
        s = line.strip()
        if s.startswith("```"):
            continue
        heading = re.match(r"^#{1,6}\s*(.*)$", s)
        if heading:
            title = heading.group(1).strip()
            skip_section = bool(BANNED_SKIP_SECTION_RE.search(title))
            level = _banned_level(title)
            if level is not None:
                section_level = level
            cols = None
            continue
        if skip_section or not s:
            if not s:
                cols = None
            continue
        if is_row(s):
            cells = cells_of(s)
            if all(BANNED_SEPARATOR_RE.match(c) for c in cells if c) or not any(cells):
                continue
            nxt = lines[idx + 1].strip() if idx + 1 < len(lines) else ""
            if cols is None and is_row(nxt) and all(BANNED_SEPARATOR_RE.match(c) for c in cells_of(nxt) if c):
                # 表头：按列名定位各列，找不到词语列时取第一列
                def find(pattern: str) -> Optional[int]:
                    return next((i for i, c in enumerate(cells) if re.search(pattern, c)), None)
                term_col = find(r"词|表述|用语|提法")
                cols = {"term": term_col if term_col is not None else 0, "level": find(r"级别|等级|类别|类型"),
                        "reason": find(r"原因|理由|来源|依据"), "suggestion": find(r"替换|建议|改为|替代"),
                        "scope": find(r"适用|范围")}
                continue
            spec = cols or {"term": 0, "level": None, "reason": None, "suggestion": None, "scope": None}

            def cell(key: str) -> str:
                i = spec.get(key)
                return cells[i] if i is not None and i < len(cells) else ""

            level = _banned_level(cell("level"))
            banned = level if level is not None else (section_level if section_level is not None else True)
            for term in _split_terms(cell("term")):
                table_terms.append(BannedTerm(term, _term_regex(term), banned, cell("reason"),
                                              cell("suggestion"), cell("scope"), source))
            continue
        item = BANNED_LIST_ITEM_RE.match(line)
        if item:
            parts = re.split(r"[：:（(—|，,；;]", item.group(1), maxsplit=1)
            head, tail = parts[0], (parts[1] if len(parts) > 1 else "")
            level = _banned_level(tail)
            banned = level if level is not None else (section_level if section_level is not None else True)
            reason = re.sub(r"[（(【]?(?:禁用|慎用)[）)】]?", "", tail).strip("）)】 　：:，,；;").strip()
            for term in _split_terms(head):
                list_terms.append(BannedTerm(term, _term_regex(term), banned, reason, "", "", source))
    return table_terms if table_terms else list_terms


def resolve_doc_type(doc: Document, given: Optional[str]) -> Tuple[str, str]:
    if given and given.strip():
        return given.strip(), "参数"
    inferred = doc.infer_doc_type()
    if inferred:
        return inferred, "稿首表头推断" if doc.has_qianbao_header() else "标题推断"
    return "其他", "默认"


def parse_allow(value: Optional[Any]) -> List[str]:
    """解析 --allow：逗号或空格分隔的规则 ID，统一为大写。"""
    if not value:
        return []
    items = value if isinstance(value, (list, tuple, set)) else re.split(r"[,，\s]+", str(value))
    return [str(v).strip().upper() for v in items if str(v).strip()]


def parse_words(value: Optional[Any]) -> List[str]:
    """解析 --evidence-allow：逗号、顿号或空格分隔的评价词。"""
    if not value:
        return []
    items = value if isinstance(value, (list, tuple, set)) else re.split(r"[,，、\s]+", str(value))
    return [str(v).strip() for v in items if str(v).strip()]


def _apply_allow(issues: List[Issue], allow: Sequence[str], linter: Linter, notes: List[str]) -> List[Issue]:
    """--allow：本单位制度明文规定的写法，命中时降为 P2 并注明，不计入 --strict。"""
    known = ({r.id for r in linter.rules} | set(linter.proc_meta) | {p.id for p in linter.patterns})
    unknown = [a for a in allow if a not in known]
    if unknown:
        notes.append(f"--allow 中的规则 {'、'.join(unknown)} 不存在，已忽略。")
    wanted = set(allow) & known
    changed: List[str] = []
    for issue in issues:
        if issue.id in wanted:
            issue.level = "P2"
            if not issue.message.startswith("按本单位制度保留："):
                issue.message = "按本单位制度保留：" + issue.message
            changed.append(issue.id)
    if changed:
        notes.append(f"按 --allow 降为 P2 的问题 {len(changed)} 处（{'、'.join(sorted(set(changed)))}），"
                     "依据本单位制度保留，不计入 --strict。")
        issues.sort(key=lambda x: (LEVEL_RANK.get(x.level, 9), x.line, x.start, x.id))
    return issues


# ---------------------------------------------------------------------------
# 基线参照（--baseline）：与公开公文语料基线 shared/style/baseline.md 比较，只提示不判错
# ---------------------------------------------------------------------------
BASELINE_CANDIDATES = (
    ("references", "_shared", "style", "baseline.json"),   # skill 内的同步副本
    ("style", "baseline.json"),                            # 仓库 shared/scripts/ 旁的 shared/style/
)
BASELINE_YISHI_RE = re.compile(r"(?:^|(?<=[，。；：！？、）」”\s]))一是")
# 机检文书类型 → 基线文种；基线中“工作总结”是政府信息公开年报、“报告”是国务院法定报告，不能代表一般稿件，不比较
BASELINE_TYPE_MAP = {
    "通知": "通知", "意见": "意见", "批复": "批复", "函": "函", "通报": "通报",
    "实施方案": "方案", "贯彻落实方案": "方案", "整改方案": "方案", "讲话稿": "讲话",
}
BASELINE_ID_META = (
    ("BSL-001", "length", "篇幅"),
    ("BSL-002", "structure", "一级标题数"),
    ("BSL-003", "style", "“一是”组数"),
    ("BSL-004", "style", "顿号密度"),
)


@dataclass
class BaselineRef:
    """--baseline 读取的结果：data 为 None 表示没有可用基线，note 说明原因或来源。"""
    data: Optional[Dict[str, Any]]
    note: str = ""


def load_baseline(explicit: Optional[str] = None) -> BaselineRef:
    """explicit 为 None 或 "auto" 时按脚本相对路径查找；找不到、读不了都不报错，只给提示。"""
    paths: List[Path] = []
    if explicit and explicit != "auto":
        paths.append(Path(explicit))
    else:
        paths.extend(SCRIPT_DIR.parent.joinpath(*parts) for parts in BASELINE_CANDIDATES)
    for path in paths:
        if not path.is_file():
            continue
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, ValueError) as exc:
            return BaselineRef(None, f"基线文件 {path} 读取失败（{exc}），已跳过基线参照。")
        if not isinstance(data, dict) or not isinstance(data.get("groups"), dict):
            return BaselineRef(None, f"基线文件 {path} 格式不符，已跳过基线参照。")
        return BaselineRef(data, f"基线参照来自 {path.name}（语料 {data.get('corpus', {}).get('docs', '?')} 篇）。")
    return BaselineRef(None, "未找到基线文件（references/_shared/style/baseline.json 或 shared/style/baseline.json），已跳过基线参照。")


def _baseline_group(doc: Document, doc_type: str) -> Tuple[Optional[str], str]:
    """返回 (基线文种, 跳过原因)。"""
    title = doc.paras[doc.title[0]].text if doc.title else (doc.paras[0].text if doc.paras else "")
    if doc_type == "通知":
        return ("通知（印发转发）" if re.search(r"(?:印发|转发|批转)", title) else "通知"), ""
    if doc_type in REPORT_TYPES and "政府工作报告" in title:
        return "政府工作报告", ""
    if doc_type in BASELINE_TYPE_MAP:
        return BASELINE_TYPE_MAP[doc_type], ""
    if doc_type == "总结":
        return None, "基线里的工作总结是政府信息公开年报，格式化程度高，不能代表一般总结，未比较。"
    if doc_type in REPORT_TYPES:
        return None, "基线里的报告是国务院法定报告（3 个机关），不能代表一般工作报告，未比较。"
    return None, f"基线未覆盖“{doc_type}”，未比较。"


def baseline_hints(doc: Document, doc_type: str, ref: BaselineRef) -> Tuple[List[Issue], List[str], Optional[Dict[str, Any]]]:
    """篇幅、一级标题数、“一是”组数、顿号密度偏离基线 P10～P90 区间时给 P2 提示。"""
    if ref.data is None:
        return [], [ref.note] if ref.note else [], None
    notes = [ref.note] if ref.note else []
    group_name, why = _baseline_group(doc, doc_type)
    if group_name is None:
        return [], notes + [why], None
    group = ref.data["groups"].get(group_name)
    if not group:
        return [], notes + [f"基线里没有“{group_name}”这一组，未比较。"], None
    texts = [p.text for p in doc.paras]
    chars = doc.char_count
    values = {
        "BSL-001": float(chars),
        "BSL-002": float(sum(1 for t in texts if (heading_of(t) or (0, 0))[0] == 1)),
        "BSL-003": float(sum(len(BASELINE_YISHI_RE.findall(t)) for t in texts)),
        "BSL-004": sum(t.count("、") for t in texts) * 1000.0 / chars if chars else 0.0,
    }
    stats = {
        "BSL-001": group.get("chars") or {},
        "BSL-002": (group.get("headings") or {}).get("h1_per_doc") or {},
        "BSL-003": (group.get("yishi") or {}).get("sets_per_doc") or {},
        "BSL-004": group.get("dunhao_per_1000") or {},
    }
    units = {"BSL-001": "字", "BSL-002": "个", "BSL-003": "组", "BSL-004": "个/千字"}
    # 太短的稿件，“一是”组数和顿号密度没有比较意义
    skip = {"BSL-003", "BSL-004"} if chars < 300 else set()
    low_sample = bool(group.get("low_sample") or group.get("low_orgs"))
    issues: List[Issue] = []
    summary: Dict[str, Any] = {}
    n = group.get("n", "?")
    for rid, check, label in BASELINE_ID_META:
        st = stats[rid]
        if rid in skip or not st.get("n") or "p90" not in st:
            continue
        low, high, val = st.get("p10"), st["p90"], values[rid]
        fmt = (lambda x: f"{x:.1f}") if rid == "BSL-004" else (lambda x: f"{x:g}")
        summary[label] = {"value": round(val, 1), "p10": low, "median": st.get("median"), "p90": high}
        side = ""
        if val > high:
            side = "高于"
        elif low is not None and val < low:
            side = "低于"
        if not side:
            continue
        bound = high if side == "高于" else low
        edge = "P90" if side == "高于" else "P10"
        issues.append(Issue(
            id=rid, level="P2", check=check, category="基线参照", gate="文风关", line=0,
            excerpt=f"{label} {fmt(val)}{units[rid]}",
            message=(f"{label} {fmt(val)}{units[rid]}，{side}同文种（{group_name}，{n} 篇）{edge} {fmt(bound)}{units[rid]}"
                     f"（中位数 {fmt(st.get('median', 0))}，P10～P90 为 {fmt(low if low is not None else 0)}～{fmt(high)}）"),
            suggestion="仅作参照：篇幅、层次以上级要求、既有模板和事项复杂程度为准，用户指定的篇幅以用户为准，有理由就不改",
            basis="公开公文语料基线 references/_shared/style/baseline.md（真实公文的分布，不是规定）"
                  + ("；该文种样本偏少，只作方向参考" if low_sample else ""),
        ))
    info = {"group": group_name, "n": n, "metrics": summary}
    if not issues:
        notes.append(f"基线参照：篇幅、一级标题数、“一是”组数、顿号密度均在同文种（{group_name}）P10～P90 内。")
    return issues, notes, info


def lint_document(doc: Document, doc_type: Optional[str] = None, checks: Optional[str] = None,
                  linter: Optional[Linter] = None, notes: Optional[List[str]] = None,
                  banned: Optional[Sequence[BannedTerm]] = None,
                  allow: Optional[Sequence[str]] = None,
                  evidence_allow: Optional[Any] = None,
                  baseline: Optional[BaselineRef] = None) -> Dict[str, Any]:
    """banned 为 None 表示没有指定禁用词表；allow 为降级规则 ID 清单（见 --allow）；
    evidence_allow 为评价词白名单（见 --evidence-allow，EVD 类规则不对其中的词报告）。"""
    linter = linter or Linter()
    previous_allow = linter.evidence_allow
    linter.evidence_allow = set(parse_words(evidence_allow))
    try:
        return _lint_document(doc, doc_type, checks, linter, notes, banned, allow, baseline)
    finally:
        linter.evidence_allow = previous_allow


def _lint_document(doc: Document, doc_type: Optional[str], checks: Optional[str], linter: Linter,
                   notes: Optional[List[str]], banned: Optional[Sequence[BannedTerm]],
                   allow: Optional[Sequence[str]], baseline: Optional[BaselineRef] = None) -> Dict[str, Any]:
    check_list = parse_checks(checks)
    dtype, source = resolve_doc_type(doc, doc_type)
    all_notes = list(notes or [])
    if dtype not in DOC_TYPES:
        all_notes.append(f"文书类型“{dtype}”不在统一词表中，按文种限定的规则不会生效。")
    issues = linter.lint(doc, dtype, set(check_list), banned)
    if banned is not None:
        if "style" not in check_list:
            all_notes.append("未检查 style 类，禁用词表没有比对。")
        elif not banned:
            all_notes.append("禁用词表中没有识别出词条（格式见单位画像模板第三部分：表格第一列为词语，“级别”列填禁用或慎用）。")
        else:
            n_ban = sum(1 for t in banned if t.banned)
            skipped = sum(1 for t in banned if not linter._scope_applies(t.scope, doc, dtype))
            msg = f"已比对本单位禁用词表 {len(banned)} 条（禁用 {n_ban} 条、慎用 {len(banned) - n_ban} 条）"
            if skipped:
                msg += f"，其中 {skipped} 条的适用范围不含“{dtype}”，未比对"
            all_notes.append(msg + "。")
    baseline_info: Optional[Dict[str, Any]] = None
    if baseline is not None:
        bsl_issues, bsl_notes, baseline_info = baseline_hints(doc, dtype, baseline)
        if bsl_issues:
            issues = sorted(issues + bsl_issues, key=lambda x: (LEVEL_RANK.get(x.level, 9), x.line, x.start, x.id))
        all_notes.extend(bsl_notes)
    if allow:
        issues = _apply_allow(issues, parse_allow(allow), linter, all_notes)
    summary: Dict[str, Any] = {lv: sum(1 for i in issues if i.level == lv) for lv in LEVELS}
    # 待补、待核实标记单独计数（开发规范第 8 节），不论是否检查 marker 类
    summary["markers"] = {
        "待补": sum(len(TODO_MARK_RE.findall(p.text)) for p in doc.paras),
        "待核实": sum(len(VERIFY_MARK_RE.findall(p.text)) for p in doc.paras),
    }
    if doc.truncated_line is not None:
        all_notes.append(f"第 {doc.truncated_line} {doc.unit}为【起草说明】，其后内容未检查。")
    if "marker" in check_list:
        todo = sum(1 for i in issues if i.id == "MRK-001")
        verify = sum(1 for i in issues if i.id == "MRK-002")
        bad = sum(1 for i in issues if i.id == "MRK-003")
        if todo or verify or bad:
            all_notes.append(
                f"正文含【待补】{todo} 处、【待核实】{verify} 处"
                + (f"、格式不规范的标记 {bad} 处" if bad else "")
                + "；【待补】定稿前补齐，【待核实】建议移至起草说明的待核实清单。")
    body_chars = doc.body_char_count()
    if "length" in check_list:
        all_notes.append(f"篇幅口径：标题之后、【起草说明】之前的正文，不含空白和版记，约 {body_chars} 字。")
    all_notes.extend(linter.warnings)
    result: Dict[str, Any] = {
        "file": doc.file,
        "doc_type": dtype,
        "doc_type_source": source,
        "char_count": doc.char_count,
        "body_char_count": body_chars,
        "summary": summary,
        "issues": [i.to_dict() for i in issues],
        "checks": check_list,
        "notes": all_notes,
        "engine_version": VERSION,
    }
    if baseline_info is not None:
        result["baseline"] = baseline_info
    return result


def lint_text(text: str, doc_type: Optional[str] = None, checks: Optional[str] = None,
              file_name: str = "<text>", linter: Optional[Linter] = None,
              banned: Optional[Sequence[BannedTerm]] = None, allow: Optional[Sequence[str]] = None,
              evidence_allow: Optional[Any] = None, baseline: Optional[BaselineRef] = None) -> Dict[str, Any]:
    """检查一段文本（供其他脚本和测试调用）。"""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    lines = [(i, line, False) for i, line in enumerate(text.split("\n"), start=1)]
    return lint_document(Document(file_name, lines), doc_type, checks, linter, banned=banned, allow=allow,
                         evidence_allow=evidence_allow, baseline=baseline)


def lint_file(path: Path, doc_type: Optional[str] = None, checks: Optional[str] = None,
              linter: Optional[Linter] = None, banned: Optional[Sequence[BannedTerm]] = None,
              allow: Optional[Sequence[str]] = None, evidence_allow: Optional[Any] = None,
              baseline: Optional[BaselineRef] = None) -> Dict[str, Any]:
    real = long_path(path)  # Windows 下超过 260 字符的路径
    if not real.is_file():
        raise InputError(f"文件不存在：{path}")
    lines, notes = read_lines(real)
    unit = "段" if path.suffix.lower() == ".docx" else "行"
    doc = Document(path.name, lines, unit)
    doc.docx_tables = list(getattr(lines, "tables", []))
    return lint_document(doc, doc_type, checks, linter, notes, banned, allow, evidence_allow, baseline)


def render_text(result: Dict[str, Any]) -> str:
    unit = "段" if str(result["file"]).lower().endswith(".docx") else "行"
    s = result["summary"]
    src = "" if result["doc_type_source"] == "参数" else f"，{result['doc_type_source']}"
    out = [
        f"机检结果：{result['file']}（文书类型：{result['doc_type']}{src}；正文 {result['char_count']} 字）",
        f"P0 {s['P0']} 处｜P1 {s['P1']} 处｜P2 {s['P2']} 处",
    ]
    markers = s.get("markers") or {}
    if markers.get("待补") or markers.get("待核实"):
        out.append(f"待补 {markers.get('待补', 0)} 处、待核实 {markers.get('待核实', 0)} 处；待补定稿前补齐，待核实建议移至起草说明")
    if not result["issues"]:
        out.append("")
        out.append("未发现问题。机检只覆盖可规则化的项目，内容、提法、行文规则仍需人工审读。")
    for issue in result["issues"]:
        out.append("")
        out.append(f"[{issue['level']}] 第{issue['line']}{unit}  {issue['id']}  {issue['category']}：{issue['message']}")
        line = f"    原文：{issue['excerpt']}"
        if issue.get("context") and len(issue["excerpt"]) < 8:
            line += f"（上下文：{issue['context']}）"
        out.append(line)
        if issue.get("suggestion"):
            out.append(f"    建议：{issue['suggestion']}")
        if issue.get("basis"):
            out.append(f"    依据：{issue['basis']}")
    if result["notes"]:
        out.append("")
        out.extend(f"提示：{n}" for n in result["notes"])
    return "\n".join(out)


class _ChineseHelpFormatter(argparse.HelpFormatter):
    def add_usage(self, usage, actions, groups, prefix=None):  # type: ignore[override]
        return super().add_usage(usage, actions, groups, prefix="用法：")


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        prog="gongwen_lint.py",
        description="公文机检：残留、标点、数字、格式要素、结构、AI 味、防御性、隐私与保密、法规、减负、待补标记、篇幅。",
        epilog="退出码：0 正常；--strict 且存在 P0 时为 1；输入错误为 2。docx 输入时“行号”为段落序号。",
        add_help=False,
        formatter_class=_ChineseHelpFormatter,
    )
    parser._positionals.title = "位置参数"
    parser._optionals.title = "选项"
    parser.add_argument("file", help="待检查文件（.txt、.md、.docx），遇到【起草说明】行即停止读取")
    parser.add_argument("--doc-type", default=None,
                        help="文书类型，如 通知、请示、报告、总结、讲话稿、对照检查材料、整改方案；"
                             "不填时按标题推断，推断不出按“其他”")
    parser.add_argument("--checks", default="all",
                        help="检查类别，逗号分隔：all（默认）、" + "、".join(CHECK_ORDER))
    parser.add_argument("--json", action="store_true", help="输出 JSON")
    parser.add_argument("--strict", action="store_true", help="存在 P0 时返回 1")
    parser.add_argument("--banned", default=None, metavar="禁用词.md",
                        help="本单位禁用词表（一般为 .gongwen/禁用词.md）：“禁用”命中报 P1（BAN-001），"
                             "“慎用”命中报 P2（BAN-002），归入 style 类；文件不存在时返回 2")
    parser.add_argument("--allow", default=None, metavar="规则ID,规则ID",
                        help="本单位制度明文规定的写法（如企业发文字号用方括号时写 FMT-001），"
                             "列出的规则命中时降为 P2 并注明“按本单位制度保留”，不计入 --strict")
    parser.add_argument("--evidence-allow", default=None, metavar="词,词",
                        help="评价词白名单（EVD 类检查）：列出的评价词（如“显著”“大幅”）不再要求同句或相邻句有数字、对比证据；"
                             "固定短语白名单在 lint_rules.json 的 evidence.whitelist")
    parser.add_argument("--baseline", nargs="?", const="auto", default=None, metavar="baseline.json",
                        help="与公开公文语料基线比较篇幅、一级标题数、“一是”组数、顿号密度，偏离 P10～P90 的报 P2 提示（BSL-001～004），"
                             "只提示不判错；不带路径时按脚本相对路径查找，找不到则跳过并提示")
    parser.add_argument("-h", "--help", action="help", help="显示帮助并退出")
    return parser


def main(argv: Optional[Sequence[str]] = None) -> int:
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8")  # type: ignore[attr-defined]
        except Exception:
            pass
    parser = build_parser()
    args = parser.parse_args(argv)
    try:
        banned = load_banned(Path(args.banned)) if args.banned else None
        baseline = load_baseline(args.baseline) if args.baseline is not None else None
        result = lint_file(Path(args.file), args.doc_type, args.checks, banned=banned, allow=parse_allow(args.allow),
                           evidence_allow=args.evidence_allow, baseline=baseline)
    except InputError as exc:
        print(f"错误：{exc}", file=sys.stderr)
        return 2
    if args.json:
        print(json.dumps(result, ensure_ascii=False, indent=2))
    else:
        print(render_text(result))
    if args.strict and result["summary"]["P0"] > 0:
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
