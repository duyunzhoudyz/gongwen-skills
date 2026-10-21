#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""检查已有 docx 是否符合 GB/T 9704-2012。

检查项：纸张、页边距、文档网格、正文和标题的字体字号、行距、首行缩进、层次标题字体、页码、
发文字号（上行文居左空一字）、成文日期、主送机关、附件说明、抄送、印发日期、残留标记等；
v1.1 版式：命令（令）、附件另面编排、联合行文版头与联署署名、预印红头纸套打（--overprint）。
--json 输出与 docs/开发规范.md 第 8 节的机检 JSON 一致，级别 P0/P1/P2。

退出码：默认 0；--strict 时存在 P0 返回 1；输入错误或缺少依赖返回 2。
"""
from __future__ import annotations

import argparse
import json
import math
import re
import sys
from pathlib import Path
from typing import Dict, List, Optional, Sequence

sys.path.insert(0, str(Path(__file__).resolve().parent))

import gw_fonts as F  # noqa: E402
from gw_layout import (ATTACH_MARK_RE, DATE_CN_RE, MARKERS_RE, ROLE_LABELS, SECRET_WORDS_RE, UPWARD_TYPES,  # noqa: E402
                       doc_kind, draft_note_index, has_brief_header, heading_split, is_heading, is_pipe_residue,
                       is_table_caption, parse_brief, parse_lines, text_width_chars)

CJK_RE = re.compile(r"[一-鿿]")
# 序号位与 gw_layout.DOCNO_PAT 一致：阿拉伯数字，或尚未编号的“【待补：序号】”
DOCNO_SEQ = r"(?:\d+|【待补[^】]*】)"
DOCNO_BRACKET_RE = re.compile(r"[一-鿿]{1,12}[\[［(（【﹝{｛]\s*(?:19|20)\d{2}\s*[\]］)）】﹞}｝]\s*第?\s*" + DOCNO_SEQ + r"\s*号")
DOCNO_SHORT_YEAR_RE = re.compile(r"〔\d{2}〕")
DOCNO_DI_RE = re.compile(r"〔\d{4}〕\s*第\s*" + DOCNO_SEQ + r"\s*号")
DOCNO_ZERO_RE = re.compile(r"〔\d{4}〕\s*0\d*\s*号")
DATE_ZERO_RE = re.compile(r"\d{4}年0[1-9]月|月0[1-9]日")
PRINT_OK_RE = re.compile(r"\d{4}年[1-9]\d?月[1-9]\d?日印发$")
PAGE_TEXT_RE = re.compile(r"^—\s?[1-9]\d*\s?—$")
ORDINAL_HINT = "依次用 一、（一）1.（1）"
# (写法, 问题说明, 修改建议)。“1、”不判断所在层级：序号后改用“.”，直接位于一级标题下的应改为“（一）”
BAD_ORDINALS = [
    (re.compile(r"^\d+、"), "阿拉伯数字序数后应用“.”，不用“、”",
     "序号后应用“.”，如“1.”；如属第二层（直接位于“一、”下），应为“（一）”"),
    (re.compile(r"^（[一二三四五六七八九十]+）[、，.．]"), "“（一）”后不加标点", ORDINAL_HINT),
    (re.compile(r"^[一二三四五六七八九十]+[.．，]"), "一级序数后用顿号“一、”", ORDINAL_HINT),
    (re.compile(r"^\((\d+|[一二三四五六七八九十]+)\)"), "序数括号应为全角“（）”", ORDINAL_HINT),
]
LAYOUT_NAMES = ("红头", "信函", "纪要", "命令", "白头", "材料", "简报")  # --layout 可选值
LEVEL_FONT = {"h1": "黑体", "h2": "楷体", "h3": "仿宋", "h4": "仿宋", "body": "仿宋", "attendance": "仿宋"}
LEVEL_NAME = {"h1": "一级标题", "h2": "二级标题", "h3": "三级标题", "h4": "四级标题", "body": "正文"}
ALLOW_PREFIX = "按本单位制度保留："
# DOCX-016 使用非标准字体：字体严格模式，只认标准字体或同款别名，缺字体时先安装，不得用其他字体代替
NON_STANDARD_HINT = "改为“{std}”；本机没有该字体时先下载安装，不要用其他字体代替（运行 python scripts/gw_fonts.py 查看获取和安装方法）"
NON_STANDARD_BASIS = "GB/T 9704-2012 7.3.1、7.3.3；gbt9704.md 第九部分（字体名称，实践）"


class Collector:
    def __init__(self) -> None:
        self.issues: List[dict] = []
        self._agg: Dict[str, dict] = {}

    def add(self, iid: str, level: str, category: str, message: str, *, line: Optional[int] = None,
            excerpt: str = "", suggestion: str = "", basis: str = "") -> None:
        self.issues.append({"id": iid, "level": level, "category": category, "line": line,
                            "excerpt": excerpt[:40], "message": message, "suggestion": suggestion, "basis": basis})

    def agg(self, key: str, iid: str, level: str, category: str, message: str, *, line: Optional[int],
            excerpt: str, suggestion: str, basis: str) -> None:
        """同类问题合并为一条，记录段数和第一处位置。"""
        if key in self._agg:
            self._agg[key]["count"] += 1
            return
        item = {"id": iid, "level": level, "category": category, "line": line, "excerpt": excerpt[:40],
                "message": message, "suggestion": suggestion, "basis": basis, "count": 1}
        self._agg[key] = item
        self.issues.append(item)

    def finish(self) -> List[dict]:
        order = {"P0": 0, "P1": 1, "P2": 2}
        for item in self.issues:
            count = item.pop("count", None)
            if count and count > 1:
                item["message"] += f"（共 {count} 段，此处为第一处）"
        return sorted(self.issues, key=lambda x: (order[x["level"]], x["line"] or 0))


def _cls_ok(cls: str, expected: str) -> bool:
    return cls == expected


def _footer_info(footer_el) -> dict:
    """页脚中的页码域、文本、对齐方式和字体。"""
    from docx.oxml.ns import qn
    info = {"has_page": False, "text": "", "align": None, "fonts": set(), "sizes": set()}
    for p in footer_el.iter(qn("w:p")):
        instr = "".join(t.text or "" for t in p.iter(qn("w:instrText")))
        simple = [f.get(qn("w:instr")) or "" for f in p.iter(qn("w:fldSimple"))]
        if "PAGE" in instr.upper() or any("PAGE" in s.upper() for s in simple):
            info["has_page"] = True
            info["text"] = "".join(t.text or "" for t in p.iter(qn("w:t"))).strip()
            jc = p.find(qn("w:pPr") + "/" + qn("w:jc"))
            info["align"] = jc.get(qn("w:val")) if jc is not None else "left"
            for r in p.iter(qn("w:r")):
                rf = r.find(qn("w:rPr") + "/" + qn("w:rFonts"))
                if rf is not None and (rf.get(qn("w:eastAsia")) or rf.get(qn("w:ascii"))):
                    info["fonts"].add(rf.get(qn("w:eastAsia")) or rf.get(qn("w:ascii")))
                sz = r.find(qn("w:rPr") + "/" + qn("w:sz"))
                if sz is not None:
                    info["sizes"].add(int(sz.get(qn("w:val"))) / 2)
            break
    return info


def _part_text(part_el) -> str:
    from docx.oxml.ns import qn
    return "".join(t.text or "" for t in part_el.iter(qn("w:t")))


def parse_allow(value) -> List[str]:
    """解析 --allow：逗号或空格分隔的检查项 ID，统一为大写（与 gongwen_lint.py 的 --allow 一致）。"""
    if not value:
        return []
    items = value if isinstance(value, (list, tuple, set)) else re.split(r"[,，\s]+", str(value))
    return [str(v).strip().upper() for v in items if str(v).strip()]


def check_document(path: Path, layout: Optional[str] = None, params: Optional[dict] = None,
                   allow=None, unit_fonts=None, overprint: bool = False) -> dict:
    """检查 docx，返回机检 JSON 结构。

    allow：本单位制度明文规定、与国标不同的写法对应的检查项（如企业发文字号用方括号的 DOCX-050），
    命中时降为 P2 并注明“按本单位制度保留”，不计入 --strict。
    unit_fonts：按单位规定指定标准字体（“仿宋=方正仿宋_GBK”字符串或 {角色: 字体名}）；指定后该角色使用单位字体视为合规，
    使用国标字体按非单位规定字体报 DOCX-016。角色或字体名不合法时抛 ValueError。
    overprint：文档是预印红头纸套打稿（不含红色版头），按套打要求检查（DOCX-150、151），版式缺省按红头。
    """
    import gw_docx as D
    from docx import Document
    from docx.oxml.ns import qn

    params = params or D.load_params()
    params = F.apply_unit_fonts(params, F.parse_unit_fonts(unit_fonts, params))
    unit_over = params["fonts"].get("unit_override", {})
    doc = Document(str(path))
    col = Collector()
    tol = float(D.pv(params, "page.margin_tolerance_mm"))
    line_pt = float(D.pv(params, "grid.line_spacing_pt"))
    body_pt = float(D.pv(params, "page.body_height_mm")) * 72 / 25.4
    lines_req = int(D.pv(params, "grid.lines_per_page"))
    chars_req = int(D.pv(params, "grid.chars_per_line"))
    std = params["fonts"]["standard"]

    # ---- 页面 ----
    for s_idx, section in enumerate(doc.sections, start=1):
        sect = section._sectPr
        pgsz = sect.find(qn("w:pgSz"))
        w = int(pgsz.get(qn("w:w"))) if pgsz is not None else 12240
        h = int(pgsz.get(qn("w:h"))) if pgsz is not None else 15840
        w_mm, h_mm = D.twips_to_mm(w), D.twips_to_mm(h)
        if abs(w_mm - 210) > tol or abs(h_mm - 297) > tol:
            col.add("DOCX-001", "P0", "页面", f"第 {s_idx} 节纸张为 {w_mm:.0f}×{h_mm:.0f} mm，不是 A4",
                    excerpt=f"{w_mm:.0f}×{h_mm:.0f} mm", suggestion="纸张设为 A4（210 mm×297 mm）", basis="GB/T 9704-2012 5.1")
        pgmar = sect.find(qn("w:pgMar"))
        if pgmar is not None:
            for attr, label, key, lvl, basis in (
                    ("top", "上边距（天头）", "margin_top_mm", "P0", "GB/T 9704-2012 5.2.1"),
                    ("left", "左边距（订口）", "margin_left_mm", "P0", "GB/T 9704-2012 5.2.1"),
                    ("bottom", "下边距", "margin_bottom_mm", "P1", "GB/T 9704-2012 5.2.1（由版心推算）"),
                    ("right", "右边距", "margin_right_mm", "P1", "GB/T 9704-2012 5.2.1（由版心推算）")):
                val = pgmar.get(qn(f"w:{attr}")) or pgmar.get(qn("w:start" if attr == "left" else "w:end"))
                if val is None:
                    continue
                mm = D.twips_to_mm(abs(int(val)))
                want = float(D.pv(params, f"page.{key}"))
                if abs(mm - want) > tol:
                    col.add("DOCX-002" if lvl == "P0" else "DOCX-003", lvl, "页面",
                            f"{label}为 {mm:.1f} mm，应为 {want:g} mm", excerpt=f"{label} {mm:.1f} mm",
                            suggestion=f"{label}设为 {want / 10:g} 厘米", basis=basis)
        grid = sect.find(qn("w:docGrid"))
        gtype = grid.get(qn("w:type")) if grid is not None else None
        if gtype == "linesAndChars":
            pgmar_l = int(pgmar.get(qn("w:left")) or 0) if pgmar is not None else 0
            pgmar_r = int(pgmar.get(qn("w:right")) or 0) if pgmar is not None else 0
            width_pt = (w - pgmar_l - pgmar_r) / 20
            pitch = 16 + int(grid.get(qn("w:charSpace")) or 0) / 4096
            chars = math.floor(width_pt / pitch + 1e-6) if pitch > 0 else 0
            if chars != chars_req:
                col.add("DOCX-004", "P2", "页面", f"文档网格每行约 {chars} 字，标准为每行 {chars_req} 字",
                        excerpt=f"每行 {chars} 字", suggestion="页面设置→文档网格→指定行和字符网格，每行 28 字",
                        basis="GB/T 9704-2012 5.2.3")
        elif s_idx == 1:
            col.add("DOCX-004", "P2", "页面", "未设置字符网格，每行字数取决于字体，可能不是 28 字",
                    excerpt="无字符网格", suggestion="页面设置→文档网格→指定行和字符网格，每行 28 字、每页 22 行",
                    basis="GB/T 9704-2012 5.2.3")
    first_grid = doc.sections[0]._sectPr.find(qn("w:docGrid"))
    grid_pitch = None
    if first_grid is not None and first_grid.get(qn("w:type")) in ("lines", "linesAndChars"):
        grid_pitch = int(first_grid.get(qn("w:linePitch")) or 0) / 20

    # ---- 段落角色 ----
    paras = list(doc.paragraphs)
    texts = [D.direct_text(p) or p.text for p in paras]
    # 【起草说明】之后不是正文，不参与识别和检查（output-format.md 第二部分）
    cut = draft_note_index(texts)
    cut_el = paras[cut]._p if cut is not None else None
    if cut is not None:
        col.add("DOCX-082", "P1", "残留", "文档中有【起草说明】，其后内容不是正文，未检查；成稿中应删除", line=cut + 1,
                excerpt=texts[cut], suggestion="删除【起草说明】及其后的全部内容",
                basis="references/_shared/style/output-format.md 第二部分")
        paras, texts = paras[:cut], texts[:cut]
    # 附件另面编排（7.3.7）：首行只有“附件”“附件1”的段落起是附件页，主文识别到它之前为止，附件页另查
    att_starts = [i for i, t in enumerate(texts) if ATTACH_MARK_RE.match(t.strip())]
    att_paras: List = []
    att_texts: List[str] = []
    att_offset = len(paras)
    if att_starts:
        att_offset = att_starts[0]
        att_paras, att_texts = paras[att_offset:], texts[att_offset:]
        paras, texts = paras[:att_offset], texts[:att_offset]
    # 正文表格（不含版记）：识别时在表格位置插入占位，表格上一行的表题不按正文检查
    tables = [(a, el) for a, el in D.body_tables(doc) if cut is None or a < cut]
    red_flags = [D.para_is_red(p, doc) for p in paras]
    letter_issuer = None
    sec0 = doc.sections[0]
    if sec0.different_first_page_header_footer:
        htxt = _part_text(sec0.first_page_header._element).strip()
        if htxt:
            letter_issuer = sec0.first_page_header
    # 简报：开头有简报名称加期号行或编印行的，报头、报尾要素不当作发文机关标志、正文来查
    brief = layout == "简报" or (layout is None and letter_issuer is None and has_brief_header(texts, red_flags))
    anchors = [a for a, _ in tables]
    parsed = (parse_brief(texts, red=red_flags, tables_after=anchors, force=layout == "简报") if brief
              else parse_lines(texts, red=red_flags, tables_after=anchors))
    kind, doc_type = doc_kind(str(parsed.fields.get("标题") or ""))
    if brief:
        layout, doc_type = "简报", "简报"
    if layout is None:
        roles = parsed.roles()
        if letter_issuer is not None:
            layout = "信函"
        elif "issuer" in roles:
            issuer = parsed.find("issuer")[0].text
            layout = "纪要" if issuer.endswith("纪要") else ("命令" if "lingno" in roles else "红头")
        elif overprint or any(D.table_mark(t._tbl) == D.JOINT_HEAD_MARK for t in doc.tables):
            layout = "红头"  # 套打稿、联合行文版头（表格）里没有可识别的标志段
        elif kind == "法定":
            layout = "白头"
        else:
            layout = "材料"
    if not parsed.find("title") and layout not in ("纪要", "命令"):
        col.add("DOCX-023", "P2", "结构", "未识别到标题，标题相关检查已跳过", excerpt=texts[0] if texts else "",
                suggestion="标题一般以“关于”开头或以文种结尾", basis="GB/T 9704-2012 7.3.1")

    for it in parsed.items:
        if it.role in ("table", "table_caption") or it.role.startswith("brief_"):
            continue  # 表格另查；表题居中、不缩进，不按正文检查；简报报头、报尾另查
        p = paras[it.src]
        line = it.src + 1
        text = texts[it.src]
        role = it.role
        raw_text = p.text
        if role in ("body", "h1", "h2", "h3", "h4"):
            _check_body_para(col, doc, p, role, it.split if role != "body" else -1, line, text, std,
                             line_pt, body_pt, lines_req, grid_pitch, unit_roles=frozenset(unit_over))
            for rx, msg, hint in BAD_ORDINALS:
                if rx.match(text):
                    col.add("DOCX-033", "P1", "格式要素", f"层次序数写法不规范：{msg}", line=line, excerpt=text,
                            suggestion=hint, basis="GB/T 9704-2012 7.3.3；GB/T 15834-2011 附录 B.3")
                    break
        elif role == "title":
            _check_title(col, doc, p, line, text, std, unit_roles=frozenset(unit_over))
        elif role == "recipient":
            fl = D.para_first_line_chars(p, doc)
            left = D.para_left_chars(p, doc)
            if abs(fl) > 0.3 or abs(left) > 0.3 or raw_text[:1] in (" ", "　"):
                col.add("DOCX-040", "P1", "段落", "主送机关应居左顶格", line=line, excerpt=text,
                        suggestion="首行缩进和左缩进设为 0", basis="GB/T 9704-2012 7.3.2")
            if text.rstrip().endswith(":"):
                col.add("DOCX-041", "P1", "标点", "主送机关后应用全角冒号", line=line, excerpt=text,
                        suggestion=text.rstrip()[:-1] + "：", basis="GB/T 9704-2012 7.3.2")
        elif role == "attachment":
            if re.match(r"^附\s*件\s*:", text):
                col.add("DOCX-042", "P1", "标点", "“附件”后应用全角冒号", line=line, excerpt=text,
                        suggestion="附件：", basis="GB/T 9704-2012 7.3.4")
            if text and text[-1] in "。；，、.;,":
                col.add("DOCX-042", "P1", "标点", "附件名称后不加标点", line=line, excerpt=text,
                        suggestion=text[:-1], basis="GB/T 9704-2012 7.3.4")
        elif role == "date":
            compact = text.replace(" ", "").replace("　", "")
            if DATE_ZERO_RE.search(compact):
                col.add("DOCX-043", "P0", "数字", "成文日期月、日不编虚位", line=line, excerpt=text,
                        suggestion=re.sub(r"(?<=年)0|(?<=月)0", "", compact), basis="GB/T 9704-2012 7.3.5.4")
            if DATE_CN_RE.match(compact):
                col.add("DOCX-044", "P1", "数字", "成文日期应使用阿拉伯数字", line=line, excerpt=text,
                        suggestion="如“2026年10月16日”", basis="GB/T 9704-2012 7.3.5.4")
            if D.para_align(p, doc) != "right" and D.para_left_chars(p, doc) < 8:
                col.add("DOCX-046", "P2", "段落", "成文日期应排在右侧，按是否盖章右空四字或二字", line=line,
                        excerpt=text, suggestion="运行 fix_docx.py 统一排署名和日期", basis="GB/T 9704-2012 7.3.5")
        elif role == "note":
            if not (text.startswith(("（", "(")) and text.endswith(("）", ")"))):
                col.add("DOCX-045", "P2", "格式要素", "附注应加圆括号", line=line, excerpt=text,
                        suggestion=f"（{text}）", basis="GB/T 9704-2012 7.3.6")
        elif role in ("docno", "signer"):
            if re.search(r"签发人\s*:", text):
                col.add("DOCX-054", "P1", "标点", "“签发人”后应用全角冒号", line=line, excerpt=text,
                        suggestion="签发人：", basis="GB/T 9704-2012 7.2.6")
            if role == "signer" and layout in ("红头", "纪要"):
                col.add("DOCX-162", "P1", "段落", "签发人另起了一行，应与发文字号同一行（居右空一字，上行文发文字号居左空一字）",
                        line=line, excerpt=text, suggestion="抽取全文后用 build_docx.py 重排版头（见 SKILL.md 第四步）",
                        basis="GB/T 9704-2012 7.2.5、7.2.6")
            if role == "docno" and doc_type in UPWARD_TYPES and layout == "红头":
                off = D.leading_offset_chars(p, doc)
                centered = D.para_align(p, doc) in ("center", "right") and not p.text.startswith("\t")
                if centered or abs(off - 1) > 0.5:
                    col.add("DOCX-056", "P1", "段落",
                            f"{doc_type}属上行文，发文字号应居左空一字，与最后一个签发人姓名处在同一行"
                            + ("（现为居中或居右）" if centered else f"（现左空 {off:g} 字）"),
                            line=line, excerpt=text, suggestion="发文字号设为左对齐、左缩进 1 字符，签发人居右空一字排在同一行；"
                            "运行 fix_docx.py 可补足左缩进", basis="GB/T 9704-2012 7.2.5、7.2.6")
        elif role == "issuer":
            if layout in ("红头", "纪要", "命令") and D.para_align(p, doc) != "center":
                col.add("DOCX-161", "P1", "段落", "发文机关标志应居中", line=line, excerpt=text,
                        suggestion="标志居中编排；整个版头有问题的，抽取全文后用 build_docx.py 重排",
                        basis="GB/T 9704-2012 7.2.4")
            if not D.para_is_red(p, doc):
                col.add("DOCX-055", "P1", "格式要素", "发文机关标志应为红色", line=line, excerpt=text,
                        suggestion="字体颜色设为红色", basis="GB/T 9704-2012 7.2.4")

    if layout in ("红头", "纪要") and not overprint:
        _check_header_rule(col, doc, paras, parsed)
    if layout in ("红头", "白头", "纪要"):
        _check_banji(col, doc, paras, parsed)
    if layout == "简报":
        _check_brief_frame(col, doc, paras, parsed)
    if layout == "命令":
        _check_order(col, doc, paras, parsed, line_pt, chars_req)
    _check_joint(col, doc, params)
    if overprint:
        _check_overprint(col, doc, paras, parsed, params, line_pt)
    if att_starts or parsed.find("attachment"):
        _check_attachment_pages(col, doc, parsed, att_paras, att_texts, att_offset, std, line_pt, body_pt, lines_req,
                                grid_pitch, unit_over, chars_req)

    # ---- 纪要出席名单前空一行（10.3） ----
    if layout == "纪要" or doc_type == "纪要":
        _check_attendance_gap(col, doc, paras, parsed, line_pt)

    # ---- 正文表格、未转换的 Markdown 表格行 ----
    _check_tables(col, doc, [el for _, el in tables])
    for i, text in enumerate(texts):
        if is_pipe_residue(text):
            col.agg("pipe", "DOCX-100", "P1", "结构", "正文中有未转换的 Markdown 表格行（竖线、短横线原样成段）",
                    line=i + 1, excerpt=text.strip(), suggestion="用 build_docx.py 重新生成（支持简单管道表格），"
                    "或在 Word 中插入表格", basis="references/_shared/style/output-format.md 第一部分")

    # ---- 全文文本类检查（正文、表格、页眉） ----
    all_paras = list(D.iter_all_paragraphs(doc))
    date_lines = {it.src for it in parsed.items if it.role in ("date", "print")}
    body_index = {p._p: i for i, p in enumerate(paras)}
    header_texts = []
    for section in doc.sections:
        for hdr in (section.header, section.first_page_header, section.even_page_header):
            if not hdr.is_linked_to_previous:
                header_texts.append(_part_text(hdr._element))
    scanned: List[str] = []  # 【起草说明】之前的正文段落和表格文字，用于统计待补、待核实标记
    for p in all_paras:
        if p._p is cut_el:
            break  # 【起草说明】及其后的段落、表格不检查
        text = p.text
        if not text.strip():
            continue
        scanned.append(text)
        line = body_index.get(p._p)
        ln = line + 1 if line is not None else None
        _check_text(col, text, ln, in_signoff=(line in date_lines) if line is not None else False)
        stripped = text.strip()
        if re.match(r"^抄\s*送", stripped):
            if re.match(r"^抄\s*送\s*:", stripped):
                col.add("DOCX-060", "P1", "标点", "“抄送”后应用全角冒号", line=ln, excerpt=stripped,
                        suggestion="抄送：", basis="GB/T 9704-2012 7.4.2")
            if not stripped.endswith("。") and layout != "简报":  # 简报报尾无国标，不要求句号
                col.add("DOCX-060", "P1", "标点", "抄送最后一个机关名称后标句号", line=ln, excerpt=stripped,
                        suggestion=stripped + "。", basis="GB/T 9704-2012 7.4.2")
        if stripped.endswith("印发") and not MARKERS_RE.search(stripped):
            if not PRINT_OK_RE.search(stripped.replace(" ", "").replace("\t", "").replace("　", "")):
                col.add("DOCX-061", "P1", "数字", "印发日期应写作“××××年×月×日印发”，月日不编虚位", line=ln,
                        excerpt=stripped, suggestion="如“2026年10月16日印发”", basis="GB/T 9704-2012 7.4.3")
    for text in header_texts:
        if text.strip():
            _check_text(col, text, None, in_signoff=False)
    first_texts = "".join(texts[:6]) + "".join(header_texts)
    if SECRET_WORDS_RE.search(first_texts) or parsed.find("secret"):
        col.add("DOCX-081", "P2", "保密", "文档含密级或“内部资料”等标识，只能在涉密或内网环境处理，不得上传联网 AI 工具",
                excerpt=(parsed.find("secret")[0].text if parsed.find("secret") else ""),
                suggestion="确认处理环境符合保密规定", basis="references/_shared/guardrails.md 第一条")

    # ---- 页码 ----
    _check_page_numbers(col, doc, layout)

    total = sum(len(re.sub(r"\s", "", t)) for t in texts)
    issues = col.finish()
    notes: List[str] = []
    allow_ids = parse_allow(allow)
    if allow_ids:
        issues = _apply_allow(issues, allow_ids, notes)
    summary = {"P0": 0, "P1": 0, "P2": 0}
    for item in issues:
        summary[item["level"]] += 1
    # 待补、待核实标记单独计数（含表格中的），与 gongwen_lint.py 的 summary.markers 一致
    markers = {"待补": 0, "待核实": 0}
    for t in scanned:
        for m in MARKERS_RE.finditer(t):
            markers["待核实" if m.group(0).startswith("【待核实") else "待补"] += 1
    summary["markers"] = markers
    if unit_over:
        notes += [f"按单位规定使用字体：{line}" for line in F.unit_notes({"unit_fonts": unit_over})]
    return {"file": Path(path).name, "doc_type": doc_type or "其他", "layout": layout, "char_count": total,
            "summary": summary, "issues": issues, "notes": notes, "unit_fonts": dict(unit_over)}


def _apply_allow(issues: List[dict], allow: Sequence[str], notes: List[str]) -> List[dict]:
    """--allow 列出的检查项命中时降为 P2，消息前加“按本单位制度保留：”。"""
    unknown = [a for a in allow if a not in KNOWN_IDS]
    if unknown:
        notes.append(f"--allow 中的检查项 {'、'.join(unknown)} 不存在，已忽略。")
    wanted = set(allow) & KNOWN_IDS
    changed: List[str] = []
    for item in issues:
        if item["id"] in wanted:
            item["level"] = "P2"
            if not item["message"].startswith(ALLOW_PREFIX):
                item["message"] = ALLOW_PREFIX + item["message"]
            changed.append(item["id"])
    if changed:
        notes.append(f"按 --allow 降为 P2 的问题 {len(changed)} 处（{'、'.join(sorted(set(changed)))}），"
                     "依据本单位制度保留，不计入 --strict；交付说明中写明与国标的差异。")
        order = {"P0": 0, "P1": 1, "P2": 2}
        issues.sort(key=lambda x: (order[x["level"]], x["line"] or 0))
    return issues


def _blank_before(doc, paras, idx: int, line_pt: float) -> bool:
    """第 idx 段前是否空了一行：前一段是空段且行高不低于半行，或本段段前距不少于大半行。"""
    import gw_docx as D
    from docx.oxml.ns import qn
    ppr = paras[idx]._p.pPr
    sp = ppr.find(qn("w:spacing")) if ppr is not None else None
    if sp is not None and int(sp.get(qn("w:before")) or 0) >= line_pt * 20 * 0.8:
        return True
    if idx == 0 or paras[idx - 1].text.strip():
        return False
    rule, val = D.para_line_spacing(paras[idx - 1], doc)
    return not (rule == "exact" and val is not None and val < line_pt * 0.5)


def _bottom_rule_color(p) -> Optional[str]:
    """段落下边框的颜色（无下边框返回 None）。"""
    from docx.oxml.ns import qn
    ppr = p._p.pPr
    bdr = ppr.find(qn("w:pBdr")) if ppr is not None else None
    bottom = bdr.find(qn("w:bottom")) if bdr is not None else None
    if bottom is None or bottom.get(qn("w:val")) in (None, "nil", "none"):
        return None
    return (bottom.get(qn("w:color")) or "auto").upper()


def _check_header_rule(col: Collector, doc, paras, parsed) -> None:
    """红头、纪要：发文字号下应有红色分隔线（7.2.7）；前 15 段和联合行文版头之外找不到红色横线的报 DOCX-160。"""
    import gw_docx as D
    has_head = parsed.find("issuer") or parsed.find("docno") or any(
        D.table_mark(t._tbl) == D.JOINT_HEAD_MARK for t in doc.tables)
    if not has_head:
        return
    for p in paras[:15]:
        color = _bottom_rule_color(p)
        if (color and D.is_red(color)) or (_has_top_rule(p) and D.is_red(_top_rule_color(p))):
            return
    first = (parsed.find("issuer") or parsed.find("docno"))
    col.add("DOCX-160", "P1", "格式要素", "版头缺红色分隔线（发文字号下 4 mm 处应有与版心等宽的红色实线）",
            line=first[0].src + 1 if first else None, excerpt=first[0].text if first else "版头",
            suggestion="抽取全文后用 build_docx.py 重排版头（见 SKILL.md 第四步）", basis="GB/T 9704-2012 7.2.7")


def _top_rule_color(p) -> str:
    from docx.oxml.ns import qn
    top = p._p.pPr.find(qn("w:pBdr")).find(qn("w:top"))
    return (top.get(qn("w:color")) or "auto").upper()


def _check_banji(col: Collector, doc, paras, parsed) -> None:
    """版记（7.4）：抄送、印发行用版记表格，首末粗线、行间细线，位于最后一页版心底部；普通段落、无线或不在页底的报 DOCX-163。"""
    import gw_docx as D
    from docx.oxml.ns import qn
    reasons: List[str] = []
    first = None
    for p in doc.paragraphs:  # 含附件页之后的段落
        t = p.text.strip()
        if re.match(r"^抄\s*送\s*[：:]", t) or (t.endswith("印发") and re.search(r"\d{4}\s*年", t)):
            first = first or type("X", (), {"src": paras.index(p) if p in paras else -1, "text": t})
            if not (_has_top_rule(p) or _bottom_rule_color(p)):
                reasons.append("抄送、印发行是没有分隔线的普通段落")
                break
    for tbl in doc.tables:
        el = tbl._tbl
        if not D.is_banji_table(el):
            continue
        tblpr = el.find(qn("w:tblPr"))
        bd = tblpr.find(qn("w:tblBorders")) if tblpr is not None else None
        edge = [bd.find(qn("w:" + s)) if bd is not None else None for s in ("top", "bottom")]
        if any(e is None or e.get(qn("w:val")) in (None, "nil", "none") for e in edge):
            reasons.append("版记表格缺首末分隔线")
        if tblpr is None or tblpr.find(qn("w:tblpPr")) is None:
            reasons.append("版记表格没有固定在最后一页版心底部")
    if reasons:
        col.add("DOCX-163", "P1", "格式要素", "版记不合要求：" + "；".join(dict.fromkeys(reasons)) +
                "（版记应有首末粗线、行间细线，位于最后一页版心底部）", line=first.src + 1 if first and first.src >= 0 else None,
                excerpt=first.text if first else "版记", suggestion="抽取全文后用 build_docx.py 重排（见 SKILL.md 第四步）",
                basis="GB/T 9704-2012 7.4.1")


def _has_top_rule(p) -> bool:
    from docx.oxml.ns import qn
    ppr = p._p.pPr
    bdr = ppr.find(qn("w:pBdr")) if ppr is not None else None
    top = bdr.find(qn("w:top")) if bdr is not None else None
    return top is not None and top.get(qn("w:val")) not in (None, "nil", "none")


def _check_brief_frame(col: Collector, doc, paras, parsed) -> None:
    """简报的报头、报尾。简报没有国家标准格式，这里只查常见做法：报头有简报名称（红色居中）、期号、
    编印单位和日期，其下有红色分隔线；报尾（有的话）上方有分隔线。报头、报尾要素不当作正文检查。"""
    import gw_docx as D
    from docx.oxml.ns import qn
    from gw_layout import BRIEF_SEP_RE
    basis = "实践（简报无国家标准格式，以本单位往期简报为准）"
    by = {it.role: it for it in parsed.items if it.role in ("brief_name", "brief_issue", "brief_editor")}
    missing = [label for role, label in (("brief_name", "简报名称"), ("brief_issue", "期号"),
                                         ("brief_editor", "编印单位和日期")) if role not in by]
    if missing:
        col.add("DOCX-110", "P1", "结构", "简报报头缺少：" + "、".join(missing), excerpt="报头",
                suggestion="报头依次为简报名称、期号行、编印单位与日期同一行，缺的先写【待补】；运行 build_docx.py --layout 简报 重排",
                basis=basis)
    name = by.get("brief_name")
    if name is not None:
        p = paras[name.src]
        if not D.para_is_red(p, doc) or D.para_align(p, doc) != "center":
            col.add("DOCX-111", "P1", "格式要素", "简报名称应为红色、居中", line=name.src + 1, excerpt=name.text,
                    suggestion="简报名称设为红色、居中、大字号", basis=basis)
    head = [paras[it.src] for it in by.values()]
    if head and not any(_bottom_rule_color(p) is not None and D.is_red(_bottom_rule_color(p)) for p in head):
        last = by.get("brief_editor") or by.get("brief_issue") or name
        col.add("DOCX-112", "P1", "格式要素", "简报报头下没有红色分隔线", line=last.src + 1, excerpt=last.text,
                suggestion="在编印单位和日期行下加红色下边框作分隔线", basis=basis)
    for role in ("brief_issue", "brief_editor"):
        it = by.get(role)
        if it is None:
            continue
        p = paras[it.src]
        fonts = [(D.run_font(r, p, doc), len(r.text)) for r in p.runs if r.text.strip()]
        sizes = [(D.run_size(r, p, doc), len(r.text)) for r in p.runs if r.text.strip()]
        font, size = _majority(fonts), _majority(sizes)
        if (font is not None and F.font_class(font) != "仿宋") or (size is not None and abs(size - 16) > 0.25):
            col.add("DOCX-114", "P2", "字体", f"{ROLE_LABELS[role]}一般用三号仿宋，现为“{font}”{size} 磅",
                    line=it.src + 1, excerpt=it.text, suggestion="设为三号仿宋", basis=basis)
    # 报尾：浮动表格，或报尾各行是普通段落
    foot_ok = None
    for tbl in doc.tables:
        text = "".join(t.text or "" for t in tbl._tbl.iter(qn("w:t"))).strip()
        if D.is_banji_table(tbl._tbl) and re.match(r"^(?:报|送|发|抄\s*送)\s*[：:]|^[（(]?共印", text):
            top = tbl._tbl.find(qn("w:tblPr") + "/" + qn("w:tblBorders") + "/" + qn("w:top"))
            foot_ok = top is not None and top.get(qn("w:val")) not in ("nil", "none")
            excerpt = text[:20]
            break
    else:
        foots = [it for it in parsed.items if it.role == "brief_foot"]
        if foots:
            first = foots[0].src
            j = first - 1
            while j >= 0 and not paras[j].text.strip() and _bottom_rule_color(paras[j]) is None:
                j -= 1
            foot_ok = _has_top_rule(paras[first]) or (j >= 0 and (
                _bottom_rule_color(paras[j]) is not None or bool(BRIEF_SEP_RE.match(paras[j].text.strip()))))
            excerpt = foots[0].text
    if foot_ok is False:
        col.add("DOCX-113", "P1", "格式要素", "简报报尾上方没有分隔线", excerpt=excerpt,
                suggestion="报尾（报、送、发、抄送、印数）与正文之间加一条分隔线", basis=basis)


def _gap_lines(doc, paras, a: int, b: int, line_pt: float) -> float:
    """第 a 段与第 b 段之间空了几行：中间空段的行高、a 的段后距、b 的段前距，折合成行。"""
    import gw_docx as D
    from docx.oxml.ns import qn

    def spacing(p, attr: str) -> float:
        ppr = p._p.pPr
        sp = ppr.find(qn("w:spacing")) if ppr is not None else None
        return int(sp.get(qn(f"w:{attr}")) or 0) / 20 if sp is not None else 0.0

    total = spacing(paras[a], "after") + spacing(paras[b], "before")
    for j in range(a + 1, b):
        rule, val = D.para_line_spacing(paras[j], doc)
        total += (val if rule in ("exact", "atLeast") and val else line_pt) + spacing(paras[j], "before") + spacing(paras[j], "after")
    return total / line_pt


def _right_space_chars(doc, p, text: str, chars_req: int) -> float:
    """成文日期右空几字：右对齐时取右缩进，否则 每行字数 − 左缩进 − 日期宽度。"""
    import gw_docx as D
    from docx.oxml.ns import qn
    left = D.para_left_chars(p, doc)
    if D.para_align(p, doc) == "right":
        ppr = p._p.pPr
        ind = ppr.find(qn("w:ind")) if ppr is not None else None
        right = 0.0
        if ind is not None:
            v = ind.get(qn("w:rightChars")) or ind.get(qn("w:endChars"))
            if v is not None and int(v):
                right = int(v) / 100
            else:
                v = ind.get(qn("w:right")) or ind.get(qn("w:end"))
                right = int(v) / 320 if v is not None else 0.0
        return right
    return chars_req - left - text_width_chars(text)


def _check_order(col: Collector, doc, paras, parsed, line_pt: float, chars_req: int) -> None:
    """命令（令）格式（10.2、7.3.5.3）：标志为发文机关全称加“令”或“命令”，居中、红色；标志下空二行居中编令号，
    令号下空二行编正文；签发人职务、签名章和成文日期按 7.3.5.3（成文日期右空四字，签名章下空一行）。"""
    import gw_docx as D
    basis = "GB/T 9704-2012 10.2"
    issuer = parsed.find("issuer")
    lingno = parsed.find("lingno")
    if not issuer:
        col.add("DOCX-120", "P1", "格式要素", "命令（令）格式缺少发文机关标志", excerpt="版头",
                suggestion="标志为发文机关全称加“命令”或“令”字，红色小标宋，居中", basis=basis)
    else:
        it = issuer[0]
        if not it.text.endswith("令"):
            col.add("DOCX-120", "P1", "格式要素", "命令（令）格式的发文机关标志应为发文机关全称加“命令”或“令”字",
                    line=it.src + 1, excerpt=it.text, suggestion="如“某市人民政府令”", basis=basis)
        if D.para_align(paras[it.src], doc) != "center":
            col.add("DOCX-121", "P1", "段落", "命令（令）的发文机关标志应居中", line=it.src + 1, excerpt=it.text,
                    suggestion="标志居中编排", basis=basis)
    if not lingno:
        col.add("DOCX-122", "P1", "格式要素", "命令（令）格式缺少令号（标志下空二行居中编令号）", excerpt="令号",
                suggestion="标志下空二行居中编令号，如“第12号”", basis=basis)
        return
    ln = lingno[0]
    if D.para_align(paras[ln.src], doc) != "center":
        col.add("DOCX-122", "P1", "段落", "令号应居中", line=ln.src + 1, excerpt=ln.text, suggestion="令号居中编排", basis=basis)
    if issuer:
        gap = _gap_lines(doc, paras, issuer[0].src, ln.src, line_pt)
        if abs(gap - 2) > 0.5:
            col.add("DOCX-123", "P1", "段落", f"标志与令号之间应空二行，现为 {gap:.1f} 行", line=ln.src + 1, excerpt=ln.text,
                    suggestion="标志下空二行编令号", basis=basis)
    body = [it for it in parsed.items if it.role in ("body", "h1", "h2", "h3", "h4") and it.src > ln.src]
    if body:
        gap = _gap_lines(doc, paras, ln.src, body[0].src, line_pt)
        if abs(gap - 2) > 0.5:
            col.add("DOCX-124", "P1", "段落", f"令号与正文之间应空二行，现为 {gap:.1f} 行", line=body[0].src + 1,
                    excerpt=body[0].text, suggestion="令号下空二行编正文", basis=basis)
    dates = parsed.find("date")
    if dates:
        d = dates[0]
        space = _right_space_chars(doc, paras[d.src], d.text, chars_req)
        if abs(space - 4) > 0.8:
            col.add("DOCX-125", "P1", "段落", f"成文日期应右空四字，现约右空 {space:.1f} 字", line=d.src + 1,
                    excerpt=d.text, suggestion="成文日期右空四字（签名章下空一行）", basis="GB/T 9704-2012 7.3.5.3")
        sigs = parsed.find("signature")
        if sigs:
            gap = _gap_lines(doc, paras, sigs[-1].src, d.src, line_pt)
            if abs(gap - 1) > 0.5:
                col.add("DOCX-126", "P2", "段落", f"签发人职务（签名章）与成文日期之间应空一行，现为 {gap:.1f} 行",
                        line=d.src + 1, excerpt=d.text, suggestion="签名章下空一行编成文日期", basis="GB/T 9704-2012 7.3.5.3")
    else:
        col.add("DOCX-125", "P1", "格式要素", "命令（令）缺少成文日期", excerpt="成文日期",
                suggestion="签名章下空一行，右空四字编成文日期", basis="GB/T 9704-2012 7.3.5.3")


def _check_joint(col: Collector, doc, params: dict) -> None:
    """联合行文版头（7.2.4）和联署署名（7.3.5.1）。只查本工具生成的表格（tblCaption 标记）。"""
    import gw_docx as D
    from docx.oxml.ns import qn
    seal_mm = float(D.pv(params, "elements.joint.seal_diameter_mm")) if "joint" in params["elements"] else 42.0
    width_tw = D.ctx_text_width_pt(doc) * 20
    for tbl in doc.tables:
        mark = D.table_mark(tbl._tbl)
        if mark == D.JOINT_HEAD_MARK:
            row = tbl.rows[0]
            if len(row.cells) < 2:
                continue
            left, right = row.cells[0], row.cells[1]
            names = [p for p in left.paragraphs if p.text.strip()]
            excerpt = "、".join(p.text.strip() for p in names)
            if len(names) < 2:
                col.add("DOCX-140", "P1", "格式要素", "联合行文版头应并列至少两个联署机关名称，主办机关在前",
                        excerpt=excerpt, suggestion="联署机关名称自上而下排列，主办机关在前", basis="GB/T 9704-2012 7.2.4")
            not_red = [r for p in names for r in p.runs if r.text.strip() and not D.is_red(D.run_color(r, p, doc))]
            if not_red:
                col.add("DOCX-140", "P1", "格式要素", "联合行文的发文机关标志应为红色", excerpt=excerpt,
                        suggestion="机关名称和“文件”二字设为红色小标宋", basis="GB/T 9704-2012 7.2.4")
            tcpr = right._tc.tcPr
            valign = tcpr.find(qn("w:vAlign")) if tcpr is not None else None
            if "文件" not in right.text or valign is None or valign.get(qn("w:val")) != "center":
                col.add("DOCX-141", "P1", "格式要素", "“文件”二字应置于联署机关名称右侧，并以名称为准上下居中",
                        excerpt=right.text.strip() or "文件", suggestion="名称右侧单元格内放“文件”，垂直居中",
                        basis="GB/T 9704-2012 7.2.4")
        elif mark == D.JOINT_SEAL_MARK:
            grid = [int(c.get(qn("w:w")) or 0) for c in tbl._tbl.iter(qn("w:gridCol"))]
            excerpt = "、".join(c.text.strip() for r in tbl.rows for c in r.cells if c.text.strip())
            if grid and min(grid) < D.mm_to_twips(seal_mm) - 20:
                col.add("DOCX-143", "P1", "格式要素",
                        f"联署署名每格宽约 {D.twips_to_mm(min(grid)):.0f} mm，小于印章直径 {seal_mm:g} mm，印章会相交或相切后超出",
                        excerpt=excerpt, suggestion="每行并排的署名减少，使印章互不相交", basis="GB/T 9704-2012 7.3.5.1")
            if sum(grid) > width_tw + 57:
                col.add("DOCX-144", "P1", "页面", "联署署名表格超出版心，印章两端不得超出版心", excerpt=excerpt,
                        suggestion="表格宽度不超过版心 15.6 cm", basis="GB/T 9704-2012 7.3.5.1")


def _check_overprint(col: Collector, doc, paras, parsed, params: dict, line_pt: float) -> None:
    """预印红头纸套打：不应输出红色版头要素；版头留白不少于发文机关标志所占高度，使发文字号落在预印纸版头之下。"""
    import gw_docx as D
    from docx.oxml.ns import qn
    red_hits = [p for p in paras if p.text.strip() and D.para_is_red(p, doc)]
    rules = [p for p in paras if (_bottom_rule_color(p) is not None and D.is_red(_bottom_rule_color(p)))]
    heads = [t for t in doc.tables if D.table_mark(t._tbl) == D.JOINT_HEAD_MARK]
    if red_hits or rules or heads:
        col.add("DOCX-150", "P1", "格式要素", "预印红头纸套打稿不应输出红色发文机关标志和红色分隔线（预印纸上已有）",
                excerpt=(red_hits[0].text if red_hits else "红色版头"),
                suggestion="用 build_docx.py --overprint 重新生成；或确认该稿不是套打", basis="实践（预印红头纸套打）")
    docno = parsed.find("docno")
    if not docno:
        return
    src = docno[0].src
    top = 0.0
    for p in paras[:src + 1]:
        ppr = p._p.pPr
        sp = ppr.find(qn("w:spacing")) if ppr is not None else None
        top += int(sp.get(qn("w:before")) or 0) / 20 if sp is not None else 0.0
        if p is paras[src]:
            break
        rule, val = D.para_line_spacing(p, doc)
        top += val if rule in ("exact", "atLeast") and val else line_pt
    top_mm = top * 25.4 / 72
    need = float(D.pv(params, "elements.overprint.min_docno_top_mm")) if "overprint" in params["elements"] else 44.0
    if top_mm < need - 1.5:
        col.add("DOCX-151", "P1", "页面",
                f"套打稿发文字号上沿距版心上边缘约 {top_mm:.0f} mm，小于预印红头所占高度约 {need:g} mm，会压到预印版头",
                line=src + 1, excerpt=docno[0].text, suggestion="版头区域留空（标志 35 mm 起加标志行高），"
                "或在版式信息写“套打版头高度”后重排", basis="GB/T 9704-2012 7.2.4、7.2.5（位置与正式版头一致）；留白高度为实践值")


def _has_page_break_before(doc, paras, idx: int) -> bool:
    """第 idx 段是否另起一页：段前分页、上一段含分页符或分节符。"""
    from docx.oxml.ns import qn
    p = paras[idx]
    ppr = p._p.pPr
    if ppr is not None:
        pb = ppr.find(qn("w:pageBreakBefore"))
        if pb is not None and pb.get(qn("w:val")) not in ("0", "false"):
            return True
    if idx > 0:
        prev = paras[idx - 1]._p
        for br in prev.iter(qn("w:br")):
            if br.get(qn("w:type")) == "page":
                return True
        if prev.pPr is not None and prev.pPr.find(qn("w:sectPr")) is not None:
            return True
    return False


def _norm_name(text: str) -> str:
    return re.sub(r"[\s　]", "", re.sub(r"^\d{1,2}\s*[.．、]\s*|^附\s*件\s*[：:]\s*", "", text.strip()))


def _check_attachment_pages(col: Collector, doc, parsed, att_paras, att_texts, offset: int, std: dict, line_pt: float,
                            body_pt: float, lines_req: int, grid_pitch, unit_over, chars_req: int) -> None:
    """附件另面编排（7.3.7）：另面，“附件”或“附件1”三号黑体顶格在第一行，附件标题居中在第三行；
    页码与正文连续；正文的附件说明与附件页的顺序号、标题一致。"""
    import gw_docx as D
    from docx.oxml.ns import qn
    basis = "GB/T 9704-2012 7.3.7"
    starts = [i for i, t in enumerate(att_texts) if ATTACH_MARK_RE.match(t.strip())]
    descr = [_norm_name(x) for x in (parsed.fields.get("附件") or [])]
    if not starts:
        return
    # 页码连续：后续各节不重新起页码
    for k, section in enumerate(doc.sections[1:], start=2):
        pg = section._sectPr.find(qn("w:pgNumType"))
        if pg is not None and pg.get(qn("w:start")) is not None:
            col.add("DOCX-134", "P1", "页码", f"第 {k} 节重新起了页码，附件与正文一起装订时页码应连续编排",
                    excerpt=f"页码从 {pg.get(qn('w:start'))} 起", suggestion="删除“起始页码”设置，续前节编号",
                    basis="GB/T 9704-2012 7.5")
    titles: List[str] = []
    labels: List[str] = []
    ends = starts[1:] + [len(att_texts)]
    for s, end in zip(starts, ends):
        idx = offset + s  # 在整篇段落列表中的下标
        p = att_paras[s]
        label = att_texts[s].strip()
        labels.append(re.sub(r"[\s　]", "", label))
        fonts = [(D.run_font(r, p, doc), len(r.text)) for r in p.runs if r.text.strip()]
        sizes = [(D.run_size(r, p, doc), len(r.text)) for r in p.runs if r.text.strip()]
        font, size = _majority(fonts), _majority(sizes)
        bad = []
        if font is not None and F.font_class(font) != "黑体":
            bad.append(f"字体为“{font}”")
        if size is not None and abs(size - 16) > 0.25:
            bad.append(f"字号为 {size} 磅")
        if D.para_align(p, doc) not in ("left", "both") or D.leading_offset_chars(p, doc) > 0.3:
            bad.append("未顶格")
        if bad:
            col.add("DOCX-130", "P1", "格式要素", "附件标识应为三号黑体、顶格：" + "，".join(bad), line=idx + 1, excerpt=label,
                    suggestion="“附件”或“附件1”设为三号黑体，顶格编排在版心左上角第一行", basis=basis)
        if not _has_page_break_before(doc, att_paras, s):
            col.add("DOCX-131", "P1", "页面", "附件应另面编排，附件标识前没有分页", line=idx + 1, excerpt=label,
                    suggestion="附件标识所在段设“段前分页”，使附件另起一页", basis=basis)
        j = s + 1
        while j < end and not att_texts[j].strip():
            j += 1
        if j >= end:
            col.add("DOCX-132", "P1", "格式要素", "附件缺少标题（应居中编排在版心第三行）", line=idx + 1, excerpt=label,
                    suggestion="标识下空一行，居中编附件标题", basis=basis)
            titles.append("")
            continue
        gap = _gap_lines(doc, att_paras, s, j, line_pt)
        tp = att_paras[j]
        titles.append(_norm_name(tp.text))
        if abs(gap - 1) > 0.5 or D.para_align(tp, doc) != "center":
            why = []
            if abs(gap - 1) > 0.5:
                why.append(f"与标识相隔 {gap:.1f} 行，应空一行")
            if D.para_align(tp, doc) != "center":
                why.append("未居中")
            col.add("DOCX-132", "P1", "格式要素", "附件标题应居中编排在版心第三行：" + "，".join(why), line=offset + j + 1,
                    excerpt=tp.text.strip(), suggestion="标识在第一行，空一行，标题居中在第三行", basis=basis)
        for k in range(j + 1, end):
            text = att_texts[k]
            if not text.strip():
                continue
            q = att_paras[k]
            if D.para_align(q, doc) == "center" and is_table_caption(text.strip()):
                continue
            level = is_heading(text.strip())
            _check_body_para(col, doc, q, level or "body", heading_split(text.strip()) if level else -1, offset + k + 1,
                             text, std, line_pt, body_pt, lines_req, grid_pitch, unit_roles=frozenset(unit_over))
    # 附件说明与附件对应
    if not descr:
        col.add("DOCX-133", "P1", "格式要素", "有附件页，但正文没有附件说明（正文后应有“附件：”说明，与附件对应）",
                excerpt=labels[0], suggestion="在正文下空一行，左空二字编“附件：附件名称”", basis=basis)
        return
    n = len(starts)
    expect = ["附件" + str(i + 1) for i in range(n)] if n > 1 else None
    if expect is not None and labels != expect:
        col.add("DOCX-133", "P1", "格式要素", "多个附件的顺序号应依次为 附件1、附件2……，现为：" + "、".join(labels),
                excerpt=labels[0], suggestion="附件标识的顺序号与附件说明一致", basis=basis)
    if len(descr) != n or any(d != t for d, t in zip(descr, titles)):
        col.add("DOCX-133", "P1", "格式要素", "附件说明与附件页的标题或个数不一致：说明为“" + "；".join(descr) +
                "”，附件页为“" + "；".join(titles) + "”", excerpt=labels[0],
                suggestion="附件顺序号和标题与正文的附件说明保持一致", basis=basis)


def _check_attendance_gap(col: Collector, doc, paras, parsed, line_pt: float) -> None:
    """纪要：第一个出席（请假、列席）段前应空一行；排在开头、前面没有正文的不查。"""
    items = [it for it in parsed.items if it.src >= 0 or it.role == "table"]
    for k, it in enumerate(items):
        if it.role != "attendance":
            continue
        prev = items[k - 1].role if k > 0 else None
        if prev in ("body", "h1", "h2", "h3", "h4", "attachment") and not _blank_before(doc, paras, it.src, line_pt):
            col.add("DOCX-047", "P1", "段落", "纪要出席名单前应空一行", line=it.src + 1, excerpt=paras[it.src].text,
                    suggestion="在第一个“出席”段前加一个空行，或运行 fix_docx.py", basis="GB/T 9704-2012 10.3")
        break


def _check_tables(col: Collector, doc, tables: Sequence) -> None:
    """正文表格：宽度不超过版心；文字一般用仿宋或宋体，小四至四号（国标未规定，按实践提示）。"""
    import gw_docx as D
    from docx.oxml.ns import qn
    from docx.table import Table
    width_tw = D.ctx_text_width_pt(doc) * 20
    for n, tbl in enumerate(tables, start=1):
        grid = [int(c.get(qn("w:w")) or 0) for c in tbl.iter(qn("w:gridCol"))]
        tblw = tbl.find(qn("w:tblPr") + "/" + qn("w:tblW"))
        w = sum(grid)
        if tblw is not None and tblw.get(qn("w:type")) == "dxa":
            w = max(w, int(tblw.get(qn("w:w")) or 0))
        if w > width_tw + 57:  # 容差约 1 mm
            col.add("DOCX-101", "P1", "页面", f"第 {n} 张表格宽约 {D.twips_to_mm(w):.0f} mm，超出版心（156 mm）",
                    excerpt=f"表格 {n}", suggestion="表格宽度设为不超过 15.6 厘米，或缩小列宽",
                    basis="GB/T 9704-2012 5.2.1（版心）")
        fonts: List[tuple] = []
        sizes: List[tuple] = []
        seen = set()
        for row in Table(tbl, doc._body).rows:
            for cell in row.cells:
                if cell._tc in seen:
                    continue  # 合并单元格会重复返回
                seen.add(cell._tc)
                for p in cell.paragraphs:
                    for r in p.runs:
                        k = len(r.text.strip())
                        if k:
                            fonts.append((D.run_font(r, p, doc), k))
                            sizes.append((D.run_size(r, p, doc), k))
        font, size = _majority(fonts), _majority(sizes)
        bad_font = font is not None and F.font_class(font) not in ("仿宋", "宋体")
        bad_size = size is not None and not 11.75 <= size <= 14.25
        if bad_font or bad_size:
            col.add("DOCX-102", "P2", "字体", f"第 {n} 张表格文字为“{font}”{size} 磅，一般用仿宋或宋体小四至四号",
                    excerpt=f"表格 {n}", suggestion="表格文字设为仿宋或宋体，小四（12 磅）至四号（14 磅），表头加粗",
                    basis="实践（国标未规定表格字号，按本单位规定）")


# check_document 可能报出的检查项，--allow 只接受这些 ID
KNOWN_IDS = {f"DOCX-{n:03d}" for n in (1, 2, 3, 4, 10, 11, 12, 13, 14, 15, 16, 20, 21, 22, 23, 24, 30, 31, 32, 33,
                                     40, 41, 42, 43, 44, 45, 46, 47, 50, 51, 52, 53, 54, 55, 60, 61, 70, 71, 72,
                                     73, 74, 80, 81, 82, 90, 56, 100, 101, 102, 110, 111, 112, 113, 114,
                                     120, 121, 122, 123, 124, 125, 126, 130, 131, 132, 133, 134,
                                     140, 141, 143, 144, 150, 151)}


def _majority(values: List[tuple]) -> Optional[object]:
    """values: (值, 字数)。返回字数最多的值。"""
    weight: Dict[object, int] = {}
    for v, n in values:
        weight[v] = weight.get(v, 0) + n
    if not weight:
        return None
    return max(weight.items(), key=lambda kv: kv[1])[0]


def _check_body_para(col: Collector, doc, p, role: str, split: int, line: int, text: str, std: dict,
                     line_pt: float, body_pt: float, lines_req: int, grid_pitch: Optional[float],
                     unit_roles=frozenset()) -> None:
    import gw_docx as D
    expected = LEVEL_FONT[role]
    end = len(text) if split is None or split < 0 else split
    fonts: List[tuple] = []
    sizes: List[tuple] = []
    colors: List[tuple] = []
    pos = 0
    for run in p.runs:
        rt = run.text
        n = len(CJK_RE.findall(rt))
        in_head = pos < end
        pos += len(rt)
        if not n:
            continue
        if role != "body" and not in_head:
            continue  # 标题与正文同段时，只按标题部分判断层次标题字体
        fonts.append((D.run_font(run, p, doc), n))
        sizes.append((D.run_size(run, p, doc), n))
        colors.append((D.run_color(run, p, doc), n))
    name = LEVEL_NAME.get(role, "正文")
    main_font = _majority(fonts)
    if main_font is not None:
        cls = F.font_class(main_font)
        if expected in unit_roles and F.same_font(main_font, std.get(expected)):
            pass  # 本单位制度规定的字体（--unit-fonts），视为合规，不再按字体类别判断
        elif role == "body" and cls != "仿宋":
            col.agg("font-body", "DOCX-010", "P0", "字体", f"正文字体为“{main_font}”，应为仿宋体（仿宋_GB2312）",
                    line=line, excerpt=text, suggestion="正文设为 3 号仿宋_GB2312（本机没有时先下载安装，不要用其他字体代替）",
                    basis="GB/T 9704-2012 7.3.3")
        elif role != "body" and cls != expected:
            iid = {"h1": "DOCX-030", "h2": "DOCX-031"}.get(role, "DOCX-032")
            col.agg(f"font-{role}", iid, "P1", "字体", f"{name}字体为“{main_font}”，应为{expected}", line=line,
                    excerpt=text, suggestion=f"{name}设为{std.get(expected, expected)}", basis="GB/T 9704-2012 7.3.3")
        elif not F.same_font(main_font, std.get(expected)):
            # 字体类别对，但不是标准字体本身（如“仿宋”不是“仿宋_GB2312”）：严格模式不接受替代字体
            col.agg(f"alt-{role}-{main_font}", "DOCX-016", "P1", "字体",
                    f"{name}使用非标准字体“{main_font}”，应为“{std.get(expected)}”", line=line, excerpt=text,
                    suggestion=NON_STANDARD_HINT.format(std=std.get(expected)), basis=NON_STANDARD_BASIS)
        if role == "body" and len({F.font_class(f) for f, _ in fonts}) > 1:
            col.agg("font-mixed", "DOCX-010", "P2", "字体", "正文段落中夹有非仿宋字体的文字", line=line, excerpt=text,
                    suggestion="核对是否有意强调；一般正文统一用仿宋", basis="GB/T 9704-2012 7.3.3")
    main_size = _majority(sizes)
    if main_size is not None and abs((main_size or 0) - 16) > 0.25:
        lvl = "P0" if role == "body" else "P1"
        col.agg(f"size-{role}", "DOCX-011", lvl, "字号", f"{name}字号为 {main_size} 磅，应为三号（16 磅）", line=line,
                excerpt=text, suggestion="字号设为三号", basis="GB/T 9704-2012 7.3.3")
    main_color = _majority(colors)
    if main_color not in (None, "AUTO", "000000"):
        col.agg("color", "DOCX-015", "P1", "字体", f"{name}文字颜色为 {main_color}，应为黑色", line=line, excerpt=text,
                suggestion="字体颜色设为黑色或自动", basis="GB/T 9704-2012 5.2.4")
    # 行距
    rule, val = D.para_line_spacing(p, doc)
    ok = True
    per_page = None
    if rule in ("exact", "atLeast"):
        per_page = math.floor(body_pt / val + 1e-6) if val else 0
        ok = per_page == lines_req
    else:
        if grid_pitch and D.para_snap(p, doc) and val <= 1.0:
            per_page = math.floor(body_pt / grid_pitch + 1e-6)
            ok = per_page == lines_req
        else:
            ok = False
    if not ok:
        desc = f"固定值 {val:g} 磅，每页约 {per_page} 行" if rule in ("exact", "atLeast") else f"{val:g} 倍行距"
        col.agg("spacing", "DOCX-012", "P1", "段落", f"正文行距为{desc}，每页应为 {lines_req} 行", line=line,
                excerpt=text, suggestion=f"行距设为固定值 {line_pt:g} 磅，或用文档网格每页 {lines_req} 行",
                basis="GB/T 9704-2012 5.2.3")
    # 首行缩进
    if p.text[:1] in (" ", "　", "\t"):
        col.agg("space-indent", "DOCX-014", "P1", "段落", "段首用空格缩进", line=line, excerpt=p.text,
                suggestion="删除段首空格，改为首行缩进 2 字符", basis="GB/T 9704-2012 7.3.3")
    else:
        fl = D.para_first_line_chars(p, doc)
        if abs(fl - 2) > 0.25:
            col.agg("indent", "DOCX-013", "P1", "段落", f"{name}首行缩进为 {fl:g} 字，应为 2 字", line=line,
                    excerpt=text, suggestion="首行缩进设为 2 字符", basis="GB/T 9704-2012 7.3.3")


def _check_title(col: Collector, doc, p, line: int, text: str, std: dict, unit_roles=frozenset()) -> None:
    import gw_docx as D
    fonts = [(D.run_font(r, p, doc), len(r.text)) for r in p.runs if r.text.strip()]
    sizes = [(D.run_size(r, p, doc), len(r.text)) for r in p.runs if r.text.strip()]
    main_font = _majority(fonts)
    cls = F.font_class(main_font)
    std_title = std.get("小标宋")
    if "小标宋" in unit_roles and F.same_font(main_font, std_title):
        pass  # 本单位制度规定的标题字体（--unit-fonts），视为合规
    elif cls in ("小标宋", "宋体") and not F.same_font(main_font, std_title):
        # 华文中宋、宋体（含加粗）、方正小标宋_GBK 等都不是方正小标宋简体，严格模式不接受替代字体
        col.add("DOCX-016", "P1", "字体", f"标题使用非标准字体“{main_font}”，应为“{std_title}”（宋体、华文中宋、宋体加粗等不能代替）",
                line=line, excerpt=text, suggestion=NON_STANDARD_HINT.format(std=std_title) + "；该字体为商业字体，需单位授权",
                basis=NON_STANDARD_BASIS)
    elif cls != "小标宋":
        col.add("DOCX-020", "P1", "字体", f"标题字体为“{main_font}”，应为小标宋体", line=line, excerpt=text,
                suggestion=f"标题设为 2 号{std.get('小标宋')}", basis="GB/T 9704-2012 7.3.1")
    main_size = _majority(sizes)
    if main_size is None or abs(main_size - 22) > 0.25:
        col.add("DOCX-021", "P1", "字号", f"标题字号为 {main_size} 磅，应为二号（22 磅）", line=line, excerpt=text,
                suggestion="标题字号设为二号", basis="GB/T 9704-2012 7.3.1")
    if D.para_align(p, doc) != "center":
        col.add("DOCX-022", "P1", "段落", "标题应居中", line=line, excerpt=text, suggestion="标题居中排列",
                basis="GB/T 9704-2012 7.3.1")
    # 回行后某一行过短（如只剩“总结”二字）：按二号字估算每行字数，模拟自动折行
    size = main_size or 22
    max_chars = max(1, int(D.ctx_text_width_pt(doc) // size))
    for seg in text.split("\n"):
        width = text_width_chars(seg.strip())
        if width <= 0:
            continue
        tail = width % max_chars if width > max_chars else width
        if 0 < tail < 3 and text_width_chars(text.replace("\n", "")) >= 3:
            col.add("DOCX-024", "P2", "段落", "标题回行后有一行少于 3 字，排列不匀", line=line, excerpt=text,
                    suggestion="在词意完整处手动换行，上行不短于下行，排成梯形或菱形", basis="GB/T 9704-2012 7.3.1")
            break


def _check_text(col: Collector, text: str, line: Optional[int], *, in_signoff: bool) -> None:
    compact = text.strip()
    for m in DOCNO_BRACKET_RE.finditer(compact):
        col.add("DOCX-050", "P0", "格式要素", "发文字号年份应用六角括号“〔〕”", line=line, excerpt=m.group(0),
                suggestion=re.sub(r"[\[［(（【﹝{｛](\s*\d{4}\s*)[\]］)）】﹞}｝]", r"〔\1〕", m.group(0)),
                basis="GB/T 9704-2012 7.2.5")
    for m in DOCNO_SHORT_YEAR_RE.finditer(compact):
        col.add("DOCX-051", "P0", "格式要素", "发文字号年份应标全称", line=line, excerpt=m.group(0),
                suggestion="如“〔2026〕”", basis="GB/T 9704-2012 7.2.5")
    for m in DOCNO_DI_RE.finditer(compact):
        col.add("DOCX-052", "P0", "格式要素", "发文字号顺序号不加“第”字", line=line, excerpt=m.group(0),
                suggestion=re.sub(r"第\s*", "", m.group(0)), basis="GB/T 9704-2012 7.2.5")
    for m in DOCNO_ZERO_RE.finditer(compact):
        col.add("DOCX-053", "P0", "格式要素", "发文字号顺序号不编虚位", line=line, excerpt=m.group(0),
                suggestion=re.sub(r"〕\s*0+", "〕", m.group(0)), basis="GB/T 9704-2012 7.2.5")
    if not in_signoff:
        for m in DATE_ZERO_RE.finditer(compact.replace(" ", "")):
            col.add("DOCX-090", "P1", "数字", "日期月、日一般不编虚位", line=line, excerpt=compact,
                    suggestion="如“2026年1月5日”", basis="GB/T 9704-2012 7.3.5.4；GB/T 15835-2011")
            break
    for m in MARKERS_RE.finditer(compact):
        col.add("DOCX-080", "P2", "残留", "文中仍有待补或待核实标记，补齐后再定稿", line=line, excerpt=m.group(0),
                suggestion="补齐信息并删除标记", basis="references/_shared/guardrails.md 第二条")


def _check_page_numbers(col: Collector, doc, layout: str) -> None:
    from docx.oxml.ns import qn
    section = doc.sections[0]
    settings = doc.settings.element
    even_odd = settings.find(qn("w:evenAndOddHeaders")) is not None and \
        settings.find(qn("w:evenAndOddHeaders")).get(qn("w:val")) not in ("0", "false")
    odd = _footer_info(section.footer._element) if not section.footer.is_linked_to_previous else {"has_page": False}
    if not odd.get("has_page"):
        col.add("DOCX-070", "P1", "页码", "页脚中没有页码", excerpt="无页码", suggestion="插入“— 1 —”格式页码",
                basis="GB/T 9704-2012 7.5")
        return
    if not PAGE_TEXT_RE.match(odd["text"]):
        col.add("DOCX-071", "P1", "页码", "页码格式应为“— 1 —”（数字两侧各一条一字线）", excerpt=odd["text"],
                suggestion="— 1 —", basis="GB/T 9704-2012 7.5")
    fonts = {F.font_class(f) for f in odd["fonts"]}
    if (fonts and fonts != {"宋体"}) or (odd["sizes"] and odd["sizes"] != {14.0}):
        col.add("DOCX-072", "P2", "页码", "页码一般用 4 号半角宋体", excerpt=odd["text"],
                suggestion="页码设为宋体四号", basis="GB/T 9704-2012 7.5")
    if not even_odd:
        col.add("DOCX-073", "P2", "页码", "未设置奇偶页不同，双页页码无法居左", excerpt=odd["text"],
                suggestion="页面设置→版式→奇偶页不同；单页居右空一字，双页居左空一字", basis="GB/T 9704-2012 7.5")
    else:
        even = _footer_info(section.even_page_footer._element) if not section.even_page_footer.is_linked_to_previous \
            else {"has_page": False}
        if odd.get("align") != "right" or not even.get("has_page") or even.get("align") not in ("left", None):
            col.add("DOCX-073", "P2", "页码", "单页页码应居右空一字，双页页码居左空一字", excerpt=odd["text"],
                    suggestion="运行 fix_docx.py 重设页码", basis="GB/T 9704-2012 7.5")
    if layout == "信函":
        first = section.first_page_footer
        has_first = section.different_first_page_header_footer and not first.is_linked_to_previous
        if not has_first or _footer_info(first._element).get("has_page"):
            col.add("DOCX-074", "P1", "页码", "信函格式首页不显示页码", excerpt="首页页码",
                    suggestion="设置首页不同，首页页脚不放页码", basis="GB/T 9704-2012 10.1")


def render_text(result: dict) -> str:
    s = result["summary"]
    out = [f"文件：{result['file']}（版式：{result.get('layout')}，文种：{result['doc_type']}）",
           f"结论：P0 {s['P0']} 项，P1 {s['P1']} 项，P2 {s['P2']} 项。" + ("可以交付。" if s["P0"] == 0 else "存在 P0，修改后再交付。")]
    out += [f"说明：{n}" for n in result.get("notes", [])]
    for i, it in enumerate(result["issues"], start=1):
        loc = f"第 {it['line']} 段" if it.get("line") else "全文"
        out.append(f"{i}. [{it['level']}] {it['message']}（{loc}：{it['excerpt']}）")
        if it.get("suggestion"):
            out.append(f"   建议：{it['suggestion']}；依据：{it['basis']}")
    return "\n".join(out)


def main(argv: Optional[Sequence[str]] = None) -> int:
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    parser = argparse.ArgumentParser(description="检查 docx 是否符合 GB/T 9704-2012《党政机关公文格式》。",
                                     epilog="退出码：默认 0；--strict 时存在 P0 返回 1；输入错误返回 2。")
    parser.add_argument("file", help="待检查的 .docx 文件")
    parser.add_argument("--layout", help="版式：红头、信函、纪要、命令、白头、材料、简报；不写时自动判断")
    parser.add_argument("--overprint", action="store_true",
                        help="文档是预印红头纸套打稿：不应有红色版头，版头区域应留空（DOCX-150、151）；版式缺省按红头")
    parser.add_argument("--json", action="store_true", help="以机检 JSON 输出（开发规范第 8 节）")
    parser.add_argument("--strict", action="store_true", help="存在 P0 时返回 1")
    parser.add_argument("--allow", metavar="DOCX-050",
                        help="本单位制度明文规定、与国标不同的写法对应的检查项（逗号分隔，如 DOCX-050 方括号文号），"
                             "命中时降为 P2、注明“按本单位制度保留”，不计入 --strict")
    parser.add_argument("--unit-fonts", metavar="角色=字体名,…", help=F.UNIT_FONTS_HELP)
    args = parser.parse_args(argv)
    try:
        import docx  # noqa: F401
    except ImportError:
        print("缺少 python-docx。请先安装：pip install -U python-docx（需要 1.2.0 及以上版本）")
        return 2
    try:
        unit_fonts = F.parse_unit_fonts(args.unit_fonts, F.load_params())
    except ValueError as exc:
        print(f"错误：{exc}")
        return 2
    path = Path(args.file)
    if not path.is_file() or path.suffix.lower() != ".docx":
        print(f"错误：找不到 docx 文件 {path}")
        return 2
    if args.layout and args.layout not in LAYOUT_NAMES:
        print(f"错误：无法识别的版式：{args.layout}。可选：{'、'.join(LAYOUT_NAMES)}")
        return 2
    try:
        result = check_document(path, layout=args.layout, allow=args.allow, unit_fonts=unit_fonts,
                                overprint=args.overprint)
    except Exception as exc:  # 文件损坏、不是 docx 等
        print(f"错误：无法读取 {path}：{exc}")
        return 2
    if args.json:
        print(json.dumps(result, ensure_ascii=False, indent=2))
    else:
        print(render_text(result))
    if args.strict and result["summary"]["P0"]:
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
