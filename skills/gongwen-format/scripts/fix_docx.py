#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""按 GB/T 9704-2012 修复已有 docx 的格式\u3002

按段落文本识别角色\uff08标题\u3001主送\u3001各级层次标题\u3001正文\u3001附件说明\u3001署名\u3001日期等\uff09\uff0c统一页面设置\u3001
字体字号\u3001行距\u3001缩进和页码\u3002只改格式\uff0c不改文字\uff08段首用于缩进的空格除外\uff1b加 --keep-text 时连这类空格也不动\uff09\u3002
输出新文件\uff0c不修改原文件\uff1b同时输出一份修复报告\uff08Markdown\uff09\uff0c并对新文件运行 check_docx.py\u3002

退出码\uff1a0 成功\uff1b2 输入错误或缺少依赖\u3002
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Dict, List, Optional, Sequence

sys.path.insert(0, str(Path(__file__).resolve().parent))

import gw_fonts as F  # noqa: E402
from gw_layout import (ATTACH_MARK_RE, BODY_ROLES, ROLE_LABELS, UPWARD_TYPES, doc_kind,  # noqa: E402
                       draft_note_index, has_brief_header, heading_levels, heading_split, is_heading, parse_brief, parse_lines,
                       summarize)

# 两个非空段落之间应保留的空行数（键为前一段角色、后一段角色；None 表示任意）
GAP_RULES = [
    (("issuer", None), 2), (("docno", None), 2), (("title", None), 1), (("subtitle", None), 1),
    (("title_note", None), 1), (("lingno", None), 2), (("brief_editor", None), 1), ((None, "attachment"), 1), ((None, "signature"), -1), ((None, "date"), -1),
]


def _gap_wanted(prev_role: Optional[str], next_role: Optional[str], seal: bool) -> int:
    for (a, b), n in GAP_RULES:
        if (a is None or a == prev_role) and (b is None or b == next_role):
            if n == -1:
                return 2 if seal else 1
            return n
    return 0


# 要改文字才能消除的检查项（标点、数字、发文字号写法、序数写法、附注括号、段首空格缩进、残留标记等）：
# --keep-text 时只在报告中列出，不改
TEXT_ISSUE_IDS = {"DOCX-014", "DOCX-033", "DOCX-041", "DOCX-042", "DOCX-043", "DOCX-044", "DOCX-045", "DOCX-050",
                  "DOCX-051", "DOCX-052", "DOCX-053", "DOCX-054", "DOCX-060", "DOCX-061", "DOCX-080", "DOCX-082",
                  "DOCX-090"}
GAP_BEFORE_ATTENDANCE = ("body", "h1", "h2", "h3", "h4", "attachment")


def _first_attendance(parsed):
    """第一个出席\uff08请假\u3001列席\uff09段\uff1b它前面是正文\u3001层次标题或附件说明时才需要空一行\uff0c否则返回 None\u3002"""
    items = [it for it in parsed.items if it.src >= 0]
    for k, it in enumerate(items):
        if it.role == "attendance":
            return it if k > 0 and items[k - 1].role in GAP_BEFORE_ATTENDANCE else None
    return None


def _is_empty(p) -> bool:
    from docx.oxml.ns import qn
    if p.text.strip():
        return False
    el = p._p
    for tag in ("w:drawing", "w:pict", "w:object", "w:fldChar", "w:sectPr"):
        if el.find(".//" + qn(tag)) is not None:
            return False
    return True


def _snapshot_page(section) -> Dict[str, str]:
    from docx.oxml.ns import qn
    out: Dict[str, str] = {}
    sect = section._sectPr
    for tag, attrs in (("pgSz", ("w", "h")), ("pgMar", ("top", "bottom", "left", "right", "footer")),
                       ("docGrid", ("type", "linePitch", "charSpace"))):
        el = sect.find(qn(f"w:{tag}"))
        for a in attrs:
            out[f"{tag}.{a}"] = el.get(qn(f"w:{a}")) if el is not None else None
    return out


def _describe_page_change(before: Dict[str, str], after: Dict[str, str]) -> List[str]:
    import gw_docx as D
    labels = {"pgSz.w": "纸张宽", "pgSz.h": "纸张高", "pgMar.top": "上边距", "pgMar.bottom": "下边距",
              "pgMar.left": "左边距", "pgMar.right": "右边距", "pgMar.footer": "页脚距边界"}
    out = []
    for key, label in labels.items():
        if before.get(key) != after.get(key):
            b = f"{D.twips_to_mm(int(before[key])):.1f}" if before.get(key) else "未设置"
            a = f"{D.twips_to_mm(int(after[key])):.1f}"
            out.append(f"{label} {b} → {a} mm")
    if any(before.get(k) != after.get(k) for k in ("docGrid.type", "docGrid.linePitch", "docGrid.charSpace")):
        out.append("文档网格改为每页 22 行\u3001每行 28 字")
    return out


def _all_texts(doc) -> List[str]:
    """正文段落和表格单元格的全部文字\uff08按文档顺序\uff09\uff0c用于核对\u201c文字没有被改动\u201d\u3002"""
    import gw_docx as D
    return [p.text for p in D.iter_all_paragraphs(doc)]


def fix_document(src: Path, dst: Path, *, layout: Optional[str] = None,
                 seal: Optional[bool] = None, installed: Optional[set] = None, params: Optional[dict] = None,
                 unit_fonts=None, keep_text: bool = False) -> dict:
    """修复 src\uff0c保存到 dst\uff08调用方保证 dst 不覆盖已有文件\uff09\u3002返回报告字典\u3002

    字体严格模式\uff1a各角色一律改为标准字体名\uff1binstalled\uff08已安装字体名集合\uff0c空集合表示无法检测\uff09只用于核对\uff0c
    缺字体时在报告中给出字体核对结果\uff08font_status\u3001missing_fonts\u3001font_notice 等\uff09\uff0c不替代\u3001不降级\u3002
    unit_fonts\uff1a按单位规定指定标准字体\uff08\u201c仿宋=方正仿宋_GBK\u201d字符串或 {角色: 字体名}\uff09\uff0c指定后该角色改为单位字体名\uff0c严格模式不变\uff1b
    角色或字体名不合法时抛 ValueError\u3002
    keep_text\uff1a不改任何文字\uff08用户要求\u201c内容别动\u201d\u201c一个字都别改\u201d时用\uff09\u3002只调页面\u3001字体字号\u3001行距\u3001缩进\u3001页码\uff1b
    段首用于缩进的空格也保留\uff08该段首行缩进设为 0\uff0c缩进由空格承担\uff09\uff1b标点\u3001文字问题只在报告中列出\u3002
    保存前核对全部文字与原文件一致\uff0c不一致时抛 ValueError 且不保存\u3002
    """
    import gw_docx as D
    from docx import Document
    from docx.oxml.ns import qn
    from gw_format import (apply_role, ensure_upward_docno_indent, keep_signoff_together, layout_signature_block,
                           offset_leading_spaces, strip_leading_spaces)

    params = params or D.load_params()
    params = F.apply_unit_fonts(params, F.parse_unit_fonts(unit_fonts, params))
    doc = Document(str(src))
    before_texts = _all_texts(doc) if keep_text else []
    report: dict = {"input": str(src), "output": str(dst), "page": [], "roles": {}, "notes": [], "warnings": [],
                    "untouched": [], "keep_text": keep_text}
    method = ""
    if installed is None:
        installed, method = F.installed_fonts()
    fonts, check = F.resolve_fonts(params, installed, method)
    report["fonts"] = fonts
    report.update(check)
    ctx = D.Ctx(params)

    # ---- 版式判断 ----
    paras = list(doc.paragraphs)
    texts = [D.direct_text(p) or p.text for p in paras]
    # 【起草说明】之后不是正文，不排版、保持原样（output-format.md 第二部分）
    cut = draft_note_index(texts)
    after_cut: set = set()  # 【起草说明】之后的表格元素
    if cut is not None:
        cut_el = paras[cut]._p
        seen = False
        for child in doc.element.body.iterchildren():
            seen = seen or child is cut_el
            if seen and child.tag == qn("w:tbl"):
                after_cut.add(child)
        report["untouched"].append(f"第 {cut + 1} 段为【起草说明】，从该段起共 {len(paras) - cut} 段未排版；"
                                   "起草说明不是正文\uff0c成稿中应删除")
        paras, texts = paras[:cut], texts[:cut]
    # 附件另面编排：首行只有“附件”“附件1”的段落起是附件页，主文识别到它之前为止，附件页另行调整
    att_starts = [i for i, t in enumerate(texts) if ATTACH_MARK_RE.match(t.strip())]
    att_paras, att_texts = [], []
    if att_starts:
        att_paras, att_texts = paras[att_starts[0]:], texts[att_starts[0]:]
        paras, texts = paras[:att_starts[0]], texts[:att_starts[0]]
    # 正文表格（不含版记）在识别时作占位，表格上一行的表题按表题排，不当成署名或正文
    anchors = [a for a, _ in D.body_tables(doc) if cut is None or a < cut]
    red_flags = [D.para_is_red(p, doc) for p in paras]
    sec0 = doc.sections[0]
    is_letter = bool(D.letter_head_text(sec0))
    brief = layout == "简报" or (layout is None and not is_letter and has_brief_header(texts, red_flags))
    parsed = (parse_brief(texts, red=red_flags, tables_after=anchors, force=layout == "简报") if brief
              else parse_lines(texts, red=red_flags, tables_after=anchors))
    kind, doc_type = doc_kind(str(parsed.fields.get("标题") or ""))
    if layout is None:
        if brief:
            layout = "简报"
        elif is_letter:
            layout = "信函"
        elif parsed.find("issuer"):
            layout = ("纪要" if parsed.find("issuer")[0].text.endswith("纪要")
                      else "命令" if parsed.find("lingno") else "红头")
        elif any(D.table_mark(t._tbl) == D.JOINT_HEAD_MARK for t in doc.tables):
            layout = "红头"  # 联合行文版头在表格里，没有可识别的标志段
        else:
            layout = "白头" if kind == "法定" else "材料"
    is_letter = is_letter or layout == "信函"
    if seal is None:
        seal = bool(params["seal_default"].get(layout, False))
    report["layout"] = layout
    report["seal"] = seal
    report["summary"] = summarize(parsed)
    report["warnings"] += [w for w in parsed.warnings if "空格缩进" not in w]

    # ---- 页面 ----
    D.setup_document(doc, ctx)
    for section in doc.sections:
        before = _snapshot_page(section)
        D.setup_section(section, ctx, letter=is_letter)
        report["page"] += _describe_page_change(before, _snapshot_page(section))

    # ---- 段落 ----
    role_of: Dict[int, str] = {}
    stats: Dict[str, Dict[str, int]] = {}
    stripped = 0
    kept_space = 0

    def _apply(p, role: str, split: int = -1):
        """按角色设格式\uff1bkeep_text 时段首空格保留\uff0c该段首行缩进设为 0\uff08缩进由空格承担\uff09\uff0c不改文字\u3002"""
        nonlocal kept_space
        changes = apply_role(p, role, ctx, split=split)
        if keep_text and role in BODY_ROLES and p.text[:1] in (" ", "　", "\t"):
            if D.set_para(p, reset_indent=True):
                changes.add("段落格式")
            kept_space += 1
        elif keep_text and role in ("recipient", "salutation") and offset_leading_spaces(p):
            changes.add("段落格式")  # 顶格的称呼、主送机关保留行首空格，用负左缩进抵消
        return changes

    for it in parsed.items:
        if it.role == "table":
            continue  # 正文表格不调整（见下文“表格”）
        p = paras[it.src]
        role_of[it.src] = it.role
        # 发文字号行的行首制表符用于“居左空一字”，版头要素不删行首空白
        if not keep_text and it.role not in ("docno", "signer", "issuer", "copy_no", "secret", "urgency", "brief_name",
                                             "brief_issue", "brief_editor") and strip_leading_spaces(p):
            stripped += 1
        changes = _apply(p, it.role, it.split)
        if it.role == "docno" and layout == "红头" and doc_type in UPWARD_TYPES and ensure_upward_docno_indent(p, doc, ctx):
            changes.add("段落格式")
            report["notes"].append(f"{doc_type}属上行文，发文字号已设为居左空一字（7.2.5）：" + p.text.strip())
        bucket = stats.setdefault(it.role, {"段数": 0})
        bucket["段数"] += 1
        for c in changes:
            bucket[c] = bucket.get(c, 0) + 1
    if stripped:
        report["notes"].append(f"删除段首空格缩进 {stripped} 段，缩进改由“首行缩进 2 字符”实现")
    if keep_text:
        report["notes"].append("按 --keep-text\uff1a未改动任何文字\uff08含段首空格\u3001标点\u3001数字\u3001文字写法\uff09\uff0c只调整格式"
                               + (f"；{kept_space} 段用空格缩进的，保留空格并把首行缩进设为 0，缩进仍由空格承担" if kept_space else ""))

    # 空行：结构性空行保留，其余压到 1 磅
    nonempty = [i for i in range(len(paras)) if not _is_empty(paras[i])]
    compressed = 0
    normalized = 0
    for k, i in enumerate(nonempty + [len(paras)]):
        prev_i = nonempty[k - 1] if k > 0 else -1
        gap = list(range(prev_i + 1, i))
        if not gap or (prev_i < 0 and layout in ("红头", "纪要")):
            continue  # 红头文件首行前的空段可能用于定位发文机关标志，保持原样
        want = _gap_wanted(role_of.get(prev_i), role_of.get(i), seal) if prev_i >= 0 and i < len(paras) else 0
        if prev_i >= 0 and i < len(paras) and role_of.get(prev_i) in BODY_ROLES and role_of.get(i) in BODY_ROLES:
            want = 0
        if role_of.get(i) == "attendance" and role_of.get(prev_i) in GAP_BEFORE_ATTENDANCE:
            want = 1  # 纪要出席名单在正文下空一行（10.3）
        if layout == "简报" and role_of.get(i) == "title" and role_of.get(prev_i) not in (None, "title") \
                and not str(role_of.get(prev_i)).startswith("brief_"):
            want = 1  # 简报两篇之间空一行
        for n, j in enumerate(gap):
            if n < want:
                if D.set_para(paras[j], line_exact=ctx.line_pt, before_pt=0, after_pt=0, reset_indent=False):
                    normalized += 1
            else:
                if D.set_para(paras[j], line_exact=1, before_pt=0, after_pt=0, reset_indent=False, snap=False):
                    compressed += 1
    if compressed:
        report["notes"].append(f"多余空行 {compressed} 个已压缩为 1 磅高（未删除，可在 Word 中显示编辑标记查看）")

    def empties_before(idx: int) -> int:
        n = 0
        j = idx - 1
        while j >= 0 and _is_empty(paras[j]):
            n += 1
            j -= 1
        return min(n, 2 if seal else 1)

    # 标题后空一行、附件说明前空一行：没有空段时用段前段后距补足
    titles = [it for it in parsed.items if it.role in ("title", "subtitle", "title_note")]
    if layout == "简报":  # 每篇标题后空一行
        for k, it in enumerate(titles):
            nxt = titles[k + 1].src if k + 1 < len(titles) else None
            if it.src + 1 < len(paras) and not _is_empty(paras[it.src + 1]) and nxt != it.src + 1:
                D.set_para(paras[it.src], after_pt=ctx.line_pt, before_pt=0, reset_indent=False)
    elif titles:
        last = titles[-1].src
        if last + 1 < len(paras) and not _is_empty(paras[last + 1]):
            D.set_para(paras[last], after_pt=ctx.line_pt, before_pt=0, reset_indent=False)
    atts = parsed.find("attachment")
    if atts and atts[0].src > 0 and not _is_empty(paras[atts[0].src - 1]):
        D.set_para(paras[atts[0].src], before_pt=ctx.line_pt, after_pt=0, reset_indent=False)
    for it in atts[:-1]:
        D.set_para(paras[it.src], reset_indent=False, keep_next=True)  # 附件说明各行不分页
    # 纪要出席名单前空一行：前面没有空段时用段前距补足
    first_att = _first_attendance(parsed)
    if first_att is not None and first_att.src > 0 and not _is_empty(paras[first_att.src - 1]):
        D.set_para(paras[first_att.src], before_pt=ctx.line_pt, after_pt=0, reset_indent=False)
        report["notes"].append("纪要出席名单前已空一行\uff08段前距一行\uff0cGB/T 9704-2012 10.3\uff09")

    sigs = [paras[it.src] for it in parsed.find("signature")]
    dates = parsed.find("date")
    date_p = paras[dates[0].src] if dates else None
    joint_seal = any(D.table_mark(t._tbl) == D.JOINT_SEAL_MARK for t in doc.tables)
    if layout == "命令":
        report["untouched"].append("命令\uff08令\uff09的签发人职务\u3001签名章位置和成文日期\uff087.3.5.3\uff09未调整\uff1b"
                                   "需要按标准排时用 build_docx.py --layout 命令 重排")
    elif joint_seal:
        report["untouched"].append("联署署名表格和成文日期位置\uff087.3.5.1\uff09未调整\uff0c请人工核对印章对应关系")
    elif sigs or date_p is not None:
        first_idx = (parsed.find("signature") or dates)[0].src
        need = (2 if seal else 1) - empties_before(first_idx)
        per_para = layout_signature_block(sigs, date_p, ctx, seal=seal, before_lines=max(0, need))
        roles_seq = ["signature"] * len(sigs) + (["date"] if date_p is not None else [])
        for role, changes in zip(roles_seq, per_para):
            bucket = stats.setdefault(role, {"段数": 0})
            if "段落格式" in changes:
                bucket["段落格式"] = bucket.get("段落格式", 0) + 1
        if any(per_para):
            report["notes"].append("署名和成文日期已按" + ("加盖印章\uff087.3.5.1\uff09" if seal else "不加盖印章\uff087.3.5.2\uff09") + "重新排列")
        notes = parsed.find("note")
        last_idx = notes[0].src if notes else (dates[0].src if dates else parsed.find("signature")[-1].src)
        att_first = atts[0].src if atts and atts[-1].src < first_idx else None
        if keep_signoff_together(paras, first_idx, last_idx, attach_first_idx=att_first):
            report["notes"].append("正文最后一段至落款已设\u201c与下段同页\u201d\uff0c落款不会单独成页")

    for role, bucket in stats.items():
        report["roles"][ROLE_LABELS.get(role, role)] = bucket

    # ---- 附件页（7.3.7）：标识三号黑体顶格并另页，标题居中，正文同主文 ----
    if att_paras:
        starts = [i for i, t in enumerate(att_texts) if ATTACH_MARK_RE.match(t.strip())]
        for s, end in zip(starts, starts[1:] + [len(att_texts)]):
            apply_role(att_paras[s], "attach_label", ctx)
            D.set_para(att_paras[s], reset_indent=False, page_break_before=True)
            j = s + 1
            while j < end and not att_texts[j].strip():
                j += 1
            if j < end:
                apply_role(att_paras[j], "title", ctx)
            att_levels, _ = heading_levels(att_texts[j + 1:end])
            for k in range(j + 1, end):
                t = att_texts[k].strip()
                if not t:
                    continue
                if not keep_text:
                    strip_leading_spaces(att_paras[k])
                level = att_levels.get(k - j - 1) if att_levels else is_heading(t)
                _apply(att_paras[k], level or "body", heading_split(t) if level else -1)
        report["notes"].append(f"附件另面编排：{len(starts)} 个附件页的标识、标题、正文格式已按 7.3.7 调整；"
                               "标识与标题之间要空一行\uff0c脚本不增删空行\uff0c请核对")

    # ---- 表格 ----
    other_tables = 0
    for table in doc.tables:
        if table._tbl in after_cut or D.is_frame_table(table._tbl):
            continue  # 起草说明中的附表、联合行文版头和联署署名表格不处理
        cells_text = "".join(c.text for row in table.rows for c in row.cells)
        if "抄送" in cells_text or cells_text.strip().endswith("印发") or (
                layout == "简报" and D.is_banji_table(table._tbl)):
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        if p.text.strip():
                            apply_role(p, "cc", ctx)
            report["notes"].append(("简报报尾表格" if layout == "简报" else "版记表格") + "已统一为 4 号仿宋")
        else:
            other_tables += 1
    if other_tables:
        report["untouched"].append(f"正文中的表格 {other_tables} 个未调整（国标未规定表格字号，按本单位要求处理）")

    # ---- 页眉、页脚、页码 ----
    # 原稿页码在页眉或页脚里的，连同其他杂项一并清除，页脚按 7.5 重设（否则旧页码残留在页眉）
    removed = D.clear_headers_footers(doc, letter=is_letter)
    for section in doc.sections:
        D.setup_page_numbers(section, ctx, first_page_blank=False)
    if removed:
        report["notes"].append("已清除原页眉\u3001页脚中的内容\uff08页码和杂项\uff09\uff1a" + "\uff1b".join(removed)
                               + "\u3002公文不用页眉\uff0c页码按标准重设在页脚\uff1b其中有需要保留的文字\uff0c请自行在 Word 中补回")
    report["notes"].append("页码已重设为 4 号宋体\u201c\u2014 1 \u2014\u201d\uff0c单页居右空一字\u3001双页居左空一字" +
                           ("\uff1b信函首页页眉\uff08发文机关标志\u3001红色双线\uff09和首页页脚双线保持原样\uff0c首页不显示页码" if is_letter else ""))
    report["stripped"] = stripped
    report["compressed"] = compressed
    if keep_text and _all_texts(doc) != before_texts:
        raise ValueError("--keep-text 模式下检测到文字被改动\uff0c已中止\uff0c未保存")
    dst.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(dst))
    return report


def render_report(report: dict) -> str:
    out = ["# 修复报告", "", f"- 原文件：{report['input']}（未修改）", f"- 新文件：{report['output']}",
           f"- 版式判断：{report['layout']}；署名按{'加盖印章' if report['seal'] else '不加盖印章'}排列"]
    status = F.status_line(report)
    if status:
        out.append(f"- 字体：文件中一律写入标准字体名；本机检测：{status}")
    out += [f"- 按单位规定使用：{line}" for line in F.unit_notes(report)]
    out.append("")
    notice = F.render_notice(report, markdown=True)  # 缺字体或无法检测字体时，醒目提示放在最前面（引用块，不占章节编号）
    if notice:
        out += [notice, ""]
    nums = iter("一二三四五六七八九十")  # 各部分按实际出现的顺序连续编号
    out.append(f"## {next(nums)}、识别结果（请核对）")
    out += [f"{i}. {s}" for i, s in enumerate(report.get("summary", []), start=1)] or ["\uff08无\uff09"]
    out.append("")
    out.append(f"## {next(nums)}、页面设置")
    out += [f"- {s}" for s in report["page"]] or ["- 原设置已符合\uff0c未改动"]
    out.append("")
    out.append(f"## {next(nums)}、段落与文字")
    out.append("| 角色 | 段数 | 段落格式 | 字体 | 字号 | 颜色 |")
    out.append("|---|---|---|---|---|---|")
    for role, b in report["roles"].items():
        out.append(f"| {role} | {b.get('段数', 0)} | {b.get('段落格式', 0)} | {b.get('字体', 0)} | {b.get('字号', 0)} | {b.get('颜色', 0)} |")
    out.append("")
    out.append("表中数字为实际发生改动的段数\u3002")
    out.append("")
    if report.get("notes"):
        out.append(f"## {next(nums)}、其他处理")
        out += [f"- {n}" for n in report["notes"]]
        out.append("")
    if report.get("untouched") or report.get("warnings"):
        out.append(f"## {next(nums)}、未处理和提示")
        out += [f"- {n}" for n in report.get("untouched", []) + report.get("warnings", [])]
        out.append("")
    chk = report.get("check")
    if report.get("keep_text") and chk:
        text_ids = [it for it in chk["issues"] if it["id"] in TEXT_ISSUE_IDS]
        out.append(f"## {next(nums)}、文字类问题（按 --keep-text 未改动，需人工处理）")
        out += [f"- [{it['level']}] {it['message']}（{('第 %d 段' % it['line']) if it.get('line') else '全文'}："
                f"{it['excerpt']}）建议：{it['suggestion']}" for it in text_ids] or ["- 未发现"]
        out.append("")
    if chk:
        s = chk["summary"]
        out.append(f"## {next(nums)}、修复后检查")
        out.append(f"P0 {s['P0']} 项，P1 {s['P1']} 项，P2 {s['P2']} 项。" + ("可以交付\u3002" if s["P0"] == 0 else "仍有 P0\uff0c需人工处理\u3002"))
        for it in chk["issues"]:
            loc = f"第 {it['line']} 段" if it.get("line") else "全文"
            out.append(f"- [{it['level']}] {it['message']}（{loc}：{it['excerpt']}）建议：{it['suggestion']}")
    return "\n".join(out) + "\n"


def main(argv: Optional[Sequence[str]] = None) -> int:
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    parser = argparse.ArgumentParser(description="按 GB/T 9704-2012 修复已有 docx 的格式\uff0c输出新文件和修复报告\u3002",
                                     epilog="示例\uff1apython scripts/fix_docx.py 旧稿.docx -o 旧稿-修复.docx")
    parser.add_argument("file", help="待修复的 .docx\uff08不会被修改\uff09")
    parser.add_argument("-o", "--output", help="输出路径\uff1b默认\u201c原文件名-修复.docx\u201d\uff0c已存在时加 -v2\u3001-v3")
    parser.add_argument("--report", help="修复报告路径\uff1b默认与输出文件同名加\u201c-修复报告.md\u201d")
    parser.add_argument("--layout", help="版式\uff1a红头\u3001信函\u3001纪要\u3001命令\u3001白头\u3001材料\u3001简报\uff1b不写时自动判断")
    parser.add_argument("--keep-text", action="store_true",
                        help="不改任何文字\uff1a用户要求\u201c内容别动\u201d\u201c一个字都别改\u201d时用\u3002只调页面\u3001字体字号\u3001行距\u3001缩进\u3001页码\uff0c"
                             "段首空格\u3001标点\u3001文字问题只在报告中列出\uff1b保存前核对全部文字与原文件一致")
    parser.add_argument("--seal", choices=["是", "否"], help="署名是否按加盖印章排列\uff1b不写时法定公文按加盖处理")
    parser.add_argument("--json", action="store_true", help="以 JSON 输出报告")
    parser.add_argument("--allow", metavar="DOCX-050", help="修复后检查时按本单位制度保留的检查项\uff08逗号分隔\uff09\uff0c命中降为 P2")
    parser.add_argument("--unit-fonts", metavar="角色=字体名,…", help=F.UNIT_FONTS_HELP)
    args = parser.parse_args(argv)
    try:
        import docx  # noqa: F401
    except ImportError:
        print("缺少 python-docx\u3002请先安装\uff1apip install -U python-docx\uff08需要 1.2.0 及以上版本\uff09")
        return 2
    import gw_docx as D
    from check_docx import LAYOUT_NAMES, check_document

    src = Path(args.file)
    if not src.is_file() or src.suffix.lower() != ".docx":
        print(f"错误：找不到 docx 文件 {src}")
        return 2
    if args.layout and args.layout not in LAYOUT_NAMES:
        print(f"错误：无法识别的版式：{args.layout}。可选：{'\u3001'.join(LAYOUT_NAMES)}")
        return 2
    out = Path(args.output) if args.output else src.with_name(src.stem + "-修复.docx")
    if out.is_dir():  # -o 给的是已有文件夹：在文件夹内按默认名输出
        out = out / (src.stem + "-修复.docx")
    if out.resolve() == src.resolve():
        out = src.with_name(src.stem + "-修复.docx")
    final = D.unique_path(out)
    seal = None if args.seal is None else args.seal == "是"
    try:
        unit_fonts = F.parse_unit_fonts(args.unit_fonts, F.load_params())
    except ValueError as exc:
        print(f"错误：{exc}")
        return 2
    try:
        report = fix_document(src, final, layout=args.layout, seal=seal, unit_fonts=unit_fonts,
                              keep_text=args.keep_text)
    except Exception as exc:
        print(f"错误：无法处理 {src}：{exc}")
        return 2
    report["check"] = check_document(final, layout=report["layout"], allow=args.allow, unit_fonts=unit_fonts)
    rpath = Path(args.report) if args.report else final.with_name(final.stem + "-修复报告.md")
    rpath = D.unique_path(rpath)
    rpath.parent.mkdir(parents=True, exist_ok=True)  # --report 指定的文件夹不存在时先建好
    rpath.write_text(render_report(report), encoding="utf-8")
    report["report"] = str(rpath)
    if args.json:
        print(json.dumps(report, ensure_ascii=False, indent=2))
    else:
        s = report["check"]["summary"]
        notice = F.render_notice(report)  # 缺字体或无法检测字体时，醒目提示放在最前面
        if notice:
            print(notice + "\n")
        print(f"已输出：{final}\n修复报告：{rpath}\n修复后检查：P0 {s['P0']} 项，P1 {s['P1']} 项，P2 {s['P2']} 项")
    return 0


if __name__ == "__main__":
    sys.exit(main())
