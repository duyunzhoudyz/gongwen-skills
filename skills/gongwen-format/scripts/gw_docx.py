# -*- coding: utf-8 -*-
"""公文 docx 公共函数：参数读取、页面设置、段落与字体格式、页码、版记表格、属性回读。

build_docx.py、check_docx.py、fix_docx.py、annotate_docx.py 共用。依赖 python-docx（1.2.0 及以上）。
"""
from __future__ import annotations

import copy
import json
import re
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Callable, Dict, Iterator, List, Optional, Sequence, Tuple

from docx import Document  # noqa: F401  由调用方捕获 ImportError
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SCRIPT_DIR = Path(__file__).resolve().parent
PARAMS_PATH = SCRIPT_DIR.parent / "assets" / "gbt9704_params.json"
INSTALL_HINT = "缺少 python-docx。请先安装：pip install -U python-docx（需要 1.2.0 及以上版本）"


def setup_console() -> None:
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8")
        except Exception:
            pass


# ---------------------------------------------------------------------------
# 参数与单位
# ---------------------------------------------------------------------------

def load_params(path: Optional[Path] = None) -> dict:
    with open(path or PARAMS_PATH, encoding="utf-8") as fh:
        return json.load(fh)


def pv(params: dict, dotted: str):
    """读取 {"value": x, "basis": ...} 形式的参数值。"""
    node = params
    for key in dotted.split("."):
        node = node[key]
    if isinstance(node, dict) and "value" in node:
        return node["value"]
    return node


def mm_to_twips(mm: float) -> int:
    return int(round(mm * 1440 / 25.4))


def pt_to_twips(pt: float) -> int:
    return int(round(pt * 20))


def twips_to_mm(tw: float) -> float:
    return tw * 25.4 / 1440


def size_pt(params: dict, name) -> float:
    if isinstance(name, (int, float)):
        return float(name)
    return float(pv(params, f"sizes_pt.{name}"))


def color_hex(params: dict, name: Optional[str]) -> Optional[str]:
    if not name:
        return None
    return pv(params, f"colors.{name}")


def unique_path(path: Path) -> Path:
    """输出路径已存在时依次加 -v2、-v3，不覆盖已有文件。"""
    if not path.exists():
        return path
    stem, suffix = path.stem, path.suffix
    m = re.match(r"^(.*)-v(\d+)$", stem)
    base, start = (m.group(1), int(m.group(2)) + 1) if m else (stem, 2)
    k = start
    while True:
        cand = path.with_name(f"{base}-v{k}{suffix}")
        if not cand.exists():
            return cand
        k += 1


# ---------------------------------------------------------------------------
# 按 schema 顺序插入子元素
# ---------------------------------------------------------------------------

PPR_ORDER = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr",
             "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap",
             "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd",
             "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc",
             "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr",
             "sectPr", "pPrChange"]
RPR_ORDER = ["rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "dstrike", "outline",
             "shadow", "emboss", "imprint", "noProof", "snapToGrid", "vanish", "webHidden", "color", "spacing",
             "w", "kern", "position", "sz", "szCs", "highlight", "u", "effect", "bdr", "shd", "fitText",
             "vertAlign", "rtl", "cs", "em", "lang", "eastAsianLayout", "specVanish", "oMath"]
SECT_ORDER = ["headerReference", "footerReference", "footnotePr", "endnotePr", "type", "pgSz", "pgMar",
              "paperSrc", "pgBorders", "lnNumType", "pgNumType", "cols", "formProt", "vAlign", "noEndnote",
              "titlePg", "textDirection", "bidi", "rtlGutter", "docGrid", "printerSettings", "sectPrChange"]
TBLPR_ORDER = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize",
               "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar",
               "tblLook", "tblCaption", "tblDescription", "tblPrChange"]
BORDER_ORDER = ["top", "left", "start", "bottom", "right", "end", "insideH", "insideV", "between", "bar"]


def _local(tag: str) -> str:
    return tag.split("}")[-1] if "}" in tag else tag


def get_child(parent, tag: str):
    return parent.find(qn(f"w:{tag}"))


def get_or_add(parent, tag: str, order: Sequence[str]):
    """取得子元素，没有时按 schema 顺序插入。"""
    found = parent.find(qn(f"w:{tag}"))
    if found is not None:
        return found
    el = OxmlElement(f"w:{tag}")
    idx = order.index(tag) if tag in order else len(order)
    for child in parent:
        name = _local(child.tag)
        if name in order and order.index(name) > idx:
            child.addprevious(el)
            return el
    parent.append(el)
    return el


def remove_child(parent, tag: str) -> bool:
    found = parent.find(qn(f"w:{tag}"))
    if found is not None:
        parent.remove(found)
        return True
    return False


def set_attrs(el, attrs: Dict[str, Optional[str]], clear: Sequence[str] = ()) -> bool:
    """设置属性，值为 None 的删除；返回是否有变化。"""
    changed = False
    for name in clear:
        key = qn(f"w:{name}")
        if key in el.attrib and name not in attrs:
            del el.attrib[key]
            changed = True
    for name, value in attrs.items():
        key = qn(f"w:{name}")
        if value is None:
            if key in el.attrib:
                del el.attrib[key]
                changed = True
        elif el.get(key) != str(value):
            el.set(key, str(value))
            changed = True
    return changed


# ---------------------------------------------------------------------------
# 字体与字号
# ---------------------------------------------------------------------------

@dataclass
class Ctx:
    """排版上下文：参数。字体严格模式：写入文档的字体一律是标准字体名，没有可替换字体的入口。"""
    params: dict

    def font(self, role: str) -> str:
        """字体角色（仿宋、楷体、黑体、小标宋、宋体）对应的标准字体名。"""
        return self.params["fonts"]["standard"].get(role, role)

    def el(self, name: str) -> dict:
        return self.params["elements"][name]

    def size(self, name) -> float:
        return size_pt(self.params, name)

    @property
    def line_pt(self) -> float:
        return float(pv(self.params, "grid.line_spacing_pt"))

    @property
    def text_width_twips(self) -> int:
        p = self.params
        return (mm_to_twips(pv(p, "page.width_mm")) - mm_to_twips(pv(p, "page.margin_left_mm"))
                - mm_to_twips(pv(p, "page.margin_right_mm")))


def set_rpr_font(rpr, name: str) -> bool:
    rfonts = get_or_add(rpr, "rFonts", RPR_ORDER)
    attrs = {"ascii": name, "hAnsi": name, "eastAsia": name, "cs": name, "hint": "eastAsia",
             "asciiTheme": None, "hAnsiTheme": None, "eastAsiaTheme": None, "cstheme": None}
    return set_attrs(rfonts, attrs)


def set_rpr(rpr, *, font: Optional[str] = None, size: Optional[float] = None, bold: Optional[bool] = None,
            color: Optional[str] = None, black: bool = False) -> bool:
    """设置 rPr；返回是否有变化。black=True 时把非黑色文字改为黑色。"""
    changed = False
    if font:
        changed |= set_rpr_font(rpr, font)
    if size is not None:
        half = str(int(round(size * 2)))
        for tag in ("sz", "szCs"):
            changed |= set_attrs(get_or_add(rpr, tag, RPR_ORDER), {"val": half})
    if bold is not None:
        for tag in ("b", "bCs"):
            if bold:
                el = get_or_add(rpr, tag, RPR_ORDER)
                changed |= set_attrs(el, {"val": None})
            else:
                el = get_child(rpr, tag)
                if el is not None and el.get(qn("w:val")) not in ("0", "false"):
                    rpr.remove(el)
                    changed = True
    if color:
        changed |= set_attrs(get_or_add(rpr, "color", RPR_ORDER), {"val": color, "themeColor": None,
                                                                     "themeShade": None, "themeTint": None})
    elif black:
        el = get_child(rpr, "color")
        if el is not None and (el.get(qn("w:val")) or "auto").upper() not in ("AUTO", "000000"):
            rpr.remove(el)
            changed = True
    return changed


def format_run(run, *, font: Optional[str] = None, size: Optional[float] = None, bold: Optional[bool] = None,
               color: Optional[str] = None, black: bool = False) -> bool:
    rpr = run._r.get_or_add_rPr()
    return set_rpr(rpr, font=font, size=size, bold=bold, color=color, black=black)


def add_run(paragraph, text: str, ctx: Ctx, font_role: str, size_name, *, color: Optional[str] = None,
            bold: bool = False):
    run = paragraph.add_run(text)
    format_run(run, font=ctx.font(font_role), size=ctx.size(size_name), bold=bold or None,
               color=color_hex(ctx.params, color) if color else None)
    return run


# ---------------------------------------------------------------------------
# 段落格式
# ---------------------------------------------------------------------------

JC_MAP = {"left": "left", "center": "center", "right": "right", "justify": "both", "both": "both"}


def set_para(paragraph, *, align: Optional[str] = None, char_pt: float = 16.0,
             first_line: Optional[float] = None, hanging: Optional[float] = None,
             left: Optional[float] = None, right: Optional[float] = None,
             line_exact: Optional[float] = None, line_single: bool = False,
             before_pt: Optional[float] = None, after_pt: Optional[float] = None,
             outline: Optional[int] = None, snap: Optional[bool] = None,
             reset_indent: bool = True, chars: bool = True,
             keep_next: Optional[bool] = None, keep_lines: Optional[bool] = None,
             page_break_before: Optional[bool] = None) -> bool:
    """设置段落格式（缩进单位为字）。返回是否有变化。

    chars=True 时同时写 *Chars（按字符网格）和 twips 两套属性；字号与正文不同的段落
    （版记、页码）传 chars=False，只写 twips，避免 Word 按正文字宽换算。
    """
    ppr = paragraph._p.get_or_add_pPr()
    changed = False
    if align:
        changed |= set_attrs(get_or_add(ppr, "jc", PPR_ORDER), {"val": JC_MAP[align]})
    if reset_indent or any(v is not None for v in (first_line, hanging, left, right)):
        ind = get_or_add(ppr, "ind", PPR_ORDER)
        unit = char_pt * 20
        attrs: Dict[str, Optional[str]] = {
            "start": None, "end": None, "startChars": None, "endChars": None,
            "left": str(int(round((left or 0) * unit))), "leftChars": str(int(round((left or 0) * 100))),
            "right": str(int(round((right or 0) * unit))), "rightChars": str(int(round((right or 0) * 100))),
            "firstLine": None, "firstLineChars": None, "hanging": None, "hangingChars": None,
        }
        if first_line:
            attrs["firstLine"] = str(int(round(first_line * unit)))
            attrs["firstLineChars"] = str(int(round(first_line * 100)))
        elif hanging:
            attrs["hanging"] = str(int(round(hanging * unit)))
            attrs["hangingChars"] = str(int(round(hanging * 100)))
        else:
            attrs["firstLine"] = "0"
            attrs["firstLineChars"] = "0"
        if not left:
            attrs["leftChars"] = "0"
        if not right:
            attrs["rightChars"] = "0"
        if not chars or hanging:
            # 悬挂缩进只写 twips：部分软件对 hangingChars 的换算与 Word 不一致
            for key in ("leftChars", "rightChars", "firstLineChars", "hangingChars"):
                attrs[key] = None
        changed |= set_attrs(ind, attrs)
    if line_exact is not None or line_single or before_pt is not None or after_pt is not None:
        sp = get_or_add(ppr, "spacing", PPR_ORDER)
        attrs = {"beforeLines": None, "afterLines": None, "beforeAutospacing": None, "afterAutospacing": None}
        if line_exact is not None:
            attrs.update({"line": str(pt_to_twips(line_exact)), "lineRule": "exact"})
        elif line_single:
            attrs.update({"line": "240", "lineRule": "auto"})
        attrs["before"] = str(pt_to_twips(before_pt or 0))
        attrs["after"] = str(pt_to_twips(after_pt or 0))
        changed |= set_attrs(sp, attrs)
    if outline is not None:
        changed |= set_attrs(get_or_add(ppr, "outlineLvl", PPR_ORDER), {"val": str(outline)})
    for tag, flag in (("keepNext", keep_next), ("keepLines", keep_lines), ("pageBreakBefore", page_break_before)):
        if flag is None:
            continue
        el = get_child(ppr, tag)
        if flag and (el is None or el.get(qn("w:val")) in ("0", "false")):
            if el is not None:
                ppr.remove(el)
            get_or_add(ppr, tag, PPR_ORDER)
            changed = True
        elif not flag and el is not None:
            ppr.remove(el)
            changed = True
    if snap is not None:
        if snap:
            el = get_child(ppr, "snapToGrid")
            if el is not None:
                ppr.remove(el)
                changed = True
        else:
            changed |= set_attrs(get_or_add(ppr, "snapToGrid", PPR_ORDER), {"val": "0"})
    return changed


def set_border(paragraph, side: str, *, val: str = "single", sz_eighths: int = 8, color: str = "000000",
               space_pt: int = 0) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    bdr = get_or_add(ppr, "pBdr", PPR_ORDER)
    el = get_or_add(bdr, side, BORDER_ORDER)
    set_attrs(el, {"val": val, "sz": str(sz_eighths), "space": str(space_pt), "color": color})


def set_tabs(paragraph, stops: Sequence[Tuple[int, str]]) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    old = get_child(ppr, "tabs")
    if old is not None:
        ppr.remove(old)
    tabs = get_or_add(ppr, "tabs", PPR_ORDER)
    for pos, kind in stops:
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), kind)
        tab.set(qn("w:pos"), str(pos))
        tabs.append(tab)


def add_break(run) -> None:
    run._r.append(OxmlElement("w:br"))


# ---------------------------------------------------------------------------
# 文档级设置
# ---------------------------------------------------------------------------

def setup_document(doc, ctx: Ctx) -> None:
    """统一文档默认字体、Normal 样式、奇偶页设置。"""
    body_font = ctx.font("仿宋")
    styles_el = doc.styles.element
    defaults = styles_el.find(qn("w:docDefaults"))
    if defaults is None:
        defaults = OxmlElement("w:docDefaults")
        styles_el.insert(0, defaults)
    rpr_default = defaults.find(qn("w:rPrDefault"))
    if rpr_default is None:
        rpr_default = OxmlElement("w:rPrDefault")
        defaults.insert(0, rpr_default)
    rpr = rpr_default.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        rpr_default.append(rpr)
    set_rpr(rpr, font=body_font, size=ctx.size("三号"))
    set_attrs(get_or_add(rpr, "lang", RPR_ORDER), {"val": "en-US", "eastAsia": "zh-CN", "bidi": "ar-SA"})

    normal = doc.styles["Normal"]
    nrpr = normal.element.get_or_add_rPr()
    set_rpr(nrpr, font=body_font, size=ctx.size("三号"))
    nppr = normal.element.get_or_add_pPr()
    set_attrs(get_or_add(nppr, "widowControl", PPR_ORDER), {"val": "0"})
    set_attrs(get_or_add(nppr, "autoSpaceDE", PPR_ORDER), {"val": "0"})
    set_attrs(get_or_add(nppr, "autoSpaceDN", PPR_ORDER), {"val": "0"})
    sp = get_or_add(nppr, "spacing", PPR_ORDER)
    set_attrs(sp, {"line": str(pt_to_twips(ctx.line_pt)), "lineRule": "exact", "before": "0", "after": "0"})
    set_attrs(get_or_add(nppr, "jc", PPR_ORDER), {"val": "both"})
    doc.settings.odd_and_even_pages_header_footer = True


def setup_section(section, ctx: Ctx, *, letter: bool = False) -> None:
    p = ctx.params
    sect = section._sectPr
    pgsz = get_or_add(sect, "pgSz", SECT_ORDER)
    set_attrs(pgsz, {"w": str(mm_to_twips(pv(p, "page.width_mm"))), "h": str(mm_to_twips(pv(p, "page.height_mm"))),
                     "orient": None, "code": "9"})
    pgmar = get_or_add(sect, "pgMar", SECT_ORDER)
    header_mm = pv(p, "elements.letter.issuer_top_mm") if letter else pv(p, "page.header_distance_mm")
    set_attrs(pgmar, {
        "top": str(mm_to_twips(pv(p, "page.margin_top_mm"))),
        "bottom": str(mm_to_twips(pv(p, "page.margin_bottom_mm"))),
        "left": str(mm_to_twips(pv(p, "page.margin_left_mm"))),
        "right": str(mm_to_twips(pv(p, "page.margin_right_mm"))),
        "header": str(mm_to_twips(header_mm)),
        "footer": str(mm_to_twips(pv(p, "page.footer_distance_mm"))),
        "gutter": "0",
    })
    grid = get_or_add(sect, "docGrid", SECT_ORDER)
    set_attrs(grid, {"type": pv(p, "grid.doc_grid_type"), "linePitch": str(pv(p, "grid.doc_grid_line_pitch_twips")),
                     "charSpace": str(pv(p, "grid.doc_grid_char_space"))})


def clear_part(part_obj) -> None:
    """清空页眉或页脚内容，只留一个空段落。"""
    el = part_obj._element
    paras = el.findall(qn("w:p"))
    for child in list(el):
        if _local(child.tag) in ("p", "tbl", "sdt") and (not paras or child is not paras[0]):
            el.remove(child)
    if paras:
        p0 = paras[0]
        for child in list(p0):
            if _local(child.tag) != "pPr":
                p0.remove(child)


def add_page_field(paragraph, ctx: Ctx) -> None:
    """写入“— PAGE —”，4 号宋体。"""
    pn = ctx.el("page_number")
    font = ctx.font(pn["font"])
    size = ctx.size(pn["size"])
    dash = pn["dash"]

    def mk(text=None, fld=None, instr=None):
        run = paragraph.add_run()
        format_run(run, font=font, size=size)
        if fld:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), fld)
            run._r.append(el)
        elif instr:
            el = OxmlElement("w:instrText")
            el.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            el.text = instr
            run._r.append(el)
        else:
            run.text = text
        return run

    mk(f"{dash} ")
    mk(fld="begin")
    mk(instr=" PAGE ")
    mk(fld="separate")
    mk("1")
    mk(fld="end")
    mk(f" {dash}")


def setup_page_numbers(section, ctx: Ctx, *, first_page_blank: bool = False) -> None:
    """单页页码居右空一字，双页居左空一字；first_page_blank 时首页不显示页码。"""
    pn = ctx.el("page_number")
    char_pt = ctx.size(pn["size"])
    after = float(pv(ctx.params, "page.page_number_space_after_pt"))
    side = pn["side_chars"]
    for footer, align in ((section.footer, "right"), (section.even_page_footer, "left")):
        footer.is_linked_to_previous = False
        clear_part(footer)
        para = footer.paragraphs[0]
        set_para(para, align=align, char_pt=char_pt, left=side if align == "left" else None,
                 right=side if align == "right" else None, line_single=True, before_pt=0, after_pt=after,
                 snap=False, chars=False)
        add_page_field(para, ctx)
    if first_page_blank:
        section.different_first_page_header_footer = True
        first = section.first_page_footer
        first.is_linked_to_previous = False
        clear_part(first)


# ---------------------------------------------------------------------------
# 版记表格
# ---------------------------------------------------------------------------

def _border_el(tag: str, val: str, sz: int = 0, color: str = "000000"):
    el = OxmlElement(f"w:{tag}")
    el.set(qn("w:val"), val)
    if val != "nil":
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
    return el


def add_banji_table(doc, ctx: Ctx, rows: List[List[Tuple[str, str]]], *, with_lines=True):
    """在文末加版记表格，浮动并对齐版心下边缘。

    rows：每行是若干段落 (类型, 文本)，类型为 cc（抄送）、main（版记主送）、print（印发）。
    with_lines：True 为首末粗线加行间细线，False 无线，"top" 只在表上方加一条粗线（简报报尾）。
    """
    bj = ctx.el("banji")
    char_pt = ctx.size(bj["size"])
    side = bj["side_chars"]
    width = ctx.text_width_twips
    thick = int(round(bj["thick_line_pt"] * 8))
    thin = int(round(bj["thin_line_pt"] * 8))
    table = doc.add_table(rows=len(rows), cols=1)
    tbl = table._tbl
    tblpr = tbl.tblPr
    for child in list(tblpr):
        if _local(child.tag) in ("tblStyle", "tblW", "tblLook", "tblBorders", "jc", "tblLayout", "tblCellMar", "tblInd"):
            tblpr.remove(child)
    tblppr = get_or_add(tblpr, "tblpPr", TBLPR_ORDER)
    set_attrs(tblppr, {"leftFromText": "0", "rightFromText": "0", "topFromText": "0", "bottomFromText": "0",
                       "vertAnchor": "margin", "horzAnchor": "margin", "tblpXSpec": "center", "tblpYSpec": "bottom"})
    set_attrs(get_or_add(tblpr, "tblOverlap", TBLPR_ORDER), {"val": "never"})
    set_attrs(get_or_add(tblpr, "tblW", TBLPR_ORDER), {"w": str(width), "type": "dxa"})
    set_attrs(get_or_add(tblpr, "jc", TBLPR_ORDER), {"val": "center"})
    borders = get_or_add(tblpr, "tblBorders", TBLPR_ORDER)
    for tag in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if with_lines == "top" and tag == "top":
            borders.append(_border_el(tag, "single", thick))
        elif with_lines is True and tag in ("top", "bottom"):
            borders.append(_border_el(tag, "single", thick))
        elif with_lines is True and tag == "insideH":
            borders.append(_border_el(tag, "single", thin))
        else:
            borders.append(_border_el(tag, "nil"))
    set_attrs(get_or_add(tblpr, "tblLayout", TBLPR_ORDER), {"type": "fixed"})
    mar = get_or_add(tblpr, "tblCellMar", TBLPR_ORDER)
    for tag in ("left", "right"):
        el = OxmlElement(f"w:{tag}")
        el.set(qn("w:w"), "0")
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    grid = tbl.tblGrid
    for col in grid.findall(qn("w:gridCol")):
        col.set(qn("w:w"), str(width))
    font = ctx.font(bj["font"])
    for r_idx, row_items in enumerate(rows):
        cell = table.rows[r_idx].cells[0]
        tcpr = cell._tc.get_or_add_tcPr()
        tcw = tcpr.find(qn("w:tcW"))
        if tcw is None:
            tcw = OxmlElement("w:tcW")
            tcpr.insert(0, tcw)
        tcw.set(qn("w:w"), str(width))
        tcw.set(qn("w:type"), "dxa")
        first = True
        for kind, text in row_items:
            para = cell.paragraphs[0] if first else cell.add_paragraph()
            first = False
            if kind in ("cc", "main"):
                m_label = re.match(r"^[^：:]{1,4}[：:]", text)
                label_w = _text_width(m_label.group(0)) if m_label else 3  # “抄送：”三个字，“报：”两个字
                set_para(para, align="left", char_pt=char_pt, left=side + label_w, hanging=label_w, right=side,
                         line_exact=ctx.line_pt, before_pt=0, after_pt=0, snap=False, chars=False)
                run = para.add_run(text)
            else:
                set_para(para, align="left", char_pt=char_pt, left=side, right=None, line_exact=ctx.line_pt,
                         before_pt=0, after_pt=0, snap=False, chars=False)
                set_tabs(para, [(width - int(round(side * char_pt * 20)), "right")])
                run = para.add_run(text)
            format_run(run, font=font, size=char_pt)
    # 浮动表格后面需要一个段落作锚点，行高压到最小，避免多出空白页
    anchor = doc.add_paragraph()
    set_para(anchor, line_exact=1, before_pt=0, after_pt=0, snap=False)
    return table


# ---------------------------------------------------------------------------
# 正文表格（Markdown 管道表格转换而来）
# ---------------------------------------------------------------------------

NUMERIC_CELL_RE = re.compile(r"^[-+±]?[\d.,，%％‰～~—\-/（）()万亿元个人次项件家所座台]*\d[\d.,，%％‰～~—\-/（）()万亿元个人次项件家所座台]*$")


def _text_width(text: str) -> float:
    """按字估算宽度，与 gw_layout.text_width_chars 口径一致（此处不依赖 gw_layout）。"""
    full = "—…“”‘’"
    return sum(0.5 if ord(ch) < 0x2E80 and ch not in full else 1.0 for ch in text)


def column_widths_twips(header: Sequence[str], rows: Sequence[Sequence[str]], total: int) -> List[int]:
    """按各列最长文字分配列宽，合计等于 total（版心宽）。每列至少 2 字宽，长列封顶，避免挤压短列。"""
    cols = len(header)
    weights = []
    for c in range(cols):
        texts = [header[c]] + [r[c] for r in rows if c < len(r)]
        longest = max((_text_width(line) for t in texts for line in (t or "").split("\n")), default=0.0)
        weights.append(min(max(longest, 2.0), 16.0))
    s = sum(weights) or 1.0
    widths = [int(total * w / s) for w in weights]
    widths[-1] += total - sum(widths)
    return widths


def add_body_table(doc, ctx: Ctx, header: Sequence[str], rows: Sequence[Sequence[str]],
                   aligns: Sequence[Optional[str]], *, font_role: str = "仿宋", size_name="四号"):
    """在文末加一张正文表格：与版心同宽、居中，全框线，表头加粗、居中并在跨页时重复，单元格垂直居中。

    单元格段落不对齐文档网格、行距为字号的 1.5 倍、无缩进；数字格居中，文字格居左，分隔行写了对齐方式的从其规定。
    """
    tb = ctx.el("table")
    width = ctx.text_width_twips
    cols = len(header)
    widths = column_widths_twips(header, rows, width)
    table = doc.add_table(rows=1 + len(rows), cols=cols)
    tbl = table._tbl
    tblpr = tbl.tblPr
    for child in list(tblpr):
        if _local(child.tag) in ("tblStyle", "tblW", "tblLook", "tblBorders", "jc", "tblLayout", "tblInd"):
            tblpr.remove(child)
    set_attrs(get_or_add(tblpr, "tblW", TBLPR_ORDER), {"w": str(width), "type": "dxa"})
    set_attrs(get_or_add(tblpr, "jc", TBLPR_ORDER), {"val": "center"})
    borders = get_or_add(tblpr, "tblBorders", TBLPR_ORDER)
    sz = int(round(float(tb["border_pt"]) * 8))
    for tag in ("top", "left", "bottom", "right", "insideH", "insideV"):
        borders.append(_border_el(tag, "single", sz))
    set_attrs(get_or_add(tblpr, "tblLayout", TBLPR_ORDER), {"type": "fixed"})
    for col, w in zip(tbl.tblGrid.findall(qn("w:gridCol")), widths):
        col.set(qn("w:w"), str(w))
    font = ctx.font(font_role)
    size = ctx.size(size_name)
    for r_idx, values in enumerate([list(header)] + [list(r) for r in rows]):
        row = table.rows[r_idx]
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))  # 一行不跨页拆开
        if r_idx == 0:
            trpr.append(OxmlElement("w:tblHeader"))  # 表格跨页时重复表头
        for c_idx, cell in enumerate(row.cells):
            tcpr = cell._tc.get_or_add_tcPr()
            tcw = tcpr.find(qn("w:tcW"))
            if tcw is None:
                tcw = OxmlElement("w:tcW")
                tcpr.insert(0, tcw)
            tcw.set(qn("w:w"), str(widths[c_idx]))
            tcw.set(qn("w:type"), "dxa")
            valign = OxmlElement("w:vAlign")
            valign.set(qn("w:val"), "center")
            tcpr.append(valign)
            text = values[c_idx] if c_idx < len(values) else ""
            if r_idx == 0:
                align = "center"
            elif aligns[c_idx] if c_idx < len(aligns) else None:
                align = aligns[c_idx]
            else:
                align = "center" if NUMERIC_CELL_RE.match(text.replace(" ", "")) else "left"
            para = cell.paragraphs[0]
            set_para(para, align=align, char_pt=size, line_exact=round(size * 1.5, 1), before_pt=0, after_pt=0,
                     snap=False, chars=False)
            lines = text.split("\n") if text else [""]
            for k, line in enumerate(lines):
                run = para.add_run(line)
                format_run(run, font=font, size=size, bold=True if (r_idx == 0 and tb["header_bold"]) else None)
                if k < len(lines) - 1:
                    add_break(run)
    return table


# 本工具生成的版头、落款表格用 tblCaption 作标记，回读时据此识别，不当作正文表格或版记
JOINT_HEAD_MARK = "gw:联合行文版头"
JOINT_SEAL_MARK = "gw:联署署名"


def table_mark(tbl_el) -> Optional[str]:
    tblpr = tbl_el.find(qn("w:tblPr"))
    cap = tblpr.find(qn("w:tblCaption")) if tblpr is not None else None
    return cap.get(qn("w:val")) if cap is not None else None


def is_frame_table(tbl_el) -> bool:
    """联合行文版头表格、联署署名表格（本工具生成，不是正文表格，也不是版记）。"""
    return table_mark(tbl_el) in (JOINT_HEAD_MARK, JOINT_SEAL_MARK)


def _plain_table(doc, ctx: Ctx, widths: Sequence[int], rows: int, mark: str):
    """无框线、定宽、居中的版面表格，标记为 mark。单元格左右边距为 0。"""
    table = doc.add_table(rows=rows, cols=len(widths))
    tbl = table._tbl
    tblpr = tbl.tblPr
    for child in list(tblpr):
        if _local(child.tag) in ("tblStyle", "tblW", "tblLook", "tblBorders", "jc", "tblLayout", "tblCellMar", "tblInd"):
            tblpr.remove(child)
    set_attrs(get_or_add(tblpr, "tblW", TBLPR_ORDER), {"w": str(sum(widths)), "type": "dxa"})
    set_attrs(get_or_add(tblpr, "jc", TBLPR_ORDER), {"val": "center"})
    borders = get_or_add(tblpr, "tblBorders", TBLPR_ORDER)
    for tag in ("top", "left", "bottom", "right", "insideH", "insideV"):
        borders.append(_border_el(tag, "nil"))
    set_attrs(get_or_add(tblpr, "tblLayout", TBLPR_ORDER), {"type": "fixed"})
    mar = get_or_add(tblpr, "tblCellMar", TBLPR_ORDER)
    for tag in ("left", "right"):
        el = OxmlElement(f"w:{tag}")
        el.set(qn("w:w"), "0")
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    set_attrs(get_or_add(tblpr, "tblCaption", TBLPR_ORDER), {"val": mark})
    for col, w in zip(tbl.tblGrid.findall(qn("w:gridCol")), widths):
        col.set(qn("w:w"), str(w))
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for cell, w in zip(row.cells, widths):
            tcpr = cell._tc.get_or_add_tcPr()
            tcw = tcpr.find(qn("w:tcW"))
            if tcw is None:
                tcw = OxmlElement("w:tcW")
                tcpr.insert(0, tcw)
            tcw.set(qn("w:w"), str(w))
            tcw.set(qn("w:type"), "dxa")
    return table


def joint_header_layout(ctx: Ctx, names: Sequence[str]) -> Tuple[float, int, int]:
    """联合行文版头：返回 (字号磅, 名称列宽 twips, “文件”列宽 twips)。

    名称列宽按最长名称的字数乘字号，“文件”列宽为两个字加半字间距；两列合计不超过版心宽的 fill_ratio。
    """
    el = ctx.el("joint")
    slack = float(el["name_slack"])
    longest = max((_text_width(n) for n in names), default=1.0) * slack
    units = longest + 2 + float(el["wen_gap_ratio"])
    avail = ctx.text_width_twips / 20 * float(el["fill_ratio"])
    cap = float(el["multi_name_size_max_pt"] if len(names) >= 3 else el["name_size_max_pt"])
    size = max(float(el["name_size_min_pt"]), min(cap, avail / max(units, 1.0)))
    size = int(size * 2) / 2
    left = int(round(longest * size * 20))
    right = int(round((2 + float(el["wen_gap_ratio"])) * size * 20))
    return size, left, right


def add_joint_header(doc, ctx: Ctx, names: Sequence[str], *, wen: str = "文件", red: Optional[str] = None):
    """联合行文版头（7.2.4）：联署机关名称自上而下排列（主办机关在前），“文件”二字置于名称右侧，
    并以联署机关名称为准上下居中。用两列无框线表格实现：左列名称，右列“文件”（单元格垂直居中）。"""
    el = ctx.el("joint")
    iss = ctx.el("issuer")
    size, left_w, right_w = joint_header_layout(ctx, names)
    table = _plain_table(doc, ctx, [left_w, right_w], 1, JOINT_HEAD_MARK)
    font = ctx.font(iss["font"])
    line = round(size * float(el["line_ratio"]), 1)
    cell_l, cell_r = table.rows[0].cells
    for k, name in enumerate(names):
        para = cell_l.paragraphs[0] if k == 0 else cell_l.add_paragraph()
        set_para(para, align="center", char_pt=size, line_exact=line, before_pt=0, after_pt=0, snap=False, chars=False)
        format_run(para.add_run(name), font=font, size=size, color=red)
    tcpr = cell_r._tc.get_or_add_tcPr()
    valign = OxmlElement("w:vAlign")
    valign.set(qn("w:val"), "center")
    tcpr.append(valign)
    para = cell_r.paragraphs[0]
    set_para(para, align="center", char_pt=size, left=float(el["wen_gap_ratio"]), line_exact=line, before_pt=0,
             after_pt=0, snap=False, chars=False)
    format_run(para.add_run(wen), font=font, size=size, color=red)
    return table


def joint_seal_columns(ctx: Ctx, count: int) -> int:
    """联署署名每行并排几个：不超过 seal_cols_max，使每格宽不小于印章直径（印章互不相交）。"""
    el = ctx.el("joint")
    return max(1, min(count, int(el["seal_cols_max"])))


def add_joint_signatures(doc, ctx: Ctx, names: Sequence[str]):
    """联合行文加盖印章（7.3.5.1）：署名按机关顺序排列，每个机关名称占一格并居中，印章一一对应下压署名。
    每行最多并排 seal_cols_max 个，每格宽不小于印章直径，印章互不相交，两端不超出版心。"""
    sg = ctx.el("signature_seal")
    k = joint_seal_columns(ctx, len(names))
    rows = -(-len(names) // k)
    col_w = ctx.text_width_twips // k
    table = _plain_table(doc, ctx, [col_w] * k, rows, JOINT_SEAL_MARK)
    font = ctx.font(sg["font"])
    size = ctx.size(sg["size"])
    for idx in range(rows * k):
        cell = table.rows[idx // k].cells[idx % k]
        para = cell.paragraphs[0]
        set_para(para, align="center", char_pt=size, line_exact=ctx.line_pt, before_pt=0, after_pt=0, snap=False,
                 chars=False, keep_next=True)
        if idx < len(names):
            format_run(para.add_run(names[idx]), font=font, size=size)
    return table


def is_banji_table(tbl_el) -> bool:
    """版记表格：浮动定位（本工具生成的版记），或含“抄送”、以“印发”结尾。"""
    if is_frame_table(tbl_el):
        return False
    tblpr = tbl_el.find(qn("w:tblPr"))
    if tblpr is not None and tblpr.find(qn("w:tblpPr")) is not None:
        return True
    text = "".join(t.text or "" for t in tbl_el.iter(qn("w:t"))).strip()
    return "抄送" in text or text.endswith("印发") or bool(re.match(r"^(?:报|送|发)[：:]", text))


def body_tables(doc) -> List[Tuple[int, object]]:
    """正文中的表格（不含版记表格），按文档顺序返回 (表格前一段在 doc.paragraphs 中的下标, w:tbl 元素)。

    表格在最前面时下标为 -1。供 check、fix 识别段落角色时插入表格占位（gw_layout.parse_lines 的 tables_after）。
    """
    out: List[Tuple[int, object]] = []
    count = 0
    for child in doc.element.body.iterchildren():
        tag = _local(child.tag)
        if tag == "p":
            count += 1
        elif tag == "tbl" and not is_banji_table(child) and not is_frame_table(child):
            out.append((count - 1, child))
    return out


def ctx_text_width_pt(doc) -> float:
    """读取第一节的版心宽度（磅）。"""
    sect = doc.sections[0]._sectPr
    pgsz, pgmar = sect.find(qn("w:pgSz")), sect.find(qn("w:pgMar"))
    try:
        w = int(pgsz.get(qn("w:w")))
        left = int(pgmar.get(qn("w:left")) or pgmar.get(qn("w:start")) or 0)
        right = int(pgmar.get(qn("w:right")) or pgmar.get(qn("w:end")) or 0)
        return (w - left - right) / 20
    except (AttributeError, TypeError, ValueError):
        return 442.25


# ---------------------------------------------------------------------------
# 回读：样式链上的有效属性
# ---------------------------------------------------------------------------

def _style_chain(style) -> Iterator:
    seen = 0
    while style is not None and seen < 20:
        yield style
        style = style.base_style
        seen += 1


def _doc_default_rpr(doc):
    d = doc.styles.element.find(qn("w:docDefaults"))
    if d is None:
        return None
    r = d.find(qn("w:rPrDefault"))
    return None if r is None else r.find(qn("w:rPr"))


def _doc_default_ppr(doc):
    d = doc.styles.element.find(qn("w:docDefaults"))
    if d is None:
        return None
    r = d.find(qn("w:pPrDefault"))
    return None if r is None else r.find(qn("w:pPr"))


def _ea_font(rpr) -> Optional[str]:
    if rpr is None:
        return None
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        return None
    theme = rf.get(qn("w:eastAsiaTheme"))
    if theme:
        return f"主题字体（{theme}）"
    return rf.get(qn("w:eastAsia"))


def _size(rpr) -> Optional[float]:
    if rpr is None:
        return None
    sz = rpr.find(qn("w:sz"))
    if sz is None or sz.get(qn("w:val")) is None:
        return None
    return int(sz.get(qn("w:val"))) / 2


def _bold(rpr) -> Optional[bool]:
    if rpr is None:
        return None
    b = rpr.find(qn("w:b"))
    if b is None:
        return None
    return b.get(qn("w:val")) not in ("0", "false")


def _color(rpr) -> Optional[str]:
    if rpr is None:
        return None
    c = rpr.find(qn("w:color"))
    if c is None:
        return None
    return (c.get(qn("w:val")) or "auto").upper()


def effective_run(run, paragraph, doc, getter: Callable):
    value = getter(run._r.rPr)
    if value is not None:
        return value
    try:
        rstyle = run.style
    except Exception:
        rstyle = None
    for st in _style_chain(rstyle):
        value = getter(st.element.rPr)
        if value is not None:
            return value
    for st in _style_chain(paragraph.style):
        value = getter(st.element.rPr)
        if value is not None:
            return value
    return getter(_doc_default_rpr(doc))


def run_font(run, paragraph, doc) -> Optional[str]:
    return effective_run(run, paragraph, doc, _ea_font)


def run_size(run, paragraph, doc) -> Optional[float]:
    return effective_run(run, paragraph, doc, _size)


def run_bold(run, paragraph, doc) -> bool:
    return bool(effective_run(run, paragraph, doc, _bold))


def run_color(run, paragraph, doc) -> str:
    return effective_run(run, paragraph, doc, _color) or "AUTO"


def is_red(color: Optional[str]) -> bool:
    """红色判断：R 分量高、G 和 B 分量低（FF0000、E60000、C00000 等）。"""
    if not color or len(color) != 6:
        return False
    try:
        r, g, b = int(color[0:2], 16), int(color[2:4], 16), int(color[4:6], 16)
    except ValueError:
        return False
    return r >= 0xC0 and g <= 0x40 and b <= 0x40


def para_is_red(paragraph, doc) -> bool:
    runs = [r for r in paragraph.runs if r.text.strip()]
    return bool(runs) and all(is_red(run_color(r, paragraph, doc)) for r in runs)


def effective_ppr(paragraph, doc, getter: Callable):
    value = getter(paragraph._p.pPr)
    if value is not None:
        return value
    for st in _style_chain(paragraph.style):
        value = getter(st.element.pPr)
        if value is not None:
            return value
    return getter(_doc_default_ppr(doc))


def _jc(ppr) -> Optional[str]:
    if ppr is None:
        return None
    jc = ppr.find(qn("w:jc"))
    return None if jc is None else jc.get(qn("w:val"))


def para_align(paragraph, doc) -> str:
    return effective_ppr(paragraph, doc, _jc) or "left"


def para_first_line_chars(paragraph, doc, char_pt: float = 16.0) -> Optional[float]:
    """首行缩进（字）；悬挂缩进返回负数；未设置返回 0。"""
    def getter(ppr):
        if ppr is None:
            return None
        ind = ppr.find(qn("w:ind"))
        if ind is None:
            return None
        for attr, sign in (("firstLineChars", 1), ("hangingChars", -1)):
            v = ind.get(qn(f"w:{attr}"))
            if v is not None and int(v) != 0:
                return sign * int(v) / 100
        for attr, sign in (("firstLine", 1), ("hanging", -1)):
            v = ind.get(qn(f"w:{attr}"))
            if v is not None and int(v) != 0:
                return sign * int(v) / (char_pt * 20)
        if any(ind.get(qn(f"w:{a}")) is not None for a in ("firstLine", "firstLineChars", "hanging", "hangingChars")):
            return 0.0
        return None
    return effective_ppr(paragraph, doc, getter) or 0.0


def para_left_chars(paragraph, doc, char_pt: float = 16.0) -> float:
    def getter(ppr):
        if ppr is None:
            return None
        ind = ppr.find(qn("w:ind"))
        if ind is None:
            return None
        for attr in ("leftChars", "startChars"):
            v = ind.get(qn(f"w:{attr}"))
            if v is not None and int(v) != 0:
                return int(v) / 100
        for attr in ("left", "start"):
            v = ind.get(qn(f"w:{attr}"))
            if v is not None:
                return int(v) / (char_pt * 20)
        return None
    return effective_ppr(paragraph, doc, getter) or 0.0


def para_tab_stops(paragraph, doc) -> List[Tuple[str, int]]:
    """段落直接设置的制表位：[(类型, 位置 twips)]，按位置排序。"""
    ppr = paragraph._p.pPr
    tabs = ppr.find(qn("w:tabs")) if ppr is not None else None
    out: List[Tuple[str, int]] = []
    if tabs is not None:
        for t in tabs.findall(qn("w:tab")):
            try:
                out.append((t.get(qn("w:val")) or "left", int(t.get(qn("w:pos")))))
            except (TypeError, ValueError):
                continue
    return sorted(out, key=lambda x: x[1])


def leading_offset_chars(paragraph, doc, char_pt: float = 16.0) -> float:
    """段落第一个字距版心左边缘的字数：左缩进、首行缩进，加上行首的制表符或空格占位。

    行首是制表符时取其后第一个左对齐制表位（没有设置时按 Word 默认 420 twips 的整数倍）；
    半角空格按 0.5 字、全角空格按 1 字计。
    """
    unit = char_pt * 20
    base = para_left_chars(paragraph, doc, char_pt) + max(0.0, para_first_line_chars(paragraph, doc, char_pt) or 0.0)
    text = paragraph.text
    if text.startswith("\t"):
        here = base * unit
        stops = [pos for kind, pos in para_tab_stops(paragraph, doc) if kind in ("left", "start", "num") and pos > here + 1]
        pos = min(stops) if stops else (int(here // 420) + 1) * 420
        return pos / unit
    return base + sum(1.0 if ch == "　" else 0.5 for ch in text[:len(text) - len(text.lstrip(" 　"))])


def para_line_spacing(paragraph, doc) -> Tuple[Optional[str], Optional[float]]:
    """返回 (lineRule, 数值)；exact/atLeast 时数值为磅，auto 时为倍数。"""
    def getter(ppr):
        if ppr is None:
            return None
        sp = ppr.find(qn("w:spacing"))
        if sp is None or sp.get(qn("w:line")) is None:
            return None
        rule = sp.get(qn("w:lineRule")) or "auto"
        line = int(sp.get(qn("w:line")))
        return (rule, line / 20 if rule in ("exact", "atLeast") else line / 240)
    return effective_ppr(paragraph, doc, getter) or ("auto", 1.0)


def para_snap(paragraph, doc) -> bool:
    def getter(ppr):
        if ppr is None:
            return None
        el = ppr.find(qn("w:snapToGrid"))
        if el is None:
            return None
        return el.get(qn("w:val")) not in ("0", "false")
    v = effective_ppr(paragraph, doc, getter)
    return True if v is None else v


# ---------------------------------------------------------------------------
# run 拆分与文本定位
# ---------------------------------------------------------------------------

SPLITTABLE = {"rPr", "t", "tab", "br", "cr", "noBreakHyphen", "softHyphen"}


def run_splittable(run) -> bool:
    return all(_local(child.tag) in SPLITTABLE for child in run._r)


def split_run(run, offset: int):
    """在 offset 处把 run 拆成两个（复制 rPr），返回后半个 run；不能拆时返回 None。"""
    from docx.text.run import Run
    text = run.text
    if offset <= 0 or offset >= len(text) or not run_splittable(run):
        return None
    new_r = copy.deepcopy(run._r)
    run._r.addnext(new_r)
    new_run = Run(new_r, run._parent)
    run.text = text[:offset]
    new_run.text = text[offset:]
    return new_run


def runs_for_range(paragraph, start: int, end: int):
    """把段落中 [start, end) 区间对应的文字拆成完整的 run，返回 (首 run, 末 run)；失败返回 None。

    只使用段落的直接 run（不含超链接内的 run）。
    """
    if start >= end:
        return None
    # 先拆开始位置
    for boundary in (start, end):
        pos = 0
        for run in list(paragraph.runs):
            length = len(run.text)
            if pos < boundary < pos + length:
                if split_run(run, boundary - pos) is None:
                    return None
                break
            pos += length
    first = last = None
    pos = 0
    for run in paragraph.runs:
        length = len(run.text)
        if length and pos >= start and pos + length <= end:
            if first is None:
                first = run
            last = run
        pos += length
    if first is None:
        return None
    return first, last


def direct_text(paragraph) -> str:
    return "".join(run.text for run in paragraph.runs)


def iter_all_paragraphs(doc) -> Iterator:
    """正文段落和表格中的段落（按文档顺序）。"""
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    seen = set()  # 合并单元格会重复返回同一单元格，按元素去重
    for block in doc.element.body.iterchildren():
        tag = _local(block.tag)
        if tag == "p":
            yield Paragraph(block, doc._body)
        elif tag == "tbl":
            table = Table(block, doc._body)
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        if para._p in seen:
                            continue
                        seen.add(para._p)
                        yield para
