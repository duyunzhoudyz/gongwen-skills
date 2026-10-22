#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把公文正文\uff08.txt/.md\uff09排成符合 GB/T 9704-2012 的 docx\u3002

版式\uff1a红头\uff08文件格式\uff09\u3001信函\u3001纪要\u3001命令\uff08令\uff09\u3001白头\uff08无版头\uff09\u3001材料\uff08只有标题和正文\uff09\u3001简报\uff08报头\u3001各篇\u3001报尾\uff09\u3002
另有\uff1a联合行文版头\uff08发文机关标志写多个机关\uff09\u3001预印红头纸套打\uff08--overprint\uff09\u3001附件另面编排\uff08正文文件中的\u201c【附件1】标题\u201d分界行\uff09\u3002
未指定版式时\uff0c法定公文用白头\uff0c常用材料用材料版式\uff1b版式信息或正文有简报报头的用简报版式\u3002
生成后自动回读检查\uff08check_docx.py\uff09\uff0c报告识别结果\u3001字体情况和检查结论\u3002

退出码\uff1a0 成功\uff1b2 输入错误或缺少依赖\u3002
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Tuple

sys.path.insert(0, str(Path(__file__).resolve().parent))

import gw_fonts as F  # noqa: E402
from gw_verify import norm  # noqa: E402
from gw_layout import (BODY_ROLES, BRIEF_META_KEYS, LINGHAO_RE, TABLE_MARK, UPWARD_TYPES, Item, Parsed,  # noqa: E402
                       break_title, clean_line, doc_kind, extract_attachments, extract_pipe_tables, has_brief_header,
                       heading_levels, heading_split, is_heading, is_table_caption, norm_issue, norm_total, normalize_date,
                       parse_brief, parse_lines, read_body_text, summarize, text_width_chars)

_DUN = "\u3001"  # 顿号；f-string 表达式内不能出现反斜杠（Python < 3.12）

LAYOUTS = {"红头": "红头", "文件": "红头", "文件格式": "红头", "红头文件": "红头",
           "信函": "信函", "信函格式": "信函", "函件": "信函",
           "纪要": "纪要", "纪要格式": "纪要",
           "命令": "命令", "令": "命令", "命令格式": "命令", "令格式": "命令", "命令\uff08令\uff09": "命令",
           "白头": "白头", "白头文件": "白头",
           "材料": "材料", "普通材料": "材料",
           "简报": "简报", "工作简报": "简报", "简报格式": "简报"}
META_KEYS = ("版式", "份号", "密级", "保密期限", "紧急程度", "发文机关标志", "发文字号", "签发人", "标题",
             "主送机关", "附件", "发文机关署名", "成文日期", "印章", "附注", "抄送", "印发机关", "印发日期",
             "三级标题加粗", "表格字体", "表格字号", "unit_fonts", "单位规定字体",
             "令号", "签发人职务", "签发人姓名", "套打", "套打版头高度", "联合行文", "留占位") + BRIEF_META_KEYS
BRIEF_IGNORED = ("发文机关标志", "发文字号", "签发人", "主送机关", "附件", "发文机关署名", "成文日期", "印章", "附注",
                 "印发机关", "份号", "密级", "保密期限", "紧急程度")
TABLE_FONTS = ("仿宋", "宋体")
TABLE_SIZES = ("四号", "小四")
HEAD_ROLES = ("copy_no", "secret", "urgency", "issuer", "docno", "signer", "lingno")
TITLE_ROLES = ("title", "subtitle", "title_note")
TAIL_ROLES = ("attachment", "signature", "date", "note", "cc", "print", "banji_main")
LEGAL_LAYOUTS = ("红头", "白头", "信函")


def _as_list(value) -> List[str]:
    if value is None or value == "":
        return []
    if isinstance(value, (list, tuple)):
        return [str(v).strip() for v in value if str(v).strip()]
    return [s.strip() for s in re.split(r"[\n]", str(value)) if s.strip()]


def _truthy(value) -> Optional[bool]:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return value
    return str(value).strip() in ("是", "加盖", "盖章", "true", "True", "yes", "1")


def _same(a: str, b: str) -> bool:
    return re.sub(r"[\s　]", "", a or "") == re.sub(r"[\s　]", "", b or "")


class BuildError(Exception):
    pass


class Builder:
    """按版式逐段生成文档\u3002"""

    def __init__(self, ctx, layout: str, keep_text: bool = False):
        import gw_docx as D
        from docx import Document
        self.D = D
        self.ctx = ctx
        self.layout = layout
        self.keep_text = keep_text
        self.restored: List[str] = []  # keep_text\uff1a已还原段首空格的原文行\uff0c保存前逐行核对
        self.doc = Document()
        D.setup_document(self.doc, ctx)
        self.section = self.doc.sections[0]
        D.setup_section(self.section, ctx, letter=(layout == "信函"))

    # 基本段落 ------------------------------------------------------------
    def blank(self, n: int = 1) -> None:
        for _ in range(n):
            p = self.doc.add_paragraph()
            self.D.set_para(p, line_exact=self.ctx.line_pt, before_pt=0, after_pt=0)

    def para(self, role: str, text: str, split: int = -1, raw: Optional[str] = None):
        """raw\uff1akeep_text 时传入原文行\uff1b正文类段落保留段首空格\uff08缩进由空格承担\uff0c首行缩进设为 0\uff09\uff0c不改任何文字\u3002"""
        from gw_format import apply_role, offset_leading_spaces
        lead = 0
        if self.keep_text and raw is not None and role in BODY_ROLES + ("salutation", "recipient") and raw != text                 and clean_line(raw) == text:
            lead = raw.index(text)
            if lead:
                self.restored.append(raw)
                text, split = raw, (split + lead if split >= 0 else split)
        p = self.doc.add_paragraph()
        p.add_run(text)
        apply_role(p, role, self.ctx, split=split)
        if lead and role in ("salutation", "recipient"):
            offset_leading_spaces(p)  # 称呼顶格\uff1a空格保留\uff0c用等宽负缩进抵消
        elif lead:
            self.D.set_para(p, reset_indent=True)
        return p

    def title(self, lines: List[str]):
        from gw_format import apply_role
        p = self.doc.add_paragraph()
        for i, line in enumerate(lines):
            run = p.add_run(line)
            if i < len(lines) - 1:
                self.D.add_break(run)
        apply_role(p, "title", self.ctx)
        return p

    # 版头 ----------------------------------------------------------------
    def issuer_size(self, text: str) -> float:
        el = self.ctx.el("issuer")
        width_pt = self.ctx.text_width_twips / 20
        chars = max(1.0, text_width_chars(text))
        size = min(float(el["size_max_pt"]), width_pt * float(el["fill_ratio"]) / chars)
        size = max(float(el["size_min_pt"]), size)
        return int(size * 2) / 2

    def issuer(self, text: str, top_from_body_mm: float, used_lines: int):
        D, ctx = self.D, self.ctx
        el = ctx.el("issuer")
        size = self.issuer_size(text)
        p = self.doc.add_paragraph()
        run = p.add_run(text)
        D.format_run(run, font=ctx.font(el["font"]), size=size, color=D.color_hex(ctx.params, el["color"]))
        before = max(0.0, top_from_body_mm * 72 / 25.4 - used_lines * ctx.line_pt)
        D.set_para(p, align="center", line_exact=round(size * float(el["line_ratio"]), 1), before_pt=round(before, 1),
                   after_pt=0, snap=False)
        return p

    def red_rule_under(self, p) -> None:
        D, ctx = self.D, self.ctx
        rule = ctx.el("header_rule")
        D.set_border(p, "bottom", sz_eighths=int(round(rule["width_pt"] * 8)),
                     color=D.color_hex(ctx.params, rule["color"]), space_pt=int(round(rule["gap_mm"] * 72 / 25.4)))

    def joint_header(self, names: List[str], top_from_body_mm: float, used_lines: int) -> None:
        """联合行文版头\uff087.2.4\uff09\uff1a先留出标志上边缘距版心上边缘的空白\uff0c再排名称与\u201c文件\u201d的表格\u3002"""
        D, ctx = self.D, self.ctx
        before = max(0.0, top_from_body_mm * 72 / 25.4 - used_lines * ctx.line_pt)
        spacer = self.doc.add_paragraph()
        D.set_para(spacer, line_exact=1, before_pt=round(max(before - 1, 0.0), 1), after_pt=0, snap=False)
        D.add_joint_header(self.doc, ctx, names, red=D.color_hex(ctx.params, ctx.el("issuer")["color"]))

    def overprint_block(self, issuer_text: str, joint_names: Optional[List[str]], top_from_body_mm: float,
                        used_lines: int, header_mm: Optional[float]) -> float:
        """预印红头纸套打\uff1a不输出发文机关标志\uff0c留出与正式版头同样高的空白\u3002

        默认高度 = 标志上边缘距版心上边缘\uff087.2.4\uff09＋标志行高\uff08按标志文字估算\uff0c没有文字时按 overprint.issuer_default_size_pt\uff09\uff1b
        版式信息\u201c套打版头高度\u201d\uff08毫米\uff0c发文字号上沿距版心上边缘\uff09给定时\uff0c按预印纸实际尺寸留白\u3002
        返回发文字号前空行之前已占的高度\uff08磅\uff0c自版心上边缘起\uff09\u3002"""
        D, ctx = self.D, self.ctx
        lp = ctx.line_pt
        blank = int(ctx.el("docno")["blank_lines_before"])
        if header_mm:
            total = header_mm * 72 / 25.4
            line = 1.0
            before = max(0.0, total - used_lines * lp - blank * lp - line)
        else:
            iss = ctx.el("issuer")
            if joint_names:
                size, _, _ = D.joint_header_layout(ctx, joint_names)
                line = len(joint_names) * round(size * float(ctx.el("joint")["line_ratio"]), 1)
            else:
                size = self.issuer_size(issuer_text) if issuer_text else float(ctx.el("overprint")["issuer_default_size_pt"])
                line = round(size * float(iss["line_ratio"]), 1)
            before = max(0.0, top_from_body_mm * 72 / 25.4 - used_lines * lp)
        p = self.doc.add_paragraph()
        D.set_para(p, line_exact=round(line, 1), before_pt=round(before, 1), after_pt=0, snap=False)
        return used_lines * lp + before + line

    def docno(self, docno: str, signers: List[str]):
        D, ctx = self.D, self.ctx
        el = ctx.el("docno")
        sg = ctx.el("signer")
        p = self.doc.add_paragraph()
        if signers:
            width = ctx.text_width_twips
            one = int(round(ctx.size(el["size"]) * 20))
            # 多个签发人\uff1a前面的行只排姓名\uff0c最后一行与发文字号同行\uff087.2.6\uff09
            rows = [signers[i:i + 2] for i in range(0, len(signers), 2)]
            for row in rows[:-1]:
                extra = self.doc.add_paragraph()
                p._p.addprevious(extra._p)
                D.set_para(extra, align="right", right=sg["right_chars"], line_exact=ctx.line_pt, before_pt=0, after_pt=0)
                add = extra.add_run("　".join(row))
                D.format_run(add, font=ctx.font(sg["name_font"]), size=ctx.size(sg["size"]))
            # 上行文发文字号居左空一字\uff087.2.5\uff09\uff1a左缩进一字\uff1b签发人居右空一字\uff087.2.6\uff09\uff1a右对齐制表位
            D.set_para(p, align="left", left=float(el["left_chars_upward"]), line_exact=ctx.line_pt, before_pt=0,
                       after_pt=0)
            D.set_tabs(p, [(width - one * int(sg["right_chars"]), "right")])
            r0 = p.add_run(docno)
            D.format_run(r0, font=ctx.font(el["font"]), size=ctx.size(el["size"]))
            r1 = p.add_run("\t签发人\uff1a")
            D.format_run(r1, font=ctx.font(sg["label_font"]), size=ctx.size(sg["size"]))
            r2 = p.add_run("　".join(rows[-1]))
            D.format_run(r2, font=ctx.font(sg["name_font"]), size=ctx.size(sg["size"]))
        else:
            D.set_para(p, align=el["align"], line_exact=ctx.line_pt, before_pt=0, after_pt=0)
            run = p.add_run(docno)
            D.format_run(run, font=ctx.font(el["font"]), size=ctx.size(el["size"]))
        return p

    def letter_head(self, issuer: str) -> None:
        """信函格式\uff1a首页页眉放发文机关标志和红色双线\uff08上粗下细\uff09\uff0c首页页脚放红色双线\uff08上细下粗\uff09\u3002"""
        D, ctx = self.D, self.ctx
        le = ctx.el("letter")
        el = ctx.el("issuer")
        red = D.color_hex(ctx.params, "red")
        ext = (le["rule_length_mm"] - D.pv(ctx.params, "page.body_width_mm")) / 2 * 72 / 25.4 / 16
        thick = int(round(le["thick_pt"] * 8))
        thin = int(round(le["thin_pt"] * 8))
        section = self.section
        section.different_first_page_header_footer = True
        header = section.first_page_header
        header.is_linked_to_previous = False
        D.clear_part(header)
        p1 = header.paragraphs[0]
        size = self.issuer_size(issuer)
        run = p1.add_run(issuer)
        D.format_run(run, font=ctx.font(el["font"]), size=size, color=red)
        D.set_para(p1, align="center", left=-ext, right=-ext, line_exact=round(size * float(el["line_ratio"]), 1),
                   before_pt=0, after_pt=0, snap=False, chars=False)
        D.set_border(p1, "bottom", sz_eighths=thick, color=red, space_pt=int(round(le["rule_gap_mm"] * 72 / 25.4)))
        p2 = header.add_paragraph()
        D.set_para(p2, left=-ext, right=-ext, line_exact=le["double_gap_pt"], before_pt=0,
                   after_pt=round(ctx.size("三号") * 7 / 8, 1), snap=False, chars=False)
        D.set_border(p2, "bottom", sz_eighths=thin, color=red, space_pt=0)
        footer = section.first_page_footer
        footer.is_linked_to_previous = False
        D.clear_part(footer)
        f1 = footer.paragraphs[0]
        D.set_para(f1, left=-ext, right=-ext, line_exact=le["double_gap_pt"], before_pt=0, after_pt=0, snap=False,
                   chars=False)
        D.set_border(f1, "bottom", sz_eighths=thin, color=red, space_pt=0)
        f2 = footer.add_paragraph()
        D.set_para(f2, left=-ext, right=-ext, line_exact=le["double_gap_pt"], before_pt=0, after_pt=0, snap=False,
                   chars=False)
        D.set_border(f2, "bottom", sz_eighths=thick, color=red, space_pt=0)


def merge_fields(parsed: Parsed, meta: dict, warnings: List[str]) -> Dict[str, object]:
    """版式信息优先\uff0c正文识别结果补缺\u3002两者不一致时给出提示\u3002"""
    fields: Dict[str, object] = dict(parsed.fields)
    for key in META_KEYS:
        if key in meta and meta[key] not in (None, "", []):
            old = parsed.fields.get(key)
            new = meta[key]
            if key == "期号" and old and norm_issue(old) != norm_issue(new):
                warnings.append(f'版式信息中的"期号"与正文识别结果不一致，已采用版式信息：{new}（正文中为：{old}）')
            elif old and key in ("标题", "主送机关", "成文日期", "简报名称", "编印单位") and not _same(str(old), str(new)):
                warnings.append(f'版式信息中的"{key}"与正文识别结果不一致，已采用版式信息：{new}（正文中为：{old}）')
            fields[key] = new
    return fields


def decide_layout(cli_layout: Optional[str], meta: dict, parsed: Parsed, report: dict) -> str:
    raw = cli_layout or meta.get("版式")
    if raw:
        layout = LAYOUTS.get(str(raw).strip())
        if not layout:
            raise BuildError(f"无法识别的版式\uff1a{raw}\u3002可选\uff1a红头\u3001信函\u3001纪要\u3001命令\u3001白头\u3001材料\u3001简报")
        report["layout_reason"] = "命令行指定" if cli_layout else "版式信息指定"
        return layout
    if parsed.kind == "简报":
        report["layout_reason"] = "未指定版式\uff0c版式信息或正文中有简报报头\uff0c按简报排版"
        return "简报"
    issuers = parsed.find("issuer")
    if issuers:
        layout = "纪要" if issuers[0].text.endswith("纪要") else "红头"
        report["layout_reason"] = "未指定版式\uff0c正文中有发文机关标志\uff0c按" + layout + "排版"
        return layout
    kind = parsed.kind
    title = meta.get("标题") or parsed.fields.get("标题")
    if meta.get("标题"):
        kind = doc_kind(str(meta["标题"]))[0]
    if kind in ("法定", "内部"):
        report["layout_reason"] = "未指定版式\uff0c按法定公文默认白头"
        layout = "白头"
    else:
        report["layout_reason"] = "未指定版式\uff0c按常用材料默认材料版式" if title else "未识别到标题\uff0c按材料版式处理"
        layout = "材料"
    if meta.get("发文机关标志"):
        report["warnings"].append("提供了发文机关标志但未指定版式\uff0c已按" + layout + "排版\uff1b需要红头请在版式信息中写 \"版式\": \"红头\"")
    if kind == "内部":
        report["warnings"].append("签报一般用本单位统一的签报单格式\uff0c本次按白头排版\uff0c请对照单位模板调整")
    return layout


def build(lines: Sequence[str], meta: dict, *, layout: Optional[str] = None,
          installed: Optional[set] = None, params: Optional[dict] = None, unit_fonts=None,
          overprint: Optional[bool] = None, keep_text: bool = False, placeholders: Optional[bool] = None,
          auto_fix_text: bool = False):
    """生成文档\uff0c返回 (Document, report)\u3002

    排版不改内容\u3001不补内容\uff1a原稿和版式信息里没有的要素一律不写进文档\uff08不补版记\u3001不插【待补】\u3001不按附件页生成附件说明\uff09\uff0c
    只在 report["missing"] 列出\uff0c由用户决定\u3002placeholders=True\uff08版式信息"留占位"为"是"\uff09时才为缺的要素写【待补】占位\uff0c
    用于用户明确要求"缺的要素帮我留出位置"\u3002

    keep_text\uff1a不改任何文字\uff08用户要求"内容别动""一个字都别改"时用\uff09\uff1a不补主送机关冒号\u3001不给附注加括号\u3001不去段首空格\u3001
    不转换日期写法\uff1b这类问题只列入 report["text_issues"]\u3002保存前核对输入的每一行文字都在文档中\uff0c不一致抛出 BuildError\u3002

    auto_fix_text\uff1a自动修正格式性文字问题\uff08补主送机关全角冒号\u3001给附注加括号\u3001转换日期写法 2026-10-16→2026年10月16日\uff09\u3002
    默认关闭\uff08False\uff09\uff1a只在 report["text_issues"] 中报告\uff0c不改动文字\uff0c保持逐字保真\u3002
    keep_text 优先于 auto_fix_text\u3002

    overprint\uff1a预印红头纸套打\uff08只红头版式\uff09\uff0c不输出红色版头\uff0c按预印纸留白\uff1bNone 时看版式信息"套打"\u3002

    字体严格模式\uff1a文档中一律写入标准字体名\uff1binstalled\uff08已安装字体名集合\uff0c空集合表示无法检测\uff09只用于核对\uff0c
    缺字体时在 report 中给出字体核对结果\uff08font_status\u3001missing_fonts\u3001font_notice 等\uff09\uff0c不替代\u3001不降级\u3002
    unit_fonts\uff1a按单位规定指定标准字体\uff08"仿宋=方正仿宋_GBK"字符串或 {角色: 字体名}\uff09\uff0c与版式信息中的 unit_fonts 合并\uff0c
    同一角色以本参数\uff08命令行\uff09为准\uff1b指定后该角色写入单位字体名\uff0c严格模式不变\u3002
    """
    import gw_docx as D
    from gw_format import keep_signoff_together, layout_signature_block

    params = params or D.load_params()
    try:
        meta_raw = next((meta[k] for k in ("unit_fonts", "单位规定字体") if meta.get(k) not in (None, "", {}, [])), None)
        params = F.apply_unit_fonts(params, {**F.parse_unit_fonts(meta_raw, params), **F.parse_unit_fonts(unit_fonts, params)})
    except ValueError as exc:
        raise BuildError(str(exc))
    report: dict = {"warnings": [], "placeholders": [], "missing": [], "detected": [], "notes": []}
    fill = bool(placeholders) if placeholders is not None else bool(_truthy(meta.get("留占位")))
    report["fill_placeholders"] = fill
    # 附件另面编排\uff1a"【附件1】标题"分界行之后的内容拆成附件页\uff0c其前是主文
    lines, attach_secs, attach_notes = extract_attachments(lines)
    report["notes"] += attach_notes
    report["keep_text"] = keep_text
    report["text_issues"] = []
    # 正文中的 Markdown 管道表格先换成占位行\uff0c识别结构后在原位置生成 Word 表格
    lines, tables, table_notes = extract_pipe_tables(lines)
    report["warnings"] += table_notes
    main_lines = list(lines)  # keep_text 核对用\uff1a主文各行\uff08表格已换成占位行\uff09
    main_tables = list(tables)
    explicit = LAYOUTS.get(str(layout or meta.get("版式") or "").strip())
    brief_mode = explicit == "简报" or (not explicit and (has_brief_header(lines) or any(
        meta.get(k) not in (None, "", []) for k in BRIEF_META_KEYS)))
    parsed = parse_brief(lines, force=explicit == "简报") if brief_mode else parse_lines(lines)
    report["warnings"] += parsed.warnings
    if keep_text:  # 解析阶段"已去掉行首空格"的提示不适用\uff1akeep_text 保留段首空格
        report["warnings"] = [w for w in report["warnings"] if not w.startswith("部分段落用空格缩进\uff0c排版时已去掉")]
    report["detected"] = parsed.detected
    lay = decide_layout(layout, meta, parsed, report)
    report["layout"] = lay
    op_on = bool(overprint) if overprint is not None else bool(_truthy(meta.get("套打")))
    if op_on and lay != "红头":
        report["warnings"].append(f'预印红头纸套打只支持红头版式\uff0c{lay}版式已忽略该选项')
        op_on = False
    report["overprint_on"] = op_on
    if attach_secs and lay == "简报":
        report["warnings"].append("简报版式不支持附件另面编排\uff0c附件分界行及其后的内容未排入")
        attach_secs = []

    if _truthy(meta.get("三级标题加粗")):
        params = json.loads(json.dumps(params))
        params["elements"]["h3"]["bold"] = True
        report["notes"].append("按版式信息\uff0c三级标题加粗")

    method = ""
    if installed is None:
        installed, method = F.installed_fonts()
    fonts, check = F.resolve_fonts(params, installed, method)
    report["fonts"] = fonts
    report.update(check)
    ctx = D.Ctx(params)

    fields = merge_fields(parsed, meta, report["warnings"])
    title_text = str(fields.get("标题") or "")
    kind, doc_type = doc_kind(title_text)
    report["doc_type"] = doc_type
    seal_meta = _truthy(meta.get("印章"))
    seal = seal_meta if seal_meta is not None else bool(params["seal_default"][lay])
    report["seal"] = seal

    # 标题下的发文字号行\uff08识别时已认作发文字号\uff0c不进正文\uff09\uff1a说明去向\uff1b版式信息中有发文字号时以版式信息为准
    title_src = [it.src for it in parsed.items if it.role == "title"]
    for it in parsed.items:
        if it.role != "docno" or not title_src or it.src < title_src[0]:
            continue
        if lay not in ("红头", "信函", "纪要"):
            report["notes"].append(f'正文标题下的发文字号"{it.text}"已从正文移出\uff1b{lay}版式不排发文字号\uff0c需要时请改用红头或信函版式')
        elif meta.get("发文字号") not in (None, "", []):
            report["notes"].append(f'正文标题下的文号行"{it.text}"已从正文删除\uff0c版头采用版式信息中的发文字号"{meta["发文字号"]}"')
        else:
            report["notes"].append(f'正文标题下的发文字号"{it.text}"已移到版头发文字号位置\uff0c正文中不再保留')

    def need(key: str, label: Optional[str] = None) -> str:
        """取要素\uff1a有就返回\uff1b没有时默认返回空串\u3001不写进文档\uff0c只记入 report["missing"]\uff1b留占位模式才写【待补】\u3002"""
        val = fields.get(key)
        if isinstance(val, list):
            val = "".join(val)
        if val:
            return str(val)
        if fill:
            mark = f"【待补\uff1a{label or key}】"
            report["placeholders"].append(f'{label or key}\uff1a版式信息和正文中都没有\uff0c已按要求写入 {mark}\uff0c请补齐')
            return mark
        report["missing"].append(f'{label or key}\uff1a原稿和版式信息中都没有\uff0c未写入文档\uff08排版不补内容\uff09\uff0c请用户决定是否提供')
        return ""

    # 日期写法\uff1a版式信息中的 2026-10-16 写法转为"2026年10月16日"
    # 默认只报告不改\uff08保持逐字保真\uff09\uff0c--auto-fix-text 时自动转换\uff1b--keep-text 时同样不改
    for key in ("成文日期", "印发日期"):
        if key in meta and isinstance(fields.get(key), str):
            new, changed = normalize_date(str(fields[key]))
            if changed:
                if auto_fix_text and not keep_text:
                    fields[key] = new
                    report["notes"].append(f'{key}已按 GB/T 9704-2012 写成"{new}"')
                else:
                    report["text_issues"].append(f'{key}"{fields[key]}"写法不规范\uff0c应为"{new}"\uff087.3.5.4\uff09\uff0c按默认不改文字未改\uff1b如需自动转换请加 --auto-fix-text')

    b = Builder(ctx, lay, keep_text)
    if lay == "简报":
        _build_brief(b, parsed, fields, meta, tables, report, need)
        parsed.fields.update(fields)  # 识别结果报告含版式信息中的报头\u3001报尾要素
        D.setup_page_numbers(b.section, ctx, first_page_blank=False)
        report["fields"] = {k: v for k, v in fields.items() if v not in (None, "", [])}
        report["fields"]["版式"] = lay
        report["summary"] = summarize(parsed)
        res_check, doc_texts = _text_check(b, main_lines, main_tables, attach_secs, parsed, meta, fields, [])
        _finish_text_check(b, res_check, doc_texts, keep_text, report)
        return b.doc, report
    used_lines = 0
    head_lines: List[Tuple[str, str]] = []
    if fields.get("份号") and lay != "信函":
        head_lines.append(("copy_no", str(fields["份号"])))
    if fields.get("密级"):
        secret = str(fields["密级"])
        if fields.get("保密期限"):
            secret += "★" + str(fields["保密期限"])
        head_lines.append(("secret", secret))
        report["warnings"].append("保密提示\uff1a版式信息含密级\u3002涉密文件只能在涉密环境中处理\uff0c不得输入联网 AI 工具")
    if fields.get("紧急程度"):
        head_lines.append(("urgency", str(fields["紧急程度"])))

    signers = _as_list(fields.get("签发人"))
    issuer_names = _as_list(fields.get("发文机关标志"))
    issuer_text = issuer_names[0] if issuer_names else ""
    joint = False
    if len(issuer_names) >= 2:
        if lay == "红头":
            joint = True
            issuer_names = [n[:-2] if n.endswith("文件") else n for n in issuer_names]
            issuer_text = issuer_names[0]
            report["notes"].append("联合行文版头\uff1a" + "\u3001".join(issuer_names) + "\uff0c主办机关在前\uff0c\u201c文件\u201d二字置于名称右侧并上下居中"
                                   "\uff08GB/T 9704-2012 7.2.4\uff09\uff1b请确认第一个是主办机关")
        else:
            report["warnings"].append(f'{lay}版式只用主办机关标志\uff0c已取"{issuer_text}"\uff1b联合行文的多机关版头只支持红头版式')
    body_items = [it for it in parsed.items if it.role not in HEAD_ROLES + TITLE_ROLES + TAIL_ROLES + ("recipient",)]
    title_lines = [it.text for it in parsed.items if it.role == "title"]
    if meta.get("标题"):
        title_lines = [s for s in str(meta["标题"]).split("\n") if s.strip()]
        if not parsed.find("title") and body_items and _same(body_items[0].text, str(meta["标题"])):
            body_items = body_items[1:]
    sub_items = [it for it in parsed.items if it.role in ("subtitle", "title_note")]

    # ---- 版头 ----
    if lay == "红头":
        for role, text in head_lines:
            b.para(role, text)
            used_lines += 1
        if not issuer_text:
            issuer_text = need("发文机关标志")
            issuer_names = [issuer_text]
        top_mm = ctx.el("issuer")["top_from_body_mm"]
        header_mm = None
        if op_on:
            raw_h = meta.get("套打版头高度")
            try:
                header_mm = float(raw_h) if raw_h not in (None, "") else None
            except (TypeError, ValueError):
                report["warnings"].append(f'版式信息"套打版头高度"应为毫米数\uff0c"{raw_h}"未采用\uff0c按默认留白')
            used_h = b.overprint_block(issuer_text if not issuer_text.startswith("【待补") else "",
                                       issuer_names if joint else None, top_mm, used_lines, header_mm)
            docno_top = (used_h + ctx.el("docno")["blank_lines_before"] * ctx.line_pt) * 25.4 / 72
            report["overprint"] = {"docno_top_mm": round(docno_top, 1), "custom": header_mm is not None}
            report["notes"].append(
                f"预印红头纸套打\uff1a不输出红色发文机关标志和红色分隔线\uff0c发文字号上沿距版心上边缘约 {docno_top:.0f} mm"
                + ("\uff08按版式信息\u201c套打版头高度\u201d\uff09" if header_mm else "\uff08与正式红头版头一致\uff0c实践值\uff09")
                + "\uff1b预印纸版头尺寸不同的\uff0c用\u201c套打版头高度\u201d调整\uff0c正式打印前先用预印纸试印核对")
        elif joint:
            b.joint_header(issuer_names, top_mm, used_lines)
        else:
            b.issuer(issuer_text, top_mm, used_lines)
        b.blank(ctx.el("docno")["blank_lines_before"])
        if doc_type in UPWARD_TYPES and not signers:
            signers = [x for x in [need("签发人", "签发人姓名")] if x]
            report["warnings"].append(f'{doc_type}属上行文\uff0c红头版式须标注签发人\uff08GB/T 9704-2012 7.2.6\uff09')
        docno_text = need("发文字号")
        if signers:
            _warn_docno_width(docno_text, signers, ctx, report)
        p = b.docno(docno_text, signers)
        if op_on:
            # 不画红线\uff0c但留出红线占的空间\uff08线与文字间距加线宽\uff09\uff0c使标题等位置与正式版头一致
            rule = ctx.el("header_rule")
            D.set_para(p, reset_indent=False, after_pt=round(int(round(rule["gap_mm"] * 72 / 25.4)) + float(rule["width_pt"]), 1))
        else:
            b.red_rule_under(p)
        b.blank(ctx.el("title")["blank_lines_after_rule"])
    elif lay == "信函":
        if not issuer_text:
            issuer_text = need("发文机关标志")
        b.letter_head(issuer_text)
        docno = need("发文字号")
        p = b.doc.add_paragraph()
        copy_no = str(fields.get("份号") or "")
        D.set_para(p, align="left" if copy_no else "right", line_exact=ctx.line_pt, before_pt=0, after_pt=0)
        if copy_no:
            D.set_tabs(p, [(ctx.text_width_twips, "right")])
            r = p.add_run(copy_no)
            D.format_run(r, font=ctx.font("仿宋"), size=ctx.size("三号"))
        r = p.add_run(("\t" if copy_no else "") + docno)
        D.format_run(r, font=ctx.font("仿宋"), size=ctx.size("三号"))
        for role, text in head_lines:
            b.para(role, text)
        b.blank(ctx.el("letter")["blank_lines_before_title"])
    elif lay == "纪要":
        for role, text in head_lines:
            b.para(role, text)
            used_lines += 1
        if not issuer_text and title_text.endswith("纪要"):
            issuer_text = title_text
            title_lines = []
            report["notes"].append("未提供纪要标志\uff0c已用标题\u201c" + title_text + "\u201d作纪要标志\uff0c不再另排标题")
        if not issuer_text:
            issuer_text = need("发文机关标志", "纪要标志")
        p = b.issuer(issuer_text, ctx.el("minutes")["top_from_body_mm"], used_lines)
        if fields.get("发文字号"):
            b.blank(ctx.el("minutes")["blank_lines_before_number"])
            p = b.docno(str(fields["发文字号"]), [])
        b.red_rule_under(p)
        b.blank(ctx.el("title")["blank_lines_after_rule"])
    elif lay == "命令":
        # 命令\uff08令\uff09格式\uff0810.2\uff09\uff1a标志\uff08发文机关全称加"令"或"命令"\uff09居中\u3001红色小标宋\uff0c上边缘距版心上边缘 20 mm\uff1b
        # 标志下空二行居中编令号\uff0c令号下空二行编正文\uff1b没有标题\u3001红色分隔线和主送机关
        od = ctx.el("order")
        for role, text in head_lines:
            b.para(role, text)
            used_lines += 1
        if not issuer_text:
            issuer_text = need("发文机关标志", "发文机关标志\uff08发文机关全称加\u201c令\u201d或\u201c命令\u201d\uff09")
        elif not issuer_text.endswith("令"):
            report["warnings"].append(f'命令\uff08令\uff09格式的发文机关标志应为发文机关全称加"命令"或"令"字\uff0c现为"{issuer_text}"\uff0810.2\uff09')
        b.issuer(issuer_text, od["top_from_body_mm"], used_lines)
        b.blank(od["blank_lines_before_number"])
        # 正文文件首行写的"第12号"被当作标题推断时\uff0c改认作令号
        if not fields.get("令号") and len(title_lines) == 1 and LINGHAO_RE.match(title_lines[0].strip()):
            for it in parsed.items:
                if it.role == "title":
                    it.role = "lingno"
            parsed.fields["令号"] = fields["令号"] = title_lines[0].strip()
            parsed.fields.pop("标题", None)
            fields.pop("标题", None)
            report["warnings"] = [w for w in report["warnings"] if not w.startswith(f'标题"{title_lines[0].strip()}"是按首行推断')]
            report["detected"] = [d for d in report["detected"] if not d.startswith("标题\uff1a首行较短")]
            report["detected"].append("令号\uff1a命令\uff08令\uff09格式中独占一行的\u201c第\u00d7号\u201d")
            title_lines = []
        number = str(fields.get("令号") or "").strip()
        if not number and body_items and LINGHAO_RE.match(body_items[0].text):
            number = body_items[0].text
            body_items = body_items[1:]
        if re.fullmatch(r"\d+", number):
            number = f"第{number}号"
        if not number:
            number = need("令号")
        b.para("lingno", number)
        b.blank(od["blank_lines_after_number"])
        if title_lines:
            report["warnings"].append("命令\uff08令\uff09格式没有标题\uff0810.2\uff09\uff0c识别出的\u201c" + "".join(title_lines) + "\u201d已按正文排\uff0c请核对")
            body_items = [Item("body", "".join(title_lines))] + body_items
            title_lines = []
        report["doc_type"] = "命令"
    else:
        for role, text in head_lines:
            b.para(role, text)
        if head_lines:
            b.blank(1)
        if lay == "白头" and (fields.get("发文字号") or issuer_text):
            report["warnings"].append("白头版式不排发文机关标志和发文字号\uff1b需要时请改用红头或信函版式")

    # ---- 标题 ----
    if len(title_lines) == 1:
        # 超长标题按词意在中点附近回行\uff0c上行不短于下行\uff087.3.1\uff09\uff1b用户自己分了行的不动
        max_chars = int(ctx.text_width_twips / 20 // ctx.size(ctx.el("title")["size"]))
        lines_auto = break_title(title_lines[0], max_chars)
        if len(lines_auto) > 1:
            title_lines = lines_auto
            report["notes"].append("标题较长\uff0c已分行排列\uff1a" + " / ".join(lines_auto) +
                                   "\uff08GB/T 9704-2012 7.3.1\uff09\uff1b要改断行位置\uff0c在版式信息\u201c标题\u201d中用换行符 \\n 分行")
    if title_lines:
        b.title(title_lines)
        for it in sub_items:
            b.para(it.role, it.text)
    elif lay != "命令" and (lay != "纪要" or not issuer_text) and not any("未识别到标题" in w for w in report["warnings"]):
        report["warnings"].append("没有标题\u3002如正文首行是标题但未被识别\uff0c请在版式信息中填写\u201c标题\u201d")
    if title_lines or sub_items:
        b.blank(ctx.el("title")["blank_lines_after"])

    # ---- 主送机关 ----
    recipient = fields.get("主送机关")
    if recipient:
        rec = "\u3001".join(_as_list(recipient)) if isinstance(recipient, list) else str(recipient)
        if body_items and _same(body_items[0].text.rstrip("\uff1a:"), rec.rstrip("\uff1a:")):
            body_items = body_items[1:]
        src_rec = next((it for it in parsed.items if it.role == "recipient"), None)
        if keep_text:
            if src_rec is not None and "主送机关" not in meta:
                rec = src_rec.text  # 保持原文的冒号写法
            if not rec.endswith(("\uff1a", ":")):
                report["text_issues"].append(f'主送机关"{rec}"末尾没有全角冒号\uff087.3.2\uff09\uff0c按 --keep-text 未补')
        elif not rec.endswith(("\uff1a", ":")):
            if auto_fix_text:
                rec += "\uff1a"
            else:
                report["text_issues"].append(f'主送机关"{rec}"末尾没有全角冒号\uff087.3.2\uff09\uff0c按默认不改文字未补\uff1b如需自动补充请加 --auto-fix-text')
        b.para("recipient", rec, raw=main_lines[src_rec.src] if keep_text and src_rec is not None and "主送机关" not in meta
               and 0 <= src_rec.src < len(main_lines) else None)  # keep_text\uff1a行首空格保留\uff0c用负缩进抵消成顶格

    # ---- 版式信息已给出的内容\uff0c正文里的同一行不再重复排出\uff08否则同一行文字会排两次\uff09----
    deduped: List[str] = []
    if meta:
        supplied = set(_supplied_texts(meta))
        kept_items = []
        for it in body_items:
            if it.role == "body" and norm(it.text) in supplied:
                deduped.append(it.text)
                report["notes"].append(f'正文中的"{it.text.strip()}"与版式信息中给出的要素是同一内容\uff0c已只排一次\uff08避免文字重复\uff09')
            else:
                kept_items.append(it)
        body_items = kept_items

    # ---- 正文\uff08含表格\u3001表题\uff09----
    table_queue = list(tables)
    table_font, table_size = _table_style(meta, report)
    prev_role = None
    for k, it in enumerate(body_items):
        if it.role == "table":
            if prev_role == "table":
                b.blank(1)  # 两张表格紧挨着会被 Word 合成一张\uff0c中间留一个空段
            tb = table_queue.pop(0) if table_queue else None
            if tb is not None:
                D.add_body_table(b.doc, ctx, tb.header, tb.rows, tb.aligns, font_role=table_font, size_name=table_size)
        else:
            if it.role == "attendance" and k > 0 and prev_role != "attendance":
                b.blank(1)  # 纪要出席名单在正文下空一行\uff08GB/T 9704-2012 10.3\uff09
            b.para(it.role, it.text, split=it.split,
                   raw=main_lines[it.src] if keep_text and 0 <= it.src < len(main_lines) else None)
        prev_role = it.role
    if keep_text and b.restored:
        report["notes"].append(f'按 --keep-text\uff1a{len(b.restored)} 段用空格缩进的正文保留了段首空格\uff0c首行缩进设为 0\uff0c缩进仍由空格承担')
    if tables:
        report["notes"].append(f'正文中的 {len(tables)} 张 Markdown 表格已转为 Word 表格\uff08全框线\u3001表头加粗\u3001'
                               f"{table_font}{table_size}\uff0c与版心同宽\uff09\uff1b不支持合并单元格\uff0c需要合并的请在 Word 中处理")

    # ---- 附件说明 ----
    attachments = _as_list(fields.get("附件"))
    attach_titles: List[str] = []
    for sec in attach_secs:
        if sec.title:
            attach_titles.append(sec.title)
        elif fill:
            mark = f"【待补\uff1a{sec.label}标题】"
            report["placeholders"].append(f'{sec.label}标题\uff1a附件分界行中没有写标题\uff0c已按要求写入 {mark}\uff0c请补齐')
            attach_titles.append(mark)
        else:
            report["missing"].append(f'{sec.label}标题\uff1a附件分界行中没有写标题\uff0c附件页只排标识\uff0c未补标题\uff0c请用户决定')
            attach_titles.append("")
    if attach_secs:
        if not attachments and (keep_text or not fill):
            report["missing"].append("附件说明\uff087.3.4\uff09\uff1a正文没有\u201c附件\uff1a\u201d说明\uff0c未按附件页标题补写\uff08排版不补内容\uff09\uff0c请用户决定是否补\u201c附件\uff1a\u201d行")
        elif not attachments:
            attachments = [t for t in attach_titles if t]
            report["notes"].append("正文没有附件说明\uff0c已按要求按附件页的标题生成\uff087.3.4\u30017.3.7\uff1a附件说明与附件对应\uff09")
        else:
            wanted = [re.sub(r"[\s　]", "", re.sub(r"^\d{1,2}\s*[.．\u3001]\s*|^附\s*件\s*[\uff1a:]\s*", "", a)) for a in attachments]
            actual = [re.sub(r"[\s　]", "", t) for t in attach_titles]
            if wanted != actual:
                report["warnings"].append("附件说明与附件页不对应\uff087.3.7\uff09\uff1a说明为\u201c" + "\uff1b".join(attachments) +
                                          "\u201d\uff0c附件页为\u201c" + "\uff1b".join(attach_titles) + "\u201d\uff0c请核对顺序号和名称")
    if attachments:
        b.blank(ctx.el("attachment")["blank_lines_before"])
        names = [re.sub(r"^\d{1,2}\s*[.．\u3001]\s*", "", a) for a in attachments]
        names = [re.sub(r"^附\s*件\s*[\uff1a:]\s*", "", a) for a in names]
        if keep_text:  # 附件说明各行照原文排\uff0c不改序号和格式
            verbatim = [it.text for it in parsed.items if it.role == "attachment"] if "附件" not in meta else attachments
            att_ps = [b.para("attachment", t) for t in verbatim]
        elif len(names) == 1:
            att_ps = [b.para("attachment", "附件\uff1a" + names[0])]
        else:
            att_ps = [b.para("attachment", "附件\uff1a1. " + names[0])]
            for i, name in enumerate(names[1:], start=2):
                att_ps.append(b.para("attachment", f'{i}. {name}'))
        for ap in att_ps[:-1]:
            D.set_para(ap, reset_indent=False, keep_next=True)  # 附件说明各行不分页
        for name in names:
            if name and name[-1] in "\u3002\uff1b\uff0c\u3001.;,":
                report["warnings"].append(f'附件名称"{name}"末尾有标点\uff0c应删去\uff08GB/T 9704-2012 7.3.4\uff09')

    # ---- 署名\u3001成文日期 ----
    sigs = _as_list(fields.get("发文机关署名"))
    date_text = str(fields.get("成文日期") or "")
    if lay in LEGAL_LAYOUTS:
        if not sigs:
            sigs = [x for x in [need("发文机关署名")] if x]
        if not date_text:
            date_text = need("成文日期")
    if lay == "命令":
        signoff_first, signoff_last = _order_signoff(b, fields, sigs, date_text, need, report)
    elif seal and len(sigs) > 1 and date_text and (joint or _truthy(meta.get("联合行文"))):
        # 联合行文加盖印章\uff087.3.5.1\uff09\uff1a署名按机关顺序排列\u3001各占一格并居中\uff0c印章一一对应下压署名
        el = ctx.el("signature_seal")
        b.blank(el["blank_lines_before"])
        D.add_joint_signatures(b.doc, ctx, sigs)
        date_p = b.para("date", date_text)
        width_chars = float(D.pv(ctx.params, "grid.chars_per_line"))
        right = float(ctx.el("joint")["date_right_chars"])
        D.set_para(date_p, align="left", left=round(max(0.0, width_chars - right - text_width_chars(date_text)), 2),
                   line_exact=ctx.line_pt, before_pt=0, after_pt=0, snap=True, keep_lines=True)
        signoff_first = signoff_last = date_p
        report["notes"].append(
            f"联合行文加盖印章\uff1a{len(sigs)} 个署名按机关顺序各占一格并居中\uff08每行最多 {D.joint_seal_columns(ctx, len(sigs))} 个\uff0c"
            "格宽不小于印章直径 4.2 cm\uff0c实践值\uff09\uff0c印章一一对应下压署名\uff087.3.5.1\uff09\uff1b加盖位置请人工核对")
    elif sigs or date_text:
        el = ctx.el("signature_seal" if seal else "signature_noseal")
        b.blank(el["blank_lines_before"])
        sig_ps = [b.para("signature", s) for s in sigs]
        date_p = b.para("date", date_text) if date_text else None
        layout_signature_block(sig_ps, date_p, ctx, seal=seal)
        signoff_first = (sig_ps or [date_p])[0]
        signoff_last = date_p if date_p is not None else sig_ps[-1]
    else:
        signoff_first = signoff_last = None
        if seal and len(sigs) > 1:
            report["warnings"].append("联合行文加盖印章时\uff0c署名和印章须一一对应排列\uff087.3.5.1\uff09\uff0c本工具只做近似排列\uff0c请人工调整")
    note = fields.get("附注")
    if note:
        note = str(note)
        if not (note.startswith(("\uff08", "(")) and note.endswith(("\uff09", ")"))):
            if keep_text:
                report["text_issues"].append(f'附注"{note}"没有加圆括号\uff087.3.6\uff09\uff0c按 --keep-text 未加')
            elif auto_fix_text and "附注" in meta:
                note = f"\uff08{note}\uff09"
            elif "附注" in meta:
                report["text_issues"].append(f'附注"{note}"没有加圆括号\uff087.3.6\uff09\uff0c按默认不改文字未加\uff1b如需自动加括号请加 --auto-fix-text')
        note_p = b.para("note", note)
        if signoff_last is not None:
            signoff_last = note_p
    if signoff_first is not None:
        # 落款不单独成页\uff1a正文最后一段\u3001空行\u3001署名\u3001日期\uff08附注\uff09设"与下段同页"
        elements = [p._p for p in b.doc.paragraphs]
        all_paras = b.doc.paragraphs
        keep_signoff_together(all_paras, elements.index(signoff_first._p), elements.index(signoff_last._p),
                              attach_first_idx=elements.index(att_ps[0]._p) if attachments else None)

    # ---- 附件另面编排\uff087.3.7\uff09\uff1a在版记之前\uff0c页码与正文连续 ----
    if attach_secs:
        _build_attach_pages(b, attach_secs, attach_titles, meta, report)

    # ---- 版记 ----
    cc = fields.get("抄送")
    cc_text = ""
    if cc:
        if isinstance(cc, list):
            cc_text = "抄送\uff1a" + "\uff0c".join(_as_list(cc)) + "\u3002"
        else:
            cc_text = str(cc) if str(cc).startswith("抄送") else "抄送\uff1a" + str(cc)
            if keep_text and "抄送" not in meta:  # 照原文\uff0c不改"抄送\uff1a"写法
                cc_text = next((it.text for it in parsed.items if it.role == "cc"), cc_text)
            if not cc_text.endswith("\u3002"):
                report["warnings"].append("抄送最后一个机关名称后应标句号\uff08GB/T 9704-2012 7.4.2\uff09")
    rows: List[List[Tuple[str, str]]] = []
    if cc_text:
        rows.append([("cc", cc_text)])
    has_print = bool(fields.get("印发机关") or fields.get("印发日期"))
    if lay == "信函":
        if has_print:
            report["warnings"].append("信函格式的版记不加印发机关和印发日期\uff08GB/T 9704-2012 10.1\uff09\uff0c已省略")
        if rows:
            D.add_banji_table(b.doc, ctx, rows, with_lines=False)
    elif lay == "材料":
        if rows or has_print:
            report["warnings"].append("材料版式不排版记\uff0c已省略抄送和印发信息")
    elif lay == "红头" or rows or has_print:
        if lay == "红头" or has_print or fill:
            org, date = need("印发机关"), need("印发日期")
        else:  # 白头\u3001纪要只有抄送时\uff0c印发机关不是必需要素\uff0c不记缺失
            org = date = ""
        if org or date:  # 原稿没有印发机关\u3001印发日期的不写\uff08不补版记\uff09
            rows.append([("print", f'{org}\t{date}印发' if date else org)])
        if rows:
            D.add_banji_table(b.doc, ctx, rows, with_lines=True)

    # 文档以表格结尾时补一个 1 磅高的空段\uff0cWord 要求表格后有段落
    last = [el for el in b.doc.element.body.iterchildren() if not el.tag.endswith("}sectPr")]
    if last and last[-1].tag.endswith("}tbl"):
        anchor = b.doc.add_paragraph()
        D.set_para(anchor, line_exact=1, before_pt=0, after_pt=0, snap=False)

    # ---- 页码 ----
    D.setup_page_numbers(b.section, ctx, first_page_blank=(lay == "信函"))
    if lay == "信函":
        # setup_page_numbers 会清空首页页脚\uff0c重画下双线
        b.letter_head(issuer_text)

    report["fields"] = {k: v for k, v in fields.items() if v not in (None, "", [])}
    report["fields"]["版式"] = lay
    report["summary"] = summarize(parsed)
    res_check, doc_texts = _text_check(b, main_lines, main_tables, attach_secs, parsed, meta, fields, deduped)
    _finish_text_check(b, res_check, doc_texts, keep_text, report)
    return b.doc, report


# 版式信息中与要素无关的键\uff08不是写进文档的文字\uff09
NON_TEXT_KEYS = ("版式", "套打", "套打版头高度", "联合行文", "印章", "三级标题加粗", "表格字体", "表格字号", "unit_fonts",
                 "单位规定字体", "留占位")


def _flat(value) -> List[str]:
    if isinstance(value, dict):
        return [x for v in value.values() for x in _flat(v)]
    if isinstance(value, (list, tuple)):
        return [x for v in value for x in _flat(v)]
    return [str(value)] if value not in (None, "") else []


def _supplied_texts(meta: dict) -> List[str]:
    """版式信息中给出的要素文字\uff08规范化后\uff09\uff0c正文里的同一行不再重复排出\u3002"""
    from gw_verify import norm
    out: List[str] = []
    for key, value in meta.items():
        if key in NON_TEXT_KEYS or key in BRIEF_META_KEYS:
            continue
        for v in _flat(value):
            out.append(norm(v))
            if key in ("成文日期", "印发日期"):
                out.append(norm(normalize_date(v)[0]))
            if key == "主送机关":
                out.append(norm(v.rstrip("\uff1a:")))
    org, date = meta.get("印发机关"), meta.get("印发日期")
    if org and date:
        out.append(norm(f"{org}{normalize_date(str(date))[0]}印发"))
    cc = meta.get("抄送")
    if cc:
        out.append(norm("抄送\uff1a" + "\uff0c".join(_flat(cc))) if isinstance(cc, list) else norm(str(cc)))
    return [x for x in out if x]


def _text_check(b: Builder, main_lines: Sequence[str], main_tables, attach_secs, parsed: Parsed, meta: dict,
                fields: dict, deduped: Sequence[str]):
    """核对文档全文与输入全文\uff08正文各行\u3001表格单元格\u3001附件页\uff0c加版式信息中的要素\uff09\uff1a双向\u3001按顺序\uff0c规范化后逐字比对\u3002"""
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph
    from gw_verify import compare_text, norm
    doc_texts = [Paragraph(p, b.doc).text for p in b.doc.element.body.iter(qn("w:p"))]
    inputs: List[str] = []

    def collect(lines, tables) -> None:
        queue = list(tables)
        for x in lines:
            c = clean_line(x)
            if c == TABLE_MARK:
                tb = queue.pop(0) if queue else None
                if tb is not None:
                    inputs.extend(cell for row in [tb.header] + tb.rows for cell in row if cell)
            elif c:
                inputs.append(c)

    collect(main_lines, main_tables)
    labels: List[str] = []
    for sec in attach_secs:
        labels.append(sec.label)
        if sec.title:
            inputs.append(sec.title)
        body, tbs, _ = extract_pipe_tables(sec.lines)
        collect(body, tbs)
    for text in deduped:  # 与版式信息重复\u3001已只排一次的行\uff0c不再计入输入
        for k, x in enumerate(inputs):
            if norm(x) == norm(text):
                del inputs[k]
                break
    allowed = [v for key, val in {**{k: v for k, v in fields.items()}, **{k: v for k, v in meta.items()
               if k not in NON_TEXT_KEYS}}.items() if key != "版式" for v in _flat(val)] + labels
    movable = [it.text for it in parsed.items if it.role in HEAD_ROLES + ("cc", "print", "banji_main")]
    return compare_text(doc_texts, inputs, allowed=allowed, movable=movable), doc_texts


def _finish_text_check(b: Builder, check_result, doc_texts: Sequence[str], keep_text: bool, report: dict) -> None:
    """写入报告的\u201c文字核对\u201d\uff1bkeep_text 时不一致就中止\uff0c不保存\u3002"""
    problems = list(check_result.problems)
    if keep_text:
        problems += [f'少字\uff1a已还原段首空格的行"{r.strip()[:30]}"与文档不一致' for r in b.restored if r not in doc_texts]
    report["text_check"] = {"keep_text": keep_text, "ok": not problems, "chars": check_result.chars,
                            "problems": problems}
    if keep_text and problems:
        raise BuildError("--keep-text 模式下核对发现文档文字与输入不一致\uff0c已中止\uff0c未保存\u3002不一致处\uff1a"
                         + "\uff1b".join(problems[:5]) + (f'\uff1b……共 {len(problems)} 处" if len(problems) > 5 else "'))


def _order_signoff(b: Builder, fields: Dict[str, object], sigs: List[str], date_text: str, need, report: dict):
    """命令\uff08令\uff09的签发人职务\u3001签名章位置和成文日期\uff0810.2\uff0c格式按 7.3.5.3\uff09\u3002

    正文下空二行\uff0c签名章右空四字\uff0c签名章左空二字标签发人职务\uff1b签名章下空一行\uff0c成文日期右空四字\u3002
    签名章占位宽按三个字估算\uff08实践值\uff09\uff1a职务右缘距占位左缘二字\uff0c姓名\uff08版式信息给了才排\uff09排在占位处\uff0c不代替签名章\u3002"""
    from gw_format import keep_signoff_together
    D, ctx = b.D, b.ctx
    od = ctx.el("order")
    post = str(fields.get("签发人职务") or "").strip()
    name = str(fields.get("签发人姓名") or "").strip()
    if not post and sigs:
        parts = re.split(r"[\s　]+", sigs[0].strip(), maxsplit=1)
        post = parts[0]
        name = name or (parts[1] if len(parts) > 1 else "")
        if len(sigs) > 1:
            report["warnings"].append("命令\uff08令\uff09只有一位签发人\uff0c多出的署名行未排入\uff1a" + "\uff1b".join(sigs[1:]))
    if not post:
        post = need("签发人职务")
    if not date_text:
        date_text = need("成文日期")
    width = float(D.pv(ctx.params, "grid.chars_per_line"))
    slot_left = width - float(od["sign_right_chars"]) - float(od["name_slot_chars"])
    post_w = text_width_chars(post)
    left = max(0.0, slot_left - float(od["post_gap_chars"]) - post_w)
    b.blank(od["sign_blank_lines_before"])
    text = post + ("　" * int(od["post_gap_chars"]) + name if name else "")
    sig_p = b.para("signature", text)
    D.set_para(sig_p, align="left", left=round(left, 2), line_exact=ctx.line_pt, before_pt=0, after_pt=0, snap=True,
               keep_lines=True)
    b.blank(od["date_blank_lines"])
    date_p = b.para("date", date_text)
    D.set_para(date_p, align="left", left=round(max(0.0, width - float(od["date_right_chars"]) - text_width_chars(date_text)), 2),
               line_exact=ctx.line_pt, before_pt=0, after_pt=0, snap=True, keep_lines=True)
    report["notes"].append("命令\uff08令\uff09签发人职务\u3001签名章\u3001成文日期按 7.3.5.3 排\uff1a签名章位置在职务右侧\uff08右空四字\uff0c占位约三字\uff0c实践值\uff09"
                           + ("\uff0c姓名为排印占位\uff0c不代替签名章" if name else "\uff0c已留白\uff0c请签署或加盖签名章"))
    return sig_p, date_p


def _build_attach_pages(b: Builder, secs, titles: List[str], meta: dict, report: dict) -> None:
    """附件另面编排\uff087.3.7\uff09\uff1a每个附件另起一页\uff1b第一行顶格\u201c附件\u201d或\u201c附件1\u201d\uff08三号黑体\uff09\uff0c第三行居中附件标题\uff0c
    其下空一行排附件正文\uff08格式同正文\uff0c支持层次标题和简单表格\uff09\u3002页码沿用同一节\uff0c与正文连续\u3002"""
    D, ctx = b.D, b.ctx
    ap = ctx.el("attach_page")
    table_font, table_size = _table_style(meta, report)
    max_chars = int(ctx.text_width_twips / 20 // ctx.size(ctx.el("title")["size"]))
    for sec, title in zip(secs, titles):
        label = b.para("attach_label", sec.label)
        D.set_para(label, reset_indent=False, page_break_before=True)
        b.blank(ap["title_blank_lines_before"])
        tlines = [title]
        auto = break_title(title, max_chars)
        if len(auto) > 1:
            tlines = auto
            report["notes"].append(f'{sec.label}标题较长\uff0c已分行排列\uff1a" + " / '.join(auto))
        b.title(tlines)
        b.blank(ap["blank_lines_after_title"])
        body_lines, tables, notes = extract_pipe_tables(sec.lines)
        report["warnings"] += [f'{sec.label}\uff1a{n}' for n in notes]
        queue = list(tables)
        raws = [x for x in body_lines if clean_line(x)]
        items = [clean_line(x) for x in raws]
        prev = None
        att_levels, _ = heading_levels(items)
        for k, text in enumerate(items):
            if text == TABLE_MARK:
                if prev == "table":
                    b.blank(1)
                tb = queue.pop(0) if queue else None
                if tb is not None:
                    D.add_body_table(b.doc, ctx, tb.header, tb.rows, tb.aligns, font_role=table_font, size_name=table_size)
                prev = "table"
                continue
            if k + 1 < len(items) and items[k + 1] == TABLE_MARK and is_table_caption(text):
                b.para("table_caption", text)
                prev = "table_caption"
                continue
            level = att_levels.get(k) if att_levels else is_heading(text)
            b.para(level or "body", text, split=heading_split(text) if level else -1, raw=raws[k])
            prev = level or "body"
        if not items:
            report["warnings"].append(f'{sec.label}没有正文内容')
    report["notes"].append(f'附件另面编排\uff1a{len(secs)} 个附件各另起一页\uff0c在版记之前\uff0c页码与正文连续\uff087.3.7\u30017.5\uff09\uff1b'
                           "附件说明与附件页标题已对应")


def _build_brief(b: Builder, parsed: Parsed, fields: Dict[str, object], meta: dict, tables, report: dict, need) -> None:
    """简报\uff1a报头\uff08简报名称\u3001期号行\u3001编印单位和日期行\u3001红色分隔线\uff09\u3001各篇标题与正文\u3001报尾\u3002

    简报无国家标准格式\uff0c参数取 gbt9704_params.json 的 elements.brief\uff08常见做法\uff0c各单位以往期样式为准\uff09\u3002
    报头要素缺失不写\uff08留占位模式才写【待补】\uff09\uff0c记入缺失要素\uff1b报尾\uff08报\u3001送\u3001发\u3001抄送\u3001印数\uff09有就排\uff0c没有就不排\u3002
    """
    D, ctx = b.D, b.ctx
    el = ctx.el("brief")
    red = D.color_hex(ctx.params, el["name_color"])
    ignored = [k for k in BRIEF_IGNORED if meta.get(k) not in (None, "", [])]
    if ignored:
        report["warnings"].append("简报版式不排" + "\u3001".join(ignored) + "\uff0c已忽略\uff1b需要时请在 Word 中手工加注")

    # 简报名称\uff1a红色居中\uff0c字号随名称长短取小初到二号
    name = need("简报名称")
    width_pt = ctx.text_width_twips / 20
    size = max(float(el["name_size_min_pt"]), min(float(el["name_size_max_pt"]),
               width_pt * float(el["name_fill_ratio"]) / max(1.0, text_width_chars(name))))
    size = int(size * 2) / 2
    if text_width_chars(name) * size > width_pt:
        report["warnings"].append(f'简报名称"{name}"较长\uff0c一行排不下\uff0c会回行\uff1b请核对名称或缩短')
    p = b.doc.add_paragraph()
    D.format_run(p.add_run(name), font=ctx.font(el["name_font"]), size=size, color=red)
    D.set_para(p, align="center", line_exact=round(size * float(el["name_line_ratio"]), 1), before_pt=0, after_pt=0,
               snap=False, keep_next=True)

    # 期号行\uff1a第12期\uff0c有总期号时写成"第12期\uff08总第68期\uff09"
    issue = norm_issue(fields.get("期号")) if fields.get("期号") else need("期号")
    fields["期号"] = issue
    if fields.get("总期号"):
        fields["总期号"] = norm_total(fields["总期号"])
        issue += "\uff08" + fields["总期号"] + "\uff09"
    p = b.doc.add_paragraph()
    D.format_run(p.add_run(issue), font=ctx.font(el["issue_font"]), size=ctx.size(el["issue_size"]))
    D.set_para(p, align="center", line_exact=ctx.line_pt, before_pt=float(el["issue_before_pt"]), after_pt=0,
               snap=False, keep_next=True)

    # 编印行\uff1a左编印单位\u3001右日期\uff0c同一行\uff1b行下红色分隔线
    editor = str(fields["编印单位"]) if fields.get("编印单位") else need("编印单位")
    if editor and not editor.startswith("【待补") and not editor.endswith(("编", "编印", "编发")):
        editor += "编"
    date = str(fields["印发日期"]) if fields.get("印发日期") else need("印发日期")
    if text_width_chars(editor) + text_width_chars(date) + 1 > float(D.pv(ctx.params, "grid.chars_per_line")):
        report["warnings"].append("编印单位较长\uff0c与日期排不进同一行\uff0c日期会被挤到下一行\uff1b请缩短或改用简称")
    p = b.doc.add_paragraph()
    D.format_run(p.add_run(editor + ("\t" + date if date else "")), font=ctx.font(el["editor_font"]),
                 size=ctx.size(el["editor_size"]))
    D.set_para(p, align="left", line_exact=ctx.line_pt, before_pt=float(el["editor_before_pt"]), after_pt=0,
               snap=False, keep_next=True)
    D.set_tabs(p, [(ctx.text_width_twips, "right")])
    D.set_border(p, "bottom", sz_eighths=int(round(el["rule_width_pt"] * 8)), color=red,
                 space_pt=int(round(el["rule_gap_mm"] * 72 / 25.4)))
    b.blank(el["blank_lines_after_rule"])

    # 各篇\uff1a标题\uff08连续几行都算标题\uff09\u3001题注\u3001层次标题\u3001正文\u3001表格
    table_queue = list(tables)
    table_font, table_size = _table_style(meta, report)
    items = [it for it in parsed.items if not it.role.startswith("brief_")]
    max_chars = int(ctx.text_width_twips / 20 // ctx.size(ctx.el("title")["size"]))
    prev = None
    for it in items:
        role = it.role
        if role == "title":
            if prev is not None and prev != "title":
                b.blank(1)  # 两篇之间空一行
            parts = [t for t in re.split(r"[\s　]+", it.text) if t]  # 标题中间的空格当作换行处
            if len(parts) > 1 and all(text_width_chars(t) <= max_chars for t in parts):
                lines = parts
            else:
                lines = break_title("".join(parts), max_chars)
            if len(lines) > 1 and len(parts) == 1:
                report["notes"].append("标题较长\uff0c已分行排列\uff1a" + " / ".join(lines) + "\uff08要改断行位置\uff0c请把标题拆成两行写\uff09")
            D.set_para(b.title(lines), reset_indent=False, keep_next=True)
        else:
            if prev in ("title", "title_note") and role != "title_note":
                b.blank(1)  # 标题下空一行\uff0c与标题\u3001正文连在一起
                D.set_para(b.doc.paragraphs[-1], reset_indent=False, keep_next=True)
            if role == "table":
                if prev == "table":
                    b.blank(1)
                tb = table_queue.pop(0) if table_queue else None
                if tb is not None:
                    D.add_body_table(b.doc, ctx, tb.header, tb.rows, tb.aligns, font_role=table_font, size_name=table_size)
            else:
                para = b.para(role, it.text, split=it.split)
                if role == "title_note":
                    D.set_para(para, reset_indent=False, keep_next=True)
        prev = role
    if tables:
        report["notes"].append(f'正文中的 {len(tables)} 张 Markdown 表格已转为 Word 表格\uff08全框线\u3001表头加粗\u3001'
                               f"{table_font}{table_size}\uff0c与版心同宽\uff09\uff1b不支持合并单元格\uff0c需要合并的请在 Word 中处理")
    if not any(it.role == "title" for it in items):
        report["warnings"].append("没有识别到篇目标题\uff1b每篇标题应是不带句末标点的独立短行")

    # 报尾\uff1a版心底部一条分隔线\uff0c其下排报\u3001送\u3001发\u3001抄送和印数\uff1b有就排\uff0c没有就不排
    foot: List[Tuple[str, str]] = []
    for key in ("报", "送", "发", "抄送"):
        val = fields.get(key)
        if val in (None, "", []):
            continue
        text = "\uff0c".join(_as_list(val)) if isinstance(val, list) else str(val).strip()
        foot.append(("cc", text if re.match(rf"^{key}\s*[：:]", text) else f"{key}：{text}"))
    if fields.get("印数"):
        n = str(fields["印数"]).strip()
        m = re.search(r"\d+", n)
        foot.append(("print", n if n.startswith(("\uff08", "(")) else f'\uff08共印{m.group(0) if m else n}份\uff09'))
    if foot:
        D.add_banji_table(b.doc, ctx, [foot], with_lines="top")
    else:
        report["notes"].append("没有报送范围\uff08报\u3001送\u3001发\u3001抄送\uff09和印数\uff0c未排报尾\uff1b需要时在版式信息中补\u201c报\u201d\u201c送\u201d\u201c发\u201d\u201c印数\u201d")


def _warn_docno_width(docno: str, signers: List[str], ctx, report: dict) -> None:
    """上行文发文字号与签发人同一行\uff1a估算宽度\uff0c超出版心时提示\uff08Word 会把签发人挤到下一行\uff09\u3002"""
    import gw_docx as D
    per_line = int(D.pv(ctx.params, "grid.chars_per_line"))
    avail = per_line - int(ctx.el("docno")["left_chars_upward"]) - int(ctx.el("signer")["right_chars"])
    last_row = signers[-2:] if len(signers) % 2 == 0 else signers[-1:]
    need = text_width_chars(docno) + text_width_chars("签发人\uff1a" + "　".join(last_row)) + 1
    if need > avail:
        report["warnings"].append(
            f'发文字号"{docno}"与签发人同一行约需 {need:g} 字\uff0c超过可排的 {avail} 字\uff0c签发人会被挤到下一行\uff1b'
            "序号未定的请写成\u201c代字〔年份〕【待补\uff1a序号】号\u201d\uff0c不要用整段长占位\uff08GB/T 9704-2012 7.2.5\u30017.2.6\uff09")


def _table_style(meta: dict, report: dict) -> Tuple[str, str]:
    """表格字体字号\uff1a默认仿宋四号\uff1b版式信息可写\u201c表格字体\u201d\uff08仿宋\u3001宋体\uff09\u3001\u201c表格字号\u201d\uff08四号\u3001小四\uff09\u3002"""
    font, size = "仿宋", "四号"
    for key, allowed in (("表格字体", TABLE_FONTS), ("表格字号", TABLE_SIZES)):
        val = str(meta.get(key) or "").strip()
        if not val:
            continue
        if val in allowed:
            if key == "表格字体":
                font = val
            else:
                size = val
        else:
            report["warnings"].append(f'版式信息"{key}"只能写{_DUN.join(allowed)}\uff0c"{val}"未采用\uff0c按默认处理')
    return font, size


def render_text(report: dict) -> str:
    out: List[str] = []
    notice = F.render_notice(report)  # 缺字体或无法检测字体时\uff0c醒目提示放在最前面
    if notice:
        out += [notice, ""]
    out.append(f"输出文件\uff1a{report.get('output', '')}")
    out.append(f"版式\uff1a{report['layout']}\uff08{report.get('layout_reason', '')}\uff09\uff1b印章\uff1a{'加盖' if report.get('seal') else '不加盖'}")
    nums = iter("一二三四五六七八九十")  # 各部分按实际出现的顺序连续编号
    out.append(f"{next(nums)}\u3001识别结果\uff08请核对\uff09")
    for i, line in enumerate(report.get("summary", []), start=1):
        out.append(f"{i}. {line}")
    for line in report.get("detected", []):
        out.append(f"   依据\uff1a{line}")
    out.append(f"{next(nums)}\u3001字体")
    if report.get("fonts"):
        out.append("   文件规定一套标准字体：" + "；".join(f"{k}：{v}" for k, v in report["fonts"].items()))
    out += [f"   按单位规定使用\uff1a{line}" for line in F.unit_notes(report)]
    status = F.status_line(report)
    if status:
        out.append("   本机检测\uff1a" + status)
    sections = [("缺失要素\uff08原稿没有\uff0c未写入文档\uff0c由用户决定\uff09", report.get("missing", [])),
                ("待补事项", report.get("placeholders", [])),
                ("文字类问题\uff08需人工处理或加 --auto-fix-text 自动修正\uff0c--keep-text 时只列不改\uff09", report.get("text_issues", [])), ("提示", report.get("warnings", []) + report.get("notes", []))]
    for head, items in sections:
        if items:
            out.append(f"{next(nums)}\u3001{head}")
            for i, w in enumerate(items, start=1):
                out.append(f"{i}. {w}")
    tc = report.get("text_check")
    if tc:
        if tc["keep_text"] and tc["ok"]:
            out.append(f"文字核对\uff1a输入全文与文档全文已双向\u3001按顺序规范化比对\uff0c逐字一致\uff08{tc['chars']} 字\uff0c无多字\u3001少字\u3001乱序\u3001重复\uff09")
        elif tc["ok"]:
            out.append("文字核对\uff1a默认模式\uff08不改文字\uff09\uff0c与输入逐字比对未发现差异\uff08如需自动修正格式性文字问题\uff0c可加 --auto-fix-text\uff1b内容别动请加 --keep-text\uff09")
        else:
            out.append("文字核对\uff1a与输入有差异\uff08含允许的格式性处理\uff0c如补冒号\u3001加括号\u3001日期写法\uff09\uff0c不得声称逐字一致\uff1a" + "\uff1b".join(tc["problems"][:6]))
    chk = report.get("check")
    if chk:
        s = chk["summary"]
        out.append(f"{next(nums)}\u3001回读检查\uff1aP0 {s['P0']} 项\uff0cP1 {s['P1']} 项\uff0cP2 {s['P2']} 项")
        for issue in chk["issues"]:
            if issue["level"] in ("P0", "P1"):
                out.append(f"   [{issue['level']}] {issue['message']}\uff1a{issue['excerpt']}")
    return "\n".join(out)


def main(argv: Optional[Sequence[str]] = None) -> int:
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    parser = argparse.ArgumentParser(description="把公文正文排成符合 GB/T 9704-2012 的 docx\u3002",
                                     epilog="示例\uff1apython scripts/build_docx.py 正文.txt --meta 版式信息.json -o 成稿.docx")
    parser.add_argument("input", help="正文文件\uff08.txt 或 .md\uff09\uff0c遇到【起草说明】行停止读取")
    parser.add_argument("--meta", help="版式信息 JSON\uff08键用中文\uff1a版式\u3001发文字号\u3001标题\u3001主送机关等\uff09")
    parser.add_argument("-o", "--output", help="输出 docx 路径\uff1b默认输出到正文同目录\uff1a正文.txt 对应\u201c成稿.docx\u201d\uff0c正文-\u00d7\u00d7.txt "
                                               "对应\u201c成稿-\u00d7\u00d7.docx\u201d\uff0c其他文件用原文件名\uff08如 送审稿.txt 对应\u201c送审稿.docx\u201d\uff09\uff1b"
                                               "已存在时加 -v2\u3001-v3")
    parser.add_argument("--layout", help="版式\uff1a红头\u3001信函\u3001纪要\u3001命令\u3001白头\u3001材料\u3001简报\uff08优先于版式信息\uff09")
    parser.add_argument("--overprint", action="store_true",
                        help="预印红头纸套打\uff08只红头版式\uff09\uff1a不输出红色发文机关标志和红色分隔线\uff0c按预印纸留出版头区域\uff1b"
                             "版头高度不同时在版式信息写\u201c套打版头高度\u201d\uff08毫米\uff09")
    parser.add_argument("--json", action="store_true", help="以 JSON 输出报告")
    parser.add_argument("--keep-text", action="store_true",
                        help="内容别动\uff1a不补主送机关冒号\u3001不给附注加括号\u3001不去段首空格\u3001不改任何文字\uff0c只排版\uff1b"
                             "文字类问题只在报告中列出\uff1b保存前核对文字与输入一致\uff0c不一致则中止")
    parser.add_argument("--auto-fix-text", action="store_true",
                        help="自动修正格式性文字问题\uff1a补主送机关全角冒号\u3001给附注加圆括号\u3001转换日期写法\uff082026-10-16→2026年10月16日\uff09\uff1b"
                             "默认不开启\uff0c只在报告中提示\uff0c保持逐字保真\uff1b--keep-text 优先")
    parser.add_argument("--placeholders", action="store_true",
                        help="缺的要素写【待补】占位\uff08只在用户明确要求\u201c缺的要素帮我留出位置\u201d时用\uff09\uff1b默认不写\uff0c只在报告\u201c缺失要素\u201d中列出")
    parser.add_argument("--no-check", action="store_true", help="生成后不自动回读检查")
    parser.add_argument("--allow", metavar="DOCX-050", help="回读检查时按本单位制度保留的检查项\uff08逗号分隔\uff09\uff0c命中降为 P2")
    parser.add_argument("--unit-fonts", metavar="角色=字体名,…", help=F.UNIT_FONTS_HELP + "\uff08优先于版式信息中的 unit_fonts\uff09")
    args = parser.parse_args(argv)

    try:
        import docx  # noqa: F401
    except ImportError:
        print("缺少 python-docx\u3002请先安装\uff1apip install -U python-docx\uff08需要 1.2.0 及以上版本\uff09")
        return 2
    import gw_docx as D

    src = Path(args.input)
    if not src.is_file():
        print(f"错误\uff1a找不到正文文件 {src}")
        return 2
    if src.suffix.lower() not in (".txt", ".md", ".markdown"):
        print("错误\uff1a正文文件须为 .txt 或 .md")
        return 2
    meta: dict = {}
    if args.meta:
        try:
            meta = json.loads(Path(args.meta).read_text(encoding="utf-8-sig"))
        except (OSError, ValueError) as exc:  # ValueError 含 JSON 格式错误和非 UTF-8 编码
            print(f"错误\uff1a版式信息 JSON 读取失败\uff1a{exc}")
            return 2
        if not isinstance(meta, dict):
            print("错误\uff1a版式信息 JSON 顶层必须是对象")
            return 2
    try:
        lines, read_notes = read_body_text(src)
    except (OSError, ValueError) as exc:  # 无法读取或无法识别编码
        print(f"错误\uff1a正文文件读取失败\uff1a{exc}")
        return 2
    try:
        doc, report = build(lines, meta, layout=args.layout, unit_fonts=args.unit_fonts,
                            overprint=True if args.overprint else None, keep_text=args.keep_text,
                            auto_fix_text=args.auto_fix_text,
                            placeholders=True if args.placeholders else None)
    except BuildError as exc:
        print(f"错误\uff1a{exc}")
        return 2
    report["notes"] += read_notes
    unknown = [k for k in meta if k not in META_KEYS]
    if unknown:
        report["warnings"].append("版式信息中有未识别的键\uff0c已忽略\uff1a" + "\u3001".join(unknown))
    # 默认输出名\uff1a正文.txt → 成稿.docx\uff1b正文-单位总结.txt → 成稿-单位总结.docx\uff08output-format.md 的多版本命名\uff09\uff1b
    # 其他文件用原文件名\uff0c如 送审稿.txt → 送审稿.docx\uff0c避免不同稿本都输出成"成稿"
    if src.stem == "正文":
        default_name = "成稿.docx"
    elif src.stem.startswith("正文-"):
        default_name = "成稿-" + src.stem[3:] + ".docx"
    else:
        default_name = src.stem + ".docx"
    out = Path(args.output) if args.output else src.with_name(default_name)
    if out.is_dir():  # -o 给的是已有文件夹\uff1a在文件夹内按默认名输出
        out = out / default_name
    final = D.unique_path(out)
    if final != out:
        report["notes"].append(f'{out.name} 已存在\uff0c未覆盖\uff0c另存为 {final.name}')
    final.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(final))
    report["input"] = str(src)
    report["output"] = str(final)
    if not args.no_check:
        from check_docx import check_document
        report["check"] = check_document(final, layout=report["layout"], allow=args.allow,
                                         unit_fonts={r: v["unit_font"] for r, v in report["unit_fonts"].items()},
                                         overprint=bool(report.get("overprint_on")))
    if args.json:
        print(json.dumps(report, ensure_ascii=False, indent=2))
    else:
        print(render_text(report))
    return 0


if __name__ == "__main__":
    sys.exit(main())
