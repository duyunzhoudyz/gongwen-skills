# -*- coding: utf-8 -*-
"""skills/gongwen-draft/scripts/draft_check.py 的单元测试。

运行：python -m unittest discover -s tests/draft -t .

samples/ 下 pass_*.txt 是必过标本（不得报出任何问题），fail_*.txt 是必错标本
（必须报出 EXPECTED_FAIL 中列出的规则）。
"""
from __future__ import annotations

import io
import json
import shutil
import sys
import tempfile
import unittest
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[2]
SCRIPTS_DIR = REPO_ROOT / "skills" / "gongwen-draft" / "scripts"
SAMPLES = Path(__file__).resolve().parent / "samples"
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

import draft_check  # noqa: E402

try:
    import docx  # type: ignore
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

# 必过标本：文件名 -> --doc-type（空字符串表示按标题推断）
EXPECTED_PASS = {
    "pass_qingshi.txt": "请示",
    "pass_huiyi.txt": "",
    "pass_pifu.txt": "批复",
    "pass_han.txt": "函",
    "pass_zhuanfa.txt": "通知",
    "pass_tonggao.txt": "通告",
    "pass_qianbao.txt": "签报",
    "pass_tongbao.txt": "通报",
    # 评审回归：会议时间小标题单独成行
    "pass_huiyi_fenhang.txt": "通知",
    # 评审回归：“研究决定，拟……”是正当的内部决策表述
    "pass_qingshi_juece.txt": "请示",
    # 评审回归：“可视……酌情调整”是授权，不是表态含糊
    "pass_pifu_shouquan.txt": "批复",
    # 评审回归：事由中的“印发”不是印发通知
    "pass_tongzhi_yinfa_guanli.txt": "通知",
    # 评审回归：批转这一层保留，不算层层转发
    "pass_zhuanfa_pizhuan.txt": "通知",
    # 第二轮评审回归：“根据我局党组决定精神”是引用已有决定，不报 DRF-011
    "pass_qingshi_genju.txt": "请示",
    # 第二轮评审回归：只有一项请求的请示，事项本身就是意见，不报 DRF-012
    "pass_qingshi_danyi.txt": "请示",
    # 第二轮评审回归：供水企业检修停水的告知不算强制措施
    "pass_tongzhi_jianxiu.txt": "通知",
    # 第二轮评审回归：标题折成两行，按标题推断文种
    "pass_title_zhehang.txt": "",
    # 第二轮评审回归：转发标题第二行以“关于”开头
    "pass_zhuanfa_zhehang.txt": "通知",
}

# 必错标本：文件名 -> (--doc-type, 必须出现的规则号及级别)
EXPECTED_FAIL = {
    "fail_huiyi.txt": ("通知", {"DRF-050": "P0", "DRF-051": "P0", "DRF-055": "P2"}),
    "fail_qingshi.txt": ("请示", {"DRF-006": "P2", "DRF-010": "P2", "DRF-011": "P1", "DRF-013": "P2"}),
    "fail_pifu.txt": ("批复", {"DRF-040": "P1", "DRF-041": "P1", "DRF-042": "P1"}),
    "fail_han.txt": ("函", {"DRF-030": "P2", "DRF-031": "P1", "DRF-032": "P2"}),
    "fail_zhuanfa.txt": ("通知", {"DRF-001": "P1", "DRF-002": "P1", "DRF-003": "P1", "DRF-060": "P1"}),
    "fail_qianbao.txt": ("签报", {"DRF-080": "P1", "DRF-081": "P2"}),
    "fail_tonggao.txt": ("通告", {"DRF-090": "P2", "DRF-091": "P1", "DRF-092": "P2", "DRF-093": "P2", "DRF-094": "P1"}),
    "fail_gonggao.txt": ("公告", {"DRF-100": "P2"}),
    "fail_tongbao.txt": ("通报", {"DRF-070": "P2"}),
    "fail_title.txt": ("请示", {"DRF-004": "P1"}),
    "fail_baogao.txt": ("报告", {"DRF-020": "P2"}),
    "fail_yinfa.txt": ("通知", {"DRF-062": "P1"}),
    # 只有回执截止时间、没有会议时间
    "fail_huiyi_huizhi.txt": ("通知", {"DRF-050": "P0"}),
    # 签报写成“主送：院长、分管副院长、财务科”
    "fail_qianbao_zhusong.txt": ("签报", {"DRF-080": "P1"}),
    # 第二轮评审回归：两个方案没有倾向
    "fail_qingshi_fangan.txt": ("请示", {"DRF-012": "P2"}),
    # 第二轮评审回归：标题没有“关于”的社区通告；用停水迫使居民履行
    "fail_tonggao_wuguanyu.txt": ("通告", {"DRF-094": "P1", "DRF-095": "P1"}),
    # 第二轮评审回归：居民通知中“对其家庭停止供水”
    "fail_tongzhi_tingshui.txt": ("通知", {"DRF-095": "P1"}),
    # 第二轮评审回归：会议时间只有日期没有钟点
    "fail_huiyi_shike.txt": ("通知", {"DRF-054": "P2"}),
}


def run_main(*args: str):
    out, err = io.StringIO(), io.StringIO()
    with redirect_stdout(out), redirect_stderr(err):
        code = draft_check.main(list(args))
    return code, out.getvalue(), err.getvalue()


def check_file(name: str, doc_type: str) -> dict:
    args = [str(SAMPLES / name), "--json"]
    if doc_type:
        args += ["--doc-type", doc_type]
    code, out, _err = run_main(*args)
    if code != 0:
        raise AssertionError(f"{name} 退出码 {code}")
    return json.loads(out)


class SampleTest(unittest.TestCase):

    def test_samples_exist(self):
        for name in list(EXPECTED_PASS) + list(EXPECTED_FAIL):
            self.assertTrue((SAMPLES / name).is_file(), name)

    def test_must_pass_samples_report_nothing(self):
        for name, doc_type in EXPECTED_PASS.items():
            with self.subTest(sample=name):
                report = check_file(name, doc_type)
                found = [(i["id"], i["level"], i["message"]) for i in report["issues"]]
                self.assertEqual(found, [], f"{name} 不应报出问题")

    def test_must_fail_samples_report_expected_rules(self):
        for name, (doc_type, expected) in EXPECTED_FAIL.items():
            with self.subTest(sample=name):
                report = check_file(name, doc_type)
                levels = {i["id"]: i["level"] for i in report["issues"]}
                for rule, level in expected.items():
                    self.assertIn(rule, levels, f"{name} 应报出 {rule}")
                    self.assertEqual(levels[rule], level, f"{name} 的 {rule} 级别")

    def test_title_mismatch(self):
        lines = draft_check.split_doc  # 保持引用，避免误删
        self.assertTrue(callable(lines))
        report = check_file("pass_han.txt", "请示")
        ids = [i["id"] for i in report["issues"]]
        self.assertIn("DRF-005", ids)

    def test_two_line_title_is_merged(self):
        report = check_file("pass_title_zhehang.txt", "")
        self.assertEqual(report["doc_type"], "请示")
        lines = draft_check.read_lines(SAMPLES / "pass_zhuanfa_zhehang.txt")
        doc = draft_check.split_doc(lines)
        self.assertEqual(doc.title.text, "某区教育局转发省教育厅关于做好冬季校园安全工作的通知")
        self.assertEqual(doc.zhusong.text, "各中小学、幼儿园，局属各单位：")

    def test_issuer_without_guanyu(self):
        issuer = draft_check.Checker.issuer_of
        self.assertEqual(issuer("某社区居民委员会通告"), "某社区居民委员会")
        self.assertEqual(issuer("某县人民政府禁止在学校周边设摊的通告"), "某县人民政府")
        self.assertEqual(issuer("某县人民政府关于汛期禁止在某河道采砂的通告"), "某县人民政府")

    def test_infer_type_from_title(self):
        report = check_file("pass_huiyi.txt", "")
        self.assertEqual(report["doc_type"], "通知")
        self.assertTrue(any("推断" in n for n in report["notes"]))


class OutputFormatTest(unittest.TestCase):

    def test_json_schema_matches_lint_interface(self):
        report = check_file("fail_qingshi.txt", "请示")
        for key in ("file", "doc_type", "char_count", "summary", "issues"):
            self.assertIn(key, report)
        self.assertEqual(set(report["summary"]["markers"]), {"待补", "待核实"})
        for issue in report["issues"]:
            for key in ("id", "level", "category", "line", "excerpt", "message", "suggestion", "basis"):
                self.assertIn(key, issue)
            self.assertIn(issue["level"], ("P0", "P1", "P2"))
            self.assertEqual(issue["check"], "draft")

    def test_issues_sorted_by_level(self):
        report = check_file("fail_huiyi.txt", "通知")
        levels = [i["level"] for i in report["issues"]]
        self.assertEqual(levels, sorted(levels))

    def test_markers_counted(self):
        report = check_file("pass_han.txt", "函")
        self.assertEqual(report["summary"]["markers"]["待补"], 4)
        self.assertEqual(report["summary"]["P0"], 0)

    def test_text_output(self):
        code, out, _ = run_main(str(SAMPLES / "fail_pifu.txt"), "--doc-type", "批复")
        self.assertEqual(code, 0)
        self.assertIn("DRF-041", out)
        self.assertIn("P0 0 条", out)


class ExitCodeTest(unittest.TestCase):

    def setUp(self) -> None:
        self.tmp = Path(tempfile.mkdtemp(prefix="draft_check_"))

    def tearDown(self) -> None:
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_strict_returns_1_on_p0(self):
        code, _, _ = run_main(str(SAMPLES / "fail_huiyi.txt"), "--doc-type", "通知", "--strict")
        self.assertEqual(code, 1)

    def test_strict_returns_0_without_p0(self):
        code, _, _ = run_main(str(SAMPLES / "fail_pifu.txt"), "--doc-type", "批复", "--strict")
        self.assertEqual(code, 0)

    def test_missing_file_returns_2(self):
        code, _, err = run_main(str(self.tmp / "不存在.txt"))
        self.assertEqual(code, 2)
        self.assertIn("未找到文件", err)

    def test_unsupported_suffix_returns_2(self):
        path = self.tmp / "a.pdf"
        path.write_bytes(b"%PDF")
        code, _, err = run_main(str(path))
        self.assertEqual(code, 2)
        self.assertIn("不支持", err)

    def test_empty_file_returns_2(self):
        path = self.tmp / "空.txt"
        path.write_text("\n\n", encoding="utf-8")
        code, _, _ = run_main(str(path))
        self.assertEqual(code, 2)

    def test_no_arguments_returns_2(self):
        code, _, _ = run_main()
        self.assertEqual(code, 2)

    def test_gb18030_file_is_readable(self):
        path = self.tmp / "gbk.txt"
        path.write_bytes((SAMPLES / "fail_pifu.txt").read_text(encoding="utf-8").encode("gb18030"))
        code, out, _ = run_main(str(path), "--doc-type", "批复", "--json")
        self.assertEqual(code, 0)
        self.assertIn("DRF-041", out)

    def test_stops_at_draft_note(self):
        # pass_huiyi.txt 在【起草说明】之后写了"近期"，不得影响检查
        report = check_file("pass_huiyi.txt", "通知")
        self.assertEqual(report["issues"], [])

    @unittest.skipUnless(HAS_DOCX, "未安装 python-docx")
    def test_docx_input(self):
        document = docx.Document()
        for line in (SAMPLES / "fail_huiyi.txt").read_text(encoding="utf-8").splitlines():
            document.add_paragraph(line)
        document.add_paragraph("【起草说明】")
        document.add_paragraph("会议地点：第一会议室")
        path = self.tmp / "会议通知.docx"
        document.save(str(path))
        code, out, _ = run_main(str(path), "--doc-type", "通知", "--json")
        self.assertEqual(code, 0)
        ids = [i["id"] for i in json.loads(out)["issues"]]
        self.assertIn("DRF-051", ids)  # 起草说明之后的地点不算


class HelpTest(unittest.TestCase):

    def test_help_is_chinese(self):
        out = io.StringIO()
        with redirect_stdout(out), self.assertRaises(SystemExit) as ctx:
            draft_check.main(["--help"])
        self.assertEqual(ctx.exception.code, 0)
        text = out.getvalue()
        for word in ("用法：", "位置参数", "选项", "显示帮助并退出"):
            self.assertIn(word, text)
        for word in ("usage:", "positional arguments", "options:", "show this help message"):
            self.assertNotIn(word, text)

    def test_bad_argument_message_is_chinese(self):
        err = io.StringIO()
        with redirect_stderr(err), self.assertRaises(SystemExit) as ctx:
            draft_check.main(["--fold-title", "某某的通知", "--issuer", "某局", "--verb", "印发"])
        self.assertEqual(ctx.exception.code, 2)
        self.assertIn("错误：", err.getvalue())
        self.assertIn("用法：", err.getvalue())


class UnknownTypeTest(unittest.TestCase):

    def test_unknown_doc_type_only_checks_title(self):
        lines = [draft_check.Line(1, "某局2026年工作总结"), draft_check.Line(3, "一、主要工作")]
        dtype, issues, notes = draft_check.check_text_lines(lines, "总结")
        self.assertEqual(dtype, "总结")
        self.assertEqual(issues, [])
        self.assertTrue(any("法定文种范围" in n for n in notes))

    def test_header_lines_skipped(self):
        lines = [
            draft_check.Line(1, "某发〔2026〕5号"),
            draft_check.Line(2, "签发人：张某"),
            draft_check.Line(4, "某镇人民政府关于拨付资金的请示"),
            draft_check.Line(6, "县人民政府："),
        ]
        doc = draft_check.split_doc(lines)
        self.assertEqual(doc.title.text, "某镇人民政府关于拨付资金的请示")
        self.assertEqual(doc.zhusong.text, "县人民政府：")


class FoldTitleTest(unittest.TestCase):

    def fold(self, title: str, issuer: str, verb: str = "转发") -> dict:
        return draft_check.fold_forward_title(title, issuer, verb)

    def test_three_level_chain_folds_to_original(self):
        result = self.fold(
            "某市人民政府办公室关于转发省人民政府办公厅关于做好秋冬季森林防火工作的通知的通知",
            "某县人民政府办公室",
        )
        self.assertEqual(result["recommended"], "某县人民政府办公室转发省人民政府办公厅关于做好秋冬季森林防火工作的通知")
        # 备选为不写本机关名称的写法（计划 D19），省去被转发标题中的“关于”
        self.assertEqual(result["alternative"], "关于转发省人民政府办公厅做好秋冬季森林防火工作的通知")
        self.assertEqual(result["skipped_levels"], ["某市人民政府办公室"])

    def test_opinion_drops_inner_de(self):
        # 与国办函〔2021〕132号标题写法一致
        result = self.fold("司法部关于审理政府信息公开行政复议案件若干问题的指导意见", "国务院办公厅")
        self.assertEqual(result["recommended"], "国务院办公厅转发司法部关于审理政府信息公开行政复议案件若干问题指导意见的通知")
        self.assertEqual(result["skipped_levels"], [])

    def test_merged_notice_form_is_parsed(self):
        result = self.fold("广东省人民政府办公厅转发国务院办公厅关于加强政府网站域名管理的通知", "某市人民政府办公室")
        self.assertEqual(result["original_title"], "国务院办公厅关于加强政府网站域名管理的通知")
        self.assertEqual(result["recommended"], "某市人民政府办公室转发国务院办公厅关于加强政府网站域名管理的通知")

    def test_printed_document_keeps_book_marks(self):
        result = self.fold("省人民政府办公厅关于印发《某某行动方案》的通知", "某县人民政府办公室")
        self.assertEqual(result["recommended"], "某县人民政府办公室转发省人民政府办公厅关于印发《某某行动方案》的通知")

    def test_statute_without_issuer_gets_book_marks(self):
        result = self.fold("《某省安全生产条例》", "某县应急管理局")
        self.assertEqual(result["recommended"], "某县应急管理局转发《某省安全生产条例》的通知")
        self.assertTrue(result["notes"])

    def test_subject_containing_forward_word_is_not_a_chain(self):
        result = self.fold("市政府办公室关于做好文件转发工作的通知", "某县人民政府办公室")
        self.assertEqual(result["skipped_levels"], [])
        self.assertEqual(result["recommended"], "某县人民政府办公室转发市政府办公室关于做好文件转发工作的通知")

    def test_batch_verb(self):
        result = self.fold("某区教育局关于加强课后服务管理的意见", "某区人民政府", "批转")
        self.assertEqual(result["recommended"], "某区人民政府批转某区教育局关于加强课后服务管理意见的通知")

    def test_statute_kinds_keep_book_marks(self):
        # 评审回归：最长匹配的“管理办法”“实施细则”“实施办法”也要加书名号
        for name in ("某市城市绿化管理办法", "某省物业管理条例实施细则", "某市城市更新实施办法"):
            with self.subTest(name=name):
                result = self.fold(f"《{name}》", "某区城市管理局")
                self.assertEqual(result["recommended"], f"某区城市管理局转发《{name}》的通知")
                result = self.fold(name, "某区城市管理局")
                self.assertEqual(result["recommended"], f"某区城市管理局转发《{name}》的通知")

    def test_input_book_marks_are_kept(self):
        result = self.fold("《某某行动方案》", "某县人民政府办公室")
        self.assertEqual(result["recommended"], "某县人民政府办公室转发《某某行动方案》的通知")

    def test_batch_middle_level_is_kept(self):
        result = self.fold("某市人民政府关于批转市发展改革委2026年重点项目计划的通知", "某县人民政府")
        self.assertEqual(result["recommended"], "某县人民政府转发某市人民政府关于批转市发展改革委2026年重点项目计划的通知")
        self.assertEqual(result["skipped_levels"], [])
        self.assertTrue(any("保留了" in n for n in result["notes"]))

    def test_government_forwarding_department_is_kept(self):
        result = self.fold("国务院办公厅转发国家发展改革委关于推进城乡供水一体化意见的通知", "某省人民政府办公厅")
        self.assertEqual(result["recommended"],
                         "某省人民政府办公厅转发国务院办公厅转发国家发展改革委关于推进城乡供水一体化意见的通知")
        self.assertEqual(result["forwarded_title"], "国务院办公厅转发国家发展改革委关于推进城乡供水一体化意见的通知")

    def test_pass_through_level_above_kept_level_is_skipped(self):
        result = self.fold("某县人民政府办公室转发某市人民政府办公室转发市发展改革委关于推进城乡供水一体化实施意见的通知",
                           "某镇人民政府")
        self.assertEqual(result["skipped_levels"], ["某县人民政府办公室"])
        self.assertEqual(result["recommended"],
                         "某镇人民政府转发某市人民政府办公室转发市发展改革委关于推进城乡供水一体化实施意见的通知")

    def test_missing_original_issuer_keeps_guanyu(self):
        result = self.fold("关于做好秋冬季森林防火工作的通知", "某县人民政府办公室")
        self.assertEqual(result["recommended"], "某县人民政府办公室转发关于做好秋冬季森林防火工作的通知")
        self.assertEqual(result["alternative"], "")
        self.assertTrue(any("原发文机关" in n for n in result["notes"]))

    def test_merged_form_keeps_inner_notice_kind(self):
        result = self.fold("某市人民政府办公室转发省人民政府办公厅关于印发某某行动方案的通知", "某县人民政府办公室")
        self.assertEqual(result["original_title"], "省人民政府办公厅关于印发某某行动方案的通知")
        self.assertEqual(result["recommended"], "某县人民政府办公室转发省人民政府办公厅关于印发某某行动方案的通知")

    def test_merged_outer_and_separate_inner_chain(self):
        result = self.fold("某市人民政府办公室转发省政府办公厅关于转发国务院办公厅关于X的意见的通知", "某县人民政府办公室")
        self.assertEqual(result["skipped_levels"], ["某市人民政府办公室", "省政府办公厅"])
        self.assertEqual(result["original_title"], "国务院办公厅关于X的意见")

    def test_kept_level_title_passes_title_check(self):
        result = self.fold("国务院办公厅转发国家发展改革委关于推进城乡供水一体化意见的通知", "某省人民政府办公厅")
        lines = [draft_check.Line(1, result["recommended"]), draft_check.Line(3, "各市人民政府："),
                 draft_check.Line(5, "现将《……》转发给你们，请结合实际认真贯彻执行。")]
        _dtype, issues, _ = draft_check.check_text_lines(lines, "通知")
        self.assertEqual([i.id for i in issues], [])

    def test_cli_fold(self):
        code, out, _ = run_main("--fold-title", "省人民政府办公厅关于做好秋冬季森林防火工作的通知",
                                "--issuer", "某县人民政府办公室")
        self.assertEqual(code, 0)
        self.assertIn("推荐标题：某县人民政府办公室转发省人民政府办公厅关于做好秋冬季森林防火工作的通知", out)

    def test_cli_fold_without_issuer(self):
        # D19：不传 --issuer 时给出“关于转发……的通知”写法
        code, out, _ = run_main("--fold-title", "省人民政府办公厅关于做好秋冬季森林防火工作的通知")
        self.assertEqual(code, 0)
        self.assertIn("推荐标题：关于转发省人民政府办公厅做好秋冬季森林防火工作的通知", out)
        self.assertNotIn("备选标题", out)

    def test_cli_fold_requires_title(self):
        code, _, _ = run_main("--fold-title", "  ", "--issuer", "某局")
        self.assertEqual(code, 2)

    def test_without_issuer_folds_chain(self):
        result = self.fold("某市人民政府办公室关于转发省人民政府办公厅关于做好秋冬季森林防火工作的通知的通知", "")
        self.assertEqual(result["recommended"], "关于转发省人民政府办公厅做好秋冬季森林防火工作的通知")
        self.assertEqual(result["alternative"], "")
        self.assertEqual(result["skipped_levels"], ["某市人民政府办公室"])
        self.assertTrue(any("本机关名称" in n for n in result["notes"]))

    def test_without_issuer_keeps_single_guanyu(self):
        cases = {
            "司法部关于审理政府信息公开行政复议案件若干问题的指导意见":
                "关于转发司法部审理政府信息公开行政复议案件若干问题指导意见的通知",
            "《某省安全生产条例》": "关于转发《某省安全生产条例》的通知",
            "关于做好秋冬季森林防火工作的通知": "关于转发做好秋冬季森林防火工作的通知",
        }
        for title, expected in cases.items():
            with self.subTest(title=title):
                result = self.fold(title, "")
                self.assertEqual(result["recommended"], expected)
                self.assertEqual(result["recommended"].count("关于"), 1)

    def test_without_issuer_title_passes_title_check(self):
        result = self.fold("某市人民政府办公室关于转发省人民政府办公厅关于做好秋冬季森林防火工作的通知的通知", "")
        lines = [draft_check.Line(1, result["recommended"]), draft_check.Line(3, "各乡镇人民政府："),
                 draft_check.Line(5, "现将《……》转发给你们，请结合实际认真贯彻执行。")]
        _dtype, issues, _ = draft_check.check_text_lines(lines, "通知")
        self.assertEqual(issues, [])

    def test_issuer_followed_by_guanyu_forward_is_flagged(self):
        # D19：标题以本机关名称开头时直接写“转发”，不加“关于”
        body = [draft_check.Line(3, "各中小学："), draft_check.Line(5, "现将《……》转发给你们，请结合实际认真贯彻执行。")]
        flagged = "某区教育局关于转发省教育厅做好冬季校园安全工作的通知"
        _dtype, issues, _ = draft_check.check_text_lines([draft_check.Line(1, flagged)] + body, "通知")
        self.assertEqual([(i.id, i.level) for i in issues if i.id == "DRF-001"], [("DRF-001", "P2")])
        for title in ("关于转发省教育厅做好冬季校园安全工作的通知",
                      "某区教育局转发省教育厅关于做好冬季校园安全工作的通知",
                      "某市人民政府关于批转市发展改革委2026年重点项目计划的通知"):
            with self.subTest(title=title):
                _dtype, issues, _ = draft_check.check_text_lines([draft_check.Line(1, title)] + body, "通知")
                self.assertEqual([i.id for i in issues if i.id == "DRF-001"], [])

    def test_folded_title_passes_title_check(self):
        result = self.fold(
            "某市人民政府办公室关于转发省人民政府办公厅关于做好秋冬季森林防火工作的通知的通知",
            "某县人民政府办公室",
        )
        lines = [draft_check.Line(1, result["recommended"]), draft_check.Line(3, "各乡镇人民政府："),
                 draft_check.Line(5, "现将《……》转发给你们，请结合实际认真贯彻执行。")]
        _dtype, issues, _ = draft_check.check_text_lines(lines, "通知")
        self.assertEqual(issues, [])


if __name__ == "__main__":
    unittest.main()
