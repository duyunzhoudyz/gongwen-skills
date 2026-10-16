# -*- coding: utf-8 -*-
"""skills/gongwen-grassroots/scripts/notice_check.py 的单元测试。

运行：python -m unittest discover -s tests/grassroots -t .
必过标本：规范的居民通知、倡议书、应急通知、村务公开（含 7 天低保初审公示）、单独的低保初审公示、
群消息（含按文件名识别）、发给村（社区）的工作通知、没有标题直接写称呼的提示、倡导语气的防火提示、
中文数字日期和“星期X（M月D日）”写对的，不报 P0、P1。
必错标本：星期写错（含中文数字日期、星期在前）、日期不存在、缺要素、倡议书强制措辞、群消息超长、
公开期限不足（含标题带起止日期、含低保子项、末尾空行后的“公示时间”）、低保公示期不足 7 天、
公示含病情、预警无来源、只有金额没有电话、村委会通知取消分红资格、签责任书、“不得进山”（NTC-018）。
期限恰好 15 日、低保恰好 7 日的报 NTC-019（P2）。
"""
from __future__ import annotations

import io
import json
import shutil
import sys
import tempfile
import unittest
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[2]
SCRIPTS_DIR = REPO_ROOT / "skills" / "gongwen-grassroots" / "scripts"
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

import notice_check  # noqa: E402

try:
    import docx  # type: ignore
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False


# ---------------------------------------------------------------------------
# 必过标本
# ---------------------------------------------------------------------------

PASS_NOTICE = """停水通知

各位居民：

因主管道检修，自来水公司定于10月20日（星期二）8:00至17:00停水，影响3号楼至8号楼。

停水期间，小区东门有两辆供水车，请自带水桶取水。来水后可能有短时浑水，放一会儿就好。

家里有老人行动不便的，请打社区电话，志愿者送水上门。

联系人：王某，电话：0571-86000000

某社区居民委员会
2026年10月16日
"""

PASS_INITIATIVE = """文明养犬倡议书

各位居民朋友：

上个月，社区和物业共接到遛狗不牵绳的投诉12起。为了让大家住得更安心，我们倡议：

一、出门遛狗请牵好犬绳，主动避让老人和孩子。

二、犬只粪便请随手清理，各楼栋门口有拾便袋。

三、电梯里请抱起小型犬，大型犬请等下一趟。

让我们一起做文明养犬的好邻居。

某社区居民委员会　某小区业主委员会
2026年10月16日
"""

PASS_EMERGENCY = """关于做好10月17日至18日强降雨防范应对工作的紧急通知

各村（社区），镇属各单位：

县气象台10月16日16时发布暴雨橙色预警信号。现就做好本轮强降雨防范应对工作通知如下：

一、立即开展隐患巡查。各村要组织人员对地质灾害隐患点、山塘水库开展巡查，发现险情立即报告镇防汛办。

二、提前转移危险区域人员。后山村、柳溪村要在10月17日12时前完成地质灾害隐患点住户转移。

三、严格值班值守。镇村两级实行24小时值班，值班电话：0571-86000001。

某镇防汛抗旱指挥部办公室
2026年10月16日17时
"""

PASS_WUGONGKAI = """某村2026年第三季度村务公开

一、村民代表会议决定事项
8月20日村民代表会议决定修建村东头便民桥，预算15万元，资金来源待定。

二、财务收支情况（7月1日至9月30日）
本期支出合计8040.00元，明细见公开栏附表。

公开时间：2026年10月10日至10月25日
监督电话：0571-86000002

某村民委员会
2026年10月10日
"""

PASS_GROUP = """【停水提醒】10月20日（星期二）8:00至17:00，3号楼至8号楼停水检修。小区东门有供水车，请自带水桶。来水后可能有短时浑水，放一会儿就好。有困难请打社区电话0571-86000000。某社区居委会"""

PASS_WUGONGKAI_WITH_DIBAO = """某村2026年第三季度村务公开

一、财务收支情况
本期支出合计8040.00元，明细见公开栏附表。

二、低保初审公示
经镇政府初审，拟将1组王某（家庭人口2人）纳入最低生活保障。
公示时间：10月10日至10月16日

公开时间：2026年10月10日至10月25日
监督电话：村务监督委员会 0571-86000002

某村民委员会
2026年10月10日
"""

PASS_DIBAO_STANDALONE = """低保初审对象公示

各位村民：

经镇政府初审，拟将1组王某（家庭人口2人）纳入最低生活保障。公示期为2026年10月10日至10月16日，共7天。对公示对象有异议的，请在公示期内反映。

监督电话：村务监督委员会 0571-86000002；镇民政办 0571-86000003

某村民委员会
2026年10月10日
"""

PASS_NO_TITLE = """各位居民：

10月20日（星期二）8:00至17:00，3号楼至8号楼停水检修。小区东门有供水车，请自带水桶。联系电话：0571-86000000。

某社区居民委员会
2026年10月16日
"""

PASS_INTERNAL = """关于做好国庆节前安全检查工作的通知

各村（社区），镇属各单位：

为做好国庆节前安全生产工作，各村（社区）要对辖区内的燃气使用场所、经营性自建房、在建工地、人员密集场所开展一次全面排查，排查情况于9月25日前报镇安监站，联系电话：0571-86000003。

某镇人民政府
2026年9月18日
"""


# ---------------------------------------------------------------------------
# 必错标本
# ---------------------------------------------------------------------------

FAIL_WEEKDAY = """停水通知

各位居民：

10月20日（星期一）8:00至17:00停水检修，影响3号楼至8号楼。联系电话：0571-86000000。

某社区居民委员会
2026年10月16日
"""

FAIL_NO_DATE_EXISTS = """接种通知

各位居民：

2月30日（星期一）上午在社区卫生服务站接种疫苗。联系电话：0571-86000000。

某社区居民委员会
2026年2月10日
"""

FAIL_MISSING_ELEMENTS = """关于开展垃圾分类宣传活动的通知

社区将开展垃圾分类集中宣传活动，请大家积极参加。
"""

FAIL_INITIATIVE = """文明养犬倡议书

各位居民：

严禁不牵绳遛狗，违者一律通报批评，并在业主群曝光。

某社区居民委员会
2026年10月16日
"""

FAIL_GROUP = "【通知】" + "各位居民请注意，社区将于10月20日（星期二）开展消防安全检查，请大家提前清理楼道杂物。" * 8 + "电话：0571-86000000"

FAIL_WUGONGKAI_DAYS = """某村2026年第三季度村务公开

一、财务收支情况
本期支出合计8040.00元。

公示期为7天。
监督电话：0571-86000002

某村民委员会
2026年10月10日
"""

FAIL_WUGONGKAI_RANGE = """某村2026年第三季度村务公开

一、财务收支情况
本期支出合计8040.00元。

公开时间：2026年10月10日至10月16日
监督电话：0571-86000002

某村民委员会
2026年10月10日
"""

FAIL_WUGONGKAI_DIBAO_SHORT_DOC = PASS_WUGONGKAI_WITH_DIBAO.replace(
    "公开时间：2026年10月10日至10月25日", "公开时间：2026年10月10日至10月19日")

FAIL_WUGONGKAI_TITLE_RANGE = """某村第三季度（7月1日至9月30日）村务公开

一、财务收支情况
本期支出合计8040.00元。

公开时间：2026年10月10日至10月19日
监督电话：0571-86000002

某村民委员会
2026年10月10日
"""

FAIL_DIBAO_SHORT = PASS_DIBAO_STANDALONE.replace("10月16日，共7天", "10月14日，共5天")

FAIL_AMOUNT_NOT_PHONE = """便民桥建设告知书

各位村民：

村东头便民桥定于11月开工，总投资1500000元，施工期间请绕行村西道路。

某村民委员会
2026年10月16日
"""

FAIL_PUBLICITY_SENSITIVE = """低保初审对象公示

各位村民：

经乡镇初审，拟将以下家庭纳入最低生活保障：1组王某，家庭人口2人，患尿毒症长期透析，家庭住址1组12号。

监督电话：0571-86000002

某村民委员会
2026年10月10日
"""

FAIL_PREWARN_NO_SOURCE = """暴雨天气安全提示

各位村民：

明天有暴雨橙色预警，请大家不要去河边，联系电话：0571-86000001。

某村民委员会
2026年10月16日
"""

FAIL_PREWARN_NO_TIME = """暴雨天气安全提示

各位村民：

县气象台发布暴雨橙色预警信号，请大家不要去河边，联系电话：0571-86000001。

某村民委员会
2026年10月16日
"""

FAIL_JARGON = """垃圾分类宣传活动通知

广大居民群众：

兹定于10月25日（星期日）上午9点在社区广场开展垃圾分类宣传活动，届时请务必准时参加，配合做好相关事宜。联系电话：0571-86000000。

某社区居民委员会
2026年10月16日
"""

# 评测 9 材料中 legal 类漏报的三句（村委会落款）
FAIL_VILLAGE_COERCION = """关于加强秋冬季森林防火工作的通知

全体村民：

一、外来人员进山一律登记，未登记的不得进山。

二、各户要签订森林防火责任书，拒不签订的，取消其年度村集体分红资格。

联系电话：0571-86000001

某村民委员会
2026年10月15日
"""

# 倡导语气的必过标本
PASS_FIRE_ADVICE = """秋冬森林防火提示

各位村民：

县人民政府公布的森林防火期从11月1日开始。防火期内请不要在山上烧荒、烧秸秆，上坟请用鲜花代替烧纸。大风天请不要进山。发现火情请马上打119，或打村委会电话0571-86000001。

某村民委员会
2026年10月16日
"""

# 小节之后空一行写在末尾的“公示时间”，是整份公开的期限
FAIL_WUGONGKAI_TAIL_PERIOD = """某村2026年第三季度村务公开

一、财务收支情况
本期支出合计8040.00元，明细见公开栏附表。

二、低保初审公示
经镇政府初审，拟将1组王某（家庭人口2人）纳入最低生活保障。

公示时间：10月10日至10月20日

监督电话：0571-86000002

某村民委员会
2026年10月10日
"""

FAIL_NO_TIME_ONLY_SIGNOFF_DATE = """关于开展垃圾分类宣传活动的通知

各位居民：

社区将开展垃圾分类集中宣传活动，请大家积极参加。联系电话：0571-86000000。

某社区居民委员会
2026年10月16日
"""


class NoticeCheckTestCase(unittest.TestCase):

    def setUp(self) -> None:
        self.tmp = Path(tempfile.mkdtemp(prefix="notice_check_test_"))

    def tearDown(self) -> None:
        shutil.rmtree(self.tmp, ignore_errors=True)

    def write(self, name: str, text: str, encoding: str = "utf-8") -> Path:
        path = self.tmp / name
        path.write_bytes(text.encode(encoding))
        return path

    def run_main(self, *args: str):
        out, err = io.StringIO(), io.StringIO()
        with redirect_stdout(out), redirect_stderr(err):
            code = notice_check.main(list(args))
        return code, out.getvalue(), err.getvalue()

    def check(self, text: str, *extra: str, name: str = "正文.txt") -> dict:
        path = self.write(name, text)
        code, out, _err = self.run_main(str(path), "--json", *extra)
        self.assertEqual(code, 0)
        return json.loads(out)

    @staticmethod
    def ids(result: dict, level: str = None):
        return [i["id"] for i in result["issues"] if level is None or i["level"] == level]

    def assert_clean(self, result: dict) -> None:
        serious = [i for i in result["issues"] if i["level"] in ("P0", "P1")]
        self.assertEqual(serious, [], json.dumps(serious, ensure_ascii=False, indent=2))


class PassSamplesTest(NoticeCheckTestCase):

    def test_resident_notice(self):
        result = self.check(PASS_NOTICE)
        self.assertEqual(result["kind"], "通知")
        self.assertEqual(result["audience"], "群众")
        self.assert_clean(result)
        self.assertEqual(self.ids(result), [])

    def test_initiative(self):
        result = self.check(PASS_INITIATIVE)
        self.assertEqual(result["kind"], "倡议书")
        self.assert_clean(result)

    def test_emergency_work_notice(self):
        result = self.check(PASS_EMERGENCY)
        self.assertEqual(result["kind"], "应急通知")
        self.assertEqual(result["audience"], "工作对象")
        self.assert_clean(result)

    def test_wugongkai(self):
        result = self.check(PASS_WUGONGKAI)
        self.assertEqual(result["kind"], "公示")
        self.assert_clean(result)

    def test_group_message(self):
        result = self.check(PASS_GROUP, "--kind", "群消息")
        self.assertEqual(result["kind"], "群消息")
        self.assert_clean(result)

    def test_wugongkai_with_dibao_subitem(self):
        result = self.check(PASS_WUGONGKAI_WITH_DIBAO)
        self.assertEqual(result["kind"], "公示")
        self.assertNotIn("NTC-012", self.ids(result))
        self.assertNotIn("NTC-017", self.ids(result))
        self.assert_clean(result)

    def test_dibao_standalone_with_supervision_committee(self):
        # 监督电话写“村务监督委员会”，不能当成村（居）务公开按 15 日判断
        result = self.check(PASS_DIBAO_STANDALONE)
        self.assertNotIn("NTC-012", self.ids(result))
        self.assertNotIn("NTC-017", self.ids(result))
        self.assert_clean(result)

    def test_no_title_salutation_first(self):
        result = self.check(PASS_NO_TITLE)
        self.assertNotIn("NTC-001", self.ids(result))
        self.assert_clean(result)

    def test_internal_notice_skips_plain_language(self):
        result = self.check(PASS_INTERNAL)
        self.assertEqual(result["audience"], "工作对象")
        self.assertNotIn("NTC-007", self.ids(result))
        self.assertNotIn("NTC-008", self.ids(result))
        self.assert_clean(result)

    def test_todo_placeholders_count_as_present(self):
        text = PASS_NOTICE.replace("0571-86000000", "【待补：联系电话】")
        result = self.check(text)
        self.assertNotIn("NTC-003", self.ids(result))

    def test_draft_note_is_ignored(self):
        text = PASS_NOTICE + "\n【起草说明】\n一、所作假设\n1. 兹定于、务必等词语只出现在说明里。\n"
        result = self.check(text)
        self.assertEqual(self.ids(result), [])
        self.assertTrue(any("起草说明" in n for n in result["notes"]))

    def test_correct_saturday(self):
        text = PASS_NOTICE.replace("10月20日（星期二）", "11月7日（星期六）")
        result = self.check(text)
        self.assertNotIn("NTC-006", self.ids(result))

    def test_advice_tone_is_not_coercion(self):
        # “请不要进山”是倡导语气，不报 NTC-018
        result = self.check(PASS_FIRE_ADVICE)
        self.assertNotIn("NTC-018", self.ids(result))
        self.assert_clean(result)

    def test_internal_notice_skips_coercion_check(self):
        # 镇发村的工作通知不是面向群众的文书，NTC-018 不查（转嫁责任由 gongwen_lint jianfu 类查）
        text = PASS_INTERNAL.replace("开展一次全面排查", "开展一次全面排查，并签订安全生产责任书")
        result = self.check(text)
        self.assertNotIn("NTC-018", self.ids(result))

    def test_chinese_numeral_date_correct(self):
        text = PASS_NOTICE.replace("10月20日（星期二）", "十月二十日（星期二）")
        result = self.check(text)
        self.assertNotIn("NTC-006", self.ids(result))

    def test_weekday_before_date_correct(self):
        text = PASS_NOTICE.replace("10月20日（星期二）", "星期二（10月20日）")
        result = self.check(text)
        self.assertNotIn("NTC-006", self.ids(result))
        self.assertFalse(any("星期未核对" in n for n in result["notes"]))

    def test_group_message_by_filename(self):
        result = self.check(PASS_GROUP, name="正文-群消息.txt")
        self.assertEqual(result["kind"], "群消息")
        self.assertEqual(result["kind_source"], "文件名推断")
        for issue_id in ("NTC-001", "NTC-004", "NTC-005"):
            self.assertNotIn(issue_id, self.ids(result))
        self.assert_clean(result)

    def test_dibao_eight_day_range_is_ok(self):
        # 10月10日至10月17日：开始当日不计入时正好 7 天，不报
        text = PASS_DIBAO_STANDALONE.replace("10月16日，共7天", "10月17日")
        result = self.check(text)
        self.assertNotIn("NTC-017", self.ids(result))
        self.assertNotIn("NTC-019", self.ids(result))


class FailSamplesTest(NoticeCheckTestCase):

    def test_weekday_mismatch_is_p0(self):
        result = self.check(FAIL_WEEKDAY)
        self.assertIn("NTC-006", self.ids(result, "P0"))
        issue = [i for i in result["issues"] if i["id"] == "NTC-006"][0]
        self.assertIn("星期二", issue["message"])

    def test_nonexistent_date_is_p0(self):
        result = self.check(FAIL_NO_DATE_EXISTS)
        self.assertIn("NTC-006", self.ids(result, "P0"))

    def test_missing_elements(self):
        result = self.check(FAIL_MISSING_ELEMENTS)
        found = set(self.ids(result, "P1"))
        for issue_id in ("NTC-001", "NTC-002", "NTC-003", "NTC-004", "NTC-005"):
            self.assertIn(issue_id, found)

    def test_signoff_date_is_not_event_time(self):
        result = self.check(FAIL_NO_TIME_ONLY_SIGNOFF_DATE)
        self.assertIn("NTC-002", self.ids(result, "P1"))

    def test_initiative_coercive_words(self):
        result = self.check(FAIL_INITIATIVE)
        words = {i["excerpt"] for i in result["issues"] if i["id"] == "NTC-009"}
        self.assertTrue({"严禁", "一律", "通报批评", "曝光"} <= words, words)

    def test_group_message_too_long(self):
        result = self.check(FAIL_GROUP, "--kind", "群消息")
        self.assertIn("NTC-010", self.ids(result, "P1"))

    def test_wugongkai_period_days(self):
        result = self.check(FAIL_WUGONGKAI_DAYS)
        self.assertIn("NTC-012", self.ids(result, "P1"))

    def test_wugongkai_period_range(self):
        result = self.check(FAIL_WUGONGKAI_RANGE)
        issue = [i for i in result["issues"] if i["id"] == "NTC-012"]
        self.assertTrue(issue)
        self.assertIn("7 天", issue[0]["message"])

    def test_wugongkai_short_with_dibao_subitem(self):
        result = self.check(FAIL_WUGONGKAI_DIBAO_SHORT_DOC)
        issues = [i for i in result["issues"] if i["id"] == "NTC-012"]
        self.assertEqual(len(issues), 1, issues)
        self.assertIn("10 天", issues[0]["message"])
        self.assertNotIn("NTC-017", self.ids(result))

    def test_wugongkai_title_range_is_not_period(self):
        result = self.check(FAIL_WUGONGKAI_TITLE_RANGE)
        issues = [i for i in result["issues"] if i["id"] == "NTC-012"]
        self.assertEqual(len(issues), 1, issues)
        self.assertIn("10 天", issues[0]["message"])

    def test_dibao_period_too_short(self):
        result = self.check(FAIL_DIBAO_SHORT)
        self.assertIn("NTC-017", self.ids(result, "P1"))
        self.assertNotIn("NTC-012", self.ids(result))

    def test_wugongkai_only_dibao_period(self):
        text = PASS_WUGONGKAI_WITH_DIBAO.replace("公开时间：2026年10月10日至10月25日\n", "")
        result = self.check(text)
        self.assertIn("NTC-011", self.ids(result, "P2"))
        self.assertNotIn("NTC-012", self.ids(result))

    def test_amount_is_not_phone(self):
        result = self.check(FAIL_AMOUNT_NOT_PHONE)
        self.assertIn("NTC-003", self.ids(result, "P1"))

    def test_publicity_missing_period(self):
        text = PASS_WUGONGKAI.replace("公开时间：2026年10月10日至10月25日\n", "")
        result = self.check(text)
        self.assertIn("NTC-011", self.ids(result, "P1"))

    def test_publicity_sensitive_info(self):
        result = self.check(FAIL_PUBLICITY_SENSITIVE)
        words = {i["excerpt"] for i in result["issues"] if i["id"] == "NTC-013"}
        self.assertIn("透析", words)
        self.assertIn("家庭住址", words)

    def test_prewarning_without_source(self):
        result = self.check(FAIL_PREWARN_NO_SOURCE)
        self.assertIn("NTC-014", self.ids(result, "P1"))

    def test_prewarning_without_time(self):
        result = self.check(FAIL_PREWARN_NO_TIME)
        self.assertIn("NTC-015", self.ids(result, "P2"))
        self.assertNotIn("NTC-014", self.ids(result))

    def test_jargon_for_public(self):
        result = self.check(FAIL_JARGON)
        words = {i["excerpt"] for i in result["issues"] if i["id"] == "NTC-008"}
        self.assertTrue({"兹定于", "届时", "务必", "相关事宜", "广大居民群众", "配合做好"} <= words, words)

    def test_repealed_regulation(self):
        text = """秋冬森林防火提示

各位村民：

根据《森林防火条例》，10月20日（星期二）起进入森林防火期，请不要在山上烧荒。电话：0571-86000001。

某村民委员会
2026年10月16日
"""
        result = self.check(text)
        self.assertIn("NTC-016", self.ids(result, "P1"))
        clean = text.replace("《森林防火条例》", "《森林草原防灭火条例》")
        self.assertNotIn("NTC-016", self.ids(self.check(clean)))

    def test_village_coercion_is_p0(self):
        result = self.check(FAIL_VILLAGE_COERCION)
        excerpts = {i["excerpt"] for i in result["issues"] if i["id"] == "NTC-018" and i["level"] == "P0"}
        self.assertIn("不得进山", excerpts)
        self.assertIn("签订森林防火责任书", excerpts)
        self.assertTrue(any(e.startswith("取消") and "分红资格" in e for e in excerpts), excerpts)

    def test_township_coercion_is_p1(self):
        text = FAIL_VILLAGE_COERCION.replace("某村民委员会", "某镇人民政府")
        result = self.check(text)
        self.assertIn("NTC-018", self.ids(result, "P1"))
        self.assertNotIn("NTC-018", self.ids(result, "P0"))
        issue = [i for i in result["issues"] if i["id"] == "NTC-018"][0]
        self.assertIn("村（社区）的按 P0", issue["message"])

    def test_coercion_with_basis_is_p1(self):
        # 村委会转述上级政府命令：只提示核对依据，不直接定为 P0
        text = FAIL_VILLAGE_COERCION.replace(
            "一、外来人员进山一律登记，未登记的不得进山。",
            "一、根据县人民政府封山令，防火期内未登记的不得进山。")
        result = self.check(text)
        hits = [i for i in result["issues"] if i["id"] == "NTC-018" and i["excerpt"] == "不得进山"]
        self.assertEqual([i["level"] for i in hits], ["P1"])
        self.assertIn("核对依据原文", hits[0]["message"])

    def test_group_message_coercion_signed_by_community(self):
        text = "【提醒】未签订承诺书的住户，停发本年度福利。电话0571-86000000。某社区居委会"
        result = self.check(text, "--kind", "群消息")
        self.assertIn("NTC-018", self.ids(result, "P0"))

    def test_chinese_numeral_weekday_mismatch(self):
        text = FAIL_WEEKDAY.replace("10月20日（星期一）", "十月二十日（星期一）")
        result = self.check(text)
        self.assertIn("NTC-006", self.ids(result, "P0"))

    def test_weekday_before_date_mismatch(self):
        text = FAIL_WEEKDAY.replace("10月20日（星期一）", "星期一（10月20日）")
        result = self.check(text)
        self.assertIn("NTC-006", self.ids(result, "P0"))

    def test_lone_weekday_is_noted(self):
        text = PASS_NOTICE.replace("家里有老人行动不便的", "星期三起家里有老人行动不便的")
        result = self.check(text)
        self.assertTrue(any("1 处星期未核对" in n for n in result["notes"]), result["notes"])

    def test_tail_period_after_blank_line_is_doc_level(self):
        result = self.check(FAIL_WUGONGKAI_TAIL_PERIOD)
        issues = [i for i in result["issues"] if i["id"] == "NTC-012"]
        self.assertEqual(len(issues), 1, issues)
        self.assertIn("11 天", issues[0]["message"])
        self.assertNotIn("NTC-017", self.ids(result))

    def test_wugongkai_exactly_15_days_boundary(self):
        text = PASS_WUGONGKAI.replace("10月10日至10月25日", "10月10日至10月24日")
        result = self.check(text)
        self.assertIn("NTC-019", self.ids(result, "P2"))
        self.assertNotIn("NTC-012", self.ids(result))
        self.assertNotIn("NTC-019", self.ids(self.check(PASS_WUGONGKAI)))

    def test_dibao_exactly_7_days_boundary(self):
        result = self.check(PASS_DIBAO_STANDALONE)
        self.assertIn("NTC-019", self.ids(result, "P2"))
        self.assertNotIn("NTC-017", self.ids(result))

    def test_strict_exit_code(self):
        path = self.write("星期错误.txt", FAIL_WEEKDAY)
        code, _out, _err = self.run_main(str(path), "--strict")
        self.assertEqual(code, 1)
        path_ok = self.write("规范.txt", PASS_NOTICE)
        code, _out, _err = self.run_main(str(path_ok), "--strict")
        self.assertEqual(code, 0)


class BehaviourTest(NoticeCheckTestCase):

    def test_kind_detection(self):
        cases = {
            "关于做好防汛工作的紧急通知": "应急通知",
            "文明养犬倡议书": "倡议书",
            "房屋安全隐患告知书": "告知书",
            "低保初审对象公示": "公示",
            "停水通知": "通知",
        }
        for title, kind in cases.items():
            detected, _source = notice_check.detect_kind(title, None)
            self.assertEqual(detected, kind, title)
        self.assertEqual(notice_check.detect_kind("停水通知", "群消息"), ("群消息", "参数"))

    def test_cross_year_weekday(self):
        text = """元旦期间安全提示

各位居民：

1月1日（星期五）至1月3日社区服务中心值班，值班电话：0571-86000000。

某社区居民委员会
2026年12月28日
"""
        # 2027年1月1日是星期五：成文日期在12月，事项在次年1月
        result = self.check(text)
        self.assertNotIn("NTC-006", self.ids(result))

    def test_year_argument(self):
        text = """活动通知

各位居民：

10月20日（星期二）上午9点在社区广场举办活动。电话：0571-86000000。

某社区居民委员会
"""
        result = self.check(text, "--year", "2026")
        self.assertNotIn("NTC-006", self.ids(result))
        result = self.check(text, "--year", "2025")
        self.assertIn("NTC-006", self.ids(result))

    def test_text_output(self):
        path = self.write("正文.txt", FAIL_WEEKDAY)
        code, out, _err = self.run_main(str(path))
        self.assertEqual(code, 0)
        self.assertIn("NTC-006", out)
        self.assertIn("P0 1 条", out)

    def test_gb18030_input(self):
        path = self.write("gbk.txt", PASS_NOTICE, encoding="gb18030")
        code, out, _err = self.run_main(str(path), "--json")
        self.assertEqual(code, 0)
        result = json.loads(out)
        self.assertTrue(any("GB18030" in n for n in result["notes"]))

    def test_missing_file(self):
        code, _out, err = self.run_main(str(self.tmp / "不存在.txt"))
        self.assertEqual(code, 2)
        self.assertIn("未找到文件", err)

    def test_unsupported_suffix(self):
        path = self.write("正文.pdf", "内容")
        code, _out, err = self.run_main(str(path))
        self.assertEqual(code, 2)
        self.assertIn("不支持", err)

    def test_empty_body(self):
        path = self.write("空.txt", "\n\n【起草说明】\n说明\n")
        code, _out, err = self.run_main(str(path))
        self.assertEqual(code, 2)
        self.assertIn("正文为空", err)

    def test_bad_year(self):
        path = self.write("正文.txt", PASS_NOTICE)
        code, _out, _err = self.run_main(str(path), "--year", "26")
        self.assertEqual(code, 2)

    def test_no_files_written(self):
        path = self.write("正文.txt", PASS_NOTICE)
        before = sorted(p.name for p in self.tmp.iterdir())
        self.run_main(str(path), "--json")
        after = sorted(p.name for p in self.tmp.iterdir())
        self.assertEqual(before, after)

    @unittest.skipUnless(HAS_DOCX, "未安装 python-docx")
    def test_docx_input(self):
        document = docx.Document()
        for line in FAIL_WEEKDAY.split("\n"):
            document.add_paragraph(line)
        path = self.tmp / "正文.docx"
        document.save(str(path))
        code, out, _err = self.run_main(str(path), "--json")
        self.assertEqual(code, 0)
        self.assertIn("NTC-006", [i["id"] for i in json.loads(out)["issues"]])


if __name__ == "__main__":
    unittest.main()
