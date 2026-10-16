# -*- coding: utf-8 -*-
"""tests/data 的公共工具：临时目录、写 CSV 和 XLSX 标本、调用脚本入口。"""
from __future__ import annotations

import io
import json
import shutil
import sys
import tempfile
import unittest
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path
from typing import Iterable, List, Optional, Sequence

REPO_ROOT = Path(__file__).resolve().parents[2]
SCRIPTS_DIR = REPO_ROOT / "skills" / "gongwen-data" / "scripts"
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

import data_tool  # noqa: E402

try:
    import openpyxl  # type: ignore  # noqa: F401
    HAS_OPENPYXL = True
except ImportError:  # pragma: no cover
    HAS_OPENPYXL = False

try:
    import matplotlib  # type: ignore  # noqa: F401
    HAS_MPL = True
except ImportError:
    HAS_MPL = False


class DataTestCase(unittest.TestCase):

    def setUp(self) -> None:
        # 临时目录用短前缀，避免 Windows 路径超长
        self.tmp = Path(tempfile.mkdtemp(prefix="gwd_"))

    def tearDown(self) -> None:
        shutil.rmtree(self.tmp, ignore_errors=True)

    # ---- 写标本 ----
    def csv(self, name: str, rows: Iterable[Sequence], encoding: str = "utf-8") -> Path:
        path = self.tmp / name
        lines = []
        for row in rows:
            cells = []
            for v in row:
                text = "" if v is None else str(v)
                if any(ch in text for ch in ',"\n'):
                    text = '"' + text.replace('"', '""') + '"'
                cells.append(text)
            lines.append(",".join(cells))
        path.write_bytes(("\n".join(lines) + "\n").encode(encoding))
        return path

    def xlsx(self, name: str, sheets: dict, merges: Optional[dict] = None, formats: Optional[dict] = None,
             hidden_rows: Optional[dict] = None) -> Path:
        """sheets：{工作表名: [[行], ...]}；merges：{工作表名: ["A1:C1"]}；formats：{工作表名: {"B3": "0.0%"}}。"""
        wb = openpyxl.Workbook()
        first = True
        for title, rows in sheets.items():
            ws = wb.active if first else wb.create_sheet()
            first = False
            ws.title = title
            for row in rows:
                ws.append(list(row))
            for ref in (merges or {}).get(title, []):
                ws.merge_cells(ref)
            for ref, fmt in (formats or {}).get(title, {}).items():
                ws[ref].number_format = fmt
            for r in (hidden_rows or {}).get(title, []):
                ws.row_dimensions[r].hidden = True
        path = self.tmp / name
        wb.save(str(path))
        return path

    # ---- 调用 ----
    def run_tool(self, module, *args: str):
        out, err = io.StringIO(), io.StringIO()
        with redirect_stdout(out), redirect_stderr(err):
            code = module.main([str(a) for a in args])
        return code, out.getvalue(), err.getvalue()

    def run_json(self, module, *args: str, expect_code: int = 0) -> dict:
        code, out, err = self.run_tool(module, *args, "--json")
        self.assertEqual(code, expect_code, msg=f"stdout={out}\nstderr={err}")
        return json.loads(out)

    def inspect(self, path: Path, *extra: str) -> dict:
        data = self.run_json(data_tool, "inspect", str(path), *extra)
        return data["tables"][0]

    @staticmethod
    def ids(issues: List[dict]) -> List[str]:
        return [i["id"] for i in issues]

    def issue(self, issues: List[dict], issue_id: str) -> dict:
        for item in issues:
            if item["id"] == issue_id:
                return item
        self.fail(f"没有找到问题 {issue_id}，现有：{self.ids(issues)}")
        return {}
