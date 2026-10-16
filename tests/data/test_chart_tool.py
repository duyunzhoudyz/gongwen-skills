# -*- coding: utf-8 -*-
"""skills/gongwen-data/scripts/chart_tool.py 的单元测试。

运行：python -m unittest discover -s tests/data -t .
必错标本：不同单位同图、饼图含负数或比率、列名不存在；必过标本：各类图表正常生成且只用黑白灰。
"""
from __future__ import annotations

import re
import unittest
import zipfile
from pathlib import Path

from tests.data.helpers import HAS_MPL, HAS_OPENPYXL, DataTestCase

if HAS_OPENPYXL:
    import chart_tool  # noqa: E402
    import openpyxl  # type: ignore  # noqa: E402

MONTHLY = [["某区政务服务中心2026年1—9月办件量"], ["单位：件"], ["月份", "2026年", "2025年"],
           ["1月", 8320, 7905], ["2月", 6110, 6352], ["3月", 9480, 8870],
           ["注：数据来源为区政务服务平台统计"]]
TOWNS = [["各镇网办率"], ["镇街", "网办率（%）", "办件量（件）"],
         ["城关街道", 82.5, 3200], ["东湖镇", 76.1, 1800], ["西山镇", 68.4, 900], ["合计", None, 5900]]


def chart_xml(path: str) -> str:
    with zipfile.ZipFile(path) as zf:
        names = [n for n in zf.namelist() if n.startswith("xl/charts/chart")]
        return "".join(zf.read(n).decode("utf-8") for n in names)


@unittest.skipUnless(HAS_OPENPYXL, "需要 openpyxl")
class ChartTest(DataTestCase):

    def chart(self, rows, *args: str, expect_code: int = 0) -> dict:
        path = self.csv("data.csv", rows)
        return self.run_json(chart_tool, str(path), "--out", str(self.tmp / "图表.xlsx"), *args,
                             expect_code=expect_code)

    def assert_grayscale(self, xml: str) -> None:
        colors = re.findall(r'srgbClr val="([0-9A-Fa-f]{6})"', xml)
        self.assertTrue(colors, "图表中没有颜色定义")
        for c in colors:
            self.assertTrue(c[0:2].upper() == c[2:4].upper() == c[4:6].upper(), f"出现非灰色：{c}")

    def test_pass_bar_chart(self):
        data = self.chart(TOWNS, "--type", "bar", "--y", "办件量", "--sort", "desc", "--source", "区政务服务平台")
        self.assertEqual(data["categories"], ["城关街道", "东湖镇", "西山镇"])  # 合计行不入图
        self.assertEqual(data["unit"], "件")
        xml = chart_xml(data["out"])
        self.assert_grayscale(xml)
        self.assertIn('<min val="0"', xml)  # 纵轴从 0 开始
        self.assertIn('<delete val="0"', xml)  # 坐标轴不隐藏
        self.assertIn("单位：件", xml)
        ws = openpyxl.load_workbook(data["out"])["图表"]
        texts = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
        self.assertIn("数据来源：区政务服务平台", texts)
        self.assertTrue(any("合计" in w for w in data["warnings"]))

    def test_pass_line_chart_two_series(self):
        data = self.chart(MONTHLY, "--type", "line", "--y", "2026年,2025年")
        self.assertEqual(data["source"], "区政务服务平台统计")
        self.assertEqual(data["unit"], "件")
        xml = chart_xml(data["out"])
        self.assert_grayscale(xml)
        self.assertIn('prstDash val="dash"', xml)  # 第二条线用虚线区分
        self.assertEqual(len(data["series"]), 2)

    def test_line_sort_ignored_and_non_time_warning(self):
        data = self.chart(MONTHLY, "--type", "line", "--y", "2026年", "--sort", "desc")
        self.assertEqual(data["categories"], ["1月", "2月", "3月"])
        self.assertTrue(any("忽略 --sort" in w for w in data["warnings"]))
        data = self.chart(TOWNS, "--type", "line", "--y", "办件量")
        self.assertTrue(any("折线图适合按时间" in w for w in data["warnings"]))

    def test_pass_pie_and_highlight(self):
        rows = [["项目", "营业收入（万元）"], ["工程", 52360], ["设计", 8640], ["运维", 12870], ["合计", 73870]]
        data = self.chart(rows, "--type", "pie")
        self.assert_grayscale(chart_xml(data["out"]))
        self.assertEqual(len(data["categories"]), 3)
        data = self.chart(TOWNS, "--type", "barh", "--y", "办件量", "--highlight", "东湖镇")
        self.assertIn('val="262626"', chart_xml(data["out"]))
        self.assertIn('orientation val="maxMin"', chart_xml(data["out"]))

    def test_pie_too_many_categories_warns(self):
        rows = [["项目", "金额（万元）"]] + [[f"项目{i}", i + 1] for i in range(8)]
        data = self.chart(rows, "--type", "pie")
        self.assertTrue(any("超过 6 类" in w for w in data["warnings"]))

    def test_fail_mixed_units(self):
        code, _out, err = self.run_tool(chart_tool, self.csv("d.csv", TOWNS), "--type", "bar",
                                        "--y", "网办率,办件量", "--out", self.tmp / "x.xlsx")
        self.assertEqual(code, 2)
        self.assertIn("不同单位", err)

    def test_fail_pie_negative_or_rate(self):
        rows = [["项目", "利润（万元）"], ["甲", 100], ["乙", -20]]
        self.assertEqual(self.run_tool(chart_tool, self.csv("n.csv", rows), "--type", "pie",
                                       "--out", self.tmp / "x.xlsx")[0], 2)
        self.assertEqual(self.run_tool(chart_tool, self.csv("r.csv", TOWNS), "--type", "pie", "--y", "网办率",
                                       "--out", self.tmp / "y.xlsx")[0], 2)

    def test_fail_missing_column_and_file(self):
        self.assertEqual(self.run_tool(chart_tool, self.csv("d.csv", TOWNS), "--type", "bar", "--y", "不存在")[0], 2)
        self.assertEqual(self.run_tool(chart_tool, self.tmp / "none.csv", "--type", "bar")[0], 2)

    def test_missing_values_and_source(self):
        rows = [["社区", "走访户数"], ["东社区", 12], ["西社区", "—"], ["南社区", 8]]
        data = self.chart(rows, "--type", "bar")
        self.assertIsNone(data["series"][0]["values"][1])
        self.assertTrue(any("不按 0 画" in w for w in data["warnings"]))
        self.assertTrue(any("数据来源" in w for w in data["warnings"]))
        ws = openpyxl.load_workbook(data["out"])["图表"]
        texts = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
        self.assertIn("数据来源：【待补：数据来源】", texts)

    def test_never_overwrite_and_top(self):
        out = self.tmp / "图表.xlsx"
        out.write_bytes(b"old")
        data = self.chart(TOWNS, "--type", "bar", "--y", "办件量", "--sort", "desc", "--top", "2")
        self.assertTrue(data["out"].endswith("图表-v2.xlsx"))
        self.assertEqual(out.read_bytes(), b"old")
        self.assertEqual(data["categories"], ["城关街道", "东湖镇"])

    def test_default_y_and_auto_title(self):
        """不给 --y：折线图画全部同单位数值列；柱状图只画第一列并提示其余列。不给 --title：按时段＋指标拼图题。"""
        data = self.chart(MONTHLY, "--type", "line")
        self.assertEqual([s["name"] for s in data["series"]], ["2026年", "2025年"])
        self.assertEqual(data["title"], "2026年与2025年1—9月办件量对比")
        self.assertTrue(any("自动生成" in w for w in data["warnings"]))
        rows = [["某医院2026年9月科室指标汇总表"], ["科室", "出院人次", "上年同月出院人次"],
                ["心内科", 128, 132], ["骨科", 80, 81], ["儿科", 273, 283]]
        data = self.chart(rows, "--type", "barh", "--sort", "desc", "--top", "2")
        self.assertEqual([s["name"] for s in data["series"]], ["出院人次"])
        self.assertTrue(any("未指定 --y" in w and "上年同月出院人次" in w for w in data["warnings"]))
        self.assertEqual(data["title"], "2026年9月各科室出院人次（前2位）")
        # 必过：给了 --title 就照用，不提示
        data = self.chart(rows, "--type", "bar", "--y", "出院人次", "--title", "图1　出院人次")
        self.assertEqual(data["title"], "图1　出院人次")
        self.assertFalse(any("自动生成" in w for w in data["warnings"]))

    def test_fail_pie_with_blank(self):
        """必错：饼图有空值时退出，不把空值画成0。"""
        rows = [["项目", "营业收入（万元）"], ["工程", 52360], ["设计", None], ["运维", 12870]]
        code, _out, err = self.run_tool(chart_tool, self.csv("pb.csv", rows), "--type", "pie",
                                        "--out", self.tmp / "x.xlsx")
        self.assertEqual(code, 2)
        self.assertIn("设计", err)

    def test_label_text_no_scientific(self):
        from decimal import Decimal
        self.assertEqual(chart_tool.label_text(286450000.0, 0), "286450000")
        self.assertEqual(chart_tool.label_text(1234567.25, 1), "1234567.3")  # 四舍五入，不用银行家舍入
        self.assertEqual(chart_tool.label_text(float("nan"), 1), "")
        self.assertEqual(chart_tool.label_places([[Decimal("286450000"), None]], False), 0)
        self.assertEqual(chart_tool.label_places([[Decimal("82.5")]], False), 1)
        self.assertEqual(chart_tool.label_places([[Decimal("82")]], True), 1)

    def test_png_branch_with_stub_matplotlib(self):
        """本机没有 matplotlib 时，用替身模块走一遍 PNG 分支，核对数据标签不用科学计数法、空值不标。"""
        from decimal import Decimal
        from unittest import mock
        fig, ax = mock.MagicMock(), mock.MagicMock()
        plt = mock.MagicMock()
        plt.subplots.return_value = (fig, ax)
        fm = mock.MagicMock()
        fm.fontManager.ttflist = []
        mpl = mock.MagicMock()
        mpl.pyplot, mpl.font_manager = plt, fm
        with mock.patch.dict("sys.modules", {"matplotlib": mpl, "matplotlib.pyplot": plt,
                                             "matplotlib.font_manager": fm}):
            warnings = []
            out = chart_tool.write_png(self.tmp / "图.xlsx", "bar", "图题", "元", "", "", ["收入"],
                                       ["甲", "乙", "丙"], [[Decimal("286450000"), None, Decimal("1500000")]],
                                       None, warnings)
            self.assertTrue(str(out).endswith("图.png"))
            labels = ax.bar_label.call_args.kwargs["labels"]
            self.assertEqual(labels, ["286450000", "", "1500000"])
            fig.savefig.assert_called_once()
            self.assertTrue(any("中文字体" in w for w in warnings))
            ax.reset_mock()
            chart_tool.write_png(self.tmp / "线.xlsx", "line", "图题", "%", "", "", ["网办率"], ["1月", "2月"],
                                 [[Decimal("82.5"), Decimal("80")]], None, [], True)
            texts = [c.args[0] for c in ax.annotate.call_args_list]
            self.assertEqual(texts, ["82.5", "80.0"])

    def test_png_optional(self):
        data = self.chart(TOWNS, "--type", "bar", "--y", "办件量", "--png")
        if HAS_MPL:
            self.assertTrue(data["png"] and Path(data["png"]).is_file())
        else:
            self.assertIsNone(data["png"])
            self.assertTrue(any("matplotlib" in w for w in data["warnings"]))
            self.assertTrue(Path(data["out"]).is_file())

    def test_fail_personal_info_as_category(self):
        """必错：姓名、手机号列不作分类（脱敏后会重名，不脱敏又泄露），退出码 2，不生成文件。"""
        rows = [["医生姓名", "手机号", "科室", "手术台次（台）"],
                ["张三丰", "13812345678", "骨科", 40], ["张小明", "13987654321", "眼科", 35],
                ["李四", "13700001111", "儿科", 20]]
        for x in ("医生姓名", "手机号"):
            with self.subTest(x=x):
                path = self.csv("p.csv", rows)
                code, _out, err = self.run_tool(chart_tool, path, "--type", "bar", "--x", x,
                                                "--out", self.tmp / "p.xlsx")
                self.assertEqual(code, 2)
                self.assertIn("个人信息列", err)
                self.assertFalse((self.tmp / "p.xlsx").exists())
        # 默认分类列（关键列）是姓名列时也不画
        code, _out, err = self.run_tool(chart_tool, self.csv("p.csv", rows), "--type", "bar",
                                        "--out", self.tmp / "p.xlsx")
        self.assertEqual(code, 2)
        # 必过：改用科室列可以出图
        data = self.chart(rows, "--type", "bar", "--x", "科室", "--y", "手术台次")
        self.assertNotIn("13812345678", str(data))

    def test_summary_rows_not_name_masked(self):
        """必过：关键列是姓名列时，提示中的“合计”不脱敏成“合某”。"""
        rows = [["姓名", "科室", "手术台次（台）"], ["张三丰", "骨科", 40], ["李四", "眼科", 20], ["合计", None, 60]]
        data = self.chart(rows, "--type", "bar", "--x", "科室", "--y", "手术台次")
        joined = "".join(data["warnings"])
        self.assertIn("合计", joined)
        self.assertNotIn("合某", joined)

    def test_fail_marked_material_intercepted(self):
        """必错：表注写“内部资料，不得外传”的，不出图，退出码 2；用户说明后加 --marking-ok 才出图。"""
        rows = TOWNS + [["注：内部资料，不得外传"]]
        path = self.csv("m.csv", rows)
        code, _out, err = self.run_tool(chart_tool, path, "--type", "bar", "--y", "办件量",
                                        "--out", self.tmp / "m.xlsx")
        self.assertEqual(code, 2)
        self.assertIn("DT-SEC-01", err)
        self.assertFalse((self.tmp / "m.xlsx").exists())
        data = self.run_json(chart_tool, str(path), "--type", "bar", "--y", "办件量", "--marking-ok",
                             "--out", str(self.tmp / "m.xlsx"))
        self.assertTrue(Path(data["out"]).is_file())


if __name__ == "__main__":
    unittest.main()
