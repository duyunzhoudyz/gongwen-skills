# -*- coding: utf-8 -*-
"""skills/gongwen-data/scripts/data_tool.py 的单元测试。

运行：python -m unittest discover -s tests/data -t .
每组检查都有"必错标本"（必须报出问题）和"必过标本"（不得误报）。
"""
from __future__ import annotations

import json
import unittest
from decimal import Decimal

from tests.data.helpers import HAS_OPENPYXL, DataTestCase, data_tool

HEADER = ["科室", "出院人次", "上年同月出院人次", "平均住院日（天）"]


def dept_rows(total_out=None):
    rows = [["某医院2026年9月科室指标表"], HEADER,
            ["心内科", 120, 110, 8.2],
            ["骨科", 95, 100, 9.6],
            ["儿科", 60, 58, 5.1]]
    rows.append(["合计", 275 if total_out is None else total_out, 268, 7.6])
    rows.append(["注：数据来源为病案首页统计"])
    return rows


# ---------------------------------------------------------------------------
# 单元格解析
# ---------------------------------------------------------------------------

class ParseCellTest(DataTestCase):

    def test_text_forms(self):
        cases = {
            "1,234": ("textnum", Decimal("1234"), "thousands"),
            "约20": ("textnum", Decimal("20"), "approx"),
            "30多": ("textnum", Decimal("30"), "approx"),
            "3500元": ("textnum", Decimal("3500"), "unit"),
            "38.2%": ("textnum", Decimal("38.2"), "percent"),
            "(123)": ("textnum", Decimal("-123"), "accounting"),
            "１２": ("textnum", Decimal("12"), "fullwidth"),
            "￥88": ("textnum", Decimal("88"), "currency"),
        }
        for raw, (kind, value, flag) in cases.items():
            with self.subTest(raw=raw):
                cell = data_tool.parse_cell(raw)
                self.assertEqual(cell.kind, kind)
                self.assertEqual(cell.value, value)
                self.assertIn(flag, cell.flags)

    def test_null_error_and_text(self):
        for raw in ("—", "/", "无", "N/A", "暂无", "-"):
            with self.subTest(raw=raw):
                self.assertEqual(data_tool.parse_cell(raw).kind, "null")
        self.assertEqual(data_tool.parse_cell("#DIV/0!").kind, "error")
        self.assertEqual(data_tool.parse_cell("心内科").kind, "text")
        self.assertEqual(data_tool.parse_cell("很多").kind, "text")  # "多"前面不是数字，不算约数
        self.assertTrue(data_tool.parse_cell(None).is_empty)

    def test_native_numbers(self):
        cell = data_tool.parse_cell(3216.0)
        self.assertEqual(str(cell.value), "3216")  # 整数浮点不带 .0
        cell = data_tool.parse_cell(0.382, "0.0%")
        self.assertEqual(cell.value, Decimal("38.2"))
        self.assertIn("percent", cell.flags)


# ---------------------------------------------------------------------------
# 体检
# ---------------------------------------------------------------------------

class InspectTotalsTest(DataTestCase):

    def test_pass_clean_table(self):
        """必过：合计正确、单位齐全的表不报 P0、P1。"""
        t = self.inspect(self.csv("ok.csv", dept_rows()))
        self.assertEqual(t["summary"]["P0"], 0)
        self.assertEqual(t["summary"]["P1"], 0)
        self.assertEqual(t["data_rows"], 3)
        self.assertEqual(t["periods"], ["2026年9月"])
        self.assertEqual(t["source_note"], "病案首页统计")
        statuses = {c["column"]: c["status"] for c in t["checks"] if "column" in c}
        self.assertEqual(statuses["出院人次"], "ok")
        self.assertEqual(statuses["平均住院日（天）"], "skipped")
        cols = {c["name"]: c for c in t["columns"]}
        self.assertEqual(cols["出院人次"]["unit"], "人次")
        self.assertEqual(cols["平均住院日（天）"]["kind"], "avg")

    def test_fail_total_mismatch(self):
        """必错：合计与明细不符报 P0，写明相差多少。"""
        t = self.inspect(self.csv("bad.csv", dept_rows(total_out=295)))
        issue = self.issue(t["issues"], "DT-SUM-01")
        self.assertEqual(issue["level"], "P0")
        self.assertIn("相差20", issue["message"])
        self.assertIn("第6行", issue["location"])

    def test_fail_subtotal_and_total(self):
        rows = [["项目", "金额（万元）"],
                ["甲", 10], ["乙", 20], ["内科小计", 31],
                ["丙", 5], ["丁", 5], ["外科小计", 10],
                ["合计", 41]]
        t = self.inspect(self.csv("sub.csv", rows))
        sub = self.issue(t["issues"], "DT-SUM-02")
        self.assertIn("内科小计", sub["location"])
        # 合计按"各小计"核对：31＋10＝41，相符
        self.assertNotIn("DT-SUM-01", self.ids(t["issues"]))

    def test_fail_top_total_row(self):
        rows = [["单位", "人数"], ["全市", 100], ["甲区", 40], ["乙区", 50]]
        t = self.inspect(self.csv("top.csv", rows))
        self.assertIn("DT-SUM-01", self.ids(t["issues"]))

    def test_rounding_tolerance_is_p2(self):
        rows = [["项目", "金额（万元）"], ["甲", 1.33], ["乙", 1.33], ["丙", 1.33], ["合计", 4.00]]
        t = self.inspect(self.csv("round.csv", rows))
        self.assertNotIn("DT-SUM-01", self.ids(t["issues"]))
        self.assertIn("DT-SUM-03", self.ids(t["issues"]))

    def test_rounding_tolerance_capped_at_two_units(self):
        """30 行 1 位小数：差 0.1 是修约（P2）；差 1.5 超过末位 2 个单位，不能算修约（P1）；差 2.0 超出累计上限（P0）。"""
        def rows(total):
            body = [[f"项目{i:02d}", "10.1"] for i in range(1, 31)]  # 明细相加 303.0
            return [["某公司2026年研发投入"], ["项目", "研发投入（万元）"]] + body + [["合计", total]]

        t = self.inspect(self.csv("r1.csv", rows("303.1")))
        self.assertEqual(self.issue(t["issues"], "DT-SUM-03")["level"], "P2")
        t = self.inspect(self.csv("r2.csv", rows("304.5")))
        ids = self.ids(t["issues"])
        self.assertNotIn("DT-SUM-03", ids)
        issue = self.issue(t["issues"], "DT-SUM-06")
        self.assertEqual(issue["level"], "P1")
        self.assertIn("相差1.5", issue["message"])
        t = self.inspect(self.csv("r3.csv", rows("305.0")))
        self.assertEqual(self.issue(t["issues"], "DT-SUM-01")["level"], "P0")

    def test_of_which_rows_excluded(self):
        rows = [["单位", "人数"], ["甲支部", 30], ["其中：党员", 12], ["乙支部", 20], ["合计", 50]]
        t = self.inspect(self.csv("ofwhich.csv", rows))
        self.assertEqual(t["summary"]["P0"], 0)
        self.assertIn("DT-STR-06", self.ids(t["issues"]))

    def test_blank_cells_noted_when_total_matches(self):
        rows = [["科室", "上月出院人次"], ["甲", 10], ["乙", None], ["丙", 5], ["合计", 15]]
        t = self.inspect(self.csv("blank.csv", rows))
        self.assertIn("DT-NUL-01", self.ids(t["issues"]))
        self.assertIn("DT-SUM-05", self.ids(t["issues"]))
        self.assertEqual(t["summary"]["P0"], 0)


class InspectCellsTest(DataTestCase):

    def test_fail_placeholders_approx_units(self):
        rows = [["社区", "走访人次", "金额（万元）"],
                ["东社区", "—", 1.2],
                ["西社区", "约20", "3500元"],
                ["南社区", "未报", 2.0]]
        t = self.inspect(self.csv("dirty.csv", rows))
        ids = self.ids(t["issues"])
        for expected in ("DT-NUL-02", "DT-TXT-03", "DT-UNIT-01", "DT-TXT-05"):
            self.assertIn(expected, ids)
        self.assertEqual(self.issue(t["issues"], "DT-UNIT-01")["level"], "P0")

    def test_pass_csv_numbers_are_not_text_numbers(self):
        """必过：CSV 中的普通数字不报"文本型数字"。"""
        t = self.inspect(self.csv("plain.csv", [["月份", "办件量（件）"], ["1月", "8320"], ["2月", "6110"]]))
        self.assertNotIn("DT-TXT-01", self.ids(t["issues"]))

    def test_gbk_csv(self):
        path = self.csv("gbk.csv", [["科室", "人数"], ["甲", 1], ["合计", 1]], encoding="gb18030")
        t = self.inspect(path)
        self.assertEqual(t["data_rows"], 1)

    def test_duplicates_and_variants(self):
        rows = [["科室", "人数"], ["心内科", 1], ["心内科 ", 2], ["骨科", 3], ["骨科", 3]]
        t = self.inspect(self.csv("dup.csv", rows))
        ids = self.ids(t["issues"])
        self.assertIn("DT-DUP-01", ids)
        self.assertIn("DT-DUP-02", ids)
        self.assertIn("DT-DUP-03", ids)

    def test_key_option_and_grouping_column(self):
        """分组列（一半以上的值重复）不作默认关键列；--key 可以改；列名不存在时退出码 2。"""
        rows = [["类别", "科室", "人数"], ["内科", "心内科", 10], ["内科", "呼吸科", 8], ["内科", "消化科", 6],
                ["外科", "骨科", 9], ["外科", "普外科", 7], ["外科", "骨科", 9]]
        path = self.csv("g.csv", rows)
        t = self.inspect(path)
        self.assertEqual({c["name"]: c["kind"] for c in t["columns"]}["科室"], "key")
        dup = self.issue(t["issues"], "DT-DUP-01")  # 必错：骨科真的报了两次
        self.assertIn("骨科", dup["location"])
        t = self.inspect(path, "--key", "类别")
        self.assertEqual({c["name"]: c["kind"] for c in t["columns"]}["类别"], "key")
        self.assertEqual(self.run_tool(data_tool, "inspect", path, "--key", "不存在")[0], 2)

    def test_year_columns_header(self):
        rows = [["某区办件量"], ["月份", 2026, 2025], ["1月", 8320, 7905], ["2月", 6110, 6352]]
        t = self.inspect(self.csv("years.csv", rows))
        self.assertEqual(t["header_rows"], [2])
        self.assertEqual([c["name"] for c in t["columns"]], ["月份", "2026", "2025"])

    def test_privacy_columns_masked(self):
        rows = [["患者姓名", "联系电话", "住院天数"], ["张三丰", "13812345678", 5], ["李四", "13987654321", 7]]
        data = self.run_json(data_tool, "inspect", str(self.csv("p.csv", rows)))
        text = json.dumps(data, ensure_ascii=False)
        self.assertIn("DT-PRI-01", text)
        self.assertNotIn("13812345678", text)
        self.assertNotIn("张三丰", text)

    def test_phone_in_text_column(self):
        rows = [["社区", "备注", "人数"], ["东社区", "联系人王某13812345678", 3]]
        data = self.run_json(data_tool, "inspect", str(self.csv("p2.csv", rows)))
        text = json.dumps(data, ensure_ascii=False)
        self.assertIn("DT-PRI-02", text)
        self.assertNotIn("13812345678", text)


class InspectNotesTest(DataTestCase):

    ROWS = [["某医院2026年9月科室指标"], ["填报单位：质控科（中医科四级手术按不开展填报）"],
            ["科室", "出院人次", "上月出院人次", "四级手术台次"],
            ["心内科", 128, 128, 9], ["神经内科", 261, None, 23], ["中医科", 111, 109, "—"],
            ["合计", 500, 237, 32],
            ["注：数据来源为病案首页统计；神经内科上月数据未报；患者王小明（13812345678）已转院"]]

    def test_notes_and_pre_texts_output(self):
        """表注和表头前说明都输出（已脱敏），空值、占位符的建议引用表注原文。"""
        t = self.inspect(self.csv("n.csv", self.ROWS))
        self.assertEqual(t["pre_texts"], ["填报单位：质控科（中医科四级手术按不开展填报）"])
        self.assertEqual(len(t["notes"]), 1)
        self.assertIn("神经内科上月数据未报", t["notes"][0])
        self.assertNotIn("13812345678", t["notes"][0])
        blank = self.issue(t["issues"], "DT-NUL-01")
        self.assertIn("“神经内科上月数据未报”", blank["suggestion"])
        null = self.issue(t["issues"], "DT-NUL-02")
        self.assertIn("中医科四级手术按不开展填报", null["suggestion"])
        code, out, _ = self.run_tool(data_tool, "inspect", self.csv("n.csv", self.ROWS))
        self.assertIn("表注：数据来源为病案首页统计", out)
        self.assertIn("表头前说明：填报单位", out)

    def test_pass_no_note_keeps_default_suggestion(self):
        rows = [["科室", "出院人次"], ["心内科", 128], ["中医科", "—"], ["合计", 128]]
        t = self.inspect(self.csv("n2.csv", rows))
        self.assertNotIn("表注已说明", self.issue(t["issues"], "DT-NUL-02")["suggestion"])

    def test_unchecked_columns_listed(self):
        rows = [["社区", "走访人次", "覆盖率（%）"], ["东社区", "约20", 80.0], ["西社区", 30, 70.0], ["合计", 50, 75.0]]
        code, out, _ = self.run_tool(data_tool, "inspect", self.csv("u.csv", rows))
        self.assertIn("未校验：走访人次（含约数，不能精确校验）", out)
        self.assertIn("覆盖率（%）（比率或平均数不能用求和校验）", out)

    def test_names_masked_in_notes_and_summary_labels_kept(self):
        rows = [["患者姓名", "费用（元）"], ["张三丰", 100], ["李四", 200], ["合计", 400],
                ["注：张三丰费用含外院检查"]]
        data = self.run_json(data_tool, "inspect", str(self.csv("nm.csv", rows)))
        text = json.dumps(data, ensure_ascii=False)
        self.assertNotIn("张三丰", text)
        self.assertIn("张某费用含外院检查", text)
        loc = self.issue(data["tables"][0]["issues"], "DT-SUM-01")["location"]
        self.assertIn("（合计）", loc)  # 合计行标签不脱敏成“合某”
        self.assertNotIn("合某", text)


class InspectMarkingTest(DataTestCase):
    """DT-SEC-01：密级和内部标注。"""

    BODY = [["项目", "投资额（万元）"], ["甲项目", 1200], ["乙项目", 800], ["合计", 2000]]

    def test_fail_marking_before_header(self):
        """必错：表头前写“（内部资料 不得外传）”，停止处理，只报标注位置，不输出标题、数据和其他问题。"""
        rows = [["某局2026年1—9月项目统计表"], ["（内部资料 不得外传）"]] + self.BODY
        path = self.csv("m.csv", rows)
        data = self.run_json(data_tool, "inspect", str(path))
        self.assertTrue(data["stopped"])
        t = data["tables"][0]
        self.assertTrue(t["stopped"])
        self.assertEqual(t["issues"][0]["id"], "DT-SEC-01")
        self.assertEqual(t["issues"][0]["level"], "P0")
        self.assertIn("guardrails 第一条", t["issues"][0]["message"])
        self.assertEqual(t["marks"][0]["location"], "A2")
        text = json.dumps(data, ensure_ascii=False)
        for leaked in ("1200", "项目统计表", "stats", "checks"):
            self.assertNotIn(leaked, text)
        code, out, _ = self.run_tool(data_tool, "inspect", path)
        self.assertTrue(out.splitlines()[0].startswith("【停止处理】DT-SEC-01"))
        self.assertNotIn("1200", out)
        self.assertEqual(self.run_tool(data_tool, "inspect", path, "--strict")[0], 1)

    def test_fail_marking_variants(self):
        cases = {
            "note": self.BODY + [["注：内部资料，不得外传"]],
            "secret": [["秘密★1年"]] + self.BODY,
            "level": [["密级：机密"]] + self.BODY,
            "work": [["本表涉及工作秘密"]] + self.BODY,
            "internal": [["仅供内部参考"]] + self.BODY,
            "cell": [["项目", "投资额（万元）", "备注"], ["甲项目", 1200, "内部使用"], ["乙项目", 800, None]],
        }
        for name, rows in cases.items():
            with self.subTest(case=name):
                data = self.run_json(data_tool, "inspect", str(self.csv(f"{name}.csv", rows)))
                self.assertTrue(data["stopped"], name)
        # 文件名带标注
        data = self.run_json(data_tool, "inspect", str(self.csv("内部资料-项目表.csv", self.BODY)))
        self.assertTrue(data["stopped"])

    def test_pass_item_names_not_marking(self):
        """必过：“保密工作经费”“内部资料印刷费”“内部使用车辆租赁费”“是否内部使用”是项目名称，照常处理。"""
        rows = [["项目", "金额（万元）", "是否内部使用"], ["保密工作经费", 12, "是"], ["内部资料印刷费", 3, "否"],
                ["内部使用车辆租赁费", 5, "是"], ["合计", 20, None]]
        data = self.run_json(data_tool, "inspect", str(self.csv("ok.csv", rows)))
        self.assertFalse(data["stopped"])
        self.assertNotIn("DT-SEC-01", self.ids(data["tables"][0]["issues"]))

    def test_other_commands_intercepted_until_confirmed(self):
        marked = self.csv("甲2026年9月.csv", [["科室", "出院人次"], ["心内科", 120], ["注：内部资料，不得外传"]])
        plain = self.csv("乙2026年9月.csv", [["科室", "出院人次"], ["儿科", 60]])
        calls = {
            "merge": ["merge", marked, plain],
            "count": ["count", marked, "--by", "科室"],
            "phrase": ["phrase", "--type", "share", "--table", marked, "--part-col", "出院人次"],
        }
        for name, args in calls.items():
            with self.subTest(cmd=name):
                code, out, err = self.run_tool(data_tool, *args)
                self.assertEqual(code, 2)
                self.assertIn("DT-SEC-01", err)
                self.assertNotIn("120", out + err)
        # 用户说明后加 --marking-ok：照常处理，DT-SEC-01 降为 P2 并提示按内部材料管理
        data = self.run_json(data_tool, "merge", str(marked), str(plain), "--marking-ok")
        sec = self.issue(data["issues"], "DT-SEC-01")
        self.assertEqual(sec["level"], "P2")
        self.assertEqual(data["aggregate"][0]["values"]["出院人次"]["value"], "180")
        t = self.inspect(marked, "--marking-ok")
        self.assertEqual(t["issues"][0]["id"], "DT-SEC-01")

    @unittest.skipUnless(HAS_OPENPYXL, "需要 openpyxl")
    def test_fail_marking_on_other_sheet(self):
        """必错：标注写在没有读取的“封面”工作表上，读数据工作表时同样停止。"""
        path = self.xlsx("w.xlsx", {"数据": self.BODY, "封面": [["某局项目统计"], ["内部资料"]]})
        data = self.run_json(data_tool, "inspect", str(path), "--sheet", "数据")
        self.assertTrue(data["stopped"])
        self.assertIn("封面", data["tables"][0]["marks"][0]["location"])


@unittest.skipUnless(HAS_OPENPYXL, "需要 openpyxl")
class InspectXlsxTest(DataTestCase):

    def test_merged_text_numbers_formula_hidden(self):
        rows = [["某局2026年1—9月项目统计表"],
                ["项目", "投资额（万元）", "完成率"],
                ["甲项目", "1200", 0.85],
                ["乙项目", 800, 0.9],
                ["丙项目", None, 0.7],
                ["合计", "=SUM(B3:B5)", None]]
        path = self.xlsx("m.xlsx", {"Sheet1": rows}, merges={"Sheet1": ["B4:B5"]},
                         formats={"Sheet1": {"C3": "0%", "C4": "0%", "C5": "0%"}}, hidden_rows={"Sheet1": [5]})
        t = self.inspect(path)
        ids = self.ids(t["issues"])
        for expected in ("DT-MRG-01", "DT-TXT-01", "DT-ERR-02", "DT-STR-04"):
            self.assertIn(expected, ids)
        cols = {c["name"]: c for c in t["columns"]}
        self.assertEqual(cols["完成率"]["kind"], "rate")

    def test_two_level_header(self):
        rows = [["科室", "出院人次", None, "手术台次", None],
                [None, "本月", "上年同月", "本月", "上年同月"],
                ["甲", 10, 9, 3, 2],
                ["乙", 20, 18, 4, 4]]
        path = self.xlsx("h.xlsx", {"S": rows}, merges={"S": ["B1:C1", "D1:E1", "A1:A2"]})
        t = self.inspect(path)
        names = [c["name"] for c in t["columns"]]
        self.assertEqual(names, ["科室", "出院人次-本月", "出院人次-上年同月", "手术台次-本月", "手术台次-上年同月"])
        self.assertEqual(t["data_rows"], 2)

    def test_pass_merged_group_column_not_key(self):
        """必过：片区列是合并单元格，关键列取街道，填充出来的片区不报重复。"""
        rows = [["某区2026年9月各街道办件量"], ["片区", "街道", "办件量（件）"],
                ["东片", "东一街道", 120], [None, "东二街道", 80], [None, "东三街道", 60],
                ["西片", "西一街道", 90], [None, "西二街道", 70]]
        path = self.xlsx("k.xlsx", {"S": rows}, merges={"S": ["A3:A5", "A6:A7"]})
        t = self.inspect(path)
        cols = {c["name"]: c["kind"] for c in t["columns"]}
        self.assertEqual(cols["街道"], "key")
        self.assertNotIn("DT-DUP-01", self.ids(t["issues"]))
        self.assertIn("DT-MRG-01", self.ids(t["issues"]))  # 合并单元格照常提示
        # --key 指定片区时，填充出来的值也不算重复
        t = self.inspect(path, "--key", "片区")
        self.assertEqual({c["name"]: c["kind"] for c in t["columns"]}["片区"], "key")
        self.assertNotIn("DT-DUP-01", self.ids(t["issues"]))

    def test_all_sheets(self):
        path = self.xlsx("s.xlsx", {"一月": [["科室", "人数"], ["甲", 1]], "二月": [["科室", "人数"], ["甲", 2]]})
        data = self.run_json(data_tool, "inspect", str(path), "--all-sheets")
        self.assertEqual(len(data["tables"]), 2)
        self.assertIn("［二月］", data["tables"][1]["source"])


class InspectCliTest(DataTestCase):

    def test_strict_exit_codes(self):
        bad = self.csv("bad.csv", dept_rows(total_out=999))
        ok = self.csv("ok.csv", dept_rows())
        self.assertEqual(self.run_tool(data_tool, "inspect", bad, "--strict")[0], 1)
        self.assertEqual(self.run_tool(data_tool, "inspect", ok, "--strict")[0], 0)

    def test_input_errors(self):
        self.assertEqual(self.run_tool(data_tool, "inspect", self.tmp / "none.csv")[0], 2)
        xls = self.tmp / "old.xls"
        xls.write_bytes(b"\xd0\xcf\x11\xe0" + b"\x00" * 100)
        code, _out, err = self.run_tool(data_tool, "inspect", xls)
        self.assertEqual(code, 2)
        self.assertIn(".xlsx", err)
        doc = self.tmp / "a.docx"
        doc.write_bytes(b"abc")
        self.assertEqual(self.run_tool(data_tool, "inspect", doc)[0], 2)
        self.assertEqual(self.run_tool(data_tool, "bogus")[0], 2)

    def test_text_output(self):
        code, out, _ = self.run_tool(data_tool, "inspect", self.csv("bad.csv", dept_rows(total_out=295)))
        self.assertEqual(code, 0)
        self.assertIn("[P0]", out)
        self.assertIn("合计校验", out)


# ---------------------------------------------------------------------------
# 合并
# ---------------------------------------------------------------------------

class MergeTest(DataTestCase):

    def record_tables(self):
        a = self.csv("东区2026年1—9月.csv", [["东区2026年1—9月台账"], ["社区", "走访户数", "补贴金额（万元）", "覆盖率（%）"],
                                           ["东一社区", 120, 12.5, 80.1], ["东二社区", 80, 8.0, 75.0]])
        b = self.csv("西区2026年1—9月.csv", [["西区2026年1—9月台账"], ["社区", "走访户数", "补贴金额（元）", "覆盖率（%）"],
                                           ["西一社区", 60, 56000, 66.0], ["东二社区", 10, 1000, 50.0]])
        return a, b

    def test_record_merge_with_conversion(self):
        a, b = self.record_tables()
        data = self.run_json(data_tool, "merge", str(a), str(b), "--expect", "东一社区,东二社区,西一社区,北一社区")
        self.assertEqual(data["layout"], "record")
        total = data["aggregate"][0]["values"]
        # 东二社区在两张表中都出现：先不计入合计，另列含重复的合计，确认后再加
        self.assertEqual(total["走访户数"]["value"], "180")  # 120＋60
        self.assertEqual(total["走访户数"]["value_with_dups"], "270")
        self.assertFalse(total["走访户数"]["complete"])
        self.assertIn("重复报送待确认", total["走访户数"]["note"])
        self.assertEqual(total["补贴金额"]["value"], "18.1")  # 12.5＋5.6
        self.assertEqual(total["补贴金额"]["value_with_dups"], "26.2")  # 另加 8.0＋0.1
        # 两张表中东二社区报的数不同（80 与 10），不替用户挑一个“只计一次”
        self.assertNotIn("value_dedup", total["走访户数"])
        self.assertIn("数不同", total["走访户数"]["note"])
        self.assertFalse(total["覆盖率"]["complete"])
        self.assertIsNone(total["覆盖率"]["value"])
        self.assertTrue(any("56000元÷10000＝5.6万元" in c["formula"] for c in data["conversions"]))
        ids = self.ids(data["issues"])
        for expected in ("DT-MG-06", "DT-MG-08", "DT-MG-09", "DT-MG-11"):
            self.assertIn(expected, ids)
        dup = self.issue(data["issues"], "DT-MG-09")
        self.assertEqual(dup["level"], "P0")
        self.assertIn("东二社区（表A、表B）", dup["location"])
        self.assertIn("北一社区", self.issue(data["issues"], "DT-MG-11")["location"])

    def test_pass_no_duplicate_full_total(self):
        """必过：没有重复报送时，合计包含全部明细，不带"含重复"字段。"""
        a = self.csv("甲2026年9月.csv", [["科室", "出院人次"], ["心内科", 120], ["骨科", 95]])
        b = self.csv("乙2026年9月.csv", [["科室", "出院人次"], ["儿科", 60]])
        data = self.run_json(data_tool, "merge", str(a), str(b))
        entry = data["aggregate"][0]["values"]["出院人次"]
        self.assertEqual(entry["value"], "275")
        self.assertNotIn("value_with_dups", entry)
        self.assertTrue(entry["complete"])
        self.assertNotIn("DT-MG-09", self.ids(data["issues"]))

    def test_dedup_total_when_same_values(self):
        """同一科室在两张表中报了相同的数：另列只计一次的合计，用户确认后直接采用。"""
        a = self.csv("内科片区2026年9月.csv", [["科室", "出院人次"], ["心内科", 120], ["重症医学科", 30]])
        b = self.csv("外科片区2026年9月.csv", [["科室", "出院人次"], ["骨科", 95], ["重症医学科", 30]])
        data = self.run_json(data_tool, "merge", str(a), str(b))
        entry = data["aggregate"][0]["values"]["出院人次"]
        self.assertEqual(entry["value"], "215")  # 重复的两行都先不计
        self.assertEqual(entry["value_with_dups"], "275")
        self.assertEqual(entry["value_dedup"], "245")  # 重症医学科只计一次
        self.assertIn("只计一次的合计为245", entry["note"])
        self.assertIn("value_dedup", self.issue(data["issues"], "DT-MG-09")["message"])

    def test_fail_unit_conflict_not_summed(self):
        a = self.csv("a2026年9月.csv", [["单位", "服务人数（人）"], ["甲", 10]])
        b = self.csv("b2026年9月.csv", [["单位", "服务人数（人次）"], ["乙", 30]])
        data = self.run_json(data_tool, "merge", str(a), str(b))
        self.assertIn("DT-MG-05", self.ids(data["issues"]))
        entry = data["aggregate"][0]["values"]["服务人数"]
        self.assertIsNone(entry["value"])
        self.assertIn("单位冲突", entry["note"])

    def test_fail_period_conflict_excluded(self):
        a = self.csv("甲2026年1—9月.csv", [["单位", "收入（万元）"], ["甲", 10]])
        b = self.csv("乙2026年1-9月.csv", [["单位", "收入（万元）"], ["乙", 20]])
        c = self.csv("丙2026年1—8月.csv", [["单位", "收入（万元）"], ["丙", 30]])
        data = self.run_json(data_tool, "merge", str(a), str(b), str(c))
        self.assertEqual(data["excluded_tables"], ["表C"])
        self.assertEqual(data["aggregate"][0]["values"]["收入"]["value"], "30")
        self.assertEqual(self.issue(data["issues"], "DT-MG-03")["level"], "P0")
        code = self.run_tool(data_tool, "merge", a, b, c, "--strict")[0]
        self.assertEqual(code, 1)

    def test_indicator_layout(self):
        def t(name, unit, value):
            return self.csv(name, [[f"{name[:3]}经营快报"], ["指标名称", "计量单位", "本年累计"],
                                   ["营业收入", unit, value], ["资产负债率", "%", 50.0], ["从业人员", "人", 10]])
        a = t("甲公司2026年1—9月.csv", "万元", 100.5)
        b = t("乙公司2026年1—9月.csv", "元", 2000000)
        data = self.run_json(data_tool, "merge", str(a), str(b))
        self.assertEqual(data["layout"], "indicator")
        agg = {item["key"]: item["values"]["本年累计"] for item in data["aggregate"]}
        self.assertEqual(agg["营业收入"]["value"], "300.5")
        self.assertEqual(agg["从业人员"]["value"], "20")
        self.assertIsNone(agg["资产负债率"]["value"])

    @unittest.skipUnless(HAS_OPENPYXL, "需要 openpyxl")
    def test_out_never_overwrites(self):
        a, b = self.record_tables()
        out = self.tmp / "汇总.xlsx"
        out.write_bytes(b"old")
        data = self.run_json(data_tool, "merge", str(a), str(b), "--out", str(out))
        self.assertTrue(data["out"].endswith("汇总-v2.xlsx"))
        self.assertEqual(out.read_bytes(), b"old")
        import openpyxl  # type: ignore
        wb = openpyxl.load_workbook(data["out"])
        self.assertEqual(wb.sheetnames[:3], ["汇总", "明细", "问题清单"])
        # 汇总行旁写明含重复待确认
        texts = [c.value for row in wb["汇总"].iter_rows() for c in row if isinstance(c.value, str)]
        self.assertIn("说明", texts)
        self.assertTrue(any("重复报送待确认" in t for t in texts))
        csv_out = self.tmp / "汇总.csv"
        data = self.run_json(data_tool, "merge", str(a), str(b), "--out", str(csv_out))
        self.assertTrue(csv_out.is_file())

    def test_single_table_rejected(self):
        a, _ = self.record_tables()
        self.assertEqual(self.run_tool(data_tool, "merge", a)[0], 2)


# ---------------------------------------------------------------------------
# 计算与转文字
# ---------------------------------------------------------------------------

class CalcTest(DataTestCase):

    def calc(self, *args: str, cmd: str = "calc") -> dict:
        return self.run_json(data_tool, cmd, *args)

    def test_growth_formula_and_rounding(self):
        d = self.calc("--type", "growth", "--current", "6171", "--base", "5879", "--basis", "同比", "--unit", "人次")
        self.assertEqual(d["display"], "5.0%")
        self.assertEqual(d["formula"], "（6171－5879）÷5879×100%＝4.967%≈5.0%")
        self.assertEqual(d["change"]["display"], "292人次")
        # 四舍五入：0.05 进位
        d = self.calc("--type", "growth", "--current", "100.05", "--base", "100")
        self.assertEqual(d["rounded"], "0.1")
        self.assertTrue(any("未指明对比基期" in w for w in d["warnings"]))

    def test_growth_edge_cases(self):
        d = self.calc("--type", "growth", "--current", "12", "--base", "0")
        self.assertIsNone(d["value"])
        self.assertTrue(any("基期为0" in w for w in d["warnings"]))
        d = self.calc("--type", "growth", "--current", "300", "--base", "-500", "--basis", "同比")
        self.assertIsNone(d["value"])
        self.assertEqual(d["sign_case"], "扭亏为盈")
        d = self.calc("--type", "growth", "--current", "12", "--base", "8", "--unit", "人", "--basis", "环比")
        self.assertTrue(any("基数较小" in w for w in d["warnings"]))

    def test_pp_share_pull_contrib(self):
        d = self.calc("--type", "pp", "--current", "38.2%", "--base", "35")
        self.assertEqual(d["display"], "3.2个百分点")
        self.assertIn("relative_change", d)
        d = self.calc("--type", "share", "--part", "869", "--total", "3264")
        self.assertEqual(d["display"], "26.6%")
        d = self.calc("--type", "pull", "--part-current", "120", "--part-base", "100", "--total-base", "1000")
        self.assertEqual(d["display"], "2.0个百分点")
        d = self.calc("--type", "contrib", "--part-current", "120", "--part-base", "100",
                      "--total-current", "1050", "--total-base", "1000")
        self.assertEqual(d["display"], "40.0%")

    def test_completion_with_time_progress(self):
        d = self.calc("--type", "completion", "--actual", "8.6", "--target", "12", "--elapsed", "9", "--span", "12")
        self.assertEqual(d["display"], "71.7%")
        self.assertEqual(d["time_progress"]["display"], "75.0%")
        self.assertEqual(d["time_progress"]["gap_display"], "3.3个百分点")

    def test_multiple_cagr_mean_ratio(self):
        d = self.calc("--type", "multiple", "--current", "6800", "--base", "3200")
        self.assertEqual(d["display"], "2.1倍")
        self.assertEqual(d["growth"]["display"], "112.5%")
        d = self.calc("--type", "cagr", "--current", "6800", "--base", "3200", "--years", "5")
        self.assertEqual(d["display"], "16.3%")
        d = self.calc("--type", "mean", "--values", "8.2,9.6,5.1", "--unit", "天")
        self.assertEqual(d["display"], "7.6天")
        d = self.calc("--type", "wmean", "--values", "8.2,9.6", "--weights", "120,95", "--unit", "天")
        self.assertEqual(d["display"], "8.8天")
        d = self.calc("--type", "ratio", "--numerator", "61516", "--denominator", "6171", "--unit", "天")
        self.assertEqual(d["display"], "10.0天")
        self.assertIn("61516÷6171", d["formula"])

    def test_bad_inputs_exit_2(self):
        for args in (["--type", "growth", "--current", "约300", "--base", "200"],
                     ["--type", "growth", "--current", "300"],
                     ["--type", "share", "--part", "1", "--total", "0"],
                     ["--type", "wmean", "--values", "1,2", "--weights", "1"]):
            with self.subTest(args=args):
                self.assertEqual(self.run_tool(data_tool, "calc", *args)[0], 2)


class PhraseTest(DataTestCase):

    def sentences(self, *args: str):
        return self.run_json(data_tool, "phrase", *args)["sentences"]

    def test_growth_wording(self):
        s = self.sentences("--type", "growth", "--name", "出院人次", "--unit", "人次", "--current", "6171",
                           "--base", "5879", "--basis", "同比", "--period", "2026年9月")
        self.assertEqual(s[0], "2026年9月，出院6171人次，同比增长5.0%。")
        self.assertIn("比上年同期增加292人次，增长5.0%", s[1])
        s = self.sentences("--type", "growth", "--name", "办件量", "--unit", "件", "--current", "9405",
                           "--base", "9810", "--basis", "同比")
        self.assertIn("同比下降4.1%", s[0])
        for sentence in s:
            self.assertNotIn("负增长", sentence)
            self.assertNotIn("-", sentence)

    def test_subject_keeps_single_char_units(self):
        """必过：单字单位不从指标名末尾截掉，不写成"不良事12件""受理案120件"。"""
        cases = [
            (("不良事件", "件", "12", "9"), "不良事件12件"),
            (("受理案件", "件", "120", "100"), "受理案件120件"),
            (("受理案件数", "件", "120", "100"), "受理案件120件"),
            (("病例数", "例", "60", "40"), "病例60例"),
            (("出院人数", "人", "6171", "5879"), "出院6171人"),
            (("走访户数", "户", "120", "100"), "走访120户"),
            (("出院人次", "人次", "6171", "5879"), "出院6171人次"),
        ]
        for (name, unit, cur, base), head in cases:
            with self.subTest(name=name):
                s = self.sentences("--type", "growth", "--name", name, "--unit", unit, "--current", cur,
                                   "--base", base, "--basis", "同比")
                self.assertTrue(all(x.startswith(head) for x in s), s)

    def test_default_wording_warned(self):
        """必错：句子里用了"比基期""该指标""年度目标""总量"等默认叫法时，都要提示。"""
        d = self.run_json(data_tool, "phrase", "--type", "pp", "--current", "71", "--base", "62")
        self.assertIn("该指标", d["sentences"][0])
        self.assertTrue(any("未指明对比基期" in w for w in d["warnings"]))
        self.assertTrue(any("未写指标名称" in w for w in d["warnings"]))
        d = self.run_json(data_tool, "phrase", "--type", "multiple", "--name", "科研经费", "--current", "6800",
                          "--base", "3200")
        self.assertTrue(any("未指明对比基期" in w for w in d["warnings"]))
        d = self.run_json(data_tool, "phrase", "--type", "completion", "--name", "营业收入", "--actual", "8.6",
                          "--target", "12")
        self.assertTrue(any("年度目标" in w for w in d["warnings"]))
        d = self.run_json(data_tool, "phrase", "--type", "share", "--part", "869", "--total", "3264")
        self.assertTrue(any("--total-name" in w for w in d["warnings"]))
        # 必过：都写明了就不提示
        d = self.run_json(data_tool, "phrase", "--type", "pp", "--name", "网办率", "--current", "71", "--base", "62",
                          "--base-label", "上年")
        self.assertEqual(d["warnings"], [])

    def test_pp_keeps_input_places(self):
        """用户给"71%""62%"时写"71%""9个百分点"，不改成"71.0%""9.0个百分点"。"""
        s = self.sentences("--type", "pp", "--name", "网办率", "--current", "71%", "--base", "62%",
                           "--base-label", "上年")
        self.assertEqual(s[0], "网办率为71%，比上年提高9个百分点。")
        self.assertIn("由上年的62%升至71%", s[1])
        # 两个输入位数不同时补齐
        s = self.sentences("--type", "pp", "--name", "网办率", "--current", "71.5", "--base", "62",
                           "--base-label", "上年")
        self.assertEqual(s[0], "网办率为71.5%，比上年提高9.5个百分点。")
        self.assertIn("由上年的62.0%升至71.5%", s[1])

    def test_pp_wording(self):
        s = self.sentences("--type", "pp", "--name", "四级手术占比", "--current", "26.6", "--base", "24.1",
                           "--basis", "同比")
        self.assertEqual(s[0], "四级手术占比为26.6%，同比提高2.5个百分点。")
        s = self.sentences("--type", "pp", "--name", "营业收入增速", "--current", "6.2", "--base", "8.3",
                           "--basis", "同比")
        self.assertIn("同比回落2.1个百分点", s[0])

    def test_loss_and_doubling_wording(self):
        s = self.sentences("--type", "growth", "--name", "利润总额", "--unit", "万元", "--current", "300",
                           "--base", "-500", "--basis", "同比")
        self.assertIn("上年同期亏损500万元，实现扭亏为盈", s[0])
        s = self.sentences("--type", "growth", "--name", "利润总额", "--unit", "万元", "--current", "-300",
                           "--base", "-500", "--basis", "同比")
        self.assertIn("减亏200万元", s[0])
        s = self.sentences("--type", "multiple", "--name", "科研经费", "--unit", "万元", "--current", "6400",
                           "--base", "3200", "--base-label", "2021年")
        self.assertTrue(any("翻了一番" in x for x in s))
        s = self.sentences("--type", "multiple", "--name", "科研经费", "--unit", "万元", "--current", "6800",
                           "--base", "3200", "--base-label", "2021年")
        self.assertFalse(any("翻" in x for x in s))
        self.assertFalse(any("增长了2.1倍" in x for x in s))

    def test_completion_wording(self):
        s = self.sentences("--type", "completion", "--name", "营业收入", "--unit", "亿元", "--actual", "8.6",
                           "--target", "12", "--elapsed", "9", "--period", "1—9月")
        self.assertEqual(s[0], "1—9月，营业收入完成8.6亿元，完成年度目标12亿元的71.7%，慢于时间进度3.3个百分点。")

    def test_pp_rounds_to_displayed_rates(self):
        d = self.run_json(data_tool, "phrase", "--type", "pp", "--name", "四级手术占比", "--current", "26.624",
                          "--base", "24.080", "--basis", "同比")
        self.assertEqual(d["formula"], "26.6%－24.1%＝2.5个百分点")
        self.assertIn("由上年同期的24.1%升至26.6%", d["sentences"][1])

    def test_table_mode_pull_sums_to_total_growth(self):
        rows = [["某县各镇产值"], ["镇街", "本年累计（万元）", "上年同期（万元）"],
                ["甲镇", 128650, 116420], ["乙镇", 45210, 47380], ["丙镇", 63480, 55920], ["全县合计", 237340, 219720]]
        data = self.run_json(data_tool, "phrase", "--type", "pull", "--table", str(self.csv("pull.csv", rows)),
                             "--current-col", "本年累计", "--base-col", "上年同期", "--total-name", "全县产值")
        pulls = {r["name"]: Decimal(r["rounded"]) for r in data["results"]}
        self.assertEqual(set(pulls), {"甲镇", "乙镇", "丙镇"})
        self.assertIn("拉低全县产值增速", next(r for r in data["results"] if r["name"] == "乙镇")["sentences"][0])
        total_growth = (Decimal(237340) - Decimal(219720)) / Decimal(219720) * 100
        self.assertLessEqual(abs(sum(pulls.values()) - total_growth), Decimal("0.15"))
        self.assertIn("合计行", data["totals"]["total_base"])

    SCOPE_ROWS = [["某医院2026年9月科室出院人次"], ["科室", "出院人次", "上月出院人次"],
                  ["心内科", 128, 128], ["呼吸科", 329, 316], ["神经内科", 261, None], ["骨科", 80, 84],
                  ["全院合计", 798, 528], ["注：神经内科上月数据未报"]]

    def test_fail_total_row_scope_gap(self):
        """必错：合计与明细都相符，但基期列有空值，本期含神经内科、基期不含，合计行环比不能直接引用。"""
        path = self.csv("scope.csv", self.SCOPE_ROWS)
        data = self.run_json(data_tool, "phrase", "--type", "growth", "--table", str(path),
                             "--current-col", "出院人次", "--base-col", "上月出院人次", "--basis", "环比")
        total = next(r for r in data["results"] if r["row_kind"] == "total")
        self.assertEqual(total["display"], "51.1%")  # （798－528）÷528，口径不同，不能用
        self.assertTrue(any("口径不一致" in w and "神经内科" in w for w in total["warnings"]))
        self.assertEqual(total["scope_gaps"][0]["label"], "神经内科")
        for s in total["sentences"]:
            self.assertIn("【待核实：两期口径不同", s)
        same = next(r for r in data["results"] if r["row_kind"] == "same_scope")
        self.assertEqual(same["name"], "全院（不含神经内科）")
        self.assertEqual(same["inputs"], {"current": "537", "base": "528"})  # 剔除后两边同口径
        self.assertEqual(same["display"], "1.7%")
        self.assertIn("剔除神经内科", same["formula"])
        self.assertNotIn("待核实", same["sentences"][0])
        # 数据行本身照常计算，神经内科跳过并标【待补】
        neuro = next(r for r in data["results"] if r["name"] == "神经内科")
        self.assertIn("skipped", neuro)

    def test_pass_total_row_same_scope(self):
        """必过：两列明细都齐全时，合计行不报口径问题，不生成同口径结果。"""
        path = self.csv("ok.csv", dept_rows())
        data = self.run_json(data_tool, "phrase", "--type", "growth", "--table", str(path),
                             "--current-col", "出院人次", "--base-col", "上年同月出院人次", "--basis", "同比")
        total = next(r for r in data["results"] if r["row_kind"] == "total")
        self.assertFalse(any("口径" in w for w in total["warnings"]))
        self.assertFalse(any(r["row_kind"] == "same_scope" for r in data["results"]))
        self.assertNotIn("待核实", total["sentences"][0])

    def test_fail_pull_scope_gap_uses_same_scope_totals(self):
        path = self.csv("scope.csv", self.SCOPE_ROWS)
        data = self.run_json(data_tool, "phrase", "--type", "pull", "--table", str(path),
                             "--current-col", "出院人次", "--base-col", "上月出院人次", "--total-name", "全院出院人次")
        self.assertTrue(any("两期口径不同" in w for w in data["warnings"]))
        self.assertIn("同口径", data["totals"]["total_current"])
        self.assertTrue(data["totals"]["total_current"].startswith("537"))
        pulls = [Decimal(r["rounded"]) for r in data["results"] if "rounded" in r]
        # 各部分拉动之和等于同口径总量增速 1.7%
        self.assertLessEqual(abs(sum(pulls) - Decimal("1.7")), Decimal("0.15"))

    def test_table_mode_and_ledger(self):
        path = self.csv("t.csv", dept_rows(total_out=295) + [["东院区", None, 10, 1.0]])
        data = self.run_json(data_tool, "phrase", "--type", "growth", "--table", str(path),
                             "--current-col", "出院人次", "--base-col", "上年同月出院人次", "--basis", "同比",
                             "--unit", "人次")
        by_name = {r["name"]: r for r in data["results"]}
        self.assertEqual(by_name["心内科"]["display"], "9.1%")
        total = by_name["合计"]
        self.assertTrue(any("合计与明细不符" in w for w in total["warnings"]))
        self.assertIn("东院区", by_name)
        self.assertIn("skipped", by_name["东院区"])
        self.assertIn("【待补", by_name["东院区"]["sentence"])
        code, out, _ = self.run_tool(data_tool, "phrase", "--type", "share", "--table", str(path),
                                     "--part-col", "出院人次", "--md")
        self.assertEqual(code, 0)
        self.assertIn("| 序号 | 指标 | 算式 | 结果 | 文字表述 | 来源 |", out)
        # 来源列用规范标记【待补】，不出现第三种标记
        self.assertIn("【待补：表名和行号】", out)
        self.assertNotIn("【待填", out)
        # 合计行与明细不符时，占比的分母用明细求和
        self.assertIn("÷275×100%", out)

    MISMATCH_ROWS = [["某医院2026年9月科室指标"],
                     ["科室", "出院人次", "上年同月出院人次", "出院者占用总床日数"],
                     ["心内科", 120, 110, 1000], ["骨科", 95, 100, 1100], ["儿科", 60, 58, 400],
                     ["全院合计", 295, 268, 2500]]

    def test_fail_total_mismatch_uses_detail_sum(self):
        """必错：合计行与明细不符（295≠275）时，成句和底稿行不得出现表中合计，改用明细相加数并带【待核实】。"""
        path = self.csv("mis.csv", self.MISMATCH_ROWS)
        data = self.run_json(data_tool, "phrase", "--type", "growth", "--table", str(path), "--current-col", "出院人次",
                             "--base-col", "上年同月出院人次", "--basis", "同比")
        total = next(r for r in data["results"] if r["row_kind"] == "total")
        self.assertEqual(total["inputs"], {"current": "275", "base": "268"})
        self.assertEqual(total["display"], "2.6%")  # （275－268）÷268，不是（295－268）÷268＝10.1%
        self.assertEqual(total["sentences"][0],
                         "全院275人次，同比增长2.6%【待核实：出院人次表中合计为295，明细相加为275】。")
        for s in total["sentences"]:
            self.assertNotIn("295人次", s)
        self.assertTrue(any("295" in w for w in total["warnings"]))  # 原数只放进提示
        # 比值：全院平均住院日用明细相加的出院人次作分母
        data = self.run_json(data_tool, "phrase", "--type", "ratio", "--table", str(path), "--name", "平均住院日",
                             "--numerator-col", "出院者占用总床日数", "--denominator-col", "出院人次", "--unit", "天")
        total = next(r for r in data["results"] if r["row_kind"] == "total")
        self.assertEqual(total["inputs"]["denominator"], "275")
        self.assertEqual(total["display"], "9.1天")  # 2500÷275＝9.09，不是 2500÷295＝8.5
        self.assertIn("【待核实：出院人次表中合计为295，明细相加为275】", total["sentences"][0])
        # 底稿行：文字表述列不出现表中合计数（【待核实】标记之外）
        code, out, _ = self.run_tool(data_tool, "phrase", "--type", "growth", "--table", path, "--current-col",
                                     "出院人次", "--base-col", "上年同月出院人次", "--basis", "同比", "--md")
        row = next(line for line in out.splitlines() if line.startswith("| 4 |"))
        sentence = row.split("|")[5]
        self.assertNotIn("295", sentence.split("【")[0])
        self.assertIn("【待核实", sentence)

    def test_pass_total_within_rounding_uses_table_total(self):
        """必过：合计与明细只差修约误差（末位 1 个单位）时，照常用表中合计，不加【待核实】。"""
        rows = [["项目", "本年（万元）", "上年（万元）"], ["甲", 1.33, 1.2], ["乙", 1.33, 1.2], ["丙", 1.33, 1.2],
                ["合计", 4.00, 3.6]]
        data = self.run_json(data_tool, "phrase", "--type", "growth", "--table", str(self.csv("r.csv", rows)),
                             "--current-col", "本年", "--base-col", "上年", "--basis", "比上年")
        total = next(r for r in data["results"] if r["row_kind"] == "total")
        self.assertEqual(total["inputs"]["current"], "4")
        self.assertNotIn("待核实", total["sentences"][0])

    def test_share_infers_unit_from_column(self):
        """表格模式的占比、拉动与增速一样，从列名取单位，句子和底稿都带单位。"""
        rows = [["某县2026年1—9月产值"], ["镇街", "本年累计（万元）", "上年同期（万元）"],
                ["甲镇", 128650, 116420], ["乙镇", 45210, 47380], ["全县合计", 173860, 163800]]
        data = self.run_json(data_tool, "phrase", "--type", "share", "--table", str(self.csv("s.csv", rows)),
                             "--part-col", "本年累计", "--total-name", "全县产值")
        first = data["results"][0]
        self.assertEqual(first["sentences"][0], "甲镇128650万元，占全县产值的74.0%。")


class CountTest(DataTestCase):

    ROWS = [["某院三季度不良事件明细"], ["序号", "患者姓名", "事件类型", "事件分级"],
            [1, "王小明", "跌倒", "三级"], [2, "李桂芳", "用药错误", "四级"], [3, "赵建国", "跌倒", "三级"],
            [4, "陈秀英", None, "二级"], [5, "刘子涵", "压力性损伤", "三级"]]

    def test_pass_count_and_share(self):
        data = self.run_json(data_tool, "count", str(self.csv("c.csv", self.ROWS)), "--by", "事件类型")
        self.assertEqual(data["total"], 4)
        first = data["items"][0]
        self.assertEqual((first["values"], first["count"], first["share"]), (["跌倒"], 2, "50.0"))
        self.assertEqual(first["sentence"], "跌倒2件，占50.0%")
        self.assertTrue(any("未计入" in n for n in data["notes"]))  # 第4例事件类型为空
        text = json.dumps(data, ensure_ascii=False)
        self.assertNotIn("王小明", text)

    def test_crosstab_and_share_sum_note(self):
        rows = [["类别"], ["甲"], ["乙"], ["丙"]]
        data = self.run_json(data_tool, "count", str(self.csv("t.csv", rows)), "--by", "类别")
        self.assertEqual(data["share_sum"], "99.9")  # 33.3×3＝99.9，修约后提示
        self.assertTrue(any("不等于100%" in n for n in data["notes"]))
        data = self.run_json(data_tool, "count", str(self.csv("c.csv", self.ROWS)), "--by", "事件类型,事件分级")
        self.assertIn(["跌倒", "三级"], [i["values"] for i in data["items"]])

    def test_fail_sensitive_or_missing_column(self):
        path = self.csv("c.csv", self.ROWS)
        self.assertEqual(self.run_tool(data_tool, "count", path, "--by", "患者姓名")[0], 2)
        self.assertEqual(self.run_tool(data_tool, "count", path, "--by", "不存在")[0], 2)


if __name__ == "__main__":
    unittest.main()
