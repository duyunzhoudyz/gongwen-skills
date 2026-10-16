# -*- coding: utf-8 -*-
"""shared/scripts/fact_diff.py 的单元测试。

运行：python -m unittest discover -s tests/fact_diff -t .
"""
from __future__ import annotations

import io
import json
import shutil
import sys
import tempfile
import unittest
import zipfile
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[2]
SCRIPTS_DIR = REPO_ROOT / "shared" / "scripts"
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

import fact_diff  # noqa: E402

try:
    import docx  # type: ignore  # noqa: F401
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False


class FactDiffTestCase(unittest.TestCase):

    def setUp(self) -> None:
        self.tmp = Path(tempfile.mkdtemp(prefix="fact_diff_test_"))

    def tearDown(self) -> None:
        shutil.rmtree(self.tmp, ignore_errors=True)

    def write(self, name: str, text: str, encoding: str = "utf-8") -> Path:
        path = self.tmp / name
        path.write_bytes(text.encode(encoding))
        return path

    def run_main(self, *args: str):
        out, err = io.StringIO(), io.StringIO()
        with redirect_stdout(out), redirect_stderr(err):
            code = fact_diff.main(list(args))
        return code, out.getvalue(), err.getvalue()

    def diff(self, original: str, revised: str) -> dict:
        a = self.write("原稿.txt", original)
        b = self.write("改稿.txt", revised)
        code, out, _err = self.run_main(str(a), str(b), "--json")
        self.assertEqual(code, 0)
        return json.loads(out)

    @staticmethod
    def values(entries):
        return [(item["type"], item["value"]) for item in entries]


class NumberTest(FactDiffTestCase):

    def test_number_change_reported_in_both_lists(self):
        report = self.diff(
            "全年平均办件时长由3.5天压缩至1.2天，满意度达到98.6%。",
            "全年平均办件时长由3.5天压缩至1天，满意度达到98.6%。",
        )
        self.assertEqual(self.values(report["missing"]), [("数字", "1.2天")])
        self.assertEqual(self.values(report["added"]), [("数字", "1天")])
        self.assertFalse(report["consistent"])

    def test_unit_change_is_a_fact_change(self):
        report = self.diff("同比提高2.3个百分点。", "同比提高2.3%。")
        self.assertEqual(self.values(report["missing"]), [("数字", "2.3个百分点")])
        self.assertEqual(self.values(report["added"]), [("数字", "2.3%")])

    def test_amount_and_date_change(self):
        report = self.diff(
            "投入资金66万元，于2026年10月16日前完成。",
            "投入资金60万元，于2026年10月26日前完成。",
        )
        self.assertIn(("数字", "66万元"), self.values(report["missing"]))
        self.assertIn(("日期", "2026年10月16日"), self.values(report["missing"]))
        self.assertIn(("数字", "60万元"), self.values(report["added"]))
        self.assertIn(("日期", "2026年10月26日"), self.values(report["added"]))

    def test_occurrence_count_is_compared(self):
        report = self.diff(
            "一季度新增12家。三季度又新增12家。",
            "一季度和三季度各新增家数相同。",
        )
        self.assertEqual(report["missing"][0]["value"], "12家")
        self.assertEqual(report["missing"][0]["original_count"], 2)
        self.assertEqual(report["missing"][0]["revised_count"], 0)

    def test_partial_occurrence_loss(self):
        report = self.diff("新增12家，其中12家已投产。", "新增12家，均已投产。")
        self.assertEqual(self.values(report["missing"]), [("数字", "12家")])
        self.assertEqual(report["missing"][0]["revised_count"], 1)

    def test_equivalent_forms_not_reported(self):
        report = self.diff(
            "覆盖十二个乡镇、三百五十项事项、1,200余家企业，百分之三十的事项实现网办，"
            "二〇二六年十一月完成，10月16号前上报，会议9：30召开，完成第三季度目标。",
            "覆盖12个乡镇、350项事项、1200余家企业，30%的事项实现网办，"
            "2026年11月完成，10月16日前上报，会议9:30召开，完成第3季度目标。",
        )
        self.assertTrue(report["consistent"], report)

    def test_percent_with_space_equivalent(self):
        report = self.diff("增长3.5 %。", "增长3.5%。")
        self.assertTrue(report["consistent"], report)

    def test_generic_chinese_numerals_ignored(self):
        report = self.diff(
            "统一部署，进一步推进，建立一个窗口，第一时间报告。",
            "统筹部署，持续推进，设立综合窗口，立即报告。",
        )
        self.assertTrue(report["consistent"], report)


class MarkerTest(FactDiffTestCase):

    def test_marker_lost(self):
        report = self.diff(
            "开展专项检查【待补：次数】次，整改率【待核实：以台账为准】。",
            "开展专项检查，整改率【待核实：以台账为准】。",
        )
        self.assertEqual(self.values(report["missing"]), [("标记", "【待补：次数】")])
        self.assertEqual(report["added"], [])

    def test_marker_text_changed(self):
        report = self.diff("全年接诊【待补：人次】人次。", "全年接诊【待补】人次。")
        self.assertEqual(self.values(report["missing"]), [("标记", "【待补：人次】")])
        self.assertEqual(self.values(report["added"]), [("标记", "【待补】")])


class TitleQuoteDocnoTest(FactDiffTestCase):

    def test_title_content_changed(self):
        report = self.diff(
            "根据《关于深化放管服改革的实施意见》要求，制定本方案。",
            "根据《关于深化放管服改革的意见》要求，制定本方案。",
        )
        self.assertEqual(self.values(report["missing"]), [("书名号", "《关于深化放管服改革的实施意见》")])
        self.assertEqual(self.values(report["added"]), [("书名号", "《关于深化放管服改革的意见》")])

    def test_numbers_inside_title_not_double_reported(self):
        report = self.diff("印发《2026年工作要点》。", "印发《2025年工作要点》。")
        kinds = {item["type"] for item in report["missing"] + report["added"]}
        self.assertEqual(kinds, {"书名号"})

    def test_quote_removed(self):
        report = self.diff("推行“最多跑一次”改革。", "推行最多跑一次改革。")
        self.assertEqual(self.values(report["missing"]), [("引号", "“最多跑一次”")])

    def test_docno_change_and_context_insensitive(self):
        report = self.diff(
            "根据某政发〔2026〕12号文件要求。",
            "按照某政发〔2026〕21号文件要求。",
        )
        self.assertEqual(self.values(report["missing"]), [("发文字号", "某政发〔2026〕12号")])
        self.assertEqual(self.values(report["added"]), [("发文字号", "某政发〔2026〕21号")])
        same = self.diff("根据某政发〔2026〕12号文件要求。", "按照某政发〔2026〕12号文件要求。")
        self.assertTrue(same["consistent"], same)

    def test_docno_without_agency_code(self):
        report = self.diff("根据〔2026〕12号文件要求。", "按照〔2026〕12号文件要求。")
        self.assertTrue(report["consistent"], report)


class SequenceTest(FactDiffTestCase):

    def test_sequence_renumbering_not_reported(self):
        original = (
            "一、主要工作\n\n（一）政务服务\n\n1. 建立清单，共12项。\n\n（1）完成培训3次。\n\n"
            "主要做法：1.建立台账；2.完善制度；（3）加强考核。\n\n"
            "第一，压实责任；第二，强化督导。"
        )
        revised = (
            "二、主要工作\n\n（三）政务服务\n\n3. 建立清单，共12项。\n\n（2）完成培训3次。\n\n"
            "主要做法：（1）建立台账；（2）完善制度；3.加强考核。\n\n"
            "一是压实责任；二是强化督导。"
        )
        report = self.diff(original, revised)
        self.assertTrue(report["consistent"], report)

    def test_markdown_heading_sequence_not_reported(self):
        report = self.diff("## 1. 工作成效\n\n**（二）**保障措施", "一、工作成效\n\n（一）保障措施")
        self.assertTrue(report["consistent"], report)


class BoundaryTest(FactDiffTestCase):

    def test_text_after_drafting_note_ignored(self):
        original = "全年完成12项。\n【起草说明】\n一、修改说明\n1. 原稿3.5天。"
        revised = "全年完成12项。\n\n【起草说明】\n一、修改说明\n1. 将4天改为5天，删除【待补：次数】。"
        report = self.diff(original, revised)
        self.assertTrue(report["consistent"], report)
        self.assertEqual(report["original_lines"], 1)

    def test_text_before_drafting_note_still_compared(self):
        report = self.diff("完成12项。\n【起草说明】\n无", "完成13项。\n【起草说明】\n无")
        self.assertEqual(self.values(report["missing"]), [("数字", "12项")])


class RealisticRewriteTest(FactDiffTestCase):

    def test_style_only_rewrite_is_consistent(self):
        original = (
            "值得注意的是，通过一年努力，全院门急诊量达到45.6万人次，同比增长8.2%，"
            "这不仅仅是数字的提升，更是服务能力的飞跃——四级手术占比提高至21.3%。"
        )
        revised = "全院门急诊量45.6万人次，同比增长8.2%；四级手术占比提高至21.3%。"
        report = self.diff(original, revised)
        self.assertTrue(report["consistent"], report)


class CliTest(FactDiffTestCase):

    def test_strict_exit_codes(self):
        a = self.write("a.txt", "完成12项。")
        b = self.write("b.txt", "完成13项。")
        c = self.write("c.txt", "共完成12项。")
        self.assertEqual(self.run_main(str(a), str(b), "--strict")[0], 1)
        self.assertEqual(self.run_main(str(a), str(b))[0], 0)
        self.assertEqual(self.run_main(str(a), str(c), "--strict")[0], 0)

    def test_text_output_has_two_lists(self):
        a = self.write("a.txt", "完成12项。")
        b = self.write("b.txt", "完成13项。")
        code, out, _ = self.run_main(str(a), str(b))
        self.assertEqual(code, 0)
        self.assertIn("一、原稿有而改稿没有（1 项）", out)
        self.assertIn("二、改稿新增（1 项）", out)
        self.assertIn("12项", out)

    def test_input_errors_return_2(self):
        a = self.write("a.txt", "完成12项。")
        code, _out, err = self.run_main(str(a), str(self.tmp / "不存在.txt"))
        self.assertEqual(code, 2)
        self.assertIn("文件不存在", err)
        pdf = self.write("b.pdf", "x")
        self.assertEqual(self.run_main(str(a), str(pdf))[0], 2)
        self.assertEqual(self.run_main(str(a))[0], 2)

    def test_md_and_gb18030_input(self):
        a = self.write("a.md", "完成12项，投入5万元。")
        b = self.write("b.txt", "完成12项，投入5万元。", encoding="gb18030")
        code, out, err = self.run_main(str(a), str(b), "--json")
        self.assertEqual(code, 0)
        self.assertTrue(json.loads(out)["consistent"])
        self.assertIn("GB18030", err)

    @unittest.skipUnless(HAS_DOCX, "未安装 python-docx")
    def test_docx_input(self):
        import docx as docx_module  # type: ignore
        document = docx_module.Document()
        document.add_paragraph("全年完成12项，投入66万元。")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "检查次数"
        table.cell(0, 1).text = "8次"
        document.add_paragraph("【起草说明】")
        document.add_paragraph("原稿为13项。")
        path = self.tmp / "改稿.docx"
        document.save(str(path))
        a = self.write("原稿.txt", "全年完成12项，投入66万元。\n检查次数\t9次")
        code, out, _ = self.run_main(str(a), str(path), "--json")
        self.assertEqual(code, 0)
        report = json.loads(out)
        self.assertEqual(self.values(report["missing"]), [("数字", "9次")])
        self.assertEqual(self.values(report["added"]), [("数字", "8次")])


class ReviewRegressionTest(FactDiffTestCase):
    """代码审查发现的缺陷回归测试。"""

    def test_docno_context_words_not_reported(self):
        # 原问题：发文字号前的“学习”“对照”“按”等上下文被并入代字，改稿只改了上下文也报“发文字号”增删
        pairs = (
            ("认真学习某政发〔2026〕12号文件精神。", "深入学习某政发〔2026〕12号文件精神。"),
            ("请各单位对照某政发〔2026〕12号文件开展自查。", "各单位要对照某政发〔2026〕12号文件开展自查。"),
            ("按某政发〔2026〕12号文件要求执行。", "依照某政发〔2026〕12号文件要求执行。"),
            ("传达某政发〔2026〕12号文件。", "以某政发〔2026〕12号文件为依据。"),
        )
        for original, revised in pairs:
            with self.subTest(original=original):
                report = self.diff(original, revised)
                self.assertTrue(report["consistent"], report)

    def test_docno_agency_change_still_reported(self):
        report = self.diff("认真学习京政发〔2026〕12号文件精神。", "深入学习沪政发〔2026〕12号文件精神。")
        self.assertEqual(self.values(report["missing"]), [("发文字号", "京政发〔2026〕12号")])
        self.assertEqual(self.values(report["added"]), [("发文字号", "沪政发〔2026〕12号")])

    def test_chinese_fraction_compared(self):
        # 原问题：“三分之一”改成“四分之一”不报差异
        report = self.diff("超过三分之一的村完成改造，千分之五的抽检不合格。", "超过四分之一的村完成改造，千分之五的抽检不合格。")
        self.assertEqual(self.values(report["missing"]), [("数字", "1/3")])
        self.assertEqual(self.values(report["added"]), [("数字", "1/4")])
        self.assertTrue(self.diff("三分之二以上", "三分之二以上")["consistent"])

    def test_markdown_drafting_note_boundary(self):
        # 原问题：“## 【起草说明】”“**【起草说明】**”不截断，起草说明里的数字参与比对（gongwen-format 已截断）
        report = self.diff("全年完成12项。\n## 【起草说明】\n原稿为13项。", "全年完成12项。\n**【起草说明】**\n改为14项。")
        self.assertTrue(report["consistent"], report)
        self.assertEqual(report["original_lines"], 1)

    @unittest.skipUnless(HAS_DOCX, "未安装 python-docx")
    def test_docx_merged_cells_counted_once(self):
        # 原问题：python-docx 对合并单元格逐格重复返回，横向合并 3 格的“100万元”被计为 3 处
        import docx as docx_module  # type: ignore
        document = docx_module.Document()
        table = document.add_table(rows=3, cols=3)
        table.cell(0, 0).merge(table.cell(0, 2)).text = "全年预算100万元"
        table.cell(1, 0).merge(table.cell(2, 0)).text = "参与单位3家"
        table.cell(1, 1).text = "一季度"
        table.cell(1, 2).text = "25万元"
        table.cell(2, 1).text = "二季度"
        table.cell(2, 2).text = "30万元"
        path = self.tmp / "原稿.docx"
        document.save(str(path))
        txt = self.write("改稿.txt", "全年预算100万元\n参与单位3家\t一季度\t25万元\n二季度\t30万元")
        code, out, _ = self.run_main(str(path), str(txt), "--json")
        self.assertEqual(code, 0)
        report = json.loads(out)
        self.assertTrue(report["consistent"], report)


class NormalizationTest(FactDiffTestCase):
    """只改写法、数值不变的，不报差异（端到端实测：review 直接改后的复检）。"""

    def test_date_zero_padding(self):
        report = self.diff("各学校于2026年11月05日前完成自查。", "各学校于2026年11月5日前完成自查。")
        self.assertTrue(report["consistent"], report)
        self.assertTrue(self.diff("截至2026-01-05。", "截至2026-1-5。")["consistent"])

    def test_range_unit_completion(self):
        for original, revised in (("可安排经费3～5万元。", "可安排经费3万～5万元。"),
                                  ("增幅控制在15～30%。", "增幅控制在15%～30%。"),
                                  ("2026—2028年实现全覆盖。", "2026年—2028年实现全覆盖。")):
            with self.subTest(original=original):
                self.assertTrue(self.diff(original, revised)["consistent"])

    def test_range_value_change_still_reported(self):
        report = self.diff("可安排经费3～5万元。", "可安排经费3万～6万元。")
        self.assertEqual(self.values(report["missing"]), [("数字", "5万元")])
        self.assertEqual(self.values(report["added"]), [("数字", "6万元")])

    def test_non_range_numbers_untouched(self):
        report = self.diff("联系电话：0571-87654321。", "联系电话：0571-87654321。")
        self.assertTrue(report["consistent"])
        self.assertIn(("数字", "571"), [(f.kind, f.value) for f in fact_diff.extract_line("0571-87654321", 1)])


LEDGER = """# 素材台账：某科室2026年工作总结

## 二、台账表

| 序号 | 类别 | 内容 | 来源 | 状态 |
|---|---|---|---|---|
| 1 | 数据 | 2026年1—11月完成PCI 612台 | 材料1第5段 | 已确认 |
| 2 | 数据 | 出院随访430人次 | 材料1第6段 | 已确认 |
| 3 | 数据 | 现有医生18人、护士36人 | 材料2第3行 | 已确认 |
"""


class LedgerModeTest(FactDiffTestCase):
    """--ledger：素材台账.md 对正文.txt，只列正文中有、台账中没有的事实，按集合比较。"""

    def ledger(self, body: str, *extra: str):
        a = self.write("素材台账.md", LEDGER)
        b = self.write("正文.txt", body)
        return self.run_main(str(a), str(b), "--ledger", *extra)

    def test_only_added_by_set(self):
        body = ("全年完成PCI 612台，出院随访430人次。PCI 612台中急诊占一半。\n"
                "现有医生18人、护士36人，新增设备2台。")
        code, out, _ = self.ledger(body, "--json")
        self.assertEqual(code, 0)
        report = json.loads(out)
        self.assertEqual(report["mode"], "ledger")
        self.assertNotIn("missing", report)
        self.assertEqual(self.values(report["added"]), [("数字", "2台")])
        self.assertFalse(report["consistent"])

    def test_ignore_seq_and_source_columns(self):
        lines = fact_diff.ledger_lines(LEDGER.split("\n"))
        joined = "\n".join(lines)
        self.assertNotIn("材料1第5段", joined)
        self.assertIn("类别\t内容\t状态", lines)
        self.assertIn("数据\t2026年1—11月完成PCI 612台\t已确认", lines)
        self.assertEqual(len(lines), len(LEDGER.split("\n")), "行号保持不变")
        # 正文中的“第5段”“第3行”式写法不能被台账的来源列覆盖
        code, out, _ = self.ledger("详见材料1第5段。", "--json")
        self.assertEqual(self.values(json.loads(out)["added"]), [("数字", "1"), ("数字", "5段")])

    def test_strict_and_text_output(self):
        code, out, _ = self.ledger("全年完成PCI 612台。", "--strict")
        self.assertEqual(code, 0)
        self.assertIn("正文中有、台账中没有（0 项）", out)
        self.assertNotIn("原稿有而改稿没有", out)
        code, out, _ = self.ledger("全年完成PCI 613台。", "--strict")
        self.assertEqual(code, 1)
        self.assertIn("613台", out)


class SpaceQuoteWanTest(FactDiffTestCase):
    """TL-1：数字与单位之间的空白、中英文引号、含万亿的汉字数字（v0.1 评测 gongwen-report 用例 5、6 误报）。"""

    # 必过：空白
    def test_space_between_number_and_unit_ignored(self):
        pairs = (
            ("门急诊3216 人次，开展培训12 次。", "门急诊3216人次，开展培训12次。"),
            ("门急诊3216　人次，增长3.5 %。", "门急诊3216人次，增长3.5%。"),
            ("投入320 万元，覆盖1,200 余家企业。", "投入320万元，覆盖1200余家企业。"),
            ("第 3 次会议于2026 年 10 月 16 日召开。", "第3次会议于2026年10月16日召开。"),
        )
        for original, revised in pairs:
            with self.subTest(original=original):
                self.assertTrue(self.diff(original, revised)["consistent"])

    def test_space_in_ledger_mode(self):
        # 复现 RP-3：台账写“3216 人次”“12 次”，正文连写，不应报台账没有
        ledger = self.write("素材台账.md", "| 序号 | 内容 | 来源 |\n|---|---|---|\n"
                                          "| 1 | 门急诊3216 人次 | 材料1 |\n| 2 | 培训12 次 | 材料2 |\n")
        body = self.write("正文.txt", "全年门急诊3216人次，开展培训12次。")
        code, out, _ = self.run_main(str(ledger), str(body), "--ledger", "--json", "--strict")
        self.assertEqual(code, 0, out)
        self.assertEqual(json.loads(out)["added"], [])

    def test_excerpt_keeps_original_spacing(self):
        facts = fact_diff.extract_line("门急诊3216 人次，满意", 1)
        self.assertEqual([(f.value, f.pos) for f in facts], [("3216人次", 3)])
        self.assertIn("3216 人次", facts[0].excerpt)

    def test_code_before_space_not_merged(self):
        # 编号“E1”后的空格不去掉，“1”不能和“家校协同”的“家”拼成“1家”
        self.assertEqual([f.value for f in fact_diff.extract_line("指标E1 家校协同", 1)], ["1"])
        self.assertEqual([f.value for f in fact_diff.extract_line("完成PCI 612 台", 1)], ["612台"])

    def test_tab_between_cells_not_removed(self):
        # 表格各列以制表符分隔，不能把下一列的文字当成单位
        values = [f.value for f in fact_diff.extract_line("数量\t12\t个人申报", 1)]
        self.assertEqual(values, ["12"])

    # 必过：引号
    def test_chinese_and_english_quotes_equivalent(self):
        for original, revised in (("推行\"最多跑一次\"改革。", "推行“最多跑一次”改革。"),
                                  ("推行「最多跑一次」改革。", "推行“最多跑一次”改革。"),
                                  ("推行“党建+业务”模式，打造『红色驿站』。", "推行\"党建+业务\"模式，打造‘红色驿站’。")):
            with self.subTest(original=original):
                self.assertTrue(self.diff(original, revised)["consistent"])

    # 必过：汉字万、亿
    def test_chinese_wan_and_yuan(self):
        pairs = (
            ("投入资金三百二十万。", "投入资金320万元。"),
            ("投入资金320万。", "投入资金320万元。"),
            ("总投资1.5亿元。", "总投资1.5亿。"),
            ("总投资十五亿。", "总投资15亿元。"),
            ("惠及三百二十万人次。", "惠及320万人次。"),
            ("投入320万余元。", "投入320万余。"),
        )
        for original, revised in pairs:
            with self.subTest(original=original):
                self.assertTrue(self.diff(original, revised)["consistent"])

    def test_chinese_decimal(self):
        # 原问题：“一点五亿”被拆成“5亿”；“三点五个百分点”被拆成“3点”“5个百分点”
        for original, revised in (("总盘子是一点五亿。", "总盘子是1.5亿元。"),
                                  ("提高三点五个百分点。", "提高3.5个百分点。"),
                                  ("增长二点零倍。", "增长2倍。")):
            with self.subTest(original=original):
                self.assertTrue(self.diff(original, revised)["consistent"])
        report = self.diff("总盘子是一点五亿。", "总盘子是1.6亿元。")
        self.assertEqual([i["value"] for i in report["missing"]], ["1.5亿"])
        self.assertEqual([f.value for f in fact_diff.extract_line("十二点三十分开会", 1)], ["12点"])

    def test_chinese_wan_in_ledger_mode(self):
        # 复现 W1 gongwen-minutes：记录写“三百二十万”，纪要写“320万元”
        ledger = self.write("原始记录.txt", "张主任：这个项目预算三百二十万，分两年支付。")
        body = self.write("纪要.txt", "会议同意项目预算320万元，分两年支付。")
        code, out, _ = self.run_main(str(ledger), str(body), "--ledger", "--json", "--strict")
        self.assertEqual(code, 0, out)

    def test_vague_numbers_marked_and_idioms_skipped(self):
        # “数十万”“几十万人次”是概数，进第三张清单，不按“10万”比对
        for text, value in (("惠及数十万群众。", "数十万"), ("几十万人次参与。", "几十万人次"), ("数十家企业", "数十家")):
            with self.subTest(text=text):
                facts = fact_diff.extract_line(text, 1)
                self.assertEqual([(f.value, f.hint) for f in facts], [(value, "概数")])
        for text in ("十万火急。", "数十位同志", "几百"):
            with self.subTest(text=text):
                self.assertEqual([f.value for f in fact_diff.extract_line(text, 1) if not f.hint], [])
        report = self.diff("惠及数十万群众。", "惠及30万群众。")
        self.assertEqual([i["value"] for i in report["added"]], ["30万"])
        self.assertEqual([(i["value"], i["list"]) for i in report["uncertain"]], [("数十万", "missing")])

    # 必错：数值或单位不同仍报
    def test_value_change_still_reported(self):
        cases = (
            ("门急诊3216 人次。", "门急诊3217人次。", "3216人次", "3217人次"),
            ("投入资金三百二十万。", "投入资金330万元。", "320万", "330万元"),
            ("投入资金320万元。", "惠及320万人。", "320万元", "320万人"),
            ("投入资金320万元。", "投入资金320元。", "320万元", "320元"),
            ("推行\"最多跑一次\"改革。", "推行“最多跑两次”改革。", "\"最多跑一次\"", "“最多跑两次”"),
        )
        for original, revised, gone, new in cases:
            with self.subTest(original=original):
                report = self.diff(original, revised)
                self.assertIn(gone, [item["value"] for item in report["missing"]])
                self.assertIn(new, [item["value"] for item in report["added"]])
                self.assertFalse(report["consistent"])
        a = self.write("a.txt", "门急诊3216 人次。")
        b = self.write("b.txt", "门急诊3217人次。")
        self.assertEqual(self.run_main(str(a), str(b), "--strict")[0], 1)


def _docx_bytes(body_xml: str) -> bytes:
    """用标准库拼一份最小 docx（含修订标记时不依赖 python-docx）。"""
    ns = ('xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
          'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"')
    document = f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:document {ns}><w:body>' \
               f'{body_xml}<w:sectPr/></w:body></w:document>'
    content_types = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/word/document.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
        '</Types>'
    )
    rels = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
        'Target="word/document.xml"/></Relationships>'
    )
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("_rels/.rels", rels)
        archive.writestr("word/document.xml", document)
    return buffer.getvalue()


def _run(text: str) -> str:
    return f'<w:r><w:t xml:space="preserve">{text}</w:t></w:r>'


REVISED_PARAGRAPHS = (
    # 全年排查隐患[删200][插218]处
    '<w:p>' + _run("全年排查隐患")
    + '<w:del w:id="1" w:author="审稿人" w:date="2026-10-16T09:00:00Z">'
      '<w:r><w:delText>200</w:delText></w:r></w:del>'
    + '<w:ins w:id="2" w:author="审稿人" w:date="2026-10-16T09:00:00Z">' + _run("218") + '</w:ins>'
    + _run("处，") + '</w:p>'
    # 移动：moveFrom 不计，moveTo 计入
    '<w:p><w:moveFrom w:id="3" w:author="审稿人">' + _run("整改率95%") + '</w:moveFrom>'
    + _run("完成整改") + '<w:moveTo w:id="4" w:author="审稿人">' + _run("，整改率96%") + '</w:moveTo></w:p>'
    # 表格单元格中的修订
    '<w:tbl><w:tr><w:tc><w:p>' + _run("检查次数") + '</w:p></w:tc>'
    '<w:tc><w:p><w:del w:id="5" w:author="审稿人"><w:r><w:delText>8次</w:delText></w:r></w:del>'
    '<w:ins w:id="6" w:author="审稿人">' + _run("9次") + '</w:ins></w:p></w:tc></w:tr></w:tbl>'
    '<w:p>' + _run("【起草说明】") + '</w:p>'
    '<w:p>' + _run("原稿为300处。") + '</w:p>'
)


class DocxRevisionTest(FactDiffTestCase):
    """W1-PL-2：docx 按“接受全部修订”读取，含 w:ins、不含 w:del（不依赖 python-docx）。"""

    def make_docx(self, name: str = "改稿.docx", body_xml: str = REVISED_PARAGRAPHS) -> Path:
        path = self.tmp / name
        path.write_bytes(_docx_bytes(body_xml))
        return path

    def test_read_accepts_revisions(self):
        text, revisions = fact_diff.read_document(self.make_docx())
        lines = fact_diff.body_lines(text)
        self.assertEqual(lines[0], "全年排查隐患218处，")
        self.assertEqual(lines[1], "完成整改，整改率96%")
        self.assertEqual(lines[2], "检查次数\t9次")
        self.assertEqual(len(lines), 3)  # 起草说明之后不读
        self.assertEqual(revisions, 6)

    def test_compare_after_accepting_revisions(self):
        path = self.make_docx()
        same = self.write("原稿.txt", "全年排查隐患218处，\n完成整改，整改率96%\n检查次数\t9次")
        code, out, _ = self.run_main(str(same), str(path), "--json", "--strict")
        self.assertEqual(code, 0, out)
        report = json.loads(out)
        self.assertTrue(report["consistent"])
        self.assertIn("改稿.docx 含修订痕迹 6 处，已按接受修订后的文字比对", report["notes"])
        old = self.write("旧稿.txt", "全年排查隐患200处，\n完成整改，整改率95%\n检查次数\t8次")
        code, out, _ = self.run_main(str(old), str(path))
        self.assertEqual(code, 0)
        self.assertIn("已按接受修订后的文字比对", out)
        self.assertIn("200处", out)
        self.assertIn("218处", out)

    def test_docx_without_revisions_has_no_note(self):
        path = self.make_docx(body_xml='<w:p>' + _run("全年完成12项。") + '</w:p>')
        a = self.write("原稿.txt", "全年完成12项。")
        code, out, _ = self.run_main(str(a), str(path), "--json")
        self.assertEqual(json.loads(out)["notes"], [])

    def test_textbox_paragraph_not_merged_and_sdt_read(self):
        body = ('<w:p>' + _run("正文12项")
                + '<w:r><w:drawing><w:txbxContent><w:p>' + _run("框内99项") + '</w:p></w:txbxContent>'
                  '</w:drawing></w:r></w:p>'
                '<w:sdt><w:sdtContent><w:p>' + _run("控件内5家") + '</w:p></w:sdtContent></w:sdt>')
        text, _ = fact_diff.read_document(self.make_docx(body_xml=body))
        self.assertEqual(text.split("\n"), ["正文12项", "控件内5家"])

    def test_broken_docx_returns_2(self):
        bad = self.tmp / "损坏.docx"
        bad.write_bytes(b"not a zip")
        a = self.write("a.txt", "完成12项。")
        code, _out, err = self.run_main(str(a), str(bad))
        self.assertEqual(code, 2)
        self.assertIn("无法打开 docx", err)
        empty = self.tmp / "空包.docx"
        with zipfile.ZipFile(empty, "w") as archive:
            archive.writestr("readme.txt", "x")
        code, _out, err = self.run_main(str(a), str(empty))
        self.assertEqual(code, 2)
        self.assertIn("word/document.xml", err)


class OrdinalApproxTest(FactDiffTestCase):
    """W1-SP-4：结构序数、成数概数单列第三张清单，不计入 --strict，但不丢弃。"""

    SPEECH_LEDGER = "会议材料：今年全区新建养老服务站12个，服务老人3200人次。"
    SPEECH = ("同志们：今天我讲三点意见。今年全区新建养老服务站12个，服务老人3200人次，"
              "群众满意度九成以上。下面我再提五项要求，重点抓好两件事，从三个方面推进。第三点，强化保障。")

    def ledger(self, body: str, *extra: str):
        a = self.write("材料.txt", self.SPEECH_LEDGER)
        b = self.write("讲话稿.txt", body)
        return self.run_main(str(a), str(b), "--ledger", *extra)

    def test_speech_ordinals_go_to_third_list(self):
        code, out, _ = self.ledger(self.SPEECH, "--json", "--strict")
        self.assertEqual(code, 0, out)
        report = json.loads(out)
        self.assertEqual(report["added"], [])
        uncertain = {(item["value"], item["hint"]) for item in report["uncertain"]}
        self.assertEqual(uncertain, {("3点", "序数"), ("9成", "概数"), ("5项", "序数"), ("2件", "序数"),
                                     ("3个", "序数"), ("第3点", "序数")})
        self.assertTrue(all(item["list"] == "added" for item in report["uncertain"]))

    def test_third_list_in_text_output(self):
        code, out, _ = self.ledger(self.SPEECH)
        self.assertEqual(code, 0)
        self.assertIn("可能是序数或概数（请人工判断，不计入 --strict）（6 项）", out)
        self.assertIn("9成", out)
        self.assertIn("可能是概数", out)

    def test_ratio_fact_kept_for_manual_check(self):
        # “9成患者满意”是数据，不能直接丢弃：留在第三张清单
        report = self.diff("患者满意度较高。", "9成患者满意。")
        self.assertEqual(report["added"], [])
        self.assertEqual([(i["value"], i["hint"], i["list"]) for i in report["uncertain"]], [("9成", "概数", "added")])
        report = self.diff("九成患者满意。", "8成患者满意。")
        self.assertEqual({(i["value"], i["list"]) for i in report["uncertain"]}, {("9成", "missing"), ("8成", "added")})
        self.assertTrue(report["consistent"])

    def test_quantities_stay_in_main_lists(self):
        # 必错：真实数量、时刻、大于十的项不进第三张清单
        report = self.diff("会议材料。", "下午3点开会，完成12项事项，办好十件实事，提出12项要求，新建3个站点。")
        values = [(item["value"], item["revised_count"]) for item in report["added"]]
        self.assertEqual(values, [("3点", 1), ("12项", 2), ("10件", 1), ("3个", 1)])
        self.assertEqual(report["uncertain"], [])
        a = self.write("a.txt", "会议材料。")
        b = self.write("b.txt", "下午3点开会。")
        self.assertEqual(self.run_main(str(a), str(b), "--strict")[0], 1)

    def test_ordinal_and_time_counted_separately(self):
        # “讲三点”与“下午3点”不相抵：改稿把时刻删掉、加了“讲三点”，时刻的删除照常报出
        report = self.diff("下午3点开会。", "开会时我讲三点。")
        self.assertEqual([i["value"] for i in report["missing"]], ["3点"])
        self.assertEqual(report["added"], [])
        self.assertEqual([(i["value"], i["list"]) for i in report["uncertain"]], [("3点", "added")])
        self.assertFalse(report["consistent"])


class SummaryCountTest(FactDiffTestCase):
    """v1.0 评测后第 1 项：做法、措施等的纯概括数（“三条做法”“措施共5项”）不报“正文有、台账无”，
    改为 P2 提示“概括数，请核对与分项一致”；真正的业务数字（新增5项制度、培训3场次）照报。"""

    LEDGER_TEXT = "会议材料：全年完成PCI 612台。"

    def ledger(self, body: str, *extra: str):
        a = self.write("素材台账.md", self.LEDGER_TEXT)
        b = self.write("正文.txt", body)
        return self.run_main(str(a), str(b), "--ledger", *extra)

    def test_summary_counts_become_p2_hint(self):
        cases = {
            "具体做法有三条。": "3条",              # 数字在后：做法有N条
            "检查中推行分级管理，具体做法有3条。": "3条",
            "我们采取了三条做法，成效明显。": "3条",  # 数字在前：N条做法
            "总结出5项措施。": "5项",
            "措施共5项：一是排查，二是整改。": "5项",
            "主要经验有三点。": "3点",
            "从3方面入手。": "3",
            "形成五项主要措施。": "5项",
            "主要有三个问题，一是。": "3个",
        }
        for body, value in cases.items():
            code, out, _ = self.ledger(body, "--json", "--strict")
            self.assertEqual(code, 0, (body, out))
            report = json.loads(out)
            self.assertEqual(report["added"], [], body)
            self.assertEqual([(i["value"], i["hint"]) for i in report["uncertain"]], [(value, "概括数")], body)

    def test_summary_item_carries_level_and_note(self):
        code, out, _ = self.ledger("具体做法有三条。", "--json")
        item = json.loads(out)["uncertain"][0]
        self.assertEqual((item["hint"], item["level"]), ("概括数", "P2"))
        self.assertIn("概括数，请核对与分项一致", item["note"])
        code, out, _ = self.ledger("具体做法有三条。")
        self.assertEqual(code, 0)
        self.assertIn("概括数，请核对与分项一致", out)
        self.assertIn("P2", out)

    def test_real_business_numbers_still_reported(self):
        # 宁可少放行：带业务动词、不带概括词、概括词后不是句读、超过十的，一律留在“正文有、台账无”
        cases = {
            "新增5项制度。": "5项",
            "培训3场次。": "3场",
            "新增3项措施。": "3项",
            "累计出台5项措施。": "5项",
            "发现3个问题。": "3个",
            "查处问题共5个。": "5个",
            "3项任务已全部完成。": "3项",
            "完成任务3项。": "3项",
            "措施有12项。": "12项",
            "下午3点开会。": "3点",
            "推进3项重点工作。": "3项",
            "措施共5项制度。": "5项",
        }
        for body, value in cases.items():
            code, out, _ = self.ledger(body, "--json", "--strict")
            report = json.loads(out)
            self.assertEqual(code, 1, (body, out))
            self.assertIn(value, [i["value"] for i in report["added"]], body)
            self.assertEqual(report["uncertain"], [], body)

    def test_summary_and_business_number_in_same_text(self):
        body = "具体做法有三条。全年新增5项制度，培训3场次。"
        code, out, _ = self.ledger(body, "--json", "--strict")
        report = json.loads(out)
        self.assertEqual(code, 1)
        self.assertEqual([i["value"] for i in report["added"]], ["5项", "3场"])
        self.assertEqual([(i["value"], i["hint"]) for i in report["uncertain"]], [("3条", "概括数")])

    def test_summary_in_diff_mode_goes_to_third_list(self):
        report = self.diff("工作开展顺利。", "具体做法有三条。")
        self.assertEqual(report["added"], [])
        self.assertEqual([(i["value"], i["hint"]) for i in report["uncertain"]], [("3条", "概括数")])


class MultiRefTest(FactDiffTestCase):
    """W1-PL-3：--ref 追加原稿，合并作比对基准，注明出处。"""

    def setUp(self) -> None:
        super().setUp()
        self.a = self.write("甲稿.txt", "全年完成12项，投入66万元。\n【起草说明】\n甲稿另有13项。")
        self.b = self.write("乙稿.txt", "全年完成12项，新增企业3家。\n【起草说明】\n乙稿备注5家。")

    def test_merged_draft_consistent(self):
        merged = self.write("合并稿.txt", "全年完成12项，投入66万元，新增企业3家。")
        code, out, _ = self.run_main(str(self.a), str(merged), "--ref", str(self.b), "--json", "--strict")
        self.assertEqual(code, 0, out)
        report = json.loads(out)
        self.assertTrue(report["consistent"])
        self.assertEqual(report["original_files"], ["甲稿.txt", "乙稿.txt"])
        # 不加 --ref 时，乙稿的事实被报为新增
        code, out, _ = self.run_main(str(self.a), str(merged), "--json")
        self.assertEqual([i["value"] for i in json.loads(out)["added"]], ["3家"])

    def test_missing_item_names_source(self):
        merged = self.write("合并稿.txt", "全年完成12项，投入66万元，新增企业4家。")
        code, out, _ = self.run_main(str(self.a), str(merged), "--ref", str(self.b), "--json")
        report = json.loads(out)
        self.assertEqual([(i["value"], i["sources"]) for i in report["missing"]], [("3家", ["乙稿.txt"])])
        self.assertEqual(report["missing"][0]["locations"][0]["file"], "乙稿.txt")
        self.assertEqual([i["value"] for i in report["added"]], ["4家"])
        code, out, _ = self.run_main(str(self.a), str(merged), "--ref", str(self.b))
        self.assertIn("出自 乙稿.txt", out)
        self.assertIn("乙稿.txt 第 1 行", out)
        self.assertIn("甲稿.txt、乙稿.txt（共 2 份）", out)

    def test_max_count_across_drafts(self):
        # 两稿各写一次“12项”，合并稿写一次不报；合并稿写两次报新增
        once = self.write("合并稿.txt", "全年完成12项，投入66万元，新增企业3家。")
        self.assertTrue(self.diff_refs(once)["consistent"])
        twice = self.write("合并稿2.txt", "全年完成12项，投入66万元，新增企业3家。其中12项已验收。")
        report = self.diff_refs(twice)
        self.assertEqual([(i["value"], i["original_count"], i["revised_count"]) for i in report["added"]],
                         [("12项", 1, 2)])

    def diff_refs(self, merged: Path) -> dict:
        code, out, _ = self.run_main(str(self.a), str(merged), "--ref", str(self.b), "--json")
        self.assertEqual(code, 0)
        return json.loads(out)

    def test_ledger_with_refs(self):
        body = self.write("正文.txt", "全年完成12项，投入66万元，新增企业3家。")
        code, out, _ = self.run_main(str(self.a), str(body), "--ledger", "--ref", str(self.b), "--strict")
        self.assertEqual(code, 0, out)
        code, _out, _ = self.run_main(str(self.a), str(body), "--ledger", "--strict")
        self.assertEqual(code, 1)

    def test_missing_ref_returns_2(self):
        body = self.write("正文.txt", "全年完成12项。")
        code, _out, err = self.run_main(str(self.a), str(body), "--ref", str(self.tmp / "不存在.txt"))
        self.assertEqual(code, 2)
        self.assertIn("文件不存在", err)

    def test_same_name_in_different_dirs(self):
        sub = self.tmp / "sub"
        sub.mkdir()
        other = sub / "甲稿.txt"
        other.write_bytes("新增企业3家。".encode("utf-8"))
        merged = self.write("合并稿.txt", "全年完成12项，投入66万元。")
        code, out, _ = self.run_main(str(self.a), str(merged), "--ref", str(other), "--json")
        report = json.loads(out)
        self.assertEqual(len(set(report["original_files"])), 2)
        self.assertEqual(report["missing"][0]["sources"], [str(other)])


class LongPathTest(unittest.TestCase):

    @unittest.skipUnless(sys.platform.startswith("win"), "只在 Windows 下加长路径前缀")
    def test_long_path_prefixed_on_windows(self):
        long_path = Path("C:/" + "a" * 250 + "/原稿.txt")
        self.assertTrue(str(fact_diff._long_path(long_path)).startswith("\\\\?\\"))
        short_path = Path("C:/temp/原稿.txt")
        self.assertEqual(fact_diff._long_path(short_path), short_path)


class CnToIntTest(unittest.TestCase):

    def test_values(self):
        cases = {
            "十": 10, "十五": 15, "二十": 20, "三百零五": 305, "三百五": 350, "一千二百": 1200,
            "一万二千": 12000, "两千": 2000, "二〇二六": 2026,
        }
        for text, value in cases.items():
            self.assertEqual(fact_diff.cn_to_int(text), value, text)

    def test_invalid(self):
        for text in ("十四五", "三四", "七一", "一一", "百"):
            self.assertIsNone(fact_diff.cn_to_int(text), text)


if __name__ == "__main__":
    unittest.main()
