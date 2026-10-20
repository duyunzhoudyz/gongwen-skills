# -*- coding: utf-8 -*-
"""check_docx 标本测试：合规样稿必须通过；每条关键规则各造一份“必错标本”，必须报出。"""
import io
import json
import unittest
from contextlib import redirect_stdout

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

from tests.format.helpers import OUT, build_sample, build_text, check_docx

BASE_TEXT = "关于某事的通知\n\n各单位：\n\n一、工作要求\n\n现将有关事项通知如下。\n\n某市某局\n\n2026年3月5日"


def ids(result):
    return {(i["id"], i["level"]) for i in result["issues"]}


class SpecimenTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.good, _ = build_text(BASE_TEXT, "标本-合规.docx", meta={"版式": "白头"})

    def mutate(self, name, fn):
        doc = Document(str(self.good))
        fn(doc)
        path = OUT / name
        doc.save(str(path))
        return check_docx.check_document(path)

    def test_good_passes(self):
        result = check_docx.check_document(self.good)
        self.assertEqual(result["summary"]["P0"], 0)
        self.assertEqual(result["summary"]["P1"], 0)
        self.assertEqual(result["doc_type"], "通知")
        for key in ("file", "doc_type", "char_count", "summary", "issues"):
            self.assertIn(key, result)  # 与开发规范第 8 节一致

    def test_letter_paper(self):
        def fn(doc):
            doc.sections[0].page_width = Mm(216)
            doc.sections[0].page_height = Mm(279)
        self.assertIn(("DOCX-001", "P0"), ids(self.mutate("标本-纸张.docx", fn)))

    def test_top_margin(self):
        def fn(doc):
            doc.sections[0].top_margin = Mm(25.4)
        self.assertIn(("DOCX-002", "P0"), ids(self.mutate("标本-天头.docx", fn)))

    def test_right_margin_p1(self):
        def fn(doc):
            doc.sections[0].right_margin = Mm(31.8)
        self.assertIn(("DOCX-003", "P1"), ids(self.mutate("标本-右边距.docx", fn)))

    def _body(self, doc):
        return [p for p in doc.paragraphs if p.text.startswith("现将")][0]

    def test_body_font(self):
        def fn(doc):
            for r in self._body(doc).runs:
                r._r.rPr.find(qn("w:rFonts")).set(qn("w:eastAsia"), "宋体")
        self.assertIn(("DOCX-010", "P0"), ids(self.mutate("标本-正文字体.docx", fn)))

    def _set_font(self, paragraph, name):
        for r in paragraph.runs:
            r._r.rPr.find(qn("w:rFonts")).set(qn("w:eastAsia"), name)

    def _hits_016(self, result):
        return [i for i in result["issues"] if i["id"] == "DOCX-016"]

    def test_non_standard_body_font_is_p1(self):
        """DOCX-016（严格模式，P1）：字体类别对、但不是标准字体本身（仿宋、方正仿宋_GBK），消息写明应为仿宋_GB2312。"""
        for k, font in enumerate(("仿宋", "方正仿宋_GBK", "STFangsong")):
            result = self.mutate(f"标本-正文非标准字体{k}.docx", lambda doc, f=font: self._set_font(self._body(doc), f))
            hits = self._hits_016(result)
            self.assertEqual([(h["level"], h["category"]) for h in hits], [("P1", "字体")], font)
            self.assertIn(f"“{font}”", hits[0]["message"])
            self.assertIn("应为“仿宋_GB2312”", hits[0]["message"])
            self.assertIn("不要用其他字体代替", hits[0]["suggestion"])
            self.assertNotIn("缺字体时用", hits[0]["suggestion"])
            self.assertNotIn("DOCX-010", {i["id"] for i in result["issues"] if i["level"] == "P0"})

    def test_standard_font_aliases_pass(self):
        """同一款字体的英文名（FangSong_GB2312、SimHei、FZXiaoBiaoSong-B05S）不算非标准字体。"""
        def fn(doc):
            self._set_font(self._body(doc), "FangSong_GB2312")
            self._set_font([p for p in doc.paragraphs if p.text.startswith("一、")][0], "SimHei")
            self._set_font(doc.paragraphs[0], "FZXiaoBiaoSong-B05S")
        result = self.mutate("标本-字体别名.docx", fn)
        self.assertEqual(self._hits_016(result), [])
        self.assertEqual((result["summary"]["P0"], result["summary"]["P1"]), (0, 0))

    def test_non_standard_heading_font_is_p1(self):
        def fn(doc):
            self._set_font([p for p in doc.paragraphs if p.text.startswith("一、")][0], "方正黑体_GBK")
        hits = self._hits_016(self.mutate("标本-一级标题非标准字体.docx", fn))
        self.assertEqual([(h["level"], h["id"]) for h in hits], [("P1", "DOCX-016")])
        self.assertIn("一级标题", hits[0]["message"])
        self.assertIn("应为“黑体”", hits[0]["message"])

    def test_non_standard_title_font_is_p1(self):
        """标题用华文中宋、宋体（含加粗）、方正小标宋_GBK 代替方正小标宋简体：DOCX-016 P1；其他字体仍是 DOCX-020。"""
        for k, font in enumerate(("华文中宋", "宋体", "方正小标宋_GBK", "FZXiaoBiaoSong-B05")):
            result = self.mutate(f"标本-标题非标准字体{k}.docx", lambda doc, f=font: self._set_font(doc.paragraphs[0], f))
            hits = self._hits_016(result)
            self.assertEqual([(h["level"], h["category"]) for h in hits], [("P1", "字体")], font)
            self.assertIn("应为“方正小标宋简体”", hits[0]["message"])
            self.assertIn(f"“{font}”", hits[0]["message"])
            self.assertNotIn("DOCX-020", {i["id"] for i in result["issues"]})
        result = self.mutate("标本-标题其他字体.docx", lambda doc: self._set_font(doc.paragraphs[0], "微软雅黑"))
        self.assertEqual(self._hits_016(result), [])
        self.assertIn(("DOCX-020", "P1"), ids(result))

    def test_allow_demotes_non_standard_font(self):
        """--allow DOCX-016 的机制不变：本单位制度明文规定用其他字体的，降为 P2 并注明按本单位制度保留。"""
        self.mutate("标本-允许非标准字体.docx", lambda doc: self._set_font(self._body(doc), "仿宋"))
        result = check_docx.check_document(OUT / "标本-允许非标准字体.docx", allow="DOCX-016")
        hits = self._hits_016(result)
        self.assertEqual([h["level"] for h in hits], ["P2"])
        self.assertTrue(hits[0]["message"].startswith("按本单位制度保留："))
        self.assertEqual(result["summary"]["P1"], 0)

    def test_theme_font_is_wrong(self):
        def fn(doc):
            for r in self._body(doc).runs:
                r._r.rPr.find(qn("w:rFonts")).set(qn("w:eastAsiaTheme"), "minorEastAsia")
        self.assertIn(("DOCX-010", "P0"), ids(self.mutate("标本-主题字体.docx", fn)))

    def test_body_size(self):
        def fn(doc):
            for r in self._body(doc).runs:
                r.font.size = Pt(12)
        self.assertIn(("DOCX-011", "P0"), ids(self.mutate("标本-正文字号.docx", fn)))

    def test_line_spacing(self):
        def fn(doc):
            sp = self._body(doc)._p.pPr.find(qn("w:spacing"))
            sp.set(qn("w:line"), "360")
            sp.set(qn("w:lineRule"), "auto")
        self.assertIn(("DOCX-012", "P1"), ids(self.mutate("标本-行距.docx", fn)))

    def test_first_line_indent(self):
        def fn(doc):
            ind = self._body(doc)._p.pPr.find(qn("w:ind"))
            ind.set(qn("w:firstLineChars"), "0")
            ind.set(qn("w:firstLine"), "0")
        self.assertIn(("DOCX-013", "P1"), ids(self.mutate("标本-缩进.docx", fn)))

    def test_h1_font(self):
        def fn(doc):
            p = [p for p in doc.paragraphs if p.text.startswith("一、")][0]
            for r in p.runs:
                r._r.rPr.find(qn("w:rFonts")).set(qn("w:eastAsia"), "楷体_GB2312")
        self.assertIn(("DOCX-030", "P1"), ids(self.mutate("标本-一级标题.docx", fn)))

    def test_title_size(self):
        def fn(doc):
            for r in doc.paragraphs[0].runs:
                r.font.size = Pt(18)
        self.assertIn(("DOCX-021", "P1"), ids(self.mutate("标本-标题字号.docx", fn)))

    def test_missing_page_number(self):
        def fn(doc):
            for p in doc.sections[0].footer.paragraphs:
                for child in list(p._p):
                    if child.tag != qn("w:pPr"):
                        p._p.remove(child)
        self.assertIn(("DOCX-070", "P1"), ids(self.mutate("标本-无页码.docx", fn)))

    def test_bad_page_number_text(self):
        def fn(doc):
            for t in doc.sections[0].footer._element.iter(qn("w:t")):
                t.text = t.text.replace("—", "-")
        self.assertIn(("DOCX-071", "P1"), ids(self.mutate("标本-页码格式.docx", fn)))

    def test_docno_rules(self):
        cases = {"某发[2026]5号": "DOCX-050", "某发〔26〕5号": "DOCX-051", "某发〔2026〕第5号": "DOCX-052",
                 "某发〔2026〕05号": "DOCX-053"}
        for k, (docno, iid) in enumerate(cases.items()):
            path, _ = build_text(BASE_TEXT, f"标本-发文字号{k}.docx",
                                 meta={"版式": "红头", "发文机关标志": "某市某局文件", "发文字号": docno,
                                       "印发机关": "某市某局办公室", "印发日期": "2026-03-05"})
            self.assertIn((iid, "P0"), ids(check_docx.check_document(path)), docno)

    def test_date_zero_padding(self):
        path, _ = build_text(BASE_TEXT.replace("2026年3月5日", "2026年03月05日"), "标本-日期虚位.docx")
        self.assertIn(("DOCX-043", "P0"), ids(check_docx.check_document(path)))

    def test_text_rules(self):
        text = ("关于某事的通知\n\n各单位:\n\n1、工作要求\n\n现通知如下。\n\n附件：报名表。\n\n某市某局\n\n2026年3月5日\n\n"
                "抄送：市委办公室")
        path, _ = build_text(text, "标本-文字规则.docx", meta={"版式": "白头", "印发机关": "某局办公室", "印发日期": "2026-03-05"})
        found = ids(check_docx.check_document(path))
        for expect in (("DOCX-033", "P1"), ("DOCX-042", "P1"), ("DOCX-060", "P1")):
            self.assertIn(expect, found)

    def test_recipient_half_colon(self):
        def fn(doc):
            p = [p for p in doc.paragraphs if p.text == "各单位："][0]
            p.runs[-1].text = p.runs[-1].text.replace("：", ":")
        self.assertIn(("DOCX-041", "P1"), ids(self.mutate("标本-主送冒号.docx", fn)))

    def test_letter_first_page_number(self):
        path, _ = build_sample("信函", "标本-信函.docx")
        doc = Document(str(path))
        doc.sections[0].different_first_page_header_footer = False
        bad = OUT / "标本-信函首页页码.docx"
        doc.save(str(bad))
        self.assertIn(("DOCX-074", "P1"), ids(check_docx.check_document(bad, layout="信函")))

    def test_cli_strict_exit_code(self):
        def fn(doc):
            doc.sections[0].top_margin = Mm(20)
        self.mutate("标本-严格模式.docx", fn)
        buf = io.StringIO()
        with redirect_stdout(buf):
            code = check_docx.main([str(OUT / "标本-严格模式.docx"), "--json", "--strict"])
        self.assertEqual(code, 1)
        data = json.loads(buf.getvalue())
        self.assertGreaterEqual(data["summary"]["P0"], 1)
        with redirect_stdout(io.StringIO()):
            self.assertEqual(check_docx.main([str(self.good), "--strict"]), 0)
            self.assertEqual(check_docx.main([str(OUT / "不存在.docx")]), 2)

    def test_markers_are_p2(self):
        """【待补】【待核实】标记按 P2 报告（定稿提醒），不计入 P1。"""
        path, _ = build_text(BASE_TEXT.replace("现将有关事项通知如下。", "全年检查【待补：次数】次，【待核实：金额】万元。"),
                             "标本-待补标记.docx")
        result = check_docx.check_document(path)
        markers = [i for i in result["issues"] if i["id"] == "DOCX-080"]
        self.assertEqual(len(markers), 2)
        self.assertEqual({i["level"] for i in markers}, {"P2"})
        self.assertEqual(result["summary"]["P1"], 0)

    def test_ordinal_1_dun_suggestion(self):
        """“1、”直接在“二、”下时按层次应为“（一）”：建议不能只写“改为 1.”，要同时提示第二层用“（一）”。"""
        text = BASE_TEXT.replace("一、工作要求", "二、检查范围").replace("现将有关事项通知如下。", "1、校舍安全。重点检查教学楼。")
        path, _ = build_text(text, "标本-序号1顿号.docx")
        hits = [i for i in check_docx.check_document(path)["issues"] if i["id"] == "DOCX-033"]
        self.assertEqual(len(hits), 1)
        self.assertNotIn("三级", hits[0]["message"])
        self.assertIn("（一）", hits[0]["suggestion"])
        self.assertIn("1.", hits[0]["suggestion"])

    def test_cli_unknown_layout(self):
        """--layout 写成不认识的值（如“函件”“蓝头”）时报输入错误，不能按未知版式静默检查（会漏查信函首页页码）。"""
        for bad in ("函件", "蓝头"):
            buf = io.StringIO()
            with redirect_stdout(buf):
                code = check_docx.main([str(self.good), "--layout", bad])
            self.assertEqual(code, 2, bad)
            self.assertIn("无法识别的版式", buf.getvalue())

    def test_draft_note_stops_checking(self):
        """docx 中有【起草说明】时，之后的内容不参与识别和检查（output-format.md 第二部分），并提示删除。

        原问题：署名、日期不在末尾，成文日期虚位只报 P1（DOCX-090）而漏报 P0（DOCX-043）；
        起草说明里的“〔26〕”误报 DOCX-051。
        """
        def fn(doc):
            for text in ("【起草说明】", "1. 所作假设：检查次数为【待补：次数】。", "2. 原文号写成某发〔26〕第05号，待核实"):
                doc.add_paragraph(text)
            date = [p for p in doc.paragraphs if p.text == "2026年3月5日"][0]
            date.runs[0].text = "2026年03月05日"
        result = self.mutate("标本-含起草说明.docx", fn)
        found = ids(result)
        self.assertIn(("DOCX-043", "P0"), found)
        self.assertNotIn(("DOCX-051", "P0"), found)
        self.assertNotIn(("DOCX-052", "P0"), found)
        self.assertFalse(any("【待补：次数】" in i["excerpt"] for i in result["issues"]))
        notes = [i for i in result["issues"] if "起草说明" in i["message"]]
        self.assertEqual(len(notes), 1)
        self.assertEqual((notes[0]["id"], notes[0]["level"]), ("DOCX-082", "P1"))
        # 字数只统计【起草说明】之前：比合规稿只多日期中的两个“0”
        good = check_docx.check_document(self.good)
        self.assertEqual(result["char_count"], good["char_count"] + 2)



if __name__ == "__main__":
    unittest.main()
