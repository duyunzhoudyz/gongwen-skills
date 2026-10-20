# -*- coding: utf-8 -*-
"""gongwen-format 单元测试的公共函数：定位脚本、生成样稿、构造格式混乱的 docx。"""
from __future__ import annotations

import contextlib
import json
import re
import sys
import zipfile
from pathlib import Path
from unittest import mock

REPO = Path(__file__).resolve().parents[2]
SKILL = REPO / "skills" / "gongwen-format"
SCRIPTS = SKILL / "scripts"
FILES = SKILL / "evals" / "files"
OUT = Path(__file__).resolve().parent / "_out"
OUT.mkdir(parents=True, exist_ok=True)
if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))

import build_docx  # noqa: E402
import check_docx  # noqa: E402
import gw_docx  # noqa: E402
import gw_fonts  # noqa: E402
import gw_layout  # noqa: E402

# 六种版式的样稿输入：(版式, 正文文件, 版式信息文件)
SAMPLES = {
    "红头": ("红头通知-正文.txt", "红头通知-版式信息.json"),
    "信函": ("信函-正文.txt", "信函-版式信息.json"),
    "纪要": ("纪要-正文.txt", "纪要-版式信息.json"),
    "白头": ("白头请示-正文.txt", None),
    "材料": ("年终总结-正文.txt", None),
    "简报": ("简报-正文.txt", "简报-版式信息.json"),
}
# 测试中不依赖本机字体：用指定的“已安装字体名集合”（gw_fonts.norm 规范化后的写法）模拟各种情况
NO_FONTS: set = set()  # 空集合表示无法检测本机字体
FULL_FONTS = {"仿宋_gb2312", "楷体_gb2312", "黑体", "方正小标宋简体", "宋体"}  # 标准字体齐全
# 现行简体中文版 Windows 的典型情况：有黑体、宋体，以及另一款“仿宋”“楷体”，缺仿宋_GB2312、楷体_GB2312、方正小标宋简体
WIN_FONTS = {"仿宋", "楷体", "黑体", "宋体", "fangsong", "kaiti", "simhei", "simsun"}
STANDARD_FONTS = {"仿宋": "仿宋_GB2312", "楷体": "楷体_GB2312", "黑体": "黑体", "小标宋": "方正小标宋简体", "宋体": "宋体"}


@contextlib.contextmanager
def fonts_installed(installed=FULL_FONTS, method="测试用字体集合"):
    """让 gw_fonts.installed_fonts 返回指定的已安装字体集合：直接调用脚本 main 的测试用它，不依赖本机字体。"""
    with mock.patch.object(gw_fonts, "installed_fonts", return_value=(set(installed), method)):
        yield


def written_fonts(path: Path) -> set:
    """docx 各部分（正文、样式、页脚等）rFonts 中写入的 eastAsia 字体名。"""
    names: set = set()
    with zipfile.ZipFile(path) as z:
        for name in z.namelist():
            if name.startswith("word/") and name.endswith(".xml"):
                names |= set(re.findall(r'<w:rFonts[^>]*?w:eastAsia="([^"]+)"', z.read(name).decode("utf-8")))
    return names


def read_lines(name: str):
    lines, _ = gw_layout.read_body_text(FILES / name)
    return lines


def read_meta(name):
    if not name:
        return {}
    return json.loads((FILES / name).read_text(encoding="utf-8"))


def build_sample(layout: str, out_name: str, *, installed=None, meta_update=None):
    """按版式生成样稿，返回 (路径, 报告)。"""
    text_name, meta_name = SAMPLES[layout]
    meta = read_meta(meta_name)
    if meta_update:
        meta.update(meta_update)
    doc, report = build_docx.build(read_lines(text_name), meta, layout=layout,
                                   installed=NO_FONTS if installed is None else installed)
    path = OUT / out_name
    if path.exists():
        path.unlink()
    doc.save(str(path))
    return path, report


def build_text(text: str, out_name: str, meta=None, layout=None):
    doc, report = build_docx.build(text.split("\n"), meta or {}, layout=layout, installed=NO_FONTS)
    path = OUT / out_name
    if path.exists():
        path.unlink()
    doc.save(str(path))
    return path, report


def make_messy_docx(path: Path) -> Path:
    """构造一份格式混乱的 docx：Letter 纸、1 英寸边距、宋体小四、单倍行距、空格缩进、段间空行、无页码。"""
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for side in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, side, Inches(1))

    def para(text, *, bold=False, size=12, font="宋体", center=False):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = bold
        run.font.size = Pt(size)
        rpr = run._r.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = rpr.makeelement(qn("w:rFonts"), {})
            rpr.insert(0, rfonts)
        rfonts.set(qn("w:eastAsia"), font)
        rfonts.set(qn("w:ascii"), "Calibri")
        if center:
            p.alignment = 1
        return p

    para("关于2026年度某区政务服务工作情况的报告", bold=True, size=16, font="黑体")
    doc.add_paragraph()
    para("　　区人民政府：")
    doc.add_paragraph()
    para("　　2026年，区政务服务中心围绕“高效办成一件事”，推进政务服务标准化、规范化、便利化。现将有关情况报告如下。")
    doc.add_paragraph()
    para("一、主要工作", bold=True)
    para("（一）推进集成办理。梳理企业开办、不动产登记等高频事项38项，实行一窗受理、并联审批。")
    para("（二）压缩办理时限。全年平均办件时长由3.5天压缩至1.2天。")
    para("1.优化审批流程。取消证明材料21项。")
    doc.add_paragraph()
    para("二、下一步打算", bold=True)
    para("　　继续推进政务服务事项标准化，扩大“一网通办”覆盖范围。")
    doc.add_paragraph()
    para("特此报告。")
    doc.add_paragraph()
    doc.add_paragraph()
    doc.add_paragraph()
    para("某区政务服务中心")
    para("2026年11月20日")
    doc.save(str(path))
    return path
