# -*- coding: utf-8 -*-
"""已有 docx 的“抽取—重排”路线：版头无红线、签发人另行、版记无线、附件紧跟正文的混乱稿，抽取后用 build --keep-text 重排。"""
from __future__ import annotations

import json
import re
import unittest

from .helpers import OUT, build_docx, check_docx, make_messy_docx

import extract_docx

BODY = [
    "关于加强全区政务服务窗口管理的请示",
    "区人民政府：",
    "　　为进一步规范窗口服务，现提出如下意见。",
    "一、主要措施",
    "（一）统一服务标准。全面推行首问负责制。",
    "二、保障措施",
    "　　各部门要加强协同。",
    "妥否，请批示。",
    "附件：1.窗口服务规范",
    "2.考核办法",
]
SIGNOFF = ["某区政务服务中心", "2026年11月20日"]
ATTACH = ["附件1", "窗口服务规范", "　　第一条　窗口人员统一着装。", "附件2", "考核办法", "　　第一条　按季度考核。"]
TAIL = ["抄送：区委办公室，区人大办公室。", "某区政务服务中心办公室　2026年11月21日印发"]


def make_messy_header_docx(path):
    from docx import Document
    doc = Document()
    doc.add_paragraph("某区政务服务中心文件")  # 无红色、不居中
    doc.add_paragraph("某政发〔2026〕5号")
    doc.add_paragraph("签发人：张三")  # 另起一行
    for t in BODY + SIGNOFF + ATTACH + TAIL:
        doc.add_paragraph(t)
    doc.save(str(path))
    return path


def para_texts(path):
    from docx import Document
    from docx.oxml.ns import qn
    doc = Document(str(path))
    out = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for tc in t._tbl.iter(qn("w:tc")):
            out += ["".join(x.text or "" for x in p.iter(qn("w:t"))) for p in tc.findall(qn("w:p"))]
    return out


def norm(items):
    return sorted(re.sub(r"\s+", "", t) for t in items if t.strip())


class ExtractRebuildTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.src = make_messy_header_docx(OUT / "extract-messy.docx")
        cls.res = extract_docx.extract(cls.src)
        txt = OUT / "extract-正文.txt"
        txt.write_text("\n".join(cls.res["lines"]) + "\n", encoding="utf-8")
        cls.dst = OUT / "extract-rebuilt.docx"
        if cls.dst.exists():
            cls.dst.unlink()
        doc, cls.report = build_docx.build(cls.res["lines"], cls.res["meta"], keep_text=True, installed=set())
        doc.save(str(cls.dst))

    def test_meta(self):
        m = self.res["meta"]
        self.assertEqual(m["版式"], "红头")
        self.assertEqual(m["发文机关标志"], "某区政务服务中心文件")
        self.assertEqual(m["发文字号"], "某政发〔2026〕5号")
        self.assertEqual(m["签发人"], "张三")
        self.assertTrue(m["抄送"].startswith("抄送："))
        self.assertEqual(m["印发日期"], "2026年11月21日")

    def test_attachments_split(self):
        lines = self.res["lines"]
        self.assertIn("【附件1】窗口服务规范", lines)
        self.assertIn("【附件2】考核办法", lines)
        self.assertNotIn("附件1", lines)

    def test_text_unchanged(self):
        before = norm(BODY + SIGNOFF + ATTACH + TAIL + ["某区政务服务中心文件", "某政发〔2026〕5号", "签发人：张三"])
        after = para_texts(self.dst)
        # 重排后版头把文号和签发人排在同一行，印发行用制表符；比较时去掉空白并把两者拆开
        joined = "".join(re.sub(r"\s+", "", t) for t in after)
        for t in before:
            self.assertIn(t, joined, f"重排后缺少：{t}")

    def test_structure_fixed(self):
        r = check_docx.check_document(self.dst)
        ids = {i["id"] for i in r["issues"]}
        self.assertFalse(ids & {"DOCX-160", "DOCX-161", "DOCX-162", "DOCX-163", "DOCX-131"}, ids)
        self.assertEqual(r["summary"]["P0"], 0)

    def test_messy_flagged_before(self):
        r = check_docx.check_document(self.src)
        ids = {i["id"] for i in r["issues"]}
        self.assertTrue({"DOCX-160", "DOCX-161", "DOCX-162", "DOCX-163"} <= ids, ids)

    def test_attach_pages_and_banji(self):
        from docx.oxml.ns import qn
        from docx import Document
        doc = Document(str(self.dst))
        self.assertGreaterEqual(len(doc.tables), 1)
        xml = doc.element.xml
        self.assertIn('w:val="FF0000"', xml.upper().replace("W:VAL=\"FF0000\"", 'w:val="FF0000"'))
        self.assertGreaterEqual(xml.count('w:br w:type="page"') + xml.count("pageBreakBefore"), 2)

    def test_light_problems_not_structural(self):
        path = make_messy_docx(OUT / "extract-light.docx")
        r = check_docx.check_document(path)
        ids = {i["id"] for i in r["issues"]}
        self.assertFalse(ids & {"DOCX-160", "DOCX-161", "DOCX-162", "DOCX-163"}, ids)


class ExtractCliTest(unittest.TestCase):
    def test_cli_writes_files(self):
        src = make_messy_header_docx(OUT / "extract-cli.docx")
        for f in OUT.glob("正文-extract-cli*"):
            f.unlink()
        for f in OUT.glob("版式信息-extract-cli*"):
            f.unlink()
        self.assertEqual(extract_docx.main([str(src)]), 0)
        self.assertTrue((OUT / "正文-extract-cli.txt").is_file())
        meta = json.loads((OUT / "版式信息-extract-cli.json").read_text(encoding="utf-8"))
        self.assertEqual(meta["版式"], "红头")

    def test_force_layout_missing_issuer(self):
        res = extract_docx.extract(make_messy_docx(OUT / "extract-light2.docx"), "红头")
        self.assertEqual(res["meta"]["版式"], "红头")
        self.assertNotIn("发文机关标志", res["meta"])
        doc, rep = build_docx.build(res["lines"], res["meta"], keep_text=True, installed=set())
        self.assertTrue(any("发文机关标志" in p for p in rep["placeholders"]))


if __name__ == "__main__":
    unittest.main()
