# -*- coding: utf-8 -*-
"""fix_docx 修复格式混乱的 docx；annotate_docx 写入批注。"""
import hashlib
import io
import json
import unittest
from contextlib import redirect_stdout

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Mm

from tests.format.helpers import NO_FONTS, OUT, build_sample, check_docx, gw_docx, make_messy_docx

import annotate_docx  # noqa: E402  helpers 已把 scripts 加入 sys.path
import fix_docx  # noqa: E402


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def clean(path):
    for old in OUT.glob(path.stem + "*"):
        old.unlink()


class FixTest(unittest.TestCase):
    def test_messy_docx_fixed(self):
        src = OUT / "混乱稿.docx"
        clean(src)
        make_messy_docx(src)
        before = check_docx.check_document(src)
        self.assertGreater(before["summary"]["P0"], 0)
        original = digest(src)
        dst = OUT / "混乱稿-修复.docx"
        report = fix_docx.fix_document(src, dst, installed=NO_FONTS)
        self.assertEqual(digest(src), original)  # 原文件不动
        after = check_docx.check_document(dst, layout=report["layout"])
        self.assertEqual(after["summary"]["P0"], 0, after["issues"])
        self.assertEqual(after["summary"]["P1"], 0, after["issues"])
        doc = Document(str(dst))
        sec = doc.sections[0]
        self.assertAlmostEqual(sec.top_margin / Mm(1), 37, delta=0.1)
        self.assertAlmostEqual(sec.page_width / Mm(1), 210, delta=0.1)
        title = doc.paragraphs[0]
        self.assertEqual(title.alignment, 1)
        self.assertEqual(title.runs[0]._r.rPr.find(qn("w:rFonts")).get(qn("w:eastAsia")), "方正小标宋简体")
        # 文字未改动，只删了段首空格
        old_text = [p.text.strip("　 ") for p in Document(str(src)).paragraphs if p.text.strip()]
        new_text = [p.text for p in doc.paragraphs if p.text.strip()]
        self.assertEqual(old_text, new_text)
        self.assertGreaterEqual(report["stripped"], 3)
        self.assertGreater(report["compressed"], 0)
        self.assertIn("页码", "".join(report["notes"]))

    def test_cli_writes_report_and_keeps_source(self):
        src = OUT / "混乱稿CLI.docx"
        clean(src)
        make_messy_docx(src)
        buf = io.StringIO()
        with redirect_stdout(buf):
            code = fix_docx.main([str(src), "--json"])
        self.assertEqual(code, 0)
        data = json.loads(buf.getvalue())
        self.assertTrue(data["output"].endswith("混乱稿CLI-修复.docx"))
        report_text = (OUT / "混乱稿CLI-修复-修复报告.md").read_text(encoding="utf-8")
        self.assertIn("修复后检查", report_text)
        with redirect_stdout(io.StringIO()):
            fix_docx.main([str(src)])
        self.assertTrue((OUT / "混乱稿CLI-修复-v2.docx").exists())  # 不覆盖上次的输出

    def test_fix_compliant_sample_stays_compliant(self):
        """合规样稿再修一遍：仍然没有 P0，文字不变（标本测试法的“幂等”不变式）。"""
        src, _ = build_sample("红头", "幂等-红头.docx")
        dst = OUT / "幂等-红头-修复.docx"
        clean(dst)
        report = fix_docx.fix_document(src, dst, installed=NO_FONTS)
        self.assertEqual(report["layout"], "红头")
        self.assertEqual(report["page"], [])  # 页面设置无变化
        after = check_docx.check_document(dst, layout="红头")
        self.assertEqual(after["summary"]["P0"], 0, after["issues"])
        self.assertEqual([p.text for p in Document(str(src)).paragraphs], [p.text for p in Document(str(dst)).paragraphs])
        body = report["roles"].get("正文", {})
        self.assertEqual(body.get("字体", 0), 0)
        self.assertEqual(body.get("字号", 0), 0)


    def test_minutes_without_number_keeps_red_issuer(self):
        """纪要没有期号时，红色标志也要认作发文机关标志，修复时不改其字号和颜色。"""
        src, _ = build_sample("纪要", "纪要无期号.docx", meta_update={"发文字号": ""})
        before = check_docx.check_document(src)
        self.assertEqual(before["layout"], "纪要")
        self.assertEqual(before["summary"]["P0"], 0, before["issues"])
        self.assertFalse(any(i["id"] in ("DOCX-020", "DOCX-021") for i in before["issues"]))
        dst = OUT / "纪要无期号-修复.docx"
        clean(dst)
        fix_docx.fix_document(src, dst, installed=NO_FONTS)
        def issuer_run(path):
            return [p for p in Document(str(path)).paragraphs if p.text == "某理工大学校长办公会纪要"][0].runs[0]
        self.assertEqual(str(issuer_run(dst).font.color.rgb), "FF0000")
        self.assertEqual(issuer_run(dst).font.size, issuer_run(src).font.size)


    def test_upward_docno_line_kept(self):
        """上行文发文字号行：修复后发文字号与签发人仍在同一行、左缩进一字，签发人姓名仍为楷体。"""
        from tests.format.helpers import FILES, build_text
        text = (FILES / "高校-请示-正文.txt").read_text(encoding="utf-8")
        src, _ = build_text(text, "上行文.docx", meta={"版式": "红头", "发文机关标志": "某理工大学文件",
                                                     "发文字号": "某理工发〔2026〕52号", "签发人": "王某",
                                                     "印发机关": "某理工大学党政办公室", "印发日期": "2026-10-26"})
        dst = OUT / "上行文-修复.docx"
        clean(dst)
        fix_docx.fix_document(src, dst, installed=NO_FONTS)
        line = [p for p in Document(str(dst)).paragraphs if "签发人" in p.text][0]
        self.assertEqual(line.text, "某理工发〔2026〕52号	签发人：王某")
        self.assertEqual(gw_docx.para_left_chars(line, Document(str(dst))), 1.0)
        name = line.runs[-1]
        self.assertEqual(name.text, "王某")
        self.assertEqual(name._r.rPr.find(qn("w:rFonts")).get(qn("w:eastAsia")), "楷体_GB2312")


class FixRobustnessTest(unittest.TestCase):
    def test_draft_note_left_untouched(self):
        """【起草说明】之后不排版：署名、日期按落款排，起草说明各段保持原样并在报告中提示。

        原问题：署名、日期被当成正文加首行缩进 2 字，起草说明也按正文、三级标题改了格式。
        """
        from docx import Document as NewDocument
        src = OUT / "含起草说明.docx"
        clean(src)
        doc = NewDocument()
        for text in ("关于某事的通知", "", "各单位：", "", "现通知如下。", "", "某局", "", "2026年3月5日", "【起草说明】",
                     "1. 所作假设：检查次数为【待补：次数】。"):
            doc.add_paragraph(text)
        doc.save(str(src))
        dst = OUT / "含起草说明-修复.docx"
        report = fix_docx.fix_document(src, dst, installed=NO_FONTS)
        self.assertIn("成文日期：2026年3月5日", report["summary"])
        self.assertTrue(any("起草说明" in n for n in report["untouched"]), report["untouched"])
        out = Document(str(dst))
        texts = [p.text for p in out.paragraphs]
        self.assertEqual(texts, [p.text for p in Document(str(src)).paragraphs])  # 文字未改动
        date = out.paragraphs[texts.index("2026年3月5日")]
        self.assertNotEqual(date._p.pPr.find(qn("w:ind")).get(qn("w:firstLineChars")), "200")
        for p in out.paragraphs[texts.index("【起草说明】"):]:
            self.assertIsNone(p._p.pPr, p.text)  # 段落格式未动
            self.assertTrue(all(r._r.rPr is None for r in p.runs), p.text)  # 文字格式未动

    def test_report_into_missing_folder(self):
        import shutil
        folder = OUT / "新建的修复报告目录"
        if folder.exists():
            shutil.rmtree(folder)
        src = OUT / "报告目录.docx"
        clean(src)
        make_messy_docx(src)
        with redirect_stdout(io.StringIO()):
            code = fix_docx.main([str(src), "--report", str(folder / "修复报告.md")])
        self.assertEqual(code, 0)
        self.assertTrue((folder / "修复报告.md").exists())

    def test_report_sections_numbered_consecutively(self):
        """修复报告中没有“未处理和提示”部分时，“修复后检查”不能从“四、”跳到“六、”。"""
        import re
        report = {"input": "a.docx", "output": "a-修复.docx", "layout": "白头", "seal": True, "summary": [],
                  "page": [], "roles": {}, "notes": ["页码已重设"], "untouched": [], "warnings": [],
                  "font_warnings": [], "check": {"summary": {"P0": 0, "P1": 0, "P2": 0}, "issues": []}}
        heads = re.findall(r"^## ([一二三四五六七八九十]+)、", fix_docx.render_report(report), flags=re.M)
        self.assertEqual(heads, ["一", "二", "三", "四", "五"])
        report["notes"] = []
        report["warnings"] = ["提示"]
        heads = re.findall(r"^## ([一二三四五六七八九十]+)、", fix_docx.render_report(report), flags=re.M)
        self.assertEqual(heads, ["一", "二", "三", "四", "五"])

    def test_cli_unknown_layout(self):
        """--layout 写成不认识的值时报输入错误；原来按未知版式静默修复，红头文件的署名被按“不加盖印章”排列。"""
        src = OUT / "未知版式.docx"
        clean(src)
        make_messy_docx(src)
        buf = io.StringIO()
        with redirect_stdout(buf):
            code = fix_docx.main([str(src), "--layout", "红头文件"])
        self.assertEqual(code, 2)
        self.assertIn("无法识别的版式", buf.getvalue())
        self.assertFalse((OUT / "未知版式-修复.docx").exists())

    def test_output_to_existing_directory(self):
        """-o 给的是已存在的文件夹时，在文件夹内输出“原名-修复.docx”，不在旁边生成无扩展名的文件。"""
        import shutil
        folder = OUT / "修复输出目录"
        if folder.exists():
            shutil.rmtree(folder)
        for old in OUT.glob("修复输出目录-v*"):
            old.unlink()
        folder.mkdir()
        src = OUT / "修复到目录.docx"
        clean(src)
        make_messy_docx(src)
        buf = io.StringIO()
        with redirect_stdout(buf):
            code = fix_docx.main([str(src), "-o", str(folder), "--json"])
        self.assertEqual(code, 0)
        self.assertEqual(json.loads(buf.getvalue())["output"], str(folder / "修复到目录-修复.docx"))
        self.assertFalse((OUT / "修复输出目录-v2").exists())


class AnnotateTest(unittest.TestCase):
    def test_output_to_existing_directory(self):
        import shutil
        folder = OUT / "批注输出目录"
        if folder.exists():
            shutil.rmtree(folder)
        for old in OUT.glob("批注输出目录-v*"):
            old.unlink()
        folder.mkdir()
        src, _ = build_sample("白头", "批注到目录.docx")
        jpath = OUT / "核稿意见-到目录.json"
        jpath.write_text(json.dumps({"issues": [{"excerpt": "妥否", "message": "x"}]}, ensure_ascii=False),
                         encoding="utf-8")
        with redirect_stdout(io.StringIO()):
            code = annotate_docx.main([str(src), str(jpath), "-o", str(folder)])
        self.assertEqual(code, 0)
        self.assertTrue((folder / "批注到目录-批注.docx").exists())
        self.assertFalse((OUT / "批注输出目录-v2").exists())

    def test_bad_docx_exit_2(self):
        bad = OUT / "损坏文件.docx"
        bad.write_bytes(b"not a zip file")
        jpath = OUT / "核稿意见-损坏.json"
        jpath.write_text(json.dumps({"issues": [{"excerpt": "某", "message": "x"}]}, ensure_ascii=False), encoding="utf-8")
        buf = io.StringIO()
        with redirect_stdout(buf):
            code = annotate_docx.main([str(bad), str(jpath)])
        self.assertEqual(code, 2)
        self.assertIn("错误", buf.getvalue())

    def test_report_into_missing_folder(self):
        src, _ = build_sample("白头", "批注报告目录.docx")
        jpath = OUT / "核稿意见-报告目录.json"
        jpath.write_text(json.dumps({"issues": [{"excerpt": "妥否", "message": "x"}]}, ensure_ascii=False),
                         encoding="utf-8")
        folder = OUT / "批注报告-不存在"
        for old in folder.glob("*") if folder.exists() else []:
            old.unlink()
        if folder.exists():
            folder.rmdir()
        clean(OUT / "批注报告目录-批注.docx")
        with redirect_stdout(io.StringIO()):
            code = annotate_docx.main([str(src), str(jpath), "--report", str(folder / "批注报告.md")])
        self.assertEqual(code, 0)
        self.assertTrue((folder / "批注报告.md").exists())

    def test_comments_written_and_unlocated_reported(self):
        src, _ = build_sample("白头", "待批注.docx")
        issues = {"source_file": "正文.txt", "issues": [
            {"level": "P1", "gate": "文字关", "category": "标点", "excerpt": "妥否，请批示。", "message": "结尾用语规范",
             "suggestion": "保留", "basis": "条例第十五条"},
            {"level": "P0", "gate": "政策关", "category": "数据", "excerpt": "尚缺46万元", "message": "核对缺口金额",
             "suggestion": "与预算表一致", "basis": ""},
            {"level": "P2", "gate": "格式关", "category": "格式", "excerpt": "县人民政府 ：", "message": "空格容错匹配"},
            {"level": "P1", "gate": "文字关", "category": "错字", "excerpt": "文中没有这句话", "message": "定位不到"},
            {"level": "P1", "gate": "文字关", "category": "错字", "excerpt": "", "message": "空片段"},
        ]}
        jpath = OUT / "核稿意见-测试.json"
        jpath.write_text(json.dumps(issues, ensure_ascii=False), encoding="utf-8")
        dst = OUT / "待批注-批注.docx"
        clean(dst)
        original = digest(src)
        report = annotate_docx.annotate(src, annotate_docx.load_issues(jpath), dst)
        self.assertEqual(digest(src), original)
        self.assertEqual(len(report["located"]), 3)
        self.assertEqual(len(report["unlocated"]), 2)
        doc = Document(str(dst))
        comments = list(doc.comments)
        self.assertEqual(len(comments), 3)
        self.assertEqual({c.author for c in comments}, {"gongwen-review"})
        self.assertIn("【P1｜文字关｜标点】结尾用语规范", comments[0].text)
        self.assertIn("建议：保留", comments[0].text)
        # 批注范围正好覆盖 excerpt
        body = doc.element.body
        starts = body.findall(".//" + qn("w:commentRangeStart"))
        self.assertEqual(len(starts), 3)
        # 文字未改动
        self.assertEqual([p.text for p in Document(str(src)).paragraphs], [p.text for p in doc.paragraphs])

    def test_duplicate_excerpt_goes_to_next_occurrence(self):
        src, _ = build_sample("红头", "重复片段.docx")
        issues = [{"level": "P2", "excerpt": "重点检查", "message": "第一处"},
                  {"level": "P2", "excerpt": "重点检查", "message": "第二处"}]
        dst = OUT / "重复片段-批注.docx"
        clean(dst)
        report = annotate_docx.annotate(src, issues, dst)
        self.assertEqual(len(report["located"]), 2)
        self.assertIn("note", report["located"][0])

    def test_cli_and_report(self):
        src, _ = build_sample("白头", "批注CLI.docx")
        jpath = OUT / "核稿意见-CLI.json"
        jpath.write_text(json.dumps({"issues": [{"level": "P1", "excerpt": "不存在的片段", "message": "x"}]},
                                    ensure_ascii=False), encoding="utf-8")
        clean(OUT / "批注CLI-批注.docx")
        with redirect_stdout(io.StringIO()):
            code = annotate_docx.main([str(src), str(jpath)])
        self.assertEqual(code, 0)
        text = (OUT / "批注CLI-批注-批注报告.md").read_text(encoding="utf-8")
        self.assertIn("未定位的意见", text)
        self.assertIn("不存在的片段", text)

    def test_version_check(self):
        self.assertTrue(annotate_docx.supports_comments())
        self.assertGreaterEqual(annotate_docx.docx_version(), (1, 2, 0))

    def test_old_python_docx_prompts_upgrade(self):
        from unittest import mock
        buf = io.StringIO()
        with mock.patch.object(annotate_docx, "supports_comments", return_value=False), redirect_stdout(buf):
            code = annotate_docx.main(["任意.docx", "任意.json"])
        self.assertEqual(code, 2)
        self.assertIn("pip install -U python-docx", buf.getvalue())


if __name__ == "__main__":
    unittest.main()
