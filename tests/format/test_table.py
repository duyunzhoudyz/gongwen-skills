# -*- coding: utf-8 -*-
"""Markdown 管道表格转 Word 表格：识别、生成、回读检查、修复时不误认表题和署名。"""
import io
import json
import unittest
from contextlib import redirect_stdout

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

from tests.format.helpers import OUT, build_docx, build_text, check_docx, gw_docx, gw_layout

TABLE_TEXT = ("关于报送第三季度设备巡检情况的通知\n\n各分厂：\n\n现将第三季度设备巡检情况通报如下。\n\n"
              "表1　第三季度设备巡检情况表\n"
              "| 分厂 | 巡检台次 | 发现隐患（处） | 备注 |\n"
              "|---|---:|:---:|---|\n"
              "| 一分厂 | 46 | 3 | **已整改** |\n"
              "| 二分厂 | 39 | 【待补：隐患处数】 | 设备更新中 |\n\n"
              "请各分厂于下月5日前报送整改情况。\n\n某机械厂办公室\n\n2026年9月30日")


def body_tables(doc):
    return [t for t in doc.tables if t._tbl.tblPr.find(qn("w:tblpPr")) is None]


def ids(result):
    return {(i["id"], i["level"]) for i in result["issues"]}


class ExtractTest(unittest.TestCase):
    def test_extract_basic(self):
        lines, tables, notes = gw_layout.extract_pipe_tables(TABLE_TEXT.split("\n"))
        self.assertEqual(len(tables), 1)
        self.assertEqual(lines.count(gw_layout.TABLE_MARK), 1)
        tb = tables[0]
        self.assertEqual(tb.header, ["分厂", "巡检台次", "发现隐患（处）", "备注"])
        self.assertEqual(tb.aligns, [None, "right", "center", None])
        self.assertEqual(tb.rows[0], ["一分厂", "46", "3", "已整改"])  # 去掉 ** 加粗记号，文字不变
        self.assertEqual(notes, [])

    def test_ragged_rows_kept(self):
        text = "| 项目 | 数量 |\n|---|---|\n| 甲 | 1 | 多出的一格 |\n| 乙 |"
        _lines, tables, notes = gw_layout.extract_pipe_tables(text.split("\n"))
        tb = tables[0]
        self.assertEqual(len(tb.header), 3)
        self.assertEqual(tb.rows[0][2], "多出的一格")  # 不丢文字
        self.assertEqual(tb.rows[1], ["乙", "", ""])
        self.assertTrue(any("格数不一致" in n for n in notes))

    def test_no_delimiter_not_converted(self):
        text = "| 项目 | 数量 |\n| 甲 | 1 |"
        lines, tables, notes = gw_layout.extract_pipe_tables(text.split("\n"))
        self.assertEqual(tables, [])
        self.assertEqual(lines, text.split("\n"))
        self.assertTrue(any("缺少分隔行" in n for n in notes))

    def test_escaped_pipe_and_br(self):
        self.assertEqual(gw_layout.split_pipe_row(r"| a\|b | c |"), ["a|b", "c"])
        self.assertEqual(gw_layout.clean_cell("第一行<br>第二行"), "第一行\n第二行")

    def test_caption_rule(self):
        for ok in ("表1　第三季度设备巡检情况表", "附表2 人员名单", "各科室门诊量统计表（单位：人次）", "督办事项清单"):
            self.assertTrue(gw_layout.is_table_caption(ok), ok)
        for bad in ("现将情况统计如下。", "一、基本情况", "具体见下表：请核对。"):
            self.assertFalse(gw_layout.is_table_caption(bad), bad)


class BuildTableTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.path, cls.report = build_text(TABLE_TEXT, "表格-通知.docx", meta={"版式": "白头"})
        cls.doc = Document(str(cls.path))

    def test_one_table_and_no_pipe_paragraphs(self):
        self.assertEqual(len(body_tables(self.doc)), 1)
        self.assertFalse(any("|" in p.text for p in self.doc.paragraphs))
        self.assertTrue(any("Markdown 表格已转为 Word 表格" in n for n in self.report["notes"]))

    def test_table_style(self):
        tb = body_tables(self.doc)[0]
        tblpr = tb._tbl.tblPr
        self.assertEqual(tblpr.find(qn("w:jc")).get(qn("w:val")), "center")
        borders = tblpr.find(qn("w:tblBorders"))
        for tag in ("top", "left", "bottom", "right", "insideH", "insideV"):
            self.assertEqual(borders.find(qn(f"w:{tag}")).get(qn("w:val")), "single", tag)
        grid = [int(c.get(qn("w:w"))) for c in tb._tbl.tblGrid.findall(qn("w:gridCol"))]
        width = gw_docx.Ctx(gw_docx.load_params(), {}).text_width_twips
        self.assertEqual(sum(grid), width)  # 与版心同宽，不超出
        header = tb.rows[0]
        self.assertIsNotNone(header._tr.trPr.find(qn("w:tblHeader")))  # 跨页重复表头
        for cell in header.cells:
            run = cell.paragraphs[0].runs[0]
            self.assertTrue(run.bold, cell.text)
            self.assertEqual(run.font.size, Pt(14))
            self.assertEqual(run._r.rPr.find(qn("w:rFonts")).get(qn("w:eastAsia")), "仿宋_GB2312")
            self.assertEqual(cell.paragraphs[0]._p.pPr.find(qn("w:jc")).get(qn("w:val")), "center")
        data = tb.rows[1].cells
        self.assertFalse(data[0].paragraphs[0].runs[0].bold)
        self.assertEqual(data[0].paragraphs[0]._p.pPr.find(qn("w:jc")).get(qn("w:val")), "left")
        self.assertEqual(data[1].paragraphs[0]._p.pPr.find(qn("w:jc")).get(qn("w:val")), "right")  # 分隔行 ---:
        self.assertEqual(data[3].text, "已整改")
        # 单元格不对齐文档网格，避免行高被拉成 28.95 磅
        snap = data[0].paragraphs[0]._p.pPr.find(qn("w:snapToGrid"))
        self.assertEqual(snap.get(qn("w:val")), "0")

    def test_caption_and_position(self):
        body = self.doc.element.body
        children = [c for c in body.iterchildren()]
        tbl = body_tables(self.doc)[0]._tbl
        k = children.index(tbl)
        caption = children[k - 1]
        self.assertEqual("".join(t.text for t in caption.iter(qn("w:t"))), "表1　第三季度设备巡检情况表")
        cap_p = [p for p in self.doc.paragraphs if p._p is caption][0]
        self.assertEqual(cap_p._p.pPr.find(qn("w:jc")).get(qn("w:val")), "center")
        self.assertIsNotNone(cap_p._p.pPr.find(qn("w:keepNext")))
        self.assertEqual(cap_p.runs[0]._r.rPr.find(qn("w:rFonts")).get(qn("w:eastAsia")), "黑体")
        after = [p for p in self.doc.paragraphs if p._p is children[k + 1]][0]
        self.assertTrue(after.text.startswith("请各分厂"))  # 表格后不插多余空段

    def test_check_passes_and_signoff_recognized(self):
        result = check_docx.check_document(self.path)
        self.assertEqual(result["summary"]["P0"], 0, result["issues"])
        self.assertEqual(result["summary"]["P1"], 0, result["issues"])
        self.assertEqual(result["doc_type"], "通知")
        self.assertEqual(result["summary"]["markers"]["待补"], 1)  # 表格中的【待补】也计数
        fields = self.report["fields"]
        self.assertEqual(fields["发文机关署名"], ["某机械厂办公室"])
        self.assertTrue(any("表格 1 个" in s for s in self.report["summary"]))

    def test_table_last_before_signature(self):
        """表格紧挨落款时，表题不能被认成署名。"""
        text = TABLE_TEXT.replace("请各分厂于下月5日前报送整改情况。\n\n", "")
        path, report = build_text(text, "表格-紧挨落款.docx", meta={"版式": "白头"})
        self.assertEqual(report["fields"]["发文机关署名"], ["某机械厂办公室"])
        result = check_docx.check_document(path)
        self.assertEqual(result["summary"]["P0"], 0, result["issues"])
        self.assertEqual(result["summary"]["P1"], 0, result["issues"])

    def test_two_tables_kept_apart(self):
        text = ("某车间工作总结\n\n| 甲 | 乙 |\n|---|---|\n| 1 | 2 |\n\n| 丙 | 丁 |\n|---|---|\n| 3 | 4 |\n\n"
                "以上为本季度数据。")
        path, _ = build_text(text, "表格-两张.docx")
        doc = Document(str(path))
        self.assertEqual(len(body_tables(doc)), 2)
        children = list(doc.element.body.iterchildren())
        k = children.index(body_tables(doc)[0]._tbl)
        self.assertTrue(children[k + 1].tag.endswith("}p"))  # 两表之间隔一个段落，Word 不会合并

    def test_table_at_end_gets_anchor(self):
        path, _ = build_text("某车间工作总结\n\n| 甲 | 乙 |\n|---|---|\n| 1 | 2 |", "表格-结尾.docx")
        doc = Document(str(path))
        children = [c for c in doc.element.body.iterchildren() if not c.tag.endswith("}sectPr")]
        self.assertTrue(children[-1].tag.endswith("}p"))

    def test_table_font_options(self):
        path, report = build_text(TABLE_TEXT, "表格-宋体小四.docx", meta={"版式": "白头", "表格字体": "宋体", "表格字号": "小四"})
        run = body_tables(Document(str(path)))[0].rows[1].cells[0].paragraphs[0].runs[0]
        self.assertEqual(run.font.size, Pt(12))
        self.assertEqual(run._r.rPr.find(qn("w:rFonts")).get(qn("w:eastAsia")), "宋体")
        _p, bad = build_text(TABLE_TEXT, "表格-字号不认.docx", meta={"版式": "白头", "表格字号": "五号"})
        self.assertTrue(any("表格字号" in w for w in bad["warnings"]))

    def test_cli_md_input(self):
        folder = OUT / "表格命令行"
        folder.mkdir(parents=True, exist_ok=True)
        for old in folder.glob("*"):
            old.unlink()
        src = folder / "正文.md"
        src.write_text(TABLE_TEXT, encoding="utf-8")
        buf = io.StringIO()
        with redirect_stdout(buf):
            self.assertEqual(build_docx.main([str(src), "--json"]), 0)
        report = json.loads(buf.getvalue())
        self.assertEqual(report["check"]["summary"]["P1"], 0, report["check"]["issues"])
        self.assertEqual(len(body_tables(Document(report["output"]))), 1)


class CheckTableTest(unittest.TestCase):
    def test_residual_pipe_rows_p1(self):
        text = "关于某事的通知\n\n各单位：\n\n| 项目 | 数量 |\n\n|---|---|\n\n| 甲 | 1 |\n\n某市某局\n\n2026年3月5日"
        path, report = build_text(text, "标本-残留表格行.docx", meta={"版式": "白头"})
        result = check_docx.check_document(path)
        hits = [i for i in result["issues"] if i["id"] == "DOCX-100"]
        self.assertEqual(len(hits), 1)
        self.assertEqual(hits[0]["level"], "P1")
        self.assertIn("共 3 段", hits[0]["message"])

    def _mutate(self, name, fn):
        path, _ = build_text(TABLE_TEXT, "标本-表格基础.docx", meta={"版式": "白头"})
        doc = Document(str(path))
        fn(body_tables(doc)[0])
        out = OUT / name
        doc.save(str(out))
        return check_docx.check_document(out)

    def test_table_too_wide(self):
        def fn(tb):
            for col in tb._tbl.tblGrid.findall(qn("w:gridCol")):
                col.set(qn("w:w"), "3000")
        self.assertIn(("DOCX-101", "P1"), ids(self._mutate("标本-表格过宽.docx", fn)))

    def test_table_font_hint(self):
        def fn(tb):
            for row in tb.rows:
                for cell in row.cells:
                    for r in cell.paragraphs[0].runs:
                        r._r.rPr.find(qn("w:rFonts")).set(qn("w:eastAsia"), "微软雅黑")
                        r.font.size = Pt(10.5)
        result = self._mutate("标本-表格字体.docx", fn)
        self.assertIn(("DOCX-102", "P2"), ids(result))
        self.assertEqual(result["summary"]["P0"], 0)


class FixTableTest(unittest.TestCase):
    def test_fix_keeps_caption_and_signature(self):
        import fix_docx
        src, _ = build_text(TABLE_TEXT.replace("请各分厂于下月5日前报送整改情况。\n\n", ""), "表格-待修复.docx",
                            meta={"版式": "白头"})
        dst = OUT / "表格-待修复-修复.docx"
        if dst.exists():
            dst.unlink()
        report = fix_docx.fix_document(src, dst, installed=set())
        self.assertEqual(report["roles"].get("表题", {}).get("段数"), 1)
        self.assertEqual(report["roles"]["发文机关署名"]["段数"], 1)
        result = check_docx.check_document(dst)
        self.assertEqual(result["summary"]["P0"], 0, result["issues"])
        self.assertEqual(result["summary"]["P1"], 0, result["issues"])
        doc = Document(str(dst))
        cap = [p for p in doc.paragraphs if p.text.startswith("表1")][0]
        self.assertEqual(cap._p.pPr.find(qn("w:jc")).get(qn("w:val")), "center")


if __name__ == "__main__":
    unittest.main()
