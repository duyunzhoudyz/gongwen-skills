# -*- coding: utf-8 -*-
"""build_docx：五种版式样稿生成、回读关键参数、check 无 P0、字体与文件保护。"""
import inspect
import io
import json
import unittest
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Mm

from tests.format.helpers import (FILES, FULL_FONTS, NO_FONTS, OUT, STANDARD_FONTS, WIN_FONTS, build_docx,
                                  build_sample, build_text, check_docx, fonts_installed, gw_docx, written_fonts)

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def run_fonts(paragraph):
    out = []
    for r in paragraph.runs:
        rf = r._r.rPr.find(qn("w:rFonts")) if r._r.rPr is not None else None
        out.append(rf.get(qn("w:eastAsia")) if rf is not None else None)
    return out


def run_sizes(paragraph):
    return [r.font.size.pt if r.font.size else None for r in paragraph.runs]


def find_para(doc, startswith):
    for p in doc.paragraphs:
        if p.text.strip().startswith(startswith):
            return p
    raise AssertionError(f"找不到段落：{startswith}")


def footer_has_page(footer) -> bool:
    return any("PAGE" in (t.text or "") for t in footer._element.iter(W + "instrText"))


class FiveLayoutsTest(unittest.TestCase):
    """每种版式生成一份样稿，check 无 P0，并回读关键参数。"""

    @classmethod
    def setUpClass(cls):
        cls.docs = {}
        for layout, name in (("红头", "样稿-红头通知.docx"), ("信函", "样稿-信函.docx"), ("纪要", "样稿-纪要.docx"),
                             ("白头", "样稿-白头请示.docx"), ("材料", "样稿-年终总结.docx")):
            path, report = build_sample(layout, name)
            cls.docs[layout] = (path, report, Document(str(path)))

    def test_check_no_p0(self):
        for layout, (path, _report, _doc) in self.docs.items():
            result = check_docx.check_document(path, layout=layout)
            self.assertEqual(result["summary"]["P0"], 0, f"{layout}: {result['issues']}")
            self.assertEqual(result["layout"], layout)

    def test_page_setup(self):
        for layout, (_p, _r, doc) in self.docs.items():
            sec = doc.sections[0]
            self.assertEqual(int(sec._sectPr.find(qn("w:pgSz")).get(qn("w:w"))), 11906, layout)
            self.assertEqual(int(sec._sectPr.find(qn("w:pgSz")).get(qn("w:h"))), 16838, layout)
            for attr, mm in (("top_margin", 37), ("bottom_margin", 35), ("left_margin", 28), ("right_margin", 26)):
                self.assertAlmostEqual(getattr(sec, attr) / Mm(1), mm, delta=0.1, msg=f"{layout} {attr}")
            grid = sec._sectPr.find(qn("w:docGrid"))
            self.assertEqual(grid.get(qn("w:type")), "linesAndChars")
            self.assertEqual(grid.get(qn("w:linePitch")), "579")
            self.assertEqual(grid.get(qn("w:charSpace")), "-842")

    def test_body_and_title_fonts(self):
        for layout, (_p, report, doc) in self.docs.items():
            fields = report["fields"]
            if layout != "纪要" or fields.get("标题"):
                title = find_para(doc, str(fields["标题"])[:6])
                self.assertEqual(set(run_fonts(title)), {"方正小标宋简体"}, layout)
                self.assertEqual(set(run_sizes(title)), {22.0}, layout)
                self.assertEqual(title.alignment, 1, layout)  # 居中
            body = [p for p in doc.paragraphs if p.text and p.paragraph_format.first_line_indent]
            self.assertTrue(body, layout)
            p = body[0]
            self.assertIn("仿宋_GB2312", run_fonts(p) + run_fonts(body[-1]), layout)
            ind = p._p.pPr.find(qn("w:ind"))
            self.assertEqual(ind.get(qn("w:firstLineChars")), "200", layout)
            sp = p._p.pPr.find(qn("w:spacing"))
            self.assertEqual((sp.get(qn("w:line")), sp.get(qn("w:lineRule"))), ("578", "exact"), layout)
            # eastAsia 属性必须写入
            rf = p.runs[0]._r.rPr.find(qn("w:rFonts"))
            self.assertIsNotNone(rf.get(qn("w:eastAsia")))
            self.assertIsNone(rf.get(qn("w:eastAsiaTheme")))

    def test_heading_fonts(self):
        doc = self.docs["红头"][2]
        self.assertEqual(run_fonts(find_para(doc, "一、检查范围")), ["黑体"])
        h2 = find_para(doc, "（一）安全责任落实情况")
        fonts = run_fonts(h2)
        self.assertEqual(fonts[0], "楷体_GB2312")  # 标题部分
        self.assertEqual(fonts[-1], "仿宋_GB2312")  # 同段正文部分
        self.assertEqual(h2.runs[0].text, "（一）安全责任落实情况。")
        self.assertEqual(set(run_fonts(find_para(doc, "1.危险化学品企业"))), {"仿宋_GB2312"})
        recipient = find_para(doc, "各区应急管理局")
        self.assertFalse(recipient.paragraph_format.first_line_indent)

    def test_page_numbers(self):
        for layout, (_p, _r, doc) in self.docs.items():
            sec = doc.sections[0]
            self.assertTrue(doc.settings.odd_and_even_pages_header_footer, layout)
            self.assertTrue(footer_has_page(sec.footer), layout)
            self.assertTrue(footer_has_page(sec.even_page_footer), layout)
            odd = sec.footer.paragraphs[0]
            even = sec.even_page_footer.paragraphs[0]
            self.assertEqual(odd._p.pPr.find(qn("w:jc")).get(qn("w:val")), "right")
            self.assertEqual(even._p.pPr.find(qn("w:jc")).get(qn("w:val")), "left")
            self.assertEqual(odd._p.pPr.find(qn("w:ind")).get(qn("w:right")), "280")  # 空一字（4 号字 14 磅）
            self.assertEqual(even._p.pPr.find(qn("w:ind")).get(qn("w:left")), "280")
            text = "".join(t.text for t in odd._p.iter(W + "t"))
            self.assertEqual(text, "— 1 —")
            self.assertEqual(set(run_sizes(odd)), {14.0})
            self.assertEqual(set(f for f in run_fonts(odd)), {"宋体"})

    def test_red_head_elements(self):
        doc = self.docs["红头"][2]
        issuer = find_para(doc, "某市应急管理局文件")
        self.assertEqual(str(issuer.runs[0].font.color.rgb), "FF0000")
        self.assertEqual(run_fonts(issuer), ["方正小标宋简体"])
        docno = find_para(doc, "某应急发〔2026〕21号")
        bottom = docno._p.pPr.find(qn("w:pBdr")).find(qn("w:bottom"))
        self.assertEqual(bottom.get(qn("w:color")), "FF0000")
        self.assertEqual(bottom.get(qn("w:space")), "11")  # 发文字号之下 4 mm
        # 版记：浮动表格贴版心下边缘，首末粗线、中间细线
        tbl = doc.tables[-1]._tbl
        tblp = tbl.tblPr.find(qn("w:tblpPr"))
        self.assertEqual(tblp.get(qn("w:tblpYSpec")), "bottom")
        borders = tbl.tblPr.find(qn("w:tblBorders"))
        self.assertEqual(borders.find(qn("w:top")).get(qn("w:sz")), "8")
        self.assertEqual(borders.find(qn("w:insideH")).get(qn("w:sz")), "6")
        cells = [c.text for row in doc.tables[-1].rows for c in row.cells]
        self.assertEqual(cells[0], "抄送：市安全生产委员会办公室。")
        self.assertEqual(cells[1], "某市应急管理局办公室\t2026年10月28日印发")
        # 加盖印章：署名和日期两段居中、左右缩进相同（中心点一致），中心点 = 28 − 4 − 日期 7 字 ÷ 2 = 20.5，
        # 居中区间取 [13, 28]，即左缩进 13 字、右缩进 0
        date = find_para(doc, "2026年10月28日")
        sig = [p for p in doc.paragraphs if p.text == "某市应急管理局"][0]
        for p in (date, sig):
            self.assertEqual(p._p.pPr.find(qn("w:jc")).get(qn("w:val")), "center")
            ind = p._p.pPr.find(qn("w:ind"))
            self.assertEqual((ind.get(qn("w:leftChars")), ind.get(qn("w:rightChars"))), ("1300", "0"))

    def test_letter_elements(self):
        doc = self.docs["信函"][2]
        sec = doc.sections[0]
        self.assertTrue(sec.different_first_page_header_footer)
        self.assertFalse(footer_has_page(sec.first_page_footer))  # 首页不显示页码
        head = sec.first_page_header.paragraphs[0]
        self.assertEqual(head.text, "某市第一医院")
        self.assertEqual(str(head.runs[0].font.color.rgb), "FF0000")
        top_rule = head._p.pPr.find(qn("w:pBdr")).find(qn("w:bottom"))
        thin_rule = sec.first_page_header.paragraphs[1]._p.pPr.find(qn("w:pBdr")).find(qn("w:bottom"))
        self.assertGreater(int(top_rule.get(qn("w:sz"))), int(thin_rule.get(qn("w:sz"))))  # 上粗下细
        foot = sec.first_page_footer.paragraphs
        self.assertLess(int(foot[0]._p.pPr.find(qn("w:pBdr")).find(qn("w:bottom")).get(qn("w:sz"))),
                        int(foot[1]._p.pPr.find(qn("w:pBdr")).find(qn("w:bottom")).get(qn("w:sz"))))  # 上细下粗
        self.assertEqual(round(sec.header_distance / Mm(1)), 30)
        # 信函版记只有抄送，没有分隔线
        borders = doc.tables[-1]._tbl.tblPr.find(qn("w:tblBorders"))
        self.assertEqual(borders.find(qn("w:top")).get(qn("w:val")), "nil")
        self.assertNotIn("印发", "".join(c.text for row in doc.tables[-1].rows for c in row.cells))

    def test_minutes_elements(self):
        doc = self.docs["纪要"][2]
        issuer = find_para(doc, "某理工大学校长办公会纪要")
        self.assertEqual(str(issuer.runs[0].font.color.rgb), "FF0000")
        att = find_para(doc, "出席：")
        self.assertEqual(att.runs[0].text, "出席：")
        self.assertEqual(run_fonts(att)[0], "黑体")
        self.assertEqual(run_fonts(att)[-1], "仿宋_GB2312")

    def test_no_seal_material(self):
        _p, report, doc = self.docs["材料"]
        self.assertFalse(report["seal"])
        sig = [p for p in doc.paragraphs if p.text == "某能源集团有限公司人力资源部"][0]
        date = find_para(doc, "2026年12月20日")
        # 不加盖印章：署名右空二字（左缩进 28−2−14=12 字）；日期左对齐，首字右移二字（左缩进 14 字）
        self.assertEqual(sig._p.pPr.find(qn("w:jc")).get(qn("w:val")), "left")
        self.assertEqual(date._p.pPr.find(qn("w:jc")).get(qn("w:val")), "left")
        self.assertEqual(sig._p.pPr.find(qn("w:ind")).get(qn("w:leftChars")), "1200")
        self.assertEqual(date._p.pPr.find(qn("w:ind")).get(qn("w:leftChars")), "1400")

    def test_signoff_kept_with_body(self):
        """落款不单独成页：正文最后一段、空行、署名设“与下段同页”，有附注时日期也接上。"""
        for layout, last_body, tail in (("材料", "（三）开展人力资源信息系统升级。", "2026年12月20日"),
                                        ("白头", "妥否，请批示。", "（联系人：李某，联系电话：0000-87654321）")):
            doc = self.docs[layout][2]
            paras = doc.paragraphs
            texts = [p.text for p in paras]
            start, end = texts.index(last_body), texts.index(tail)
            for p in paras[start:end]:
                self.assertIsNotNone(p._p.pPr.find(qn("w:keepNext")), f"{layout}: {p.text}")
            self.assertIsNone(paras[end]._p.pPr.find(qn("w:keepNext")))
            self.assertIsNone(paras[start]._p.pPr.find(qn("w:keepLines")))  # 正文段可拆分，至少一行随落款换页
            self.assertIsNone(paras[start - 1]._p.pPr.find(qn("w:keepNext")))

    def test_attachment_lines_kept_together(self):
        doc = self.docs["红头"][2]
        first = find_para(doc, "附件：1. ")
        self.assertIsNotNone(first._p.pPr.find(qn("w:keepNext")))

    def test_long_title_breaks_evenly(self):
        doc = self.docs["材料"][2]
        title = doc.paragraphs[0]
        self.assertEqual(title.text, "某能源集团有限公司人力资源部\n2026年工作总结")
        self.assertEqual(len(title._p.findall(".//" + qn("w:br"))), 1)


class BuildBehaviourTest(unittest.TestCase):
    def test_default_layouts(self):
        _p, r1 = build_text("关于某事的请示\n\n县人民政府：\n\n现请示如下。\n\n某镇人民政府\n\n2026年1月5日", "默认-请示.docx")
        self.assertEqual(r1["layout"], "白头")
        _p, r2 = build_text("某科室2026年工作总结\n\n今年完成三项工作。", "默认-总结.docx")
        self.assertEqual(r2["layout"], "材料")

    def test_placeholders_not_invented(self):
        path, report = build_text((FILES / "请示-缺落款-正文.txt").read_text(encoding="utf-8"), "缺落款.docx",
                                  meta={"版式": "红头", "发文机关标志": "某街道办事处文件"})
        text = "\n".join(p.text for p in Document(str(path)).paragraphs)
        for mark in ("【待补：发文字号】", "【待补：签发人姓名】", "【待补：发文机关署名】", "【待补：成文日期】"):
            self.assertIn(mark, text)
        self.assertTrue(report["placeholders"])
        result = check_docx.check_document(path, layout="红头")
        self.assertEqual(result["summary"]["P0"], 0)
        self.assertTrue(any(i["id"] == "DOCX-080" for i in result["issues"]))

    def test_meta_overrides_and_iso_dates(self):
        path, report = build_text("关于某事的通知\n\n各单位：\n\n现通知如下。", "版式信息优先.docx",
                                  meta={"版式": "白头", "发文机关署名": "某市某局", "成文日期": "2026-03-05",
                                        "抄送": ["市委办公室", "市政府办公室"], "印发机关": "某市某局办公室",
                                        "印发日期": "2026-03-05", "印章": "否"})
        doc = Document(str(path))
        texts = [p.text for p in doc.paragraphs]
        self.assertIn("2026年3月5日", texts)
        cells = [c.text for row in doc.tables[-1].rows for c in row.cells]
        self.assertEqual(cells[0], "抄送：市委办公室，市政府办公室。")
        self.assertFalse(report["seal"])

    def test_cli_no_overwrite_and_json(self):
        src = FILES / "白头请示-正文.txt"
        target = OUT / "不覆盖.docx"
        for old in OUT.glob("不覆盖*.docx"):
            old.unlink()
        target.write_bytes(b"placeholder")
        buf = io.StringIO()
        with redirect_stdout(buf):
            code = build_docx.main([str(src), "-o", str(target), "--json"])
        self.assertEqual(code, 0)
        report = json.loads(buf.getvalue())
        self.assertTrue(report["output"].endswith("不覆盖-v2.docx"))
        self.assertEqual(target.read_bytes(), b"placeholder")
        self.assertEqual(report["check"]["summary"]["P0"], 0)
        with redirect_stdout(io.StringIO()):
            build_docx.main([str(src), "-o", str(target)])
        self.assertTrue((OUT / "不覆盖-v3.docx").exists())

    def test_date_longer_than_signature_no_seal(self):
        path, _ = build_text("关于某事的通知\n\n各单位：\n\n现通知如下。\n\n某局\n\n2026年10月16日", "日期较长.docx",
                             meta={"印章": "否"})
        doc = Document(str(path))
        sig = [p for p in doc.paragraphs if p.text == "某局"][0]
        date = [p for p in doc.paragraphs if p.text == "2026年10月16日"][0]
        # 日期右空二字；署名首字比日期首字左移二字：左缩进 28−2−7−2=17 字
        self.assertEqual(date._p.pPr.find(qn("w:jc")).get(qn("w:val")), "right")
        self.assertEqual(date._p.pPr.find(qn("w:ind")).get(qn("w:rightChars")), "200")
        self.assertEqual(sig._p.pPr.find(qn("w:ind")).get(qn("w:leftChars")), "1700")

    def test_user_title_breaks_respected_and_checked(self):
        path, report = build_text("某科室工作总结\n\n今年完成三项工作。", "标题手动分行.docx",
                                  meta={"标题": "某能源集团有限公司人力资源部2026年工作\n总结"})
        doc = Document(str(path))
        self.assertEqual(doc.paragraphs[0].text, "某能源集团有限公司人力资源部2026年工作\n总结")
        result = check_docx.check_document(path)
        self.assertIn("DOCX-024", {i["id"] for i in result["issues"]})
        self.assertEqual(result["summary"]["P0"], 0)

    def test_cli_default_output_name(self):
        folder = OUT / "公文" / "20261028-默认文件名"
        folder.mkdir(parents=True, exist_ok=True)
        for old in folder.glob("*"):
            old.unlink()
        src = folder / "正文-单位总结.txt"
        src.write_text((FILES / "年终总结-正文.txt").read_text(encoding="utf-8"), encoding="utf-8")
        with redirect_stdout(io.StringIO()):
            self.assertEqual(build_docx.main([str(src)]), 0)
        self.assertTrue((folder / "成稿-单位总结.docx").exists())

    def test_cli_bad_layout(self):
        with redirect_stdout(io.StringIO()):
            code = build_docx.main([str(FILES / "白头请示-正文.txt"), "--layout", "蓝头", "-o", str(OUT / "坏版式.docx")])
        self.assertEqual(code, 2)

    def test_unique_path(self):
        p = OUT / "唯一路径-v4.docx"
        for old in OUT.glob("唯一路径*.docx"):
            old.unlink()
        p.write_bytes(b"x")
        self.assertEqual(gw_docx.unique_path(p).name, "唯一路径-v5.docx")


class FontStrictModeTest(unittest.TestCase):
    """字体严格模式：文档一律写入标准字体名；缺字体时只指出缺哪些并引导安装，绝不降级、替代。"""

    NOT_STANDARD = ("仿宋", "楷体", "华文中宋", "华文仿宋", "方正小标宋_GBK", "方正仿宋_GBK")

    def assert_only_standard(self, path):
        used = written_fonts(path)
        self.assertTrue(used <= set(STANDARD_FONTS.values()), used)
        for name in self.NOT_STANDARD:
            self.assertNotIn(name, used)

    def test_missing_fonts_written_as_standard_and_reported(self):
        # 本机只有现行 Windows 自带的“仿宋”“楷体”“黑体”“宋体”：这些不是仿宋_GB2312 等，不能顶替
        path, report = build_sample("白头", "字体缺失.docx", installed=WIN_FONTS)
        self.assertEqual(report["fonts"], STANDARD_FONTS)
        self.assert_only_standard(path)
        self.assertEqual(report["font_status"], "missing")
        self.assertEqual([m["font"] for m in report["missing_fonts"]], ["仿宋_GB2312", "楷体_GB2312", "方正小标宋简体"])
        self.assertEqual([m["role"] for m in report["missing_fonts"]], ["仿宋", "楷体", "小标宋"])
        for item in report["missing_fonts"]:
            for key in ("usage", "source", "obtain", "note"):
                self.assertTrue(item["guide"][key], (item["font"], key))
        by_font = {m["font"]: m["guide"] for m in report["missing_fonts"]}
        self.assertIn("https://www.foundertype.com", by_font["方正小标宋简体"]["obtain"])
        self.assertIn("单位信息中心", by_font["仿宋_GB2312"]["obtain"])
        notice = report["font_notice"]
        self.assertEqual(notice["headline"], "字体缺失")
        self.assertIn("不是正式效果", notice["consequence"])
        self.assertIn("不能据此核对版式或打印", notice["consequence"])
        self.assertIn("不提供、不打包任何字体文件", notice["statement"])
        self.assertEqual(set(notice["install_methods"]), {"Windows", "macOS", "Linux", "安装后"})
        self.assertEqual(report["unverified_fonts"], [])
        # 报告里没有任何“已改用替代字体”的痕迹
        self.assertNotIn("font_warnings", report)
        self.assertFalse(any("改用" in n for n in report["notes"]), report["notes"])

    def test_screen_report_has_prominent_block_first(self):
        _path, report = build_sample("白头", "字体缺失屏显.docx", installed=WIN_FONTS)
        text = build_docx.render_text(dict(report, output="x.docx"))
        self.assertTrue(text.startswith("="), text[:30])
        block = text.split("输出文件：")[0]
        for must in ("【字体缺失】", "仿宋_GB2312", "楷体_GB2312", "方正小标宋简体", "不是正式效果", "不能据此核对版式或打印",
                     "https://www.foundertype.com", "为所有用户安装", "字体册", "fc-cache -f", "重启 Word/WPS",
                     "不提供、不打包任何字体文件"):
            self.assertIn(must, block)
        self.assertNotIn("黑体（用于", block)  # 黑体、宋体已装，不在缺失清单里
        self.assertIn("本机检测：缺少 3 款标准字体", text)

    def test_complete_fonts_no_notice(self):
        path, report = build_sample("白头", "字体齐全.docx", installed=FULL_FONTS)
        self.assert_only_standard(path)
        self.assertEqual(report["font_status"], "complete")
        self.assertEqual(report["missing_fonts"], [])
        self.assertIsNone(report["font_notice"])
        text = build_docx.render_text(dict(report, output="x.docx"))
        self.assertNotIn("字体缺失", text)
        self.assertIn("标准字体均已安装", text)

    def test_undetected_fonts_requires_manual_check(self):
        path, report = build_sample("白头", "字体无法检测.docx", installed=NO_FONTS)
        self.assert_only_standard(path)
        self.assertEqual(report["font_status"], "undetected")
        self.assertEqual(report["missing_fonts"], [])
        self.assertEqual([m["font"] for m in report["unverified_fonts"]],
                         ["仿宋_GB2312", "楷体_GB2312", "黑体", "方正小标宋简体", "宋体"])
        text = build_docx.render_text(dict(report, output="x.docx"))
        block = text.split("输出文件：")[0]
        self.assertIn("无法检测本机字体，需人工核对", block)
        self.assertIn("人工核对", block)
        self.assertIn("https://www.foundertype.com", block)  # 无法确认时也给出获取与安装引导

    def test_every_layout_writes_only_standard_fonts(self):
        for layout in ("红头", "信函", "纪要", "白头", "材料", "简报"):
            path, report = build_sample(layout, f"字体严格-{layout}.docx", installed=WIN_FONTS)
            self.assert_only_standard(path)
            self.assertEqual(report["font_status"], "missing", layout)

    def test_font_fallback_option_removed(self):
        src = FILES / "白头请示-正文.txt"
        buf = io.StringIO()
        with redirect_stderr(buf), self.assertRaises(SystemExit) as ctx:
            build_docx.main([str(src), "-o", str(OUT / "无降级参数.docx"), "--font-fallback"])
        self.assertEqual(ctx.exception.code, 2)
        self.assertIn("--font-fallback", buf.getvalue())
        self.assertNotIn("font_fallback", inspect.signature(build_docx.build).parameters)

    def test_cli_json_and_text_report_missing_fonts(self):
        src = FILES / "白头请示-正文.txt"
        with fonts_installed(WIN_FONTS):
            buf = io.StringIO()
            with redirect_stdout(buf):
                self.assertEqual(build_docx.main([str(src), "-o", str(OUT / "字体缺失CLI-json.docx"), "--json"]), 0)
            report = json.loads(buf.getvalue())
            self.assertEqual(report["font_status"], "missing")
            self.assertEqual({m["font"] for m in report["missing_fonts"]}, {"仿宋_GB2312", "楷体_GB2312", "方正小标宋简体"})
            self.assertEqual(report["fonts"], STANDARD_FONTS)
            self.assert_only_standard(Path(report["output"]))
            buf = io.StringIO()
            with redirect_stdout(buf):
                self.assertEqual(build_docx.main([str(src), "-o", str(OUT / "字体缺失CLI-text.docx")]), 0)
            self.assertTrue(buf.getvalue().startswith("="))
            self.assertIn("【字体缺失】", buf.getvalue())
        with fonts_installed(FULL_FONTS):
            buf = io.StringIO()
            with redirect_stdout(buf):
                self.assertEqual(build_docx.main([str(src), "-o", str(OUT / "字体齐全CLI.docx"), "--json"]), 0)
            report = json.loads(buf.getvalue())
            self.assertEqual((report["font_status"], report["missing_fonts"], report["font_notice"]),
                             ("complete", [], None))

    def test_check_does_not_flag_standard_fonts(self):
        """缺字体不影响文件本身：写入标准字体名的文档，机检没有 DOCX-016。"""
        path, _ = build_sample("红头", "字体缺失机检.docx", installed=WIN_FONTS)
        result = check_docx.check_document(path, layout="红头")
        self.assertNotIn("DOCX-016", {i["id"] for i in result["issues"]})
        self.assertEqual(result["summary"]["P0"], 0)


class BuildInputRobustnessTest(unittest.TestCase):
    """输入异常时给出提示并返回 2，不抛异常；正常的 GBK 正文、输出到目录要能处理。"""

    def setUp(self):
        import shutil
        self.folder = OUT / "输入 健壮性"
        if self.folder.exists():
            shutil.rmtree(self.folder)
        self.folder.mkdir(parents=True)

    def run_main(self, args):
        buf = io.StringIO()
        with redirect_stdout(buf):
            code = build_docx.main(args)
        return code, buf.getvalue()

    def test_gbk_body_file(self):
        src = self.folder / "正文.txt"
        src.write_bytes("关于某事的通知\r\n\r\n各单位：\r\n\r\n现通知如下。\r\n\r\n某局\r\n\r\n2026年3月5日\r\n".encode("gbk"))
        code, out = self.run_main([str(src), "--no-check", "--json"])
        self.assertEqual(code, 0, out)
        report = json.loads(out)
        self.assertEqual(report["fields"]["标题"], "关于某事的通知")
        self.assertTrue(any("GB18030" in n for n in report["notes"]))
        self.assertTrue((self.folder / "成稿.docx").exists())

    def test_undecodable_body_file(self):
        src = self.folder / "正文.txt"
        src.write_bytes("关于某事的通知".encode("utf-16"))
        code, out = self.run_main([str(src), "--no-check"])
        self.assertEqual(code, 2)
        self.assertIn("UTF-8", out)

    def test_meta_not_utf8(self):
        src = self.folder / "正文.txt"
        src.write_text("关于某事的通知\n\n各单位：\n\n现通知如下。", encoding="utf-8")
        meta = self.folder / "版式信息.json"
        meta.write_bytes(json.dumps({"版式": "白头"}, ensure_ascii=False).encode("gbk"))
        code, out = self.run_main([str(src), "--meta", str(meta), "--no-check"])
        self.assertEqual(code, 2)
        self.assertIn("版式信息 JSON 读取失败", out)

    def test_output_to_existing_directory(self):
        """-o 给的是已存在的文件夹时，在文件夹里按默认名输出，不在文件夹旁边生成无扩展名的文件。"""
        src = self.folder / "正文-单位总结.txt"
        src.write_text((FILES / "年终总结-正文.txt").read_text(encoding="utf-8"), encoding="utf-8")
        target = self.folder / "输出目录"
        target.mkdir(exist_ok=True)
        code, out = self.run_main([str(src), "-o", str(target), "--no-check", "--json"])
        self.assertEqual(code, 0, out)
        self.assertEqual(json.loads(out)["output"], str(target / "成稿-单位总结.docx"))
        self.assertTrue((target / "成稿-单位总结.docx").exists())
        self.assertFalse((self.folder / "输出目录-v2").exists())


DOCNO_UNDER_TITLE = ("关于开展2026年冬季校园安全检查的通知\n\n某区教发[2026]第12号\n\n各中小学、幼儿园：\n\n"
                     "现将有关事项通知如下。\n\n某区教育局\n\n2026年11月1日")
RED_META = {"版式": "红头", "发文机关标志": "某区教育局文件", "印发机关": "某区教育局办公室", "印发日期": "2026-11-01"}


class DocnoUnderTitleTest(unittest.TestCase):
    """AI 稿常把发文字号写在标题下：排版时移到版头发文字号位置，正文中不再保留。"""

    def test_moved_to_head(self):
        path, report = build_text(DOCNO_UNDER_TITLE, "文号在标题下.docx", meta=RED_META)
        texts = [p.text for p in Document(str(path)).paragraphs]
        self.assertEqual(texts.count("某区教发[2026]第12号"), 1, texts)
        self.assertLess(texts.index("某区教发[2026]第12号"), texts.index("关于开展2026年冬季校园安全检查的通知"))
        self.assertFalse(any("【待补：发文字号】" in t for t in texts))
        self.assertEqual(report["fields"]["发文字号"], "某区教发[2026]第12号")
        self.assertTrue(any("某区教发[2026]第12号" in n and "版头" in n for n in report["notes"]), report["notes"])

    def test_meta_wins_and_body_line_removed(self):
        meta = dict(RED_META, 发文字号="某区教发〔2026〕12号")
        path, report = build_text(DOCNO_UNDER_TITLE, "文号在标题下-版式信息.docx", meta=meta)
        texts = [p.text for p in Document(str(path)).paragraphs]
        self.assertIn("某区教发〔2026〕12号", texts)
        self.assertFalse(any("某区教发[2026]第12号" in t for t in texts), texts)
        self.assertTrue(any("某区教发[2026]第12号" in n and "版式信息" in n for n in report["notes"]), report["notes"])

    def test_white_head_reports_removal(self):
        path, report = build_text(DOCNO_UNDER_TITLE, "文号在标题下-白头.docx")
        self.assertEqual(report["layout"], "白头")
        texts = [p.text for p in Document(str(path)).paragraphs]
        self.assertFalse(any("某区教发" in t for t in texts))
        self.assertTrue(any("某区教发[2026]第12号" in n for n in report["notes"] + report["warnings"]))


class ReportNumberingTest(unittest.TestCase):
    def test_sections_numbered_consecutively(self):
        """没有待补事项时不能从“二、”跳到“四、”。"""
        import re
        report = {"layout": "材料", "summary": [], "fonts": {}, "placeholders": [], "warnings": ["提示"], "notes": [],
                  "check": {"summary": {"P0": 0, "P1": 0, "P2": 0}, "issues": []}}
        heads = re.findall(r"^([一二三四五六七八九十]+)、", build_docx.render_text(report), flags=re.M)
        self.assertEqual(heads, ["一", "二", "三", "四"])


class DefaultOutputNameTest(unittest.TestCase):
    """默认输出名：正文.txt → 成稿.docx；正文-××.txt → 成稿-××.docx；其他文件 → 同名 .docx；已存在时加 -v2。"""

    def setUp(self):
        import shutil
        self.folder = OUT / "默认输出名"
        if self.folder.exists():
            shutil.rmtree(self.folder)
        self.folder.mkdir(parents=True)

    def build(self, name):
        src = self.folder / name
        src.write_text("关于某事的通知\n\n各单位：\n\n现通知如下。", encoding="utf-8")
        with redirect_stdout(io.StringIO()):
            self.assertEqual(build_docx.main([str(src), "--no-check"]), 0)

    def test_other_stem_keeps_name(self):
        (self.folder / "送审稿.docx").write_bytes(b"placeholder")
        self.build("送审稿.txt")
        self.assertEqual((self.folder / "送审稿.docx").read_bytes(), b"placeholder")
        self.assertTrue((self.folder / "送审稿-v2.docx").exists())
        self.assertFalse((self.folder / "成稿.docx").exists())

    def test_zhengwen_maps_to_chenggao(self):
        self.build("正文.txt")
        self.assertTrue((self.folder / "成稿.docx").exists())
        self.build("正文-核改.txt")
        self.assertTrue((self.folder / "成稿-核改.docx").exists())


class MissingDocxTest(unittest.TestCase):
    """模拟没有安装 python-docx：各脚本给出安装提示、返回 2，不抛异常。"""

    def test_all_clis(self):
        import subprocess
        import sys
        from tests.format.helpers import SCRIPTS
        code = ("import sys, runpy; sys.modules['docx'] = None; script = sys.argv[1]; sys.argv = sys.argv[1:]; "
                "runpy.run_path(script, run_name='__main__')")
        sample = str(FILES / "白头请示-正文.txt")
        cases = {"build_docx.py": [sample, "-o", str(OUT / "无依赖.docx")], "check_docx.py": ["任意.docx"],
                 "fix_docx.py": ["任意.docx"], "annotate_docx.py": ["任意.docx", "任意.json"]}
        for name, args in cases.items():
            r = subprocess.run([sys.executable, "-X", "utf8", "-c", code, str(SCRIPTS / name)] + args,
                               capture_output=True, text=True, encoding="utf-8")
            self.assertEqual(r.returncode, 2, name)
            self.assertIn("pip install -U python-docx", r.stdout, name)
            self.assertNotIn("Traceback", r.stderr, name)


if __name__ == "__main__":
    unittest.main()
