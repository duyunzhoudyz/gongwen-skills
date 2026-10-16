# -*- coding: utf-8 -*-
"""v0.1 评测发现的问题（evals/results/v0.1-待修问题.md FM-1、FM-4、FM-5、FM-6）的回归测试。"""
import io
import json
import unittest
from contextlib import redirect_stdout

from docx import Document
from docx.oxml.ns import qn

from tests.format.helpers import OUT, build_sample, build_text, check_docx, gw_layout

REQUEST = ("关于申请更新实验室消防设施的请示\n\n市应急管理局：\n\n我院实验楼消防设施已使用多年，需要更新。\n\n"
           "妥否，请批示。\n\n某职业技术学院\n\n2026年4月8日\n\n（联系人：周某，联系电话：0000-55551234）")
RED_UP = {"版式": "红头", "发文机关标志": "某职业技术学院文件", "发文字号": "某职院发〔2026〕【待补：序号】号",
          "签发人": "钱某", "印发机关": "某职业技术学院办公室", "印发日期": "2026-04-08"}


def ids(result):
    return {(i["id"], i["level"]) for i in result["issues"]}


class DocnoPlaceholderTest(unittest.TestCase):
    """FM-1：发文字号序号写【待补：序号】时，识别为发文字号行（可带签发人），文种不判为“其他”。"""

    def test_parse_head_line(self):
        parsed = gw_layout.parse_lines(["某职院发〔2026〕【待补：序号】号　签发人：钱某", ""] + REQUEST.split("\n"))
        self.assertEqual(parsed.items[0].role, "docno")
        self.assertEqual(parsed.fields["发文字号"], "某职院发〔2026〕【待补：序号】号")
        self.assertEqual(parsed.fields["签发人"], "钱某")
        self.assertEqual(parsed.doc_type, "请示")

    def test_build_and_check_red_upward(self):
        path, report = build_text(REQUEST, "红头请示-序号待补.docx", meta=RED_UP)
        self.assertFalse(any("签发人会被挤到下一行" in w for w in report["warnings"]))
        doc = Document(str(path))
        line = [p for p in doc.paragraphs if "签发人" in p.text][0]
        self.assertIn("某职院发〔2026〕【待补：序号】号", line.text)
        self.assertIn("钱某", line.text)
        result = check_docx.check_document(path, layout="红头")
        self.assertEqual(result["doc_type"], "请示")
        self.assertEqual(result["summary"]["P0"], 0, result["issues"])
        self.assertEqual(result["summary"]["P1"], 0, result["issues"])
        self.assertTrue(any(i["id"] == "DOCX-080" and "序号" in i["excerpt"] for i in result["issues"]))

    def test_placeholder_docno_errors_still_reported(self):
        """序号写【待补】时，年份括号写错、加“第”字仍要报 P0。"""
        for docno, rule in (("某职院发[2026]【待补：序号】号", "DOCX-050"),
                            ("某职院发〔2026〕第【待补：序号】号", "DOCX-052")):
            with self.subTest(docno=docno):
                path, _ = build_text(REQUEST, f"红头请示-序号待补-{rule}.docx", meta=dict(RED_UP, 发文字号=docno))
                result = check_docx.check_document(path, layout="红头")
                self.assertIn((rule, "P0"), ids(result))

    def test_long_placeholder_warns(self):
        meta = dict(RED_UP, 发文字号="【待补：发文字号，待党政办公室统一编号后填写】", 签发人=["钱某", "孙某某"])
        _path, report = build_text(REQUEST, "红头请示-长占位.docx", meta=meta)
        self.assertTrue(any("签发人会被挤到下一行" in w for w in report["warnings"]), report["warnings"])


class MinutesAttendanceGapTest(unittest.TestCase):
    """FM-4：纪要出席名单前空一行；没有空行时 check 报 P1；fix 补上。"""

    @classmethod
    def setUpClass(cls):
        cls.path, _ = build_sample("纪要", "纪要-出席空行.docx")

    def test_blank_line_before_attendance(self):
        doc = Document(str(self.path))
        paras = doc.paragraphs
        k = [i for i, p in enumerate(paras) if p.text.startswith("出席：")][0]
        self.assertEqual(paras[k - 1].text, "")
        self.assertTrue(paras[k - 2].text.strip())
        sp = paras[k - 1]._p.pPr.find(qn("w:spacing"))
        self.assertEqual(sp.get(qn("w:line")), "578")  # 整行高，不是压缩的空段
        # 请假、列席紧接出席，不再空行
        self.assertTrue(paras[k + 1].text.startswith("请假："))
        result = check_docx.check_document(self.path, layout="纪要")
        self.assertNotIn("DOCX-047", {i["id"] for i in result["issues"]})

    def _without_blank(self):
        doc = Document(str(self.path))
        k = [i for i, p in enumerate(doc.paragraphs) if p.text.startswith("出席：")][0]
        blank = doc.paragraphs[k - 1]._p
        blank.getparent().remove(blank)
        bad = OUT / "标本-纪要出席无空行.docx"
        doc.save(str(bad))
        return bad

    def test_missing_blank_reported(self):
        bad = self._without_blank()
        self.assertIn(("DOCX-047", "P1"), ids(check_docx.check_document(bad, layout="纪要")))

    def test_fix_adds_blank(self):
        import fix_docx
        bad = self._without_blank()
        dst = OUT / "标本-纪要出席无空行-修复.docx"
        if dst.exists():
            dst.unlink()
        report = fix_docx.fix_document(bad, dst, installed=set())
        self.assertTrue(any("出席名单前已空一行" in n for n in report["notes"]))
        self.assertNotIn("DOCX-047", {i["id"] for i in check_docx.check_document(dst, layout="纪要")["issues"]})

    def test_not_checked_outside_minutes(self):
        text = "关于某事的通知\n\n各单位：\n\n现通知如下。\n出席：张某\n\n某市某局\n\n2026年3月5日"
        path, _ = build_text(text, "通知-含出席.docx", meta={"版式": "白头"})
        self.assertNotIn("DOCX-047", {i["id"] for i in check_docx.check_document(path)["issues"]})


class HeadingKeepNextTest(unittest.TestCase):
    """FM-5：独占一行的层次标题设“与下段同页”；与正文同段的、以句号结尾的条目不设。"""

    def test_keep_next(self):
        path, _ = build_sample("红头", "层次标题同页.docx")
        doc = Document(str(path))
        whole = [p for p in doc.paragraphs if p.text in ("一、检查范围", "二、检查重点", "三、工作要求")]
        self.assertTrue(whole)
        for p in whole:
            self.assertIsNotNone(p._p.pPr.find(qn("w:keepNext")), p.text)
        inline = [p for p in doc.paragraphs if p.text.startswith("（一）安全责任落实情况。")][0]
        self.assertIsNone(inline._p.pPr.find(qn("w:keepNext")))

    def test_fix_sets_keep_next(self):
        import fix_docx
        from tests.format.helpers import make_messy_docx
        src = make_messy_docx(OUT / "层次标题同页-乱稿.docx")
        dst = OUT / "层次标题同页-乱稿-修复.docx"
        if dst.exists():
            dst.unlink()
        fix_docx.fix_document(src, dst, installed=set())
        doc = Document(str(dst))
        for text in ("一、主要工作", "二、下一步打算"):
            p = [p for p in doc.paragraphs if p.text == text][0]
            self.assertIsNotNone(p._p.pPr.find(qn("w:keepNext")), text)

    def test_period_ended_item_not_kept(self):
        path, _ = build_text("某科室工作总结\n\n一、主要工作\n\n（一）完成设备巡检。\n\n（二）完成人员培训。", "层次标题-句号条目.docx")
        doc = Document(str(path))
        self.assertIsNotNone([p for p in doc.paragraphs if p.text == "一、主要工作"][0]._p.pPr.find(qn("w:keepNext")))
        item = [p for p in doc.paragraphs if p.text == "（一）完成设备巡检。"][0]
        self.assertIsNone(item._p.pPr.find(qn("w:keepNext")))


class AllowTest(unittest.TestCase):
    """FM-6：--allow DOCX-050 时方括号文号降为 P2、注明按本单位制度保留，不计入 --strict。"""

    @classmethod
    def setUpClass(cls):
        meta = {"版式": "红头", "发文机关标志": "某建材集团有限公司文件", "发文字号": "某建材发［2026］9号",
                "印发机关": "某建材集团有限公司办公室", "印发日期": "2026-05-12"}
        cls.path, _ = build_text("关于某事的通知\n\n各子公司：\n\n现通知如下。\n\n某建材集团有限公司\n\n2026年5月12日",
                                 "企业方括号文号.docx", meta=meta)

    def run_cli(self, *extra):
        buf = io.StringIO()
        with redirect_stdout(buf):
            code = check_docx.main([str(self.path), "--json", "--strict", *extra])
        return code, json.loads(buf.getvalue())

    def test_default_unchanged(self):
        code, data = self.run_cli()
        self.assertEqual(code, 1)
        self.assertIn(("DOCX-050", "P0"), ids(data))

    def test_allow_downgrades(self):
        code, data = self.run_cli("--allow", "docx-050")
        self.assertEqual(code, 0)
        hit = [i for i in data["issues"] if i["id"] == "DOCX-050"][0]
        self.assertEqual(hit["level"], "P2")
        self.assertTrue(hit["message"].startswith("按本单位制度保留："))
        self.assertEqual(data["summary"]["P0"], 0)
        self.assertTrue(any("不计入 --strict" in n for n in data["notes"]))

    def test_unknown_id_noted(self):
        _code, data = self.run_cli("--allow", "DOCX-050,DOCX-999")
        self.assertTrue(any("DOCX-999" in n and "不存在" in n for n in data["notes"]))

    def test_build_passes_allow_to_readback(self):
        import build_docx
        folder = OUT / "按制度保留"
        folder.mkdir(parents=True, exist_ok=True)
        for old in folder.glob("*"):
            old.unlink()
        src = folder / "正文.txt"
        src.write_text("关于某事的通知\n\n各子公司：\n\n现通知如下。\n\n某建材集团有限公司\n\n2026年5月12日", encoding="utf-8")
        meta = folder / "版式信息.json"
        meta.write_text(json.dumps({"版式": "红头", "发文机关标志": "某建材集团有限公司文件", "发文字号": "某建材发［2026］9号",
                                    "印发机关": "某建材集团有限公司办公室", "印发日期": "2026-05-12"}, ensure_ascii=False),
                        encoding="utf-8")
        buf = io.StringIO()
        with redirect_stdout(buf):
            self.assertEqual(build_docx.main([str(src), "--meta", str(meta), "--json", "--allow", "DOCX-050"]), 0)
        self.assertEqual(json.loads(buf.getvalue())["check"]["summary"]["P0"], 0)


if __name__ == "__main__":
    unittest.main()
