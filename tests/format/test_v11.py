# -*- coding: utf-8 -*-
"""v1.1 版式与缺陷修复：上行文发文字号居左空一字、fix_docx --keep-text、命令（令）、附件另面编排、
联合行文版头与联署署名、预印红头纸套打。每项都有 build 生成、check_docx 回读和变异检出的测试。"""
import io
import json
import unittest
from contextlib import redirect_stdout

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from tests.format.helpers import NO_FONTS, OUT, build_text, check_docx, gw_docx, make_messy_docx

import build_docx  # noqa: E402
import fix_docx  # noqa: E402

REQUEST = ("关于申请增加2027年硕士研究生招生计划的请示\n\n省教育厅：\n\n我校现有硕士学位授权一级学科12个，2026年招收硕士研究生860人。"
           "近三年新增国家级科研平台2个，具备扩大培养规模的条件。\n\n妥否，请批示。\n\n某理工大学\n\n2026年10月26日")
UP_META = {"版式": "红头", "发文机关标志": "某理工大学文件", "发文字号": "某理工发〔2026〕52号", "签发人": "王某",
           "印发机关": "某理工大学党政办公室", "印发日期": "2026-10-26"}
NOTICE = ("关于做好冬季消防安全检查的通知\n\n各区应急管理局：\n\n现就有关事项通知如下。\n\n一、总体要求\n\n全面排查火灾隐患。\n\n"
          "请认真贯彻落实。\n\n某市应急管理局\n\n2026年10月28日")
RED_META = {"版式": "红头", "发文机关标志": "某市应急管理局文件", "发文字号": "某应急发〔2026〕21号",
            "印发机关": "某市应急管理局办公室", "印发日期": "2026-10-28"}
ORDER = "第12号\n\n《某市某某管理办法》已经市人民政府常务会议通过，现予公布，自2027年1月1日起施行。\n\n市长　张某\n\n2026年10月1日"
ORDER_META = {"版式": "命令", "发文机关标志": "某市人民政府令"}
ATTACH_TEXT = (NOTICE + "\n\n【附件1】冬季消防安全检查表\n\n一、检查内容\n\n检查消防设施、疏散通道。\n\n表1　检查项目\n\n"
               "| 项目 | 标准 |\n|---|---|\n| 消防设施 | 完好有效 |\n\n【附件2】整改通知书（样式）\n\n兹有你单位存在隐患，限期整改。")
JOINT = ["某市应急管理局", "某市公安局"]


def ids(result):
    return {(i["id"], i["level"]) for i in result["issues"]}


def has(result, iid):
    return any(i["id"] == iid for i in result["issues"])


def docno_top_mm(path_or_doc) -> float:
    """发文字号段上沿距版心上边缘（毫米）：其前各段的段前距加行高。"""
    doc = Document(str(path_or_doc)) if not hasattr(path_or_doc, "paragraphs") else path_or_doc
    top = 0.0
    for p in doc.paragraphs:
        ppr = p._p.pPr
        sp = ppr.find(qn("w:spacing")) if ppr is not None else None
        before = int(sp.get(qn("w:before")) or 0) / 20 if sp is not None else 0.0
        if "〔" in p.text and "号" in p.text:
            return (top + before) * 25.4 / 72
        top += before + (int(sp.get(qn("w:line")) or 0) / 20 if sp is not None else 0.0)
    raise AssertionError("没有找到发文字号")


def para_with(doc, text):
    return [p for p in doc.paragraphs if text in p.text][0]


class UpwardDocnoTest(unittest.TestCase):
    """缺陷：上行文发文字号居左空一字（7.2.5），与最后一个签发人姓名同一行（7.2.6）。"""

    def test_build_indents_one_char(self):
        path, report = build_text(REQUEST, "v11-上行文.docx", meta=UP_META)
        doc = Document(str(path))
        line = para_with(doc, "签发人：")
        self.assertEqual(line.text, "某理工发〔2026〕52号\t签发人：王某")  # 同一行，无行首制表符
        self.assertEqual(gw_docx.para_left_chars(line, doc), 1.0)
        self.assertEqual(gw_docx.leading_offset_chars(line, doc), 1.0)
        result = check_docx.check_document(path, layout="红头")
        self.assertFalse(has(result, "DOCX-056"), result["issues"])
        self.assertEqual((result["summary"]["P0"], result["summary"]["P1"]), (0, 0))

    def _mutated(self, name, fn):
        path, _ = build_text(REQUEST, "v11-上行文-基础.docx", meta=UP_META)
        doc = Document(str(path))
        fn(doc, para_with(doc, "签发人："))
        out = OUT / name
        doc.save(str(out))
        return out

    def test_check_flags_flush_left(self):
        def flush(doc, p):
            gw_docx.set_para(p, align="left", left=0, line_exact=28.9, before_pt=0, after_pt=0)
        out = self._mutated("v11-上行文-顶格.docx", flush)
        result = check_docx.check_document(out, layout="红头")
        self.assertIn(("DOCX-056", "P1"), ids(result))
        self.assertIn("左空 0 字", [i for i in result["issues"] if i["id"] == "DOCX-056"][0]["message"])

    def test_check_flags_centered(self):
        def center(doc, p):
            gw_docx.set_para(p, align="center", left=0, line_exact=28.9, before_pt=0, after_pt=0)
        out = self._mutated("v11-上行文-居中.docx", center)
        self.assertIn(("DOCX-056", "P1"), ids(check_docx.check_document(out, layout="红头")))

    def test_check_accepts_two_legacy_forms(self):
        """旧版生成的行首制表符加制表位、行首一个全角空格，也算空一字。"""
        def legacy_tab(doc, p):
            gw_docx.set_para(p, align="left", left=0, line_exact=28.9, before_pt=0, after_pt=0)
            gw_docx.set_tabs(p, [(320, "left"), (8525, "right")])
            p.runs[0].text = "\t" + p.runs[0].text
        def full_space(doc, p):
            gw_docx.set_para(p, align="left", left=0, line_exact=28.9, before_pt=0, after_pt=0)
            p.runs[0].text = "　" + p.runs[0].text
        for name, fn in (("v11-上行文-制表符.docx", legacy_tab), ("v11-上行文-全角空格.docx", full_space)):
            with self.subTest(name=name):
                out = self._mutated(name, fn)
                self.assertFalse(has(check_docx.check_document(out, layout="红头"), "DOCX-056"))

    def test_not_upward_not_checked(self):
        path, _ = build_text(NOTICE, "v11-通知.docx", meta=RED_META)
        doc = Document(str(path))
        line = para_with(doc, "某应急发")
        self.assertEqual(line.paragraph_format.alignment, 1)  # 下行文仍居中
        self.assertFalse(has(check_docx.check_document(path, layout="红头"), "DOCX-056"))

    def test_fix_restores_indent_without_changing_text(self):
        out = self._mutated("v11-上行文-待修复.docx", lambda d, p: gw_docx.set_para(
            p, align="left", left=0, line_exact=28.9, before_pt=0, after_pt=0))
        dst = OUT / "v11-上行文-已修复.docx"
        if dst.exists():
            dst.unlink()
        before = [p.text for p in Document(str(out)).paragraphs]
        fix_docx.fix_document(out, dst, installed=NO_FONTS)
        doc = Document(str(dst))
        self.assertEqual(gw_docx.leading_offset_chars(para_with(doc, "签发人："), doc), 1.0)
        self.assertEqual([p.text for p in doc.paragraphs], before)
        self.assertFalse(has(check_docx.check_document(dst, layout="红头"), "DOCX-056"))


class KeepTextTest(unittest.TestCase):
    """缺陷：用户要求“一个字都别动”时，fix_docx --keep-text 只调格式，不改文字。"""

    def _messy(self, name):
        path = OUT / name
        make_messy_docx(path)
        doc = Document(str(path))
        doc.paragraphs[2].runs[0].text = "　　区人民政府:"  # 半角冒号
        doc.save(str(path))
        return path

    def _fix(self, src, dst_name, **kw):
        dst = OUT / dst_name
        if dst.exists():
            dst.unlink()
        report = fix_docx.fix_document(src, dst, installed=NO_FONTS, **kw)
        return dst, report

    def test_default_strips_indent_spaces(self):
        src = self._messy("v11-混乱-默认.docx")
        dst, report = self._fix(src, "v11-混乱-默认-修复.docx")
        self.assertNotEqual([p.text for p in Document(str(src)).paragraphs], [p.text for p in Document(str(dst)).paragraphs])

    def test_keep_text_changes_no_character(self):
        src = self._messy("v11-混乱-保留.docx")
        dst, report = self._fix(src, "v11-混乱-保留-修复.docx", keep_text=True)
        a, b = Document(str(src)), Document(str(dst))
        self.assertEqual([p.text for p in a.paragraphs], [p.text for p in b.paragraphs])
        self.assertTrue(report["keep_text"])
        self.assertTrue(any("--keep-text" in n for n in report["notes"]), report["notes"])
        # 格式照常调整：页面改为 A4，正文改为 3 号仿宋
        pgsz = b.sections[0]._sectPr.find(qn("w:pgSz"))
        self.assertEqual(round(gw_docx.twips_to_mm(int(pgsz.get(qn("w:w"))))), 210)
        body = [p for p in b.paragraphs if p.text.startswith("　　2026年")][0]
        self.assertEqual(body.runs[0].font.size.pt, 16)
        self.assertEqual(gw_docx.para_first_line_chars(body, b), 0)  # 缩进由空格承担，不叠加

    def test_cli_lists_text_issues_only_in_report(self):
        src = self._messy("v11-混乱-命令行.docx")
        out = OUT / "v11-混乱-命令行-修复.docx"
        for old in (out, out.with_name(out.stem + "-修复报告.md")):
            if old.exists():
                old.unlink()
        buf = io.StringIO()
        with redirect_stdout(buf):
            self.assertEqual(fix_docx.main([str(src), "-o", str(out), "--keep-text"]), 0)
        report = out.with_name(out.stem + "-修复报告.md").read_text(encoding="utf-8")
        self.assertIn("文字类问题（按 --keep-text 未改动，需人工处理）", report)
        self.assertIn("主送机关后应用全角冒号", report)
        self.assertEqual([p.text for p in Document(str(src)).paragraphs], [p.text for p in Document(str(out)).paragraphs])

    def test_text_change_aborts_without_saving(self):
        """保险：keep_text 时若文字被改动，中止且不保存。"""
        from unittest import mock
        src = self._messy("v11-混乱-保险.docx")
        dst = OUT / "v11-混乱-保险-修复.docx"
        if dst.exists():
            dst.unlink()
        import gw_format
        real = gw_format.apply_role

        def tamper(paragraph, role, ctx, **kw):
            if role == "body" and paragraph.runs:
                paragraph.runs[-1].text = paragraph.runs[-1].text + "。"
            return real(paragraph, role, ctx, **kw)
        with mock.patch.object(gw_format, "apply_role", tamper):
            with self.assertRaises(ValueError):
                fix_docx.fix_document(src, dst, installed=NO_FONTS, keep_text=True)
        self.assertFalse(dst.exists())


class OrderTest(unittest.TestCase):
    """命令（令）格式（10.2、7.3.5.3）。"""

    def setUp(self):
        self.path, self.report = build_text(ORDER, "v11-命令.docx", meta=ORDER_META)
        self.doc = Document(str(self.path))

    def test_structure(self):
        paras = [p.text for p in self.doc.paragraphs if p.text.strip()]
        self.assertEqual(paras[0], "某市人民政府令")
        self.assertEqual(paras[1], "第12号")
        self.assertTrue(paras[2].startswith("《某市某某管理办法》"))
        self.assertEqual(paras[-2], "市长　　张某")
        self.assertEqual(paras[-1], "2026年10月1日")
        issuer = self.doc.paragraphs[0]
        self.assertEqual(str(issuer.runs[0].font.color.rgb), "FF0000")
        self.assertEqual(issuer.paragraph_format.alignment, 1)
        # 标志上边缘距版心上边缘 20 mm（10.2）
        before = int(issuer._p.pPr.find(qn("w:spacing")).get(qn("w:before"))) / 20
        self.assertAlmostEqual(before * 25.4 / 72, 20, delta=0.2)
        self.assertEqual(self.report["layout"], "命令")
        self.assertFalse(any("标题" in s for s in self.report["summary"]), self.report["summary"])

    def test_check_clean_and_autodetect(self):
        for layout in ("命令", None):
            result = check_docx.check_document(self.path, layout=layout)
            self.assertEqual(result["layout"], "命令")
            self.assertEqual((result["summary"]["P0"], result["summary"]["P1"], result["summary"]["P2"]), (0, 0, 0), result["issues"])
            self.assertFalse(has(result, "DOCX-023"))  # 令没有标题，不报“未识别到标题”

    def test_positions(self):
        """令号下空二行正文、签名章右空四字、成文日期右空四字，职务与日期之间空一行。"""
        res = check_docx.check_document(self.path, layout="命令")
        self.assertFalse({"DOCX-123", "DOCX-124", "DOCX-125", "DOCX-126"} & {i["id"] for i in res["issues"]})
        date = para_with(self.doc, "2026年10月1日")
        space = 28 - gw_docx.para_left_chars(date, self.doc) - 6.5
        self.assertAlmostEqual(space, 4, delta=0.1)
        sig = para_with(self.doc, "市长")
        # 姓名占位 3 字，其右缘距版心右边缘 4 字：职务起点 = 28-4-3-2-职务宽
        self.assertAlmostEqual(gw_docx.para_left_chars(sig, self.doc), 28 - 4 - 3 - 2 - 2, delta=0.01)

    def test_mutations_detected(self):
        def mutate(name, fn):
            doc = Document(str(self.path))
            fn(doc)
            out = OUT / name
            doc.save(str(out))
            return check_docx.check_document(out, layout="命令")
        # 标志不是“令”结尾
        r = mutate("v11-命令-标志.docx", lambda d: setattr(d.paragraphs[0].runs[0], "text", "某市人民政府"))
        self.assertIn(("DOCX-120", "P1"), ids(r))
        # 标志不居中
        r = mutate("v11-命令-左对齐.docx", lambda d: setattr(d.paragraphs[0].paragraph_format, "alignment", 0))
        self.assertIn(("DOCX-121", "P1"), ids(r))
        # 成文日期不是右空四字
        r = mutate("v11-命令-日期.docx", lambda d: gw_docx.set_para(para_with(d, "2026年10月1日"), align="left", left=2,
                                                                  line_exact=28.9, before_pt=0, after_pt=0))
        self.assertIn(("DOCX-125", "P1"), ids(r))
        # 令号下只空一行
        def one_gap(d):
            blanks = [p for p in d.paragraphs if not p.text.strip()]
            blanks[2]._p.getparent().remove(blanks[2]._p)  # 令号后第二个空行
        r = mutate("v11-命令-空行.docx", one_gap)
        self.assertIn(("DOCX-124", "P1"), ids(r))

    def test_missing_lingno_placeholder_and_numeric(self):
        path, report = build_text("现予公布。\n\n市长　张某\n\n2026年10月1日", "v11-命令-缺令号.docx", meta=ORDER_META)
        self.assertTrue(any("令号" in s for s in report["placeholders"]))
        self.assertIn("【待补：令号】", "".join(p.text for p in Document(str(path)).paragraphs))
        path, _ = build_text("现予公布。\n\n市长　张某\n\n2026年10月1日", "v11-命令-数字令号.docx", meta=dict(ORDER_META, 令号="7"))
        self.assertIn("第7号", [p.text for p in Document(str(path)).paragraphs])

    def test_issuer_without_ling_warns(self):
        _, report = build_text(ORDER, "v11-命令-标志不规范.docx", meta=dict(ORDER_META, 发文机关标志="某市人民政府"))
        self.assertTrue(any("加“命令”或“令”" in w for w in report["warnings"]), report["warnings"])

    def test_meta_post_and_name_and_blank_slot(self):
        path, _ = build_text("现予公布。", "v11-命令-职务.docx",
                             meta=dict(ORDER_META, 令号="第3号", 签发人职务="省长", 成文日期="2026-10-01"))
        doc = Document(str(path))
        self.assertIn("省长", [p.text for p in doc.paragraphs])  # 没给姓名：只排职务，签名章位置留白
        res = check_docx.check_document(path, layout="命令")
        self.assertEqual(res["summary"]["P1"], 0, res["issues"])

    def test_fix_leaves_order_signoff(self):
        dst = OUT / "v11-命令-修复.docx"
        if dst.exists():
            dst.unlink()
        report = fix_docx.fix_document(self.path, dst, installed=NO_FONTS)
        self.assertEqual(report["layout"], "命令")
        self.assertTrue(any("命令（令）的签发人职务" in s for s in report["untouched"]))
        res = check_docx.check_document(dst, layout="命令")
        self.assertEqual((res["summary"]["P0"], res["summary"]["P1"]), (0, 0), res["issues"])


class AttachmentTest(unittest.TestCase):
    """附件另面编排（7.3.7）。"""

    def setUp(self):
        self.path, self.report = build_text(ATTACH_TEXT, "v11-附件.docx", meta=RED_META)
        self.doc = Document(str(self.path))

    def test_structure(self):
        texts = [p.text for p in self.doc.paragraphs]
        self.assertIn("附件：1. 冬季消防安全检查表", texts)  # 没写附件说明时按附件页生成
        self.assertIn("2. 整改通知书（样式）", texts)
        labels = [p for p in self.doc.paragraphs if p.text in ("附件1", "附件2")]
        self.assertEqual([p.text for p in labels], ["附件1", "附件2"])
        for p in labels:
            self.assertIsNotNone(p._p.pPr.find(qn("w:pageBreakBefore")))
            self.assertEqual(p.runs[0].font.size.pt, 16)
            self.assertEqual(p.runs[0].font.name or p.runs[0]._r.rPr.find(qn("w:rFonts")).get(qn("w:eastAsia")), "黑体")
            self.assertEqual(gw_docx.leading_offset_chars(p, self.doc), 0)
            i = self.doc.paragraphs.index(p) if p in self.doc.paragraphs else [x._p for x in self.doc.paragraphs].index(p._p)
            self.assertEqual(self.doc.paragraphs[i + 1].text, "")  # 第二行空
            self.assertEqual(self.doc.paragraphs[i + 2].paragraph_format.alignment, 1)  # 第三行标题居中
        # 版记在附件之后（文档最后一个块是版记表格的锚点段，版记表格在附件正文之后）
        body = list(self.doc.element.body.iterchildren())
        last_tbl = [i for i, el in enumerate(body) if el.tag == qn("w:tbl")][-1]
        attach2 = [i for i, el in enumerate(body) if el.tag == qn("w:p") and "整改通知书（样式）" in "".join(t.text or "" for t in el.iter(qn("w:t")))][-1]
        self.assertGreater(last_tbl, attach2)
        self.assertEqual(len(self.doc.sections), 1)  # 同一节，页码连续

    def test_check_clean(self):
        result = check_docx.check_document(self.path, layout="红头")
        self.assertEqual((result["summary"]["P0"], result["summary"]["P1"]), (0, 0), result["issues"])

    def _mutate(self, name, fn):
        doc = Document(str(self.path))
        fn(doc)
        out = OUT / name
        doc.save(str(out))
        return check_docx.check_document(out, layout="红头")

    def test_mutations_detected(self):
        # 不另面
        def no_break(d):
            p = para_with_exact(d, "附件1")
            p._p.pPr.remove(p._p.pPr.find(qn("w:pageBreakBefore")))
        self.assertIn(("DOCX-131", "P1"), ids(self._mutate("v11-附件-不分页.docx", no_break)))
        # 标识字体不对
        def bad_font(d):
            p = para_with_exact(d, "附件2")
            gw_docx.format_run(p.runs[0], font="仿宋_GB2312")
        self.assertIn(("DOCX-130", "P1"), ids(self._mutate("v11-附件-字体.docx", bad_font)))
        # 标题不居中
        def left_title(d):
            p = para_with_exact(d, "附件1")
            idx = [x._p for x in d.paragraphs].index(p._p)
            d.paragraphs[idx + 2].paragraph_format.alignment = 0
        self.assertIn(("DOCX-132", "P1"), ids(self._mutate("v11-附件-标题.docx", left_title)))
        # 标识缩进
        def indent(d):
            gw_docx.set_para(para_with_exact(d, "附件1"), align="left", left=2, line_exact=28.9, before_pt=0, after_pt=0)
        self.assertIn(("DOCX-130", "P1"), ids(self._mutate("v11-附件-缩进.docx", indent)))
        # 附件说明与附件页标题不一致
        def mismatch(d):
            p = para_with(d, "附件：1. 冬季消防安全检查表")
            p.runs[0].text = "附件：1. 另一个名称"
        self.assertIn(("DOCX-133", "P1"), ids(self._mutate("v11-附件-说明不符.docx", mismatch)))
        # 附件正文字体
        def body_font(d):
            gw_docx.format_run(para_with(d, "检查消防设施").runs[0], font="宋体")
        self.assertIn(("DOCX-010", "P0"), ids(self._mutate("v11-附件-正文字体.docx", body_font)))

    def test_page_numbers_restart_detected(self):
        def restart(d):
            sect = d.sections[0]._sectPr
            pg = OxmlElement("w:pgNumType")
            pg.set(qn("w:start"), "1")
            sect.append(pg)
            new = d.add_section()
            new._sectPr.append(OxmlElement("w:pgNumType"))
            new._sectPr.find(qn("w:pgNumType")).set(qn("w:start"), "1")
        self.assertIn(("DOCX-134", "P1"), ids(self._mutate("v11-附件-页码重起.docx", restart)))

    def test_explicit_description_mismatch_warns(self):
        text = ATTACH_TEXT.replace("请认真贯彻落实。", "请认真贯彻落实。\n\n附件：1. 别的名称\n2. 整改通知书（样式）")
        _, report = build_text(text, "v11-附件-说明不符生成.docx", meta=RED_META)
        self.assertTrue(any("附件说明与附件页不对应" in w for w in report["warnings"]), report["warnings"])

    def test_single_attachment_no_number_and_missing_title(self):
        path, report = build_text(NOTICE + "\n\n【附件】\n\n内容", "v11-附件-单个.docx", meta=RED_META)
        self.assertTrue(any("附件标题" in s for s in report["placeholders"]), report["placeholders"])
        self.assertIn("附件", [p.text for p in Document(str(path)).paragraphs])
        self.assertFalse(has(check_docx.check_document(path, layout="红头"), "DOCX-130"))

    def test_fix_formats_attachment_pages(self):
        dst = OUT / "v11-附件-修复.docx"
        if dst.exists():
            dst.unlink()
        before = [p.text for p in self.doc.paragraphs]
        report = fix_docx.fix_document(self.path, dst, installed=NO_FONTS, keep_text=True)
        self.assertEqual([p.text for p in Document(str(dst)).paragraphs], before)
        self.assertTrue(any("附件另面编排" in n for n in report["notes"]))
        res = check_docx.check_document(dst, layout="红头")
        self.assertEqual(res["summary"]["P0"], 0, res["issues"])


def para_with_exact(doc, text):
    return [p for p in doc.paragraphs if p.text == text][0]


class JointTest(unittest.TestCase):
    """联合行文版头（7.2.4）和联署署名（7.3.5.1）。"""

    META = dict(RED_META, 发文机关标志=JOINT, 发文机关署名=JOINT, 印章="是", 联合行文="是")

    def setUp(self):
        self.path, self.report = build_text(NOTICE, "v11-联合行文.docx", meta=self.META)
        self.doc = Document(str(self.path))

    def test_header_table(self):
        heads = [t for t in self.doc.tables if gw_docx.table_mark(t._tbl) == gw_docx.JOINT_HEAD_MARK]
        self.assertEqual(len(heads), 1)
        left, right = heads[0].rows[0].cells
        self.assertEqual([p.text for p in left.paragraphs], JOINT)  # 主办机关在前
        self.assertEqual(right.text, "文件")
        self.assertEqual(right._tc.tcPr.find(qn("w:vAlign")).get(qn("w:val")), "center")  # 上下居中
        for p in left.paragraphs + right.paragraphs:
            self.assertEqual(str(p.runs[0].font.color.rgb), "FF0000")
            self.assertEqual(p.paragraph_format.alignment, 1)
        self.assertEqual(len(gw_docx.body_tables(self.doc)), 0)  # 不当作正文表格
        # 发文字号仍在版头红线上，标题在其后
        texts = [p.text for p in self.doc.paragraphs if p.text.strip()]
        self.assertEqual(texts[0], "某应急发〔2026〕21号")

    def test_header_fits_width(self):
        heads = [t for t in self.doc.tables if gw_docx.table_mark(t._tbl) == gw_docx.JOINT_HEAD_MARK]
        grid = [int(c.get(qn("w:w"))) for c in heads[0]._tbl.iter(qn("w:gridCol"))]
        self.assertLessEqual(sum(grid), gw_docx.Ctx(gw_docx.load_params()).text_width_twips)
        # 最长名称（7 个字）不回行：列宽不小于 字数×字号
        size, left_w, _ = gw_docx.joint_header_layout(gw_docx.Ctx(gw_docx.load_params()), JOINT)
        self.assertGreaterEqual(left_w, 7 * size * 20)

    def test_signatures_table(self):
        seals = [t for t in self.doc.tables if gw_docx.table_mark(t._tbl) == gw_docx.JOINT_SEAL_MARK]
        self.assertEqual(len(seals), 1)
        cells = [c.text for c in seals[0].rows[0].cells]
        self.assertEqual(cells, JOINT)
        width = sum(int(c.get(qn("w:w"))) for c in seals[0]._tbl.iter(qn("w:gridCol")))
        self.assertLessEqual(width, 8845)
        self.assertGreaterEqual(min(int(c.get(qn("w:w"))) for c in seals[0]._tbl.iter(qn("w:gridCol"))),
                                gw_docx.mm_to_twips(42))
        date = para_with(self.doc, "2026年10月28日")
        self.assertAlmostEqual(28 - gw_docx.para_left_chars(date, self.doc) - 7.0, 4, delta=0.1)  # 日期右空四字

    def test_four_organs_wrap_rows(self):
        names = ["甲单位", "乙单位", "丙单位", "丁单位"]
        path, _ = build_text(NOTICE, "v11-联合行文-四个.docx", meta=dict(self.META, 发文机关标志=names, 发文机关署名=names))
        seals = [t for t in Document(str(path)).tables if gw_docx.table_mark(t._tbl) == gw_docx.JOINT_SEAL_MARK][0]
        self.assertEqual(len(seals.rows), 2)  # 每行最多 3 个，印章互不相交
        self.assertEqual([c.text for r in seals.rows for c in r.cells], names + [""] * 2)

    def test_check_clean(self):
        result = check_docx.check_document(self.path, layout="红头")
        self.assertEqual((result["summary"]["P0"], result["summary"]["P1"]), (0, 0), result["issues"])
        self.assertFalse(has(check_docx.check_document(self.path), "DOCX-102"))  # 版头表格不被当作正文表格
        self.assertEqual(check_docx.check_document(self.path)["layout"], "红头")  # 版头在表格里，自动判断仍是红头

    def _mutate(self, name, fn):
        doc = Document(str(self.path))
        fn(doc)
        out = OUT / name
        doc.save(str(out))
        return check_docx.check_document(out, layout="红头")

    def test_mutations_detected(self):
        def top_align(d):
            t = [x for x in d.tables if gw_docx.table_mark(x._tbl) == gw_docx.JOINT_HEAD_MARK][0]
            t.rows[0].cells[1]._tc.tcPr.find(qn("w:vAlign")).set(qn("w:val"), "top")
        self.assertIn(("DOCX-141", "P1"), ids(self._mutate("v11-联合-顶端.docx", top_align)))

        def black(d):
            t = [x for x in d.tables if gw_docx.table_mark(x._tbl) == gw_docx.JOINT_HEAD_MARK][0]
            t.rows[0].cells[0].paragraphs[0].runs[0].font.color.rgb = __import__("docx.shared", fromlist=["RGBColor"]).RGBColor(0, 0, 0)
        self.assertIn(("DOCX-140", "P1"), ids(self._mutate("v11-联合-黑色.docx", black)))

        def one_organ(d):
            t = [x for x in d.tables if gw_docx.table_mark(x._tbl) == gw_docx.JOINT_HEAD_MARK][0]
            p = t.rows[0].cells[0].paragraphs[1]
            p._p.getparent().remove(p._p)
        self.assertIn(("DOCX-140", "P1"), ids(self._mutate("v11-联合-单机关.docx", one_organ)))

        def narrow(d):
            t = [x for x in d.tables if gw_docx.table_mark(x._tbl) == gw_docx.JOINT_SEAL_MARK][0]
            for gc in t._tbl.iter(qn("w:gridCol")):
                gc.set(qn("w:w"), "2000")
        self.assertIn(("DOCX-143", "P1"), ids(self._mutate("v11-联合-窄格.docx", narrow)))

        def wide(d):
            t = [x for x in d.tables if gw_docx.table_mark(x._tbl) == gw_docx.JOINT_SEAL_MARK][0]
            for gc in t._tbl.iter(qn("w:gridCol")):
                gc.set(qn("w:w"), "5000")
        self.assertIn(("DOCX-144", "P1"), ids(self._mutate("v11-联合-超宽.docx", wide)))

    def test_single_organ_unchanged_and_other_layout_warns(self):
        path, _ = build_text(NOTICE, "v11-单机关.docx", meta=RED_META)
        self.assertEqual([t for t in Document(str(path)).tables if gw_docx.is_frame_table(t._tbl)], [])
        _, report = build_text(NOTICE, "v11-联合-信函.docx", meta=dict(RED_META, 版式="信函", 发文机关标志=JOINT))
        self.assertTrue(any("只用主办机关标志" in w for w in report["warnings"]), report["warnings"])

    def test_fix_keeps_frame_tables(self):
        dst = OUT / "v11-联合-修复.docx"
        if dst.exists():
            dst.unlink()
        report = fix_docx.fix_document(self.path, dst, installed=NO_FONTS, keep_text=True)
        self.assertEqual(report["layout"], "红头")
        doc = Document(str(dst))
        self.assertEqual(len([t for t in doc.tables if gw_docx.is_frame_table(t._tbl)]), 2)
        self.assertFalse(any("正文中的表格" in s for s in report["untouched"]), report["untouched"])


class OverprintTest(unittest.TestCase):
    """预印红头纸套打：不输出红色版头，按预印纸留出版头区域，发文字号位置与正式版头一致。"""

    def build_pair(self, **extra):
        formal, _ = build_text(REQUEST, "v11-正式红头.docx", meta=UP_META)
        over_meta = dict(UP_META, **extra)
        doc, report = build_docx.build(REQUEST.split("\n"), over_meta, installed=NO_FONTS, overprint=True)
        path = OUT / "v11-套打.docx"
        if path.exists():
            path.unlink()
        doc.save(str(path))
        return formal, path, report

    def test_no_red_and_same_position(self):
        formal, over, report = self.build_pair()
        doc = Document(str(over))
        xml = doc.element.xml
        self.assertNotIn('w:color w:val="FF0000"', xml)  # 没有红色文字或红线
        self.assertNotIn("某理工大学文件", "".join(p.text for p in doc.paragraphs))
        self.assertAlmostEqual(docno_top_mm(over), docno_top_mm(formal), delta=0.5)
        # 标题等后续要素的位置也一致：红线占的空间被保留
        def title_top(path):
            d = Document(str(path))
            total = 0.0
            for p in d.paragraphs:
                sp = p._p.pPr.find(qn("w:spacing")) if p._p.pPr is not None else None
                before = int(sp.get(qn("w:before")) or 0) / 20 if sp is not None else 0.0
                after = int(sp.get(qn("w:after")) or 0) / 20 if sp is not None else 0.0
                line = int(sp.get(qn("w:line")) or 0) / 20 if sp is not None else 0.0
                if "关于申请增加" in p.text:
                    return total + before
                total += before + line + after
                bd = p._p.pPr.find(qn("w:pBdr")) if p._p.pPr is not None else None
                if bd is not None:
                    total += 11 + 1.5  # 正式版头红线：线与文字间距加线宽
        self.assertAlmostEqual(title_top(over), title_top(formal), delta=1.5)
        self.assertTrue(report["overprint_on"])
        self.assertTrue(any("预印红头纸套打" in n for n in report["notes"]))

    def test_check_overprint(self):
        _, over, _ = self.build_pair()
        result = check_docx.check_document(over, overprint=True)
        self.assertEqual(result["layout"], "红头")
        self.assertEqual((result["summary"]["P0"], result["summary"]["P1"]), (0, 0), result["issues"])

    def test_formal_doc_flagged_as_overprint(self):
        formal, _, _ = self.build_pair()
        self.assertIn(("DOCX-150", "P1"), ids(check_docx.check_document(formal, overprint=True)))

    def test_too_small_reserve_flagged(self):
        _, over, _ = self.build_pair()
        doc = Document(str(over))
        for p in doc.paragraphs:  # 把版头留白压到很小
            sp = p._p.pPr.find(qn("w:spacing")) if p._p.pPr is not None else None
            if sp is not None and int(sp.get(qn("w:before")) or 0) > 1000:
                sp.set(qn("w:before"), "200")
        out = OUT / "v11-套打-留白不足.docx"
        doc.save(str(out))
        self.assertIn(("DOCX-151", "P1"), ids(check_docx.check_document(out, overprint=True)))

    def test_custom_header_height(self):
        _, over, report = self.build_pair(套打版头高度=60)
        self.assertAlmostEqual(docno_top_mm(over), 60, delta=0.6)
        self.assertTrue(report["overprint"]["custom"])
        self.assertEqual(check_docx.check_document(over, overprint=True)["summary"]["P1"], 0)

    def test_meta_flag_and_cli(self):
        path, report = build_text(REQUEST, "v11-套打-版式信息.docx", meta=dict(UP_META, 套打="是"))
        self.assertTrue(report["overprint_on"])
        folder = OUT / "套打命令行"
        folder.mkdir(parents=True, exist_ok=True)
        for old in folder.glob("*"):
            old.unlink()
        (folder / "正文.txt").write_text(REQUEST, encoding="utf-8")
        (folder / "meta.json").write_text(json.dumps(UP_META, ensure_ascii=False), encoding="utf-8")
        buf = io.StringIO()
        with redirect_stdout(buf):
            self.assertEqual(build_docx.main([str(folder / "正文.txt"), "--meta", str(folder / "meta.json"), "--overprint",
                                              "--json"]), 0)
        rep = json.loads(buf.getvalue())
        self.assertTrue(rep["overprint_on"])
        self.assertEqual(rep["check"]["summary"]["P1"], 0, rep["check"]["issues"])

    def test_only_red_layout_and_joint(self):
        _, report = build_docx.build(NOTICE.split("\n"), {"版式": "白头"}, installed=NO_FONTS, overprint=True)
        self.assertFalse(report["overprint_on"])
        self.assertTrue(any("套打" in w and "只支持红头" in w for w in report["warnings"]), report["warnings"])
        # 联合行文也可套打：不画表格，按名称行数留白
        doc, rep = build_docx.build(NOTICE.split("\n"), dict(RED_META, 发文机关标志=JOINT), installed=NO_FONTS, overprint=True)
        self.assertEqual([t for t in doc.tables if gw_docx.is_frame_table(t._tbl)], [])
        self.assertTrue(rep["overprint_on"])

    def test_layout_cli_choices(self):
        buf = io.StringIO()
        folder = OUT / "命令命令行"
        folder.mkdir(parents=True, exist_ok=True)
        for old in folder.glob("*"):
            old.unlink()
        (folder / "正文.txt").write_text(ORDER, encoding="utf-8")
        (folder / "meta.json").write_text(json.dumps({"发文机关标志": "某市人民政府令"}, ensure_ascii=False), encoding="utf-8")
        with redirect_stdout(buf):
            self.assertEqual(build_docx.main([str(folder / "正文.txt"), "--meta", str(folder / "meta.json"), "--layout", "令",
                                              "--json"]), 0)
        rep = json.loads(buf.getvalue())
        self.assertEqual(rep["layout"], "命令")
        self.assertEqual(rep["check"]["summary"]["P1"], 0, rep["check"]["issues"])


if __name__ == "__main__":
    unittest.main()
