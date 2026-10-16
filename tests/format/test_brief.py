# -*- coding: utf-8 -*-
"""简报版式（无国家标准格式，按机关常见做法）：报头、各篇、报尾的生成、回读检查、修复和缺项写【待补】。"""
import io
import shutil
import unittest
from contextlib import redirect_stdout

from docx import Document
from docx.oxml.ns import qn

from tests.format.helpers import FILES, NO_FONTS, OUT, build_sample, build_text, check_docx, gw_layout, read_meta

import build_docx  # noqa: E402  helpers 已把 scripts 加入 sys.path
import fix_docx  # noqa: E402
import preview_docx  # noqa: E402

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def fonts_of(p):
    return {r._r.rPr.find(qn("w:rFonts")).get(qn("w:eastAsia")) for r in p.runs if r.text.strip()}


def sizes_of(p):
    return {r.font.size.pt for r in p.runs if r.text.strip()}


def find(doc, startswith):
    for p in doc.paragraphs:
        if p.text.strip().startswith(startswith):
            return p
    raise AssertionError(f"找不到段落：{startswith}")


def ids(result):
    return {(i["id"], i["level"]) for i in result["issues"]}


def footer_table(doc):
    return doc.tables[-1]


def table_lines(table):
    return [p.text for c in table.rows[0].cells for p in c.paragraphs]


class BriefBuildTest(unittest.TestCase):
    """虚构简报：两篇稿件、有报尾。"""

    @classmethod
    def setUpClass(cls):
        cls.path, cls.report = build_sample("简报", "样稿-简报.docx")
        cls.doc = Document(str(cls.path))

    def test_layout_and_check(self):
        self.assertEqual(self.report["layout"], "简报")
        self.assertEqual(self.report["placeholders"], [])
        result = check_docx.check_document(self.path)  # 不指定版式，自动识别为简报
        self.assertEqual(result["layout"], "简报")
        self.assertEqual(result["doc_type"], "简报")
        self.assertEqual(result["summary"]["P0"], 0, result["issues"])
        self.assertEqual(result["summary"]["P1"], 0, result["issues"])
        self.assertEqual(result["summary"]["P2"], 0, result["issues"])
        self.assertEqual(check_docx.check_document(self.path, layout="简报")["summary"]["P0"], 0)

    def test_page_setup_same_as_others(self):
        sec = self.doc.sections[0]
        self.assertEqual(int(sec._sectPr.find(qn("w:pgSz")).get(qn("w:w"))), 11906)
        self.assertEqual(sec._sectPr.find(qn("w:docGrid")).get(qn("w:linePitch")), "579")
        self.assertTrue(any("PAGE" in (t.text or "") for t in sec.footer._element.iter(W + "instrText")))

    def test_masthead_name(self):
        name = self.doc.paragraphs[0]
        self.assertEqual(name.text, "某区农业农村局工作简报")
        self.assertEqual(str(name.runs[0].font.color.rgb), "FF0000")
        self.assertEqual(name.alignment, 1)
        self.assertEqual(fonts_of(name), {"方正小标宋简体"})
        self.assertGreaterEqual(min(sizes_of(name)), 22.0)  # 大字号，不小于标题二号

    def test_masthead_issue_and_editor(self):
        issue = self.doc.paragraphs[1]
        self.assertEqual(issue.text, "第8期（总第56期）")  # 数字写入版式信息即可，“第”“期”由脚本补
        self.assertEqual(issue.alignment, 1)
        self.assertEqual((fonts_of(issue), sizes_of(issue)), ({"仿宋_GB2312"}, {16.0}))
        editor = self.doc.paragraphs[2]
        self.assertEqual(editor.text, "某区农业农村局办公室编\t2026年11月6日")  # 编印单位左、日期右同一行；日期写法已规范
        tabs = editor._p.pPr.find(qn("w:tabs")).findall(qn("w:tab"))
        self.assertEqual([(t.get(qn("w:val")), t.get(qn("w:pos"))) for t in tabs], [("right", "8845")])  # 右对齐到版心右边缘
        rule = editor._p.pPr.find(qn("w:pBdr")).find(qn("w:bottom"))  # 红色分隔线
        self.assertEqual((rule.get(qn("w:val")), rule.get(qn("w:color"))), ("single", "FF0000"))
        self.assertEqual((fonts_of(editor), sizes_of(editor)), ({"仿宋_GB2312"}, {16.0}))

    def test_articles_follow_material_layout(self):
        t1 = find(self.doc, "某区推行春耕物资线上预订")
        t2 = find(self.doc, "某镇建立田间课堂")
        for t in (t1, t2):
            self.assertEqual((fonts_of(t), sizes_of(t), t.alignment), ({"方正小标宋简体"}, {22.0}, 1))
            self.assertIsNotNone(t._p.pPr.find(qn("w:keepNext")))  # 标题不落在页末
        self.assertEqual(t1.text, "某区推行春耕物资线上预订\n农户领料等候时间缩短")  # 标题中间的全角空格处分行
        body = find(self.doc, "某区农业服务中心今年春耕前")
        self.assertEqual((fonts_of(body), sizes_of(body)), ({"仿宋_GB2312"}, {16.0}))
        self.assertEqual(body._p.pPr.find(qn("w:ind")).get(qn("w:firstLineChars")), "200")
        h1 = find(self.doc, "一、课程安排")
        self.assertEqual(fonts_of(h1), {"黑体"})
        note = find(self.doc, "（某镇农业服务中心供稿）")  # 标题下的括号注按题注排
        self.assertEqual((fonts_of(note), note.alignment), ({"楷体_GB2312"}, 1))
        # 篇间空一行：第二篇标题前是空段
        texts = [p.text for p in self.doc.paragraphs]
        self.assertEqual(texts[texts.index(t2.text) - 1], "")

    def test_footer(self):
        tbl = footer_table(self.doc)
        self.assertEqual(table_lines(tbl), ["报：区委办公室，区政府办公室", "送：区委农村工作部",
                                            "发：各镇农业服务中心，局机关各科室", "（共印30份）"])
        pr = tbl._tbl.tblPr
        self.assertEqual(pr.find(qn("w:tblpPr")).get(qn("w:tblpYSpec")), "bottom")  # 贴版心底部
        borders = pr.find(qn("w:tblBorders"))
        self.assertEqual(borders.find(qn("w:top")).get(qn("w:val")), "single")  # 只有一条分隔线
        for side in ("bottom", "insideH", "left", "right", "insideV"):
            self.assertEqual(borders.find(qn(f"w:{side}")).get(qn("w:val")), "nil", side)
        for p in tbl.rows[0].cells[0].paragraphs:
            self.assertEqual((fonts_of(p), sizes_of(p)), ({"仿宋_GB2312"}, {14.0}))
        # “报：”两字宽，换行时悬挂缩进与标签等宽
        ind = tbl.rows[0].cells[0].paragraphs[0]._p.pPr.find(qn("w:ind"))
        self.assertEqual((ind.get(qn("w:hanging")), ind.get(qn("w:left"))), ("560", "840"))

    def test_report_lists_recognition(self):
        text = build_docx.render_text(dict(self.report, output="x.docx"))
        for word in ("简报名称：某区农业农村局工作简报", "期号：第8期", "总期号：总第56期", "篇目："):
            self.assertIn(word, text)


class BriefMissingTest(unittest.TestCase):
    def test_missing_issue_no_leaves_placeholder(self):
        meta = read_meta("简报-版式信息.json")
        del meta["期号"], meta["总期号"]
        doc, report = build_docx.build(gw_layout.read_body_text(FILES / "简报-正文.txt")[0], meta, installed=NO_FONTS)
        path = OUT / "样稿-简报-缺期号.docx"
        doc.save(str(path))
        self.assertEqual(Document(str(path)).paragraphs[1].text, "【待补：期号】")
        self.assertEqual(len(report["placeholders"]), 1)
        self.assertIn("期号", report["placeholders"][0])
        result = check_docx.check_document(path)
        self.assertEqual(result["summary"]["P0"], 0, result["issues"])
        self.assertEqual(result["summary"]["markers"]["待补"], 1)  # 待补标记单独计数、提醒补齐
        self.assertIn(("DOCX-080", "P2"), ids(result))

    def test_nothing_given_all_masthead_placeholders_and_no_footer(self):
        """只给正文、指定简报版式：报头四项都写【待补】，不编造；没有报送范围和印数就不排报尾。"""
        text = "某街道建立网格员“每日一巡”制度\n\n网格员每天巡查一遍，发现问题当场登记。"
        path, report = build_text(text, "样稿-简报-只有正文.docx", layout="简报")
        doc = Document(str(path))
        self.assertEqual([p.text for p in doc.paragraphs[:3]],
                         ["【待补：简报名称】", "【待补：期号】", "【待补：编印单位】\t【待补：印发日期】"])
        self.assertEqual(len(report["placeholders"]), 4)
        self.assertEqual(len(doc.tables), 0)  # 没有报尾
        self.assertTrue(any("未排报尾" in n for n in report["notes"]))
        result = check_docx.check_document(path)
        self.assertEqual(result["summary"]["P0"], 0, result["issues"])
        self.assertEqual(result["summary"]["markers"]["待补"], 4)

    def test_footer_only_what_is_given(self):
        text = "某街道建立网格员“每日一巡”制度\n\n网格员每天巡查一遍。"
        meta = {"版式": "简报", "简报名称": "某街道工作简报", "期号": "第3期", "编印单位": "某街道办公室编",
                "印发日期": "2026年11月6日", "抄送": ["区信访局", "区民政局"]}
        path, report = build_text(text, "样稿-简报-只有抄送.docx", meta=meta)
        doc = Document(str(path))
        self.assertEqual(table_lines(footer_table(doc)), ["抄送：区信访局，区民政局"])
        self.assertEqual(doc.paragraphs[2].text, "某街道办公室编\t2026年11月6日")  # 已写“编”的不重复加
        result = check_docx.check_document(path)
        self.assertEqual((result["summary"]["P0"], result["summary"]["P1"]), (0, 0), result["issues"])  # 抄送不要求句号
        self.assertEqual(report["placeholders"], [])


class BriefTextFrameTest(unittest.TestCase):
    """gongwen-brief 的纯文本模板（报头、各篇、报尾都写在正文里）可以直接排版。"""

    TEMPLATE = ("某局工作简报\n第12期（总第68期）\n某局办公室编　　　　　　　　2026年10月20日\n────────\n\n"
                "某局推行“一表通”　基层报表数量减少四成\n\n某局把原来的多张报表合并为一张。\n\n"
                "某街道建立网格员“每日一巡”制度\n（某街道供稿）\n\n网格员每天巡查一遍。\n────────\n"
                "报：区委办公室，区政府办公室\n送：区委组织部\n（共印30份）")

    def test_template_recognized(self):
        parsed = gw_layout.parse_brief(self.TEMPLATE.split("\n"))
        f = parsed.fields
        self.assertEqual((f["简报名称"], f["期号"], f["总期号"]), ("某局工作简报", "第12期", "总第68期"))
        self.assertEqual((f["编印单位"], f["印发日期"]), ("某局办公室编", "2026年10月20日"))
        self.assertEqual((f["报"], f["送"], f["印数"]), ("区委办公室，区政府办公室", "区委组织部", "30"))
        self.assertEqual(len(f["篇目"]), 2)

    def test_template_builds_without_meta(self):
        path, report = build_text(self.TEMPLATE, "样稿-简报-模板.docx")
        self.assertEqual(report["layout"], "简报")
        self.assertEqual(report["placeholders"], [])
        doc = Document(str(path))
        self.assertEqual([p.text for p in doc.paragraphs[:3]],
                         ["某局工作简报", "第12期（总第68期）", "某局办公室编\t2026年10月20日"])
        self.assertEqual(table_lines(footer_table(doc)), ["报：区委办公室，区政府办公室", "送：区委组织部", "（共印30份）"])
        # 报头、报尾、分隔线行不进入正文
        self.assertNotIn("────────", "".join(p.text for p in doc.paragraphs))
        result = check_docx.check_document(path)
        self.assertEqual((result["summary"]["P0"], result["summary"]["P1"]), (0, 0), result["issues"])

    def test_meta_wins_and_conflict_is_reported(self):
        path, report = build_text(self.TEMPLATE, "样稿-简报-冲突.docx", meta={"期号": "13"})
        self.assertEqual(Document(str(path)).paragraphs[1].text, "第13期（总第68期）")
        self.assertTrue(any("期号" in w and "不一致" in w for w in report["warnings"]))

    def test_other_layouts_not_taken_for_brief(self):
        minutes = "某理工大学校长办公会纪要\n第18期\n\n会议研究了三件事。"
        self.assertFalse(gw_layout.has_brief_header(minutes.split("\n")))
        self.assertFalse(gw_layout.has_brief_header("某工作总结\n\n今年开展了三项工作。".split("\n")))

    def test_ignored_fields_warn(self):
        _p, report = build_text("某街道建立制度\n\n网格员每天巡查一遍。", "样稿-简报-忽略.docx",
                                meta={"版式": "简报", "发文字号": "某发〔2026〕1号", "印章": "是"})
        self.assertTrue(any("简报版式不排" in w and "发文字号" in w for w in report["warnings"]))


class BriefCheckFixTest(unittest.TestCase):
    """版式被改乱的简报：check 报出，fix 修好；报头、报尾不当作错误。"""

    @classmethod
    def setUpClass(cls):
        cls.good, _ = build_sample("简报", "样稿-简报-标本.docx")

    def mutate(self, name, fn):
        doc = Document(str(self.good))
        fn(doc)
        path = OUT / name
        doc.save(str(path))
        return path

    def test_masthead_problems(self):
        def fn(doc):
            for r in doc.paragraphs[0].runs:
                r.font.color.rgb = None
            doc.paragraphs[0].alignment = 0
            bdr = doc.paragraphs[2]._p.pPr.find(qn("w:pBdr"))
            doc.paragraphs[2]._p.pPr.remove(bdr)
            doc.tables[-1]._tbl.tblPr.find(qn("w:tblBorders")).find(qn("w:top")).set(qn("w:val"), "nil")
        got = ids(check_docx.check_document(self.mutate("标本-简报-报头.docx", fn)))
        for item in (("DOCX-111", "P1"), ("DOCX-112", "P1"), ("DOCX-113", "P1")):
            self.assertIn(item, got)

    def test_missing_masthead_elements(self):
        def fn(doc):
            doc.paragraphs[1]._p.getparent().remove(doc.paragraphs[1]._p)  # 删去期号行
        result = check_docx.check_document(self.mutate("标本-简报-缺期号行.docx", fn), layout="简报")
        self.assertIn(("DOCX-110", "P1"), ids(result))

    def test_body_font_still_checked(self):
        def fn(doc):
            body = find(doc, "某区农业服务中心今年春耕前")
            for r in body.runs:
                r.font.size = None
                r._r.rPr.find(qn("w:rFonts")).set(qn("w:eastAsia"), "宋体")
        self.assertIn(("DOCX-010", "P0"), ids(check_docx.check_document(self.mutate("标本-简报-正文字体.docx", fn))))

    def test_fix_keeps_masthead_and_repairs_body(self):
        def fn(doc):
            body = find(doc, "某区农业服务中心今年春耕前")
            for r in body.runs:
                r._r.rPr.find(qn("w:rFonts")).set(qn("w:eastAsia"), "宋体")
            sec = doc.sections[0]
            sec.top_margin = sec.top_margin + 400000
        src = self.mutate("标本-简报-待修复.docx", fn)
        self.assertGreater(check_docx.check_document(src)["summary"]["P0"], 0)
        dst = OUT / "标本-简报-已修复.docx"
        if dst.exists():
            dst.unlink()
        report = fix_docx.fix_document(src, dst, installed=NO_FONTS)
        self.assertEqual(report["layout"], "简报")
        after = check_docx.check_document(dst, layout="简报")
        self.assertEqual((after["summary"]["P0"], after["summary"]["P1"]), (0, 0), after["issues"])
        old, new = Document(str(src)), Document(str(dst))
        self.assertEqual([p.text for p in old.paragraphs], [p.text for p in new.paragraphs])  # 文字不变
        self.assertEqual(table_lines(footer_table(old)), table_lines(footer_table(new)))
        self.assertEqual(str(new.paragraphs[0].runs[0].font.color.rgb), "FF0000")  # 报头保持原样

    def test_fix_default_layout_detection(self):
        dst = OUT / "标本-简报-自动识别.docx"
        if dst.exists():
            dst.unlink()
        self.assertEqual(fix_docx.fix_document(self.good, dst, installed=NO_FONTS)["layout"], "简报")


class BriefPreviewTest(unittest.TestCase):
    @unittest.skipUnless(preview_docx.find_soffice(), "本机没有 LibreOffice")
    def test_preview_renders_pages(self):
        """用 preview_docx 渲染样稿；预览图留在 _out/render/ 供目检（报头居中、分隔线、报尾在末页底部）。"""
        src = OUT / "样稿-简报.docx"
        if not src.exists():
            build_sample("简报", "样稿-简报.docx")
        out = OUT / "render" / "简报样稿-预览"
        shutil.rmtree(out, ignore_errors=True)
        with redirect_stdout(io.StringIO()):
            code = preview_docx.main([str(src), "-o", str(out), "--dpi", "90"])
        self.assertEqual(code, 0)
        pdf = out / "样稿-简报.pdf"
        self.assertTrue(pdf.exists())
        try:
            import pymupdf
        except ImportError:
            return
        with pymupdf.open(str(pdf)) as pages:
            self.assertEqual(len(pages), 2)
            first, last = pages[0].get_text(), pages[-1].get_text()
        first, last = "".join(first.split()), "".join(last.split())  # PDF 抽字时中英文之间可能带空格
        self.assertIn("某区农业农村局工作简报", first)
        self.assertIn("第8期（总第56期）", first)
        self.assertIn("共印30份", last)  # 报尾在末页
        self.assertNotIn("共印30份", first)


if __name__ == "__main__":
    unittest.main()
