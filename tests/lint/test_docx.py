# -*- coding: utf-8 -*-
"""docx 输入：用 python-docx 现场生成临时文件；python-docx 缺失时返回 2。"""
from __future__ import annotations

import io
import sys
import tempfile
import unittest
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path
from unittest import mock

from tests.lint.helpers import gl

try:
    import docx  # type: ignore
except ImportError:  # pragma: no cover
    docx = None


@unittest.skipIf(docx is None, "未安装 python-docx")
class DocxInputTest(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.path = Path(self._tmp.name) / "成稿.docx"
        document = docx.Document()
        document.add_paragraph("某县农业农村局关于开展农村人居环境整治的通知")      # 段 1
        document.add_paragraph("各乡镇人民政府：")                                  # 段 2
        document.add_paragraph("")                                                  # 段 3（空段）
        document.add_paragraph("根据某农发[2026]第05号文件，各乡镇要按时报送。")   # 段 4
        table = document.add_table(rows=1, cols=2)                                  # 段 5、6
        table.cell(0, 0).text = "1、序号"
        table.cell(0, 1).text = "患者王建国，男，56岁"
        document.add_paragraph("【起草说明】")                                      # 段 7
        document.add_paragraph("以上内容仅供参考。")
        document.save(str(self.path))

    def tearDown(self):
        self._tmp.cleanup()

    def test_docx_lint(self):
        result = gl.lint_file(self.path, None, "all")
        found = {(i["id"], i["line"]) for i in result["issues"]}
        self.assertEqual(result["doc_type"], "通知")
        self.assertIn(("FMT-001", 4), found)
        self.assertIn(("PRI-009", 6), found, "表格单元格也要检查隐私")
        self.assertNotIn("PUN-003", {i["id"] for i in result["issues"]}, "表格内编号不按层次序数检查")
        self.assertNotIn("DEF-001", {i["id"] for i in result["issues"]}, "【起草说明】之后不检查")
        self.assertTrue(any("第 7 段" in n for n in result["notes"]))

    def test_docx_cli_json(self):
        out = io.StringIO()
        with redirect_stdout(out):
            code = gl.main([str(self.path), "--json"])
        self.assertEqual(code, 0)
        self.assertIn("FMT-001", out.getvalue())

    def test_missing_python_docx_returns_2(self):
        err = io.StringIO()
        with mock.patch.dict(sys.modules, {"docx": None, "docx.oxml.ns": None, "docx.text.paragraph": None}):
            with redirect_stderr(err), redirect_stdout(io.StringIO()):
                code = gl.main([str(self.path)])
        self.assertEqual(code, 2)
        self.assertIn("python-docx", err.getvalue())

    def test_broken_docx_returns_2(self):
        bad = Path(self._tmp.name) / "坏文件.docx"
        bad.write_bytes(b"not a zip")
        with redirect_stderr(io.StringIO()), redirect_stdout(io.StringIO()):
            self.assertEqual(gl.main([str(bad)]), 2)


if __name__ == "__main__":
    unittest.main()
