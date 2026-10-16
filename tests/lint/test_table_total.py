# -*- coding: utf-8 -*-
"""表格合计、小计核算（NUM-023，完整性审查 G06，顶层设计 §4.15）。

只核算同一张表格里首列为“合计”“小计”“总计”“总数”的行；百分比、比率、增长率列不求和；
空单元格、【待补】单元格所在列跳过；有小计行时按分组核算。
"""
from __future__ import annotations

import tempfile
import unittest
from decimal import Decimal
from pathlib import Path

from tests.lint.helpers import RULES_PATH, gl, ids, issues_of, load_json, run

try:
    import docx  # type: ignore
except ImportError:  # pragma: no cover
    docx = None


def table(*rows: str) -> str:
    """把 “表头|甲|乙” 形式的行拼成 Markdown 表格（第一行为表头），前后各带一段正文。"""
    head, *body = rows
    cols = head.count("|") + 1
    lines = ["| " + " | ".join(head.split("|")) + " |", "|" + "---|" * cols]
    lines += ["| " + " | ".join(row.split("|")) + " |" for row in body]
    return "2026年全院服务量见下表。\n\n" + "\n".join(lines) + "\n\n以上为全年数据。\n"


def found(text: str):
    return issues_of(run(text, None, "number"), "NUM-023")


class TableTotalTest(unittest.TestCase):
    def test_correct_total_ok(self):
        text = table("项目|金额（万元）|人次", "门诊|120|300", "住院|80|200", "合计|200|500")
        self.assertEqual(found(text), [])

    def test_wrong_total_reported_as_p0(self):
        text = table("项目|金额（万元）|人次", "门诊|120|300", "住院|80|200", "合计|210|500")
        items = found(text)
        self.assertEqual(len(items), 1)
        item = items[0]
        self.assertEqual(item["level"], "P0")
        self.assertEqual(item["category"], "数字")
        self.assertIn("金额（万元）", item["message"])
        self.assertIn("210", item["message"])
        self.assertIn("200", item["message"])
        self.assertIn("合计", item["excerpt"])
        # 两列都错就报两处
        text = table("项目|金额（万元）|人次", "门诊|120|300", "住院|80|200", "总计|210|510")
        self.assertEqual(len(found(text)), 2)

    def test_total_label_variants(self):
        for label in ("总计", "总数", "**合计**", "合　计", "合计："):
            with self.subTest(label=label):
                bad = table("项目|人次", "门诊|300", "住院|200", f"{label}|600")
                good = table("项目|人次", "门诊|300", "住院|200", f"{label}|500")
                self.assertEqual(len(found(bad)), 1)
                self.assertEqual(found(good), [])

    def test_total_row_first_or_in_any_position(self):
        text = table("项目|人次", "合计|600", "门诊|300", "住院|200")
        self.assertEqual(len(found(text)), 1)
        text = table("项目|人次", "合计|500", "门诊|300", "住院|200")
        self.assertEqual(found(text), [])

    def test_subtotal_groups(self):
        head = "科室|类别|人次"
        rows = ["内科|门诊|100", "内科|住院|50", "小计||150", "外科|门诊|30", "外科|住院|20", "小计||50"]
        ok = table(head, *rows, "合计||200")
        self.assertEqual(found(ok), [])
        # 小计错
        bad_sub = table(head, *rows[:5], "小计||60", "合计||200")
        items = found(bad_sub)
        self.assertEqual(len(items), 1)
        self.assertIn("小计", items[0]["excerpt"])
        # 合计错（既不等于明细之和，也不等于小计之和）
        bad_total = table(head, *rows, "合计||250")
        items = found(bad_total)
        self.assertEqual(len(items), 1)
        self.assertIn("合计", items[0]["excerpt"])

    def test_total_equals_either_details_or_subtotals(self):
        head = "科室|类别|人次"
        # 明细之和 150+50=200，小计之和 150+60=210：合计写 210（等于小计之和）不再报合计，只报错的小计
        rows = ["内科|门诊|100", "内科|住院|50", "小计||150", "外科|门诊|30", "外科|住院|20", "小计||60"]
        items = found(table(head, *rows, "合计||210"))
        self.assertEqual(len(items), 1)
        self.assertIn("小计", items[0]["excerpt"])
        # 合计写 200（等于全部明细之和）同样算一致
        items = found(table(head, *rows, "合计||200"))
        self.assertEqual(len(items), 1)
        self.assertIn("小计", items[0]["excerpt"])

    def test_subtotal_on_top_of_group(self):
        # 小计行写在分组上方（下面紧跟明细）时，按下方的连续明细核算
        self.assertEqual(found(table("科室|人次", "小计|150", "门诊|100", "住院|50")), [])
        self.assertEqual(len(found(table("科室|人次", "小计|160", "门诊|100", "住院|50"))), 1)

    def test_ratio_columns_not_summed(self):
        text = table("项目|人次|完成率|占比（%）|同比增长|环比|较上年", "门诊|300|80%|60%|5.2%|1.1%|10%",
                     "住院|200|70%|40%|-3.0%|2.2%|20%", "合计|500|75%|100%|3.1%|1.5%|15%")
        self.assertEqual(found(text), [])
        # 数值列照常核算
        text = text.replace("| 合计 | 500 |", "| 合计 | 520 |")
        self.assertEqual(len(found(text)), 1)

    def test_percent_cells_under_neutral_header_skipped(self):
        text = table("项目|满意度", "门诊|96.5%", "住院|93%", "合计|95%")
        self.assertEqual(found(text), [])

    def test_average_and_unit_price_columns_not_summed(self):
        text = table("项目|人次|平均住院日|单价（元）", "内科|300|8.2|50", "外科|200|6.1|80", "合计|500|7.4|130")
        self.assertEqual(found(text), [])

    def test_thousand_separator_and_units(self):
        text = table("项目|预算", "设备购置|1,000万元", "基础改造|234.5万元", "合计|1,234.5万元")
        self.assertEqual(found(text), [])
        text = table("项目|预算", "设备购置|1,000万元", "基础改造|234.5万元", "合计|1,334.5万元")
        items = found(text)
        self.assertEqual(len(items), 1)
        self.assertIn("1,334.5", items[0]["message"] + items[0]["excerpt"])
        text = table("项目|服务量", "门诊|12,000人次", "住院|3,500人次", "合计|15,500人次")
        self.assertEqual(found(text), [])
        text = table("项目|服务量", "门诊|12,000人次", "住院|3,500人次", "合计|15,000人次")
        self.assertEqual(len(found(text)), 1)

    def test_mixed_units_in_column_skipped(self):
        text = table("项目|金额", "设备|1200万元", "耗材|30元", "合计|1230万元")
        self.assertEqual(found(text), [])

    def test_todo_and_blank_cells_skip_column(self):
        text = table("项目|金额（万元）|人次", "门诊|【待补：金额】|300", "住院|80|200", "合计|999|500")
        self.assertEqual(found(text), [])  # 金额列含待补，跳过；人次列正确
        text = table("项目|金额（万元）|人次", "门诊|【待补：金额】|300", "住院|80|200", "合计|999|520")
        items = found(text)
        self.assertEqual(len(items), 1)
        self.assertIn("人次", items[0]["message"])  # 人次列照查
        text = table("项目|金额（万元）", "门诊||", "住院|80", "合计|999")
        self.assertEqual(found(text), [])
        text = table("项目|金额（万元）", "门诊|120", "住院|80", "合计|【待补：合计】")
        self.assertEqual(found(text), [])
        text = table("项目|金额（万元）", "门诊|120", "住院|80", "合计|【待核实：合计】")
        self.assertEqual(found(text), [])

    def test_dash_cell_counts_as_zero(self):
        text = table("项目|人次", "门诊|300", "住院|—", "合计|300")
        self.assertEqual(found(text), [])
        text = table("项目|人次", "门诊|300", "住院|-", "合计|350")
        self.assertEqual(len(found(text)), 1)

    def test_text_cell_in_numeric_column_skips_column(self):
        text = table("项目|人次", "门诊|300", "住院|详见附件", "合计|999")
        self.assertEqual(found(text), [])

    def test_rounding_tolerance(self):
        # 分项修约后相加与合计差末位，不报
        ok = table("项目|金额（万元）", "甲|33.3", "乙|33.3", "丙|33.3", "合计|100.0")
        self.assertEqual(found(ok), [])
        ok = table("项目|金额（万元）", "甲|33.3", "乙|33.3", "丙|33.3", "合计|99.8")
        self.assertEqual(found(ok), [])
        bad = table("项目|金额（万元）", "甲|33.3", "乙|33.3", "丙|33.3", "合计|100.2")
        self.assertEqual(len(found(bad)), 1)
        bad = table("项目|金额（万元）", "甲|33.3", "乙|33.3", "丙|33.3", "合计|99.7")
        self.assertEqual(len(found(bad)), 1)

    def test_rounding_tolerance_integer_amount_vs_count(self):
        # 整数的金额列（万元）可差 1；人次等计数列不能差
        amount = table("项目|金额（万元）", "甲|33", "乙|33", "丙|33", "合计|100")
        self.assertEqual(found(amount), [])
        amount = table("项目|金额（万元）", "甲|33", "乙|33", "丙|33", "合计|101")
        self.assertEqual(len(found(amount)), 1)
        count = table("项目|人次", "甲|33", "乙|33", "丙|33", "合计|100")
        self.assertEqual(len(found(count)), 1)

    def test_tables_judged_independently(self):
        first = table("项目|人次", "门诊|300", "住院|200", "合计|500")
        second = table("项目|人次", "门诊|10", "住院|20", "合计|40")
        items = found(first + "\n" + second)
        self.assertEqual(len(items), 1)
        self.assertIn("40", items[0]["excerpt"])
        # 相邻两张表（没有空行以外的内容隔开）：上一张的明细不并入下一张
        items = found(second + "\n" + first)
        self.assertEqual(len(items), 1)

    def test_table_without_total_row_or_non_table_ignored(self):
        self.assertEqual(found(table("项目|人次", "门诊|300", "住院|200")), [])
        self.assertEqual(found("合计 500 人次，其中门诊 300，住院 100。\n"), [])
        self.assertEqual(found("| 合计 | 500 |\n"), [])  # 没有分隔行，不是表格

    def test_total_without_detail_rows_ignored(self):
        self.assertEqual(found(table("项目|人次", "合计|500")), [])

    def test_header_label_not_taken_as_total(self):
        # 首列为“合计”的表头行（极端写法）不当作合计行
        text = table("合计|人次", "门诊|300", "住院|200")
        self.assertEqual(found(text), [])

    def test_selectable_by_checks_and_default(self):
        text = table("项目|人次", "门诊|300", "住院|200", "合计|600")
        self.assertIn("NUM-023", ids(run(text, None, "number")))
        self.assertIn("NUM-023", ids(run(text)))
        self.assertNotIn("NUM-023", ids(run(text, None, "structure,format,style")))
        self.assertEqual(len(gl.CHECK_ORDER), 12)

    def test_table_cells_still_checked_by_other_rules_unchanged(self):
        # 表格行本身仍是普通段落，合计核算不影响 summary
        result = run(table("项目|人次", "门诊|300", "住院|200", "合计|600"), None, "number")
        self.assertEqual(result["summary"]["P0"], 1)

    def test_rule_registered(self):
        rules = {r["id"]: r for r in load_json(RULES_PATH)["procedural"]}
        rule = rules["NUM-023"]
        self.assertEqual((rule["check"], rule["level"], rule["category"]), ("number", "P0", "数字"))
        for phrase in ("只核算可解析为数字的单元格", "有小计行时按分组核算", "合计等于各小计之和或全部明细之和二者之一即算一致"):
            self.assertIn(phrase, rule["description"])


class ParseNumberTest(unittest.TestCase):
    def test_parse(self):
        parse = gl.parse_table_number
        self.assertEqual(parse("1,234.50万元"), (Decimal("1234.50"), 2, "万"))
        self.assertEqual(parse(" 300 人次 "), (Decimal("300"), 0, "人次"))
        self.assertEqual(parse("**12.5**"), (Decimal("12.5"), 1, ""))
        self.assertEqual(parse("95%"), (Decimal("95"), 0, "%"))
        self.assertEqual(parse("-3.0"), (Decimal("-3.0"), 1, ""))
        for text in ("", "—", "详见附件", "2026年10月", "3-5", "约300", "【待补：金额】", "12/13", "5.2.1"):
            with self.subTest(text=text):
                self.assertIsNone(parse(text))


@unittest.skipIf(docx is None, "未安装 python-docx")
class DocxTableTotalTest(unittest.TestCase):
    def build(self, rows, after_note=None):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        path = Path(tmp.name) / "成稿.docx"
        document = docx.Document()
        document.add_paragraph("2026年全院服务量见下表。")
        tab = document.add_table(rows=len(rows), cols=len(rows[0]))
        for i, row in enumerate(rows):
            for j, value in enumerate(row):
                tab.cell(i, j).text = value
        if after_note:
            document.add_paragraph("【起草说明】")
            tab2 = document.add_table(rows=3, cols=2)
            for i, row in enumerate((("项目", "人次"), ("门诊", "1"), ("合计", "9"))):
                for j, value in enumerate(row):
                    tab2.cell(i, j).text = value
        document.save(str(path))
        return path

    def lint(self, path):
        result = gl.lint_file(path, None, "number", linter=gl.Linter())
        return [i for i in result["issues"] if i["id"] == "NUM-023"]

    def test_docx_wrong_total_reported(self):
        path = self.build([("项目", "金额（万元）"), ("门诊", "120"), ("住院", "80"), ("合计", "210")])
        items = self.lint(path)
        self.assertEqual(len(items), 1)
        self.assertEqual(items[0]["level"], "P0")
        self.assertIn("合计", items[0]["excerpt"])

    def test_docx_correct_total_ok(self):
        path = self.build([("项目", "金额（万元）"), ("门诊", "120"), ("住院", "80"), ("合计", "200")])
        self.assertEqual(self.lint(path), [])

    def test_docx_blank_and_todo_cells_skip(self):
        path = self.build([("项目", "金额（万元）", "人次"), ("门诊", "", "300"), ("住院", "80", "200"),
                           ("合计", "999", "500")])
        self.assertEqual(self.lint(path), [])

    def test_docx_table_after_drafting_note_ignored(self):
        path = self.build([("项目", "人次"), ("门诊", "1"), ("合计", "1")], after_note=True)
        self.assertEqual(self.lint(path), [])


if __name__ == "__main__":
    unittest.main()
