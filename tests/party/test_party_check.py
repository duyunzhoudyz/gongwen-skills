# -*- coding: utf-8 -*-
"""skills/gongwen-party/scripts/party_check.py 的单元测试。

运行：python -m unittest discover -s tests/party -t .

每个子命令都有“必错标本”（必须报出的问题）和“必过标本”（不得误报），防止检查器空跑。
"""
from __future__ import annotations

import importlib.util
import io
import json
import shutil
import sys
import tempfile
import unittest
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[2]
SCRIPT = REPO_ROOT / "skills" / "gongwen-party" / "scripts" / "party_check.py"
EVAL_FILES = REPO_ROOT / "skills" / "gongwen-party" / "evals" / "files"


def _load():
    name = "party_check_under_test"
    if name in sys.modules:
        return sys.modules[name]
    spec = importlib.util.spec_from_file_location(name, SCRIPT)
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)  # type: ignore[union-attr]
    return module


pc = _load()

try:
    import docx  # type: ignore  # noqa: F401
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False


# ---------------------------------------------------------------------------
# 标本
# ---------------------------------------------------------------------------

# 必过：3 个问题、3 条措施，措施都有时限
MEASURES_PASS = """2026年度民主生活会个人发言提纲

按照通知要求，会前本人参加了学习。

一、对照检查情况

（一）理论学习方面。今年中心组学习有3次因下乡请假，没有补学。

（二）作风方面。到分管窗口现场只去了3次。

（三）担当作为方面。窗口流程漏了一个事项，群众跑了两趟。

二、产生问题的原因

主要是本人急于赶进度，没有把群众的难处当成自己的事。

三、整改措施

针对上述问题，本人将从以下方面整改：

（一）中心组学习因公缺席的，一周内补学，并在当月支部会上交流。

（二）从2027年1月起，每月到分管窗口现场办公1次。

（三）2027年3月底前牵头修订窗口办件流程清单，经局党组审定后执行。
"""

# 必错：3 个问题，2 条措施，其中 1 条没有时限；【起草说明】之后的内容不参与检查
MEASURES_FAIL = """个人对照检查材料

一、存在的问题

（一）理论学习方面。今年中心组学习有3次请假。

（二）作风方面。到窗口现场只去了3次。

（三）担当方面。窗口流程漏项。

二、原因分析

本人急于赶进度。

三、整改措施

（一）进一步加强理论学习，不断提高能力水平。

（二）每月到分管窗口现场办公1次。

【起草说明】
四、整改措施
（一）加强学习。
"""

OPINION_LIST = """征求意见清单
1. 分管领域项目审批慢，企业多次反映。
2. 下基层调研次数偏少。
3. 会上发言准备不充分，
有时临时拿稿。

以上意见由组织部门反馈。
"""

DRAFT_PASS = """正文……

【起草说明】
附：征求意见逐条回应表

| 序号 | 征求到的意见（照录原文） | 对应的查摆问题 | 整改措施 | 时限 |
|---|---|---|---|---|
| 1 | 分管领域项目审批慢，企业多次反映。 | 问题（三） | 2027年3月底前修订审批流程 | 2027年3月 |
| 2 | 下基层调研次数偏少。 | 问题（二） | 每月到分管单位调研1次 | 每月 |
| 3 | 会上发言准备不充分，有时临时拿稿。 | 问题（一） | 会前一天完成发言提纲 | 长期坚持 |
"""

DRAFT_FAIL = """正文……

【起草说明】
附：征求意见逐条回应表

| 序号 | 征求到的意见（照录原文） | 对应的查摆问题 | 整改措施 | 时限 |
|---|---|---|---|---|
| 1 | 工作效率有待提升 | 问题（三） | 2027年3月底前修订审批流程 | 2027年3月 |
| 2 | 下基层调研次数偏少。 | 问题（二） |  |  |
| 5 | 清单里没有的一条 | | 无 | |
"""


# 分组式征求意见汇总：分类标题下每类重新从 1 编号，共 5 条意见
GROUPED_LIST = """某单位领导班子民主生活会征求意见汇总

（一）班子建设方面
1. 班子会研究问题有时临时动议，会前沟通不够。
2. 对年轻干部的培养使用关心不够
（二）作风建设方面
1. 到基层调研提前打招呼，看不到真实情况。
2. 会议偏多，基层负责人每周要参加3到4个会。
3. 职工食堂菜品单一。
"""

GROUPED_DRAFT_PASS = """【起草说明】
附：征求意见逐条回应表

| 序号 | 征求到的意见（照录原文） | 对应的查摆问题 | 整改措施 |
|---|---|---|---|
| 1 | 班子会研究问题有时临时动议，会前沟通不够。 | 问题（一） | 2027年1月起议题会前3天送达 |
| 2 | 对年轻干部的培养使用关心不够 | 问题（二） | 2027年6月底前制定年轻干部培养计划 |
| 3 | 到基层调研提前打招呼，看不到真实情况。 | 问题（三） | 每季度不打招呼调研1次 |
| 4 | 会议偏多，基层负责人每周要参加3到4个会。 | 问题（四） | 2027年起合并同类会议 |
| 5 | 职工食堂菜品单一。 | 不属于班子集体问题 | 转交后勤部门研究，分管领导跟进反馈 |
"""

# 必错：只回应了前 4 条
GROUPED_DRAFT_FAIL = GROUPED_DRAFT_PASS.split("| 5 |")[0]


class PartyCheckTestCase(unittest.TestCase):

    def setUp(self) -> None:
        self.tmp = Path(tempfile.mkdtemp(prefix="party_check_test_"))

    def tearDown(self) -> None:
        shutil.rmtree(self.tmp, ignore_errors=True)

    def write(self, name: str, text: str, encoding: str = "utf-8") -> Path:
        path = self.tmp / name
        path.write_bytes(text.encode(encoding))
        return path

    def run_main(self, *args: str):
        out, err = io.StringIO(), io.StringIO()
        with redirect_stdout(out), redirect_stderr(err):
            code = pc.main(list(args))
        return code, out.getvalue(), err.getvalue()

    def run_json(self, *args: str) -> dict:
        code, out, err = self.run_main(*args, "--json")
        self.assertEqual(code, 0, err)
        return json.loads(out)

    @staticmethod
    def ids(report: dict):
        return [issue["id"] for issue in report["issues"]]


# ---------------------------------------------------------------------------
# measures
# ---------------------------------------------------------------------------

class MeasuresTest(PartyCheckTestCase):

    def test_pass_sample_has_no_issue(self):
        report = self.run_json("measures", str(self.write("必过.txt", MEASURES_PASS)))
        self.assertEqual(report["issues"], [])
        self.assertEqual(report["stats"]["problems"], 3)
        self.assertEqual(report["stats"]["measures"], 3)

    def test_fail_sample_reports_missing_deadline_and_count(self):
        report = self.run_json("measures", str(self.write("必错.txt", MEASURES_FAIL)))
        ids = self.ids(report)
        self.assertEqual(ids.count("MEA-001"), 1)
        self.assertIn("MEA-002", ids)
        mea1 = next(i for i in report["issues"] if i["id"] == "MEA-001")
        self.assertEqual(mea1["level"], "P2")  # 材料没有时限的不强求，只作提示
        self.assertNotIn("待补", mea1["suggestion"].replace("“待补”", "").replace("【待补】", "").replace("待补一行", ""))
        self.assertIn("进一步加强理论学习", mea1["excerpt"])
        self.assertEqual(report["stats"]["measures"], 2)

    def test_draft_notes_are_ignored(self):
        # 【起草说明】之后的“（一）加强学习。”没有时限，但不应计入
        report = self.run_json("measures", str(self.write("必错.txt", MEASURES_FAIL)))
        self.assertEqual(report["stats"]["measures_without_deadline"], 1)

    def test_yishi_enumeration(self):
        text = ("一、存在的问题\n\n一是学习不够。二是调研少。\n\n"
                "二、整改措施\n\n针对以上问题：\n\n"
                "一是每月补学一次。二是2027年6月底前走遍分管的8个村。\n")
        report = self.run_json("measures", str(self.write("一是.txt", text)))
        self.assertEqual(report["issues"], [])
        self.assertEqual(report["stats"]["problems"], 2)
        self.assertEqual(report["stats"]["measures"], 2)

    def test_yishi_without_deadline_is_reported(self):
        text = "一、整改措施\n\n一是加强学习，提高认识。二是每季度开展一次谈心谈话。\n"
        report = self.run_json("measures", str(self.write("一是缺时限.txt", text)))
        self.assertEqual(self.ids(report), ["MEA-001"])

    def test_missing_measure_section(self):
        text = "一、存在的问题\n\n（一）学习不够。\n\n二、原因分析\n\n认识不到位。\n"
        report = self.run_json("measures", str(self.write("无措施.txt", text)))
        self.assertIn("MEA-003", self.ids(report))

    def test_mixed_section(self):
        text = "一、存在问题及整改措施\n\n（一）学习不够，将加强学习。\n"
        report = self.run_json("measures", str(self.write("混写.txt", text)))
        self.assertEqual(self.ids(report), ["MEA-004"])

    def test_measure_for_problem_heading_is_not_mixed(self):
        text = ("一、存在的问题\n\n（一）学习不够。\n\n"
                "二、针对问题的整改措施\n\n（一）每月补学一次。\n")
        report = self.run_json("measures", str(self.write("针对问题.txt", text)))
        self.assertEqual(report["issues"], [])

    def test_last_year_rectification_section_is_not_measures(self):
        text = ("一、上一次民主生活会整改措施落实情况\n\n（一）已出台会前论证制度。\n\n"
                "二、对照检查情况\n\n（一）学习不够。\n\n"
                "三、整改措施\n\n（一）2027年1月起每月补学。\n")
        report = self.run_json("measures", str(self.write("上年.txt", text)))
        self.assertEqual(report["issues"], [])
        self.assertEqual(report["stats"]["measures"], 1)

    def test_deadline_variants(self):
        for phrase in ("2027年3月底前完成", "6月底前完成", "年底前完成", "三个月内完成",
                       "每季度开展一次", "立行立改", "【待补：完成时限】", "【待补：检查频次】",
                       "十月底前完成", "今年十二月完成", "到十二月完成", "明年初完成", "年中前完成"):
            with self.subTest(phrase=phrase):
                self.assertIsNotNone(pc.DEADLINE_RE.search(phrase))
        for phrase in ("进一步加强理论学习", "不断提高能力水平", "切实转变作风，多与群众交流",
                       "统一月报格式，规范台账管理", "落实年初定的联系制度", "长期坚持"):
            with self.subTest(phrase=phrase):
                self.assertIsNone(pc.DEADLINE_RE.search(phrase))

    def test_false_deadlines_are_reported(self):
        """必错：“统一月报”的“一月”、叙述中的“年初”不算时限；只写“长期坚持”单独报 P2。"""
        text = ("一、存在的问题\n\n（一）报表不统一。\n\n（二）联系点去得少。\n\n（三）学习不系统。\n\n"
                "二、整改措施\n\n（一）统一月报格式，规范台账管理。\n\n"
                "（二）走遍分管的8个村，落实年初定的联系制度。\n\n"
                "（三）长期坚持理论学习。\n")
        report = self.run_json("measures", str(self.write("伪时限.txt", text)))
        self.assertEqual(self.ids(report), ["MEA-001", "MEA-001", "MEA-006"])
        mea6 = report["issues"][2]
        self.assertEqual(mea6["level"], "P2")
        self.assertIn("长期坚持", mea6["excerpt"])
        self.assertEqual(report["stats"]["measures_without_deadline"], 3)

    def test_restated_problem_date_is_not_deadline(self):
        """必错：措施开头“针对7月12日……问题，”复述的日期不是时限；必过：引导分句之后另有时限的不报。"""
        text = ("一、存在的问题\n\n（一）汛前排查只看台账。\n\n（二）分管的村有3个没去过。\n\n"
                "二、整改措施\n\n"
                "（一）针对7月12日河堤管涌暴露的排查不实问题，今后排查一律到现场核实。\n\n"
                "（二）针对9月汛后检查发现的联系不够问题，2027年4月底前走遍分管的8个村。\n")
        report = self.run_json("measures", str(self.write("复述日期.txt", text)))
        self.assertEqual(self.ids(report), ["MEA-001"])
        self.assertEqual(report["issues"][0]["line"], 9)
        self.assertEqual(report["stats"]["measures_without_deadline"], 1)

    def test_unrecognized_structure(self):
        report = self.run_json("measures", str(self.write("散文.txt", "今年学习了很多。\n明年继续努力。\n")))
        self.assertEqual(self.ids(report), ["MEA-005"])

    def test_strict_exit_code(self):
        code, _out, _err = self.run_main("measures", str(self.write("必错.txt", MEASURES_FAIL)), "--strict")
        self.assertEqual(code, 0)  # 缺时限（MEA-001）、数量不对应均为 P2，不触发
        nomeasure = "个人对照检查材料\n\n一、存在的问题\n\n（一）理论学习方面。今年中心组学习有3次请假。\n"
        code, _out, _err = self.run_main("measures", str(self.write("无措施.txt", nomeasure)), "--strict")
        self.assertEqual(code, 1)  # 有问题无措施（MEA-003）为 P1
        code, _out, _err = self.run_main("measures", str(self.write("必过.txt", MEASURES_PASS)), "--strict")
        self.assertEqual(code, 0)

    def test_text_output(self):
        code, out, _err = self.run_main("measures", str(self.write("必错.txt", MEASURES_FAIL)))
        self.assertEqual(code, 0)
        self.assertIn("MEA-001", out)
        self.assertIn("合计", out)

    @unittest.skipUnless(HAS_DOCX, "未安装 python-docx")
    def test_docx_input(self):
        document = docx.Document()
        for line in MEASURES_FAIL.split("\n"):
            document.add_paragraph(line)
        path = self.tmp / "必错.docx"
        document.save(str(path))
        report = self.run_json("measures", str(path))
        self.assertIn("MEA-001", self.ids(report))


# ---------------------------------------------------------------------------
# opinions
# ---------------------------------------------------------------------------

class OpinionsTest(PartyCheckTestCase):

    def test_pass_sample_has_no_issue(self):
        report = self.run_json("opinions", str(self.write("清单.txt", OPINION_LIST)),
                               str(self.write("必过.txt", DRAFT_PASS)))
        self.assertEqual(report["issues"], [])
        self.assertEqual(report["stats"]["opinions"], 3)
        self.assertEqual(report["stats"]["answered"], 3)

    def test_fail_sample(self):
        report = self.run_json("opinions", str(self.write("清单.txt", OPINION_LIST)),
                               str(self.write("必错.txt", DRAFT_FAIL)))
        ids = self.ids(report)
        self.assertIn("OPN-001", ids)  # 第 3 条缺行
        self.assertIn("OPN-002", ids)  # 第 2 条措施为空
        self.assertIn("OPN-004", ids)  # 第 1 条被弱化
        self.assertIn("OPN-006", ids)  # 第 5 行清单中没有
        self.assertEqual(report["stats"]["missing"], [3])
        weakened = next(i for i in report["issues"] if i["id"] == "OPN-004")
        self.assertIn("分管领域项目审批慢", weakened["suggestion"])
        self.assertEqual(weakened["level"], "P1")  # 弱化意见违反“照录原文”的硬规则

    def test_continuation_line_and_trailing_note(self):
        opinions = pc.parse_opinions(OPINION_LIST)
        self.assertEqual([op.number for op in opinions], [1, 2, 3])
        self.assertEqual(opinions[2].text, "会上发言准备不充分，有时临时拿稿。")

    def test_pending_measure_is_p2(self):
        draft = DRAFT_PASS.replace("会前一天完成发言提纲", "【待补：整改措施】")
        report = self.run_json("opinions", str(self.write("清单.txt", OPINION_LIST)),
                               str(self.write("待补.txt", draft)))
        self.assertEqual(self.ids(report), ["OPN-003"])
        self.assertEqual(report["issues"][0]["level"], "P2")

    def test_no_table(self):
        report = self.run_json("opinions", str(self.write("清单.txt", OPINION_LIST)),
                               str(self.write("无表.txt", "正文……\n")))
        self.assertEqual(self.ids(report), ["OPN-005"])
        self.assertEqual(report["stats"]["missing"], [1, 2, 3])

    def test_chinese_numbered_list_and_table_list(self):
        cn_list = "一、分管领域项目审批慢，企业多次反映。\n二、下基层调研次数偏少。\n三、会上发言准备不充分，有时临时拿稿。\n"
        report = self.run_json("opinions", str(self.write("汉字编号.txt", cn_list)),
                               str(self.write("必过.txt", DRAFT_PASS)))
        self.assertEqual(report["issues"], [])
        table_list = ("| 序号 | 意见 |\n|---|---|\n| 1 | 分管领域项目审批慢，企业多次反映。 |\n"
                      "| 2 | 下基层调研次数偏少。 |\n| 3 | 会上发言准备不充分，有时临时拿稿。 |\n")
        report = self.run_json("opinions", str(self.write("表格清单.md", table_list)),
                               str(self.write("必过.txt", DRAFT_PASS)))
        self.assertEqual(report["issues"], [])

    def test_rows_without_number_match_by_text(self):
        draft = ("| 征求到的意见 | 整改措施 |\n|---|---|\n"
                 "| 分管领域项目审批慢，企业多次反映。 | 2027年3月底前修订流程 |\n"
                 "| 下基层调研次数偏少。 | 每月调研1次 |\n"
                 "| 会上发言准备不充分，有时临时拿稿。 | 会前一天完成提纲 |\n")
        report = self.run_json("opinions", str(self.write("清单.txt", OPINION_LIST)),
                               str(self.write("无序号.txt", draft)))
        self.assertEqual(report["issues"], [])

    def test_gb18030_list(self):
        path = self.write("清单GBK.txt", OPINION_LIST, encoding="gb18030")
        code, out, err = self.run_main("opinions", str(path), str(self.write("必过.txt", DRAFT_PASS)), "--json")
        self.assertEqual(code, 0)
        self.assertIn("GB18030", err)
        self.assertEqual(json.loads(out)["issues"], [])

    def test_empty_list_is_input_error(self):
        code, _out, err = self.run_main("opinions", str(self.write("空.txt", "\n\n")),
                                        str(self.write("必过.txt", DRAFT_PASS)))
        self.assertEqual(code, 2)
        self.assertIn("错误", err)

    def test_strict_exit_code(self):
        code, _out, _err = self.run_main("opinions", str(self.write("清单.txt", OPINION_LIST)),
                                         str(self.write("必错.txt", DRAFT_FAIL)), "--strict")
        self.assertEqual(code, 1)

    @unittest.skipUnless(HAS_DOCX, "未安装 python-docx")
    def test_docx_table(self):
        document = docx.Document()
        document.add_paragraph("【起草说明】")
        table = document.add_table(rows=4, cols=3)
        rows = [("序号", "征求到的意见（照录原文）", "整改措施"),
                ("1", "分管领域项目审批慢，企业多次反映。", "2027年3月底前修订流程"),
                ("2", "下基层调研次数偏少。", "每月调研1次"),
                ("3", "会上发言准备不充分，有时临时拿稿。", "会前一天完成提纲")]
        for r, values in enumerate(rows):
            for c, value in enumerate(values):
                table.cell(r, c).text = value
        path = self.tmp / "回应表.docx"
        document.save(str(path))
        report = self.run_json("opinions", str(self.write("清单.txt", OPINION_LIST)), str(path))
        self.assertEqual(report["issues"], [])

    def test_grouped_list_missing_one_is_reported(self):
        """必错：分组式清单（每类重新编号）共 5 条，回应表只回应前 4 条，必须报第 5 条缺失，不得错位报 OPN-004。"""
        report = self.run_json("opinions", str(self.write("分组清单.txt", GROUPED_LIST)),
                               str(self.write("缺一条.txt", GROUPED_DRAFT_FAIL)))
        self.assertEqual(self.ids(report), ["OPN-001"])
        self.assertIn("职工食堂菜品单一", report["issues"][0]["excerpt"])
        stats = report["stats"]
        self.assertEqual((stats["opinions"], stats["answered"], stats["missing"]), (5, 4, [5]))
        self.assertEqual(stats["group_titles"], 2)
        self.assertTrue(stats["renumbered"])
        code, out, _err = self.run_main("opinions", str(self.tmp / "分组清单.txt"), str(self.tmp / "缺一条.txt"))
        self.assertEqual(code, 0)
        self.assertIn("2 行分类标题", out)

    def test_grouped_list_pass(self):
        """必过：分组式清单逐条回应，回应表连续编号或沿用分类内编号都不误报。"""
        list_path = self.write("分组清单.txt", GROUPED_LIST)
        report = self.run_json("opinions", str(list_path), str(self.write("齐全.txt", GROUPED_DRAFT_PASS)))
        self.assertEqual(report["issues"], [])
        self.assertEqual(report["stats"]["answered"], 5)
        group_numbered = GROUPED_DRAFT_PASS
        for old, new in (("| 3 | 到基层", "| 1 | 到基层"), ("| 4 | 会议偏多", "| 2 | 会议偏多"),
                         ("| 5 | 职工食堂", "| 3 | 职工食堂")):
            group_numbered = group_numbered.replace(old, new)
        report = self.run_json("opinions", str(list_path), str(self.write("分类编号.txt", group_numbered)))
        self.assertEqual(report["issues"], [])

    def test_group_titles_variants(self):
        """不编号的分类标题不并入上一条意见；编号的分类标题下不编号的意见一行一条。"""
        unnumbered_titles = ("班子建设方面\n1. 班子会研究问题有时临时动议，会前沟通不够。\n"
                             "2. 对年轻干部的培养使用关心不够\n作风建设方面\n3. 职工食堂菜品单一。\n")
        opinions = pc.parse_opinions(unnumbered_titles)
        self.assertEqual([op.text for op in opinions],
                         ["班子会研究问题有时临时动议，会前沟通不够。", "对年轻干部的培养使用关心不够",
                          "职工食堂菜品单一。"])
        plain_items = ("一、班子建设方面\n班子会研究问题有时临时动议。\n对年轻干部关心不够。\n\n"
                       "二、作风建设方面\n会议偏多。\n\n以上意见由组织部门反馈。\n")
        parsed = pc.parse_opinion_list(plain_items)
        self.assertEqual([op.text for op in parsed.opinions],
                         ["班子会研究问题有时临时动议。", "对年轻干部关心不够。", "会议偏多。"])
        self.assertEqual([op.number for op in parsed.opinions], [1, 2, 3])
        self.assertEqual(parsed.group_titles, 2)
        no_numbers = "某单位征求意见汇总\n班子建设方面\n调研少。\n会议多。\n以上意见由组织部门反馈。\n"
        self.assertEqual([op.text for op in pc.parse_opinions(no_numbers)], ["调研少。", "会议多。"])

    def test_paraphrased_row_without_number_is_not_missing(self):
        """回应表没有序号栏、某条意见被改了几个字：按原文相近对应，报 OPN-004，不报缺行。"""
        draft = ("| 征求到的意见 | 整改措施 |\n|---|---|\n"
                 "| 分管领域项目审批慢，企业多次反映。 | 2027年3月底前修订流程 |\n"
                 "| 下基层调研的次数偏少。 | 每月调研1次 |\n"
                 "| 会上发言准备不充分，有时临时拿稿。 | 会前一天完成提纲 |\n")
        report = self.run_json("opinions", str(self.write("清单.txt", OPINION_LIST)),
                               str(self.write("改字.txt", draft)))
        self.assertEqual(self.ids(report), ["OPN-004"])
        self.assertIn("第 2 条", report["issues"][0]["message"])


# ---------------------------------------------------------------------------
# 命令行与输入错误
# ---------------------------------------------------------------------------

class CliTest(PartyCheckTestCase):

    def test_missing_file(self):
        code, _out, err = self.run_main("measures", str(self.tmp / "不存在.txt"))
        self.assertEqual(code, 2)
        self.assertIn("文件不存在", err)

    def test_unsupported_suffix(self):
        code, _out, err = self.run_main("measures", str(self.write("稿件.pdf", "x")))
        self.assertEqual(code, 2)
        self.assertIn("不支持", err)

    def test_no_subcommand(self):
        code, _out, _err = self.run_main()
        self.assertEqual(code, 2)

    def test_help_is_chinese(self):
        code, out, _err = self.run_main("--help")
        self.assertEqual(code, 0)
        self.assertIn("用法", out)
        self.assertIn("子命令", out)

    def test_bad_option(self):
        code, _out, _err = self.run_main("measures", "--不存在的选项")
        self.assertEqual(code, 2)


# ---------------------------------------------------------------------------
# 评测材料上的回归
# ---------------------------------------------------------------------------

class EvalFilesTest(PartyCheckTestCase):

    def test_hospital_draft_measures_have_no_deadline(self):
        """评测 5 的初稿：三条措施中两条没有时限（“每月至少自学两次”有频次）。"""
        path = EVAL_FILES / "某医院党员2026年组织生活会对照检查材料初稿.txt"
        report = self.run_json("measures", str(path))
        self.assertEqual(report["stats"]["problems"], 3)
        self.assertEqual(report["stats"]["measures"], 3)
        self.assertEqual(self.ids(report).count("MEA-001"), 2)

    def test_university_draft(self):
        """评测 17 的送审稿：措施（一）（二）没有时限，（三）有时限。"""
        path = EVAL_FILES / "某大学教工党员对照检查材料送审稿.txt"
        report = self.run_json("measures", str(path))
        self.assertEqual(self.ids(report).count("MEA-001"), 2)

    def test_eval_opinion_list_parses(self):
        opinions = pc.parse_opinions((EVAL_FILES / "某公司党委征求意见汇总.txt").read_text(encoding="utf-8"))
        self.assertEqual([op.number for op in opinions], [1, 2, 3, 4, 5, 6])


if __name__ == "__main__":
    unittest.main()
