# -*- coding: utf-8 -*-
"""skills/gongwen-hospital/scripts/hospital_check.py 的单元测试。

运行：python -m unittest discover -s tests/hospital -t .
必错标本和必过标本在 tests/hospital/specimens/ 下。
"""
from __future__ import annotations

import io
import json
import shutil
import sys
import tempfile
import unittest
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[2]
SCRIPTS_DIR = REPO_ROOT / "skills" / "gongwen-hospital" / "scripts"
SPECIMENS = Path(__file__).resolve().parent / "specimens"
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

import hospital_check as hc  # noqa: E402

try:
    import docx  # type: ignore  # noqa: F401
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False


def ids_of(result) -> list:
    return [issue["id"] for issue in result["issues"]]


class HospitalCheckBase(unittest.TestCase):

    @classmethod
    def setUpClass(cls) -> None:
        cls.ruleset = hc.load_rules()

    def setUp(self) -> None:
        self.tmp = Path(tempfile.mkdtemp(prefix="hospital_check_test_"))

    def tearDown(self) -> None:
        shutil.rmtree(self.tmp, ignore_errors=True)

    def check(self, text: str, doc_type=None, external: bool = False):
        return hc.check_text(text, doc_type, external, ruleset=self.ruleset)

    def run_main(self, *args: str):
        out, err = io.StringIO(), io.StringIO()
        with redirect_stdout(out), redirect_stderr(err):
            code = hc.main(list(args))
        return code, out.getvalue(), err.getvalue()


class RuleFileTest(HospitalCheckBase):
    """规则文件本身的完整性。"""

    def test_every_rule_has_examples_and_basis(self) -> None:
        for rule in self.ruleset.rules:
            with self.subTest(rule=rule.id):
                self.assertTrue(rule.hit_examples, "每条正则规则至少要有一个命中示例")
                self.assertTrue(rule.miss_examples, "每条正则规则至少要有一个不命中示例")
                self.assertTrue(rule.basis)
        for rule in self.ruleset.procedural:
            self.assertTrue(rule.basis)

    def test_rule_examples(self) -> None:
        for rule in self.ruleset.rules:
            doc_type = rule.types[0] if rule.types else None
            for example in rule.hit_examples:
                with self.subTest(rule=rule.id, hit=example):
                    self.assertIn(rule.id, ids_of(self.check(example, doc_type)))
            for example in rule.miss_examples:
                with self.subTest(rule=rule.id, miss=example):
                    self.assertNotIn(rule.id, ids_of(self.check(example, doc_type)))

    def test_bad_rule_files_raise(self) -> None:
        cases = {
            "not_json.json": "{",
            "bad_level.json": json.dumps({"types": [], "rules": [
                {"id": "X-1", "category": "c", "gate": "g", "level": "P9", "message": "m",
                 "basis": "b", "pattern": "a"}]}),
            "bad_regex.json": json.dumps({"types": [], "rules": [
                {"id": "X-1", "category": "c", "gate": "g", "level": "P1", "message": "m",
                 "basis": "b", "pattern": "("}]}),
            "dup_id.json": json.dumps({"types": [], "rules": [
                {"id": "X-1", "category": "c", "gate": "g", "level": "P1", "message": "m",
                 "basis": "b", "pattern": "a"},
                {"id": "X-1", "category": "c", "gate": "g", "level": "P1", "message": "m",
                 "basis": "b", "pattern": "b"}]}),
            "unknown_type.json": json.dumps({"types": ["RCA"], "rules": [
                {"id": "X-1", "category": "c", "gate": "g", "level": "P1", "message": "m",
                 "basis": "b", "pattern": "a", "types": ["讲话稿"]}]}),
        }
        for name, content in cases.items():
            path = self.tmp / name
            path.write_text(content, encoding="utf-8")
            with self.subTest(case=name):
                with self.assertRaises(hc.InputError):
                    hc.load_rules(path)
        with self.assertRaises(hc.InputError):
            hc.load_rules(self.tmp / "missing.json")


class SpecimenTest(HospitalCheckBase):
    """必错标本必须报出预期问题，必过标本必须零问题。"""

    def test_must_fail_specimen(self) -> None:
        result = hc.check_file(SPECIMENS / "必错-根因分析.txt", "RCA", ruleset=self.ruleset)
        found = ids_of(result)
        for rid in ("HOS-002", "HOS-003", "HOS-004", "HOS-006", "HOS-018", "HOS-019", "HOS-020"):
            with self.subTest(rule=rid):
                self.assertIn(rid, found)
        # 【起草说明】之后的“责任心不强”不参与检查
        self.assertEqual(found.count("HOS-018"), 1)

    def test_must_fail_specimen_without_type_skips_rca_rules(self) -> None:
        result = hc.check_file(SPECIMENS / "必错-根因分析.txt", None, ruleset=self.ruleset)
        found = ids_of(result)
        for rid in ("HOS-018", "HOS-019", "HOS-020"):
            self.assertNotIn(rid, found)
        self.assertIn("HOS-004", found)
        self.assertTrue(any("未指定 --type" in note for note in result["notes"]))

    def test_must_pass_specimen(self) -> None:
        for doc_type in ("质量分析", "绩效分析", None):
            with self.subTest(doc_type=doc_type):
                result = hc.check_file(SPECIMENS / "必过-质量分析.txt", doc_type, ruleset=self.ruleset)
                self.assertEqual(result["issues"], [], result["issues"])
                self.assertEqual({lv: result["summary"][lv] for lv in hc.LEVELS},
                                 {"P0": 0, "P1": 0, "P2": 0})
                self.assertEqual(result["summary"]["markers"], {"待补": 3, "待核实": 0})

    def test_must_pass_rca_specimen(self) -> None:
        """规范的根因分析：顿号连接的指标、“未给予处分”、起草说明中的处理意见都不报。"""
        result = hc.check_file(SPECIMENS / "必过-根因分析.txt", "RCA", ruleset=self.ruleset)
        self.assertEqual(result["issues"], [], result["issues"])

    def test_must_pass_specimen_as_external(self) -> None:
        result = hc.check_file(SPECIMENS / "必过-质量分析.txt", "对外宣传", ruleset=self.ruleset)
        self.assertEqual(result["issues"], [])
        self.assertTrue(result["external"])


class BehaviourTest(HospitalCheckBase):

    def test_external_escalates_levels(self) -> None:
        text = "该技术可保证治愈。手术成功率达100%。达到国内领先水平。"
        internal = {i["id"]: i["level"] for i in self.check(text)["issues"]}
        external = {i["id"]: i["level"] for i in self.check(text, external=True)["issues"]}
        by_type = {i["id"]: i["level"] for i in self.check(text, "对外宣传")["issues"]}
        self.assertEqual(internal, {"HOS-001": "P1", "HOS-002": "P2", "HOS-003": "P2"})
        self.assertEqual(external, {"HOS-001": "P0", "HOS-002": "P1", "HOS-003": "P1"})
        self.assertEqual(by_type, external)

    def test_surgery_names_with_genzhi_pass(self) -> None:
        text = "完成胃癌根治术86例，乳腺癌改良根治手术40例，开展根治性放疗12例。"
        self.assertNotIn("HOS-002", ids_of(self.check(text, external=True)))

    def test_joint_meeting_exempt_in_prohibition(self) -> None:
        self.assertNotIn("HOS-004", ids_of(self.check("医院严格执行议事规则，不得以党政联席会议代替党委会议研究决定重大问题。")))
        self.assertIn("HOS-004", ids_of(self.check("经党政联席会议研究，同意调整绩效方案。")))

    def test_procedural_sanzhongyida(self) -> None:
        missing = "议题名称：购置数字减影血管造影机\n本项目属于“三重一大”事项，预算1200万元。经院长办公会议讨论通过。"
        complete = missing + "\n拟提交党委会议研究决定。"
        self.assertIn("HOS-005", ids_of(self.check(missing, "议题")))
        self.assertNotIn("HOS-005", ids_of(self.check(complete, "议题")))
        # 非议题文书不做这项检查
        self.assertNotIn("HOS-005", ids_of(self.check(missing, "质量分析")))
        issue = next(i for i in self.check(missing, "议题")["issues"] if i["id"] == "HOS-005")
        self.assertEqual(issue["line"], 2)
        self.assertEqual(issue["level"], "P1")

    def test_procedural_negated_party_meeting(self) -> None:
        """“尚未提交党委会议”“未经党委会议”不算已有党委会议程序。"""
        base = "本项目属于“三重一大”事项，预算1200万元。"
        for text in ("目前尚未提交党委会议研究。",
                     "未经党委会议研究，已签订意向书。",
                     "该事项尚未提交医院党委会议。",
                     "此前没有上党委会研究。"):
            with self.subTest(negated=text):
                self.assertIn("HOS-005", ids_of(self.check(base + text, "议题")))
        for text in ("经院长办公会议讨论通过后，提交党委会议研究决定。",
                     "尚未提交党委会议研究，拟于11月提交党委会议研究决定。",
                     "未通过的事项不再上会，本事项拟提交党委会议研究决定。"):
            with self.subTest(satisfied=text):
                self.assertNotIn("HOS-005", ids_of(self.check(base + text, "议题")))

    def test_index_vs_percentage_points(self) -> None:
        self.assertIn("HOS-006", ids_of(self.check("时间消耗指数下降0.05个百分点。")))
        self.assertNotIn("HOS-006", ids_of(self.check("时间消耗指数由1.05下降到1.00，出院患者手术占比提高2.1个百分点。")))
        self.assertIn("HOS-007", ids_of(self.check("科室CMI：1.25%")))

    def test_index_rule_does_not_cross_enumeration(self) -> None:
        """顿号并列的另一个比率指标用百分点是对的，不能误报。"""
        for text in ("平均住院日缩短0.4天、四级手术比例提高1.2个百分点",
                     "平均住院日缩短0.4天、出院患者手术占比提高2.1个百分点",
                     "RW值高的病组占比提高3.2个百分点",
                     "CMI为0.94，微创手术占比提高1.9个百分点"):
            with self.subTest(text=text):
                self.assertNotIn("HOS-006", ids_of(self.check(text)))
        self.assertIn("HOS-006", ids_of(self.check("平均住院日较上年缩短0.4个百分点、手术占比提高1.2个百分点")))

    def test_index_rule_crosses_comma(self) -> None:
        """最常见的错法带逗号：“CMI为0.94，较上年提高0.05个百分点”。"""
        for text in ("CMI为0.94，较上年提高0.05个百分点",
                     "病例组合指数（CMI）为1.08，同比提高0.06个百分点",
                     "平均住院日为7.6天，较上年同期缩短0.4个百分点",
                     "CMI由1.02提高到1.08，提高了6个百分点",
                     "（一）医疗质量。CMI为1.05，较上年提高0.06个百分点。"):
            with self.subTest(hit=text):
                self.assertIn("HOS-006", ids_of(self.check(text)))
        for text in ("CMI为0.94，较上年提高0.03；手术患者并发症发生率下降0.1个百分点",
                     "CMI为1.08，门诊次均费用增幅提高2.1个百分点",
                     "平均住院日为7.6天，病床使用率下降1.2个百分点",
                     "CMI增幅较上年提高0.5个百分点"):
            with self.subTest(miss=text):
                self.assertNotIn("HOS-006", ids_of(self.check(text)))

    def test_eval7_draft_reports_index_points(self) -> None:
        """评测 7 的初稿中“CMI为1.05，较上年提高0.06个百分点”必须查出。"""
        draft = REPO_ROOT / "skills" / "gongwen-hospital" / "evals" / "files" / "某医院2025年度绩效监测分析初稿.txt"
        result = hc.check_file(draft, "绩效分析", ruleset=self.ruleset)
        self.assertIn("HOS-006", ids_of(result))
        self.assertIn("HOS-015", ids_of(result))

    def test_guokao_not_in_other_words(self) -> None:
        for text in ("按全国考核要求按时上报数据。", "赴中国考察学习。", "出国考察3人次。"):
            with self.subTest(text=text):
                self.assertNotIn("HOS-015", ids_of(self.check(text, "绩效分析")))
        self.assertIn("HOS-015", ids_of(self.check("国考评价结果为A。", "绩效分析")))

    def test_efficacy_hundred_percent_in_words(self) -> None:
        self.assertIn("HOS-002", ids_of(self.check("手术成功率达到百分之百。")))

    def test_rca_punishment_variants(self) -> None:
        """扣绩效、通报批评、受到处分都算写入处理；否定说法不报。"""
        for text in ("扣当事护士当月绩效500元，全科通报批评。",
                     "科室意见：扣当月绩效500元。",
                     "当事护士已受到处分。",
                     "对当事人予以经济处罚。"):
            with self.subTest(hit=text):
                self.assertIn("HOS-020", ids_of(self.check(text, "RCA")))
        for text in ("均已妥善处理，未给予处分",
                     "主动上报的不予处分，不扣绩效，不通报批评",
                     "处理意见另按本院程序办理，未写入本报告",
                     "扣除节假日后，平均每日上报2.1例次"):
            with self.subTest(miss=text):
                self.assertNotIn("HOS-020", ids_of(self.check(text, "RCA")))
        # 只在 RCA 中检查
        self.assertNotIn("HOS-020", ids_of(self.check("扣当月绩效500元。", "质量分析")))

    def test_rca_more_punishment_and_character_words(self) -> None:
        for text in ("对当事护士批评教育。", "在全院通报。", "对护士长诫勉谈话。",
                     "取消当事人当年评先评优资格。", "当事护士待岗培训。"):
            with self.subTest(hit=text):
                self.assertIn("HOS-020", ids_of(self.check(text, "RCA")))
        for text in ("将考核扣分与科室绩效挂钩。", "在全院通报表扬主动上报的科室。"):
            with self.subTest(miss=text):
                self.assertNotIn("HOS-020", ids_of(self.check(text, "RCA")))
        self.assertIn("HOS-018", ids_of(self.check("当班护士工作不够细致。", "RCA")))

    def test_markers_counted_in_summary(self) -> None:
        text = "出院患者【待补：人次】人次，病历甲级率【待核实：病案科确认】。\n【起草说明】\n【待补：不计入】"
        summary = self.check(text)["summary"]
        self.assertEqual(summary["markers"], {"待补": 1, "待核实": 1})
        self.assertEqual(summary["P2"], 0)

    def test_outdated_basis_exempt_in_history(self) -> None:
        self.assertIn("HOS-016", ids_of(self.check("按照《三级医院评审标准(2020年版)》要求整改。")))
        self.assertNotIn("HOS-016", ids_of(self.check("上一周期按《三级医院评审标准（2022年版）》评审。")))

    def test_guokao_only_in_formal_types(self) -> None:
        text = "2025年度国考成绩为A。"
        self.assertIn("HOS-015", ids_of(self.check(text, "绩效分析")))
        self.assertNotIn("HOS-015", ids_of(self.check(text, "述职")))
        self.assertNotIn("HOS-015", ids_of(self.check(text)))

    def test_draft_note_cutoff_variants(self) -> None:
        for note_line in ("【起草说明】", "## 【起草说明】", "**【起草说明】**"):
            with self.subTest(note=note_line):
                text = f"正文没有问题。\n{note_line}\n大手术比例为18%。"
                self.assertEqual(self.check(text)["issues"], [])

    def test_issue_fields_match_lint_interface(self) -> None:
        issue = self.check("大手术比例为18.6%")["issues"][0]
        for key in ("id", "level", "category", "line", "excerpt", "message", "suggestion", "basis", "gate"):
            self.assertIn(key, issue)
        self.assertEqual(issue["check"], "hospital")

    def test_invalid_type_raises(self) -> None:
        with self.assertRaises(hc.InputError):
            self.check("正文", "讲话稿")


class CliTest(HospitalCheckBase):

    def test_json_output_and_strict(self) -> None:
        path = self.tmp / "对外稿.txt"
        path.write_text("本院专家团队保证治愈，药到病除。", encoding="utf-8")
        code, out, _err = self.run_main(str(path), "--type", "对外宣传", "--json")
        self.assertEqual(code, 0)
        data = json.loads(out)
        self.assertEqual(data["summary"]["P0"], 2)
        code, _out, _err = self.run_main(str(path), "--type", "对外宣传", "--strict")
        self.assertEqual(code, 1)
        code, _out, _err = self.run_main(str(path), "--strict")
        self.assertEqual(code, 0)  # 内部材料只报 P1

    def test_text_output(self) -> None:
        path = self.tmp / "稿.md"
        path.write_text("# 标题\n\n全年CMI为108%。\n", encoding="utf-8")
        code, out, _err = self.run_main(str(path))
        self.assertEqual(code, 0)
        self.assertIn("HOS-007", out)
        self.assertIn("第3行", out)

    def test_input_errors_return_2(self) -> None:
        code, _out, err = self.run_main(str(self.tmp / "不存在.txt"))
        self.assertEqual(code, 2)
        self.assertIn("文件不存在", err)
        bad = self.tmp / "稿.pdf"
        bad.write_bytes(b"%PDF")
        code, _out, err = self.run_main(str(bad))
        self.assertEqual(code, 2)
        good = self.tmp / "稿.txt"
        good.write_text("正文", encoding="utf-8")
        code, _out, err = self.run_main(str(good), "--type", "讲话稿")
        self.assertEqual(code, 2)
        self.assertIn("--type", err)
        code, _out, _err = self.run_main()
        self.assertEqual(code, 2)

    def test_list_rules(self) -> None:
        code, out, _err = self.run_main("--list-rules")
        self.assertEqual(code, 0)
        self.assertIn("HOS-001", out)
        self.assertIn("HOS-005", out)

    def test_gb18030_input(self) -> None:
        path = self.tmp / "旧编码.txt"
        path.write_bytes("抗生素使用强度为38.5。".encode("gb18030"))
        code, out, _err = self.run_main(str(path), "--json")
        self.assertEqual(code, 0)
        self.assertIn("HOS-010", out)

    @unittest.skipUnless(HAS_DOCX, "未安装 python-docx")
    def test_docx_input_with_table_and_note(self) -> None:
        import docx  # type: ignore
        document = docx.Document()
        document.add_paragraph("关于住院患者跌倒事件的根本原因分析报告")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "原因"
        table.cell(0, 1).text = "当班护士粗心大意"
        document.add_paragraph("【起草说明】")
        document.add_paragraph("当班护士责任心不强")
        path = self.tmp / "根因分析.docx"
        document.save(str(path))
        code, out, _err = self.run_main(str(path), "--type", "RCA", "--json")
        self.assertEqual(code, 0)
        data = json.loads(out)
        self.assertEqual(ids_of(data), ["HOS-018"])
        self.assertEqual(data["issues"][0]["excerpt"], "粗心大意")


if __name__ == "__main__":
    unittest.main()
