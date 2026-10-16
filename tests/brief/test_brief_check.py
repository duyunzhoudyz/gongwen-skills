# -*- coding: utf-8 -*-
"""skills/gongwen-brief/scripts/brief_check.py 的单元测试。

标本放在 tests/brief/specimens/：pass-* 为必过标本（不得有 P0、P1），fail-* 为必错标本（必须报出指定规则）。

运行：python -m unittest discover -s tests/brief -t .
"""
from __future__ import annotations

import io
import json
import shutil
import sys
import tempfile
import unittest
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[2]
SCRIPTS_DIR = REPO_ROOT / "skills" / "gongwen-brief" / "scripts"
SPECIMENS = Path(__file__).resolve().parent / "specimens"
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

import brief_check  # noqa: E402

try:
    import docx  # type: ignore  # noqa: F401
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False


def ids(report: dict, level: str = "") -> list:
    return [it["id"] for it in report["issues"] if not level or it["level"] == level]


class BriefTestCase(unittest.TestCase):

    def setUp(self) -> None:
        self.tmp = Path(tempfile.mkdtemp(prefix="brief_check_test_"))

    def tearDown(self) -> None:
        shutil.rmtree(self.tmp, ignore_errors=True)

    def write(self, name: str, text: str, encoding: str = "utf-8") -> Path:
        path = self.tmp / name
        path.write_bytes(text.encode(encoding))
        return path

    def check(self, text: str, doc_type: str = "信息", source: str = None, **kwargs) -> dict:
        draft = self.write("正文.txt", text)
        sources = [self.write("材料.txt", source)] if source is not None else None
        return brief_check.build_report(draft, doc_type, sources, **kwargs)

    def run_main(self, *args: str):
        out, err = io.StringIO(), io.StringIO()
        with redirect_stdout(out), redirect_stderr(err):
            code = brief_check.main(list(args))
        return code, out.getvalue(), err.getvalue()


# ---------------------------------------------------------------------------
# 标本
# ---------------------------------------------------------------------------

class SpecimenTests(BriefTestCase):
    """必过标本没有 P0、P1；必错标本报出指定规则。"""

    PASS = [
        ("pass-信息.txt", "信息", "pass-材料.txt"),
        ("pass-简报.txt", "简报", None),
        ("pass-推文.txt", "公众号", None),
        ("pass-工作提法.txt", "新闻稿", None),
        ("pass-升学率批评.txt", "信息", None),
    ]
    FAIL = [
        ("fail-编造引语.txt", "信息", "pass-材料.txt", "BRF-003", "P0"),
        ("fail-标题数字.txt", "信息", None, "BRF-002", "P1"),
        ("fail-导语铺垫.txt", "信息", None, "BRF-005", "P1"),
        ("fail-疗效保证.txt", "新闻稿", None, "BRF-008", "P0"),
        ("fail-疗效保证.txt", "新闻稿", None, "BRF-007", "P1"),
        ("fail-简报缺要素.txt", "简报", None, "BRF-012", "P1"),
        ("fail-简报缺要素.txt", "简报", None, "BRF-014", "P2"),
        ("fail-标题党.txt", "公众号", None, "BRF-010", "P1"),
        ("fail-标题党.txt", "公众号", None, "BRF-011", "P2"),
        ("fail-升学率.txt", "新闻稿", None, "BRF-009", "P1"),
        ("fail-简报标题数字.txt", "简报", None, "BRF-002", "P1"),
        ("fail-短引语.txt", "信息", "材料-推进会记录.txt", "BRF-003", "P1"),
    ]

    def test_pass_specimens(self):
        for name, doc_type, source in self.PASS:
            with self.subTest(name=name):
                sources = [SPECIMENS / source] if source else None
                report = brief_check.build_report(SPECIMENS / name, doc_type, sources)
                self.assertEqual(report["summary"]["P0"], 0, report["issues"])
                self.assertEqual(report["summary"]["P1"], 0, report["issues"])

    def test_pass_info_in_range_and_quotes_traced(self):
        report = brief_check.build_report(SPECIMENS / "pass-信息.txt", "信息", [SPECIMENS / "pass-材料.txt"])
        low, high = report["length_range"]
        self.assertTrue(low <= report["char_count"] <= high, report["char_count"])
        self.assertNotIn("BRF-003", ids(report))
        self.assertNotIn("BRF-004", ids(report))

    def test_fail_specimens(self):
        for name, doc_type, source, rule, level in self.FAIL:
            with self.subTest(name=name, rule=rule):
                sources = [SPECIMENS / source] if source else None
                report = brief_check.build_report(SPECIMENS / name, doc_type, sources)
                self.assertIn(rule, ids(report, level), report["issues"])

    def test_fabricated_quote_only_flags_the_invented_one(self):
        report = brief_check.build_report(
            SPECIMENS / "fail-编造引语.txt", "信息", [SPECIMENS / "pass-材料.txt"])
        p0 = [it for it in report["issues"] if it["id"] == "BRF-003" and it["level"] == "P0"]
        self.assertEqual(len(p0), 1)
        self.assertIn("最暖心", p0[0]["excerpt"])

    def test_pass_hospital_surgery_names(self):
        # “胃癌根治术”“根治性切除”是手术名称，医院稿加 --unit 公立医院 也不按疗效保证处理
        report = brief_check.build_report(SPECIMENS / "pass-医院手术.txt", "新闻稿", unit="公立医院")
        self.assertEqual(report["summary"]["P0"], 0, report["issues"])
        self.assertEqual(report["summary"]["P1"], 0, report["issues"])
        self.assertNotIn("BRF-008", ids(report))

    def test_short_quote_specimen_flags_all_three(self):
        report = brief_check.build_report(
            SPECIMENS / "fail-短引语.txt", "信息", [SPECIMENS / "材料-推进会记录.txt"])
        p1 = {it["excerpt"]: it["message"] for it in report["issues"] if it["id"] == "BRF-003" and it["level"] == "P1"}
        self.assertEqual(set(p1), {"要完成。", "挺好的。", "全国人居环境示范村"})
        self.assertIn("村民王某", p1["挺好的。"])
        self.assertIn("荣誉称号", p1["全国人居环境示范村"])
        self.assertEqual(report["summary"]["P0"], 0)

    def test_bait_title_detected_even_with_exclamation(self):
        report = brief_check.build_report(SPECIMENS / "fail-标题党.txt", "公众号")
        self.assertTrue(report["title"].startswith("重磅"))
        excerpts = [it["excerpt"] for it in report["issues"] if it["id"] == "BRF-010"]
        self.assertIn("重磅", excerpts)


# ---------------------------------------------------------------------------
# 篇幅
# ---------------------------------------------------------------------------

class LengthTests(BriefTestCase):

    def test_over_limit_is_p1_for_info(self):
        body = "今年7月以来，某街道试行定时定点投放。" + "督导员现场指导居民分类。" * 80
        report = self.check("某街道试行定时定点投放\n\n" + body)
        self.assertIn("BRF-001", ids(report, "P1"))

    def test_over_limit_is_p2_for_news(self):
        body = "10月15日，某医院开展义诊。" + "医护人员为居民测量血压。" * 120
        report = self.check("某医院开展义诊\n\n" + body, "新闻稿")
        self.assertIn("BRF-001", ids(report, "P2"))
        self.assertNotIn("BRF-001", ids(report, "P1"))

    def test_custom_range_overrides_default(self):
        body = "今年7月以来，某街道试行定时定点投放，分类准确率由62%提高到85%。"
        report = self.check("某街道试行定时定点投放\n\n" + body, low=10, high=60)
        self.assertNotIn("BRF-001", ids(report))
        self.assertEqual(report["length_range"], [10, 60])

    def test_alias_ranges(self):
        text = "某镇强降雨首报\n\n10月8日凌晨2时至8时，某镇累计降雨186毫米，已转移群众127人。"
        urgent = self.check(text, "紧急信息")
        self.assertEqual(urgent["type"], "信息")
        self.assertEqual(urgent["length_range"], [None, 500])
        self.assertNotIn("BRF-001", ids(urgent))
        self.assertEqual(self.check(text, "社区信息")["length_range"], [300, 500])

    def test_title_and_draft_note_excluded_from_count(self):
        text = "标题行很长很长很长很长\n\n今年7月以来，某街道试行定时定点投放。\n\n【起草说明】\n" + "说明" * 500
        report = self.check(text, low=None, high=None)
        self.assertEqual(report["char_count"], len("今年7月以来，某街道试行定时定点投放。"))

    def test_brief_header_and_footer_excluded(self):
        report = brief_check.build_report(SPECIMENS / "pass-简报.txt", "简报")
        self.assertLess(report["char_count"], 330)
        self.assertNotIn("报：", "".join(it["excerpt"] for it in report["issues"]))


# ---------------------------------------------------------------------------
# 标题数字概括
# ---------------------------------------------------------------------------

class TitleNumberTests(BriefTestCase):

    def test_quoted_summary_matches_items(self):
        text = "某街道“三个到位”推进垃圾分类\n\n今年7月以来，某街道推进垃圾分类。\n\n一是时段到位。\n\n二是人员到位。\n\n三是宣传到位。"
        self.assertNotIn("BRF-002", ids(self.check(text)))

    def test_plain_summary_mismatch(self):
        text = "某街道推出五项举措推进垃圾分类\n\n今年7月以来，某街道推进垃圾分类。一是定时投放；二是督导员守桶；三是积分兑换。"
        report = self.check(text)
        self.assertIn("BRF-002", ids(report, "P1"))

    def test_headings_take_priority(self):
        text = "某局三项举措优化营商环境\n\n今年以来，某局优化营商环境。\n\n一、一窗通办\n\n一是整合窗口；二是压缩时限；三是网上办理；四是帮办代办。\n\n二、综合查一次\n\n三、政策直达"
        self.assertNotIn("BRF-002", ids(self.check(text)))

    def test_no_countable_items_is_p2(self):
        text = "某街道推出三项举措推进垃圾分类\n\n今年7月以来，某街道推进垃圾分类，分类准确率由62%提高到85%。"
        report = self.check(text)
        self.assertIn("BRF-002", ids(report, "P2"))

    def test_fixed_terms_and_data_not_treated_as_summary(self):
        cases = [
            "某支部开展“两个确立”主题学习\n\n10月12日，某支部开展学习。一是读原文。",
            "某区落实中央八项规定精神\n\n今年以来，某区开展检查。一是查公车。",
            "某区推行“一站式”服务\n\n今年以来，某区推行服务。一是设窗口。二是减材料。",
            "某局12项改革任务全部完成\n\n今年以来，某局完成改革任务。一是简化审批。",
            "某镇建成一个便民服务平台\n\n今年以来，某镇建成平台。",
        ]
        for text in cases:
            with self.subTest(text=text[:12]):
                self.assertNotIn("BRF-002", ids(self.check(text)))

    def test_title_summaries_parse(self):
        found = brief_check.title_summaries("某街道“四个到位”抓实安全　三大行动和5项措施同步推进")
        self.assertEqual(sorted(n for n, _ in found), [3, 4, 5])


# ---------------------------------------------------------------------------
# 引语
# ---------------------------------------------------------------------------

class QuoteTests(BriefTestCase):

    SOURCE = "居民王某说：“家门口就能吃上热乎饭，孩子们也放心了。”街道推出“暖心食堂”品牌。"

    def test_verbatim_and_excerpt_pass(self):
        text = "某社区助餐点开业\n\n9月1日，某社区助餐点开业。居民王某说：“家门口就能吃上热乎饭”。"
        report = self.check(text, source=self.SOURCE)
        self.assertNotIn("BRF-003", ids(report))

    def test_changed_word_is_p0(self):
        text = "某社区助餐点开业\n\n9月1日，某社区助餐点开业。居民王某说：“家门口就能吃上热乎饭，孩子们都放心了。”"
        report = self.check(text, source=self.SOURCE)
        self.assertIn("BRF-003", ids(report, "P0"))

    def test_spliced_quote_is_p0(self):
        text = "某社区助餐点开业\n\n9月1日开业。居民王某说：“家门口就能吃上热乎饭，暖心食堂真好。”"
        report = self.check(text, source=self.SOURCE)
        self.assertIn("BRF-003", ids(report, "P0"))

    def test_punctuation_differences_ignored(self):
        text = "某社区助餐点开业\n\n9月1日开业。居民王某说：“家门口就能吃上热乎饭 孩子们也放心了！”"
        report = self.check(text, source=self.SOURCE)
        self.assertNotIn("BRF-003", ids(report))

    def test_unknown_name_term_is_p2_not_p0(self):
        text = "某社区助餐点开业\n\n9月1日，社区推出“银龄饭桌”服务品牌，每天供应午餐。"
        report = self.check(text, source=self.SOURCE)
        self.assertIn("BRF-003", ids(report, "P2"))
        self.assertNotIn("BRF-003", ids(report, "P0"))

    def test_without_source_gives_single_p2_hint(self):
        text = "某社区助餐点开业\n\n9月1日开业。居民王某说：“很方便。”老人李某说：“菜不错。”"
        report = self.check(text)
        hints = [it for it in report["issues"] if it["id"] == "BRF-003"]
        self.assertEqual(len(hints), 1)
        self.assertEqual(hints[0]["level"], "P2")
        self.assertTrue(any("--source" in note for note in report["notes"]))

    def test_speaker_missing(self):
        text = "某社区助餐点开业\n\n9月1日开业，“家门口就能吃上热乎饭，孩子们也放心了。”"
        report = self.check(text, source=self.SOURCE)
        self.assertIn("BRF-004", ids(report, "P2"))

    def test_speaker_with_colon_or_after(self):
        for text in (
            "标题\n\n9月1日开业。李大爷竖起大拇指：“家门口就能吃上热乎饭，孩子们也放心了。”",
            "标题\n\n9月1日开业。“家门口就能吃上热乎饭，孩子们也放心了。”居民王某说。",
        ):
            with self.subTest(text=text[-8:]):
                self.assertNotIn("BRF-004", ids(self.check(text, source=self.SOURCE)))

    def test_place_name_ending_with_dao_not_speaker(self):
        # “街道”的“道”不能当作说话动词
        text = "某街道“五个一”工作法见效\n\n今年以来，某街道推行工作法。"
        report = self.check(text, source=self.SOURCE)
        self.assertNotIn("BRF-003", ids(report, "P0"))

    def test_ascii_double_quotes_traced(self):
        # 直引号里的引语同样溯源
        text = '某区优化营商环境\n\n今年以来，某区压缩审批时限。某企业负责人表示："现在来区里办事就像回家一样方便。"'
        report = self.check(text, source=self.SOURCE)
        self.assertIn("BRF-003", ids(report, "P0"))
        traced = self.check('标题\n\n9月1日开业。居民王某说："家门口就能吃上热乎饭，孩子们也放心了。"', source=self.SOURCE)
        self.assertNotIn("BRF-003", ids(traced))

    def test_single_curly_quotes_traced(self):
        text = "标题\n\n9月1日开业。居民王某说：‘家门口就能吃上热乎饭，孩子们都放心了。’"
        self.assertIn("BRF-003", ids(self.check(text, source=self.SOURCE), "P0"))

    def test_short_quote_cut_from_narrative_is_p1(self):
        source = "镇长讲话，要求年底前全镇14个村都要完成村庄清洁行动的验收。"
        report = self.check("标题\n\n10月14日开会。镇长强调：“要完成。”", source=source)
        self.assertIn("BRF-003", ids(report, "P1"))
        self.assertNotIn("BRF-003", ids(report, "P0"))

    def test_short_quote_in_record_passes(self):
        # 采访记录“王某：……”格式，冒号后的话算原话
        source = "访谈对象：居民王某，72岁\n问：助餐点怎么样？\n王某：挺好的，方便。"
        report = self.check("标题\n\n9月1日开业。居民王某说：“挺好的，方便。”", source=source)
        self.assertNotIn("BRF-003", ids(report))

    def test_speaker_not_in_source_is_p1(self):
        source = "村民反映挺好的（没有具体记录谁说了什么）。"
        report = self.check("标题\n\n10月14日开会。村民王某说：“挺好的。”", source=source)
        hits = [it for it in report["issues"] if it["id"] == "BRF-003"]
        self.assertEqual([it["level"] for it in hits], ["P1"])
        self.assertIn("村民王某", hits[0]["message"])

    def test_speaker_swapped_long_quote_is_p1(self):
        report = self.check("标题\n\n9月1日开业。老人李某说：“家门口就能吃上热乎饭，孩子们也放心了。”", source=self.SOURCE)
        self.assertIn("BRF-003", ids(report, "P1"))

    def test_honor_without_source_is_p1(self):
        source = "东坪村今年拆除乱搭乱建38处。"
        report = self.check("标题\n\n今年以来，东坪村被评为“全国人居环境示范村”。", source=source)
        self.assertIn("BRF-003", ids(report, "P1"))
        listed = self.check("标题\n\n今年以来，东坪村被评为“全国人居环境示范村”。",
                            source="东坪村被评为“全国人居环境示范村”。")
        self.assertNotIn("BRF-003", ids(listed))

    def test_marker_inside_quote_ignored(self):
        text = "标题\n\n9月1日开业。居民说：“【待补：就餐老人原话，须本人同意引用】”"
        report = self.check(text, source=self.SOURCE)
        self.assertNotIn("BRF-003", ids(report))
        self.assertNotIn("BRF-004", ids(report))


# ---------------------------------------------------------------------------
# 导语、宣传用语、语体
# ---------------------------------------------------------------------------

class LeadAndWordingTests(BriefTestCase):

    def test_lead_with_time_and_no_padding_passes(self):
        report = self.check("标题\n\n今年7月以来，某街道在12个小区试行定时定点投放。")
        self.assertNotIn("BRF-005", ids(report))
        self.assertNotIn("BRF-006", ids(report))

    def test_lead_without_time(self):
        report = self.check("标题\n\n某街道在12个小区试行定时定点投放。")
        self.assertIn("BRF-006", ids(report, "P2"))

    def test_lead_time_marker_accepted(self):
        report = self.check("标题\n\n【待补：时间】，某街道在12个小区试行定时定点投放。")
        self.assertNotIn("BRF-006", ids(report))

    def test_news_dateline_and_byline_stripped(self):
        text = "某医院进社区义诊\n（通讯员　张某）\n本网讯　为进一步提升居民健康意识，10月15日，某医院开展义诊。"
        report = self.check(text, "新闻稿")
        self.assertIn("BRF-005", ids(report, "P1"))
        self.assertNotIn("BRF-006", ids(report))

    def test_lead_check_skipped_for_wechat_time(self):
        report = self.check("标题\n\n社区助餐点开门了，60岁以上老人每份12元。", "公众号")
        self.assertNotIn("BRF-006", ids(report))

    def test_absolute_claims_level_by_channel(self):
        text = "标题\n\n今年以来，某局该做法在全省首创，填补了本地空白。"
        internal = self.check(text, "信息")
        external = self.check(text, "信息", external=True)
        self.assertIn("BRF-007", ids(internal, "P2"))
        self.assertIn("BRF-007", ids(external, "P1"))

    def test_shou_li_is_p2_and_zui_mei_not_flagged(self):
        report = self.check("标题\n\n10月9日，某医院完成本院首例手术。某护士获评全国最美护士。", "新闻稿")
        levels = {it["excerpt"]: it["level"] for it in report["issues"] if it["id"] == "BRF-007"}
        self.assertEqual(levels.get("首例"), "P2")
        self.assertNotIn("全国最美", "".join(levels))

    def test_cured_discharge_not_flagged(self):
        report = self.check("标题\n\n10月15日，患者康复出院，家属送来锦旗。", "新闻稿")
        self.assertNotIn("BRF-008", ids(report))

    def test_efficacy_internal_is_p1(self):
        report = self.check("标题\n\n今年以来，该技术治愈率达98%。", "信息")
        self.assertIn("BRF-008", ids(report, "P1"))

    def test_official_phrases_not_efficacy(self):
        # 根治欠薪是工作提法；“更有效率”指效率，不是有效率指标
        report = self.check("某区开展根治欠薪冬季专项行动\n\n11月1日，某区开展根治欠薪冬季专项行动，让维权更有效率。", "新闻稿")
        self.assertNotIn("BRF-008", ids(report))

    def test_non_medical_efficacy_word_is_p2(self):
        report = self.check("标题\n\n11月1日，某区推行“一案一专班”，整改立竿见影。", "新闻稿")
        self.assertIn("BRF-008", ids(report, "P2"))
        self.assertNotIn("BRF-008", ids(report, "P0"))
        self.assertNotIn("BRF-008", ids(report, "P1"))

    def test_medical_context_efficacy_is_p0(self):
        report = self.check("标题\n\n10月15日，某医院开展义诊，专家介绍该手术可以根治房颤，疗效立竿见影。", "新闻稿")
        self.assertEqual(ids(report, "P0").count("BRF-008"), 2)

    def test_unit_hospital_makes_whole_text_medical(self):
        text = "标题\n\n10月15日，某中心开展义诊，效果立竿见影。"
        self.assertIn("BRF-008", ids(self.check(text, "新闻稿"), "P2"))
        self.assertIn("BRF-008", ids(self.check(text, "新闻稿", unit="公立医院"), "P0"))

    def test_efficacy_phrase_in_source_downgraded(self):
        source = "医院党委开展根治医药购销领域突出问题专项行动。"
        text = "标题\n\n10月15日，某医院党委开展根治医药购销领域突出问题专项行动。"
        report = self.check(text, "新闻稿", source=source)
        self.assertNotIn("BRF-008", ids(report, "P0"))
        self.assertIn("BRF-008", ids(report, "P1"))

    def test_efficacy_rate_needs_rate_context(self):
        self.assertIn("BRF-008", ids(self.check("标题\n\n今年以来，某医院该疗法总有效率达92%。", "新闻稿"), "P0"))
        self.assertNotIn("BRF-008", ids(self.check("标题\n\n今年以来，某医院门诊流程调整后工作更有效率。", "新闻稿")))

    def test_surgery_names_not_efficacy(self):
        text = "标题\n\n10月15日，某医院完成3台腹腔镜胃癌根治术，1台根治性切除术。"
        self.assertNotIn("BRF-008", ids(self.check(text, "新闻稿", unit="公立医院")))
        claim = self.check("标题\n\n10月15日，某医院介绍该手术可以根治胃癌。", "新闻稿", unit="公立医院")
        self.assertIn("BRF-008", ids(claim, "P0"))

    def test_edu_critical_context_downgraded(self):
        text = "标题\n\n近期调研发现，部分学校片面追求升学率，仍按升学率考核教师。"
        for external in (False, True):
            with self.subTest(external=external):
                report = self.check(text, "信息", external=external)
                self.assertIn("BRF-009", ids(report, "P2"))
                self.assertNotIn("BRF-009", ids(report, "P1"))

    def test_edu_hype_internal_is_p2_external_p1(self):
        text = "标题\n\n7月，我校高考本科上线率达95%，升学率排名全市第一。"
        self.assertNotIn("BRF-009", ids(self.check(text, "信息"), "P1"))
        self.assertIn("BRF-009", ids(self.check(text, "信息", external=True), "P1"))

    def test_empty_title_words(self):
        report = self.check("某区多措并举扎实推进营商环境优化取得明显成效\n\n今年以来，某区优化营商环境。")
        words = [it["excerpt"] for it in report["issues"] if it["id"] == "BRF-016" and it["level"] == "P1"]
        self.assertEqual(words, ["多措并举", "扎实推进", "取得明显成效"])
        clean = brief_check.build_report(SPECIMENS / "pass-信息.txt", "信息")
        self.assertNotIn("BRF-016", ids(clean))

    def test_empty_title_words_in_brief_sections(self):
        text = (
            "工作简报\n第3期\n某局办公室编　2026年10月20日\n\n"
            "某局扎实推进报表清理\n\n10月以来，某局合并报表12张。\n\n报：区委办公室\n"
        )
        self.assertIn("BRF-016", ids(self.check(text, "简报"), "P1"))

    def test_zou_zai_qian_lie_flagged(self):
        report = self.check("标题\n\n今年以来，某区营商环境走在了全国前列。", "信息", external=True)
        self.assertIn("BRF-007", ids(report, "P1"))

    def test_skill_champion_not_edu_hype(self):
        report = self.check("标题\n\n10月12日，公司举办技能比武，评出焊工技术状元3名。", "新闻稿")
        self.assertNotIn("BRF-009", ids(report))

    def test_exclamation_limit(self):
        ok = self.check("标题\n\n9月1日开业。欢迎大家来就餐！", "公众号")
        bad = self.check("标题\n\n9月1日开业！欢迎大家来就餐！", "公众号")
        self.assertNotIn("BRF-011", ids(ok))
        self.assertIn("BRF-011", ids(bad))

    def test_exclamation_inside_quotes_not_counted(self):
        source = "老人李某说：“饭菜可口！价钱实惠！”"
        text = "标题\n\n9月1日开业。老人李某说：“饭菜可口！价钱实惠！”欢迎大家来就餐！"
        report = self.check(text, "公众号", source=source)
        self.assertNotIn("BRF-011", ids(report))

    def test_body_double_exclamation_is_bait(self):
        report = self.check("标题\n\n9月1日开业，欢迎大家来就餐！！", "公众号")
        self.assertIn("BRF-010", ids(report, "P1"))

    def test_slang_emo_needs_word_boundary(self):
        clean = self.check("标题\n\n10月12日，团队完成系统demo演示，会议memo已归档。", "公众号")
        self.assertNotIn("BRF-010", ids(clean))
        slang = self.check("标题\n\n10月12日开课，再也不用emo了。", "公众号")
        self.assertIn("BRF-010", ids(slang, "P1"))

    def test_quanshi_shuaixian_flagged(self):
        report = self.check("标题\n\n今年以来，某区在全市率先实现“零材料”办理。", "新闻稿")
        self.assertIn("BRF-007", ids(report, "P1"))
        plain = self.check("标题\n\n今年以来，某街道党员干部率先垂范。", "新闻稿")
        self.assertNotIn("BRF-007", ids(plain))

    def test_platform_online_rate_not_edu_hype(self):
        report = self.check("标题\n\n今年以来，政务服务平台上线率100%。", "新闻稿")
        self.assertNotIn("BRF-009", ids(report))
        hype = self.check("标题\n\n7月，我校高考本科上线率达95%。", "新闻稿")
        self.assertIn("BRF-009", ids(hype, "P1"))

    def test_place_name_with_tai_not_bait(self):
        report = self.check("太原路社区助餐点开业了\n\n9月1日，太原路社区助餐点开业。", "公众号")
        self.assertNotIn("BRF-010", ids(report))

    def test_photo_reminder_only_external(self):
        text = "标题\n\n10月15日义诊。图为医护人员为居民测血压。"
        self.assertIn("BRF-015", ids(self.check(text, "新闻稿")))
        self.assertNotIn("BRF-015", ids(self.check(text, "信息")))


# ---------------------------------------------------------------------------
# 简报
# ---------------------------------------------------------------------------

class BriefFormatTests(BriefTestCase):

    def test_issue_number_marker_accepted(self):
        text = (
            "工作简报\n第【待补：期号】期\n某局办公室编　　2026年10月20日\n\n"
            "某局推行一表通\n\n10月以来，某局推行一表通。\n\n报：区委办公室\n"
        )
        report = self.check(text, "简报")
        for rule in ("BRF-012", "BRF-013", "BRF-014"):
            self.assertNotIn(rule, ids(report))

    def test_each_article_title_checked(self):
        report = brief_check.build_report(SPECIMENS / "fail-简报标题数字.txt", "简报")
        hits = [it for it in report["issues"] if it["id"] == "BRF-002"]
        self.assertEqual(len(hits), 1)
        self.assertIn("五项", hits[0]["excerpt"])

    def test_brief_sections_split(self):
        text = (
            "工作简报\n第3期\n某局办公室编　2026年10月20日\n\n工作动态\n某局推行三项举措\n"
            "——减少基层报表\n\n10月以来，某局推行：一是合并报表；二是数据共享；三是限时报送。\n\n"
            "（某街道供稿）\n\n某街道建立两项清单\n\n10月起，某街道建立准入清单。\n\n报：区委办公室\n"
        )
        draft = brief_check.Draft(brief_check.body_lines(text.split("\n")), "简报")
        sections = draft.sections()
        self.assertEqual(len(sections), 2)
        self.assertEqual(len(sections[0].titles), 3)
        report = self.check(text, "简报")
        hits = [it for it in report["issues"] if it["id"] == "BRF-002"]
        self.assertEqual([it["level"] for it in hits], ["P2"])
        self.assertIn("两项", hits[0]["excerpt"])

    def test_editor_and_date_on_separate_lines(self):
        text = "工作简报\n第3期\n某局办公室编\n2026年10月20日\n\n某局推行一表通\n\n10月以来推行。\n\n发：各科室\n"
        report = self.check(text, "简报")
        self.assertNotIn("BRF-013", ids(report))
        self.assertNotIn("某局办公室编", [t for _, t in brief_check.Draft(
            brief_check.body_lines(text.split("\n")), "简报").paragraphs])


# ---------------------------------------------------------------------------
# 命令行与输入
# ---------------------------------------------------------------------------

class CliTests(BriefTestCase):

    def test_json_output_fields(self):
        code, out, _ = self.run_main(str(SPECIMENS / "pass-信息.txt"), "--json")
        self.assertEqual(code, 0)
        data = json.loads(out)
        for key in ("file", "type", "external", "char_count", "length_range", "title", "summary", "issues", "notes"):
            self.assertIn(key, data)
        for key in ("P0", "P1", "P2"):
            self.assertIn(key, data["summary"])

    def test_strict_returns_1_on_p0(self):
        code, _, _ = self.run_main(str(SPECIMENS / "fail-疗效保证.txt"), "--type", "新闻稿", "--strict")
        self.assertEqual(code, 1)

    def test_strict_returns_0_without_p0(self):
        code, _, _ = self.run_main(str(SPECIMENS / "fail-导语铺垫.txt"), "--strict")
        self.assertEqual(code, 0)

    def test_text_output(self):
        code, out, _ = self.run_main(str(SPECIMENS / "fail-导语铺垫.txt"))
        self.assertEqual(code, 0)
        self.assertIn("BRF-005", out)
        self.assertIn("依据", out)

    def test_missing_file_returns_2(self):
        code, _, err = self.run_main(str(self.tmp / "不存在.txt"))
        self.assertEqual(code, 2)
        self.assertIn("文件不存在", err)

    def test_unsupported_type_returns_2(self):
        code, _, err = self.run_main(str(SPECIMENS / "pass-信息.txt"), "--type", "通知")
        self.assertEqual(code, 2)
        self.assertIn("不支持的稿件类型", err)

    def test_unit_option(self):
        path = self.write("正文.txt", "标题\n\n10月15日，某中心开展义诊，效果立竿见影。")
        code, out, _ = self.run_main(str(path), "--type", "新闻稿", "--unit", "公立医院", "--json")
        self.assertEqual(code, 0)
        self.assertEqual(json.loads(out)["summary"]["P0"], 1)

    def test_type_alias(self):
        code, out, _ = self.run_main(str(SPECIMENS / "pass-推文.txt"), "--type", "公众号推文", "--json")
        self.assertEqual(code, 0)
        self.assertEqual(json.loads(out)["type"], "公众号")

    def test_unsupported_suffix_returns_2(self):
        path = self.write("正文.pdf", "x")
        code, _, err = self.run_main(str(path))
        self.assertEqual(code, 2)
        self.assertIn("不支持的文件类型", err)

    def test_empty_body_returns_2(self):
        path = self.write("正文.txt", "\n\n【起草说明】\n只有说明")
        code, _, err = self.run_main(str(path))
        self.assertEqual(code, 2)
        self.assertIn("正文为空", err)

    def test_min_greater_than_max_returns_2(self):
        code, _, _ = self.run_main(str(SPECIMENS / "pass-信息.txt"), "--min", "900", "--max", "800")
        self.assertEqual(code, 2)

    def test_gb18030_file(self):
        path = self.write("正文.txt", "标题\n\n今年7月以来，某街道试行定时定点投放。", encoding="gb18030")
        code, out, _ = self.run_main(str(path), "--json")
        self.assertEqual(code, 0)
        self.assertEqual(json.loads(out)["title"], "标题")

    def test_help_is_chinese(self):
        code, out, _ = self.run_main("--help")
        self.assertEqual(code, 0)
        self.assertIn("用法：", out)

    def test_does_not_write_files(self):
        before = sorted(p.name for p in SPECIMENS.iterdir())
        self.run_main(str(SPECIMENS / "pass-信息.txt"), "--source", str(SPECIMENS / "pass-材料.txt"))
        self.assertEqual(before, sorted(p.name for p in SPECIMENS.iterdir()))

    @unittest.skipUnless(HAS_DOCX, "未安装 python-docx")
    def test_docx_input(self):
        import docx as _docx  # type: ignore
        document = _docx.Document()
        document.add_paragraph("某医院进社区义诊")
        document.add_paragraph("10月15日，某医院开展义诊，介入治疗治愈率达98%。")
        path = self.tmp / "正文.docx"
        document.save(str(path))
        report = brief_check.build_report(path, "新闻稿")
        self.assertIn("BRF-008", ids(report, "P0"))


if __name__ == "__main__":
    unittest.main()
