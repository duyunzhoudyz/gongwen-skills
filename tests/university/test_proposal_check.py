# -*- coding: utf-8 -*-
"""skills/gongwen-university/scripts/proposal_check.py 的单元测试。

运行：python -m unittest discover -s tests/university -t .

samples/ 下有四份虚构标本：
    活页_必错.txt、申请书_必错.txt：每类问题至少命中一次；
    活页_必过.txt、申请书_必过.txt：规范写法，不得报任何问题。
"""
from __future__ import annotations

import importlib.util
import io
import json
import shutil
import sys
import tempfile
import unittest
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[2]
SCRIPT = REPO_ROOT / "skills" / "gongwen-university" / "scripts" / "proposal_check.py"
SAMPLES = Path(__file__).resolve().parent / "samples"

NAMES = "王明远,澜江理工大学"
TITLE = "数字乡村治理中的多元协同机制研究"


def load_module():
    name = "proposal_check_under_test"
    if name in sys.modules:
        return sys.modules[name]
    spec = importlib.util.spec_from_file_location(name, SCRIPT)
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)  # type: ignore[union-attr]
    return module


pc = load_module()

try:
    import docx  # type: ignore  # noqa: F401
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False


class ProposalCheckCase(unittest.TestCase):

    def setUp(self) -> None:
        self.tmp = Path(tempfile.mkdtemp(prefix="proposal_check_test_"))

    def tearDown(self) -> None:
        shutil.rmtree(self.tmp, ignore_errors=True)

    def write(self, name: str, text: str, encoding: str = "utf-8") -> Path:
        path = self.tmp / name
        path.write_bytes(text.encode(encoding))
        return path

    def run_main(self, *args: str):
        out, err = io.StringIO(), io.StringIO()
        with redirect_stdout(out), redirect_stderr(err):
            code = pc.main(list(args))
        return code, out.getvalue(), err.getvalue()

    def run_json(self, *args: str) -> dict:
        code, out, _err = self.run_main(*args, "--json")
        self.assertIn(code, (0, 1))
        return json.loads(out)

    @staticmethod
    def ids(report: dict):
        return [issue["id"] for issue in report["issues"]]


class TestMustFailSamples(ProposalCheckCase):
    """必错标本：每类问题都要报出来。"""

    def test_blind_sample_reports_every_category(self):
        report = self.run_json(str(SAMPLES / "活页_必错.txt"), "--blind", "--names", NAMES,
                               "--max-chars", "7000", "--section", "选题说明=300", "--title", TITLE)
        ids = self.ids(report)
        for expected in ("PRC-ID-001", "PRC-ID-002", "PRC-ID-003", "PRC-ID-004", "PRC-ID-007", "PRC-PRV-001",
                         "PRC-LEN-002", "PRC-TTL-001", "PRC-FIL-001", "PRC-FIL-002", "PRC-FIL-003",
                         "PRC-RES-001", "PRC-RES-002"):
            self.assertIn(expected, ids, expected)
        self.assertGreaterEqual(report["summary"]["P0"], 4)
        # 同一行的两个身份词合并为一条
        self.assertEqual(ids.count("PRC-ID-001"), 1)

    def test_output_never_repeats_identity_or_numbers(self):
        code, out, _err = self.run_main(str(SAMPLES / "活页_必错.txt"), "--blind", "--names", NAMES, "--json")
        self.assertEqual(code, 0)
        for secret in ("王明远", "澜江理工大学", "13912345678", "wmy@example.edu.cn"):
            self.assertNotIn(secret, out)
        code, text_out, _err = self.run_main(str(SAMPLES / "活页_必错.txt"), "--blind", "--names", NAMES)
        for secret in ("王明远", "澜江理工大学", "13912345678", "23BSH045", "基层数字治理"):
            self.assertNotIn(secret, text_out)


class TestRegressions(ProposalCheckCase):
    """独立评审发现的问题：复合词误报、邮箱片段外泄、空文件判为通过、前期成果自述漏报。"""

    def test_compound_words_not_self_reference(self):
        text = ("选题依据\n调查在样本单位和样本校中开展，样本人群覆盖基层干部。参考日本人口减少地区的经验，"
                "分析运行成本所占比重。政策文本院际差异和版本校对由编码员完成，保存自我校正记录。"
                "研究不讨论基本人权和资本所得分配问题。\n")
        report = self.run_json(str(self.write("复合词.txt", text)), "--blind")
        self.assertNotIn("PRC-ID-002", self.ids(report))

    def test_real_self_reference_still_reported(self):
        text = "选题依据\n前期已完成本校教师的问卷调查。同样本人也参与了访谈。我校设有研究平台。\n"
        report = self.run_json(str(self.write("自称.txt", text)), "--blind")
        hits = [i["message"] for i in report["issues"] if i["id"] == "PRC-ID-002"]
        self.assertEqual(len(hits), 3, hits)

    def test_wosuo_particle_not_self_reference(self):
        text = "选题依据\n据我所知，已有研究多为个案。我所见的文献缺少比较，我所谓协同指跨部门合作。\n"
        report = self.run_json(str(self.write("我所.txt", text)), "--blind")
        self.assertNotIn("PRC-ID-002", self.ids(report))
        # "我所在单位""我所的调查"仍是自称
        text = "选题依据\n我所在单位开展过调查。我所的调查显示协同不足。\n"
        report = self.run_json(str(self.write("我所自称.txt", text)), "--blind")
        hits = [i for i in report["issues"] if i["id"] == "PRC-ID-002"]
        self.assertEqual(len(hits), 2, hits)

    def test_blind_placeholder_downgraded(self):
        text = "研究内容\n本课题选取××省××县的12个村，访谈某某乡干部。模板提示XXX未删。\n"
        path = self.write("占位.txt", text)
        blind = [i for i in self.run_json(str(path), "--blind")["issues"] if i["id"] == "PRC-FIL-002"]
        levels = sorted(i["level"] for i in blind)
        # ××省、××县、某某降为 P2 并提示可保留；XXX 仍是 P1
        self.assertEqual(levels, ["P1", "P2", "P2", "P2"], blind)
        self.assertTrue(all("有意匿名可保留" in i["suggestion"] for i in blind if i["level"] == "P2"))
        plain = [i for i in self.run_json(str(path))["issues"] if i["id"] == "PRC-FIL-002"]
        self.assertTrue(all(i["level"] == "P1" for i in plain))

    def test_org_excerpt_starts_at_org_name(self):
        text = ("选题依据\n本研究借鉴了丁省师范大学的做法，参考甲省社会科学院和乙市大学的数据。"
                "我国的大学普遍如此。\n")
        report = self.run_json(str(self.write("机构摘录.txt", text)), "--blind")
        org = [i for i in report["issues"] if i["id"] == "PRC-ID-006"][0]
        self.assertEqual(org["excerpt"], "丁省师范大学、甲省社会科学院、乙市大学")
        self.assertIn("3 处", org["message"])

    def test_email_after_name_never_leaks(self):
        text = ("预期成果\n研究报告一份。课题组联系人王明远，邮箱wmy2026@example.edu.cn，"
                "主页 https://www.example.edu.cn/~wmy 。\n")
        path = self.write("邮箱.txt", text)
        _code, json_out, _err = self.run_main(str(path), "--blind", "--names", NAMES, "--json")
        _code, text_out, _err = self.run_main(str(path), "--blind", "--names", NAMES)
        for out in (json_out, text_out):
            for fragment in ("wmy", "@", "example", "edu.cn", "王明远"):
                self.assertNotIn(fragment, out)
        ids = self.ids(json.loads(json_out))
        for expected in ("PRC-ID-001", "PRC-ID-004", "PRC-ID-005"):
            self.assertIn(expected, ids)

    def test_number_excerpt_does_not_leak_email(self):
        path = self.write("号码.txt", "预期成果\n电话13912345678，邮箱wmy@example.edu.cn\n")
        _code, out, _err = self.run_main(str(path), "--blind", "--json")
        self.assertIn("PRC-PRV-001", out)
        for fragment in ("wmy", "example", "13912345678"):
            self.assertNotIn(fragment, out)

    def test_prior_work_reported_and_title_masked(self):
        text = "选题依据\n前期出版专著《乡村治理新论》，主持国家社会科学基金一般项目一项。\n"
        path = self.write("前期成果.txt", text)
        report = self.run_json(str(path), "--blind")
        hits = [i for i in report["issues"] if i["id"] == "PRC-ID-007"]
        self.assertEqual(len(hits), 2)
        self.assertTrue(all(i["level"] == "P2" for i in hits))
        _code, out, _err = self.run_main(str(path), "--blind")
        self.assertNotIn("乡村治理新论", out)
        # 非匿名材料不查前期成果
        self.assertNotIn("PRC-ID-007", self.ids(self.run_json(str(path))))

    def test_empty_file_is_input_error(self):
        for name, text in (("空.txt", ""), ("空白.txt", " \n　\n\n"), ("只有说明.txt", "【起草说明】\n本人补充。\n")):
            code, out, err = self.run_main(str(self.write(name, text)), "--blind")
            self.assertEqual(code, 2, name)
            self.assertIn("没有可检查的内容", err)
            self.assertNotIn("未发现形式问题", out)

    def test_application_sample(self):
        report = self.run_json(str(SAMPLES / "申请书_必错.txt"), "--title", TITLE)
        ids = self.ids(report)
        self.assertIn("PRC-TTL-001", ids)
        self.assertIn("PRC-PRV-002", ids)
        self.assertIn("PRC-FIL-003", ids)
        self.assertNotIn("PRC-ID-001", ids)  # 非匿名材料不查身份词

    def test_strict_exit_code(self):
        code, _out, _err = self.run_main(str(SAMPLES / "申请书_必错.txt"), "--title", TITLE, "--strict")
        self.assertEqual(code, 1)


class TestMustPassSamples(ProposalCheckCase):
    """必过标本：规范写法不得报任何问题。"""

    def test_blind_sample_clean(self):
        report = self.run_json(str(SAMPLES / "活页_必过.txt"), "--blind", "--names", NAMES,
                               "--max-chars", "7000", "--section", "选题说明=300", "--title", TITLE)
        self.assertEqual(report["issues"], [])
        names = [row["name"] for row in report["sections"]]
        self.assertEqual(names, ["选题说明", "选题依据", "研究内容", "创新之处", "预期成果", "参考文献"])

    def test_application_sample_clean(self):
        report = self.run_json(str(SAMPLES / "申请书_必过.txt"), "--title", TITLE)
        self.assertEqual(report["issues"], [])

    def test_strict_exit_zero_when_clean(self):
        code, _out, _err = self.run_main(str(SAMPLES / "活页_必过.txt"), "--blind", "--strict")
        self.assertEqual(code, 0)


class TestCounting(ProposalCheckCase):

    def test_count_modes(self):
        counts = pc.count_text("研究 ABC 2026年。")
        self.assertEqual(counts["hanzi"], 3)
        self.assertEqual(counts["chars"], 11)
        # 汉字3 + 句号1 + 西文词2
        self.assertEqual(counts["word"], 6)

    def test_max_chars(self):
        path = self.write("长稿.txt", "选题说明\n" + "研究" * 100 + "\n")
        report = self.run_json(str(path), "--max-chars", "150")
        self.assertIn("PRC-LEN-001", self.ids(report))
        report = self.run_json(str(path), "--max-chars", "300")
        self.assertNotIn("PRC-LEN-001", self.ids(report))

    def test_hanzi_mode_is_looser_than_chars(self):
        path = self.write("混排.txt", "选题说明\n研究，研究，研究。\n")
        chars = self.run_json(str(path), "--section", "选题说明=7", "--count", "chars")
        hanzi = self.run_json(str(path), "--section", "选题说明=7", "--count", "hanzi")
        self.assertIn("PRC-LEN-002", self.ids(chars))
        self.assertNotIn("PRC-LEN-002", self.ids(hanzi))

    def test_missing_section_is_p2(self):
        path = self.write("无栏目.txt", "研究内容\n正文。\n")
        report = self.run_json(str(path), "--section", "选题说明=300")
        issue = [i for i in report["issues"] if i["id"] == "PRC-LEN-003"][0]
        self.assertEqual(issue["level"], "P2")


class TestHeadings(ProposalCheckCase):

    def test_heading_with_note_and_serial(self):
        self.assertEqual(pc.match_heading("三、研究内容（限3000字）", pc.KNOWN_SECTIONS), ("研究内容", ""))
        self.assertEqual(pc.match_heading("（二）创新之处", pc.KNOWN_SECTIONS), ("创新之处", ""))
        self.assertEqual(pc.match_heading("1.选题依据", pc.KNOWN_SECTIONS), ("选题依据", ""))

    def test_inline_heading_needs_serial(self):
        self.assertEqual(pc.match_heading("一、选题说明：本课题研究协同问题", pc.KNOWN_SECTIONS),
                         ("选题说明", "本课题研究协同问题"))
        self.assertIsNone(pc.match_heading("研究方法：采用问卷调查和访谈", pc.KNOWN_SECTIONS))
        self.assertEqual(pc.match_heading("研究方法：采用问卷", pc.KNOWN_SECTIONS, ["研究方法"]),
                         ("研究方法", "采用问卷"))

    def test_inline_heading_without_serial_counts_into_previous_section(self):
        # 不带序号的"研究内容：……"计入上一栏目；用 --section 指定后单独统计（见 docstring"栏目识别"）
        path = self.write("同行栏目.txt", "一、选题说明：本课题研究协同问题。\n研究内容：分三部分展开。\n")
        rows = {row["name"]: row["count"] for row in self.run_json(str(path))["sections"]}
        self.assertEqual(list(rows), ["选题说明"])
        rows = {row["name"]: row["count"] for row in self.run_json(str(path), "--section", "研究内容=3000")["sections"]}
        self.assertEqual(list(rows), ["选题说明", "研究内容"])
        self.assertLess(rows["选题说明"], 15)

    def test_body_sentence_is_not_heading(self):
        self.assertIsNone(pc.match_heading("研究内容包括三个方面。", pc.KNOWN_SECTIONS))

    def test_parent_section_empty_not_reported(self):
        path = self.write("上级标题.txt", "课题设计论证\n选题依据\n已有研究不足。\n")
        report = self.run_json(str(path))
        self.assertNotIn("PRC-FIL-003", self.ids(report))


class TestScopeAndFormats(ProposalCheckCase):

    def test_stops_at_draft_note(self):
        path = self.write("带说明.txt", "选题说明\n本课题研究协同问题。\n【起草说明】\n本人补充：XXX\n")
        report = self.run_json(str(path), "--blind")
        self.assertEqual(report["issues"], [])
        self.assertEqual(report["scope"], "【起草说明】之前")

    def test_markdown_not_checked_for_md_input(self):
        text = "选题说明\n**重点**是测量。\n"
        md = self.run_json(str(self.write("稿.md", text)))
        txt = self.run_json(str(self.write("稿.txt", text)))
        self.assertNotIn("PRC-RES-002", self.ids(md))
        self.assertIn("PRC-RES-002", self.ids(txt))

    def test_gb18030_input(self):
        path = self.write("国标编码.txt", "选题说明\n本人前期研究。\n", encoding="gb18030")
        code, out, err = self.run_main(str(path), "--blind", "--json")
        self.assertEqual(code, 0)
        self.assertIn("GB18030", err)
        self.assertIn("PRC-ID-002", out)

    def test_org_names_aggregated_and_daxuesheng_ignored(self):
        path = self.write("机构.txt", "选题依据\n某省社会科学研究所和甲大学的调查显示，大学生返乡创业增多。\n另有丙省大学的研究。\n")
        report = self.run_json(str(path), "--blind")
        orgs = [i for i in report["issues"] if i["id"] == "PRC-ID-006"]
        self.assertEqual(len(orgs), 1)
        self.assertEqual(orgs[0]["level"], "P2")
        self.assertNotIn("大学生", orgs[0]["excerpt"])
        self.assertIn("3 处", orgs[0]["message"])

    def test_names_file(self):
        names = self.write("名单.txt", "王明远\n澜江理工大学\n")
        path = self.write("活页.txt", "选题依据\n王明远团队的调查显示协同不足。\n")
        report = self.run_json(str(path), "--blind", "--names-file", str(names))
        self.assertIn("PRC-ID-001", self.ids(report))

    def test_labeled_number_fields_allowed_when_not_blind(self):
        path = self.write("申请书.txt", "联系电话：13912345678\n身份证号：110101199001011234\n")
        report = self.run_json(str(path))
        self.assertEqual(report["issues"], [])

    @unittest.skipUnless(HAS_DOCX, "未安装 python-docx")
    def test_docx_input_with_table(self):
        import docx as docx_mod  # type: ignore
        document = docx_mod.Document()
        document.add_paragraph("选题依据")
        document.add_paragraph("本人前期研究显示协同不足。")
        table = document.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "课题名称"
        table.rows[0].cells[1].text = "数字乡村治理协同研究"
        path = self.tmp / "活页.docx"
        document.save(str(path))
        report = self.run_json(str(path), "--blind", "--title", TITLE)
        ids = self.ids(report)
        self.assertIn("PRC-ID-002", ids)
        self.assertIn("PRC-TTL-001", ids)


class TestInputErrors(ProposalCheckCase):

    def test_missing_file(self):
        code, _out, err = self.run_main(str(self.tmp / "不存在.txt"))
        self.assertEqual(code, 2)
        self.assertIn("文件不存在", err)

    def test_unsupported_suffix(self):
        path = self.write("稿.pdf", "x")
        code, _out, _err = self.run_main(str(path))
        self.assertEqual(code, 2)

    def test_bad_section_argument(self):
        path = self.write("稿.txt", "选题说明\n内容\n")
        for bad in ("选题说明", "选题说明=abc", "=300", "选题说明=0"):
            code, _out, _err = self.run_main(str(path), "--section", bad)
            self.assertEqual(code, 2, bad)

    def test_bad_max_chars(self):
        path = self.write("稿.txt", "内容\n")
        code, _out, _err = self.run_main(str(path), "--max-chars", "0")
        self.assertEqual(code, 2)

    def test_help(self):
        code, out, _err = self.run_main("--help")
        self.assertEqual(code, 0)
        self.assertIn("用法", out)


if __name__ == "__main__":
    unittest.main()
