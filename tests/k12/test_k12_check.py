# -*- coding: utf-8 -*-
"""skills/gongwen-k12/scripts/k12_check.py 的单元测试。

运行：python -m unittest discover -s tests/k12 -t .

samples/ 下的标本：
    comments_must_fail.txt、parents_must_fail.txt  必错标本：每类问题都要报出；
    parents_must_fail_lines.txt                    必错标本：每行一句不写主语、语序变化的违规写法，逐行都要报出；
    parents_must_fail_speech.txt                   必错标本：送审的家长会讲话稿（同评测第12条材料）；
    comments_must_pass.txt、parents_must_pass.txt  必过标本：不得报出任何问题；
    comments_facts.md                              必过评语对应的脱敏事实表；
    students_lines.txt                             含姓名、手机号、学号的虚构事实表，用于脱敏测试。
"""
from __future__ import annotations

import csv
import io
import json
import shutil
import sys
import tempfile
import unittest
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[2]
SCRIPTS_DIR = REPO_ROOT / "skills" / "gongwen-k12" / "scripts"
SAMPLES = Path(__file__).resolve().parent / "samples"
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

import k12_check  # noqa: E402

try:
    import openpyxl  # type: ignore  # noqa: F401
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

try:
    import docx  # type: ignore  # noqa: F401
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

NAMES = ("王小明", "李华", "欧阳晓晓", "小明", "晓晓")


class K12TestCase(unittest.TestCase):

    def setUp(self) -> None:
        self.tmp = Path(tempfile.mkdtemp(prefix="k12_check_test_"))

    def tearDown(self) -> None:
        shutil.rmtree(self.tmp, ignore_errors=True)

    def write(self, name: str, text: str) -> Path:
        path = self.tmp / name
        path.write_text(text, encoding="utf-8")
        return path

    def copy(self, name: str) -> Path:
        target = self.tmp / name
        shutil.copyfile(SAMPLES / name, target)
        return target

    def run_main(self, *args: str):
        out, err = io.StringIO(), io.StringIO()
        with redirect_stdout(out), redirect_stderr(err):
            code = k12_check.main(list(args))
        return code, out.getvalue(), err.getvalue()

    def run_json(self, *args: str) -> dict:
        code, out, err = self.run_main(*args, "--json")
        self.assertIn(code, (0, 1), err)
        return json.loads(out)

    @staticmethod
    def ids(report: dict, level: str = "") -> list:
        return [i["id"] for i in report["issues"] if not level or i["level"] == level]


class CommentsTest(K12TestCase):

    def test_must_pass_sample_has_no_issue(self):
        report = self.run_json("comments", str(SAMPLES / "comments_must_pass.txt"),
                               "--facts", str(SAMPLES / "comments_facts.md"))
        self.assertEqual(report["entry_count"], 8, "【起草说明】之后的内容不应计入")
        self.assertEqual(report["issues"], [], report["issues"])

    def test_must_fail_sample_reports_each_problem(self):
        report = self.run_json("comments", str(SAMPLES / "comments_must_fail.txt"))
        found = set(self.ids(report))
        for rule in ("PY-001", "PY-002", "PY-003", "PY-004", "PY-005", "PY-006",
                     "PY-007", "PY-009", "PY-010", "PY-011", "PY-014"):
            self.assertIn(rule, found, rule)
        by_label = {}
        for item in report["issues"]:
            by_label.setdefault(item["label"], set()).add(item["id"])
        self.assertIn("PY-003", by_label["S03"])
        self.assertIn("PY-006", by_label["S05"])
        self.assertTrue(any(p["a"] == "S01" and p["b"] == "S02" for p in report["similar_pairs"]))

    def test_exam_rank_is_p0_and_competition_rank_is_not_reported(self):
        path = self.write("评语.txt",
                          "S01：你在校运会跳远比赛中获得第一名，赛前每天练习起跳，老师看在眼里。希望你把这股劲用到英语朗读上。\n"
                          "S02：你期中考试语文考了年级前三名，老师为你高兴。希望你下学期多读整本书，写作会更有底气。\n")
        report = self.run_json("comments", str(path), "--min-chars", "0")
        s01 = [i for i in report["issues"] if i["label"] == "S01"]
        s02 = [i for i in report["issues"] if i["label"] == "S02" and i["id"] == "PY-003"]
        self.assertEqual(s01, [])
        self.assertEqual(s02[0]["level"], "P0")

    def test_high_school_stage_lowers_rank_level(self):
        path = self.write("评语.txt", "S01：你期中考试数学考了全班第一名，课上常常帮同学讲题。希望你把错题本坚持下去，物理的计算会更稳。\n")
        report = self.run_json("comments", str(path), "--stage", "高中")
        rank = [i for i in report["issues"] if i["id"] == "PY-003"]
        self.assertEqual(rank[0]["level"], "P1")

    def test_strict_returns_1_on_p0(self):
        code, _out, _err = self.run_main("comments", str(SAMPLES / "comments_must_fail.txt"), "--strict")
        self.assertEqual(code, 1)
        code, _out, _err = self.run_main("comments", str(SAMPLES / "comments_must_pass.txt"), "--strict")
        self.assertEqual(code, 0)

    def test_marker_entry_is_p2_and_counted(self):
        path = self.write("评语.txt",
                          "S01：【待补：该生本学期的两三件具体表现】\n"
                          "S02：你每天早上第一个到教室开窗通风，还主动给班里的绿萝浇水。希望你上课多举手，把想法说出来。\n")
        report = self.run_json("comments", str(path))
        self.assertIn("PY-013", self.ids(report, "P2"))
        self.assertEqual(report["summary"]["markers"]["待补"], 1)
        self.assertEqual(report["summary"]["P0"], 0)

    def test_marker_text_does_not_trigger_person_check(self):
        path = self.write("评语.txt",
                          "S01：美术课上你画的《家乡的桥》贴在了走廊展板上，很多同学停下来细看。"
                          "【待补：该生本学期另一件具体表现】做数学计算时，试着每算完一题就换一种方法验一遍。\n")
        report = self.run_json("comments", str(path), "--min-chars", "0")
        self.assertNotIn("PY-010", self.ids(report))

    def test_facts_reconciliation(self):
        facts = self.write("事实表.md", "| 编号 | 具体表现 |\n|---|---|\n| S01 | 担任两次升旗手；竹编笔筒进展柜 |\n"
                                       "| S02 | 合唱领唱 |\n| S03 | 书写进步 |\n")
        comments = self.write("评语.txt",
                              "S01：这学期你3次担任升旗手，站姿挺拔，劳动课做的竹编笔筒被放进了学校展柜。希望你上课再专注一些。\n"
                              "S02：合唱比赛你担任领唱，声音清亮，还在《童心向党》演出中独唱。希望你按时交作业，做事更有计划。\n"
                              "S04：你这学期进步很大，老师为你高兴，希望你继续保持，在新学期里取得更好的成绩。\n")
        report = self.run_json("comments", str(comments), "--facts", str(facts))
        messages = " ".join(i["message"] for i in report["issues"])
        self.assertIn("“3”在该生的事实表中找不到", messages)        # 事实表写的是“两次”
        self.assertIn("童心向党", messages)
        self.assertIn("事实表中有 S03，评语中没有", messages)
        self.assertIn("事实表中没有 S04", messages)

    def test_chinese_numerals_in_facts_match_digits(self):
        facts = self.write("事实表.md", "| 编号 | 具体表现 |\n|---|---|\n| S01 | 读完十二本课外书 |\n")
        comments = self.write("评语.txt", "S01：这学期你读完了12本课外书，读书笔记写得认真，常给同学推荐好书。希望你把读到的好词好句用进作文。\n")
        report = self.run_json("comments", str(comments), "--facts", str(facts))
        self.assertNotIn("PY-016", self.ids(report))

    def test_table_input_with_comment_column(self):
        path = self.write("评语.md", "| 编号 | 评语 |\n|---|---|\n"
                                     "| S01 | 你是班上的后进生，要多努力，下学期要按时完成作业，把字写工整。 |\n"
                                     "| S02 | 你在科技节上做的水火箭飞得最远，讲解时条理清楚。希望你写实验记录时把数据记全。 |\n")
        report = self.run_json("comments", str(path))
        self.assertEqual(report["entry_count"], 2)
        self.assertIn("PY-002", self.ids(report))

    def test_label_on_its_own_line(self):
        path = self.write("评语.txt", "期末评语\n\nS01\n你每天坚持跳绳一千下，体育课上的进步大家都看到了。希望你把这份坚持用到背诵古诗上。\n\n"
                                      "S02\n你的手抄报《节约用水》被贴在了走廊展板上，配色很用心。希望你做题时把草稿写整齐。\n")
        report = self.run_json("comments", str(path), "--min-chars", "30")
        self.assertEqual([e["label"] for e in report["entries"]], ["S01", "S02"])
        self.assertEqual(report["issues"], [])

    def test_labour_words_and_negations_are_not_insults(self):
        path = self.write("评语.txt",
                          "S01：你在垃圾分类活动中当了宣传员，还把废物利用的点子讲给大家听。你常说“笨鸟先飞”，每天提前到校读书。\n"
                          "S02：你一点也不懒，大扫除时抢着搬笨重的桌子，从不偷懒。希望你上课多举手。\n"
                          "S03：你太懒了，是个笨孩子，每天都把垃圾丢在座位旁。\n")
        report = self.run_json("comments", str(path), "--min-chars", "0")
        by_label = {}
        for item in report["issues"]:
            by_label.setdefault(item["label"], set()).add(item["id"])
        self.assertFalse({"PY-001", "PY-002"} & by_label.get("S01", set()), by_label.get("S01"))
        self.assertFalse({"PY-001", "PY-002"} & by_label.get("S02", set()), by_label.get("S02"))
        self.assertIn("PY-001", by_label["S03"])            # “笨孩子”仍报贬损
        self.assertIn("PY-002", by_label["S03"])            # “太懒了”仍报标签化

    def test_shared_sentence_is_reported(self):
        tail = "希望你在新的学期里继续努力争取更大进步"
        body = ["你给班级图书角捐了五本书，还做了借阅登记表。",
                "你在拔河比赛中当第一棒，喊口号最响。",
                "你帮生病的同学整理了一周的笔记。",
                "你把教室后面的植物角照顾得很好。"]
        text = "\n".join(f"S0{i}：{b}{tail}。" for i, b in enumerate(body, start=1))
        report = self.run_json("comments", str(self.write("评语.txt", text)), "--min-chars", "0")
        self.assertIn("PY-008", self.ids(report))

    def test_no_entries_is_input_error(self):
        code, _out, err = self.run_main("comments", str(self.write("空.txt", "\n\n")))
        self.assertEqual(code, 2)
        self.assertIn("没有识别到评语", err)

    def test_excerpt_masks_id_number(self):
        report = self.run_json("comments", str(SAMPLES / "comments_must_fail.txt"))
        dumped = json.dumps(report, ensure_ascii=False)
        self.assertNotIn("110101201301010011", dumped)


class ParentsTest(K12TestCase):

    def test_must_pass_sample_has_no_issue(self):
        report = self.run_json("parents", str(SAMPLES / "parents_must_pass.txt"))
        self.assertEqual(report["issues"], [], report["issues"])

    def test_must_fail_sample_reports_each_problem(self):
        report = self.run_json("parents", str(SAMPLES / "parents_must_fail.txt"))
        found = set(self.ids(report))
        for rule in ("JZ-001", "JZ-002", "JZ-003", "JZ-004", "JZ-005", "JZ-006", "JZ-007",
                     "JZ-008", "JZ-009", "JZ-010", "JZ-011", "JZ-012", "JZ-013"):
            self.assertIn(rule, found, rule)
        p0 = set(self.ids(report, "P0"))
        self.assertTrue({"JZ-001", "JZ-003", "JZ-006", "JZ-008", "JZ-010"} <= p0, p0)

    def test_each_line_of_unlabelled_violations_is_reported(self):
        report = self.run_json("parents", str(SAMPLES / "parents_must_fail_lines.txt"))
        by_line = {}
        for item in report["issues"]:
            by_line.setdefault(item["line"], set()).add(item["id"])
        expected = {1: "JZ-001", 2: "JZ-003", 3: "JZ-014", 4: "JZ-003", 5: "JZ-002"}
        for line, rule in expected.items():
            self.assertIn(rule, by_line.get(line, set()), (line, rule, by_line))
        self.assertIn("JZ-012", by_line[3])                 # 点名并写“拖了全班后腿”
        levels = {(i["line"], i["id"]): i["level"] for i in report["issues"]}
        self.assertEqual(levels[(1, "JZ-001")], "P0")
        self.assertEqual(levels[(2, "JZ-003")], "P0")
        self.assertEqual(levels[(3, "JZ-014")], "P0")

    def test_speech_for_review_reports_every_problem(self):
        report = self.run_json("parents", str(SAMPLES / "parents_must_fail_speech.txt"))
        found = set(self.ids(report))
        for rule in ("JZ-001", "JZ-002", "JZ-003", "JZ-005", "JZ-006", "JZ-007", "JZ-011", "JZ-012",
                     "JZ-013", "JZ-014"):
            self.assertIn(rule, found, rule)
        self.assertTrue({"JZ-001", "JZ-003", "JZ-006", "JZ-014"} <= set(self.ids(report, "P0")))

    def test_rank_words_in_dates_and_counts_are_not_reported(self):
        path = self.write("信.txt", "七年级第一学期的家长开放日另行通知。年级第一次家长会安排在开学第三周。"
                                    "全区第一批试点学校名单已公布。学期倒数第二周安排期末复习。")
        report = self.run_json("parents", str(path))
        self.assertNotIn("JZ-003", self.ids(report))

    def test_negated_statements_are_not_reported(self):
        path = self.write("信.txt", "学校不要求家长批改作业。考试结果不排名、不公布。课后服务不讲新课。"
                                    "学校不统一订购教辅资料。教师不得推荐学生参加校外培训班。"
                                    "期末成绩不会发到家长群，由班主任单独告知。不需要家长检查作业并签字。")
        self.assertEqual(self.run_json("parents", str(path))["issues"], [])

    def test_teacher_grading_is_not_reported(self):
        path = self.write("信.txt", "老师会当天批改作业，请家长督促孩子按时休息。请家长在回执上签字后交回。")
        self.assertEqual(self.run_json("parents", str(path))["issues"], [])

    def test_competition_rank_and_group_mentions_are_not_reported(self):
        path = self.write("讲话.txt", "本学期我班在校运动会上获得年级第二名。个别同学课间追跑，有过两次打架，会后我会单独沟通。"
                                      "请同学们不要打架。")
        report = self.run_json("parents", str(path))
        self.assertNotIn("JZ-003", self.ids(report))
        self.assertNotIn("JZ-012", self.ids(report))

    def test_high_school_stage(self):
        path = self.write("信.txt", "请家长每天检查孩子作业并签字。高一年级不布置书面作业的说法不适用。")
        report = self.run_json("parents", str(path), "--stage", "高中")
        levels = {i["id"]: i["level"] for i in report["issues"]}
        self.assertEqual(levels.get("JZ-001"), "P1")
        self.assertNotIn("JZ-004", levels)

    def test_fee_with_basis_in_paragraph(self):
        path = self.write("信.txt", "根据某区教育局《关于课后服务收费的通知》（某教〔2026〕5号），课后服务每生每月收费150元，家长自愿申请。")
        self.assertNotIn("JZ-007", self.ids(self.run_json("parents", str(path))))
        path2 = self.write("信2.txt", "请于周五前缴纳资料费每人40元，资料为《期末冲刺》。")
        self.assertIn("JZ-007", self.ids(self.run_json("parents", str(path2))))

    def test_text_after_draft_note_is_ignored(self):
        report = self.run_json("parents", str(SAMPLES / "parents_must_pass.txt"))
        self.assertNotIn("JZ-001", self.ids(report))

    @unittest.skipUnless(HAS_DOCX, "未安装 python-docx")
    def test_docx_input(self):
        import docx  # type: ignore
        d = docx.Document()
        d.add_paragraph("各位家长：")
        d.add_paragraph("期末考试结束后，学校将在家长群公布成绩排名。")
        path = self.tmp / "信.docx"
        d.save(str(path))
        self.assertIn("JZ-003", self.ids(self.run_json("parents", str(path))))


class MaskTest(K12TestCase):

    def assert_no_names(self, text: str) -> None:
        for name in NAMES:
            self.assertNotIn(name, text, name)

    def test_mask_lines_and_unmask_roundtrip(self):
        src = self.copy("students_lines.txt")
        code, out, err = self.run_main("mask", str(src))
        self.assertEqual(code, 0, err)
        self.assert_no_names(out)                                   # 屏幕输出不含姓名
        masked = (self.tmp / "students_lines-脱敏.txt").read_text(encoding="utf-8")
        self.assert_no_names(masked)
        self.assertIn("S01：", masked)
        self.assertIn("S01和S03是同桌", masked)                    # 名字（小明、晓晓）也替换
        self.assertNotIn("13812345678", masked)
        self.assertNotIn("20260102", masked)
        self.assertIn("1500米", masked)                            # 普通数字不动
        self.assertIn("S03", out)                                   # 敏感信息提示只给编号
        with (self.tmp / "students_lines-对照表.csv").open(encoding="utf-8-sig", newline="") as fh:
            rows = list(csv.reader(fh))
        self.assertEqual(rows[0], ["编号", "姓名"])
        self.assertEqual(rows[1], ["S01", "王小明"])
        comments = self.write("评语.txt", "S01：你认真负责。\nS03：你书写进步很大。\n")
        code, out, err = self.run_main("unmask", str(comments), str(self.tmp / "students_lines-对照表.csv"))
        self.assertEqual(code, 0, err)
        self.assert_no_names(out)
        restored = (self.tmp / "评语-姓名版.txt").read_text(encoding="utf-8")
        self.assertIn("王小明：你认真负责。", restored)
        self.assertIn("欧阳晓晓：你书写进步很大。", restored)

    def test_teacher_header_lines_are_not_students(self):
        src = self.copy("students_lines.txt")                       # 开头两行是“班主任：李老师”等表头
        code, out, err = self.run_main("mask", str(src), "--json")
        self.assertEqual(code, 0, err)
        self.assertEqual(json.loads(out)["students"], 3)
        masked = (self.tmp / "students_lines-脱敏.txt").read_text(encoding="utf-8")
        self.assertIn("班主任：李老师", masked)
        self.assertIn("任课教师：数学王老师、英语张老师", masked)

    def test_drop_sensitive_replaces_family_and_health_sentences(self):
        src = self.write("情况.txt", "1. 陈雨桐：科学课记录豌豆苗高度。需改进：默写错得多。父母离异，跟奶奶生活，心理老师说有焦虑情绪。\n"
                                     "2. 刘子涵：合唱比赛领唱。受过警告处分。\n")
        code, out, err = self.run_main("mask", str(src), "--drop-sensitive")
        self.assertEqual(code, 0, err)
        masked = (self.tmp / "情况-脱敏.txt").read_text(encoding="utf-8")
        for secret in ("离异", "奶奶", "焦虑", "警告", "陈雨桐", "刘子涵"):
            self.assertNotIn(secret, masked)
        self.assertIn("【已略去：家庭、健康信息】", masked)
        self.assertIn("科学课记录豌豆苗高度", masked)                 # 其他句子保留
        self.assertIn("已略去：家庭信息 1 处、健康信息 1 处、处分信息 1 处", out)
        self.assertNotIn("陈雨桐", out)

    def test_without_drop_sensitive_hint_names_the_option(self):
        src = self.write("情况.txt", "1. 陈雨桐：科学课记录豌豆苗高度。父母离异。\n2. 刘子涵：合唱比赛领唱。\n")
        code, out, err = self.run_main("mask", str(src))
        self.assertEqual(code, 0, err)
        self.assertIn("--drop-sensitive", out)
        self.assertIn("离异", (self.tmp / "情况-脱敏.txt").read_text(encoding="utf-8"))  # 不加选项时保留，只提示

    def test_given_name_next_to_other_words_is_replaced_with_hint(self):
        src = self.write("情况.txt", "1. 张明月：明月当空的夜晚写了一首小诗。\n2. 刘子涵：合唱比赛领唱。\n")
        code, out, err = self.run_main("mask", str(src))
        self.assertEqual(code, 0, err)
        masked = (self.tmp / "情况-脱敏.txt").read_text(encoding="utf-8")
        self.assertNotIn("明月", masked)                             # 隐私优先，照常替换
        self.assertIn("明月当空", out)                               # 提示人工核对

    def test_mask_does_not_overwrite(self):
        src = self.copy("students_lines.txt")
        self.run_main("mask", str(src))
        self.run_main("mask", str(src))
        self.assertTrue((self.tmp / "students_lines-脱敏-v2.txt").exists())
        self.assertTrue((self.tmp / "students_lines-对照表-v2.csv").exists())

    def test_mask_csv_table_drops_sensitive_columns(self):
        src = self.write("名单.csv", "姓名,学号,具体表现,家长电话\n张小雨,20260301,运动会跳绳第一名；张小雨帮同学补课,13900001111\n"
                                     "刘洋,20260302,作文进步,13900002222\n")
        code, out, err = self.run_main("mask", str(src))
        self.assertEqual(code, 0, err)
        masked = (self.tmp / "名单-脱敏.md").read_text(encoding="utf-8")
        for secret in ("张小雨", "刘洋", "小雨", "20260301", "13900001111"):
            self.assertNotIn(secret, masked)
        self.assertIn("| 编号 |", masked)
        self.assertIn("S01帮同学补课", masked)

    def test_mask_markdown_table_keeps_other_text(self):
        src = self.write("事实.md", "# 四年级一班\n\n| 姓名 | 表现 |\n|---|---|\n| 赵子涵 | 担任班长 |\n| 钱多多 | 爱画画 |\n\n说明：子涵和多多都参加了科技节。\n")
        code, _out, err = self.run_main("mask", str(src))
        self.assertEqual(code, 0, err)
        masked = (self.tmp / "事实-脱敏.md").read_text(encoding="utf-8")
        self.assertIn("# 四年级一班", masked)
        self.assertIn("| S01 | 担任班长 |", masked)
        self.assertIn("S01和S02都参加了科技节", masked)

    @unittest.skipUnless(HAS_OPENPYXL, "未安装 openpyxl")
    def test_mask_xlsx(self):
        import openpyxl  # type: ignore
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["学生姓名", "具体表现"])
        ws.append(["孙悦", "主动打扫卫生"])
        ws.append(["周一鸣", "数学竞赛获奖"])
        path = self.tmp / "名单.xlsx"
        wb.save(str(path))
        code, _out, err = self.run_main("mask", str(path))
        self.assertEqual(code, 0, err)
        masked = (self.tmp / "名单-脱敏.md").read_text(encoding="utf-8")
        self.assertNotIn("孙悦", masked)
        self.assertNotIn("周一鸣", masked)

    def test_mask_without_names_is_input_error(self):
        src = self.write("无姓名.txt", "这学期班级开展了三次主题班会。\n")
        code, _out, err = self.run_main("mask", str(src))
        self.assertEqual(code, 2)
        self.assertIn("没有识别到学生姓名", err)
        self.assertFalse(any(self.tmp.glob("*-脱敏*")))

    def test_unmask_rejects_bad_mapping(self):
        comments = self.write("评语.txt", "S01：你好。\n")
        bad = self.write("表.csv", "a,b\n1,2\n")
        code, _out, err = self.run_main("unmask", str(comments), str(bad))
        self.assertEqual(code, 2)


class CliTest(K12TestCase):

    def test_missing_file_returns_2(self):
        code, _out, err = self.run_main("parents", str(self.tmp / "不存在.txt"))
        self.assertEqual(code, 2)
        self.assertIn("找不到文件", err)

    def test_unsupported_suffix_returns_2(self):
        code, _out, _err = self.run_main("parents", str(self.write("信.pdf", "x")))
        self.assertEqual(code, 2)

    def test_no_subcommand_returns_2(self):
        code, _out, _err = self.run_main()
        self.assertEqual(code, 2)

    def test_help_is_chinese(self):
        code, out, _err = self.run_main("--help")
        self.assertEqual(code, 0)
        self.assertIn("用法：", out)
        self.assertIn("子命令", out)

    def test_text_output(self):
        code, out, _err = self.run_main("parents", str(SAMPLES / "parents_must_fail.txt"))
        self.assertEqual(code, 0)
        self.assertIn("[P0] JZ-001", out)


if __name__ == "__main__":
    unittest.main()
