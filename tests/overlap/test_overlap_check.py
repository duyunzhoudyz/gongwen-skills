# -*- coding: utf-8 -*-
"""shared/scripts/overlap_check.py 的单元测试。

运行：python -m unittest discover -s tests/overlap -t .

标本（tests/overlap/samples/）：
    上级文件.txt：虚构的省级专项整治通知；
    照搬稿.txt：必错标本，大段照抄上级文件，重合率必须超过 30%，--strict 返回 1；
    转化稿.txt：必过标本，任务已转化为本地措施，只在书名号内引用文件名称，--strict 返回 0。
"""
from __future__ import annotations

import io
import json
import shutil
import sys
import tempfile
import time
import unittest
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[2]
SCRIPTS_DIR = REPO_ROOT / "shared" / "scripts"
SAMPLES = Path(__file__).resolve().parent / "samples"
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

import overlap_check  # noqa: E402

try:
    import docx  # type: ignore  # noqa: F401
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

# 互不相同的汉字，用于拼出不会偶然重合的文字
POOL = [chr(code) for code in range(0x4E00, 0x4E00 + 3000)]


def unique_text(start: int, length: int) -> str:
    return "".join(POOL[start:start + length])


class OverlapTestCase(unittest.TestCase):

    def setUp(self) -> None:
        self.tmp = Path(tempfile.mkdtemp(prefix="overlap_test_"))

    def tearDown(self) -> None:
        shutil.rmtree(self.tmp, ignore_errors=True)

    def write(self, name: str, text: str, encoding: str = "utf-8") -> Path:
        path = self.tmp / name
        path.write_bytes(text.encode(encoding))
        return path

    def run_main(self, *args: str):
        out, err = io.StringIO(), io.StringIO()
        with redirect_stdout(out), redirect_stderr(err):
            code = overlap_check.main(list(args))
        return code, out.getvalue(), err.getvalue()

    def check(self, draft: str, *refs: str, min_len: int = 15) -> dict:
        draft_path = self.write("稿件.txt", draft)
        ref_paths = [self.write(f"参照{i}.txt", text) for i, text in enumerate(refs, start=1)]
        code, out, err = self.run_main(str(draft_path), *map(str, ref_paths),
                                       "--min-len", str(min_len), "--json")
        self.assertEqual(code, 0, err)
        return json.loads(out)


class SpecimenTest(OverlapTestCase):
    """必错标本与必过标本。"""

    def run_specimen(self, name: str):
        return self.run_main(str(SAMPLES / name), str(SAMPLES / "上级文件.txt"), "--json", "--strict")

    def test_copied_specimen_must_fail(self):
        code, out, _err = self.run_specimen("照搬稿.txt")
        report = json.loads(out)
        self.assertEqual(code, 1)
        ref = report["references"][0]
        self.assertTrue(ref["over_threshold"])
        self.assertGreater(ref["ratio"], 0.6)
        lines = [item["line"] for item in report["paragraphs"]]
        for expected in (6, 9, 10, 11, 12, 15):
            self.assertIn(expected, lines)
        target = [f for f in ref["fragments"] if f["draft_line"] == 6]
        self.assertEqual(len(target), 1)
        self.assertEqual(target[0]["ref_line"], 8)
        self.assertIn("重大隐患整改率达到100%", target[0]["text"])

    def test_copied_specimen_ignores_text_after_boundary(self):
        _code, out, _err = self.run_specimen("照搬稿.txt")
        report = json.loads(out)
        for fragment in report["references"][0]["fragments"]:
            self.assertLess(fragment["draft_line"], 17)
        self.assertTrue(any("第 17 行为【起草说明】" in note for note in report["notes"]))

    def test_localized_specimen_must_pass(self):
        code, out, _err = self.run_specimen("转化稿.txt")
        report = json.loads(out)
        self.assertEqual(code, 0)
        self.assertFalse(report["over_threshold"])
        self.assertLess(report["references"][0]["ratio"], 0.1)
        self.assertEqual(report["paragraphs"], [])


class NormalizationTest(OverlapTestCase):

    def test_punctuation_and_whitespace_ignored(self):
        ref = "坚持问题导向，强化源头治理，推动隐患动态清零。"
        draft = "坚持问题导向、强化源头治理 推动隐患动态清零"
        report = self.check(draft, ref)
        self.assertEqual(report["references"][0]["overlap_chars"], report["char_count"])
        self.assertEqual(report["char_count"], 20)

    def test_fullwidth_and_case_normalized(self):
        ref = "到2026年底完成abc系统改造任务并全部通过验收"
        draft = "到２０２６年底完成ＡＢＣ系统改造任务并全部通过验收"
        report = self.check(draft, ref)
        self.assertEqual(report["references"][0]["ratio"], 1.0)

    def test_book_title_not_counted(self):
        title = "某省人民政府办公厅关于开展城镇燃气安全专项整治行动的通知"
        ref = f"{title}\n各地要抓好落实。"
        draft = f"为贯彻落实《{title}》，结合本市实际，制定本方案。"
        report = self.check(draft, ref)
        self.assertEqual(report["references"][0]["overlap_chars"], 0)
        self.assertEqual(report["char_count"], len("为贯彻落实结合本市实际制定本方案"))

    def test_marker_not_counted(self):
        piece = unique_text(0, 20)
        draft = f"{piece}【待补：本单位落实措施和牵头科室】"
        report = self.check(draft, piece + unique_text(100, 10))
        self.assertEqual(report["char_count"], 20)
        self.assertEqual(report["references"][0]["ratio"], 1.0)

    def test_excluded_span_breaks_fragment(self):
        left, right = unique_text(0, 10), unique_text(10, 10)
        ref = left + right
        draft = f"{left}《某某文件》{right}"
        report = self.check(draft, ref)
        self.assertEqual(report["references"][0]["overlap_chars"], 0)

    def test_fragment_does_not_cross_draft_paragraphs(self):
        ref = unique_text(0, 24)
        draft = unique_text(0, 12) + "\n" + unique_text(12, 12)
        report = self.check(draft, ref)
        self.assertEqual(report["references"][0]["overlap_chars"], 0)

    def test_merged_reference_paragraphs_still_found(self):
        ref = unique_text(0, 12) + "。\n" + unique_text(12, 12) + "。"
        draft = unique_text(0, 24)
        report = self.check(draft, ref)
        fragments = report["references"][0]["fragments"]
        self.assertEqual(len(fragments), 1)
        self.assertEqual((fragments[0]["ref_line"], fragments[0]["ref_end_line"]), (1, 2))
        self.assertEqual(fragments[0]["length"], 24)


class MinLenTest(OverlapTestCase):

    def test_default_min_len_15(self):
        common = unique_text(0, 14)
        draft = unique_text(100, 10) + common + unique_text(200, 10)
        ref = unique_text(300, 10) + common + unique_text(400, 10)
        self.assertEqual(self.check(draft, ref)["references"][0]["overlap_chars"], 0)
        report = self.check(draft, ref, min_len=10)
        self.assertEqual(report["references"][0]["overlap_chars"], 14)
        self.assertEqual(report["references"][0]["fragments"][0]["length"], 14)

    def test_exactly_min_len_reported(self):
        common = unique_text(0, 15)
        report = self.check(unique_text(100, 5) + common, common + unique_text(200, 5))
        self.assertEqual(report["references"][0]["overlap_chars"], 15)

    def test_invalid_min_len(self):
        draft = self.write("稿件.txt", "测试")
        ref = self.write("参照.txt", "测试")
        for value in ("2", "0", "-5", "abc", "5000"):
            code, _out, _err = self.run_main(str(draft), str(ref), "--min-len", value)
            self.assertEqual(code, 2, value)


class RatioTest(OverlapTestCase):

    def test_ratio_is_overlap_over_draft(self):
        common = unique_text(0, 20)
        draft = common + unique_text(100, 20)
        report = self.check(draft, common + unique_text(500, 30))
        ref = report["references"][0]
        self.assertEqual(report["char_count"], 40)
        self.assertEqual(ref["overlap_chars"], 20)
        self.assertEqual(ref["ratio"], 0.5)
        self.assertEqual(ref["percent"], 50.0)

    def test_threshold_is_strictly_greater_than_30_percent(self):
        draft30 = unique_text(0, 30) + unique_text(1000, 70)
        report = self.check(draft30, unique_text(0, 30))
        self.assertEqual(report["references"][0]["percent"], 30.0)
        self.assertFalse(report["over_threshold"])
        draft31 = unique_text(0, 31) + unique_text(1000, 69)
        draft_path = self.write("稿件31.txt", draft31)
        ref_path = self.write("参照31.txt", unique_text(0, 31))
        code, out, _err = self.run_main(str(draft_path), str(ref_path), "--json", "--strict")
        self.assertEqual(code, 1)
        self.assertTrue(json.loads(out)["over_threshold"])
        code, _out, _err = self.run_main(str(draft_path), str(ref_path), "--json")
        self.assertEqual(code, 0)

    def test_multiple_references_and_combined(self):
        a, b, c = unique_text(0, 20), unique_text(20, 20), unique_text(40, 20)
        draft = a + "\n" + b + "\n" + c + unique_text(500, 20)
        report = self.check(draft, a + b, b + unique_text(800, 5))
        first, second = report["references"]
        self.assertEqual(first["overlap_chars"], 40)
        self.assertEqual(second["overlap_chars"], 20)
        # b 在两份参照中都出现，合计时只算一次
        self.assertEqual(report["combined"]["overlap_chars"], 40)
        self.assertEqual(report["combined"]["ratio"], 0.5)


class FragmentTest(OverlapTestCase):

    def test_quoted_fragment_flagged(self):
        quote = "严禁以罚代管严禁层层加码严禁过度留痕增加负担"
        ref = "工作要求：" + quote + "。"
        draft = "执行“" + quote + "”的要求。\n" + quote + "。"
        fragments = self.check(draft, ref)["references"][0]["fragments"]
        self.assertEqual([f["quoted"] for f in fragments], [True, False])

    def test_halfwidth_double_quotes_flagged(self):
        quote = "严禁以罚代管严禁层层加码严禁过度留痕增加负担"
        ref = "工作要求：" + quote + "。"
        draft = '执行"' + quote + '"的要求。\n' + quote + "。"
        fragments = self.check(draft, ref)["references"][0]["fragments"]
        self.assertEqual([f["quoted"] for f in fragments], [True, False])

    def test_unpaired_halfwidth_quote_not_flagged(self):
        # 同一行只有一个半角双引号、引号不跨行配对：片段不算在引号内
        quote = "严禁以罚代管严禁层层加码严禁过度留痕增加负担"
        ref = "工作要求：" + quote + "。"
        draft = '执行"' + quote + "的要求。\n其他内容\"另起一行。"
        fragments = self.check(draft, ref)["references"][0]["fragments"]
        self.assertEqual([f["quoted"] for f in fragments], [False])

    def test_line_numbers_and_original_text(self):
        ref = "第一行无关内容\n第二行" + "对辖区内燃气经营企业逐一开展检查，重点检查经营许可"
        draft = "标题\n\n本单位要对辖区内燃气经营企业逐一开展检查、重点检查经营许可。"
        fragment = self.check(draft, ref)["references"][0]["fragments"][0]
        self.assertEqual(fragment["draft_line"], 3)
        self.assertEqual(fragment["ref_line"], 2)
        self.assertEqual(fragment["text"], "对辖区内燃气经营企业逐一开展检查、重点检查经营许可")
        self.assertEqual(fragment["ref_text"], "对辖区内燃气经营企业逐一开展检查，重点检查经营许可")

    def test_repeated_reference_gives_one_fragment(self):
        piece = unique_text(0, 20)
        ref = (piece + "。\n") * 5
        fragments = self.check(piece, ref)["references"][0]["fragments"]
        self.assertEqual(len(fragments), 1)
        self.assertEqual(fragments[0]["length"], 20)


class ParagraphTest(OverlapTestCase):

    def test_high_overlap_paragraph_listed(self):
        common = unique_text(0, 30)
        draft = "（一）" + common + unique_text(100, 10) + "\n" + unique_text(200, 40) + unique_text(0, 16)
        report = self.check(draft, common)
        paragraphs = report["paragraphs"]
        self.assertEqual([p["line"] for p in paragraphs], [1])
        self.assertEqual(paragraphs[0]["overlap_chars"], 30)
        self.assertEqual(paragraphs[0]["references"], [1])
        # 第 2 行只有 16/56 重合，不列为高重合段落
        self.assertEqual(report["references"][0]["overlap_chars"], 46)


class InputTest(OverlapTestCase):

    def test_missing_file(self):
        ref = self.write("参照.txt", "内容")
        code, _out, err = self.run_main(str(self.tmp / "不存在.txt"), str(ref))
        self.assertEqual(code, 2)
        self.assertIn("文件不存在", err)

    def test_unsupported_suffix(self):
        draft = self.write("稿件.pdf", "内容")
        ref = self.write("参照.txt", "内容")
        code, _out, err = self.run_main(str(draft), str(ref))
        self.assertEqual(code, 2)
        self.assertIn("不支持的文件类型", err)

    def test_directory_rejected(self):
        ref = self.write("参照.txt", "内容")
        code, _out, _err = self.run_main(str(self.tmp), str(ref))
        self.assertEqual(code, 2)

    def test_reference_required(self):
        draft = self.write("稿件.txt", "内容")
        code, _out, _err = self.run_main(str(draft))
        self.assertEqual(code, 2)

    def test_help_returns_zero(self):
        code, out, _err = self.run_main("--help")
        self.assertEqual(code, 0)
        self.assertIn("用法", out)

    def test_gb18030_file(self):
        piece = "全面排查整治城镇燃气安全风险隐患并建立清单"
        draft = self.write("稿件.txt", piece, encoding="gb18030")
        ref = self.write("参照.txt", piece)
        code, out, err = self.run_main(str(draft), str(ref), "--json")
        self.assertEqual(code, 0)
        self.assertIn("GB18030", err)
        self.assertEqual(json.loads(out)["references"][0]["ratio"], 1.0)

    def test_empty_draft(self):
        report = self.check("【待补：全部内容】\n《某某通知》", "随便一段参照文字随便一段参照文字随便一段")
        self.assertEqual(report["char_count"], 0)
        self.assertEqual(report["references"][0]["ratio"], 0.0)
        self.assertTrue(report["notes"])

    def test_short_reference_note(self):
        report = self.check(unique_text(0, 30), "很短的参照")
        self.assertEqual(report["references"][0]["overlap_chars"], 0)
        self.assertTrue(any("不足 15 字" in note for note in report["notes"]))

    @unittest.skipUnless(HAS_DOCX, "未安装 python-docx")
    def test_docx_input(self):
        import docx  # type: ignore
        document = docx.Document()
        document.add_paragraph("某市燃气安全专项整治实施方案")
        document.add_paragraph("对辖区内燃气经营企业逐一开展检查，重点检查经营许可、气瓶充装情况。")
        document.add_paragraph("【起草说明】")
        document.add_paragraph("各地要对辖区内燃气经营企业逐一开展检查，这段不参与比对。")
        path = self.tmp / "稿件.docx"
        document.save(str(path))
        code, out, _err = self.run_main(str(path), str(SAMPLES / "上级文件.txt"), "--json")
        self.assertEqual(code, 0)
        report = json.loads(out)
        fragments = report["references"][0]["fragments"]
        self.assertEqual([f["draft_line"] for f in fragments], [2])
        self.assertTrue(any("第 3 行为【起草说明】" in note for note in report["notes"]))


class OutputTest(OverlapTestCase):

    def test_text_output(self):
        code, out, _err = self.run_main(str(SAMPLES / "照搬稿.txt"), str(SAMPLES / "上级文件.txt"))
        self.assertEqual(code, 0)
        self.assertIn("重合率", out)
        self.assertIn("超过 30%", out)
        self.assertIn("高重合段落（本段一半以上的文字与参照重合）", out)
        self.assertIn("是否属于照搬照抄或雷同由人工判断", out)

    def test_json_fields(self):
        report = self.check(unique_text(0, 20), unique_text(0, 20))
        for key in ("file", "char_count", "min_len", "threshold_percent", "references",
                    "combined", "paragraphs", "over_threshold", "notes"):
            self.assertIn(key, report)
        ref = report["references"][0]
        for key in ("file", "char_count", "overlap_chars", "ratio", "percent", "over_threshold", "fragments"):
            self.assertIn(key, ref)
        fragment = ref["fragments"][0]
        for key in ("length", "draft_line", "draft_end_line", "ref_line", "ref_end_line",
                    "text", "ref_text", "quoted"):
            self.assertIn(key, fragment)


class PerformanceTest(OverlapTestCase):

    def test_long_documents(self):
        lines, pos = [], 0
        while pos < 2900:
            lines.append(unique_text(pos, 40) + "。")
            pos += 40
        ref_text = "\n".join(lines * 20)  # 约 5.8 万字，大量重复
        draft_text = "\n".join(lines[:30] * 10)
        draft = self.write("稿件.txt", draft_text)
        ref = self.write("参照.txt", ref_text)
        started = time.time()
        code, out, _err = self.run_main(str(draft), str(ref), "--json")
        self.assertEqual(code, 0)
        self.assertLess(time.time() - started, 20)
        self.assertEqual(json.loads(out)["references"][0]["ratio"], 1.0)


if __name__ == "__main__":
    unittest.main()
