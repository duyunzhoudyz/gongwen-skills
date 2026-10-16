# -*- coding: utf-8 -*-
"""skills/gongwen-polish/scripts/polish_stat.py 的单元测试。

运行：python -m unittest discover -s tests/polish -t .

每个子命令都有“必错标本”（脚本必须报出问题）和“必过标本”（脚本不得误报）。
BodyTest 中与 fact_diff 联动的用例需要先运行 python tools/sync_shared.py --skill gongwen-polish。
"""
from __future__ import annotations

import io
import json
import shutil
import sys
import tempfile
import unittest
import zipfile
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[2]
SCRIPTS_DIR = REPO_ROOT / "skills" / "gongwen-polish" / "scripts"
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

import polish_stat  # noqa: E402

try:
    import fact_diff  # type: ignore  # noqa: E402  同步进本 skill 的共享脚本
    HAS_FACT_DIFF = True
except ImportError:
    HAS_FACT_DIFF = False

try:
    import gongwen_lint  # type: ignore  # noqa: E402  同步进本 skill 的共享机检脚本
    HAS_LINT = True
except ImportError:
    HAS_LINT = False

try:
    import docx  # type: ignore  # noqa: F401
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

REPORT = """某局关于2026年安全生产工作情况的报告

县政府：
今年以来，我局高度重视安全生产工作。全年开展专项检查36次，排查隐患218处。

一、主要做法
（一）压实责任。与12个乡镇签订安全生产责任书。
（二）开展检查。全年开展专项检查36次，整改率【待补：整改率】。

二、存在问题
部分企业主体责任落实不到位。

特此报告。
【起草说明】
一、所作假设
1. 起草说明里的文字很长很长很长，不计入篇幅。
"""

DRAFT_A = """某学院2026年工作总结

一、人才培养
全年招收本科生420人，毕业生就业率91.2%。开设实践课程18门，组织学生参加学科竞赛获省级奖项7项。

二、科学研究
获批省级科研项目5项，到账经费320万元。
"""

DRAFT_B = """某学院2026年工作总结

一、人才培养
今年招收本科生420人，毕业生去向落实率91.2%。新建校外实习基地4个。

二、科学研究
获批省级科研项目5项，发表论文46篇。
"""

MERGED_OK = """某学院2026年工作总结

一、人才培养
全年招收本科生420人，毕业生就业率91.2%。开设实践课程18门，组织学生参加学科竞赛获省级奖项7项。新建校外实习基地4个。

二、科学研究
获批省级科研项目5项，到账经费320万元。发表论文46篇。
"""

MERGED_BAD = MERGED_OK.replace(
    "发表论文46篇。",
    "发表论文46篇。学院获评全省教学管理先进单位，社会反响热烈。",
)

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

DOC_XML = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="{W}"><w:body>
<w:p><w:r><w:t>某街道关于垃圾分类工作的汇报</w:t></w:r></w:p>
<w:p/>
<w:p><w:r><w:t>今年以来，</w:t></w:r><w:commentRangeStart w:id="0"/><w:r><w:t>我街道垃圾分类工作取得显著成效</w:t></w:r><w:commentRangeEnd w:id="0"/><w:r><w:commentReference w:id="0"/></w:r><w:r><w:t>。</w:t></w:r></w:p>
<w:p><w:r><w:t>下一步，我们将继续加强宣传。</w:t></w:r><w:r><w:commentReference w:id="1"/></w:r></w:p>
<w:p><w:ins w:id="9" w:author="X"><w:r><w:t>新增一句。</w:t></w:r></w:ins><w:del w:id="10" w:author="X"><w:r><w:delText>删去一句。</w:delText></w:r></w:del></w:p>
</w:body></w:document>
"""

REVISED_DOC_XML = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="{W}"><w:body>
<w:p><w:r><w:t>某公司关于安全生产工作情况的汇报</w:t></w:r></w:p>
<w:p><w:r><w:t>全年排查隐患</w:t></w:r><w:del w:id="1" w:author="X"><w:r><w:delText>200</w:delText></w:r></w:del><w:ins w:id="2" w:author="X"><w:r><w:t>218</w:t></w:r></w:ins><w:r><w:t>处。</w:t></w:r></w:p>
<w:tbl><w:tr><w:tc><w:p><w:r><w:t>整改</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>210处</w:t></w:r></w:p></w:tc></w:tr></w:tbl>
<w:sectPr/>
</w:body></w:document>
"""

COMMENTS_XML = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:comments xmlns:w="{W}">
<w:comment w:id="0" w:author="王某某" w:date="2026-10-12T09:30:00Z"><w:p><w:r><w:t>太虚，数据呢？</w:t></w:r></w:p></w:comment>
<w:comment w:id="1" w:author="王某某" w:date="2026-10-12T09:31:00Z"><w:p><w:r><w:t>措施要具体</w:t></w:r></w:p></w:comment>
</w:comments>
"""


class PolishStatCase(unittest.TestCase):

    def setUp(self) -> None:
        self.tmp = Path(tempfile.mkdtemp(prefix="polish_stat_test_"))

    def tearDown(self) -> None:
        shutil.rmtree(self.tmp, ignore_errors=True)

    def write(self, name: str, text: str, encoding: str = "utf-8") -> Path:
        path = self.tmp / name
        path.write_bytes(text.encode(encoding))
        return path

    def make_docx_zip(self, name: str, comments: bool = True, document: str = DOC_XML) -> Path:
        path = self.tmp / name
        with zipfile.ZipFile(path, "w") as zf:
            zf.writestr("[Content_Types].xml", "<Types/>")
            zf.writestr("word/document.xml", document)
            if comments:
                zf.writestr("word/comments.xml", COMMENTS_XML)
        return path

    def run_main(self, *args: str):
        out, err = io.StringIO(), io.StringIO()
        with redirect_stdout(out), redirect_stderr(err):
            code = polish_stat.main(list(args))
        return code, out.getvalue(), err.getvalue()


class CountTest(PolishStatCase):

    def test_title_and_draft_note_excluded(self) -> None:
        report = polish_stat.count_report(self.write("a.txt", REPORT))
        self.assertEqual(report["title"], "某局关于2026年安全生产工作情况的报告")
        body = REPORT.split("【起草说明】")[0].split("\n", 1)[1]
        expected = len("".join(body.split()))
        self.assertEqual(report["char_count"], expected)
        self.assertEqual(report["markers"], {"待补": 1, "待核实": 0})

    def test_sections_split_by_level_one_heading(self) -> None:
        report = polish_stat.count_report(self.write("a.txt", REPORT))
        headings = [s["heading"] for s in report["sections"]]
        self.assertEqual(headings, ["开头", "一、主要做法", "二、存在问题"])
        self.assertEqual(sum(s["char_count"] for s in report["sections"]), report["char_count"])

    def test_include_title(self) -> None:
        without = polish_stat.count_report(self.write("a.txt", REPORT))["char_count"]
        with_title = polish_stat.count_report(self.write("b.txt", REPORT), include_title=True)["char_count"]
        self.assertEqual(with_title - without, len("某局关于2026年安全生产工作情况的报告"))

    def test_head_lines_skipped_before_title(self) -> None:
        text = "某县应急管理局文件\n某应急发〔2026〕5号\n关于开展汛期安全检查的通知\n各乡镇：\n现将有关事项通知如下。\n"
        report = polish_stat.count_report(self.write("h.txt", text))
        self.assertEqual(report["title"], "关于开展汛期安全检查的通知")
        self.assertEqual(report["char_count"], len("各乡镇：现将有关事项通知如下。"))

    def test_sentence_first_line_not_title(self) -> None:
        text = "全年开展专项检查36次。\n排查隐患218处。\n"
        report = polish_stat.count_report(self.write("s.txt", text))
        self.assertEqual(report["title"], "")
        self.assertEqual(report["char_count"], len("全年开展专项检查36次。排查隐患218处。"))

    # 必错标本：两行折行标题，第二行不得算进正文，篇幅与机检一致
    def test_two_line_title_matches_lint(self) -> None:
        text = ("某县应急管理局关于开展2026年汛期\n安全隐患排查整治工作的通知\n"
                "各乡镇人民政府：\n现将有关事项通知如下。\n")
        path = self.write("t2.txt", text)
        report = polish_stat.count_report(path)
        self.assertEqual(report["title"], "某县应急管理局关于开展2026年汛期安全隐患排查整治工作的通知")
        self.assertEqual(report["char_count"], len("各乡镇人民政府：现将有关事项通知如下。"))
        if HAS_LINT:
            lines, _ = gongwen_lint.read_lines(path)
            self.assertEqual(report["char_count"], gongwen_lint.Document(path.name, lines).body_char_count())

    # 必过标本：标题已以文种结尾的，下一行短句不当作标题续行
    def test_title_ending_with_doc_type_not_extended(self) -> None:
        text = "关于开展汛期安全检查的通知\n加强汛期安全检查\n现将有关事项通知如下。\n"
        report = polish_stat.count_report(self.write("t3.txt", text))
        self.assertEqual(report["title"], "关于开展汛期安全检查的通知")
        self.assertEqual(report["char_count"], len("加强汛期安全检查现将有关事项通知如下。"))

    # 必错标本：超出目标±10%
    def test_over_target_reported_and_strict_fails(self) -> None:
        path = self.write("a.txt", REPORT)
        total = polish_stat.count_report(path)["char_count"]
        target = int(total / 1.5)
        report = polish_stat.count_report(path, target=target)
        self.assertEqual(report["target"]["status"], "超出")
        self.assertGreater(report["target"]["diff"], 0)
        code, out, _ = self.run_main("count", str(path), "--target", str(target), "--strict")
        self.assertEqual(code, 1)
        self.assertIn("超出", out)

    def test_under_target_reported(self) -> None:
        path = self.write("a.txt", REPORT)
        total = polish_stat.count_report(path)["char_count"]
        report = polish_stat.count_report(path, target=total * 2)
        self.assertEqual(report["target"]["status"], "不足")

    # 必过标本：在±10% 以内
    def test_within_target_passes_strict(self) -> None:
        path = self.write("a.txt", REPORT)
        total = polish_stat.count_report(path)["char_count"]
        target = int(round(total * 1.05))
        code, _, _ = self.run_main("count", str(path), "--target", str(target), "--strict")
        self.assertEqual(code, 0)
        report = polish_stat.count_report(path, target=target)
        self.assertEqual(report["target"]["status"], "在范围内")
        self.assertTrue(all("target" in s for s in report["sections"]))

    # 必错标本：用户说“3000字以内”，3001字必须报超出（±10%口径会放过）
    def test_upper_limit_exceeded_by_one(self) -> None:
        path = self.write("u.txt", "排" * 3000 + "。\n")
        self.assertEqual(polish_stat.count_report(path)["char_count"], 3001)
        report = polish_stat.count_report(path, target=3000, upper=True)
        self.assertEqual(report["target"]["status"], "超出")
        self.assertEqual(report["target"]["high"], 3000)
        code, out, _ = self.run_main("count", str(path), "--target", "3000", "--max", "--strict")
        self.assertEqual(code, 1)
        self.assertIn("不超过 3000 字", out)
        # 对照：±10%口径下同一稿件在范围内
        self.assertEqual(polish_stat.count_report(path, target=3000)["target"]["status"], "在范围内")

    # 必过标本：上限以内、不少于七成
    def test_upper_limit_within(self) -> None:
        path = self.write("u.txt", "排" * 2999 + "。\n")
        code, out, _ = self.run_main("count", str(path), "--target", "3000", "--max", "--strict")
        self.assertEqual(code, 0)
        self.assertIn("在范围内", out)

    def test_upper_limit_far_below_warns(self) -> None:
        path = self.write("u.txt", "排" * 1865 + "。\n")
        report = polish_stat.count_report(path, target=3000, upper=True)
        self.assertEqual(report["target"]["status"], "偏少")
        self.assertIn("必保内容", report["target"]["advice"])
        code, _, _ = self.run_main("count", str(path), "--target", "3000", "--max", "--strict")
        self.assertEqual(code, 0)  # 偏少只提示，不算失败

    def test_max_requires_target(self) -> None:
        code, _, err = self.run_main("count", str(self.write("a.txt", REPORT)), "--max")
        self.assertEqual(code, 2)
        self.assertIn("--target", err)

    def test_gb18030_file(self) -> None:
        path = self.write("g.txt", REPORT, encoding="gb18030")
        self.assertEqual(polish_stat.count_report(path)["markers"]["待补"], 1)

    def test_json_output(self) -> None:
        code, out, _ = self.run_main("count", str(self.write("a.txt", REPORT)), "--json")
        self.assertEqual(code, 0)
        data = json.loads(out)
        self.assertIn("char_count", data)
        self.assertIsNone(data["target"])

    def test_invalid_target(self) -> None:
        code, _, err = self.run_main("count", str(self.write("a.txt", REPORT)), "--target", "0")
        self.assertEqual(code, 2)
        self.assertIn("正整数", err)

    @unittest.skipUnless(HAS_DOCX, "未安装 python-docx")
    def test_docx_input(self) -> None:
        import docx as docx_mod  # type: ignore
        document = docx_mod.Document()
        for line in REPORT.splitlines():
            document.add_paragraph(line)
        path = self.tmp / "a.docx"
        document.save(str(path))
        txt = polish_stat.count_report(self.write("a.txt", REPORT))["char_count"]
        self.assertEqual(polish_stat.count_report(path)["char_count"], txt)


class CompareTest(PolishStatCase):

    def test_changes_detected(self) -> None:
        revised = REPORT.replace("今年以来，我局高度重视安全生产工作。", "") \
                        .replace("（二）开展检查。全年开展专项检查36次，整改率【待补：整改率】。\n", "") \
                        .replace("部分企业主体责任落实不到位。", "部分企业主体责任落实不到位，3家企业未建立隐患台账。")
        report = polish_stat.compare_report(self.write("a.txt", REPORT), self.write("b.txt", revised))
        summary = report["paragraph_summary"]
        self.assertEqual(summary["修改"], 2)
        self.assertEqual(summary["删除"], 1)
        self.assertEqual(summary["新增"], 0)
        self.assertLess(report["revised_chars"], report["original_chars"])
        self.assertEqual(report["revised_markers"]["待补"], 0)

    # 必过标本：只改标点、空白的段落不算修改
    def test_whitespace_and_punct_only_is_unchanged(self) -> None:
        revised = REPORT.replace("县政府：", "县政府:").replace("一、主要做法", "一、 主要做法")
        report = polish_stat.compare_report(self.write("a.txt", REPORT), self.write("b.txt", revised))
        self.assertEqual(report["changes"], [])

    def test_added_and_removed_sections(self) -> None:
        revised = REPORT.replace("二、存在问题\n部分企业主体责任落实不到位。\n", "二、下一步打算\n加强督导。\n")
        report = polish_stat.compare_report(self.write("a.txt", REPORT), self.write("b.txt", revised))
        originals = [row["original"] for row in report["sections"]]
        reviseds = [row["revised"] for row in report["sections"]]
        self.assertIn("二、存在问题", originals)
        self.assertIn("二、下一步打算", reviseds)
        # 内容完全不同的段落记为删除和新增，不硬配成“修改”
        self.assertEqual(report["paragraph_summary"]["修改"], 0)
        self.assertEqual(report["paragraph_summary"]["删除"], 2)
        self.assertEqual(report["paragraph_summary"]["新增"], 2)

    # 必错标本：大幅压缩时，不把无关段落配成“修改”
    def test_heavy_compression_pairs_by_content(self) -> None:
        original = REPORT.replace(
            "今年以来，我局高度重视安全生产工作。全年开展专项检查36次，排查隐患218处。",
            "当前安全生产形势复杂严峻，做好安全生产工作意义重大。\n"
            "我局坚持人民至上、生命至上，统筹发展和安全。\n"
            "今年以来，我局高度重视安全生产工作。全年开展专项检查36次，排查隐患218处。",
        )
        revised = REPORT.replace("今年以来，我局高度重视安全生产工作。", "")
        revised = revised.replace("一、主要做法", "二、主要做法")
        report = polish_stat.compare_report(self.write("a.txt", original), self.write("b.txt", revised))
        pairs = [(c["original"][:4], c["revised"][:4]) for c in report["changes"] if c["type"] == "修改"]
        self.assertIn(("今年以来", "全年开展"), pairs)
        self.assertIn(("一、主要", "二、主要"), pairs)
        deleted = [c["original"][:4] for c in report["changes"] if c["type"] == "删除"]
        self.assertEqual(sorted(deleted), sorted(["当前安全", "我局坚持"]))
        self.assertEqual(report["paragraph_summary"]["新增"], 0)

    # 必过标本：调整顺序（如批注“放前面”）记为移动，不算新增
    def test_moved_paragraph(self) -> None:
        line = "（一）压实责任。与12个乡镇签订安全生产责任书。\n"
        revised = REPORT.replace(line, "").replace("县政府：\n", "县政府：\n" + line)
        report = polish_stat.compare_report(self.write("a.txt", REPORT), self.write("b.txt", revised))
        summary = report["paragraph_summary"]
        self.assertEqual(summary["移动"], 1)
        self.assertEqual(summary["新增"], 0)
        self.assertEqual(summary["删除"], 0)

    # 必过标本：一段拆成几段，记为拆分，不算新增
    def test_split_paragraph(self) -> None:
        revised = REPORT.replace(
            "今年以来，我局高度重视安全生产工作。全年开展专项检查36次，排查隐患218处。",
            "今年以来，我局高度重视安全生产工作。\n全年开展专项检查36次，排查隐患218处。",
        )
        report = polish_stat.compare_report(self.write("a.txt", REPORT), self.write("b.txt", revised))
        summary = report["paragraph_summary"]
        self.assertEqual(summary["新增"], 0)
        self.assertEqual(summary["修改"] + summary["拆分"], 2)

    # 必错标本：原稿中找不到来源的段落记为新增
    def test_new_paragraph_flagged(self) -> None:
        revised = REPORT.replace("特此报告。", "我局安全生产工作获评全省先进单位。\n特此报告。")
        report = polish_stat.compare_report(self.write("a.txt", REPORT), self.write("b.txt", revised))
        added = [c["revised"] for c in report["changes"] if c["type"] == "新增"]
        self.assertEqual(len(added), 1)
        self.assertIn("全省先进", added[0])

    # 必过标本：压缩时一段只留小标题句，不得报“新增”
    def test_heading_sentence_kept_not_added(self) -> None:
        long_para = "（一）进一步压实安全生产责任。坚持党政同责、一岗双责、齐抓共管、失职追责，与12个乡镇签订安全生产责任书。"
        original = REPORT.replace("（一）压实责任。与12个乡镇签订安全生产责任书。", long_para)
        revised = original.replace(long_para, "（一）进一步压实安全生产责任。")
        report = polish_stat.compare_report(self.write("a.txt", original), self.write("b.txt", revised))
        summary = report["paragraph_summary"]
        self.assertEqual(summary["新增"], 0)
        self.assertEqual(summary["删除"], 0)
        self.assertEqual(summary["修改"], 1)
        item = [c for c in report["changes"] if c["type"] == "修改"][0]
        self.assertEqual(item["inserted"], [])

    # 必错标本：原段整段保留、句内追加无数字的新荣誉，必须列出多出的文字
    def test_appended_fact_listed(self) -> None:
        revised = REPORT.replace("部分企业主体责任落实不到位。",
                                 "部分企业主体责任落实不到位，我局获评全省安全生产先进单位。")
        report = polish_stat.compare_report(self.write("a.txt", REPORT), self.write("b.txt", revised))
        items = [c for c in report["changes"] if c["type"] == "修改"]
        self.assertEqual(len(items), 1)
        self.assertTrue(any("全省安全生产先进单位" in f for f in items[0]["inserted"]))
        self.assertGreaterEqual(items[0]["inserted_chars"], 12)
        code, out, _ = self.run_main("compare", str(self.tmp / "a.txt"), str(self.tmp / "b.txt"))
        self.assertIn("多出的文字", out)

    # 必错标本：只改一个字的状态升级（拟→已），“多出的文字”列不出来，必须报状态/强度词变化
    def test_status_upgrade_single_char_flagged(self) -> None:
        original = REPORT.replace("部分企业主体责任落实不到位。", "全县燃气管道改造工程拟于11月启动。")
        revised = REPORT.replace("部分企业主体责任落实不到位。", "全县燃气管道改造工程11月已启动。")
        report = polish_stat.compare_report(self.write("a.txt", original), self.write("b.txt", revised))
        items = [c for c in report["changes"] if c["type"] == "修改"]
        self.assertEqual(len(items), 1)
        self.assertEqual(items[0]["inserted"], [])
        changes = {c["word"]: (c["original"], c["revised"]) for c in items[0]["status_changes"]}
        self.assertEqual(changes, {"拟": (1, 0), "已": (0, 1)})
        self.assertEqual(report["status_changed"], 1)
        code, out, _ = self.run_main("compare", str(self.tmp / "a.txt"), str(self.tmp / "b.txt"))
        self.assertIn("状态/强度词变化：已 0→1；拟 1→0", out)

    # 必错标本：义务强度改动（可以→必须）
    def test_obligation_strength_change_flagged(self) -> None:
        original = REPORT.replace("部分企业主体责任落实不到位。", "各企业可以自行组织应急演练。")
        revised = REPORT.replace("部分企业主体责任落实不到位。", "各企业必须自行组织应急演练。")
        report = polish_stat.compare_report(self.write("a.txt", original), self.write("b.txt", revised))
        items = [c for c in report["changes"] if c["type"] == "修改"]
        self.assertEqual(len(items), 1)
        changes = {c["word"]: (c["original"], c["revised"]) for c in items[0]["status_changes"]}
        self.assertEqual(changes, {"必须": (0, 1), "可以": (1, 0)})

    # 必过标本：“重要”“要求”之类的复合词和【待核实】标记里的词不算状态/强度词
    def test_status_words_ignore_compounds_and_markers(self) -> None:
        original = REPORT.replace("部分企业主体责任落实不到位。", "实验室改造已基本完成，按上级要求推进。")
        revised = REPORT.replace(
            "部分企业主体责任落实不到位。",
            "实验室改造已基本完成【待核实：甲稿为基本完成，乙稿为全面完成】，这是重要基础，按上级要求推进。")
        report = polish_stat.compare_report(self.write("a.txt", original), self.write("b.txt", revised))
        items = [c for c in report["changes"] if c["type"] == "修改"]
        self.assertEqual(len(items), 1)
        self.assertEqual(items[0]["status_changes"], [])
        self.assertEqual(report["status_changed"], 0)

    def test_cli_text(self) -> None:
        code, out, _ = self.run_main("compare", str(self.write("a.txt", REPORT)),
                                     str(self.write("b.txt", REPORT)))
        self.assertEqual(code, 0)
        self.assertIn("修改 0，移动 0，拆分 0，删除 0，新增 0", out)
        self.assertNotIn("状态/强度词", out)


class TraceTest(PolishStatCase):

    def trace(self, merged: str):
        return polish_stat.trace_report(
            self.write("m.txt", merged),
            [self.write("甲.txt", DRAFT_A), self.write("乙.txt", DRAFT_B)],
        )

    # 必过标本：每句都来自甲稿或乙稿
    def test_all_sentences_traced(self) -> None:
        report = self.trace(MERGED_OK)
        self.assertEqual(report["unmatched"], 0)
        self.assertEqual(report["partial"], 0)
        sources = {row["sentence"][:6]: row["source"] for row in report["rows"]}
        self.assertEqual(sources["新建校外实习"], "乙稿")
        self.assertEqual(sources["获批省级科研"], "甲稿")
        code, _, _ = self.run_main("trace", str(self.tmp / "m.txt"), str(self.tmp / "甲.txt"),
                                   str(self.tmp / "乙.txt"), "--strict")
        self.assertEqual(code, 0)

    # 必错标本：合并稿新增了两稿都没有的内容
    def test_new_sentence_flagged(self) -> None:
        report = self.trace(MERGED_BAD)
        self.assertEqual(report["unmatched"], 1)
        flagged = [r for r in report["rows"] if r["source"] == "未找到来源"]
        self.assertIn("先进单位", flagged[0]["sentence"])
        code, out, _ = self.run_main("trace", str(self.tmp / "m.txt"), str(self.tmp / "甲.txt"),
                                     str(self.tmp / "乙.txt"), "--strict")
        self.assertEqual(code, 1)
        self.assertIn("未找到来源", out)

    # 必错标本：整句来自甲稿，句内追加了两稿都没有的无数字新内容
    def test_clause_appended_flagged(self) -> None:
        merged = MERGED_OK.replace("获批省级科研项目5项，到账经费320万元。",
                                   "获批省级科研项目5项，到账经费320万元，获评全省教学管理先进单位。")
        report = self.trace(merged)
        self.assertEqual(report["unmatched"], 0)
        self.assertEqual(report["partial"], 1)
        flagged = [r for r in report["rows"] if r["uncovered_clauses"]]
        self.assertIn("先进单位", flagged[0]["uncovered_clauses"][0])
        code, out, _ = self.run_main("trace", str(self.tmp / "m.txt"), str(self.tmp / "甲.txt"),
                                     str(self.tmp / "乙.txt"), "--strict")
        self.assertEqual(code, 1)
        self.assertIn("句内找不到来源的小句", out)

    def test_sentence_combined_from_two_drafts(self) -> None:
        merged = MERGED_OK.replace(
            "获批省级科研项目5项，到账经费320万元。发表论文46篇。",
            "获批省级科研项目5项，到账经费320万元，发表论文46篇。",
        )
        report = self.trace(merged)
        self.assertEqual(report["unmatched"], 0)
        self.assertGreaterEqual(report["by_source"]["多稿合成"] + report["by_source"]["甲稿"], 1)

    def test_unused_source_sentences_listed(self) -> None:
        merged = MERGED_OK.replace("新建校外实习基地4个。", "")
        report = self.trace(merged)
        unused = [row["sentence"] for row in report["unused"]]
        self.assertTrue(any("实习基地" in s for s in unused))

    def test_markers_ignored(self) -> None:
        merged = MERGED_OK.replace("到账经费320万元。", "到账经费【待补：到账经费】万元。")
        report = self.trace(merged)
        self.assertEqual(report["unmatched"], 0)

    def test_too_many_sources(self) -> None:
        paths = [str(self.write(f"s{i}.txt", DRAFT_A)) for i in range(7)]
        code, _, err = self.run_main("trace", str(self.write("m.txt", MERGED_OK)), *paths)
        self.assertEqual(code, 2)
        self.assertIn("最多", err)


class CommentsTest(PolishStatCase):

    def test_comments_extracted_without_author(self) -> None:
        report = polish_stat.comments_report(self.make_docx_zip("c.docx"))
        self.assertEqual(report["comment_count"], 2)
        first, second = report["comments"]
        self.assertEqual(first["comment"], "太虚，数据呢？")
        self.assertEqual(first["anchor"], "我街道垃圾分类工作取得显著成效")
        self.assertEqual(first["paragraph"], 2)
        self.assertEqual(first["date"], "2026-10-12")
        # 没有圈定范围的批注，原文取所在段落
        self.assertEqual(second["anchor"], "下一步，我们将继续加强宣传。")
        self.assertEqual(second["paragraph"], 3)
        self.assertEqual(report["revision_marks"], 2)
        self.assertNotIn("王某某", json.dumps(report, ensure_ascii=False))

    def test_no_comments(self) -> None:
        report = polish_stat.comments_report(self.make_docx_zip("n.docx", comments=False))
        self.assertEqual(report["comment_count"], 0)
        code, out, _ = self.run_main("comments", str(self.tmp / "n.docx"))
        self.assertEqual(code, 0)
        self.assertIn("没有找到批注", out)

    def test_not_docx(self) -> None:
        code, _, err = self.run_main("comments", str(self.write("a.txt", REPORT)))
        self.assertEqual(code, 2)
        self.assertIn(".docx", err)

    def test_broken_docx(self) -> None:
        path = self.write("bad.docx", "这不是压缩包")
        code, _, err = self.run_main("comments", str(path))
        self.assertEqual(code, 2)
        self.assertIn("无法打开", err)


class RevisionTest(PolishStatCase):

    # 必错标本：带修订痕迹的 docx，插入的“218”不能丢，删除的“200”不能读进来
    def test_revisions_accepted_in_count_and_body(self) -> None:
        path = self.make_docx_zip("r.docx", comments=False, document=REVISED_DOC_XML)
        report = polish_stat.count_report(path)
        self.assertEqual(report["revisions"], 2)
        self.assertEqual(report["char_count"], len("全年排查隐患218处。整改210处"))
        text, info = polish_stat.body_text([path])
        self.assertIn("全年排查隐患218处。", text)
        self.assertNotIn("200", text)
        self.assertIn("整改\t210处", text)
        self.assertEqual(info["files"][0]["revisions"], 2)
        code, out, _ = self.run_main("count", str(path))
        self.assertEqual(code, 0)
        self.assertIn("接受全部修订", out)

    def test_comments_note_on_revisions(self) -> None:
        report = polish_stat.comments_report(self.make_docx_zip("c.docx"))
        self.assertIn("body", report["note"])
        plain = REVISED_DOC_XML.replace(
            '<w:del w:id="1" w:author="X"><w:r><w:delText>200</w:delText></w:r></w:del>', "").replace(
            '<w:ins w:id="2" w:author="X"><w:r><w:t>218</w:t></w:r></w:ins>', "<w:r><w:t>218</w:t></w:r>")
        clean = polish_stat.comments_report(self.make_docx_zip("n.docx", document=plain))
        self.assertEqual(clean["revision_marks"], 0)
        self.assertNotIn("body", clean["note"])

    def test_compare_docx_with_txt(self) -> None:
        path = self.make_docx_zip("r.docx", comments=False, document=REVISED_DOC_XML)
        txt = self.write("r.txt", "某公司关于安全生产工作情况的汇报\n全年排查隐患218处。\n整改\t210处\n")
        report = polish_stat.compare_report(path, txt)
        self.assertEqual(report["changes"], [])


class BodyTest(PolishStatCase):

    ORIGINAL = ("某社区关于垃圾分类工作的汇报\n分类准确率由62%提高到81%。\n【起草说明】\n"
                "一、所作假设\n1. 原稿起草说明，含2025年数据。\n")
    MATERIAL = "台账摘要\n志愿者46人，投放点12个。\n"

    # 必过标本：原稿带【起草说明】，拼在后面的材料不能被截掉
    def test_concat_keeps_material_after_draft_note(self) -> None:
        a, b = self.write("原稿.txt", self.ORIGINAL), self.write("材料.txt", self.MATERIAL)
        text, info = polish_stat.body_text([a, b])
        self.assertIn("志愿者46人", text)
        self.assertNotIn("【起草说明】", text)
        self.assertNotIn("2025年", text)
        self.assertTrue(info["files"][0]["draft_note_removed"])

    @unittest.skipUnless(HAS_FACT_DIFF, "未找到 fact_diff.py（先运行同步脚本）")
    def test_concat_then_ledger_has_no_false_additions(self) -> None:
        a, b = self.write("原稿.txt", self.ORIGINAL), self.write("材料.txt", self.MATERIAL)
        info = polish_stat.write_body([a, b], self.tmp / "参照合并.txt")
        revised = self.write("正文.txt", "某社区关于垃圾分类工作的汇报\n分类准确率由62%提高到81%，"
                                         "志愿者46人，投放点12个。\n")
        report = fact_diff.build_report(Path(str(info["out"])), revised, ledger=True)
        self.assertEqual(report["added"], [])
        # 对照：直接拼接（不去掉起草说明）时，材料中的事实会被报为新增
        raw = self.write("直接拼接.txt", self.ORIGINAL + self.MATERIAL)
        self.assertNotEqual(fact_diff.build_report(raw, revised, ledger=True)["added"], [])

    def test_out_does_not_overwrite(self) -> None:
        a = self.write("原稿.txt", self.ORIGINAL)
        existing = self.write("参照合并.txt", "已有内容")
        code, out, _ = self.run_main("body", str(a), "--out", str(existing))
        self.assertEqual(code, 0)
        self.assertEqual(existing.read_text(encoding="utf-8"), "已有内容")
        written = self.tmp / "参照合并-v2.txt"
        self.assertTrue(written.exists())
        self.assertIn("62%", written.read_text(encoding="utf-8"))
        self.assertIn("参照合并-v2.txt", out)

    def test_out_suffix_checked(self) -> None:
        code, _, err = self.run_main("body", str(self.write("原稿.txt", self.ORIGINAL)),
                                     "--out", str(self.tmp / "x.docx"))
        self.assertEqual(code, 2)
        self.assertIn(".txt", err)


class CliTest(PolishStatCase):

    def test_missing_file(self) -> None:
        code, _, err = self.run_main("count", str(self.tmp / "不存在.txt"))
        self.assertEqual(code, 2)
        self.assertIn("未找到文件", err)

    def test_unsupported_suffix(self) -> None:
        code, _, err = self.run_main("count", str(self.write("a.pdf", "x")))
        self.assertEqual(code, 2)
        self.assertIn("只支持", err)

    def test_no_command(self) -> None:
        code, _, _ = self.run_main()
        self.assertEqual(code, 2)

    def test_help_exit_zero(self) -> None:
        code, out, _ = self.run_main("count", "--help")
        self.assertEqual(code, 0)
        self.assertIn("目标字数", out)

    def test_no_files_written(self) -> None:
        path = self.write("a.txt", REPORT)
        before = sorted(p.name for p in self.tmp.iterdir())
        self.run_main("count", str(path), "--target", "50")
        self.run_main("compare", str(path), str(path))
        self.run_main("body", str(path))
        after = sorted(p.name for p in self.tmp.iterdir())
        self.assertEqual(before, after)


if __name__ == "__main__":
    unittest.main()
