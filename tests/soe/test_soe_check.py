# -*- coding: utf-8 -*-
"""skills/gongwen-soe/scripts/soe_check.py 的单元测试。

运行：python -m unittest discover -s tests/soe -t .
必错标本和必过标本在 tests/soe/specimens/ 下。
"""
from __future__ import annotations

import datetime as dt
import io
import json
import re
import shutil
import sys
import tempfile
import unittest
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[2]
SKILL_DIR = REPO_ROOT / "skills" / "gongwen-soe"
SCRIPTS_DIR = SKILL_DIR / "scripts"
SPECIMENS = Path(__file__).resolve().parent / "specimens"
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

import soe_check as sc  # noqa: E402

try:
    import docx  # type: ignore  # noqa: F401
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

AS_OF = dt.date(2026, 9, 30)


def ids_of(result) -> list:
    return [issue["id"] for issue in result["issues"]]


def serious(result) -> list:
    return [issue for issue in result["issues"] if issue["level"] in ("P0", "P1")]


class SoeCheckBase(unittest.TestCase):

    @classmethod
    def setUpClass(cls) -> None:
        cls.ruleset = sc.load_rules()

    def setUp(self) -> None:
        self.tmp = Path(tempfile.mkdtemp(prefix="soe_check_test_"))

    def tearDown(self) -> None:
        shutil.rmtree(self.tmp, ignore_errors=True)

    def check(self, text: str, doc_type=None, external: bool = False):
        return sc.check_text(text, doc_type, external, ruleset=self.ruleset, as_of=AS_OF)

    def check_specimen(self, name: str, doc_type: str, external: bool = False):
        path = next(SPECIMENS.glob(name + ".*"))
        return sc.check_file(path, doc_type, external, ruleset=self.ruleset, as_of=AS_OF)

    def run_main(self, *args: str):
        out, err = io.StringIO(), io.StringIO()
        with redirect_stdout(out), redirect_stderr(err):
            code = sc.main(list(args))
        return code, out.getvalue(), err.getvalue()


class RuleFileTest(SoeCheckBase):
    """规则文件本身的完整性，以及每条正则规则的命中、不命中例句。"""

    def test_rules_load(self) -> None:
        self.assertEqual(len(self.ruleset.rules), 16)
        self.assertEqual(set(self.ruleset.procedural), set(sc.PROCEDURAL_IDS))
        self.assertIn("改革台账", self.ruleset.types)
        self.assertIn("ESG报告", self.ruleset.external_types)

    def test_every_rule_has_examples(self) -> None:
        for rule in self.ruleset.rules:
            with self.subTest(rule=rule.id):
                self.assertTrue(rule.hit_examples, "缺少必中例句")
                self.assertTrue(rule.miss_examples, "缺少必不中例句")
                self.assertTrue(rule.basis)

    def _type_for(self, rule) -> str:
        return rule.types[0] if rule.types else None

    def test_hit_examples(self) -> None:
        for rule in self.ruleset.rules:
            for example in rule.hit_examples:
                with self.subTest(rule=rule.id, example=example):
                    result = self.check(example, self._type_for(rule), external=rule.external_only)
                    self.assertIn(rule.id, ids_of(result))

    def test_miss_examples(self) -> None:
        for rule in self.ruleset.rules:
            for example in rule.miss_examples:
                with self.subTest(rule=rule.id, example=example):
                    result = self.check(example, self._type_for(rule), external=rule.external_only)
                    self.assertNotIn(rule.id, ids_of(result))

    def test_bad_rule_file(self) -> None:
        bad = self.tmp / "rules.json"
        bad.write_text("{not json", encoding="utf-8")
        with self.assertRaises(sc.InputError):
            sc.load_rules(bad)
        bad.write_text(json.dumps({"types": ["议题材料"], "rules": [
            {"id": "X-1", "category": "c", "gate": "g", "level": "P9", "message": "m", "basis": "b", "pattern": "a"}
        ]}, ensure_ascii=False), encoding="utf-8")
        with self.assertRaises(sc.InputError):
            sc.load_rules(bad)


class SpecimenTest(SoeCheckBase):
    """必错标本必须报出预期问题；必过标本不得有 P0、P1。"""

    CASES = (
        ("议题材料", "议题材料", False, {"SOE-001", "SOE-003", "SOE-004", "SOE-008", "SOE-101", "SOE-102", "SOE-107"}),
        ("董事会决议", "决议", False, {"SOE-104"}),
        ("董事会决议-票数写法", "决议", False, {"SOE-104"}),
        ("总经理办公会纪要", "纪要", False, {"SOE-007", "SOE-108", "SOE-109"}),
        ("ESG报告", "ESG报告", True, {"SOE-008", "SOE-009", "SOE-010", "SOE-011", "SOE-012", "SOE-013",
                                     "SOE-106", "SOE-107"}),
        ("融合材料", "融合材料", False, {"SOE-016"}),
        ("经营分析", "经营分析", False, {"SOE-014", "SOE-015", "SOE-107"}),
        ("合规报告", "合规报告", False, {"SOE-005", "SOE-010"}),
        ("改革台账", "改革台账", False, {"SOE-105"}),
    )

    def test_must_fail(self) -> None:
        for name, doc_type, external, expected in self.CASES:
            with self.subTest(specimen=name):
                result = self.check_specimen("必错-" + name, doc_type, external)
                self.assertTrue(expected.issubset(set(ids_of(result))),
                                f"缺少：{expected - set(ids_of(result))}")
                self.assertTrue(serious(result) or result["summary"]["P2"], "必错标本没有报出问题")

    def test_must_pass(self) -> None:
        for name, doc_type, external, _expected in self.CASES:
            with self.subTest(specimen=name):
                result = self.check_specimen("必过-" + name, doc_type, external)
                self.assertEqual(serious(result), [], json.dumps(result["issues"], ensure_ascii=False))

    def test_vote_format_details(self) -> None:
        result = self.check_specimen("必错-董事会决议-票数写法", "决议")
        messages = " ".join(i["message"] for i in result["issues"])
        self.assertIn("票数之和 8", messages)
        self.assertIn("未超过全体董事 9 人的半数", messages)
        self.assertNotIn("未写表决结果", messages)
        self.assertEqual(result["summary"]["P0"], 4)

    def test_resolution_details(self) -> None:
        result = self.check_specimen("必错-董事会决议", "决议")
        messages = " ".join(i["message"] for i in result["issues"])
        self.assertIn("未超过全体董事 9 人的半数", messages)
        self.assertIn("票数之和 4", messages)
        self.assertIn("全票", messages)
        self.assertEqual(result["summary"]["P0"], 4)

    def test_ledger_details(self) -> None:
        result = self.check_specimen("必错-改革台账", "改革台账")
        messages = [i["message"] for i in result["issues"]]
        self.assertTrue(any("有序推进" in m for m in messages))
        self.assertTrue(any("没有佐证材料" in m for m in messages))
        self.assertTrue(any("已过" in m and "按计划推进" in m for m in messages))
        self.assertTrue(any("牵头部门为空" in m for m in messages))
        self.assertTrue(any("滞后任务没有写原因" in m for m in messages))
        self.assertEqual(result["as_of"], "2026-09-30")

    def test_draft_note_cut(self) -> None:
        # 必过-合规报告 的【起草说明】中写了“零违规”，不应报出
        result = self.check_specimen("必过-合规报告", "合规报告")
        self.assertNotIn("SOE-010", ids_of(result))


class ResolutionTest(SoeCheckBase):

    HEAD = "会议应出席董事7人，实际出席董事7人。\n"

    def test_attendance_not_majority(self) -> None:
        result = self.check("会议应出席董事7人，实际出席董事3人。\n审议通过《关于A的议案》。\n表决结果：同意3票，反对0票，弃权0票。", "决议")
        self.assertTrue(any("会议不能举行" in i["message"] for i in result["issues"]))

    def test_actual_more_than_should(self) -> None:
        result = self.check("会议应出席董事5人，实际出席董事6人。\n表决结果：同意6票，反对0票，弃权0票。", "决议")
        self.assertTrue(any("多于应出席" in i["message"] for i in result["issues"]))

    def test_chinese_numerals_and_alt_format(self) -> None:
        text = "会议应出席董事七人，实际出席董事七人。\n审议通过《关于A的议案》。\n表决结果：7票同意，0票反对，0票弃权。"
        self.assertEqual(serious(self.check(text, "决议")), [])
        text_bad = "会议应出席董事七人，实际出席董事七人。\n审议通过《关于A的议案》。\n表决结果：6票同意，0票反对，0票弃权。"
        self.assertIn("SOE-104", ids_of(self.check(text_bad, "决议")))

    def test_recusal_counted(self) -> None:
        ok = self.HEAD + "审议通过《关于关联交易的议案》。关联董事张某、李某回避表决。\n表决结果：同意5票，反对0票，弃权0票。"
        self.assertEqual(serious(self.check(ok, "决议")), [])
        bad = self.HEAD + "审议通过《关于关联交易的议案》。关联董事张某回避表决。\n表决结果：同意5票，反对0票，弃权0票。"
        self.assertIn("SOE-104", ids_of(self.check(bad, "决议")))

    def test_recusal_without_count(self) -> None:
        text = self.HEAD + "审议通过《关于关联交易的议案》。有关联关系的董事回避表决。\n表决结果：同意5票，反对0票，弃权0票。"
        result = self.check(text, "决议")
        self.assertTrue(any(i["level"] == "P2" and "回避人数" in i["message"] for i in result["issues"]))

    def test_recusal_majority_of_unrelated_directors(self) -> None:
        # 9 名董事，5 名关联董事回避，4 名无关联董事中 3 票同意：依法通过（《公司法》第一百八十五条）
        text = ("会议应出席董事9人，实际出席董事9人。\n审议通过《关于关联交易的议案》。关联董事5人回避表决。\n"
                "表决结果：同意3票，反对1票，弃权0票。")
        self.assertEqual(serious(self.check(text, "决议")), [])
        bad = text.replace("同意3票，反对1票", "同意2票，反对2票")
        self.assertTrue(any("无关联关系董事 4 人的半数" in i["message"] for i in self.check(bad, "决议")["issues"]))

    def test_unrelated_directors_less_than_three(self) -> None:
        text = ("会议应出席董事7人，实际出席董事7人。\n审议通过《关于关联交易的议案》。关联董事5人回避表决。\n"
                "表决结果：同意2票，反对0票，弃权0票。")
        result = self.check(text, "决议")
        self.assertTrue(any(i["level"] == "P1" and "不足三人" in i["message"] for i in result["issues"]))

    def test_recusal_unknown_count_no_pass_judgement(self) -> None:
        text = ("会议应出席董事9人，实际出席董事9人。\n审议通过《关于关联交易的议案》。有关联关系的董事回避表决。\n"
                "表决结果：同意4票，反对0票，弃权0票。")
        result = self.check(text, "决议")
        self.assertFalse(any("半数" in i["message"] for i in result["issues"]))
        self.assertTrue(any(i["level"] == "P2" and "回避人数" in i["message"] for i in result["issues"]))

    def test_vote_formats(self) -> None:
        for votes in ("同意票8票，反对票1票，弃权票0票", "同意：8票；反对：1票；弃权：0票", "同意8票，反对1票",
                      "8票同意，1票反对，0票弃权"):
            with self.subTest(votes=votes):
                text = "会议应出席董事9人，实际出席董事9人。\n审议通过《关于A的议案》。\n表决结果：" + votes + "。"
                self.assertEqual(serious(self.check(text, "决议")), [])
                bad = text.replace("8", "4", 1)
                self.assertIn("SOE-104", [i["id"] for i in serious(self.check(bad, "决议"))])

    def test_failed_item_not_flagged(self) -> None:
        text = self.HEAD + "审议《关于B的议案》。\n表决结果：同意3票，反对4票，弃权0票。该议案未通过。"
        self.assertEqual(serious(self.check(text, "决议")), [])

    def test_missing_counts(self) -> None:
        result = self.check("某公司董事会决议\n审议通过《关于A的议案》。", "决议")
        messages = " ".join(i["message"] for i in result["issues"])
        self.assertIn("未写明应出席", messages)
        self.assertIn("未写表决结果", messages)

    def test_to_int(self) -> None:
        self.assertEqual(sc.to_int("7"), 7)
        self.assertEqual(sc.to_int("七"), 7)
        self.assertEqual(sc.to_int("十"), 10)
        self.assertEqual(sc.to_int("十二"), 12)
        self.assertEqual(sc.to_int("二十一"), 21)
        self.assertIsNone(sc.to_int("若干"))


class OtherProceduralTest(SoeCheckBase):

    def test_yian_elements(self) -> None:
        result = self.check("关于A的议案\n各位董事：\n公司拟向某子公司增资。", "议案")
        messages = " ".join(i["message"] for i in result["issues"])
        self.assertIn("提请审议事项", messages)
        self.assertIn("决策程序履行情况", messages)
        good = ("关于A的议案\n各位董事：\n一、决策程序履行情况\n2026年10月20日公司党委会前置研究讨论，同意提交董事会审议。\n"
                "二、提请审议事项\n同意公司向某子公司增资2.4亿元。\n以上议案，请审议。")
        self.assertNotIn("SOE-103", ids_of(self.check(good, "议案")))

    def test_fusion_two_skins(self) -> None:
        text = "某支部党建经验\n党支部开展主题党日活动，党员参与积极。\n公司全年发电量86亿千瓦时。"
        self.assertIn("SOE-110", ids_of(self.check(text, "融合材料")))
        no_data = "某支部党建经验\n党支部开展主题党日活动，党员参与积极。"
        result = self.check(no_data, "融合材料")
        self.assertTrue(any("没有结果数据" in i["message"] for i in result["issues"]))

    def test_minutes_intro_line_skipped(self) -> None:
        text = "总经理办公会纪要\n现将会议议定事项纪要如下：\n一、关于A\n会议同意方案，由综合部牵头，10月31日前完成。"
        self.assertNotIn("SOE-108", ids_of(self.check(text, "纪要")))

    def test_soe108_is_hint_without_placeholder_advice(self) -> None:
        text = "总经理办公会纪要\n一、关于A\n会议同意修订值班制度。"
        issues = [i for i in self.check(text, "纪要")["issues"] if i["id"] == "SOE-108"]
        self.assertEqual(len(issues), 1)
        self.assertEqual(issues[0]["level"], "P2")
        self.assertNotIn("标【待补】", issues[0]["suggestion"])
        self.assertIn("起草说明", issues[0]["suggestion"])

    def test_ledger_without_table(self) -> None:
        result = self.check("改革任务进展：已完成5项。", "改革台账")
        self.assertTrue(any("没有找到" in i["message"] for i in result["issues"]))

    def test_ledger_rate_over_100(self) -> None:
        text = ("| 序号 | 任务 | 牵头部门 | 完成时限 | 完成率 | 状态 | 佐证材料 |\n|---|---|---|---|---|---|---|\n"
                "| 1 | A | 综合部 | 2026年12月 | 120% | 按计划推进 | 无 |")
        self.assertTrue(any("超过 100%" in i["message"] for i in self.check(text, "改革台账")["issues"]))

    def test_parse_deadline(self) -> None:
        self.assertEqual(sc.parse_deadline("2026年6月30日"), dt.date(2026, 6, 30))
        self.assertEqual(sc.parse_deadline("2026年2月"), dt.date(2026, 2, 28))
        self.assertEqual(sc.parse_deadline("2026-09-30"), dt.date(2026, 9, 30))
        self.assertEqual(sc.parse_deadline("2026年底"), dt.date(2026, 12, 31))
        self.assertIsNone(sc.parse_deadline("长期"))
        self.assertIsNone(sc.parse_deadline("2026年13月"))

    def test_parse_deadline_without_year(self) -> None:
        self.assertEqual(sc.parse_deadline("8月31日", 2026), dt.date(2026, 8, 31))
        self.assertEqual(sc.parse_deadline("9月底", 2026), dt.date(2026, 9, 30))
        self.assertEqual(sc.parse_deadline("2026年二季度"), dt.date(2026, 6, 30))
        self.assertEqual(sc.parse_deadline("三季度末", 2026), dt.date(2026, 9, 30))
        self.assertEqual(sc.parse_deadline("上半年", 2026), dt.date(2026, 6, 30))
        self.assertEqual(sc.parse_deadline("年底", 2026), dt.date(2026, 12, 31))
        self.assertIsNone(sc.parse_deadline("8月31日"))
        self.assertIsNone(sc.parse_deadline("13月5日", 2026))

    def test_ledger_deadline_edge_cases(self) -> None:
        text = ("| 序号 | 任务 | 牵头部门 | 完成时限 | 状态 | 佐证材料 | 备注 |\n|---|---|---|---|---|---|---|\n"
                "| 1 | A | 综合部 | 2026年9月30日 | 按计划推进 | 无 |  |\n"
                "| 2 | B | 综合部 | 8月31日 | 按计划推进 | 无 |  |\n"
                "| 3 | C | 综合部 | 2026年二季度 | 未启动 | 无 |  |\n"
                "| 4 | D | 综合部 | 2026年8月31日 | 有序推进 | 无 |  |\n"
                "| 5 | E | 综合部 | 2026年10月31日 | 按计划推进 | 无 |  |\n"
                "| 6 | F | 综合部 | 2026年6月30日 | 滞后 | 无 | 因审批未完成，新时限为11月 |")
        issues = self.check(text, "改革台账")["issues"]
        overdue = [i for i in issues if "状态仍为" in i["message"]]
        self.assertEqual(len(overdue), 4, json.dumps(overdue, ensure_ascii=False))
        self.assertTrue(any("已到" in i["message"] for i in overdue))
        self.assertTrue(any("有序推进" in i["message"] and "不在约定口径" in i["message"] for i in issues))

    def test_minutes_guarantee_and_items(self) -> None:
        text = ("总经理办公会纪要\n2026年11月18日，总经理主持召开总经理办公会，研究仓储中心扩建、担保等事项。\n"
                "一、关于港区仓储中心扩建\n会议原则同意项目建设方向，要求战略投资部比选选址，"
                "【待补：完成时限】前按程序提交公司党委会前置研究讨论后，提请董事会审议。\n"
                "二、关于为参股公司提供担保\n会议同意为某参股公司提供担保2000万元，由财务部办理相关手续，12月10日前完成。")
        result = self.check(text, "纪要")
        self.assertIn("SOE-007", ids_of(result))
        lines_109 = [i["line"] for i in result["issues"] if i["id"] == "SOE-109"]
        self.assertTrue(lines_109 and min(lines_109) >= 5, lines_109)
        marked = text.replace("12月10日前完成。",
                              "12月10日前完成【待核实：本事项是否在董事会授权范围内；是否须先提交党委会前置研究讨论】。")
        marked_ids = ids_of(self.check(marked, "纪要"))
        self.assertNotIn("SOE-007", marked_ids)
        self.assertNotIn("SOE-109", marked_ids)

    def test_money_units_standard_amount_excluded(self) -> None:
        # 制度标准金额（门槛）不计入金额单位统计（W1-SO-3）
        cases = (
            "5000万元以上的投资须提交党委会前置研究。本项目投资1.2亿元。",
            "单笔3000万元（含）以上的对外投资，应经党委会前置研究讨论。本项目投资2.6亿元。",
            "投资额超过3000万元的项目，按规定报集团公司审批。本项目投资2.6亿元。",
            "大额度资金标准为单笔2000万元及以上。本项目借款1.5亿元。",
        )
        for text in cases:
            with self.subTest(text=text):
                self.assertNotIn("SOE-107", ids_of(self.check(text, "议题材料")))

    def test_money_units_still_reported(self) -> None:
        cases = (
            "收入12.3亿元、利润8600万元。",
            "应收账款余额超过8600万元，营业收入12.3亿元。",
            "项目投资2000万元以上，营业收入1.5亿元。",
            "5000万元以上的投资须提交党委会前置研究。本项目投资1.2亿元，其中设备购置4500万元。",
        )
        for text in cases:
            with self.subTest(text=text):
                self.assertIn("SOE-107", ids_of(self.check(text, "议题材料")))

    def test_money_units_specimens(self) -> None:
        passed = self.check_specimen("必过-议题材料-制度标准金额", "议题材料")
        self.assertNotIn("SOE-107", ids_of(passed))
        self.assertEqual(serious(passed), [], json.dumps(passed["issues"], ensure_ascii=False))
        failed = self.check_specimen("必错-议题材料-金额单位", "议题材料")
        self.assertIn("SOE-107", ids_of(failed))

    def test_soe005_board_decides_after_party(self) -> None:
        self.assertNotIn("SOE-005", ids_of(self.check("党委会研究后，董事会审议批准增资方案。")))
        self.assertIn("SOE-005", ids_of(self.check("党委会研究后决定增资1.5亿元。")))

    def test_ratio_with_base_is_p1(self) -> None:
        result = self.check("6月末资产负债率较年初下降1.4%。", "经营分析")
        self.assertEqual([i["level"] for i in result["issues"] if i["id"] == "SOE-014"], ["P1"])

    def test_external_only_rules_need_external(self) -> None:
        text = "公司全年实现营业收入128.6亿元，投标报价为每吨386元，上述数据未公开。"
        self.assertFalse({"SOE-011", "SOE-012", "SOE-013"} & set(ids_of(self.check(text, "经营分析"))))
        result = self.check(text, "对外材料")
        self.assertTrue({"SOE-011", "SOE-012", "SOE-013"} <= set(ids_of(result)))
        self.assertTrue(result["external"])

    def test_external_level_upgrade(self) -> None:
        text = "该技术填补了国内空白。"
        internal = self.check(text, "融合材料")
        external = self.check(text, "融合材料", external=True)
        self.assertEqual([i["level"] for i in internal["issues"] if i["id"] == "SOE-008"], ["P1"])
        self.assertEqual([i["level"] for i in external["issues"] if i["id"] == "SOE-008"], ["P0"])

    def test_no_type_runs_general_rules_only(self) -> None:
        result = self.check("项目已先行开工。\n会议决定投资5亿元。")
        self.assertIn("SOE-001", ids_of(result))
        self.assertNotIn("SOE-007", ids_of(result))
        self.assertTrue(any("未指定 --type" in n for n in result["notes"]))


class CliTest(SoeCheckBase):

    def test_json_output_and_strict(self) -> None:
        path = next(SPECIMENS.glob("必错-议题材料.*"))
        code, out, _err = self.run_main(str(path), "--type", "议题材料", "--json")
        self.assertEqual(code, 0)
        data = json.loads(out)
        for key in ("file", "doc_type", "char_count", "summary", "issues"):
            self.assertIn(key, data)
        code, _out, _err = self.run_main(str(path), "--type", "议题材料", "--strict")
        self.assertEqual(code, 1)

    def test_strict_pass(self) -> None:
        path = next(SPECIMENS.glob("必过-议题材料.*"))
        code, out, _err = self.run_main(str(path), "--type", "议题材料", "--strict")
        self.assertEqual(code, 0)
        self.assertIn("未发现行业规则问题", out)

    def test_as_of_option(self) -> None:
        path = next(SPECIMENS.glob("必错-改革台账.*"))
        code, out, _err = self.run_main(str(path), "--type", "改革台账", "--as-of", "2026-01-01", "--json")
        self.assertEqual(code, 0)
        data = json.loads(out)
        self.assertEqual(data["as_of"], "2026-01-01")
        self.assertFalse(any("已过" in i["message"] for i in data["issues"]))
        code, _out, _err = self.run_main(str(path), "--type", "改革台账", "--as-of", "2026/01/01")
        self.assertEqual(code, 2)

    def test_input_errors(self) -> None:
        code, _out, err = self.run_main(str(self.tmp / "不存在.txt"))
        self.assertEqual(code, 2)
        self.assertIn("文件不存在", err)
        bad = self.tmp / "a.pdf"
        bad.write_text("x", encoding="utf-8")
        self.assertEqual(self.run_main(str(bad))[0], 2)
        good = self.tmp / "a.txt"
        good.write_text("正文", encoding="utf-8")
        self.assertEqual(self.run_main(str(good), "--type", "不存在的类型")[0], 2)
        self.assertEqual(self.run_main()[0], 2)

    def test_list_rules_and_help(self) -> None:
        code, out, _err = self.run_main("--list-rules")
        self.assertEqual(code, 0)
        self.assertIn("SOE-104", out)
        code, out, _err = self.run_main("--help")
        self.assertEqual(code, 0)
        self.assertIn("用法", out)

    def test_gb18030_input(self) -> None:
        path = self.tmp / "gbk.txt"
        path.write_bytes("项目已先行开工。".encode("gb18030"))
        code, out, _err = self.run_main(str(path), "--json")
        self.assertEqual(code, 0)
        self.assertIn("SOE-001", out)

    def test_does_not_write_files(self) -> None:
        path = self.tmp / "a.txt"
        path.write_text("项目已先行开工。", encoding="utf-8")
        before = sorted(p.name for p in self.tmp.iterdir())
        self.run_main(str(path), "--json")
        self.assertEqual(sorted(p.name for p in self.tmp.iterdir()), before)


class EvalIsolationTest(unittest.TestCase):
    """评测材料中的关键数字和日期不得出现在 SKILL.md 和 references/ 中，避免模板示例替评测作答。"""

    TOKEN_RE = re.compile(r"\d+(?:\.\d+)?\s*(?:亿元|万元)|(?<![\d.])\d+\.\d+(?![\d.])|\d+(?:\.\d+)?\s*[%％]"
                          r"|\d{1,2}月\d{1,2}日|同意\s*\d+\s*票|\d+\s*天|\d+\s*名")
    # 通用的期末日期、交易所披露截止日和工具阈值，不算重合
    ALLOWED = {"3月31日", "6月30日", "9月30日", "12月31日", "4月30日", "1月1日", "30%", "100%"}

    @staticmethod
    def _norm(text: str) -> str:
        return re.sub(r"\s", "", text).replace("％", "%")

    def _eval_tokens(self) -> dict:
        texts = [(f.name, f.read_text(encoding="utf-8")) for f in (SKILL_DIR / "evals" / "files").iterdir() if f.is_file()]
        data = json.loads((SKILL_DIR / "evals" / "evals.json").read_text(encoding="utf-8"))
        texts += [(f"evals.json 第{e['id']}条", e["prompt"]) for e in data["evals"]]
        tokens: dict = {}
        for name, text in texts:
            for m in self.TOKEN_RE.finditer(text):
                token = self._norm(m.group(0))
                if token not in self.ALLOWED:
                    tokens.setdefault(token, name)
        return tokens

    def test_eval_data_not_in_references(self) -> None:
        tokens = self._eval_tokens()
        self.assertTrue(tokens)
        files = [SKILL_DIR / "SKILL.md"] + [p for p in (SKILL_DIR / "references").rglob("*.md") if "_shared" not in p.parts]
        found = []
        for path in files:
            text = self._norm(path.read_text(encoding="utf-8"))
            for token, source in tokens.items():
                for m in re.finditer(re.escape(token), text):
                    s, e = m.start(), m.end()
                    if s and (text[s - 1].isdigit() or text[s - 1] == "."):
                        continue
                    if token[-1].isdigit() and e < len(text) and (text[e].isdigit() or text[e] == "."):
                        continue
                    found.append(f"{path.relative_to(SKILL_DIR)}：{token}（来自 {source}）")
                    break
        self.assertEqual(found, [])


@unittest.skipUnless(HAS_DOCX, "未安装 python-docx")
class DocxTest(SoeCheckBase):

    def test_docx_ledger_table(self) -> None:
        import docx as _docx  # type: ignore
        document = _docx.Document()
        document.add_paragraph("某集团改革任务台账")
        table = document.add_table(rows=3, cols=5)
        rows = [("序号", "改革任务", "完成时限", "状态", "佐证材料"),
                ("1", "完善授权制度", "2026年6月30日", "按计划推进", "无"),
                ("2", "专业化整合", "2026年12月", "已完成", "无")]
        for r, values in enumerate(rows):
            for c, value in enumerate(values):
                table.cell(r, c).text = value
        document.add_paragraph("【起草说明】")
        document.add_paragraph("项目已先行开工。")
        path = self.tmp / "台账.docx"
        document.save(str(path))
        result = sc.check_file(path, "改革台账", ruleset=self.ruleset, as_of=AS_OF)
        messages = " ".join(i["message"] for i in result["issues"])
        self.assertIn("已过", messages)
        self.assertIn("没有佐证材料", messages)
        self.assertNotIn("SOE-001", ids_of(result))


if __name__ == "__main__":
    unittest.main()
