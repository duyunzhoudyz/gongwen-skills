# -*- coding: utf-8 -*-
"""shared/scripts/md2xlsx.py 的单元测试。

运行：python -m unittest discover -s tests/rectify -t .
"""
from __future__ import annotations

import importlib.util
import io
import shutil
import sys
import tempfile
import unittest
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path
from unittest import mock

REPO_ROOT = Path(__file__).resolve().parents[2]
SCRIPTS_DIR = REPO_ROOT / "shared" / "scripts"


def _load(name: str, path: Path):
    """按文件路径加载，避免与其他 skill 的同名副本（scripts/md2xlsx.py）混淆。"""
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


md2xlsx = _load("rectify_tests_md2xlsx", SCRIPTS_DIR / "md2xlsx.py")

try:
    import openpyxl  # noqa: F401
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

LEDGER_MD = """# 某局巡察整改台账

| 编号 | 问题表现 | 整改措施 | 完成时限 | 整改状态 | 金额（万元） |
| :---: | --- | --- | --- | --- | ---: |
| 01 | 部分科室**高值耗材**领用登记不全 | 1. 全面盘点<br>2. 上线系统 | 2026年12月31日前 | 持续推进 | 125.6 |
| 2 | 联系电话 138****0000 | **立即整改** | 2026-11-30 | 已完成 | 1.10 |
| 3 | =1+1 \\| 竖线 | [制度文件](http://example.invalid) | 长期坚持 | 长期坚持 | 3216 |
"""


class Md2xlsxCase(unittest.TestCase):

    def setUp(self) -> None:
        self.tmp = Path(tempfile.mkdtemp(prefix="md2xlsx_test_"))

    def tearDown(self) -> None:
        shutil.rmtree(self.tmp, ignore_errors=True)

    def write(self, name: str, text: str) -> Path:
        path = self.tmp / name
        path.write_text(text, encoding="utf-8")
        return path

    def run_main(self, *args: str):
        out, err = io.StringIO(), io.StringIO()
        with redirect_stdout(out), redirect_stderr(err):
            code = md2xlsx.main(list(args))
        return code, out.getvalue(), err.getvalue()


class ParseTest(Md2xlsxCase):

    def test_basic_table_with_title_and_aligns(self):
        tables = md2xlsx.parse_tables(LEDGER_MD)
        self.assertEqual(len(tables), 1)
        table = tables[0]
        self.assertEqual(table.title, "某局巡察整改台账")
        self.assertEqual(table.header[0], "编号")
        self.assertEqual(table.aligns[0], "center")
        self.assertEqual(table.aligns[5], "right")
        self.assertIsNone(table.aligns[1])
        self.assertEqual(len(table.rows), 3)

    def test_cell_cleanup(self):
        rows = md2xlsx.parse_tables(LEDGER_MD)[0].rows
        self.assertEqual(rows[0][1], "部分科室高值耗材领用登记不全")  # 行内加粗去掉
        self.assertEqual(rows[0][2], "1. 全面盘点\n2. 上线系统")  # <br> 转换行
        self.assertEqual(rows[1][2], "立即整改")  # 整格加粗去掉
        self.assertEqual(rows[2][1], "=1+1 | 竖线")  # 转义竖线
        self.assertEqual(rows[2][2], "制度文件")  # 链接只留文字

    def test_masked_numbers_kept(self):
        self.assertEqual(md2xlsx.clean_cell("138****0000"), "138****0000")
        self.assertEqual(md2xlsx.clean_cell("110101********1234"), "110101********1234")
        self.assertEqual(md2xlsx.clean_cell("张**、李**"), "张**、李**")

    def test_table_without_outer_pipes(self):
        text = "序号 | 事项\n--- | ---\n1 | 盘点\n"
        tables = md2xlsx.parse_tables(text)
        self.assertEqual(len(tables), 1)
        self.assertEqual(tables[0].rows, [["1", "盘点"]])

    def test_ragged_rows_are_padded_and_truncated(self):
        text = "| a | b | c |\n|---|---|---|\n| 1 |\n| 1 | 2 | 3 | 4 |\n"
        rows = md2xlsx.parse_tables(text)[0].rows
        self.assertEqual(rows, [["1", "", ""], ["1", "2", "3"]])

    def test_code_fence_skipped(self):
        text = "```\n| a | b |\n| - | - |\n| 1 | 2 |\n```\n\n| x | y |\n|---|---|\n| 3 | 4 |\n"
        tables = md2xlsx.parse_tables(text)
        self.assertEqual(len(tables), 1)
        self.assertEqual(tables[0].header, ["x", "y"])

    def test_not_a_table(self):
        # 列数不一致、水平线、普通竖线文字都不是表格
        self.assertEqual(md2xlsx.parse_tables("a | b\n---\n1 | 2\n"), [])
        self.assertEqual(md2xlsx.parse_tables("标题\n---\n正文"), [])
        self.assertEqual(md2xlsx.parse_tables("说明：甲|乙\n下一行"), [])

    def test_table_ends_at_blank_line(self):
        text = "| a | b |\n|---|---|\n| 1 | 2 |\n\n| 3 | 4 |\n"
        tables = md2xlsx.parse_tables(text)
        self.assertEqual(len(tables), 1)
        self.assertEqual(len(tables[0].rows), 1)

    def test_title_detection(self):
        cases = {
            "## 整改任务清单\n": "整改任务清单",
            "表1 问题清单\n": "表1 问题清单",
            "附：整改台账\n": "整改台账",
            "**责任清单**\n": "责任清单",  # 加粗的表题去掉标记
            "**以下内容仅供核对**\n": None,  # 不以"表""清单""台账"结尾的普通段落不作表题
            "以下是普通说明文字。\n": None,
        }
        for prefix, expected in cases.items():
            with self.subTest(prefix=prefix):
                table = md2xlsx.parse_tables(prefix + "| a | b |\n|---|---|\n| 1 | 2 |\n")[0]
                self.assertEqual(table.title, expected)

    def test_to_number(self):
        self.assertEqual(md2xlsx.to_number("3216"), 3216)
        self.assertEqual(md2xlsx.to_number("125.6"), 125.6)
        for text in ("01", "1.10", "2026-11-30", "66.7%", "1234567890123", "-0", "约3", ""):
            with self.subTest(text=text):
                self.assertIsNone(md2xlsx.to_number(text))

    def test_display_width(self):
        self.assertEqual(md2xlsx.display_width("ab"), 2)
        self.assertEqual(md2xlsx.display_width("整改"), 4)

    def test_sanitize_sheet_name(self):
        used = []
        first = md2xlsx.sanitize_sheet_name("整改:台账/2026", used, 1)
        self.assertEqual(first, "整改台账2026")
        used.append(first)
        self.assertEqual(md2xlsx.sanitize_sheet_name("整改:台账/2026", used, 2), "整改台账2026(2)")
        self.assertLessEqual(len(md2xlsx.sanitize_sheet_name("长" * 40, [], 1)), 31)
        self.assertEqual(md2xlsx.sanitize_sheet_name("", [], 3), "表3")

    def test_read_gb18030(self):
        path = self.tmp / "gbk.md"
        path.write_bytes("| 编号 | 问题 |\n|---|---|\n| 1 | 盘点 |\n".encode("gb18030"))
        self.assertIn("盘点", md2xlsx.read_text(path))


@unittest.skipUnless(HAS_OPENPYXL, "需要 openpyxl")
class ConvertTest(Md2xlsxCase):

    def convert(self, text: str, **kwargs):
        src = self.write("台账.md", text)
        target, tables, names = md2xlsx.convert(src, self.tmp / "台账.xlsx", **kwargs)
        return target, tables, names

    def test_styles_and_values(self):
        target, _, names = self.convert(LEDGER_MD)
        wb = openpyxl.load_workbook(target)
        ws = wb[names[0]]
        self.assertEqual(ws["A1"].value, "某局巡察整改台账")
        self.assertIn("A1:F1", [str(r) for r in ws.merged_cells.ranges])
        self.assertTrue(ws["A1"].font.bold)
        self.assertEqual(ws["A1"].font.name, "宋体")
        header = ws["B2"]
        self.assertEqual(header.value, "问题表现")
        self.assertTrue(header.font.bold)
        self.assertEqual(header.fill.fgColor.rgb[-6:], "D9D9D9")
        self.assertEqual(header.border.left.style, "thin")
        self.assertEqual(ws.freeze_panes, "A3")
        self.assertEqual(ws.print_title_rows, "$2:$2")
        self.assertEqual(ws.page_setup.orientation, "landscape")
        # 编号列保留文本（01 不变成 1）；金额列能转数值的转数值
        self.assertEqual(ws["A3"].value, "01")
        self.assertEqual(ws["A4"].value, "2")
        self.assertEqual(ws["F3"].value, 125.6)
        self.assertEqual(ws["F4"].value, "1.10")
        self.assertEqual(ws["F5"].value, 3216)
        # 以 = 开头的文字按文本写入
        self.assertEqual(ws["B5"].value, "=1+1 | 竖线")
        self.assertEqual(ws["B5"].data_type, "s")
        self.assertTrue(ws["C3"].alignment.wrap_text)
        self.assertEqual(ws["C3"].value, "1. 全面盘点\n2. 上线系统")
        self.assertGreater(ws.column_dimensions["B"].width, ws.column_dimensions["A"].width)

    def test_title_override_and_no_title(self):
        target, _, names = self.convert(LEDGER_MD, title="关于落实巡察反馈意见的整改台账")
        ws = openpyxl.load_workbook(target)[names[0]]
        self.assertEqual(ws["A1"].value, "关于落实巡察反馈意见的整改台账")
        target2, _, names2 = md2xlsx.convert(self.write("无题.md", "| a | b |\n|---|---|\n| 1 | 2 |\n"),
                                             self.tmp / "无题.xlsx")
        ws2 = openpyxl.load_workbook(target2)[names2[0]]
        self.assertEqual(ws2["A1"].value, "a")
        self.assertEqual(ws2.freeze_panes, "A2")
        self.assertEqual(ws2.page_setup.orientation, None)

    def test_all_tables(self):
        text = "## 问题清单\n| 编号 | 问题 |\n|---|---|\n| 1 | 甲 |\n\n## 责任清单\n| 编号 | 责任部门 |\n|---|---|\n| 1 | 办公室 |\n"
        target, tables, names = self.convert(text, all_tables=True)
        self.assertEqual(names, ["问题清单", "责任清单"])
        wb = openpyxl.load_workbook(target)
        self.assertEqual(wb.sheetnames, ["问题清单", "责任清单"])
        target1, tables1, _ = md2xlsx.convert(self.tmp / "台账.md", self.tmp / "单表.xlsx")
        self.assertEqual(len(tables1), 1)

    def test_no_overwrite(self):
        src = self.write("台账.md", LEDGER_MD)
        out = self.tmp / "台账.xlsx"
        out.write_bytes(b"existing")
        first, _, _ = md2xlsx.convert(src, out)
        second, _, _ = md2xlsx.convert(src, out)
        self.assertEqual(first.name, "台账-v2.xlsx")
        self.assertEqual(second.name, "台账-v3.xlsx")
        self.assertEqual(out.read_bytes(), b"existing")

    def test_cli_success(self):
        src = self.write("台账.md", LEDGER_MD)
        code, out, err = self.run_main(str(src), str(self.tmp / "out.xlsx"), "--title", "整改台账")
        self.assertEqual(code, 0, err)
        self.assertIn("已生成", out)
        self.assertTrue((self.tmp / "out.xlsx").is_file())


class CliErrorTest(Md2xlsxCase):

    def test_missing_input(self):
        code, _, err = self.run_main(str(self.tmp / "无.md"), str(self.tmp / "o.xlsx"))
        self.assertEqual(code, 2)
        self.assertIn("未找到输入文件", err)

    def test_bad_output_suffix(self):
        src = self.write("t.md", LEDGER_MD)
        code, _, err = self.run_main(str(src), str(self.tmp / "o.xls"))
        self.assertEqual(code, 2)
        self.assertIn(".xlsx", err)

    @unittest.skipUnless(HAS_OPENPYXL, "需要 openpyxl")
    def test_no_table(self):
        src = self.write("t.md", "# 只有标题\n\n正文没有表格。\n")
        code, _, err = self.run_main(str(src), str(self.tmp / "o.xlsx"))
        self.assertEqual(code, 2)
        self.assertIn("未找到 Markdown 表格", err)

    def test_openpyxl_missing(self):
        src = self.write("t.md", LEDGER_MD)
        with mock.patch.dict(sys.modules, {"openpyxl": None}):
            code, _, err = self.run_main(str(src), str(self.tmp / "o.xlsx"))
        self.assertEqual(code, 2)
        self.assertIn("pip install openpyxl", err)

    def test_help(self):
        code, out, _ = self.run_main("--help")
        self.assertEqual(code, 0)
        self.assertIn("用法", out)

    def test_bad_arguments(self):
        code, _, _ = self.run_main()
        self.assertEqual(code, 2)


if __name__ == "__main__":
    unittest.main()
