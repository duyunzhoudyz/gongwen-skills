# -*- coding: utf-8 -*-
"""skills/gongwen-rectify/scripts/rectify_check.py 的单元测试。

必过标本：fixtures/台账-标准.md + fixtures/报告-必过.txt，不报任何问题；
必错标本：fixtures/台账-标准.md + fixtures/报告-必错.txt，报出约定的 P0、P1。

运行：python -m unittest discover -s tests/rectify -t .
"""
from __future__ import annotations

import datetime as dt
import importlib.util
import io
import json
import shutil
import sys
import tempfile
import unittest
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[2]
SCRIPTS_DIR = REPO_ROOT / "skills" / "gongwen-rectify" / "scripts"
FIXTURES = Path(__file__).resolve().parent / "fixtures"


def _load(name: str, path: Path):
    """按文件路径加载，避免与其他 skill 的同名脚本混淆。"""
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


rc = _load("rectify_tests_rectify_check", SCRIPTS_DIR / "rectify_check.py")

try:
    import openpyxl  # noqa: F401
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

try:
    import docx  # type: ignore  # noqa: F401
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

LEDGER_HEAD = (
    "| 编号 | 问题表现 | 整改措施 | 责任领导 | 责任部门 | 完成时限 | 成效标准 | 整改状态 | 佐证材料 |\n"
    "|---|---|---|---|---|---|---|---|---|\n"
)


def row(no, status, evidence="文件", problem=None, deadline="2026年12月31日"):
    problem = problem or f"问题表现{no}"
    return f"| {no} | {problem} | 措施 | 分管领导 | 办公室 | {deadline} | 标准 | {status} | {evidence} |\n"


class RectifyCase(unittest.TestCase):

    def setUp(self) -> None:
        self.tmp = Path(tempfile.mkdtemp(prefix="rectify_test_"))

    def tearDown(self) -> None:
        shutil.rmtree(self.tmp, ignore_errors=True)

    def write(self, name: str, text: str) -> Path:
        path = self.tmp / name
        path.write_text(text, encoding="utf-8")
        return path

    def check(self, ledger_text: str, report_text=None, **kwargs) -> dict:
        ledger = self.write("台账.md", ledger_text)
        report = self.write("报告.txt", report_text) if report_text is not None else None
        return rc.build_report(ledger, report, **kwargs)

    @staticmethod
    def codes(result, level=None):
        return [i["id"] for i in result["issues"] if level is None or i["level"] == level]

    def run_main(self, *args: str):
        out, err = io.StringIO(), io.StringIO()
        with redirect_stdout(out), redirect_stderr(err):
            code = rc.main(list(args))
        return code, out.getvalue(), err.getvalue()


class StatusTest(unittest.TestCase):

    def test_canonical_and_synonyms(self):
        self.assertEqual(rc.normalize_status("已完成"), ("已完成", ""))
        self.assertEqual(rc.normalize_status("已整改")[0], "已完成")
        self.assertEqual(rc.normalize_status("正在推进")[0], "持续推进")
        self.assertEqual(rc.normalize_status("已立行立改")[0], "立行立改")
        self.assertEqual(rc.normalize_status("需长期坚持")[0], "长期坚持")
        self.assertEqual(rc.normalize_status("尚未启动")[0], "未启动")
        self.assertTrue(rc.normalize_status("正在推进")[1].startswith("synonym"))

    def test_with_notes(self):
        self.assertEqual(rc.normalize_status("已完成（11月20日）")[0], "已完成")
        self.assertEqual(rc.normalize_status("持续推进，预计12月底完成")[0], "持续推进")
        self.assertEqual(rc.normalize_status("**已完成**")[0], "已完成")

    def test_nonstandard_unknown_empty_marker(self):
        self.assertEqual(rc.normalize_status("基本完成"), (None, "nonstandard:基本完成"))
        self.assertEqual(rc.normalize_status("持续整改")[1], "nonstandard:持续整改")
        self.assertEqual(rc.normalize_status("看情况")[1], "unknown")
        self.assertEqual(rc.normalize_status("")[1], "empty")
        self.assertEqual(rc.normalize_status("【待补：状态】")[1], "marker")
        # "已完成阶段性目标"不能当作已完成
        self.assertIsNone(rc.normalize_status("已完成阶段性目标")[0])

    def test_aggregate(self):
        self.assertEqual(rc.aggregate_status(["已完成", "立行立改"]), "已完成")
        self.assertEqual(rc.aggregate_status(["立行立改", "立行立改"]), "立行立改")
        self.assertEqual(rc.aggregate_status(["已完成", "持续推进"]), "持续推进")
        self.assertEqual(rc.aggregate_status(["已完成", "长期坚持"]), "长期坚持")
        self.assertEqual(rc.aggregate_status(["未启动", "未启动"]), "未启动")
        self.assertEqual(rc.aggregate_status(["未启动", "已完成"]), "持续推进")
        self.assertIsNone(rc.aggregate_status(["已完成", None]))

    def test_normalize_id(self):
        self.assertEqual(rc.normalize_id("01"), "1")
        self.assertEqual(rc.normalize_id("问题3"), "3")
        self.assertEqual(rc.normalize_id("３－１"), "3-1")
        self.assertEqual(rc.normalize_id("xc-03"), "XC-3")
        self.assertEqual(rc.normalize_id("3.0"), "3")
        self.assertEqual(rc.normalize_id("第5项"), "5")

    def test_parse_deadline(self):
        self.assertEqual(rc.parse_deadline("2026年12月31日前"), dt.date(2026, 12, 31))
        self.assertEqual(rc.parse_deadline("2026年11月底"), dt.date(2026, 11, 30))
        self.assertEqual(rc.parse_deadline("2026年底"), dt.date(2026, 12, 31))
        self.assertEqual(rc.parse_deadline("2026-09-30"), dt.date(2026, 9, 30))
        self.assertEqual(rc.parse_deadline("1. 2026年11月15日前\n2. 2026年12月31日前"), dt.date(2026, 12, 31))
        self.assertIsNone(rc.parse_deadline("长期坚持"))

    def test_cn_to_int(self):
        self.assertEqual(rc.cn_to_int("12"), 12)
        self.assertEqual(rc.cn_to_int("十二"), 12)
        self.assertEqual(rc.cn_to_int("二十"), 20)
        self.assertEqual(rc.cn_to_int("三"), 3)


class FixtureTest(RectifyCase):
    """必过、必错标本。"""

    def test_must_pass(self):
        result = rc.build_report(FIXTURES / "台账-标准.md", FIXTURES / "报告-必过.txt",
                                 FIXTURES / "反馈原文.txt", dt.date(2026, 10, 31))
        self.assertEqual(result["issues"], [], json.dumps(result["issues"], ensure_ascii=False, indent=1))
        summary = result["ledger"]["summary"]
        self.assertEqual(summary["total"], 5)
        self.assertEqual(summary["done"], 2)
        self.assertEqual(summary["done_rate"], 40.0)

    def test_must_fail(self):
        result = rc.build_report(FIXTURES / "台账-标准.md", FIXTURES / "报告-必错.txt")
        p0 = self.codes(result, "P0")
        p1 = self.codes(result, "P1")
        self.assertEqual(p0.count("RC-101"), 2)  # 问题3、问题4写成已完成
        self.assertIn("RC-201", p0)  # 完成数 4 高于台账 2
        self.assertIn("RC-202", p0)  # 完成率 80% 高于 40%
        self.assertIn("RC-203", p0)  # "整改工作已全部完成"
        for code in ("RC-102", "RC-105", "RC-106", "RC-108", "RC-204"):
            self.assertIn(code, p1, code)
        overclaimed = {i["item"] for i in result["issues"] if i["id"] == "RC-101"}
        self.assertEqual(overclaimed, {"3", "4"})  # "该问题已整改完毕。""该问题已完成。"
        self.assertEqual(result["summary"]["P0"], len(p0))
        # 统计句"已有4个问题完成整改"、全称句"5个问题均已整改到位"
        excerpts = " ".join(i["excerpt"] for i in result["issues"] if i["level"] == "P0")
        self.assertGreaterEqual(p0.count("RC-201"), 2, excerpts)
        self.assertGreaterEqual(p0.count("RC-203"), 2, excerpts)
        self.assertIn("已有4个问题完成整改", excerpts)
        self.assertIn("均已整改到位", excerpts)

    def test_must_fail_feedback(self):
        """台账删了反馈原文的数字和引号内容：RC-302。"""
        ledger = self.write("台账.md", (FIXTURES / "台账-标准.md").read_text(encoding="utf-8")
                            .replace("，超支金额共0.38万元", "")
                            .replace("政府采购“化整为零”规避公开招标，", ""))
        result = rc.build_report(ledger, None, FIXTURES / "反馈原文.txt")
        lost = {i["item"]: i["message"] for i in result["issues"] if i["id"] == "RC-302"}
        self.assertEqual(set(lost), {"1", "3"})
        self.assertIn("超支金额共0.38万元", lost["1"])
        self.assertIn("化整为零", lost["3"])


class ReportCompareTest(RectifyCase):

    LEDGER = LEDGER_HEAD + row(1, "已完成") + row(2, "持续推进", "") + row(3, "长期坚持")

    def test_future_and_negated_claims_are_not_done(self):
        report = (
            "问题1：已完成整改。\n"
            "问题2：制度修订正在推进，预计12月31日前完成整改。\n"
            "问题2：尚未完成整改，确保年底前全部整改到位。\n"
            "问题3：取得阶段性成效，需长期坚持。\n"
        )
        result = self.check(self.LEDGER, report)
        self.assertEqual(self.codes(result, "P0"), [])

    def test_past_claim_with_qian_is_done(self):
        result = self.check(self.LEDGER, "问题1：已完成整改。\n问题2：已于10月20日前完成整改。\n问题3：需长期坚持。\n")
        self.assertIn("RC-101", self.codes(result, "P0"))

    def test_partial_subunits_not_flagged(self):
        report = "问题1：已完成整改。\n问题2：已督促3个科室完成整改，其余2个科室正在推进。\n问题3：需长期坚持。\n"
        result = self.check(self.LEDGER, report)
        self.assertNotIn("RC-101", self.codes(result))

    def test_declared_status(self):
        report = "问题1：整改状态：已完成。\n问题2：整改状态：已完成。\n问题3：整改状态：长期坚持。\n"
        result = self.check(self.LEDGER, report)
        items = [i["item"] for i in result["issues"] if i["id"] == "RC-101"]
        self.assertEqual(items, ["2"])

    def test_longterm_hiding_ongoing(self):
        result = self.check(self.LEDGER, "问题1：已完成整改。\n问题2：需长期坚持。\n问题3：需长期坚持。\n")
        self.assertIn("RC-103", self.codes(result, "P1"))

    def test_grouped_mentions_share_status(self):
        report = "问题1已完成整改。问题2、问题3已全部整改到位。\n"
        result = self.check(self.LEDGER, report)
        items = sorted(i["item"] for i in result["issues"] if i["id"] == "RC-101")
        self.assertEqual(items, ["2", "3"])

    def test_multiple_items_in_one_line(self):
        result = self.check(self.LEDGER, "问题1已完成整改；问题2正在推进；问题3需长期坚持。\n")
        self.assertEqual(self.codes(result, "P0"), [])
        self.assertEqual(self.codes(result, "P1"), [])

    def test_block_continues_to_next_paragraph_until_heading(self):
        report = (
            "（一）问题2：制度不健全。\n"
            "整改情况：修订稿已提交审议。该问题已完成整改。\n"
            "二、下一步打算\n"
            "已完成整改的问题要巩固成果。\n"
        )
        result = self.check(self.LEDGER, report)
        self.assertEqual([i["item"] for i in result["issues"] if i["id"] == "RC-101"], ["2"])

    def test_quoted_problem_is_not_a_status_claim(self):
        ledger = LEDGER_HEAD + row(1, "未启动", "", problem="上次巡察指出的问题仍有1个社区未整改到位")
        report = "问题1：上次巡察指出的问题仍有1个社区未整改到位。该问题尚未启动，计划于11月启动。\n"
        self.assertEqual(self.check(ledger, report)["issues"], [])

    def test_plan_mode(self):
        report = "问题1：已完成整改。\n问题2：措施【待补：整改措施】。\n问题3：建立长效机制。\n"
        normal = self.check(self.LEDGER, report)
        self.assertIn("RC-104", self.codes(normal, "P2"))
        ledger = self.write("台账.md", self.LEDGER)
        plan = rc.build_report(ledger, self.write("方案.txt", report), plan=True)
        self.assertNotIn("RC-104", self.codes(plan))
        overclaim = rc.build_report(ledger, self.write("方案2.txt", report + "问题2：已完成整改。\n"), plan=True)
        self.assertIn("RC-101", self.codes(overclaim, "P0"))

    def test_missing_status_sentence(self):
        result = self.check(self.LEDGER, "问题1：已退还款项。\n问题2：正在推进。\n问题3：需长期坚持。\n")
        self.assertIn("RC-104", self.codes(result, "P2"))

    def test_omitted_items_level(self):
        # 已完成条目没提到只提示（进展报告可不写）；未完成条目没提到为 P1
        result = self.check(self.LEDGER, "问题3：需长期坚持。\n")
        levels = {i["item"]: i["level"] for i in result["issues"] if i["id"] == "RC-105"}
        self.assertEqual(levels, {"1": "P2", "2": "P1"})

    def test_no_ids_in_report(self):
        result = self.check(self.LEDGER, "我单位高度重视整改工作。\n")
        self.assertIn("RC-107", self.codes(result))

    def test_draft_note_ignored(self):
        report = "问题1：已完成整改。\n问题2：正在推进。\n问题3：需长期坚持。\n【起草说明】\n问题2已完成整改。\n"
        result = self.check(self.LEDGER, report)
        self.assertEqual(self.codes(result, "P0"), [])

    def test_sub_items_aggregate(self):
        ledger = LEDGER_HEAD + row("3-1", "已完成") + row("3-2", "持续推进", "")
        result = self.check(ledger, "问题3：已完成整改。\n")
        self.assertIn("RC-101", self.codes(result, "P0"))
        ledger_done = LEDGER_HEAD + row("3-1", "已完成") + row("3-2", "立行立改")
        result2 = self.check(ledger_done, "问题3：已完成整改。\n")
        self.assertEqual(self.codes(result2, "P0"), [])
        self.assertNotIn("RC-105", self.codes(result2))

    def test_alpha_ids(self):
        ledger = LEDGER_HEAD + row("XC-01", "持续推进", "") + row("SJ-02", "已完成")
        result = self.check(ledger, "XC-01：已完成整改。\nSJ-02：已完成整改。\n")
        self.assertEqual([i["item"] for i in result["issues"] if i["id"] == "RC-101"], ["XC-1"])


class IdFormatTest(RectifyCase):
    """带前缀的编号、同前缀的独立编号、列举合写的编号：不能漏报 P0。"""

    def p0_items(self, result):
        return sorted(i["item"] for i in result["issues"] if i["id"] == "RC-101" and i["level"] == "P0")

    def test_parent_key(self):
        self.assertEqual(rc.parent_key("3-1"), "3")
        self.assertEqual(rc.parent_key("4.2"), "4")
        self.assertEqual(rc.parent_key("XC-3-1"), "XC-3")
        self.assertIsNone(rc.parent_key("XC-3"))
        self.assertIsNone(rc.parent_key("巡-1"))
        self.assertIsNone(rc.parent_key("12"))

    def test_mixed_source_prefixes(self):
        """ledger.md 推荐的"巡-1""审-1"：报告写成已完成而台账持续推进，报 P0。"""
        ledger = (LEDGER_HEAD + row("巡-1", "持续推进", "") + row("巡-2", "已完成")
                  + row("审-1", "持续推进", "") + row("审-2", "已完成"))
        report = ("问题巡-1：制度修订稿已提交审议，该问题已完成整改。\n"
                  "巡-2：已完成整改。\n"
                  "问题审-1：该问题已完成整改。\n"
                  "审-2：已完成整改。\n")
        result = self.check(ledger, report)
        self.assertEqual(self.p0_items(result), ["审-1", "巡-1"])
        self.assertNotIn("RC-107", self.codes(result))
        self.assertNotIn("RC-105", self.codes(result))
        self.assertNotIn("groups", result["ledger"]["summary"])

    def test_self_check_prefix(self):
        """writing.md 推荐的"自-1"：自-1 至自-5，报告漏掉自-5、把自-4 写成已完成。"""
        ledger = (LEDGER_HEAD + row("自-1", "已完成") + row("自-2", "立行立改") + row("自-3", "持续推进", "")
                  + row("自-4", "持续推进", "") + row("自-5", "未启动", ""))
        report = "自-1：已完成整改。\n自-2：已立行立改。\n自-3：正在推进。\n自-4：该问题已整改完毕。\n"
        result = self.check(ledger, report)
        self.assertEqual(self.p0_items(result), ["自-4"])
        omitted = {i["item"]: i["level"] for i in result["issues"] if i["id"] == "RC-105"}
        self.assertEqual(omitted, {"自-5": "P1"})
        self.assertEqual(result["ledger"]["summary"]["total"], 5)

    def test_same_prefix_ids_are_separate_items(self):
        """XC-01 至 XC-03 是三个问题，不是 XC 的子项；只写 XC-01 时另两条报漏报。"""
        ledger = LEDGER_HEAD + row("XC-01", "已完成") + row("XC-02", "持续推进", "") + row("XC-03", "持续推进", "")
        result = self.check(ledger, "XC-01：已完成整改。\n")
        self.assertNotIn("groups", result["ledger"]["summary"])
        omitted = {i["item"]: i["level"] for i in result["issues"] if i["id"] == "RC-105"}
        self.assertEqual(omitted, {"XC-2": "P1", "XC-3": "P1"})
        code, out, _ = self.run_main(str(self.tmp / "台账.md"), str(self.tmp / "报告.txt"))
        self.assertNotIn("未发现问题", out)

    def test_prefixed_sub_items(self):
        ledger = LEDGER_HEAD + row("XC-1-1", "已完成") + row("XC-1-2", "持续推进", "") + row("XC-2", "已完成")
        result = self.check(ledger, "问题XC-1：已完成整改。\nXC-2：已完成整改。\n")
        self.assertEqual(self.p0_items(result), ["XC-1"])
        self.assertEqual(result["ledger"]["summary"]["groups"]["total"], 2)

    def test_enumerated_ids_share_status(self):
        """"问题1、4、5已完成整改"：4、5 台账为持续推进，报 P0，不能降为漏报。"""
        ledger = (LEDGER_HEAD + row(1, "已完成") + row(2, "持续推进", "") + row(3, "持续推进", "")
                  + row(4, "持续推进", "") + row(5, "持续推进", ""))
        result = self.check(ledger, "问题1、4、5已完成整改，问题2、3正在推进。\n")
        self.assertEqual(self.p0_items(result), ["4", "5"])
        self.assertNotIn("RC-105", self.codes(result))
        self.assertEqual(result["report"]["mentioned"], ["1", "4", "5", "2", "3"])

    def test_enumeration_variants(self):
        ledger = LEDGER_HEAD + "".join(row(n, "持续推进", "") for n in range(1, 6))
        cases = {
            "问题1和2已完成整改；问题3、问题4、问题5正在推进。\n": ["1", "2"],
            "问题2至4已完成整改；问题1和5正在推进。\n": ["2", "3", "4"],
            "问题1至问题3已完成整改，问题4、问题5正在推进。\n": ["1", "2", "3"],
            "第1、4、5项问题已完成整改，第2项问题和第3项问题正在推进。\n": ["1", "4", "5"],
            "台账第2、3项已完成整改；问题1、4、5正在推进。\n": ["2", "3"],
        }
        for report, expected in cases.items():
            with self.subTest(report=report):
                result = self.check(ledger, report)
                self.assertEqual(self.p0_items(result), expected, json.dumps(result["issues"], ensure_ascii=False))
                self.assertNotIn("RC-105", self.codes(result))
                self.assertNotIn("RC-201", self.codes(result))  # "5项问题"是编号，不是统计数

    def test_quantities_after_list_word_are_not_ids(self):
        ledger = LEDGER_HEAD + "".join(row(n, "持续推进", "") for n in range(1, 4))
        report = ("问题1和3个科室的盘点正在推进。\n问题2与12月底前要完成的系统部署正在推进。\n"
                  "问题3及2026年的预算调整正在推进。\n")
        result = self.check(ledger, report)
        self.assertEqual(result["report"]["mentioned"], ["1", "2", "3"])
        self.assertNotIn("RC-106", self.codes(result))


class PlanAndPeriodTest(RectifyCase):
    """方案中的计划句、"届时"、本期完成数：不误报。"""

    def test_plan_clause_without_yi_is_p2(self):
        ledger = LEDGER_HEAD + row(1, "持续推进", "")
        sentence = "问题1：城建办牵头，2026年12月31日前再完工5部，完成整改。\n"
        plan = self.check(ledger, sentence, plan=True)
        self.assertEqual(self.codes(plan, "P0"), [])
        self.assertIn("RC-101", self.codes(plan, "P2"))
        # 方案里写了"已"的仍是 P0
        past = self.check(ledger, "问题1：该问题已完成整改。\n", plan=True)
        self.assertIn("RC-101", self.codes(past, "P0"))

    def test_jieshi_is_future(self):
        ledger = LEDGER_HEAD + row(1, "持续推进", "")
        for sentence in ("问题1：系统正在部署，预计11月30日上线，届时该问题即可完成整改。\n",
                         "问题1：补签手续后方可完成整改，目前正在推进。\n"):
            with self.subTest(sentence=sentence):
                self.assertNotIn("RC-101", self.codes(self.check(ledger, sentence)))

    def test_period_count_not_compared_with_cumulative(self):
        ledger = LEDGER_HEAD + row(1, "已完成") + row(2, "已完成") + row(3, "持续推进", "")
        items = "问题1：已完成整改。\n问题2：已完成整改。\n问题3：正在推进。\n"
        for lead in ("本月新完成整改1个问题。", "本期完成整改1个。"):
            with self.subTest(lead=lead):
                self.assertNotIn("RC-201", self.codes(self.check(ledger, lead + "\n" + items)))
        for lead in ("截至本月，已完成整改3个。", "截至本月已完成整改3个。", "本月新完成整改5个问题。"):
            with self.subTest(lead=lead):
                self.assertIn("RC-201", self.codes(self.check(ledger, lead + "\n" + items), "P0"))


class ClaimPhrasingTest(RectifyCase):
    """常见的"已完成"说法不能漏报，名词性用法不能误报。台账：已完成 2、持续推进 3、长期坚持 1。"""

    LEDGER = (LEDGER_HEAD + row(1, "已完成") + row(2, "已完成") + row(3, "持续推进", "")
              + row(4, "持续推进", "") + row(5, "持续推进", "") + row(6, "长期坚持"))
    REST = "问题4：正在推进。\n问题5：正在推进。\n问题6：需长期坚持。\n"

    def report(self, item3: str, lead: str = "") -> dict:
        text = lead + "\n问题1：已完成整改。\n问题2：已完成整改。\n问题3：" + item3 + "\n" + self.REST
        return self.check(self.LEDGER, text)

    def test_item_claims_are_p0(self):
        for phrase in ("该问题已完成。", "该问题已整改。", "该问题已整改完毕。", "整改完毕。",
                       "整改措施已落实到位。", "该问题已解决。", "已完成。", "本项已整改，凭证已归档。"):
            with self.subTest(phrase=phrase):
                items = [i["item"] for i in self.report(phrase)["issues"]
                         if i["id"] == "RC-101" and i["level"] == "P0"]
                self.assertEqual(items, ["3"])

    def test_statistic_claims_are_p0(self):
        for sentence in ("反馈的6个问题，已有5个完成整改。", "已完成整改的有5个。", "共整改完成5个问题。",
                         "截至目前，整改完成5个。", "6个问题已整改5个。", "已有5个问题整改到位。"):
            with self.subTest(sentence=sentence):
                result = self.report("正在推进。", sentence)
                self.assertIn("RC-201", self.codes(result, "P0"))

    def test_all_done_claims_are_p0(self):
        for sentence in ("6个问题均已完成整改。", "反馈的6个问题，均已整改到位。", "所有问题已整改完毕。",
                         "全部问题均已完成整改。", "所有问题都已整改。"):
            with self.subTest(sentence=sentence):
                self.assertIn("RC-203", self.codes(self.report("正在推进。", sentence), "P0"))

    def test_noun_and_partial_usages_are_not_p0(self):
        for phrase in ("已完成整改方案制定，制度修订正在推进。", "已完成整改措施第1项，第2项正在推进。",
                       "已完成整改台账建立。该问题正在推进。", "已明确整改到位的标准，该问题正在推进。",
                       "已整改到位的有2个社区，其余正在推进。", "涉及的3个科室均已完成整改，该问题正在推进。",
                       "19处隐患已整改完成12处，该问题正在推进。", "盘点已完成，系统部署正在推进。",
                       "部分科室已完成整改，该问题正在推进。"):
            with self.subTest(phrase=phrase):
                result = self.report(phrase)
                self.assertEqual(self.codes(result, "P0"), [], json.dumps(result["issues"], ensure_ascii=False))

    def test_subset_each_claim_is_not_all_done(self):
        result = self.report("正在推进。", "检查期间立行立改的2个问题均已整改到位。")
        self.assertNotIn("RC-203", self.codes(result))

    def test_mixed_status_is_p1(self):
        result = self.report("已完成整改。该问题正在推进，预计12月31日前完成。")
        levels = [i["level"] for i in result["issues"] if i["id"] == "RC-101"]
        self.assertEqual(levels, ["P1"])
        # 明确写了"整改状态：已完成"的仍是 P0
        result2 = self.report("整改状态：已完成。该问题正在推进。")
        self.assertIn("RC-101", self.codes(result2, "P0"))

    def test_plan_goal_sentence(self):
        ledger = LEDGER_HEAD + row(1, "未启动", "") + row(2, "持续推进", "")
        goal = "一、整改目标\n到2026年12月31日，巡察反馈问题全部整改到位，长效机制基本建立。\n问题1：修订制度。\n问题2：补正案卷。\n"
        path = self.write("台账.md", ledger)
        for plan in (True, False):
            with self.subTest(plan=plan):
                result = rc.build_report(path, self.write("方案.txt", goal), plan=plan)
                self.assertNotIn("RC-203", self.codes(result))
        no_date = "巡察反馈问题全部整改到位。\n问题1：修订制度。\n问题2：补正案卷。\n"
        plan_result = rc.build_report(path, self.write("方案2.txt", no_date), plan=True)
        self.assertIn("RC-203", self.codes(plan_result, "P2"))
        self.assertNotIn("RC-203", self.codes(plan_result, "P0"))
        report_result = rc.build_report(path, self.write("报告.txt", no_date))
        self.assertIn("RC-203", self.codes(report_result, "P0"))
        past = rc.build_report(path, self.write("方案3.txt", "巡察反馈问题已全部整改到位。\n问题1：修订制度。\n"), plan=True)
        self.assertIn("RC-203", self.codes(past, "P0"))


class StatisticsTest(RectifyCase):

    LEDGER = (LEDGER_HEAD + row(1, "已完成") + row(2, "立行立改") + row(3, "持续推进", "")
              + row(4, "长期坚持"))

    def stat_issues(self, sentence):
        report = sentence + "\n问题1：已完成整改。\n问题2：已立行立改。\n问题3：正在推进。\n问题4：需长期坚持。\n"
        return self.check(self.LEDGER, report)

    def test_correct_sentence(self):
        result = self.stat_issues("反馈的4个问题，已完成整改2个（其中立行立改1个），持续推进1个，需长期坚持1个，整改完成率50.0%。")
        self.assertEqual(result["issues"], [])

    def test_chinese_numerals(self):
        result = self.stat_issues("反馈的四个问题，已完成整改两个，持续推进一个，需长期坚持一个，整改完成率50%。")
        self.assertEqual(result["issues"], [])

    def test_overclaim_count_and_rate(self):
        result = self.stat_issues("反馈的4个问题，已完成整改3个，持续推进1个，整改完成率75%。")
        p0 = self.codes(result, "P0")
        self.assertIn("RC-201", p0)
        self.assertIn("RC-202", p0)

    def test_undercount_is_p1(self):
        result = self.stat_issues("反馈的5个问题，已完成整改1个，整改完成率25%。")
        self.assertIn("RC-201", self.codes(result, "P1"))
        self.assertIn("RC-202", self.codes(result, "P1"))
        self.assertNotIn("RC-201", self.codes(result, "P0"))

    def test_total_before_done_count(self):
        # "4个问题已完成整改2个"：4 是总数，不能当完成数
        result = self.stat_issues("截至目前，4个问题已完成整改2个。")
        self.assertEqual(self.codes(result), [])

    def test_all_done_claim(self):
        result = self.stat_issues("整改工作已全部完成。")
        self.assertIn("RC-203", self.codes(result, "P0"))
        result2 = self.stat_issues("确保年底前全部整改到位。")
        self.assertNotIn("RC-203", self.codes(result2))

    def test_basic_completion(self):
        result = self.stat_issues("整改工作基本完成。")
        self.assertIn("RC-204", self.codes(result, "P1"))

    def test_subunit_numbers_not_counted(self):
        result = self.stat_issues("已完成3个科室的盘点，正在推进5个病区的系统部署。")
        self.assertNotIn("RC-201", self.codes(result))

    def test_rate_rounding(self):
        ledger = LEDGER_HEAD + row(1, "已完成") + row(2, "持续推进", "") + row(3, "持续推进", "")
        report = "问题1：已完成整改。\n问题2：正在推进。\n问题3：正在推进。\n整改完成率33.3%。\n"
        self.assertEqual(self.check(ledger, report)["issues"], [])
        bad = report.replace("33.3%", "33.5%")
        self.assertIn("RC-202", self.codes(self.check(ledger, bad), "P0"))


class LedgerSelfCheckTest(RectifyCase):

    def test_duplicate_empty_nonstandard_evidence(self):
        ledger = (LEDGER_HEAD + row(1, "已完成", "") + row(1, "已完成") + row(2, "")
                  + row(3, "基本完成") + row(4, "正在推进", ""))
        result = self.check(ledger)
        codes = self.codes(result)
        self.assertIn("RC-001", codes)  # 编号重复
        self.assertIn("RC-002", codes)  # 状态为空
        self.assertIn("RC-003", self.codes(result, "P1"))  # 基本完成
        self.assertIn("RC-003", self.codes(result, "P2"))  # 正在推进 → 建议统一为持续推进
        self.assertIn("RC-004", codes)  # 已完成没有佐证

    def test_required_field_empty(self):
        ledger = LEDGER_HEAD + "| 1 | 问题 |  | 分管领导 |  | 2026年12月31日 | 标准 | 未启动 |  |\n"
        result = self.check(ledger)
        issue = [i for i in result["issues"] if i["id"] == "RC-005"][0]
        self.assertIn("整改措施", issue["message"])
        self.assertIn("责任部门", issue["message"])

    def test_marker_counts(self):
        ledger = LEDGER_HEAD + "| 1 | 问题 | 【待补：整改措施】 | 分管领导 | 办公室 | 【待补：完成时限】 | 【待核实：目标值】 | 未启动 |  |\n"
        result = self.check(ledger)
        self.assertEqual(result["summary"]["markers"], {"待补": 2, "待核实": 1})
        self.assertNotIn("RC-005", self.codes(result))

    def test_overdue_only_with_date(self):
        ledger = LEDGER_HEAD + row(1, "持续推进", "", deadline="2026年9月30日") + row(2, "已完成", deadline="2026年9月30日")
        self.assertNotIn("RC-006", self.codes(self.check(ledger)))
        result = self.check(ledger, as_of=dt.date(2026, 10, 16))
        overdue = [i["item"] for i in result["issues"] if i["id"] == "RC-006"]
        self.assertEqual(overdue, ["1"])

    def test_missing_evidence_column_note(self):
        ledger = ("| 编号 | 问题表现 | 整改状态 |\n|---|---|---|\n| 1 | 甲 | 已完成 |\n| 2 | 乙 | 已完成 |\n")
        result = self.check(ledger)
        self.assertEqual(self.codes(result, "P2").count("RC-004"), 1)

    def test_continuation_rows(self):
        ledger = LEDGER_HEAD + row(1, "已完成") + "|  |  | 措施二 |  | 财务科 | 2026年12月31日 | 标准 | 持续推进 |  |\n"
        result = self.check(ledger)
        self.assertEqual(result["ledger"]["summary"]["status_count"]["持续推进"], 1)
        self.assertEqual(result["ledger"]["summary"]["total"], 1)

    def test_header_variants(self):
        ledger = "| 序号 | 反馈问题 | 完成情况 |\n|---|---|---|\n| 1 | 甲 | 已完成 |\n"
        result = self.check(ledger)
        self.assertEqual(result["ledger"]["summary"]["done"], 1)

    def test_vague_deadline(self):
        ledger = (LEDGER_HEAD + row(1, "长期坚持", deadline="长期") + row(2, "持续推进", "", deadline="尽快")
                  + row(3, "持续推进", "", deadline="11月30日前")
                  + row(4, "长期坚持", deadline="长期坚持，2026年10月31日前首次研究")
                  + row(5, "已完成", deadline="近期") + row(6, "未启动", "", deadline="【待补：完成时限】"))
        result = self.check(ledger)
        self.assertEqual([i["item"] for i in result["issues"] if i["id"] == "RC-007"], ["1", "2"])
        self.assertIn("RC-007", self.codes(result, "P1"))

    def test_escaped_pipe_in_cell(self):
        ledger = LEDGER_HEAD + row(1, "已完成", problem="A\\|B 两类票据混放")
        result = self.check(ledger)
        self.assertEqual(result["ledger"]["summary"]["total"], 1)
        loaded = rc.load_ledger(self.write("台账2.md", ledger))
        self.assertEqual(loaded.items["1"].fields["problem"], "A|B 两类票据混放")

    def test_short_name_column_not_used_as_problem(self):
        ledger = ("| 编号 | 问题来源 | 问题简称 | 整改状态 |\n|---|---|---|---|\n"
                  "| 1 | 巡察反馈第1条 | 耗材账实不符 | 未启动 |\n")
        feedback = self.write("反馈.txt", "1. 部分科室高值耗材领用登记不全，账实不符。\n")
        result = self.check(ledger, feedback_path=feedback)
        self.assertEqual(self.codes(result, "P1"), [])
        self.assertIn("RC-301", self.codes(result, "P2"))  # 没有问题表现列，提示增加

    def test_feedback_weakening(self):
        ledger = (LEDGER_HEAD + row(1, "已完成", problem="抽查的30宗执法案卷中有9宗现场检查记录缺少当事人签字")
                  + row(2, "已完成", problem="执法规范化水平有待提升"))
        feedback = self.write("反馈.txt", "1. 行政执法不规范，抽查的30宗执法案卷中有9宗现场检查记录缺少当事人签字。\n"
                                          "2. 违规发放津贴补贴共计12.4万元。\n")
        result = self.check(ledger, feedback_path=feedback)
        items = [i["item"] for i in result["issues"] if i["id"] == "RC-301"]
        self.assertEqual(items, ["2"])
        self.assertNotIn("RC-302", self.codes(result))


class FeedbackDeletionTest(RectifyCase):
    """问题表现删数字、截短：RC-302。"""

    def feedback_issues(self, problems, feedback_text):
        ledger = LEDGER_HEAD + "".join(row(no, "未启动", "", problem=p) for no, p in problems)
        result = self.check(ledger, feedback_path=self.write("反馈.txt", feedback_text))
        return result, {i["item"]: i for i in result["issues"] if i["id"] == "RC-302"}

    def test_deleted_numbers(self):
        _, lost = self.feedback_issues(
            [(1, "部分科室高值耗材领用登记不全，账实不符")],
            "部分科室高值耗材领用登记不全，账实不符，2025年盘亏金额12.6万元。\n")
        self.assertIn("2025年盘亏金额12.6万元", lost["1"]["message"])
        self.assertEqual(lost["1"]["level"], "P1")

    def test_truncated_to_first_clause(self):
        result, lost = self.feedback_issues(
            [(1, "工程项目管理不严"), (2, "违规发放津贴补贴共计12.4万元")],
            "1. 工程项目管理不严，2025年有1个项目未经招标直接委托施工，合同金额48.6万元。\n"
            "2. 违规发放津贴补贴共计12.4万元。\n")
        self.assertEqual(set(lost), {"1"})
        self.assertIn("未经招标直接委托施工", lost["1"]["message"])
        self.assertIn("等 2 处", lost["1"]["message"])
        self.assertNotIn("RC-301", self.codes(result))  # 截短不是改写，原文里找得到

    def test_deleted_quote_and_qualifier(self):
        _, lost = self.feedback_issues(
            [(3, "2025年有3批次办公设备采购拆分合同，合计金额96万元"), (4, "科室耗材领用登记不全")],
            "3. 政府采购\"化整为零\"规避公开招标，2025年有3批次办公设备采购拆分合同，合计金额96万元。\n"
            "4. 部分科室耗材领用登记不全。\n")
        self.assertIn("化整为零", lost["3"]["message"])
        self.assertIn("部分科室耗材领用登记不全", lost["4"]["message"])

    def test_verbatim_and_split_sub_items_pass(self):
        ledger = (LEDGER_HEAD
                  + row("4-1", "未启动", "", problem="预算执行不严格，2025年有3个项目未履行调整审批程序")
                  + row("4-2", "未启动", "", problem="部分科室超预算支出共计8.2万元")
                  + row(5, "未启动", "", problem="违规发放津贴补贴，以“加班补助”名义发放补贴12.4万元"))
        feedback = self.write("反馈.txt",
                              "（四）预算执行不严格，2025年有3个项目未履行调整审批程序；部分科室超预算支出共计8.2万元。\n"
                              "（五）违规发放津贴补贴，以\"加班补助\"名义发放补贴12.4万元。\n"
                              "请于2个月内报送整改方案。\n")
        result = self.check(ledger, feedback_path=feedback)
        self.assertEqual([i for i in result["issues"] if i["id"] in ("RC-301", "RC-302")], [])

    def test_split_feedback(self):
        entries = rc.split_feedback("市委巡察组反馈意见\n1. 甲问题，涉及\n3个项目。\n二、乙问题。\n（3）丙问题。丁情况。\n")
        self.assertEqual(entries, [(None, "市委巡察组反馈意见"), (1, "甲问题，涉及3个项目。"),
                                   (2, "乙问题。"), (3, "丙问题。丁情况。")])


@unittest.skipUnless(HAS_OPENPYXL, "需要 openpyxl")
class XlsxTest(RectifyCase):

    def test_xlsx_with_title_row_and_merged_cells(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["某公司审计整改台账"])
        ws.append(["序号", "问题", "整改措施", "责任部门", "完成时限", "整改状态", "佐证材料"])
        ws.append([1, "问题甲", "措施一", "财务部", dt.datetime(2026, 11, 30), "已完成", "凭证"])
        ws.append([None, None, "措施二", "审计部", "2026-12-31", "持续推进", None])
        ws.append([2.0, "问题乙", "措施", "办公室", "2026-10-01", "已完成", "文件"])
        ws.merge_cells("A3:A4")
        path = self.tmp / "台账.xlsx"
        wb.save(path)
        report = self.write("报告.txt", "问题1：已完成整改。\n问题2：已完成整改。\n")
        result = rc.build_report(path, report)
        self.assertEqual(result["ledger"]["summary"]["total"], 2)
        self.assertEqual([i["item"] for i in result["issues"] if i["id"] == "RC-101"], ["1"])

    def test_md2xlsx_roundtrip(self):
        md2xlsx = _load("rectify_tests_md2xlsx_copy", SCRIPTS_DIR / "md2xlsx.py")  # 同步到 skill 的副本
        src = FIXTURES / "台账-标准.md"
        target, _, _ = md2xlsx.convert(src, self.tmp / "台账.xlsx")
        result = rc.build_report(target, FIXTURES / "报告-必过.txt")
        self.assertEqual(result["issues"], [])


@unittest.skipUnless(HAS_DOCX, "需要 python-docx")
class DocxTest(RectifyCase):

    def test_docx_report(self):
        document = docx.Document()
        document.add_paragraph("问题1：已完成整改。")
        document.add_paragraph("问题2：已完成整改。")
        document.add_paragraph("【起草说明】")
        document.add_paragraph("问题3：已完成整改。")
        path = self.tmp / "报告.docx"
        document.save(path)
        ledger = self.write("台账.md", LEDGER_HEAD + row(1, "已完成") + row(2, "持续推进", "") + row(3, "持续推进", ""))
        result = rc.build_report(ledger, path)
        self.assertEqual([i["item"] for i in result["issues"] if i["id"] == "RC-101"], ["2"])


class CliTest(RectifyCase):

    def test_json_and_strict(self):
        code, out, _ = self.run_main(str(FIXTURES / "台账-标准.md"), str(FIXTURES / "报告-必错.txt"),
                                     "--json", "--strict")
        self.assertEqual(code, 1)
        data = json.loads(out)
        for key in ("ledger", "report", "feedback", "as_of", "summary", "issues"):
            self.assertIn(key, data)
        self.assertGreater(data["summary"]["P0"], 0)
        first = data["issues"][0]
        for key in ("id", "level", "where", "category", "item", "line", "excerpt", "message", "suggestion"):
            self.assertIn(key, first)

    def test_strict_pass(self):
        code, out, _ = self.run_main(str(FIXTURES / "台账-标准.md"), str(FIXTURES / "报告-必过.txt"), "--strict")
        self.assertEqual(code, 0)
        self.assertIn("未发现问题", out)

    def test_ledger_only(self):
        code, out, _ = self.run_main(str(FIXTURES / "台账-标准.md"))
        self.assertEqual(code, 0)
        self.assertIn("完成率 40.0%", out)

    def test_input_errors(self):
        no_status = self.write("无状态.md", "| 编号 | 问题 |\n|---|---|\n| 1 | 甲 |\n")
        no_table = self.write("无表.md", "没有表格\n")
        cases = [
            [str(self.tmp / "不存在.md")],
            [str(no_status)],
            [str(no_table)],
            [str(self.write("台账.csv", "a,b"))],
            [str(FIXTURES / "台账-标准.md"), str(self.tmp / "不存在.txt")],
            [str(FIXTURES / "台账-标准.md"), "--date", "2026/10/16"],
        ]
        for args in cases:
            with self.subTest(args=args):
                code, _, err = self.run_main(*args)
                self.assertEqual(code, 2)
                self.assertIn("错误", err)

    def test_help(self):
        code, out, _ = self.run_main("--help")
        self.assertEqual(code, 0)
        self.assertIn("用法", out)


if __name__ == "__main__":
    unittest.main()
