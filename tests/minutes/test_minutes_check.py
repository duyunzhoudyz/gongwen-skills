# -*- coding: utf-8 -*-
"""skills/gongwen-minutes/scripts/minutes_check.py 的单元测试。

运行：python -m unittest discover -s tests/minutes -t .

每条检查规则都有"必错标本"（必须报出）和"必过标本"（不得误报），防止检查器空跑或误伤。
"""
from __future__ import annotations

import importlib.util
import io
import json
import shutil
import sys
import tempfile
import unittest
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[2]
SCRIPT_PATH = REPO_ROOT / "skills" / "gongwen-minutes" / "scripts" / "minutes_check.py"


def _load():
    name = "minutes_check_under_test"
    if name in sys.modules:
        return sys.modules[name]
    spec = importlib.util.spec_from_file_location(name, SCRIPT_PATH)
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)  # type: ignore[union-attr]
    return module


mc = _load()

try:
    import docx  # type: ignore  # noqa: F401
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

try:
    import openpyxl  # type: ignore  # noqa: F401
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


# 必过标本：一份写法完整的纪要（虚构）
GOOD_MINUTES = """某医院院长办公会议纪要

2026年10月14日下午，院长张某在行政楼三楼会议室主持召开院长办公会，研究CT设备采购等事项。现将会议议定事项纪要如下：

一、关于CT设备采购

会议听取了设备处关于CT设备采购方案的汇报。

会议议定：

（一）原则同意按方案二采购CT设备，预算320万元。由设备处会同放射科核对设备配置清单后，于11月15日前按程序提交党委会研究。

二、关于门诊预约挂号改革

会议听取了门诊部的汇报。会议认为，改革方案总体可行。

会议议定：

（一）同意《门诊预约挂号改革方案》，留足老年患者现场号源。由门诊部牵头、信息科配合，11月底前上线。

（二）预约爽约限制措施暂缓实施，待分时段预约运行三个月后再提交研究。

三、关于调整夜班补贴标准

会议议定：

（一）调整夜班补贴标准事项暂缓，待年底职代会召开后再研究。

（二）由财务科测算调整夜班补贴所需资金，于【待补：完成时限】前完成。

会议要求，各科室要按“我觉得能办就马上办”的态度抓好落实。

出席：张某、李某、王某

列席：赵某、陈某

【起草说明】
一、所作假设
1. 会议同意由某科负责，我觉得可以。（起草说明不参与检查）
"""

GOOD_LIST = """| 序号 | 议定事项 | 纪要位置 | 责任单位 | 配合单位 | 完成时限 | 办理状态 | 备注 |
|---|---|---|---|---|---|---|---|
| 1 | 核对CT设备配置清单后提交党委会研究 | 一、（一） | 设备处 | 放射科 | 11月15日前 | 待办理 | |
| 2 | 门诊预约挂号改革方案上线，留足老年患者现场号源 | 二、（一） | 门诊部 | 信息科 | 2026年11月底前 | 待办理 | |
| 3 | 测算调整夜班补贴所需资金 | 三、（二） | 财务科 | | 【待补：完成时限】 | 待办理 | |
"""


# 必过标本：按评测材料写成的纪要和督办清单（虚构），覆盖评测材料中的真实时限写法
EVAL1_MINUTES = """某医院院长办公会议纪要

2026年10月14日下午，院长张卫东在行政楼三楼会议室主持召开第18次院长办公会，研究CT设备采购等4项议题。现将会议议定事项纪要如下：

一、关于CT设备采购

会议听取了设备处关于CT设备采购方案的汇报。

会议议定：

（一）原则同意按方案二采购CT设备，所需经费从2026年设备购置预算中列支。由设备处会同放射科核对设备配置清单，明确高级后处理功能的需求和费用【待补：完成时限】，按程序提交党委会研究。

二、关于门诊预约挂号改革

会议议定：

（一）同意《门诊预约挂号改革方案》，留足老年患者现场号源；预约爽约限制措施暂缓实施【待补：暂缓原因或再议条件】。由门诊部牵头、信息科配合，11月底前上线。

三、关于调整夜班补贴标准

会议议定：

（一）调整夜班补贴标准事项暂缓，待年底职代会召开后再研究。由财务科先行测算调整所需资金【待补：完成时限】。

出席：张卫东、李建华、王晓红、刘志强、马丽

请假：黄斌

列席：赵刚、陈静、孙立、周敏、吴海、郑强、钱峰
"""

EVAL1_LIST = """| 序号 | 议定事项 | 纪要位置 | 责任单位 | 配合单位 | 完成时限 | 办理状态 | 备注 |
|---|---|---|---|---|---|---|---|
| 1 | 核对CT设备配置清单，按程序提交党委会研究 | 一、（一） | 设备处 | 放射科 | 【待补：完成时限】 | 待办理 | |
| 2 | 门诊预约挂号改革方案上线，留足老年患者现场号源 | 二、（一） | 门诊部 | 信息科 | 11月底前 | 待办理 | |
| 3 | 测算调整夜班补贴所需资金 | 三、（一） | 财务科 | | 【待补：完成时限】 | 待办理 | |
"""

EVAL3_MINUTES = """某大学校长办公会议纪要

2026年10月15日上午，校长孙启明在行政楼二楼第一会议室主持召开校长办公会，研究新专业申报等3项议题。现将会议议定事项纪要如下：

一、关于智能制造工程本科专业申报

会议议定：

（一）原则同意机械学院申报智能制造工程本科专业。由教务处会同机械学院补充师资补充计划，尽快提交学术委员会审议【待补：具体完成时限】，学术委员会审议通过后按程序报党委常委会。

二、关于东区6栋学生宿舍安装空调

会议议定：

（一）同意采用租赁方式为东区6栋学生宿舍安装空调，由后勤处负责，寒假前完成安装。收费标准由后勤处先拿出方案，听取学生代表意见，按规定备案后再执行。

三、关于实验室安全隐患整改

会议议定：

（一）剩余8处隐患由资产与实验室管理处负责督促整改，月底前全部整改完成。

会议强调，各学院院长是实验室安全第一责任人。

出席：孙启明、周文、郑海燕、冯立新

列席：杨帆、蒋伟、韩雪、沈红
"""

EVAL3_LIST = """| 序号 | 议定事项 | 纪要位置 | 责任单位 | 配合单位 | 完成时限 | 办理状态 | 备注 |
|---|---|---|---|---|---|---|---|
| 1 | 补充师资补充计划，提交学术委员会审议 | 一、（一） | 教务处 | 机械学院 | 【待补：具体完成时限】 | 待办理 | |
| 2 | 采用租赁方式为东区6栋学生宿舍安装空调 | 二、（一） | 后勤处 | | 寒假前 | 待办理 | |
| 3 | 督促剩余8处实验室安全隐患整改 | 三、（一） | 资产与实验室管理处 | | 月底前 | 待办理 | |
"""

EVAL8_MINUTES = """流感疫苗接种工作协调会议纪要

10月16日上午，区卫生健康委员会分管副主任【待补：姓名】在区卫健委五楼会议室主持召开流感疫苗接种工作协调会。现将会议议定事项纪要如下：

会议议定：

一、学校集中接种由区疾控中心牵头、区教育局配合，知情同意书11月10日前收齐，11月下旬开始接种。

二、60岁以上老年人接种由各社区卫生服务中心负责，结合老年人体检一并开展，每周向区疾控中心报送一次接种进度。

三、社区卫生服务中心周末加开接种门诊事项另行研究，提交下次会议研究。

出席：【待补：出席人员名单】
"""

EVAL8_LIST = """| 序号 | 议定事项 | 纪要位置 | 责任单位 | 配合单位 | 完成时限 | 办理状态 | 备注 |
|---|---|---|---|---|---|---|---|
| 1 | 收齐学校集中接种知情同意书 | 一、 | 区疾控中心 | 区教育局 | 11月10日前 | 待办理 | |
| 2 | 开始学校集中接种 | 一、 | 区疾控中心 | 区教育局 | 11月下旬 | 待办理 | |
| 3 | 老年人接种，每周报送接种进度 | 二、 | 各社区卫生服务中心 | | 每周 | 待办理 | |
"""

EVAL11_MINUTES = """某集团总经理办公会议纪要

10月14日下午，总经理吴志远主持召开总经理办公会，研究物流园二期项目等事项。现将会议议定事项纪要如下：

一、关于物流园二期项目

经党委会前置研究，会议同意建设物流园二期项目，总投资1.2亿元，由投资发展部牵头，今年12月开工。财务部按年度安排资金。

二、关于三季度经营情况

会议听取了财务部关于三季度经营情况的汇报。

会议要求，各子公司确保完成全年目标，具体措施各子公司10月底前报财务部。

出席：吴志远、王建、刘洋、郭明
"""

EVAL11_LIST = """| 序号 | 议定事项 | 纪要位置 | 责任单位 | 配合单位 | 完成时限 | 办理状态 | 备注 |
|---|---|---|---|---|---|---|---|
| 1 | 建设物流园二期项目 | 一、 | 投资发展部 | 财务部 | 今年12月开工 | 待办理 | |
| 2 | 报送确保完成全年目标的具体措施 | 二、会议要求第1条 | 各子公司 | | 10月底前 | 待办理 | |
"""


def items_of(text: str):
    return mc.extract_items(mc.body_lines(text.split("\n")))


def issue_ids(text: str):
    items = items_of(text)
    body = mc.body_lines(text.split("\n"))
    return [i.id for i in mc.check_items(items) + mc.check_body(body, items)]


def list_issue_ids(minutes: str, table_md: str):
    items = items_of(minutes)
    lines = table_md.strip().split("\n")
    header = mc._split_md_row(lines[0])
    rows = [mc._split_md_row(line) for line in lines[2:]]
    return [i.id for i in mc.match_list(items, mc._rows_to_table(header, rows))]


def wrap(*paragraphs: str) -> str:
    """把若干段落包成带出席名单的最小纪要。"""
    return "\n\n".join(("某局局务会议纪要",) + paragraphs + ("出席：张某、李某",)) + "\n"


class ExtractTest(unittest.TestCase):

    def test_good_minutes_items(self):
        items = items_of(GOOD_MINUTES)
        decided = [it for it in items if it.kind == "议定"]
        self.assertEqual(len(decided), 5)
        self.assertEqual([it.position for it in decided], ["一、（一）", "二、（一）", "二、（二）", "三、（一）", "三、（二）"])
        self.assertEqual([it.decision for it in decided], ["原则同意", "同意", "暂缓", "暂缓", "决定"])

    def test_requirement_kind(self):
        items = items_of(GOOD_MINUTES)
        reqs = [it for it in items if it.kind == "要求"]
        self.assertEqual(len(reqs), 1)
        self.assertEqual(reqs[0].decision, "会议要求")
        self.assertIn("会议要求第1条", reqs[0].position)

    def test_heard_and_opinion_are_not_items(self):
        items = items_of(wrap("会议听取了财务科的汇报。", "会议认为，方案总体可行。", "会议强调，要严格执行采购程序。"))
        self.assertEqual(items, [])

    def test_inline_yishi_split(self):
        items = items_of(wrap("会议议定：一是同意由办公室牵头修订值班制度，11月底前完成；二是由财务科负责测算经费，12月10日前完成。"))
        self.assertEqual(len(items), 2)
        self.assertEqual(items[0].owner, "办公室")
        self.assertEqual(items[1].owner, "财务科")
        self.assertEqual(items[1].deadline, "12月10日前")

    def test_single_topic_numbered_items(self):
        text = wrap("会议议定以下事项：", "一、同意由综合科牵头修订会议制度，12月底前完成。", "二、暂缓调整办公用房，待新楼验收后再研究。")
        items = items_of(text)
        self.assertEqual(len(items), 2)
        self.assertEqual(items[0].position, "一、")
        self.assertEqual(items[1].decision, "暂缓")

    def test_single_topic_heading_like_item_without_period(self):
        items = items_of(wrap("会议议定如下：", "一、同意由综合科牵头修订会议制度", "二、暂缓调整办公用房"))
        self.assertEqual([it.decision for it in items], ["同意", "暂缓"])

    def test_sentence_mode(self):
        items = items_of(wrap("经研究，会议同意由院办公室牵头制定会议管理办法，于12月31日前印发。"))
        self.assertEqual(len(items), 1)
        self.assertEqual(items[0].owner, "院办公室")
        self.assertEqual(items[0].deadline, "12月31日前")

    def test_decision_after_process_cue(self):
        items = items_of(wrap("会议听取了设备处关于CT采购的汇报，原则同意采购方案，由设备处按会议意见修改后，11月底前提交党委会。",
                              "（一）关于值班制度修订事项。会议同意由办公室牵头修订，12月底前完成。"))
        self.assertEqual([it.decision for it in items], ["原则同意", "同意"])
        self.assertTrue(items[0].text.startswith("原则同意"))
        self.assertEqual(items[1].owner, "办公室")

    def test_process_cue_without_decision_is_not_item(self):
        self.assertEqual(items_of(wrap("会议听取了财务科关于暂缓支付的汇报。")), [])

    def test_decision_heading_style(self):
        text = wrap("一、会议概况", "会议听取了各科室的汇报。", "二、议定事项",
                    "（一）由办公室牵头修订值班制度，11月底前完成。", "（二）由财务科测算经费，12月10日前完成。",
                    "三、会议要求", "（一）由各科室于每月5日前报送进展。")
        items = items_of(text)
        self.assertEqual([it.kind for it in items], ["议定", "议定", "要求"])
        self.assertEqual(items[0].position, "二、（一）")
        self.assertEqual(items[2].owner, "各科室")

    def test_review_and_pass_sentence(self):
        items = items_of(wrap("会议审议并原则通过了《会议管理办法》，由办公室按会议意见修改后于12月底前印发。"))
        self.assertEqual(len(items), 1)
        self.assertEqual(items[0].decision, "原则同意")

    def test_attendance_ends_block(self):
        text = "某局局务会议纪要\n\n会议议定：\n\n（一）同意由办公室牵头，11月底前完成。\n\n出席：张某\n\n列席：李某\n"
        self.assertEqual(len(items_of(text)), 1)

    def test_draft_note_ignored(self):
        text = wrap("会议同意由办公室牵头，11月底前完成。") + "\n【起草说明】\n会议同意由某处负责。\n"
        self.assertEqual(len(items_of(text)), 1)

    def test_classify(self):
        cases = {
            "会议原则同意该方案": "原则同意",
            "会议原则上通过该方案": "原则同意",
            "会议不同意该申请": "不同意",
            "会议决定暂缓实施": "暂缓",
            "该事项另行研究": "另行研究",
            "提交下次会议研究": "另行研究",
            "方案退回财务科修改后重新提交": "退回修改",
            "会议同意该方案": "同意",
            "明确由办公室牵头": "决定",
            "知情同意书11月10日前收齐": "决定",
            "爽约限制暂不实施": "暂缓",
            "会议同意《关于暂缓收费的报告》": "同意",
            "会议同意暂不调整收费标准": "同意",
        }
        for text, expected in cases.items():
            with self.subTest(text=text):
                self.assertEqual(mc.classify(text), expected)

    def test_owner_patterns(self):
        cases = {
            "由院办公室牵头制定办法": "院办公室",
            "由设备处压缩采购预算": "设备处",
            "门诊部牵头，信息科配合": "门诊部",
            "会议同意门诊部牵头实施": "门诊部",
            "责任单位：财务科，11月底前完成": "财务科",
            "由分管副院长负责协调": "分管副院长",
            "由信息中心负责系统改造": "信息中心",
            "由财务科、审计科负责核查": "财务科、审计科",
        }
        for text, expected in cases.items():
            with self.subTest(text=text):
                self.assertEqual(mc.find_owner(text), expected)

    def test_owner_not_found_in_quotes(self):
        self.assertEqual(mc.find_owner("同意《由办公室牵头的方案》"), "")

    def test_cooperate(self):
        self.assertEqual(mc.find_cooperate("由门诊部牵头、信息科配合"), "信息科")
        self.assertEqual(mc.find_cooperate("由设备处会同放射科核对清单"), "放射科")

    def test_deadline_patterns(self):
        cases = {
            "11月底前完成": "11月底前",
            "于2026年11月30日前完成": "2026年11月30日前",
            "每月报送进展": "每月",
            "3个工作日内完成": "3个工作日内",
            "三个工作日内完成": "三个工作日内",
            "年底前完成": "年底前",
            "下周五前报送": "下周五前",
            "长期坚持": "长期坚持",
            # 评测材料中的真实写法
            "今年12月开工": "今年12月",
            "12月份启动": "12月份",
            "月底前全部整改完": "月底前",
            "周五前把情况报给主任": "周五前",
            "下次会议前完成": "下次会议前",
            "节后第一周完成": "节后第一周",
            "11月下旬开始接种": "11月下旬",
            "明年3月前开工": "明年3月前",
        }
        for text, expected in cases.items():
            with self.subTest(text=text):
                self.assertEqual(mc.find_deadline(text), expected)

    def test_deadline_no_false_hit(self):
        for text in ("岁月中留下的经验", "上个月召开了会议", "由办公室牵头修订制度"):
            with self.subTest(text=text):
                self.assertEqual(mc.find_deadline(text), "")

    def test_owner_extra_patterns(self):
        cases = {
            "具体措施各子公司10月底前报财务部": "各子公司",
            "由社区居委会上门做工作": "社区居委会",
            "社区居委会负责通知居民": "社区居委会",
        }
        for text, expected in cases.items():
            with self.subTest(text=text):
                self.assertEqual(mc.find_owner(text), expected)

    def test_split_yishi_ignores_inner_shi(self):
        self.assertEqual(mc.split_yishi("同意该方案，一是要抓紧"), ["同意该方案，一是要抓紧"])

    def test_lead_in_with_meeting_name(self):
        items = items_of(wrap("经党委会前置研究，会议同意建设物流园二期项目，由投资发展部牵头，今年12月开工。"))
        self.assertEqual([it.decision for it in items], ["同意"])
        self.assertEqual(items[0].deadline, "今年12月")


class MixedDecisionTest(unittest.TestCase):
    """一段中有几类结论：必须拆开，各自判定（必错标本）；单一结论不得误拆（必过标本）。"""

    def decided(self, *paragraphs):
        return [it for it in items_of(wrap("会议议定：", *paragraphs)) if it.kind == "议定"]

    def test_agree_and_postpone_semicolon(self):
        items = self.decided("会议同意门诊预约挂号改革方案，由门诊部牵头、信息科配合，11月底前上线；爽约限制暂缓实施。")
        self.assertEqual([it.decision for it in items], ["同意", "暂缓"])
        self.assertEqual((items[0].owner, items[0].deadline), ("门诊部", "11月底前"))

    def test_agree_and_postpone_comma(self):
        items = self.decided("会议同意门诊预约挂号改革方案，留足老年患者现场号源，爽约限制措施暂缓实施，"
                             "由门诊部牵头、信息科配合，11月底前上线。")
        self.assertEqual([it.decision for it in items], ["同意", "暂缓"])
        self.assertEqual((items[0].owner, items[0].deadline), ("门诊部", "11月底前"))

    def test_owner_after_postpone_sentence_belongs_to_agree(self):
        items = self.decided("（一）同意《门诊预约挂号改革方案》，留足老年患者现场号源；预约爽约限制措施暂缓实施。"
                             "由门诊部牵头、信息科配合，11月底前上线。")
        self.assertEqual([it.decision for it in items], ["同意", "暂缓"])
        self.assertEqual(items[0].owner, "门诊部")
        self.assertEqual([it.position for it in items], ["（一）", "（一）"])

    def test_agree_and_later(self):
        items = self.decided("会议同意由办公室牵头修订值班制度，11月底前完成；值班补贴标准另行研究。")
        self.assertEqual([it.decision for it in items], ["同意", "另行研究"])

    def test_agree_and_disagree(self):
        items = self.decided("会议同意设立便民门诊，由门诊部牵头，12月底前开诊；不同意增设夜间门诊。")
        self.assertEqual([it.decision for it in items], ["同意", "不同意"])

    def test_postpone_with_preparation_task(self):
        items = self.decided("会议决定暂缓调整夜班补贴标准，由财务科先行测算所需资金，待年底职代会召开后再提交研究。")
        self.assertEqual([it.decision for it in items], ["暂缓", "决定"])
        self.assertEqual(items[1].owner, "财务科")
        self.assertIn("待年底职代会", items[0].text)

    def test_preparation_task_needs_deadline(self):
        ids = issue_ids(wrap("会议议定：", "调整夜班补贴标准事项暂缓，待年底职代会召开后再研究。由财务科先行测算所需资金。"))
        self.assertIn("MIN-002", ids)
        ids = issue_ids(wrap("会议议定：", "调整夜班补贴标准事项暂缓，待年底职代会召开后再研究。由财务科先行测算所需资金【待补：完成时限】。"))
        self.assertEqual(ids, [])

    def test_mixed_paragraph_list_reconcile(self):
        minutes = wrap("会议同意门诊预约挂号改革方案，由门诊部牵头、信息科配合，11月底前上线；爽约限制暂缓实施，待分时段预约运行三个月后再研究。")
        empty = "| 序号 | 议定事项 | 责任单位 | 完成时限 |\n|---|---|---|---|\n| 1 | 修订值班制度 | 办公室 | 11月底前 |"
        self.assertIn("MIN-104", list_issue_ids(minutes, empty))
        md = mc.render_list(items_of(minutes))
        self.assertIn("门诊部", md)
        self.assertNotIn("暂缓", md)
        self.assertEqual(list_issue_ids(minutes, md), [])

    def test_single_decision_not_split(self):
        cases = (
            "会议同意由办公室牵头修订，12月底前完成；修订后的制度报党委会审议通过后印发。",
            "预约爽约限制措施暂缓实施，待分时段预约运行三个月后再提交研究；下次会议再议。",
            "夜班补贴调整事项由财务科进一步测算后，提交下次会议研究。",
            "会议决定暂缓由后勤处实施的停车收费改革，待职工意见征集完成后再研究。",
            "会议同意《关于暂缓收费的报告》，由办公室牵头，11月底前印发。",
        )
        for text in cases:
            with self.subTest(text=text):
                self.assertEqual(len(self.decided(text)), 1)

    def test_postpone_owner_inside_not_listed(self):
        md = mc.render_list(items_of(wrap("会议决定暂缓由后勤处实施的停车收费改革，待职工意见征集完成后再研究。")))
        self.assertNotIn("后勤处", md)


class ItemCheckTest(unittest.TestCase):
    """必错标本与必过标本。"""

    def test_good_minutes_pass(self):
        self.assertEqual(issue_ids(GOOD_MINUTES), [])

    def test_min001_missing_owner(self):
        self.assertIn("MIN-001", issue_ids(wrap("会议同意修订值班制度，11月底前完成。")))

    def test_min001_002_are_hints_without_placeholder_advice(self):
        """材料没有的责任单位、时限不强求：只报 P2，建议不再要求在正文插【待补】。"""
        issues = [i for i in mc.check_items(items_of(wrap("会议同意修订值班制度。")))
                  if i.id in ("MIN-001", "MIN-002")]
        self.assertEqual(sorted(i.id for i in issues), ["MIN-001", "MIN-002"])
        for i in issues:
            self.assertEqual(i.level, "P2")
            self.assertNotIn("写【待补", i.suggestion)
            self.assertIn("起草说明", i.suggestion)

    def test_min001_marker_passes(self):
        self.assertNotIn("MIN-001", issue_ids(wrap("会议同意修订值班制度，由【待补：责任单位】牵头，11月底前完成。")))

    def test_min002_missing_deadline(self):
        self.assertIn("MIN-002", issue_ids(wrap("会议同意由办公室牵头修订值班制度。")))

    def test_min002_vague_deadline(self):
        text = wrap("会议同意由办公室牵头尽快修订值班制度。")
        items = items_of(text)
        issues = mc.check_items(items)
        self.assertEqual([i.id for i in issues], ["MIN-002"])
        self.assertIn("尽快", issues[0].message)

    def test_min002_marker_passes(self):
        self.assertNotIn("MIN-002", issue_ids(wrap("会议同意由办公室牵头尽快修订值班制度【待补：具体完成时限】。")))

    def test_min003_principle_without_revision(self):
        self.assertIn("MIN-003", issue_ids(wrap("会议原则同意该方案，由办公室牵头，11月底前实施。")))

    def test_min003_principle_with_revision_passes(self):
        self.assertNotIn("MIN-003", issue_ids(wrap("会议原则同意该方案，由办公室按会议意见修改完善后，11月底前实施。")))

    def test_min004_postpone_without_reason(self):
        ids = issue_ids(wrap("会议决定暂缓调整夜班补贴。"))
        self.assertIn("MIN-004", ids)
        self.assertNotIn("MIN-001", ids)
        self.assertNotIn("MIN-002", ids)

    def test_min004_postpone_with_condition_passes(self):
        self.assertNotIn("MIN-004", issue_ids(wrap("会议决定暂缓调整夜班补贴，待职代会召开后再研究。")))

    def test_exempt_decisions_need_no_owner(self):
        ids = issue_ids(wrap("会议不同意该申请。", "会议议定：停车收费事项另行研究。"))
        self.assertNotIn("MIN-001", ids)
        self.assertNotIn("MIN-002", ids)

    def test_min005_colloquial(self):
        self.assertIn("MIN-005", issue_ids(wrap("会议同意由办公室牵头，咱们11月底前弄完。")))

    def test_min005_quoted_passes(self):
        self.assertNotIn("MIN-005", issue_ids(wrap("会议同意由办公室牵头整理“我觉得”类口头表述，11月底前完成。")))

    def test_min006_no_items(self):
        self.assertIn("MIN-006", issue_ids(wrap("会议听取了财务科的汇报。")))

    def test_min007_missing_attendance(self):
        self.assertIn("MIN-007", issue_ids("某局局务会议纪要\n\n会议同意由办公室牵头，11月底前完成。\n"))

    def test_min007_not_for_record(self):
        self.assertNotIn("MIN-007", issue_ids("某局局务会议记录\n\n会议同意由办公室牵头，11月底前完成。\n"))

    def test_min008_speaker_tag(self):
        self.assertIn("MIN-008", issue_ids(wrap("发言人3提出，会议同意由办公室牵头，11月底前完成。")))


class ListMatchTest(unittest.TestCase):

    def test_good_list_passes(self):
        self.assertEqual(list_issue_ids(GOOD_MINUTES, GOOD_LIST), [])

    def test_min101_row_not_in_minutes(self):
        table = GOOD_LIST + "| 4 | 职工停车按科室分配车位试点 | | 后勤处 | | 11月30日前 | 待办理 | |\n"
        self.assertIn("MIN-101", list_issue_ids(GOOD_MINUTES, table))

    def test_min102_owner_mismatch(self):
        table = GOOD_LIST.replace("| 门诊部 | 信息科 |", "| 门诊部、医务部 | 信息科 |")
        self.assertEqual(list_issue_ids(GOOD_MINUTES, table), ["MIN-102"])

    def test_min103_deadline_mismatch(self):
        table = GOOD_LIST.replace("2026年11月底前", "12月底前")
        self.assertEqual(list_issue_ids(GOOD_MINUTES, table), ["MIN-103"])

    def test_min104_item_missing(self):
        table = "\n".join(GOOD_LIST.strip().split("\n")[:-1])
        self.assertEqual(list_issue_ids(GOOD_MINUTES, table), ["MIN-104"])

    def test_min105_missing_columns(self):
        table = "| 序号 | 议定事项 | 责任单位 |\n|---|---|---|\n| 1 | 门诊预约挂号改革 | 门诊部 |"
        self.assertEqual(list_issue_ids(GOOD_MINUTES, table), ["MIN-105"])

    def test_min106_list_adds_deadline(self):
        table = GOOD_LIST.replace("| 【待补：完成时限】 |", "| 11月20日前 |")
        self.assertEqual(list_issue_ids(GOOD_MINUTES, table), ["MIN-106"])

    def test_min106_list_adds_owner(self):
        minutes = wrap("会议同意修订值班制度，由【待补：责任单位】牵头，11月底前完成。")
        table = ("| 序号 | 议定事项 | 责任单位 | 完成时限 |\n|---|---|---|---|\n"
                 "| 1 | 修订值班制度 | 办公室 | 11月底前 |")
        self.assertEqual(list_issue_ids(minutes, table), ["MIN-106"])

    def test_min107_list_missing_value(self):
        table = GOOD_LIST.replace("| 11月15日前 |", "| 【待补：完成时限】 |")
        self.assertEqual(list_issue_ids(GOOD_MINUTES, table), ["MIN-107"])

    def test_requirement_row_matches(self):
        minutes = wrap("会议同意由办公室牵头修订值班制度，11月底前完成。",
                       "会议要求，由财务科牵头开展经费自查，12月15日前报送结果。")
        table = ("| 序号 | 议定事项 | 责任单位 | 完成时限 |\n|---|---|---|---|\n"
                 "| 1 | 修订值班制度 | 办公室 | 11月底前 |\n"
                 "| 2 | 开展经费自查并报送结果 | 财务科 | 12月15日前 |")
        self.assertEqual(list_issue_ids(minutes, table), [])

    def test_render_list_round_trip(self):
        md = mc.render_list(items_of(GOOD_MINUTES))
        self.assertIn("| 门诊部 | 信息科 | 11月底前 | 待办理 |", md)
        self.assertNotIn("暂缓", md)
        self.assertEqual(list_issue_ids(GOOD_MINUTES, md), [])

    def test_eval_minutes_pass(self):
        """评测材料写法：时限都能识别，不报 MIN-002；照纪要写的督办清单不报 MIN-106 等对账问题。"""
        cases = (
            ("eval1", EVAL1_MINUTES, EVAL1_LIST),
            ("eval3", EVAL3_MINUTES, EVAL3_LIST),
            ("eval8", EVAL8_MINUTES, EVAL8_LIST),
            ("eval11", EVAL11_MINUTES, EVAL11_LIST),
        )
        for name, minutes, table in cases:
            with self.subTest(name=name):
                self.assertNotIn("MIN-002", issue_ids(minutes))
                self.assertEqual(list_issue_ids(minutes, table), [])
                self.assertEqual(list_issue_ids(minutes, mc.render_list(items_of(minutes))), [])

    def test_eval1_items(self):
        decided = [it for it in items_of(EVAL1_MINUTES) if it.kind == "议定"]
        self.assertEqual([it.decision for it in decided], ["原则同意", "同意", "暂缓", "暂缓", "决定"])
        self.assertEqual(decided[1].owner, "门诊部")
        self.assertEqual(decided[4].owner, "财务科")

    def test_eval_deadline_in_list_not_p0(self):
        """清单照纪要写"今年12月开工""月底前"，不得报 MIN-106、MIN-103。"""
        ids = list_issue_ids(EVAL11_MINUTES, EVAL11_LIST.replace("今年12月开工", "2026年12月开工"))
        self.assertEqual(ids, [])
        ids = list_issue_ids(EVAL3_MINUTES, EVAL3_LIST)
        self.assertNotIn("MIN-106", ids)

    def test_render_list_marks_missing(self):
        md = mc.render_list(items_of(wrap("会议同意修订值班制度。")))
        self.assertIn("【待补：责任单位】", md)
        self.assertIn("【待补：完成时限】", md)
        self.assertIn("纪要未写责任单位", md)


LIST_HEAD = ("| 序号 | 议定事项 | 纪要位置 | 责任单位 | 配合单位 | 完成时限 | 办理状态 | 备注 |\n"
             "|---|---|---|---|---|---|---|---|\n")
# 一条中牵头、配合各一个单位
LEAD_COOP = wrap("会议议定：", "（一）同意档案数字化外包方案。由档案室牵头、信息中心配合，于6月底前完成。")
# 一条中两个单位各有任务、各有时限
TWO_TASKS = wrap("会议议定：", "（一）同意档案室改造方案，由后勤科于12月10日前完成，财务科11月20日前拨付经费。")


class TaskPairTest(unittest.TestCase):
    """对账比的是识别出的牵头、配合单位和各任务的时限，不是"在原文中出现过"。"""

    def test_split_tasks(self):
        tasks = mc.split_tasks(items_of(TWO_TASKS)[0].text)
        self.assertEqual([(t.owner, t.deadlines) for t in tasks],
                         [("后勤科", ["12月10日前"]), ("财务科", ["11月20日前"])])

    def test_lead_and_coop_list_passes(self):
        table = LIST_HEAD + "| 1 | 档案数字化外包 | （一） | 档案室 | 信息中心 | 6月底前 | 待办理 | |"
        self.assertEqual(list_issue_ids(LEAD_COOP, table), [])

    def test_swapped_lead_and_coop(self):
        """必错：牵头和配合对调。"""
        table = LIST_HEAD + "| 1 | 档案数字化外包 | （一） | 信息中心 | 档案室 | 6月底前 | 待办理 | |"
        ids = list_issue_ids(LEAD_COOP, table)
        self.assertIn("MIN-102", ids)
        self.assertIn("MIN-108", ids)

    def test_coop_written_as_owner(self):
        """必错：把配合单位并进责任单位栏。"""
        table = LIST_HEAD + "| 1 | 档案数字化外包 | （一） | 档案室、信息中心 | | 6月底前 | 待办理 | |"
        self.assertEqual(list_issue_ids(LEAD_COOP, table), ["MIN-102"])

    def test_coop_not_in_minutes(self):
        """必错：清单新增纪要这一条没有的配合单位。"""
        table = LIST_HEAD + "| 1 | 档案数字化外包 | （一） | 档案室 | 财务科 | 6月底前 | 待办理 | |"
        self.assertEqual(list_issue_ids(LEAD_COOP, table), ["MIN-108"])

    def test_deadline_of_other_task(self):
        """必错：时限挂到同一条中其他单位的任务上。"""
        table = LIST_HEAD + "| 1 | 档案室改造 | （一） | 后勤科 | | 11月20日前 | 待办理 | |"
        self.assertEqual(list_issue_ids(TWO_TASKS, table), ["MIN-103"])

    def test_two_tasks_list_passes(self):
        table = (LIST_HEAD
                 + "| 1 | 档案室改造 | （一） | 后勤科 | | 12月10日前 | 待办理 | |\n"
                 + "| 2 | 拨付档案室改造经费 | （一） | 财务科 | | 11月20日前 | 待办理 | |")
        self.assertEqual(list_issue_ids(TWO_TASKS, table), [])

    def test_emit_splits_two_tasks(self):
        md = mc.render_list(items_of(TWO_TASKS))
        self.assertIn("| 后勤科 |  | 12月10日前 |", md)
        self.assertIn("| 财务科 |  | 11月20日前 |", md)
        self.assertNotIn("12月10日前、11月20日前", md)
        self.assertEqual(list_issue_ids(TWO_TASKS, md), [])

    def test_marked_task_deadline_not_filled_by_list(self):
        """必错：这项任务的时限在纪要中标了待补，清单却借用同一条中其他任务的时限。"""
        minutes = wrap("会议议定：", "（一）同意由档案室牵头，6月底前完成扫描；由信息中心【待补：完成时限】完成系统对接。")
        table = (LIST_HEAD
                 + "| 1 | 完成扫描 | （一） | 档案室 | | 6月底前 | 待办理 | |\n"
                 + "| 2 | 完成系统对接 | （一） | 信息中心 | | 6月底前 | 待办理 | |")
        self.assertEqual(list_issue_ids(minutes, table), ["MIN-106"])

    def test_meeting_body_is_not_task_owner(self):
        """必过："由党委会研究决定"不另起任务，时限仍归办公室。"""
        minutes = wrap("会议议定：", "（一）原则同意《会议室管理办法》，由办公室修改完善后，由党委会研究决定，年底前印发。")
        table = LIST_HEAD + "| 1 | 修改完善会议室管理办法 | （一） | 办公室 | | 年底前 | 待办理 | |"
        self.assertEqual(list_issue_ids(minutes, table), [])


class NoFollowTest(unittest.TestCase):
    """任职、维持现状、审议通过等无需落实的结论（必过标本），以及不带"会议"的暂缓段落。"""

    def test_keep_status_quo_is_not_postpone(self):
        ids = issue_ids(wrap("会议同意暂不调整收费标准。"))
        self.assertNotIn("MIN-004", ids)
        self.assertNotIn("MIN-001", ids)
        self.assertNotIn("MIN-002", ids)
        self.assertNotIn("同意暂不", mc.render_list(items_of(wrap("会议同意暂不调整收费标准。"))))

    def test_appointment_not_marked(self):
        minutes = wrap("会议议定：", "（一）同意张某某任档案室主任，组织人事部按程序公示。")
        self.assertEqual(issue_ids(minutes), [])
        md = mc.render_list(items_of(minutes))
        self.assertNotIn("【待补", md)
        self.assertIn("任职事项，请确认是否列入", md)
        empty = LIST_HEAD + "| 1 | 修订值班制度 | | 办公室 | | 11月底前 | 待办理 | |"
        self.assertNotIn("MIN-104", list_issue_ids(minutes, empty))

    def test_task_with_owner_still_checked(self):
        """必错：有责任单位的审议通过事项仍要写时限。"""
        self.assertIn("MIN-002", issue_ids(wrap("会议审议通过了《会议室管理办法》，由办公室印发。")))

    def test_bare_postpone_under_topic(self):
        text = ("某中心主任办公会议纪要\n\n一、关于档案数字化\n\n会议议定：\n\n（一）同意由档案室牵头，6月底前完成。\n\n"
                "二、关于公务用车更新\n\n更新公务用车事项暂缓，待明年预算批复后再研究。由办公室先行摸清车况【待补：完成时限】。\n\n"
                "出席：张某\n")
        decided = [it for it in items_of(text) if it.kind == "议定"]
        self.assertEqual([it.decision for it in decided], ["同意", "暂缓", "决定"])
        self.assertEqual(decided[2].owner, "办公室")

    def test_narrative_without_section_not_item(self):
        self.assertEqual(items_of(wrap("二期资金暂缓拨付的情况已向区里报告。")), [])


class CliTest(unittest.TestCase):

    def setUp(self) -> None:
        self.tmp = Path(tempfile.mkdtemp(prefix="minutes_check_"))

    def tearDown(self) -> None:
        shutil.rmtree(self.tmp, ignore_errors=True)

    def write(self, name: str, text: str) -> Path:
        path = self.tmp / name
        path.write_text(text, encoding="utf-8")
        return path

    def run_cli(self, *args: str):
        out, err = io.StringIO(), io.StringIO()
        with redirect_stdout(out), redirect_stderr(err):
            code = mc.main(list(args))
        return code, out.getvalue(), err.getvalue()

    def test_json_output(self):
        path = self.write("纪要.txt", GOOD_MINUTES)
        code, out, _ = self.run_cli(str(path), "--json")
        self.assertEqual(code, 0)
        data = json.loads(out)
        self.assertEqual(data["summary"]["P0"], 0)
        self.assertEqual(len([i for i in data["items"] if i["kind"] == "议定"]), 5)

    def test_strict_exit_code(self):
        minutes = self.write("纪要.txt", GOOD_MINUTES)
        bad = self.write("清单.md", GOOD_LIST.replace("2026年11月底前", "12月底前"))
        code, _, _ = self.run_cli(str(minutes), "--list", str(bad), "--strict")
        self.assertEqual(code, 1)
        code, _, _ = self.run_cli(str(minutes), "--list", str(bad))
        self.assertEqual(code, 0)

    def test_missing_file(self):
        code, _, err = self.run_cli(str(self.tmp / "不存在.txt"))
        self.assertEqual(code, 2)
        self.assertIn("未找到文件", err)

    def test_unsupported_suffix(self):
        path = self.write("纪要.pdf", "x")
        code, _, _ = self.run_cli(str(path))
        self.assertEqual(code, 2)

    def test_list_without_table(self):
        minutes = self.write("纪要.txt", GOOD_MINUTES)
        lst = self.write("清单.md", "没有表格")
        code, _, err = self.run_cli(str(minutes), "--list", str(lst))
        self.assertEqual(code, 2)
        self.assertIn("没有找到", err)

    def test_emit_list_does_not_overwrite(self):
        minutes = self.write("纪要.txt", GOOD_MINUTES)
        existing = self.write("督办清单.md", "用户原有内容")
        code, out, _ = self.run_cli(str(minutes), "--emit-list", str(existing))
        self.assertEqual(code, 0)
        self.assertEqual(existing.read_text(encoding="utf-8"), "用户原有内容")
        v2 = self.tmp / "督办清单-v2.md"
        self.assertTrue(v2.is_file())
        self.assertIn("门诊部", v2.read_text(encoding="utf-8"))
        self.assertIn("督办清单-v2.md", out)

    def test_location_label_for_whole_text(self):
        path = self.write("空纪要.txt", "")
        code, out, _ = self.run_cli(str(path))
        self.assertEqual(code, 0)
        self.assertIn("全文", out)
        self.assertNotIn("清单：", out)

    def test_bad_argument(self):
        code, _, _ = self.run_cli("--no-such-option")
        self.assertEqual(code, 2)

    @unittest.skipUnless(HAS_OPENPYXL, "未安装 openpyxl")
    def test_xlsx_list(self):
        import openpyxl  # type: ignore
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["某医院院长办公会议定事项督办清单"])
        lines = GOOD_LIST.strip().split("\n")
        ws.append(mc._split_md_row(lines[0]))
        for line in lines[2:]:
            ws.append(mc._split_md_row(line))
        path = self.tmp / "清单.xlsx"
        wb.save(str(path))
        minutes = self.write("纪要.txt", GOOD_MINUTES)
        code, out, _ = self.run_cli(str(minutes), "--list", str(path), "--json")
        self.assertEqual(code, 0)
        self.assertEqual(json.loads(out)["summary"], {"items": 6, "P0": 0, "P1": 0, "P2": 0})

    @unittest.skipUnless(HAS_DOCX, "未安装 python-docx")
    def test_docx_input(self):
        import docx  # type: ignore
        document = docx.Document()
        for line in GOOD_MINUTES.split("\n"):
            document.add_paragraph(line)
        path = self.tmp / "纪要.docx"
        document.save(str(path))
        code, out, _ = self.run_cli(str(path), "--json")
        self.assertEqual(code, 0)
        self.assertEqual(json.loads(out)["summary"]["items"], 6)


if __name__ == "__main__":
    unittest.main()
