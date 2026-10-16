# -*- coding: utf-8 -*-
"""Markdown 表格转 Excel。

源文件位于仓库 shared/scripts/，各 skill 中的副本由 tools/sync_shared.py 生成，请在源文件处修改。
维护者：gongwen-rectify。接口见 docs/开发规范.md 8.4 节。
调用方：gongwen-rectify（整改台账）、gongwen-implement（任务分解表）、gongwen-minutes（督办清单）。

用法：
    python scripts/md2xlsx.py <输入.md> <输出.xlsx> [--title 表题] [--all-tables]

表格识别（按 GitHub Flavored Markdown 的表格写法）：
    表头行含竖线 |，下一行是分隔行（如 | --- | :---: |），分隔行的列数与表头相同；
    表体到空行或不含竖线的行为止；代码块（``` 或 ~~~ 围起来的部分）中的表格不转换；
    单元格中的 \\| 是竖线本身，<br> 转为换行；加粗（**合计**）和整格代码（`A-01`）去掉标记；
    [文字](链接) 只保留文字。138****0000 这类打码写法原样保留。

默认只转换第一张表；--all-tables 时每张表各占一个工作表。

表题：--title 指定的表题用于第一张表；其余表（或未指定 --title 时）取表格上方最近一行，
    该行是 Markdown 标题（# 开头），或以"表""附表""附件""附："开头，或以"表""清单""台账"结尾时作为表题，
    否则不加表题行。

样式（公文表格）：
    字体宋体；表题 14 磅加粗，跨列居中；表头 11 磅加粗、浅灰底色、居中；
    全部单元格加细边框，自动列宽（最宽 50 个字符宽），自动换行，按内容估算行高；
    冻结表头（表头以下滚动）；打印时每页重复表头，A4 纸，列多时横向并缩放到一页宽。

单元格类型：
    纯整数、小数（转换后写法不变的，如 125.6、3216）存为数值，便于求和；
    列名含"编号""序号""代码""号码""电话""证件"的列、带前导 0 的数（01）、末尾带 0 的小数（1.10）、
    超过 12 位的数字串、日期、百分数一律存为文本，避免 Excel 改写；
    以 = 开头的文字存为文本，不当作公式。

输出文件已存在时，自动改名为 <名称>-v2.xlsx、-v3.xlsx……，不覆盖已有文件。

退出码：0 成功；2 输入错误（文件不存在、没有表格、输出扩展名不是 .xlsx、未安装 openpyxl 等）。
只依赖标准库和 openpyxl；openpyxl 未安装时给出安装提示并返回 2。
"""
from __future__ import annotations

import argparse
import html
import math
import re
import sys
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path
from typing import List, Optional, Sequence, Tuple

FONT_NAME = "宋体"
TITLE_SIZE = 14
BODY_SIZE = 11
HEADER_FILL = "D9D9D9"
MAX_COL_WIDTH = 50
MIN_COL_WIDTH = 6
LINE_HEIGHT = 15.0  # 宋体 11 磅每行约 15 磅
TITLE_ROW_HEIGHT = 30.0
SHEET_NAME_MAX = 31
INVALID_SHEET_CHARS = re.compile(r"[\[\]:*?/\\：]")

# 文本列：列名含这些字样的列不做数值转换
TEXT_COLUMN_KEYWORDS = ("编号", "序号", "代码", "号码", "电话", "证件", "文号", "工号", "学号", "账号")

FENCE_RE = re.compile(r"^\s{0,3}(`{3,}|~{3,})")
DELIM_CELL_RE = re.compile(r"^:?-+:?$")
HEADING_RE = re.compile(r"^\s{0,3}#{1,6}\s+(.*?)\s*#*\s*$")
BR_RE = re.compile(r"<br\s*/?>", re.IGNORECASE)
LINK_RE = re.compile(r"!?\[([^\]\n]*)\]\([^)\n]*\)")
WRAP_MARK_RE = re.compile(r"^(\*\*|__)(?![*_\s])(.+?)(?<![*_\s])\1$", re.DOTALL)
WRAP_CODE_RE = re.compile(r"^`([^`]+)`$")
# 行内加粗 **文字**：前后不能紧挨星号或字母数字，文字不能以空白、星号、标点开头，避免误伤 138****0000 这类打码
INLINE_BOLD_RE = re.compile(
    r"(?<![*0-9A-Za-z])\*\*(?![\s*，。、；：！？,.;:!?）)」』”’])(.+?)(?<![\s*])\*\*(?![*0-9A-Za-z])"
)
INT_RE = re.compile(r"^-?(?:0|[1-9]\d{0,11})$")
DEC_RE = re.compile(r"^-?(?:0|[1-9]\d{0,11})\.\d{1,6}$")


class InputError(Exception):
    """输入错误：退出码 2。"""


@dataclass
class MdTable:
    """一张 Markdown 表格。line 为表头所在行号（从 1 起）。"""

    header: List[str]
    rows: List[List[str]]
    aligns: List[Optional[str]]
    title: Optional[str] = None
    line: int = 0
    extra: dict = field(default_factory=dict)

    @property
    def width(self) -> int:
        return len(self.header)


# ---------------------------------------------------------------------------
# 解析
# ---------------------------------------------------------------------------

def split_row(line: str) -> List[str]:
    """按未转义的竖线拆分一行，去掉首尾竖线。"""
    text = line.strip()
    cells: List[str] = []
    buf: List[str] = []
    i = 0
    while i < len(text):
        ch = text[i]
        if ch == "\\" and i + 1 < len(text) and text[i + 1] == "|":
            buf.append("\\|")  # 保留转义，清洗单元格时再还原
            i += 2
            continue
        if ch == "|":
            cells.append("".join(buf))
            buf = []
        else:
            buf.append(ch)
        i += 1
    cells.append("".join(buf))
    if text.startswith("|"):
        cells = cells[1:]
    if text.endswith("|") and not text.endswith("\\|") and cells:
        cells = cells[:-1]
    return [c.strip() for c in cells]


def has_pipe(line: str) -> bool:
    return "|" in line.replace("\\|", "")


def parse_delimiter(line: str) -> Optional[List[Optional[str]]]:
    """分隔行返回各列对齐方式（left、center、right 或 None），不是分隔行返回 None。"""
    if not has_pipe(line) and not line.strip().startswith(("-", ":")):
        return None
    cells = split_row(line)
    if not cells:
        return None
    aligns: List[Optional[str]] = []
    for cell in cells:
        cell = cell.replace(" ", "")
        if not DELIM_CELL_RE.match(cell):
            return None
        if cell.startswith(":") and cell.endswith(":"):
            aligns.append("center")
        elif cell.endswith(":"):
            aligns.append("right")
        elif cell.startswith(":"):
            aligns.append("left")
        else:
            aligns.append(None)
    return aligns


def clean_cell(text: str) -> str:
    """去掉单元格中的 Markdown 标记，保留文字。"""
    value = BR_RE.sub("\n", text)
    value = value.replace("\\|", "|")
    value = LINK_RE.sub(lambda m: m.group(1), value)
    value = html.unescape(value).replace(" ", " ")
    lines = []
    for part in value.split("\n"):
        part = part.strip()
        match = WRAP_MARK_RE.match(part)
        if match:
            part = match.group(2).strip()
        match = WRAP_CODE_RE.match(part)
        if match:
            part = match.group(1).strip()
        part = INLINE_BOLD_RE.sub(lambda m: m.group(1), part)
        lines.append(part)
    return "\n".join(lines).strip()


def _plain_title(line: str) -> Optional[str]:
    """表格上方一行能作表题时返回表题文字。"""
    text = line.strip()
    if not text or has_pipe(text):
        return None
    match = HEADING_RE.match(text)
    if match:
        text = clean_cell(match.group(1))
    else:
        text = clean_cell(text)
        if not (
            re.match(r"^(表|附表|附件|附[：:])", text)
            or re.search(r"(表|清单|台账)[）)]?[：:]?$", text)
        ):
            return None
    text = re.sub(r"^附[：:]\s*", "", text)
    text = text.rstrip("：:").strip()
    if not text or len(text) > 60:
        return None
    return text


def parse_tables(text: str) -> List[MdTable]:
    """从 Markdown 文本中找出全部表格。"""
    lines = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    tables: List[MdTable] = []
    in_fence: Optional[str] = None
    last_text_line: Optional[str] = None  # 表格上方最近的非空行
    i = 0
    while i < len(lines):
        line = lines[i]
        fence = FENCE_RE.match(line)
        if in_fence:
            if fence and fence.group(1)[0] == in_fence[0] and len(fence.group(1)) >= len(in_fence):
                in_fence = None
            i += 1
            continue
        if fence:
            in_fence = fence.group(1)
            i += 1
            continue
        if has_pipe(line) and i + 1 < len(lines):
            header = split_row(line)
            aligns = parse_delimiter(lines[i + 1])
            if aligns is not None and len(aligns) == len(header) and (any(header) or len(header) > 1):
                rows: List[List[str]] = []
                j = i + 2
                while j < len(lines) and lines[j].strip() and has_pipe(lines[j]):
                    if FENCE_RE.match(lines[j]):
                        break
                    cells = split_row(lines[j])
                    if len(cells) < len(header):
                        cells += [""] * (len(header) - len(cells))
                    rows.append([clean_cell(c) for c in cells[: len(header)]])
                    j += 1
                title = _plain_title(last_text_line) if last_text_line else None
                tables.append(MdTable(
                    header=[clean_cell(h) for h in header],
                    rows=rows,
                    aligns=aligns,
                    title=title,
                    line=i + 1,
                ))
                last_text_line = None
                i = j
                continue
        if line.strip():
            last_text_line = line
        i += 1
    return tables


def read_text(path: Path) -> str:
    """读取 UTF-8 文本（兼容 BOM），失败时尝试 GB18030。"""
    data = path.read_bytes()
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            return data.decode("gb18030")
        except UnicodeDecodeError as exc:
            raise InputError(f"无法按 UTF-8 或 GB18030 读取文件：{path}") from exc


# ---------------------------------------------------------------------------
# 单元格与版面计算
# ---------------------------------------------------------------------------

def display_width(text: str) -> int:
    """估算显示宽度：中文等全角字符算 2，其余算 1。"""
    width = 0
    for ch in text:
        width += 2 if unicodedata.east_asian_width(ch) in ("W", "F") else 1
    return width


def is_text_column(name: str) -> bool:
    return any(key in name for key in TEXT_COLUMN_KEYWORDS)


def to_number(value: str):
    """能无损转为数值的返回 int 或 float，否则返回 None。"""
    text = value.strip()
    if text == "-0":
        return None
    if INT_RE.match(text):
        return int(text)
    if DEC_RE.match(text):
        number = float(text)
        # 末尾带 0 的小数（1.10）转换后写法会变，保留文本
        if repr(number) == text or ("%g" % number) == text:
            return number
    return None


def column_widths(table: MdTable) -> List[float]:
    widths: List[float] = []
    for col in range(table.width):
        texts = [table.header[col]] + [row[col] for row in table.rows]
        longest = 0
        for text in texts:
            for part in text.split("\n"):
                longest = max(longest, display_width(part))
        header_width = display_width(table.header[col]) + 2  # 表头加粗略宽
        widths.append(float(min(MAX_COL_WIDTH, max(MIN_COL_WIDTH, longest + 2, header_width))))
    return widths


def estimate_lines(text: str, col_width: float) -> int:
    usable = max(1.0, col_width - 1.5)
    total = 0
    for part in (text or "").split("\n"):
        total += max(1, math.ceil(display_width(part) / usable))
    return max(1, total)


def sanitize_sheet_name(name: str, used: List[str], index: int) -> str:
    base = INVALID_SHEET_CHARS.sub("", name or "").strip().strip("'") or f"表{index}"
    base = base[:SHEET_NAME_MAX]
    candidate = base
    n = 2
    lowered = [u.lower() for u in used]
    while candidate.lower() in lowered:
        suffix = f"({n})"
        candidate = base[: SHEET_NAME_MAX - len(suffix)] + suffix
        n += 1
    return candidate


def unique_path(path: Path) -> Path:
    """已存在时依次尝试 -v2、-v3……，不覆盖已有文件。"""
    if not path.exists():
        return path
    n = 2
    while True:
        candidate = path.with_name(f"{path.stem}-v{n}{path.suffix}")
        if not candidate.exists():
            return candidate
        n += 1


# ---------------------------------------------------------------------------
# 写入
# ---------------------------------------------------------------------------

def _require_openpyxl():
    try:
        import openpyxl  # noqa: F401
    except ImportError as exc:
        raise InputError("未安装 openpyxl，请先运行：pip install openpyxl") from exc
    return openpyxl


def write_table(ws, table: MdTable, title: Optional[str]) -> None:
    """把一张表写入工作表 ws，套用公文表格样式。"""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    body_font = Font(name=FONT_NAME, size=BODY_SIZE)
    header_font = Font(name=FONT_NAME, size=BODY_SIZE, bold=True)
    title_font = Font(name=FONT_NAME, size=TITLE_SIZE, bold=True)
    fill = PatternFill(fill_type="solid", fgColor=HEADER_FILL)

    ncols = max(1, table.width)
    widths = column_widths(table)
    row = 1
    if title:
        cell = ws.cell(row=1, column=1, value=title)
        cell.data_type = "s"
        cell.font = title_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if ncols > 1:
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
        ws.row_dimensions[1].height = TITLE_ROW_HEIGHT
        row = 2

    header_row = row
    header_lines = 1
    for col, name in enumerate(table.header, start=1):
        cell = ws.cell(row=row, column=col, value=name)
        cell.data_type = "s"
        cell.font = header_font
        cell.fill = fill
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        header_lines = max(header_lines, estimate_lines(name, widths[col - 1]))
    ws.row_dimensions[row].height = header_lines * LINE_HEIGHT + 6

    text_cols = [is_text_column(name) for name in table.header]
    for values in table.rows:
        row += 1
        lines = 1
        for col, value in enumerate(values, start=1):
            number = None if text_cols[col - 1] else to_number(value)
            cell = ws.cell(row=row, column=col)
            if number is not None:
                cell.value = number
            else:
                cell.value = value
                cell.data_type = "s"  # 以 = 开头的文字也按文本写入
            cell.font = body_font
            cell.border = border
            align = table.aligns[col - 1] if col - 1 < len(table.aligns) else None
            if align is None:
                if text_cols[col - 1]:
                    align = "center"
                elif number is not None:
                    align = "right"
                else:
                    align = "left"
            cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
            lines = max(lines, estimate_lines(value, widths[col - 1]))
        ws.row_dimensions[row].height = lines * LINE_HEIGHT + 6

    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    # 冻结表头：表头以下滚动
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1).coordinate
    # 打印：A4，每页重复表头，列多时横向并缩放到一页宽
    ws.print_title_rows = f"{header_row}:{header_row}"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    if ncols > 5 or sum(widths) > 100:
        ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True


def convert(
    input_path: Path,
    output_path: Path,
    title: Optional[str] = None,
    all_tables: bool = False,
) -> Tuple[Path, List[MdTable], List[str]]:
    """转换主函数，返回 (实际输出路径, 已转换的表, 工作表名)。出错时抛出 InputError。"""
    if not input_path.is_file():
        raise InputError(f"未找到输入文件：{input_path}")
    if output_path.suffix.lower() != ".xlsx":
        raise InputError(f"输出文件须为 .xlsx：{output_path}")
    openpyxl = _require_openpyxl()

    tables = parse_tables(read_text(input_path))
    if not tables:
        raise InputError(
            "未找到 Markdown 表格。表格需要表头行和分隔行，例如：\n"
            "| 编号 | 问题 |\n| --- | --- |"
        )
    selected = tables if all_tables else tables[:1]
    if title:
        selected[0].title = title.strip()

    wb = openpyxl.Workbook()
    used: List[str] = []
    for index, table in enumerate(selected, start=1):
        ws = wb.active if index == 1 else wb.create_sheet()
        name = sanitize_sheet_name(table.title or "", used, index)
        used.append(name)
        ws.title = name
        write_table(ws, table, table.title)

    target = unique_path(output_path)
    try:
        if not target.parent.is_dir():  # 已存在的目录不再 mkdir（Windows 超长路径下 mkdir 会报错）
            target.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(target))
    except OSError as exc:
        raise InputError(f"无法写入输出文件：{target}（{exc}）") from exc
    return target, selected, used


# ---------------------------------------------------------------------------
# 入口
# ---------------------------------------------------------------------------

def _setup_console() -> None:
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8")
        except Exception:
            pass


class _ChineseHelpFormatter(argparse.RawDescriptionHelpFormatter):
    def add_usage(self, usage, actions, groups, prefix=None):
        return super().add_usage(usage, actions, groups, "用法：" if prefix is None else prefix)


def make_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        prog="md2xlsx.py",
        description="把 Markdown 文件中的表格转成公文样式的 Excel（宋体、表头加粗加底色、全边框、"
                    "自动列宽和换行、冻结表头）。默认只转第一张表。",
        epilog="输出文件已存在时自动加 -v2、-v3 后缀，不覆盖。退出码：0 成功；2 输入错误或未安装 openpyxl。",
        formatter_class=_ChineseHelpFormatter,
        add_help=False,
    )
    parser._optionals.title = "选项"
    parser._positionals.title = "参数"
    parser.add_argument("-h", "--help", action="help", help="显示本帮助信息并退出")
    parser.add_argument("input", help="输入的 Markdown 文件（.md 或 .txt）")
    parser.add_argument("output", help="输出的 Excel 文件（.xlsx）")
    parser.add_argument("--title", help="表题，写在表格上方并跨列居中（用于第一张表）")
    parser.add_argument("--all-tables", action="store_true", help="转换全部表格，每张表各占一个工作表")
    return parser


def main(argv: Optional[Sequence[str]] = None) -> int:
    _setup_console()
    parser = make_parser()
    try:
        args = parser.parse_args(argv)
    except SystemExit as exc:
        return 0 if exc.code == 0 else 2
    try:
        target, tables, names = convert(
            Path(args.input), Path(args.output), title=args.title, all_tables=args.all_tables
        )
    except InputError as exc:
        print(f"错误：{exc}", file=sys.stderr)
        return 2
    print(f"已生成：{target}")
    if target != Path(args.output):
        print(f"（{args.output} 已存在，未覆盖，另存为 {target.name}）")
    for name, table in zip(names, tables):
        print(f"工作表“{name}”：{table.width} 列，{len(table.rows)} 行（不含表头）")
    return 0


if __name__ == "__main__":
    sys.exit(main())
